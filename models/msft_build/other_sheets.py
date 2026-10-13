"""Rewire the Bull-Base-Bear, DCF and Charts tabs of the MLM template (current layout: DCF with live valuation date, fade period,
implied ROIC, target-IRR entry price, reverse DCF and Mauboussin EV/NOPAT; Bull-Base-Bear with buyback rows) to the MSFT Model rows /
columns; keep the template chart XML. The template's DCF_old tab is not carried.

Column mapping: the template's 2025A base column becomes MSFT FY2026A (fiscal year ended 30-Jun-2026, ~3.5 months before the valuation
date) and the template's 2026E-2030E become MSFT FY2027E-FY2031E, so the DCF's explicit years start in FY2027E and net debt is taken
at 30-Jun-2026."""
import re, zipfile, shutil, copy
from openpyxl.utils import get_column_letter as L
from fw import R, TEMPLATE, ANN, NOTE

# template Model row -> MSFT row key
MAP = {158: 'is_rev', 164: 'r_cash', 166: 'is_oi', 182: 'is_dda', 186: 'sh_dil', 211: 'r_adj', 212: 'r_adj_m', 219: 'r_ebit', 221: 'r_adj',
       229: 'e_adjni', 231: 'e_eps', 238: 'gm_ai', 245: 'gm_ebitda_m', 247: 'gm_op_m', 260: 'cf_wc', 262: 'cf_capex', 263: 'cf_acq',
       280: 'fcf', 282: 'fcf_m', 290: 'bb_px', 291: 'bb_cash', 292: 'bb_sh', 331: 'bs_nci', 408: 'ra_ic', 411: 'ra_roic', 422: 'rn_ronta',
       428: 'lv_nd', 433: 'v_px', 434: 'v_sh'}
CMAP = {'BV': ANN[2026].c, 'CA': ANN[2027].c, 'CB': ANN[2028].c, 'CC': ANN[2029].c, 'CD': ANN[2030].c, 'CE': ANN[2031].c, 'CF': NOTE}
REF = re.compile(r"Model!(\$?)([A-Z]{1,3})(\$?)(\d+)")
HIST = [ANN[y].c for y in range(2006, 2027)]

def remap(f):
    def rep(m):
        c, r = m.group(2), int(m.group(4))
        c2 = CMAP.get(c, c)
        r2 = R[MAP[r]] if r in MAP else r
        if c == 'CF' and r == 2: r2 = 2
        return f'Model!{m.group(1)}{c2}{m.group(3)}{r2}'
    return REF.sub(rep, f)

def hist_stat(f):
    """LT historical AVERAGE / MIN / MAX over the template's annual columns -> MSFT FY2006-FY2026 annual columns."""
    m = re.match(r'=IFERROR\((AVERAGE|MIN|MAX)\(Model!C(\d+),', f)
    if not m: return None
    fn, r = m.group(1), int(m.group(2))
    rr = R[MAP[r]]
    return f'=IFERROR({fn}({",".join(f"Model!{c}{rr}" for c in HIST)}),"")'

def fy_dates(f):
    f = re.sub(r'DATE\((\d{4}),12,31\)', lambda m: f'DATE({int(m.group(1)) + 1},6,30)', f)
    f = re.sub(r'DATE\((\d{4})-1,12,31\)', lambda m: f'DATE({int(m.group(1)) + 1}-1,6,30)', f)
    return f

TEXT = {
    'Bull-Base-Bear': {
        'B2': 'Microsoft Corporation (MSFT) — Bull / Base / Bear Case P&L Summary',
        'B3': ('Adjusted figures (company "As Adjusted" definitions; adjusted EBITDA is a model definition) · USD millions · fiscal years ended 30 June · Live panel linked to '
               'Model — toggle scenario via Model!CK2. Bull and Bear panels are value snapshots generated from the Model with the switch set to each case. FY2027E = Q1/27E '
               'calibrated to the Q1 FY27 outlook (8-K 2-Sep-2026) + Q2–Q4/27E carrying the Q1 calibration.'),
        'G26': 30, 'H26': 30, 'S26': 34, 'T26': 34, 'AD26': 24, 'AE26': 24, 'AF26': 24,
    },
    'DCF': {
        'B2': 'DCF VALUATION — MICROSOFT CORPORATION (MSFT)',
        'C5': 0.0425, 'D5': 'Assumption (~US 10Y Treasury yield, Oct-2026; government rate sources not reachable from the build environment); update to market',
        'D6': 'Assumption: US implied ERP ~4.5–5.5%',
        'C7': 1.05, 'D7': 'Assumption: mega-cap software / cloud peers (ORCL, CRM, SAP, AMZN, GOOGL) ~0.9–1.2; capex intensity of the AI build-out adds operating leverage',
        'C9': 0.045, 'D9': 'Assumption: AAA-rated issuer; finance-lease discount rate 4.5% (FY2026 10-K); senior notes carry ~1.2% net of capitalised interest',
        'C10': 0.19, 'D10': 'Model FY2028E+ effective tax rate (FY26 19.4%; Q1 FY27 guided ~20%)',
        'C12': 0.02, 'D12': 'Debt $40.3bn + finance leases $66.6bn vs ~$4.0trn market cap ($541 × ~7.45bn diluted shares): D/V ~2–3%',
        'D15': 'Periods = (30-Jun of fiscal year − today) / 365.25; cash flows assumed at fiscal year-end (30 June); FY2026A (ended 30-Jun-2026) is the base year',
        'B18': 'Adj. EBIT (= adjusted operating income; amortization and SBC are expensed)',
        'B21': 'plus D&A (depreciation, amortization & other)',
        'B22': '− Capex incl. finance-lease assets and acquisitions (M&A lever 2028E–31E)',
        'B26': 'Fade-period growth (2032E–36E; 10.0% → 6.0% linear)', 'I26': 0.10, 'M26': 0.06, 'C27': 0.035,
        'D32': ("Microsoft's operating income already deducts amortization of acquired intangibles and stock-based compensation (neither is adjusted out). D&A is added back "
                'in row 21. Row 22 deducts capital expenditures INCLUDING finance-lease assets (datacenters) because finance-lease liabilities are counted in net debt; it also '
                'includes the M&A lever in FY2028E–31E. Investment gains / losses (incl. OpenAI) are excluded from UFCF; the OpenAI stake and other equity investments '
                '(~$36bn carrying value) are not added back (conservative). The FY2027E capex peak (Q1 guidance over $50bn) depresses near-term UFCF.'),
        'D43': 'Opening IC = Model FY2026A book IC (equity + net debt incl. finance leases at 30-Jun-2026), rolled forward with capex + M&A − D&A + ΔWC.',
        'B48': 'Less: Net debt incl. finance leases + NCI (Model FY2026A year end, 30-Jun-2026)',
        'D48': 'MSFT: no noncontrolling interests; debt $40.3bn + finance leases $66.6bn − cash & short-term investments $76.8bn at the base-year end.',
        'B50': 'Diluted shares (m, Model FY2026A period end)',
        'D76': 'Rationale (Microsoft characteristics)',
        'C78': 0.12, 'D78': 'Model adjusted operating income CAGR FY2026–31E: Azure / Microsoft 365 cloud growth with margins absorbing the AI-capex depreciation ramp → ~12% NOPAT growth',
        'C79': 0.25, 'D79': 'Software and cloud platforms with high switching costs; incremental AI infrastructure returns are lower than legacy software returns (capex-heavy), so RONIC fades toward ~25%',
        'C80': 20, 'D80': 'Installed base across enterprise productivity, cloud infrastructure and developer tools sustains a long competitive-advantage period; risk from AI-capacity overbuild and pricing',
        'B90': 'VALUATION  —  Applied to FY2029E NOPAT',
        'B91': 'FY2029E Adj. EBIT (adjusted operating income)',
        'B93': 'FY2029E NOPAT',
        'B95': 'Implied Enterprise Value (FY2029E basis)',
        'B96': 'Less: FY2028E year-end net debt incl. finance leases + NCI (Model)',
        'B98': 'Diluted shares FY2029E (m, Model)',
        'B99': 'Implied share price FY2029 ($)',
        'B106': 'Plus: FY2026A year-end net debt incl. finance leases + NCI',
        'B108': 'MSFT market-implied EV/NOPAT (on FY2029E NOPAT)',
    },
    'Charts': {
        'B2': 'Microsoft Corporation — Operating Performance (FY2012–FY2026) and Model Forecast (FY2027E–FY2031E)',
        'B3': ('USDm · fiscal years ended 30 June · Linked to Model. Bars = contribution of each FY27 segment to total revenue growth (pp; restated history FY2026 only). '
               'Adjusted EBITDA = model definition; operating income = GAAP. ROIC / RONTA on adjusted operating income.'),
        'B8': 'Agents and Infra contribution to growth (pp)', 'B9': 'Devices and Consumer contribution to growth (pp)', 'B12': 'Adjusted EBITDA', 'B13': 'Operating income (GAAP)',
        'B14': 'Adj. EBITDA Margin %', 'B15': 'Operating Margin %',
    },
}
BBB_ROWS = {9: 'Revenue', 10: '   Revenue y/y %', 11: '   Agents and Infra revenue y/y %', 12: 'Adjusted EBITDA (model)',
            15: 'Operating income (GAAP)', 16: '   Operating income y/y %', 17: '   Operating margin (GAAP) %', 19: 'Adjusted Net Income',
            35: 'Adjusted EBITDA (leverage basis)', 36: '   Adjusted EBITDA % of revenue'}
CHART_TITLES = {
    'Aggregates Revenue Drivers: Shipment Volume vs. Price (ASP), y/y %': 'Revenue Growth: Agents and Infra vs. Devices and Consumer Contribution, pp',
    'Adjusted EBITDA &amp; Earnings from Operations (USDm)': 'Adjusted EBITDA &amp; Operating Income (USDm)',
}

def build_charts(ws):
    for i, y in enumerate(range(2012, 2032)):
        c = L(3 + i); mc = ANN[y].c
        ws[f'{c}7'] = y if y <= 2026 else f'{y}E'
        ws[f'{c}8'] = f'=IFERROR(Model!{mc}{R["gc_ai"]}+0,0)'
        ws[f'{c}9'] = f'=IFERROR(Model!{mc}{R["gc_dc"]}+0,0)'
        ws[f'{c}11'] = f'={c}7'
        ws[f'{c}12'] = f'=Model!{mc}{R["r_adj"]}'
        ws[f'{c}13'] = f'=Model!{mc}{R["is_oi"]}'
        ws[f'{c}14'] = f'=Model!{mc}{R["gm_ebitda_m"]}'
        ws[f'{c}15'] = f'=Model!{mc}{R["gm_op_m"]}'
        ws[f'{c}17'] = f'={c}7'
        ws[f'{c}18'] = f'=IFERROR(Model!{mc}{R["ra_roic"]}+0,NA())'
        ws[f'{c}19'] = f'=IFERROR(Model!{mc}{R["rn_ronta"]}+0,NA())'
        ws[f'{c}20'] = f'=IFERROR(Model!{mc}{R["gm_ebitda_m"]}+0,NA())'
        ws[f'{c}21'] = f'=IFERROR(Model!{mc}{R["fcf_m"]}+0,NA())'

def do_all(wb):
    for name in ('Bull-Base-Bear', 'DCF', 'Charts'):
        ws = wb[name]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and 'Model!' in c.value:
                    h = hist_stat(c.value) if name == 'Bull-Base-Bear' else None
                    c.value = h or remap(c.value)
                if isinstance(c.value, str) and 'DATE(' in c.value:
                    c.value = fy_dates(c.value)
        for addr, t in TEXT.get(name, {}).items():
            ws[addr] = t
    b = wb['Bull-Base-Bear']
    for r, t in BBB_ROWS.items():
        for c in ('B', 'N', 'Z'): b[f'{c}{r}'] = t
    for c0 in ('C', 'O', 'AA'):
        i = b[f'{c0}7'].column
        b.cell(7, i).value = '2026A'
        for k in range(1, 6): b.cell(7, i + k).value = f'{2026 + k}E'
        b.cell(7, i + 6).value = "26-'31 CAGR"
    from openpyxl.comments import Comment
    cm = Comment('P/E on adjusted diluted EPS (company definition; FY2026A $17.28). MSFT at $541 (user price, 12-Oct-2026) trades at ~31x FY2026A adjusted EPS. '
                 'Exit multiples are assumptions: Base 30x, Bull 34x, Bear 24x applied to FY2030E–31E (FY2029E–31E for Bear), reflecting slower growth as the AI '
                 'capex cycle matures.', 'MSFT model')
    cm.width, cm.height = 420, 120; b['N26'].comment = cm
    cm2 = Comment(f'Snapshot (values) of Model rows {R["bb_cash"]}, {R["bb_px"]} and {R["bb_sh"]} (Share Buyback Schedule) with Model!{NOTE}2 = Bull / Bear. '
                  'Base panel (D40:H42) is live-linked.', 'MSFT model')
    cm2.width, cm2.height = 420, 90; b['P40'].comment = cm2
    b['J2'] = 'Share price ($):'
    for c in ('B', 'N', 'Z'): b[f'{c}29'] = 'IRR (to 30-Jun of Y-1)'
    for c in ('B', 'N', 'Z'):
        b[f'{c}40'] = '   Share repurchases ($m)'; b[f'{c}41'] = '   Implied avg buyback price ($)'; b[f'{c}42'] = '   Shares repurchased (m)'
        b[f'{c}40']._style = copy.copy(b['B40']._style); b[f'{c}41']._style = copy.copy(b['B41']._style); b[f'{c}42']._style = copy.copy(b['B42']._style)
    d = wb['DCF']
    for i, col in enumerate('DEFGHIJKLM'):
        d[f'{col}17'] = f'{2027 + i}E'
    a26 = ANN[2026].c
    for i, col in enumerate('DEFGH'):
        mc = ANN[2027 + i].c
        d[f'{col}22'] = f'=Model!{mc}{R["cf_capex"]}+Model!{mc}{R["cf_acq"]}-Model!{mc}{R["x_fl_add"]}'
    d['D34'] = f'=Model!{a26}{R["ra_ic"]}'
    d['C48'] = f'=Model!{a26}{R["lv_nd"]}+Model!{a26}{R["bs_nci"]}'
    d['C50'] = f'=Model!{a26}{R["v_sh"]}'
    d['C52'] = f'=Model!{a26}{R["v_px"]}'
    d['P17'] = 'CAGR 27–31E'; d['Q17'] = 'CAGR 27–36E'
    build_charts(wb['Charts'])

def snapshot_values(wb, bull, bear):
    ws = wb['Bull-Base-Bear']
    cols = 'CDEFGHIJKL'
    for snap, off in ((bull, 'OPQRSTUVWX'), (bear, ['AA', 'AB', 'AC', 'AD', 'AE', 'AF', 'AG', 'AH', 'AI', 'AJ'])):
        for (r, c), v in snap.items():
            tgt = f'{off[cols.index(c)]}{r}'
            cur = ws[tgt].value
            if isinstance(cur, str) and cur.startswith('='): continue
            if r in (7, 26): continue
            ws[tgt] = v

def inject_charts(path):
    """Replace openpyxl-rewritten chart parts with the template chart XML (combo charts, colours, layout), retitled, ranges kept
    (Charts!C:V = FY2012-FY2031E) and cached values removed so Excel redraws from the MSFT cells."""
    tz = zipfile.ZipFile(TEMPLATE)
    tmpl = {n: tz.read(n) for n in tz.namelist() if n.startswith('xl/charts/chart')}
    tmp = path + '.tmp'
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in tmpl:
                t = tmpl[item.filename].decode('utf-8')
                for a, b2 in CHART_TITLES.items(): t = t.replace(a, b2)
                t = re.sub(r'<c:numCache>.*?</c:numCache>', '', t, flags=re.S)
                t = re.sub(r'<c:strCache>.*?</c:strCache>', '', t, flags=re.S)
                data = t.encode('utf-8')
            zout.writestr(item, data)
    shutil.move(tmp, path)
