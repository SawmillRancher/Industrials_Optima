"""Build MSFT_Model.xlsx from the MLM template (same layout, formatting and process as the prior models)."""
import sys, copy, os, json
from openpyxl.styles import Font
from fw import *
import fw
from spec import build_spec, scen_layout, SCEN_ROWS, SC
from spec2 import build_spec2

OUT = os.path.join(os.path.dirname(BASE), 'MSFT_Model.xlsx')
GQ = MAN['guidance_q1fy27']
GU = MAN['guidance_fy27_unverified']

# ====================================================================== inputs (V0)
def wc_inputs():
    q = DATA['2026']['x']
    rev = q['rev']
    return {'dso': round(q['bs_ar'] / rev * 365), 'ucr': round(q['bs_ucr'] / rev, 3), 'ap': round(q['bs_ap'] / rev, 3)}
WC = wc_inputs()
WC['note'] = (f'Forecast inputs at fiscal-year-end (30-Jun-2026) levels on FY2026 revenue (DSO {WC["dso"]}d, short-term unearned revenue {WC["ucr"]:.1%}, '
              f'accounts payable {WC["ap"]:.1%} of revenue); June is the seasonal billing peak, so year ends are compared with year ends.')
AM = MAN['amort_sched']

V0 = {
    'out_hdr': 'Q1 FY27 guidance — 29-Jul-2026 call, re-filed in 8-K 2-Sep-2026 (Item 7.01, Ex. 99.1; SEC-filed)',
    'out_note': ('Total revenue $89.85–90.95bn: Agents and Infra $75.15–75.75bn, Devices and Consumer $14.7–15.2bn (adjusted to the FY27 segments); COGS $29.6–29.8bn; '
                 'operating expenses $16.8–16.9bn; other income ~$(100)m ex OpenAI; operating margin relatively flat y/y; ETR ~20%; capex over $50bn. '
                 'Cost rows: Bull = low end, Bear = high end.'),
    'outlook': [
        ('out_ai', '  Q1 FY27 Agents and Infra revenue ($m)', tuple(GQ['ai_rev']), 'm', 'Guided $75.15–75.75bn (8-K 2-Sep-2026). Drives the Agents and Infra growth calibration.'),
        ('out_dc', '  Q1 FY27 Devices and Consumer revenue ($m)', tuple(GQ['dc_rev']), 'm', 'Guided $14.7–15.2bn. Drives the Devices and Consumer growth calibration.'),
        ('out_cogs', '  Q1 FY27 cost of revenue ($m; Bull = low end)', tuple(GQ['cogs']), 'm', 'Guided $29.6–29.8bn. With opex, sets the guided operating income for the margin calibration.'),
        ('out_opex', '  Q1 FY27 operating expenses ($m; Bull = low end)', tuple(GQ['opex']), 'm', 'Guided $16.8–16.9bn. Drives Q1/27E R&D, S&M and G&A (prior-year mix).'),
    ],
    'points': [
        ('pt_oth', '  Q1–Q4/27E other income (expense), net per quarter ($m)', GQ['other_inc'], 'm', 'Guided ~$(100)m for Q1 excluding any impact from investments in OpenAI; held for Q2–Q4 (assumption).'),
        ('pt_etr', '  FY27E effective tax rate', GQ['etr'], 'pct', 'Guided ~20% for Q1 (8-K 2-Sep-2026); held for the year (assumption). FY26 19.4%.'),
        ('pt_dps', '  FY27E quarterly dividend declared ($)', 1.00, 'ps', 'Assumption: $0.91 + ~10% (Sep-2025 raise +9.6%); the Sep-2026 declaration is not in the EDGAR filings reviewed. Update when filed.'),
        ('pt_fyg', '  FY27 revenue / operating-income growth "double-digit" — UNVERIFIED', GU['rev_growth_min'], 'pct',
         'UNVERIFIED: 29-Jul-2026 call commentary as reported by third parties (MarketBeat, Nasdaq, Motley Fool transcript); not in an SEC filing. Information row only — not used as a solve target.'),
    ],
    'liq': (60000.0, 70000.0, 80000.0),
    'liq_note': ('Cash & short-term investments $76.8bn at 30-Jun-2026 ($94.6bn a year earlier) against $40.3bn of debt and $66.6bn of finance leases. Surplus above the target is '
                 'returned through buybacks ($40.6bn left on the $60bn program); Bull = lower target (more buybacks). Assumption.'),
    'bbmin': (25000.0, 20000.0, 15000.0),
    'bbmin_note': 'FY2026 repurchases $22.3bn (program $16.7bn + $5.6bn tax withholding); FY2025 $18.4bn. Minimum per year by scenario (assumption).',
    'blocks': [
        ('AI_AZ', 'AI_AZ (Azure y/y)', 0.06, -0.10, [0.32, 0.26, 0.21, 0.17],
         'FY26 +40% restated (Q4 +42%); Q1 FY27 guided +44–45% cc; capacity-constrained AI demand, commercial RPO +84% to $678bn. Decelerating with scale (assumption).', 'pct', False),
        ('AI_M365', 'AI_M365 (Microsoft 365 cloud y/y)', 0.03, -0.05, [0.14, 0.13, 0.12, 0.11],
         'FY26 +18.5% (incl. GitHub, Security Copilot); Q1 FY27 M365 commercial cloud ~17% cc; seats +7%, ARPU from E5 / Copilot.', 'pct', False),
        ('AI_PSL', 'AI_PSL (Productivity & server licensing y/y)', 0.03, -0.04, [-0.02, -0.02, -0.03, -0.03],
         'FY26 +5.4% (lumpy on-premises server / Office transactional); Q1 FY27 guided to decline low-single digits; long-run cloud migration.', 'pct', False),
        ('AI_IND', 'AI_IND (Industry solutions y/y)', 0.03, -0.04, [0.09, 0.09, 0.08, 0.08], 'FY26 +10.5%; Q1 FY27 cloud high-single digits (Dynamics 365, LinkedIn Talent / Sales, healthcare).', 'pct', False),
        ('AI_FSS', 'AI_FSS (Frontier & support services y/y)', 0.02, -0.03, [0.04, 0.04, 0.03, 0.03], 'FY26 +6.4% (enterprise support, industry consulting).', 'pct', False),
        ('DC_SA', 'DC_SA (Search & advertising y/y)', 0.03, -0.05, [0.08, 0.08, 0.07, 0.07], 'FY26 +12% (ex-TAC +14% restated); Q1 FY27 mid- to high-single digits.', 'pct', False),
        ('DC_XBOX', 'DC_XBOX (XBOX y/y)', 0.04, -0.06, [0.02, 0.03, 0.03, 0.03], 'FY26 −7%; Q1 FY27 content & services down mid-single digits, hardware down.', 'pct', False),
        ('DC_WIN', 'DC_WIN (Windows OEM & devices y/y)', 0.03, -0.04, [0.00, 0.01, 0.01, 0.01],
         'FY26 −1.3% (Windows 10 end-of-support pull-forward); Q1 FY27 guided down low-twenties on that comparison; PC-market growth thereafter.', 'pct', False),
        ('AI_M', 'AI_M (Agents and Infra operating margin)', 0.02, -0.04, [0.470, 0.465, 0.470, 0.475],
         'FY26 50.9% (FY25 50.6%; Q4/26 49.3%); depreciation from the FY26–27 capex wave (FY26 $34bn, +56%) compresses margins before operating leverage. Assumption.', 'pct', False),
        ('DC_M', 'DC_M (Devices and Consumer operating margin)', 0.02, -0.03, [0.290, 0.290, 0.295, 0.300], 'FY26 29.6% (FY25 28.1%); Q4/26 25.1% after XBOX charges.', 'pct', False),
        ('MNA', 'M&A_spend (acquisitions, $m)', 2000.0, -3000.0, [3000.0, 3000.0, 3000.0, 3000.0],
         'Acquisition spend FY24 $69.1bn (Activision), FY25 $6.0bn, FY26 $1.7bn. Base $3bn p.a. of bolt-ons (assumption).', 'm', True),
    ],
    'guid_note': GQ['note'],
    'unver_note': GU['status'] + ': ' + GU['text'] + ' Sources: ' + '; '.join(GU['src']),
    'dep_rate': {2027: 0.17, 2028: 0.17, 2029: 0.17, 2030: 0.17, 2031: 0.17},
    'dep_note': ('FY2026: $34.3bn ÷ opening net PP&E $205.0bn = 16.7%; FY2025 16.2%. Held at 17% (assumption): server lives 2–6 years; the Q1 FY27 outlook mentions a '
                 '"useful life update" that is not quantified in the filings.'),
    'amort_sched': {2027: AM['2027'], 2028: AM['2028'], 2029: AM['2029'], 2030: AM['2030'], 2031: AM['2031']},
    'amort_note': 'FY2026 10-K schedule of intangibles held at 30-Jun-2026: FY27 $3,097m, FY28 $2,141m, FY29 $1,944m, FY30 $1,477m, FY31 $1,128m, thereafter $8,822m.',
    'sbc_pct': 0.037, 'rd_pct': 0.105, 'sm_pct': 0.075, 'ga_pct': 0.022, 'etr': 0.19,
    'capex_fy27': 200000.0,
    'capex_note': ('Capital expenditures including finance leases (Microsoft\'s capex metric): FY2026 $140.6bn ($115.9bn cash + $24.6bn finance-lease assets), FY2025 $85.1bn. '
                   'FY2027E $200bn = Q1 guidance "over $50 billion" × 4 (assumption). 2028E+: % of revenue.'),
    'fl_share': 0.175,
    'capex_pct': {2028: 0.45, 2029: 0.40, 2030: 0.35, 2031: 0.32},
    'dps_g': 0.10, 'def_tax': 0.0, 'stock_iss': 2000.0, 'iss_sh': 30.0, 'dil': 24.0,
    'repay': {2027: 9250.0, 2028: 0.0, 2029: 2001.0, 2030: 0.0, 2031: 500.0},
    'mat': {2027: 0.0, 2028: 2001.0, 2029: 0.0, 2030: 500.0, 2031: 0.0},
    'min_liq': 40000.0,
    'r_debt': 0.015, 'r_debt_note': ('FY2026 interest expense $3,051m incl. $2,547m finance-lease interest → ~$0.5bn on ~$42bn of notes (~1.2%) after capitalised interest '
                                    '(datacenter construction). 1.5% input (assumption).'),
    'r_cash': 0.035, 'r_cash_note': 'FY2026 interest & dividend income $3,301m on average cash & short-term investments of ~$86bn (~3.9%). 3.5% input (assumption).',
    'fl_prin_r': 0.067,
    'wc': WC,
    'px': MAN['price']['px'],
    'px_note': 'VALUATION — share price input $541 (user-provided, 12-Oct-2026); update to market',
    'mna_note': 'Scenario lever M&A_spend (Base $3bn p.a.). Funded from cash.',
}

# ====================================================================== compute engine
def hist_value(row, col):
    kw = row.kw; path = kw.get('data')
    if not path: return None
    if kw.get('ppa'):
        v = MAN['ppa'].get(pkey(col), {}).get(path.split('_', 1)[1])
        return v if kw.get('textdata') else num(v)
    v = xget(pkey(col), path)
    return v if kw.get('textdata') else num(v)

def first(kw, keys):
    for k in keys:
        if k in kw: return k
    return None

def compute(row, col):
    kw = row.kw; k = col.kind
    CTX.col = col
    if kw.get('ann') and col.q: return None, False
    if k in HISTK:
        if 'a26_in' in kw and col.c == A26: return kw['a26_in'], True
        if 'a26' in kw and col.c == A26: return kw['a26'](), False
        v = hist_value(row, col)
        if v is not None: return v, True
        for key in ('hist', 'all'):
            if key in kw:
                return kw[key](), False
        return None, False
    if k == 'QE':
        if row.spec2: return None, False
        if 'qe_in' in kw:
            v = kw['qe_in']; return (v[col.q] if isinstance(v, dict) else v), True
        key = first(kw, ('qe', 'fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    if k == 'E26':
        if 'e26_in' in kw: return kw['e26_in'], True
        if 'e26' in kw: return kw['e26'](), False
        if kw.get('avg') and ('qe' in kw or 'qe_in' in kw): return f'=AVERAGE({Q1}{row.r}:{QE}{row.r})', False
        if kw.get('flow') and not row.spec2 and any(x in kw for x in ('qe', 'qe_in', 'fc')): return S26(row.key), False
        key = first(kw, ('fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    if k == 'AE':
        if 'ae_in' in kw:
            v = kw['ae_in']; return (v[col.year] if isinstance(v, dict) else v), True
        key = first(kw, ('ae', 'fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    return None, False

# ====================================================================== writers
CAGR_G = {'is_rev', 'g_rev', 'ai_rev', 'dc_rev', 'ai_az', 'ai_m365', 'ai_psl', 'ai_ind', 'ai_fss', 'dc_sa', 'dc_xbox', 'dc_win', 'mx_cloud',
          'o_adj', 'g_oi', 'r_adj', 'r_ebitda', 'is_oi', 'd_dda', 'e_adjni', 'e_eps', 'cf_cfo', 'fcf', 'bs_ta', 'bs_eq', 'sh_dil', 'is_gp', 'is_rd', 'is_sm',
          'is_ga', 'd_dep', 'd_amort', 'ai_oi', 'dc_oi', 'is_ni', 'eps_dil', 'dps', 'd_capex', 's3_pbp_r', 's3_ic_r', 's3_mpc_r'}
CAGR_A = {'ai_m', 'dc_m', 'g_oim', 'o_adj_m', 'r_adj_m', 'is_gm', 'is_oi_m', 'r_ebit_m', 'gm_rev', 'gm_ai', 'gm_dc', 'gm_az', 'gm_gm', 'gm_ebitda_m',
          'gm_op_m', 'gm_net_m', 'gm_rd', 'gm_sm', 'gm_ga', 'gm_adjop_m', 'fcf_m', 'fcf_conv', 'capex_pct', 'ra_roic', 'rn_ronta', 'roe', 'lv_x',
          'd_rd', 'd_sm', 'd_ga', 'd_sbc_pct', 'wc_nwc_p', 'is_etr', 'fcf_eb', 'mx_az_p', 'd_capex_pct', 'd_dep_r', 'wc_ucr'}
A = {y: ANN[y].c for y in ANN}

def style_cagr(ws, r, on):
    src = 16 if on else 12
    for c in CAGR:
        ws[f'{c}{r}']._style = copy.copy(TM[f'{L(CI(c) + TSHIFT)}{src}']._style)

def write_row(ws, row):
    r = row.r; arch = ARCH[row.arch]
    clone_row_style(ws, arch, r)
    ws.cell(r, 2).value = row.label or None
    if row.kw.get('note'): ws[f'{NOTE}{r}'] = row.kw['note']
    style_cagr(ws, r, row.key in CAGR_G or row.key in CAGR_A)
    if row.arch in ('section', 'block', 'blank', 'sub') and not row.kw.get('data'):
        return
    if row.key in DEFS:
        write_defs(ws, row); return
    nf = NF[row.nf]
    for col in COLS:
        v, inp = compute(row, col)
        if v is None: continue
        cell = ws[f'{col.c}{r}']
        cell.value = v
        if not row.kw.get('textdata'): cell.number_format = nf
        if inp:
            set_color(cell, BLUE, bold=True if col.kind in FCK else None)
            lk = row.kw.get('labkey')
            if lk and col.kind in HISTK and isinstance(v, (int, float)) and abs(v) > 1e-9:
                t = xget(pkey(col), lk)
                if t: comment(cell, f'Company items ({pkey(col)}): {t}', 380, 110)
        elif row.arch == 'growth' and col.kind in HISTK:
            set_color(cell, GRAY)
        else:
            set_color(cell, BLACK)
    rc = CELL_COMMENTS.get(row.key, {})
    for c, t in rc.items(): comment(ws[f'{c}{r}'], t)
    if row.key in CAGR_G:
        ws[f'{CAGR[0]}{r}'] = f'=IFERROR(({A[2031]}{r}/{A[2026]}{r})^(1/5)-1,"n/m")'
        ws[f'{CAGR[1]}{r}'] = f'=IFERROR(({A[2026]}{r}/{A[2021]}{r})^(1/5)-1,"n/m")'
        ws[f'{CAGR[2]}{r}'] = f'=IFERROR(({A[2026]}{r}/{A[2016]}{r})^(1/10)-1,"n/m")'
        ws[f'{CAGR[3]}{r}'] = f'=IFERROR(({A[2026]}{r}/{A[2006]}{r})^(1/20)-1,"n/m")'
        for c in CAGR: ws[f'{c}{r}'].number_format = '0%;\\(0%\\);"n/m"'
    elif row.key in CAGR_A:
        ws[f'{CAGR[0]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2027, 2032))}),"n/m")'
        ws[f'{CAGR[1]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2022, 2027))}),"n/m")'
        ws[f'{CAGR[2]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2017, 2027))}),"n/m")'
        ws[f'{CAGR[3]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2006, 2027))}),"n/m")'
        for c in CAGR: ws[f'{c}{r}'].number_format = NF['pct']
    t = LABEL_SOURCES.get(row.key)
    if t: comment(ws.cell(r, 2), t)

REL = 'Source: Microsoft earnings releases (Form 8-K Item 2.02, Ex. 99.1) — '
LABEL_SOURCES = {
    'ai_az': ('Source: Form 8-K 2-Sep-2026 (Item 7.01, Ex. 99.1 "FY27 Segments and Investor Metrics"): product and service revenue restated to the FY27 definitions, '
              'Q1 FY25 – Q4 FY26. Earlier periods are not available on this basis.'),
    'ai_oi': 'Source: Form 8-K 2-Sep-2026 (Ex. 99.1): Agents and Infra / Devices and Consumer revenue, cost of revenue, operating expenses and operating income restated, Q1 FY25 – Q4 FY26.',
    's3_h': REL + 'segment revenue and operating income tables, as originally reported (FY2016 structure; later recasts not applied).',
    'l2_h': REL + 'segment revenue and gross margin tables (FY2014–FY2015), as originally reported.',
    'l1_h': 'Source: 10-K segment notes (FY2006–FY2012) and FY2013 earnings releases, as originally reported.',
    'is_rev': (REL + 'income statements, as originally reported (FY2006–FY2012: 10-K Item 8). FY2016–FY2017 were restated for ASC 606 / 842 in the 8-K of 3-Aug-2017; '
               'the model keeps the original figures (cell comments).'),
    'd_dda': REL + 'cash-flow statements (three-month columns); depreciation and amortization of intangibles: 10-K notes (XBRL, annual).',
    'o_pub': REL + '"GAAP to Non-GAAP" / "As Adjusted" reconciliation tables (ngaap.json); items and the basis per period are in the cell comments of the published adjusted EPS row.',
    'e_pub': REL + 'reconciliation tables (ngaap.json). FY2012–Q1 FY14 published per share only.',
    'cf_ni': REL + 'cash-flow statements (three-month columns for quarters; Q1 FY26 restated from the Q2 FY26 six-month column).',
    'bs_cash': REL + 'balance sheets (period ends).',
    'ppa_cons': 'Source: 10-K business-combination notes — LinkedIn (FY2017 10-K), Nuance (final allocation, FY2024 10-K), Activision Blizzard (FY2024 10-K).',
    'sh_basic': REL + 'income statements (weighted average shares). No stock splits since Feb-2003.',
    'dps': 'Source: XBRL CommonStockDividendsPerShareDeclared (10-K / 10-Q); FY2006–07 from the 10-K income statements.',
    'ds_total': 'Source: balance sheets; FY2026 10-K Note 10 (debt) and Note 14 (leases).',
    'fl_liab': 'Source: 10-K lease notes (XBRL FinanceLeaseLiability, RightOfUseAssetObtainedInExchangeForFinanceLeaseLiability, FinanceLeasePrincipalPayments).',
    'cal_ai_gd': 'Source: 8-K 2-Sep-2026, Item 7.01 Ex. 99.1 ("FY27 Q1 Adjusted Outlook"), which re-files the outlook given on the 29-Jul-2026 earnings call.',
    'cal_fyg': 'UNVERIFIED — not in an SEC filing. ' + V0['unver_note'],
    'd_capex': 'Source: cash-flow statements (additions to property and equipment) + 10-K lease note (finance-lease right-of-use assets obtained, XBRL).',
}
CELL_COMMENTS = {}
def comments_setup():
    CELL_COMMENTS['is_rev'] = {A[2016]: 'FY2016 as originally reported ($85,320m). Restated for ASC 606 / 842 to $91,154m (8-K 3-Aug-2017).',
                               A[2017]: 'FY2017 as originally reported ($89,950m). Restated for ASC 606 / 842 to $96,571m (8-K 3-Aug-2017); FY2018 y/y in this row therefore crosses bases.'}
    CELL_COMMENTS['is_oi'] = {A[2008]: 'FY2008 as originally reported in the FY2008 10-K ($22,492m); later 10-K comparatives show $22,271m.',
                              A[2016]: 'FY2016 operating income as originally reported ($20,182m); $26,078m restated for ASC 606.',
                              A[2017]: 'FY2017 operating income as originally reported ($22,326m); $29,025m restated for ASC 606.'}
    CELL_COMMENTS['is_tax'] = {A[2016]: 'FY2016 reflects the early adoption of ASU 2016-09 effective 1-Jul-2015; Q1–Q3 FY16 are as originally reported, so they sum to $302m more tax than the year.'}
    CELL_COMMENTS['cf_dda'] = {QC[(2026, 1)].c: MAN['corrections']['Q1-26_cf']['note']}
    CELL_COMMENTS['d_dda'] = {QC[(2026, 1)].c: 'Restated Q1 FY26 (see the cash-flow row comment): $8,147m vs $13,061m as first reported.'}
    CELL_COMMENTS['s3_pbp_r'] = {A[2016]: 'FY2016 annual segment figures (Q4 FY16 release) differ from the sum of Q1–Q3 as first reported (re-allocations within the year).'}

DEFS = {}
def defs_setup():
    q = lambda y, n: QC[(y, n)].c
    DEFS['o_def'] = {
        'C': ('FY2006–FY2011: no adjusted measures published — adjusted = GAAP (flag)', None),
        A[2012]: ('FY2012–Q1 FY14: revenue deferrals for Windows / Office upgrade offers and pre-sales; FY2012 goodwill impairment; Q3 FY13 EC fine',
                  'FY2012 adjusted figures were first presented in the Q4 FY13 release (prior-year comparative).'),
        q(2015, 1): ('FY2015: impairment, integration & restructuring (Nokia)', 'Q1–Q3 FY15 first given adjusted figures in the FY2016 releases (re-presented comparatives) — that presentation is used.'),
        q(2016, 1): ('FY2016–FY2017: Windows 10 revenue deferrals + impairment / integration / restructuring', None),
        q(2018, 1): ('FY2018+ (ASC 606): no operating adjustments except Q2 FY23 severance, hardware impairment & lease consolidation ($1,171m)', None),
    }
    DEFS['e_def'] = {
        'C': ('FY2006–FY2011: no adjusted EPS published — adjusted = GAAP (flag)', None),
        A[2012]: ('FY2012–Q1 FY14: per-share bridges only — after-tax $ = per-share effect × diluted shares (labelled)', None),
        q(2018, 2): ('FY2018–FY2022: discrete tax items adjusted out (TCJA, intangible-property transfers, India withholding-tax decision)', None),
        q(2025, 1): ('FY2025–FY2026: impact from investments in OpenAI adjusted out (net income / EPS only)',
                     'Introduced in the Q1 FY26 release; FY2025 quarters take the adjusted figures re-presented in the FY2026 releases (the FY2025 releases published none).'),
    }
    DEFS['r_def'] = {
        'C': ('Model definitions throughout (Microsoft publishes no EBITDA): EBIT = operating income; D&A = cash-flow line "Depreciation, amortization, and other"', None),
    }
def write_defs(ws, row):
    for c, (txt, cm) in DEFS[row.key].items():
        cell = ws[f'{c}{row.r}']; cell.value = txt
        cell.font = Font(name='Calibri', sz=10, b=True, i=True, color='FFC00000')
        if cm: comment(cell, cm)

def write_scen(ws):
    srcrow = {'hdr': 9, 'out_hdr': 11, 'pt_hdr': 15, 'out': 12, 'pt': 16, 'lev': 28, 'tgt': 28, 'blk': 30, 'yr': 31}
    for spec_ in SCEN_ROWS:
        r, kind = spec_[0], spec_[1]
        for ci in range(CI(SL), CI(SNOTE) + 1):
            ws.cell(r, ci)._style = copy.copy(TM.cell(srcrow[kind], ci + TSHIFT)._style)
        if kind in ('hdr', 'out_hdr', 'pt_hdr'):
            _, _, a, b, c_, d_, e = spec_
            for col, v in zip([SL, SBULL, SBASE, SBEAR, SNOTE], [a, b, c_, d_, e]):
                if v is not None: ws[f'{col}{r}'] = v
        elif kind == 'out':
            _, _, lab, hi, mid, lo, note, nf = spec_
            ws[f'{SL}{r}'] = lab
            for col, v in zip([SBULL, SBASE, SBEAR], [hi, mid, lo]):
                ws[f'{col}{r}'] = v; ws[f'{col}{r}'].number_format = NF[nf]; set_color(ws[f'{col}{r}'], BLUE)
            ws[f'{SNOTE}{r}'] = note
        elif kind == 'pt':
            _, _, lab, _, val, _, note, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBASE}{r}'] = val; ws[f'{SBASE}{r}'].number_format = NF[nf]; set_color(ws[f'{SBASE}{r}'], BLUE); ws[f'{SNOTE}{r}'] = note
        elif kind == 'tgt':
            _, _, lab, a, b, c_, note = spec_
            ws[f'{SL}{r}'] = lab
            for col, v in zip([SBULL, SBASE, SBEAR], [a, b, c_]):
                ws[f'{col}{r}'] = v; ws[f'{col}{r}'].number_format = NF['m']; set_color(ws[f'{col}{r}'], BLUE)
            ws[f'{SNOTE}{r}'] = note
        elif kind == 'blk':
            _, _, lab, db, mid, de, note, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBULL}{r}'] = db; ws[f'{SBASE}{r}'] = mid; ws[f'{SBEAR}{r}'] = de; ws[f'{SNOTE}{r}'] = note
            for col in (SBULL, SBEAR):
                ws[f'{col}{r}'].number_format = '\\+#,##0;\\-#,##0' if nf == 'm' else '\\+0.0%;\\-0.0%'
        elif kind == 'yr':
            _, _, lab, hdr, base, clamp, _, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBASE}{r}'] = base
            if clamp:
                ws[f'{SBULL}{r}'] = f'=MAX(0,{SBASE}{r}+${SBULL}${hdr})'; ws[f'{SBEAR}{r}'] = f'=MAX(0,{SBASE}{r}+${SBEAR}${hdr})'
            else:
                ws[f'{SBULL}{r}'] = f'={SBASE}{r}+${SBULL}${hdr}'; ws[f'{SBEAR}{r}'] = f'={SBASE}{r}+${SBEAR}${hdr}'
            for col in (SBULL, SBASE, SBEAR): ws[f'{col}{r}'].number_format = NF[nf]

def write_header(ws):
    for r in range(1, 8):
        clone_row_style(ws, r, r, 2, CI(LASTC))
    ws[f'{NOTE}1'] = 'SCENARIO SWITCH ▼ (Bull / Base / Bear) — drives all scenario-linked forecast drivers'
    ws[f'{NOTE}2'] = 'Base'
    comment(ws[f'{NOTE}2'], 'Scenario toggle. Choose Bull / Base / Bear. Drives: Q1 FY27 guidance end-points (segment revenue high / mid / low; COGS and opex low / mid / high), '
            '2028E–2031E product & service growth, segment margins, M&A spend, and the liquidity target and minimum for buybacks '
            f'(scenario table {SL}:{SNOTE}).')
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type='list', formula1='"Bull,Base,Bear"', allow_blank=False); ws.add_data_validation(dv); dv.add(f'{NOTE}2')
    ws['B2'] = 'Microsoft Corporation (NASDAQ: MSFT) — 3-Statement Financial Model'
    ws['B3'] = ('US GAAP as reported · USD millions (per-share data in USD; shares in millions) · Fiscal years ended 30 June (2027 = Jul-2026 – Jun-2027) · '
                'Source: Microsoft Forms 10-K / 10-Q, Form 8-K earnings releases (Ex. 99.1) and the 8-K of 2-Sep-2026 (FY27 segments, restated history, Q1 FY27 outlook), '
                'SEC EDGAR (CIK 0000789019); XBRL company facts. FY27 full-year call commentary is UNVERIFIED (information rows only).')
    ws['B4'] = ('Basis notes: as originally reported (FY2016–17 ASC 606 / 842 restatement and later segment recasts not applied; memo comments). Segment build on the FY27 structure '
                '(Agents and Infra / Devices and Consumer), restated history Q1/25–Q4/26 only; earlier structures as memo. Acquisitions: LinkedIn Dec-16, Nuance Mar-22, '
                'Activision Blizzard Oct-23. Q1/27E calibrated to the Q1 FY27 outlook; Q2–Q4/27E carry the Q1 calibration. Share price $541 (user, 12-Oct-2026).')
    ws[f'{CAGR[0]}5'] = 'Forecast'
    for c in CAGR[1:]: ws[f'{c}5'] = 'Historical'
    ws['B6'] = '(USDm)'
    for col in COLS: ws[f'{col.c}6'] = col.label
    ws[f'{NOTE}6'] = 'Modelling Notes'
    for c, a, b in zip(CAGR, ['5Y CAGR', '5Y CAGR', '10Y CAGR', '20Y CAGR'], ["26-'31", "21-'26", "16-'26", "06-'26"]):
        ws[f'{c}6'] = a; ws[f'{c}7'] = b

def set_columns(ws):
    for k in list(ws.column_dimensions.keys()): del ws.column_dimensions[k]
    ws.column_dimensions['A'].width = TM.column_dimensions['A'].width
    ws.column_dimensions['B'].width = TM.column_dimensions['B'].width
    wid = {'A': 13.0, 'Q': 9.5703125, 'QE': 11.42578125, 'E26': 14.140625, 'AE': 13.0}
    for col in COLS:
        cd = ws.column_dimensions[col.c]; cd.width = wid[col.kind]
        if col.q:
            cd.outlineLevel = 1; cd.hidden = True
    for ci in range(CI(NOTE), CI(LASTC) + 1):
        ws.column_dimensions[L(ci)].width = TM.column_dimensions[L(ci + TSHIFT)].width or 13.0
    ws.sheet_format.outlineLevelCol = 1

def main():
    import openpyxl
    wb = openpyxl.load_workbook(TEMPLATE)
    if 'DCF_old' in wb.sheetnames: del wb['DCF_old']
    ws = wb['Model']
    blank_style = copy.copy(ws['A1']._style)
    for rowc in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for c in rowc:
            c.value = None; c.comment = None; c._style = copy.copy(blank_style)
    for k in list(ws.row_dimensions.keys()):
        rd = ws.row_dimensions[k]; rd.outlineLevel = 0; rd.hidden = False
    ws.data_validations.dataValidation = []
    scen_layout(V0)
    build_spec(V0)
    n1 = len(fw.SPEC)
    build_spec2(V0)
    for i, row in enumerate(fw.SPEC): row.spec2 = i >= n1
    last = number_rows(8)
    for r in range(1, last + 1): ws.row_dimensions[r].height = 15.0
    ws.row_dimensions[2].height = 21.0
    defs_setup(); comments_setup()
    write_header(ws)
    for row in fw.SPEC: write_row(ws, row)
    write_scen(ws)
    set_columns(ws)
    ws.freeze_panes = 'C8'
    ws.sheet_view.zoomScale = 70
    import other_sheets as osh
    osh.do_all(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    osh.inject_charts(OUT)
    print('rows', fw.SPEC[-1].r, 'saved', OUT)
    return wb

RECALC = os.environ.get('XLSX_RECALC', '/mnt/skills/public/xlsx/scripts/recalc.py')
def recalc(path, timeout=600):
    import subprocess
    out = subprocess.run(['python3', RECALC, path, str(timeout)], capture_output=True, text=True, cwd=os.path.dirname(RECALC))
    try:
        return json.loads(out.stdout)
    except Exception:
        return {'raw': out.stdout[-2000:], 'err': out.stderr[-2000:]}

def scenario_snapshots():
    import openpyxl
    import other_sheets as osh
    snaps = {}
    for scen in ('Bull', 'Bear'):
        p = os.path.join(BASE, f'MSFT_{scen}.xlsx')
        wb = openpyxl.load_workbook(OUT); wb['Model'][f'{NOTE}2'] = scen; wb.save(p)
        res = recalc(p); print(scen, res.get('status'), res.get('total_errors'))
        v = openpyxl.load_workbook(p, data_only=True)['Bull-Base-Bear']
        snaps[scen] = {(r, c): v[f'{c}{r}'].value for r in range(9, 51) for c in 'CDEFGHIJKL' if v[f'{c}{r}'].value is not None}
        os.remove(p)
    wb = openpyxl.load_workbook(OUT)
    osh.snapshot_values(wb, snaps['Bull'], snaps['Bear'])
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    osh.inject_charts(OUT)

NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
def inject_values(src_xlsx, recalc_xlsx):
    """Store cached formula results (from a LibreOffice-recalculated copy) in the openpyxl file, so the workbook reads correctly in any
    viewer; fullCalcOnLoad stays set so Excel recalculates on open. Same approach as the LECO / ESAB finishing step."""
    import zipfile, shutil, openpyxl
    from lxml import etree
    wbv = openpyxl.load_workbook(recalc_xlsx, data_only=True)
    zin = zipfile.ZipFile(src_xlsx)
    wbxml = etree.fromstring(zin.read('xl/workbook.xml'))
    rels = etree.fromstring(zin.read('xl/_rels/workbook.xml.rels'))
    rid2t = {r.get('Id'): r.get('Target') for r in rels}
    sheets = {}
    for s in wbxml.iter('{%s}sheet' % NS):
        rid = s.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
        sheets['xl/' + rid2t[rid].lstrip('/').replace('xl/', '')] = s.get('name')
    tmp = src_xlsx + '.tmp'
    zout = zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED)
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename in sheets:
            ws = wbv[sheets[item.filename]]
            root = etree.fromstring(data)
            for c in root.iter('{%s}c' % NS):
                f = c.find('{%s}f' % NS)
                if f is None: continue
                v = ws[c.get('r')].value
                old = c.find('{%s}v' % NS)
                if old is not None: c.remove(old)
                if v is None: continue
                ve = etree.SubElement(c, '{%s}v' % NS)
                if isinstance(v, bool):
                    c.set('t', 'b'); ve.text = '1' if v else '0'
                elif isinstance(v, (int, float)):
                    if 't' in c.attrib: del c.attrib['t']
                    ve.text = repr(float(v)) if isinstance(v, float) else str(v)
                elif isinstance(v, str) and v.startswith('#'):
                    c.set('t', 'e'); ve.text = v
                else:
                    c.set('t', 'str'); ve.text = str(v)
            data = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
        zout.writestr(item, data)
    zout.close(); zin.close()
    shutil.move(tmp, src_xlsx)

def store_values():
    import shutil
    p = os.path.join(BASE, 'MSFT_recalc.xlsx')
    shutil.copy(OUT, p)
    res = recalc(p); print('recalc', res.get('status'), res.get('total_errors'))
    inject_values(OUT, p)
    os.remove(p)

if __name__ == '__main__':
    main()
    if '--snap' in sys.argv: scenario_snapshots()
    if '--recalc' in sys.argv: store_values()
