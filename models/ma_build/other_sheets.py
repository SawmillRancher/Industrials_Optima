"""Rewire the Bull-Base-Bear, DCF and Charts tabs of the MLM template (current layout: DCF with live valuation date, fade period,
implied ROIC, target-IRR entry price, reverse DCF and Mauboussin EV/NOPAT; Bull-Base-Bear with buyback rows) to the MA Model rows;
keep the template chart XML. The template's DCF_old tab is not carried.

Column mapping: Mastercard's calendar fiscal years and history window match the template's own columns (2025A in BV, 2026E in CA,
2027E-2030E in CB-CE, notes in CF), so only row references are remapped (template row -> MA row key). The 2026E UFCF is excluded
because net debt is taken at YE2026E (valuation date Oct-2026)."""
import re, zipfile, shutil
from openpyxl.utils import get_column_letter as L
from fw import R, TEMPLATE, ANN, NOTE

# template Model row -> MA row key
MAP = {158: 'is_rev', 164: 'r_cash', 166: 'is_oi', 182: 'is_dda', 186: 'sh_dil', 211: 'r_adj', 212: 'r_adj_m', 219: 'r_ebit', 221: 'r_adj',
       229: 'e_adjni', 231: 'e_eps', 238: 'gm_pn', 245: 'gm_ebitda_m', 247: 'gm_op_m', 260: 'cf_wc', 262: 'cf_capex', 263: 'cf_acq',
       280: 'fcf', 282: 'fcf_m', 290: 'bb_px', 291: 'bb_cash', 292: 'bb_sh', 331: 'bs_nci', 408: 'ra_ic', 411: 'ra_roic', 422: 'rn_ronta',
       428: 'lv_nd', 433: 'v_px', 434: 'v_sh'}
CMAP = {}
REF = re.compile(r"Model!(\$?)([A-Z]{1,3})(\$?)(\d+)")
HIST = [ANN[y].c for y in range(2006, 2026)]

def remap(f):
    def rep(m):
        c, r = m.group(2), int(m.group(4))
        c2 = CMAP.get(c, c)
        r2 = R[MAP[r]] if r in MAP else r
        if c == 'CF' and r == 2: r2 = 2
        return f'Model!{m.group(1)}{c2}{m.group(3)}{r2}'
    return REF.sub(rep, f)

def hist_stat(f):
    """LT historical AVERAGE / MIN / MAX over the template's annual columns -> MA FY2006-FY2025 annual columns."""
    m = re.match(r'=IFERROR\((AVERAGE|MIN|MAX)\(Model!C(\d+),', f)
    if not m: return None
    fn, r = m.group(1), int(m.group(2))
    rr = R[MAP[r]]
    return f'=IFERROR({fn}({",".join(f"Model!{c}{rr}" for c in HIST)}),"")'

TEXT = {
    'Bull-Base-Bear': {
        'B2': 'Mastercard Incorporated (MA) — Bull / Base / Bear Case P&L Summary',
        'B3': ('Adjusted figures (company definitions: adjusted operating income, adjusted EPS; adjusted EBITDA is a model definition) · USD millions · Live panel linked to Model — '
               'toggle the scenario on the Model sheet (cell CF2). Bull and Bear panels are value snapshots generated from the Model with the switch set to each case. 2026E = Q1–Q2 actual + Q3/Q4 '
               'calibrated to the FY26 outlook (net revenue growth high end of low double digits; adjusted opex growth low double digits) — outlook UNVERIFIED (earnings call, not an SEC filing).'),
        'G26': 33, 'H26': 33, 'S26': 37, 'T26': 37, 'AD26': 26, 'AE26': 26, 'AF26': 26,
    },
    'DCF': {
        'B2': 'DCF VALUATION — MASTERCARD INCORPORATED (MA)',
        'C5': 0.0425, 'D5': 'Assumption (~US 10Y Treasury yield, Oct-2026; government rate sources are not reachable from the build environment); update to market',
        'D6': 'Assumption: US implied ERP ~4.5–5.5%',
        'C7': 1.00, 'D7': 'Assumption: payment networks (V, AXP, PYPL) ~0.9–1.2; volume-linked, recurring network revenue with consumer-spending cyclicality',
        'C9': 0.047, 'D9': 'June-2026 notes priced at 4.325%–5.0% (SOFR + 0.39% floaters); A+/Aa3-rated; commercial paper 3.84% (Q2-26 10-Q Note 9)',
        'C10': 0.205, 'D10': 'Model 2027E+ effective tax rate (non-GAAP rate guided 20–21% for H2-26, UNVERIFIED)',
        'C12': 0.03, 'D12': 'Net debt ~$13bn at 30-Jun-2026 vs ~$520bn market cap ($595 × ~879m diluted shares)',
        'D15': 'Periods = (31-Dec of the year − today) / 365.25; cash flows assumed at year end. 2026E UFCF excluded (its cash flows sit in YE2026E net debt; valuation date Oct-2026).',
        'B18': 'Adj. EBIT (adjusted operating income, full-cost) less litigation provisions & restructuring (cash)',
        'B21': 'plus D&A (depreciation + amortization of capitalized software and acquired intangibles)',
        'B22': '− Capex (property, equipment & capitalized software) and acquisitions (M&A lever 2027E–30E)',
        'B26': 'Fade-period growth (2031E–35E; 9.0% → 6.0% linear)', 'I26': 0.09, 'M26': 0.06, 'C27': 0.04,
        'D32': ("Mastercard's adjusted operating income expenses share-based compensation and the amortization of acquired intangibles and capitalized software (neither is adjusted out); "
                'D&A is added back in row 21 and capex (incl. capitalized software) deducted in row 22, so the basis is consistent full cost. Litigation provisions and restructuring '
                '(cash; excluded from adjusted results) are deducted inside row 18 ($450m a year assumed from 2027E). Row 22 includes acquisitions from the M&A lever in 2027E–30E; '
                'the fade period and terminal value assume no new acquisitions. The 2026E column (incl. the BVNK consideration) is excluded: those cash flows are in YE2026E net debt.'),
        'D43': 'Opening IC = Model 2025A book IC (equity + net debt) + the BVNK consideration (2026E), rolled forward with capex + M&A − D&A + ΔWC.',
        'B48': 'Less: Net debt + NCI (Model 2026E year end)',
        'D48': 'Net debt = total debt − unrestricted cash − investments (restricted cash and customer security deposits excluded); NCI and redeemable NCI deducted (nil at YE2026E).',
        'B50': 'Diluted shares (m, Model 2026E year end)',
        'D76': 'Rationale (Mastercard characteristics)',
        'C78': 0.11, 'D78': 'Model adjusted EBIT CAGR 2026–30E: high-single-digit GDV growth, rising yield and mid-teens value-added services growth with modest margin expansion → ~11% NOPAT growth',
        'C79': 0.40, 'D79': 'A two-sided network with ~60% adjusted operating margins and capex of ~4% of revenue earns very high returns on organic investment; acquired VAS businesses (Recorded Future, BVNK) earn far less on their purchase price',
        'C80': 15, 'D80': 'Global acceptance network, brand and issuer relationships sustain a long competitive-advantage period; risks are regulation (interchange, routing — CCCA), account-to-account and stablecoin rails, and rebate competition',
        'B90': 'VALUATION  —  Applied to 2028E NOPAT',
        'B91': '2028E Adj. EBIT less litigation provisions & restructuring',
        'B108': 'MA market-implied EV/NOPAT (on 2028E NOPAT)',
    },
    'Charts': {
        'B2': 'Mastercard Incorporated — Operating Performance (2011–2025) and Model Forecast (2026E–2030E)',
        'B3': ('USDm · Linked to Model. Bars = contribution to net revenue growth (pp): payment network vs value-added services (from 2021; earlier years = total net revenue growth). '
               'Adjusted EBITDA is a model definition (adjusted operating income + D&A). ROIC on adjusted EBIT (full-cost).'),
        'B8': 'Payment network contribution to growth (pp)', 'B9': 'Value-added services contribution to growth (pp)', 'B12': 'Adjusted EBITDA (model)',
        'B13': 'Operating income (GAAP)', 'B14': 'Adj. EBITDA Margin %', 'B15': 'Operating Margin %',
    },
}
BBB_ROWS = {9: 'Net revenue', 10: '   Net revenue y/y %', 11: '   Payment network y/y %', 12: 'Adjusted EBITDA (model)',
            15: 'Operating income (GAAP)', 16: '   Operating income y/y %', 17: '   Operating margin (GAAP) %', 19: 'Adjusted Net Income',
            35: 'Adjusted EBITDA (leverage basis)', 36: '   Adjusted EBITDA % of net revenue'}
CHART_TITLES = {
    'Aggregates Revenue Drivers: Shipment Volume vs. Price (ASP), y/y %': 'Net Revenue Growth: Payment Network vs. Value-Added Services Contribution, pp',
    'Adjusted EBITDA &amp; Earnings from Operations (USDm)': 'Adjusted EBITDA (model) &amp; Operating Income (USDm)',
}

def build_charts(ws):
    for i, y in enumerate(range(2011, 2031)):
        c = L(3 + i); mc = ANN[y].c
        ws[f'{c}7'] = y if y <= 2025 else f'{y}E'
        ws[f'{c}8'] = f'=IFERROR(Model!{mc}{R["gc_pn"]}+0,0)'
        ws[f'{c}9'] = f'=IFERROR(Model!{mc}{R["gc_vas"]}+0,0)'
        ws[f'{c}11'] = f'={c}7'
        ws[f'{c}12'] = f'=Model!{mc}{R["r_adj"]}'
        ws[f'{c}13'] = f'=Model!{mc}{R["is_oi"]}'
        ws[f'{c}14'] = f'=Model!{mc}{R["gm_ebitda_m"]}'
        ws[f'{c}15'] = f'=Model!{mc}{R["gm_op_m"]}'
        ws[f'{c}17'] = f'={c}7'
        ws[f'{c}18'] = f'=IFERROR(Model!{mc}{R["ra_roic"]}+0,NA())'
        ws[f'{c}19'] = f'=IFERROR(Model!{mc}{R["rn_ronta"]}+0,NA())'
        ws[f'{c}20'] = f'=IFERROR(Model!{mc}{R["gm_ebitda_m"]}+0,NA())'
        ws[f'{c}21'] = f'=IFERROR(Model!{mc}{R["fcf_m"]}+0,NA())'

PE_NOTE = ('P/E on adjusted EPS (company definition). MA at $595 (user price, 13-Oct-2026) trades at ~30x 2026E adjusted EPS (~$19.9) and ~26x 2027E (~$23.0). '
           'Exit multiples are assumptions: Base 33x, Bull 37x, Bear 26x applied to 2029E–30E, against MA\'s ~28–40x forward range over 2016–25 (not sourced from EDGAR; '
           'market data are not reachable from the build environment).')
def do_all(wb):
    for name in ('Bull-Base-Bear', 'DCF', 'Charts'):
        ws = wb[name]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and 'Model!' in c.value:
                    h = hist_stat(c.value) if name == 'Bull-Base-Bear' else None
                    c.value = h or remap(c.value)
        for addr, t in TEXT.get(name, {}).items():
            ws[addr] = t
    b = wb['Bull-Base-Bear']
    for r, t in BBB_ROWS.items():
        for c in ('B', 'N', 'Z'): b[f'{c}{r}'] = t
    for c0 in ('C', 'O', 'AA'):
        i = b[f'{c0}7'].column
        b.cell(7, i).value = '2025A'
        for k in range(1, 6): b.cell(7, i + k).value = f'{2025 + k}E'
        b.cell(7, i + 6).value = "25-'30 CAGR"
    from openpyxl.comments import Comment
    cm = Comment(PE_NOTE, 'MA model')
    cm.width, cm.height = 420, 120; b['N26'].comment = cm
    cm2 = Comment(f'Snapshot (values) of Model rows {R["bb_cash"]}, {R["bb_px"]} and {R["bb_sh"]} (Share Buyback Schedule) with the Model sheet scenario switch (cell {NOTE}2) set to Bull / Bear. '
                  'Base panel (D40:H42) is live-linked.', 'MA model')
    cm2.width, cm2.height = 420, 90; b['P40'].comment = cm2
    b['J2'] = 'Share price ($):'
    for c in ('B', 'N', 'Z'):
        if b[f'{c}40'].value: b[f'{c}40'] = '   Share repurchases ($m; 2026E = H2)'
    d = wb['DCF']
    a25, a26 = ANN[2025].c, ANN[2026].c
    d['D31'] = '=0*D24*D30'
    d['D34'] = f'=Model!{a25}{R["ra_ic"]}+Model!{a26}{R["bv_cons"]}'
    d['C48'] = f'=Model!{a26}{R["lv_nd"]}+Model!{a26}{R["bs_nci"]}+Model!{a26}{R["bs_mez"]}'
    d['C50'] = f'=Model!{a26}{R["v_sh"]}'
    d['C52'] = f'=Model!{a26}{R["v_px"]}'
    d['B31'] = 'PV of UFCF (2026E excluded: in YE2026E net debt)'
    build_charts(wb['Charts'])

def snapshot_values(wb, bull, bear):
    ws = wb['Bull-Base-Bear']
    cols = 'CDEFGHIJKL'
    for snap, off in ((bull, 'OPQRSTUVWX'), (bear, ['AA', 'AB', 'AC', 'AD', 'AE', 'AF', 'AG', 'AH', 'AI', 'AJ'])):
        for (r, c), v in snap.items():
            tgt = f'{off[cols.index(c)]}{r}'
            cur = ws[tgt].value
            if isinstance(cur, str) and cur.startswith('='): continue
            if r in (7, 26): continue
            ws[tgt] = v

def inject_charts(path):
    """Replace openpyxl-rewritten chart parts with the template chart XML (combo charts, colours, layout), retitled, ranges kept
    (Charts!C:V = 2011-2030E) and cached values removed so Excel redraws from the MA cells."""
    tz = zipfile.ZipFile(TEMPLATE)
    tmpl = {n: tz.read(n) for n in tz.namelist() if n.startswith('xl/charts/chart')}
    tmp = path + '.tmp'
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in tmpl:
                t = tmpl[item.filename].decode('utf-8')
                for a, b2 in CHART_TITLES.items(): t = t.replace(a, b2)
                t = re.sub(r'<c:numCache>.*?</c:numCache>', '', t, flags=re.S)
                t = re.sub(r'<c:strCache>.*?</c:strCache>', '', t, flags=re.S)
                data = t.encode('utf-8')
            zout.writestr(item, data)
    shutil.move(tmp, path)
