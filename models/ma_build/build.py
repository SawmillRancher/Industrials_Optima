"""Build MA_Model.xlsx (Mastercard Incorporated) from the MLM template (same layout, formatting and process as the prior models)."""
import sys, copy, os, json
from openpyxl.styles import Font
from fw import *
import fw
from spec import build_spec, scen_layout, SCEN_ROWS, SC
from spec2 import build_spec2

OUT = os.path.join(os.path.dirname(BASE), 'MA_Model.xlsx')
G26 = MAN['guidance_fy26']

# ====================================================================== inputs (V0)
def wc_inputs():
    q = DATA['2025']['x']
    rev, gdv = q['rev'], q['gdv']
    return {'dso': round(q['bs_ar'] / rev * 365), 'set': round(q['bs_set_a'] / gdv, 3), 'setl': round(q['bs_set_l'] / gdv, 3),
            'oca': round(q['bs_oca'] / rev, 3), 'onca': round(q['bs_onca'] / rev, 3), 'ap': round(q['bs_ap'] / rev, 3),
            'accr': round(q['bs_accr'] / rev, 3), 'ocl': round(q['bs_ocl'] / rev, 3)}
WC = wc_inputs()
WC['note'] = (f'Balance-sheet drivers held at FY2025 year-end levels: DSO {WC["dso"]}d; settlement assets / obligations {WC["set"]:.3f} / {WC["setl"]:.3f} × GDV; '
              f'prepaid & other current {WC["oca"]:.1%}, other noncurrent assets {WC["onca"]:.1%}, payables {WC["ap"]:.1%}, accrued expenses {WC["accr"]:.1%}, '
              f'other current liabilities {WC["ocl"]:.1%} of net revenue (customer-incentive prepayments and accruals grow with revenue). Year ends compared with year ends.')
DB = MAN['debt']['maturities']

V0 = {
    'out_hdr': 'FY26 outlook — Q2-26 earnings call (30-Jul-2026) — UNVERIFIED (not in an SEC filing)',
    'out_note': ('Net revenue growth at the high end of a low-double-digits rate (currency-neutral, ex-acquisitions; raised), FX ~+1 ppt; adjusted opex growth low double digits '
                 '(currency-neutral, ex-acquisitions), FX 0.5–1 ppt headwind; Q3 other expense ~$(125)m; non-GAAP tax rate 20–21% in Q3 and Q4. Source: third-party call '
                 'transcripts (web search); Mastercard\'s 8-K release has no outlook. Numeric end-points are model translations of the wording.'),
    'outlook': [
        ('out_g', '  FY26 net revenue growth, reported (UNVERIFIED)', tuple(G26['model_rev_growth_reported']), 'pct', G26['model_rev_note']),
        ('out_ox', '  FY26 adjusted operating expense growth, reported (UNVERIFIED; Bull = low end)', tuple(G26['model_opex_growth_reported']), 'pct', G26['model_opex_note']),
        ('bb26', '  H2-26E share repurchases ($m)', (7800.0, 7000.0, 6000.0), 'm',
         'H1-26 $8,933m (Q2 $4.9bn); $0.7bn repurchased 1–27 Jul; $7.8bn of authorization left at 27-Jul-2026 (a new program is usually approved in December). Assumption.'),
    ],
    'points': [
        ('pt_fxrev', '  FY26 FX effect on net revenue growth (pp; memo)', 0.010, 'pct', 'Outlook: ~1 ppt FX tailwind for the year (H1-26: +2.6 ppt) — UNVERIFIED.'),
        ('pt_da', '  Q3/Q4-26E depreciation & amortization ($m per quarter)', 325.0, 'm', 'Q2-26 $309m, Q1-26 $299m; rising with capitalized software and the June-2026 property additions.'),
        ('pt_inv', '  Q3/Q4-26E investment income ($m per quarter)', 85.0, 'm', 'Q2-26 $88m, Q1-26 $81m.'),
        ('pt_int', '  Q3/Q4-26E interest expense ($m per quarter)', 255.0, 'm', 'Q2-26 $218m; + a full quarter of the $5.0bn June-2026 notes (~4.5%, ~$55m) less the Nov-2026 $750m maturity.'),
        ('pt_oth', '  Q3/Q4-26E other income (expense), net ($m per quarter)', 45.0, 'm',
         'Set so total other income (expense) ≈ −$125m in Q3 (85 − 255 + 45), per the UNVERIFIED call guidance; H1-26 $106m incl. government grants.'),
        ('pt_etr', '  Q3/Q4-26E tax rate (GAAP & adjusted)', 0.205, 'pct', 'Non-GAAP tax rate guided 20–21% for Q3 and Q4 (UNVERIFIED); H1-26 19.7%.'),
        ('pt_capex_h2', '  H2-26E capital expenditures incl. capitalized software ($m)', 700.0, 'm', 'H1-26 $813m (property $445m, software $368m); FY2025 $1,215m.'),
    ],
    'fx_q3': 0.005, 'fx_q4': 0.0,
    'lev': (0.70, 0.50, 0.30),
    'lev_note': ('Net debt (total debt − unrestricted cash − investments) $13.0bn at 30-Jun-2026 ≈ 0.6x LTM adjusted EBITDA (model); FY2023–25 0.3–0.4x. Mastercard has no stated '
                 'leverage target; it returns excess capital through buybacks (FY2025 $11.7bn). Target = leverage up to which surplus cash is returned (assumption).'),
    'blocks': [
        ('GDV', 'GDV (gross dollar volume growth, local currency)', 0.015, -0.025, [0.085, 0.080, 0.080, 0.075],
         'H1-26 +7.7% local (FY2025 +8.6%, FY2024 +9.2%); purchase volume +9–10%. Secular cash-to-card and tokenized / e-commerce shift; Bear = consumer slowdown / U.S. debit share loss.', 'pct', False),
        ('YLD', 'YLD (payment network yield change y/y)', 0.010, -0.015, [0.015, 0.015, 0.010, 0.010],
         'Yield 15.5 bps (2021) → 18.3 bps (2025): cross-border recovery (higher-yield), pricing and switching, partly offset by rising rebates & incentives (+22% in H1-26). Bear = rebate pressure / regulation (interchange / CCCA).', 'pct', False),
        ('VAS', 'VAS (value-added services and solutions growth)', 0.030, -0.050, [0.150, 0.140, 0.130, 0.120],
         'FY2025 +23% (incl. Recorded Future from Dec-2024), H1-26 +21% (+18% currency-neutral). Security, digital & authentication, consumer engagement, data & insights, open banking.', 'pct', False),
        ('OM', 'OM (adjusted operating margin)', 0.010, -0.015, [0.603, 0.608, 0.613, 0.618],
         'Adjusted operating margin 57.0% (2022) → 59.2% (2025) → 61.0% (H1-26). Modest annual expansion; investment in VAS / talent offsets operating leverage.', 'pct', False),
        ('MNA', 'M&A_spend (acquisitions, $m)', 500.0, -500.0, [1000.0, 1000.0, 1000.0, 1000.0],
         'Recorded Future $2.65bn (2024), none in 2025, BVNK $1.5bn (2026); Finicity, Ekata, Nets A2A (2020–21). Base $1.0bn p.a. (assumption).', 'm', True),
    ],
    'am_pct': 0.027, 'da_pct': 0.034, 'sbc_pct': 0.018, 'etr': 0.205, 'oth_inc': 150.0,
    'lit_q': {3: 100.0, 4: 100.0}, 'lit_ae': 450.0,
    'lit_note': ('Litigation provisions are excluded from adjusted results (special items). FY2022–25: $356m / $539m / $680m / $504m; H1-26 $82m. Forecast: $100m a quarter in H2-26 '
                 'and $450m a year from 2027E (assumption: recurring interchange / merchant / ATM litigation), deducted in the DCF as a cash cost.'),
    'capex_pct': 0.040,
    'dps_q3': MAN['dividend']['q'], 'dps_q4': 0.99, 'dps_g': 0.13,
    'def_h2': 0.0, 'def_ae': 0.0, 'ofin_h2': -50.0, 'ofin_ae': -250.0, 'iss_h2': 0.5, 'iss_ae': 2.0, 'dil': 1.0,
    'ppe_share': 0.45, 'ppe_dshare': 0.34,
    'bv_cons': MAN['deals']['bvnk']['consideration'],
    'iss_notes': {2026: 0.0, 2027: 2900.0, 2028: 5500.0, 2029: 4000.0, 2030: 2700.0},
    'repay': {2026: DB['2026'], 2027: DB['2027'], 2028: DB['2028'], 2029: DB['2029'], 2030: DB['2030']},
    'mat': {2026: DB['2027'], 2027: DB['2028'], 2028: DB['2029'], 2029: DB['2030'], 2030: 2700.0},
    'cp_jun': MAN['debt']['cp'],
    'min_cash': 12000.0,
    'min_cash_note': ('Cash-flow basis (cash + restricted cash + restricted security deposits): $13.2bn at 31-Dec-2025, $14.0bn at 30-Jun-2026; unrestricted cash has been $8–11bn since 2021. '
                      '$12.0bn assumed (≈ $9.3bn unrestricted).'),
    'r_notes': {2027: 0.039, 2028: 0.040, 2029: 0.041, 2030: 0.042},
    'wc': WC,
    'px': MAN['price']['px'],
    'px_note': 'VALUATION — share price input $595.00 (user-provided, 13-Oct-2026; ~$520bn market cap); update to market',
    'mna_note': 'Scenario lever M&A_spend (Base $1.0bn p.a. of bolt-ons). Funded from cash / commercial paper (debt schedule).',
    'rev_guid_note': G26['net_revenue'] + ' ' + G26['status'],
    'ox_guid_note': G26['opex'] + ' ' + G26['status'],
}

# ====================================================================== compute engine
def hist_value(row, col):
    kw = row.kw; path = kw.get('data')
    if not path: return None
    if kw.get('ppa'):
        v = MAN['ppa'].get(pkey(col), {}).get(path.split('_', 1)[1])
        return v if kw.get('textdata') else num(v)
    v = xget(pkey(col), path)
    return v if kw.get('textdata') else num(v)

def first(kw, keys):
    for k in keys:
        if k in kw: return k
    return None

def compute(row, col):
    kw = row.kw; k = col.kind
    CTX.col = col
    if kw.get('ann') and col.q: return None, False
    if k in HISTK:
        v = hist_value(row, col)
        if v is not None: return v, True
        for key in ('hist', 'all'):
            if key in kw:
                return kw[key](), False
        return None, False
    if k == 'QE':
        if row.spec2: return None, False
        if 'qe_in' in kw:
            v = kw['qe_in']; return (v[col.q] if isinstance(v, dict) else v), True
        key = first(kw, ('qe', 'fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    if k == 'E26':
        if 'e26_in' in kw: return kw['e26_in'], True
        if 'e26' in kw: return kw['e26'](), False
        if kw.get('avg') and ('qe' in kw or 'qe_in' in kw): return f'=AVERAGE({Q1}{row.r}:{Q4E}{row.r})', False
        if kw.get('flow') and not row.spec2 and any(x in kw for x in ('qe', 'qe_in', 'fc')): return S26(row.key), False
        key = first(kw, ('fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    if k == 'AE':
        if 'ae_in' in kw:
            v = kw['ae_in']; return (v[col.year] if isinstance(v, dict) else v), True
        key = first(kw, ('ae', 'fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    return None, False

# ====================================================================== writers
CAGR_G = {'is_rev', 'g_rev', 'pn', 'vas', 'k_gdv', 'g_adj', 'o_aoi', 'r_adj', 'r_ebitda', 'is_oi', 'is_da', 'e_adjni', 'e_eps', 'cf_cfo', 'fcf', 'bs_ta',
          'bs_teq', 'sh_dil', 'is_ga', 'is_am', 'is_ni', 'eps_dil', 'dps', 'g_ox', 'o_adjox', 'lg_dom', 'lg_xb', 'lg_tp', 'lg_oth', 'lg_reb', 'k_cards'}
CAGR_A = {'g_om', 'o_aom', 'r_adj_m', 'is_oim', 'r_ebit_m', 'gm_rev', 'gm_pn', 'gm_vas', 'gm_gdv', 'gm_aom', 'gm_ebitda_m', 'gm_op_m', 'gm_net_m', 'gm_ga',
          'gm_am', 'gm_aox', 'fcf_m', 'fcf_conv', 'capex_pct', 'ra_roic', 'roe', 'lv_x', 'd_am', 'd_da_pct', 'd_sbc_pct', 'wc_nwc_p', 'is_etr',
          'e_rate', 'fcf_eb', 'vas_mix', 'pn_yld', 'k_gdv_gl', 'k_xb', 'k_sw', 'g_cn', 'fcf_ret', 'hr_lit'}
A = {y: ANN[y].c for y in ANN}

def style_cagr(ws, r, on):
    src = 16 if on else 12
    for c in CAGR:
        ws[f'{c}{r}']._style = copy.copy(TM[f'{L(CI(c) + TSHIFT)}{src}']._style)

def write_row(ws, row):
    r = row.r; arch = ARCH[row.arch]
    clone_row_style(ws, arch, r)
    ws.cell(r, 2).value = row.label or None
    if row.kw.get('note'): ws[f'{NOTE}{r}'] = row.kw['note']
    style_cagr(ws, r, row.key in CAGR_G or row.key in CAGR_A)
    if row.arch in ('section', 'block', 'blank', 'sub') and not row.kw.get('data'):
        return
    if row.key in DEFS:
        write_defs(ws, row); return
    nf = NF[row.nf]
    for col in COLS:
        v, inp = compute(row, col)
        if v is None: continue
        cell = ws[f'{col.c}{r}']
        cell.value = v
        if not row.kw.get('textdata'): cell.number_format = nf
        if inp:
            set_color(cell, BLUE, bold=True if col.kind in FCK else None)
            lk = row.kw.get('labkey')
            if lk and col.kind in HISTK and isinstance(v, (int, float)) and abs(v) > 1e-9:
                t = xget(pkey(col), lk)
                if t: comment(cell, f'Company items ({pkey(col)}): {t}', 380, 110)
        elif row.arch == 'growth' and col.kind in HISTK:
            set_color(cell, GRAY)
        else:
            set_color(cell, BLACK)
    rc = CELL_COMMENTS.get(row.key, {})
    for c, t in rc.items(): comment(ws[f'{c}{r}'], t)
    if row.key in CAGR_G:
        ws[f'{CAGR[0]}{r}'] = f'=IFERROR(({A[2030]}{r}/{A[2025]}{r})^(1/5)-1,"n/m")'
        ws[f'{CAGR[1]}{r}'] = f'=IFERROR(({A[2025]}{r}/{A[2020]}{r})^(1/5)-1,"n/m")'
        ws[f'{CAGR[2]}{r}'] = f'=IFERROR(({A[2025]}{r}/{A[2015]}{r})^(1/10)-1,"n/m")'
        ws[f'{CAGR[3]}{r}'] = f'=IFERROR(({A[2025]}{r}/{A[2006]}{r})^(1/19)-1,"n/m")'
        for c in CAGR: ws[f'{c}{r}'].number_format = '0%;\\(0%\\);"n/m"'
    elif row.key in CAGR_A:
        ws[f'{CAGR[0]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2026, 2031))}),"n/m")'
        ws[f'{CAGR[1]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2021, 2026))}),"n/m")'
        ws[f'{CAGR[2]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2016, 2026))}),"n/m")'
        ws[f'{CAGR[3]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2006, 2026))}),"n/m")'
        for c in CAGR: ws[f'{c}{r}'].number_format = NF['pct']
    t = LABEL_SOURCES.get(row.key)
    if t: comment(ws.cell(r, 2), t)

REL = 'Source: Mastercard earnings releases (Form 8-K Item 2.02, Ex. 99.1), SEC EDGAR CIK 1141391 — '
LABEL_SOURCES = {
    'k_gdv': REL + '"Mastercard Incorporated Operating Performance" tables (GDV, all Mastercard credit, charge and debit programs, worldwide; three months / full year).',
    'k_xb': REL + 'release text ("cross-border volume growth of x% on a local currency basis"; 2010–15 "cross-border volumes"). Transcribed per period (data/_agent).',
    'pn': ('Source: Forms 10-Q / 10-K, net revenue by category: payment network / value-added services and solutions from the FY2022 10-K (FY2020–21 comparatives); 2021–22 quarters from '
           'the 2023 10-Q prior-year columns (Q1-21 = H1 − Q2; Q4 = FY − 9M). From 2023 the 10-Qs; Q4 = FY (10-K) − nine months (Q3 10-Q).'),
    'lg_dom': ('Source: Forms 10-Q / 10-K revenue tables (domestic assessments, cross-border volume fees, transaction processing fees, other revenues, rebates and incentives), as originally '
               'presented FY2008–Q3-22; FY2007–08 from the FY2009 10-K recast (the FY2006–08 10-Ks presented operations fees / assessments); FY2006 not published on this basis.'),
    'pm_dom': 'Source: Forms 10-Q / 10-K, payment network key metrics (assessments) and the rebates & incentives amount in the MD&A text, from FY2022.',
    'is_rev': REL + 'consolidated statements of operations, as originally reported (USD thousands to 2009). 2014 re-presented in 2015 (memo row).',
    'd_lit': REL + 'non-GAAP reconciliation tables ("special items"); company wording per period in the cell comments.',
    'o_pub': REL + 'Non-GAAP reconciliations (adjusted operating expenses / operating margin), published in periods with special items and every quarter from 2017–18.',
    'e_pub': REL + 'Non-GAAP reconciliations (adjusted net income / adjusted diluted EPS); FY2006–09 "excluding special items" (annual).',
    'cf_ni': REL + 'condensed consolidated statements of cash flows (year to date; discrete quarters derived, Q4 = FY − 9M).',
    'bs_cash': REL + 'consolidated balance sheets (period ends), condensed by section; "other" lines derived so totals tie.',
    'sh_basic': REL + 'statements of operations (weighted-average shares). 10-for-1 split effective 21-Jan-2014: documents filed earlier restated (shares × 10, per share ÷ 10).',
    'dps': 'Source: Forms 10-Q / 10-K, "dividends declared per share" table or equity-statement caption (data/_dps.json); split-adjusted; Q4 = FY − 9M.',
    'ds_total': 'Source: release balance sheets; Q2-26 10-Q Note 9 (Debt) and Note 10 (stockholders\' equity); FY2025 10-K.',
    'bv_hdr': 'Source: Q2-26 10-Q Note 2 (Acquisitions): BVNK Holdings Limited — agreement March 2026, $1.5bn + up to $300m contingent consideration.',
    'g_hdr': 'FY26 outlook: Q2-26 earnings call (30-Jul-2026) — UNVERIFIED: not in any SEC filing; wording from third-party transcripts (Yahoo Finance, Motley Fool, Globe and Mail, Investing.com).',
}
CELL_COMMENTS = {}
def comments_setup():
    CELL_COMMENTS['is_rev'] = {A[2014]: 'FY2014 as originally reported ($9,473m); re-presented as $9,441m in the 2015 releases (reclassification) — see memo row.'}
    CELL_COMMENTS['is_ga'] = {A[2018]: 'FY2018 from the Q4-18 release: $89m reclassified from advertising & marketing to G&A for the year; the 2018 quarters are as originally reported (Σ Q ≠ FY).'}
    CELL_COMMENTS['eps_pub'] = {A[2008]: 'Pre-split $9.41 restated / originally reported; 2008–09 EPS uses the two-class method (participating securities), so net income ÷ diluted shares differs by ~$0.007.',
                                A[2009]: 'Two-class method (FSP EITF 03-6-1, adopted 2009).'}
    CELL_COMMENTS['cf_end'] = {QC[(2026, 1)].c: MAN['corrections']['Q1-26']['src']}
    CELL_COMMENTS['k_gdv'] = {A[2025]: 'FY2025 annual table $10,632bn vs Σ quarters $10,614bn: GDV is customer-reported and revised in later tables.'}

DEFS = {}
def defs_setup():
    q = lambda y, n: QC[(y, n)].c
    DEFS['o_def'] = {
        'C': ('FY2006–09: results "excluding special items" (litigation settlements; 2006 Foundation stock contribution and IPO-related interest; 2007 settlement-related other income; 2008 customer-agreement termination gain)',
              'Annual non-GAAP figures from the Q4 releases (2006–09). Redecard gains (2007–08) were not treated as special items.'),
        A[2010]: ('FY2010: no special items — no non-GAAP figures published (adjusted = GAAP, flag)', None),
        A[2011]: ('FY2011–12: U.S. merchant litigation provision excluded ($770m 2011, $20m 2012)', None),
        q(2013, 1): ('2013–16: non-GAAP published only in quarters with special items (merchant litigation, 2015 pension termination); other quarters adjusted = GAAP (flag)',
                     'Q1–Q3-13, 2014 (incl. the Q4-14 $87m restructuring charge, which the company did not exclude), Q1-15, Q4-15, Q1-16, Q3-16 and Q3-17 show no adjusting items.'),
        q(2017, 4): ('Q4-17: U.S. Tax Act and Venezuela deconsolidation charge excluded', None),
        q(2018, 1): ('2018+: litigation provisions (incl. the European Commission fine, Q4-18) and Tax Act items; non-GAAP published every quarter', None),
        q(2019, 2): ('Q2-19+: (gains) losses on equity investments excluded (Q1-19 as originally reported; later restated)', None),
        q(2021, 3): ('Q3-21: indirect tax matter excluded', None),
        q(2022, 1): ('2022: Russia-related impacts excluded (incl. a revenue adjustment — adjusted net revenue ≠ GAAP)', None),
        q(2024, 3): ('Q3-24 / Q1-26: restructuring charges excluded', None),
    }
    DEFS['e_def'] = {
        'C': ('FY2006–09: net income / EPS excluding special items, after tax (annual)', None),
        A[2010]: ('FY2010: no adjusted EPS published (adjusted = GAAP, flag)', None),
        A[2011]: ('FY2011–16: adjusted EPS excluding litigation provisions (and the 2015 pension termination) where published', None),
        q(2017, 4): ('Q4-17–2019: U.S. Tax Act impacts excluded (discrete tax)', None),
        q(2019, 2): ('Q2-19+: equity-investment (gains) losses and their tax effect excluded', None),
    }
def write_defs(ws, row):
    for c, (txt, cm) in DEFS[row.key].items():
        cell = ws[f'{c}{row.r}']; cell.value = txt
        cell.font = Font(name='Calibri', sz=10, b=True, i=True, color='FFC00000')
        if cm: comment(cell, cm)

def write_scen(ws):
    srcrow = {'hdr': 9, 'out_hdr': 11, 'pt_hdr': 15, 'out': 12, 'pt': 16, 'lev': 28, 'blk': 30, 'yr': 31}
    for spec_ in SCEN_ROWS:
        r, kind = spec_[0], spec_[1]
        for ci in range(CI(SL), CI(SNOTE) + 1):
            ws.cell(r, ci)._style = copy.copy(TM.cell(srcrow[kind], ci + TSHIFT)._style)
        if kind in ('hdr', 'out_hdr', 'pt_hdr'):
            _, _, a, b, c_, d_, e = spec_
            for col, v in zip([SL, SBULL, SBASE, SBEAR, SNOTE], [a, b, c_, d_, e]):
                if v is not None: ws[f'{col}{r}'] = v
        elif kind == 'out':
            _, _, lab, hi, mid, lo, note, nf = spec_
            ws[f'{SL}{r}'] = lab
            for col, v in zip([SBULL, SBASE, SBEAR], [hi, mid, lo]):
                ws[f'{col}{r}'] = v; ws[f'{col}{r}'].number_format = NF[nf]; set_color(ws[f'{col}{r}'], BLUE)
            ws[f'{SNOTE}{r}'] = note
        elif kind == 'pt':
            _, _, lab, _, val, _, note, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBASE}{r}'] = val; ws[f'{SBASE}{r}'].number_format = NF[nf]; set_color(ws[f'{SBASE}{r}'], BLUE); ws[f'{SNOTE}{r}'] = note
        elif kind == 'lev':
            _, _, lab, a, b, c_, note = spec_
            ws[f'{SL}{r}'] = lab
            for col, v in zip([SBULL, SBASE, SBEAR], [a, b, c_]):
                ws[f'{col}{r}'] = v; ws[f'{col}{r}'].number_format = NF['x']
            ws[f'{SNOTE}{r}'] = note
        elif kind == 'blk':
            _, _, lab, db, mid, de, note, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBULL}{r}'] = db; ws[f'{SBASE}{r}'] = mid; ws[f'{SBEAR}{r}'] = de; ws[f'{SNOTE}{r}'] = note
            for col in (SBULL, SBEAR):
                ws[f'{col}{r}'].number_format = '\\+#,##0;\\-#,##0' if nf == 'm' else '\\+0.0%;\\-0.0%'
        elif kind == 'yr':
            _, _, lab, hdr, base, clamp, _, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBASE}{r}'] = base
            if clamp:
                ws[f'{SBULL}{r}'] = f'=MAX(0,{SBASE}{r}+${SBULL}${hdr})'; ws[f'{SBEAR}{r}'] = f'=MAX(0,{SBASE}{r}+${SBEAR}${hdr})'
            else:
                ws[f'{SBULL}{r}'] = f'={SBASE}{r}+${SBULL}${hdr}'; ws[f'{SBEAR}{r}'] = f'={SBASE}{r}+${SBEAR}${hdr}'
            for col in (SBULL, SBASE, SBEAR): ws[f'{col}{r}'].number_format = NF[nf]

def write_header(ws):
    for r in range(1, 8):
        clone_row_style(ws, r, r, 2, CI(LASTC))
    ws[f'{NOTE}1'] = 'SCENARIO SWITCH ▼ (Bull / Base / Bear) — drives all scenario-linked forecast drivers'
    ws[f'{NOTE}2'] = 'Base'
    comment(ws[f'{NOTE}2'], 'Scenario toggle. Choose Bull / Base / Bear. Drives: FY26 net revenue and adjusted opex end-points (high / mid / low of the UNVERIFIED outlook), '
            'H2-26 buybacks, 2027E–2030E GDV growth, payment network yield, value-added services growth, adjusted operating margin, M&A spend and the leverage target for 2027E+ buybacks '
            f'(scenario table {SL}:{SNOTE}).')
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type='list', formula1='"Bull,Base,Bear"', allow_blank=False); ws.add_data_validation(dv); dv.add(f'{NOTE}2')
    ws['B2'] = 'Mastercard Incorporated (NYSE: MA) — 3-Statement Financial Model'
    ws['B3'] = ('US GAAP as reported · USD millions (per-share data in USD; shares in millions; GDV in USD billions) · Fiscal years ended 31 December · Source: Mastercard Forms 10-K / 10-Q, '
                'Form 8-K earnings releases (Ex. 99.1), SEC EDGAR (CIK 1141391); XBRL company facts for cross-checks · FY26 outlook from the Q2-26 call is UNVERIFIED (not in an SEC filing)')
    ws['B4'] = ('Basis notes: as originally reported in each release (2014 net revenue re-presented in 2015 kept as memo; 2018 cost reclassification; Q1-26 ending cash per the 10-Q). '
                'One reportable segment; net revenue = payment network + value-added services and solutions (presentation from the FY2022 10-K; legacy categories in memo). 10-for-1 split Jan-2014. '
                'Cash flow includes restricted cash and security deposits from 2018. Acquisitions: Vocalink 2017, Finicity / Nets A2A / Ekata 2020–21, Recorded Future Dec-2024; BVNK pending (assumed 30-Sep-2026).')
    ws[f'{CAGR[0]}5'] = 'Forecast'
    for c in CAGR[1:]: ws[f'{c}5'] = 'Historical'
    ws['B6'] = '(USDm)'
    for col in COLS: ws[f'{col.c}6'] = col.label
    ws[f'{NOTE}6'] = 'Modelling Notes'
    for c, a, b in zip(CAGR, ['5Y CAGR', '5Y CAGR', '10Y CAGR', '19Y CAGR'], ["25-'30", "20-'25", "15-'25", "06-'25"]):
        ws[f'{c}6'] = a; ws[f'{c}7'] = b

def set_columns(ws):
    for k in list(ws.column_dimensions.keys()): del ws.column_dimensions[k]
    ws.column_dimensions['A'].width = TM.column_dimensions['A'].width
    ws.column_dimensions['B'].width = TM.column_dimensions['B'].width
    wid = {'A': 13.0, 'Q': 9.5703125, 'QE': 11.42578125, 'E26': 14.140625, 'AE': 13.0}
    for col in COLS:
        cd = ws.column_dimensions[col.c]; cd.width = wid[col.kind]
        if col.q:
            cd.outlineLevel = 1; cd.hidden = True
    for ci in range(CI(NOTE), CI(LASTC) + 1):
        ws.column_dimensions[L(ci)].width = TM.column_dimensions[L(ci + TSHIFT)].width or 13.0
    ws.sheet_format.outlineLevelCol = 1

def main():
    import openpyxl
    wb = openpyxl.load_workbook(TEMPLATE)
    if 'DCF_old' in wb.sheetnames: del wb['DCF_old']
    ws = wb['Model']
    blank_style = copy.copy(ws['A1']._style)
    for rowc in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for c in rowc:
            c.value = None; c.comment = None; c._style = copy.copy(blank_style)
    for k in list(ws.row_dimensions.keys()):
        rd = ws.row_dimensions[k]; rd.outlineLevel = 0; rd.hidden = False
    ws.data_validations.dataValidation = []
    scen_layout(V0)
    build_spec(V0)
    n1 = len(fw.SPEC)
    build_spec2(V0)
    for i, row in enumerate(fw.SPEC): row.spec2 = i >= n1
    last = number_rows(8)
    for r in range(1, last + 1): ws.row_dimensions[r].height = 15.0
    ws.row_dimensions[2].height = 21.0
    defs_setup(); comments_setup()
    write_header(ws)
    for row in fw.SPEC: write_row(ws, row)
    write_scen(ws)
    set_columns(ws)
    ws.freeze_panes = 'C8'
    ws.sheet_view.zoomScale = 70
    import other_sheets as osh
    osh.do_all(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    osh.inject_charts(OUT)
    print('rows', fw.SPEC[-1].r, 'saved', OUT)
    return wb

RECALC = os.environ.get('XLSX_RECALC', '/mnt/skills/public/xlsx/scripts/recalc.py')
def recalc(path, timeout=600):
    import subprocess
    out = subprocess.run(['python3', RECALC, path, str(timeout)], capture_output=True, text=True, cwd=os.path.dirname(RECALC))
    try:
        return json.loads(out.stdout)
    except Exception:
        return {'raw': out.stdout[-2000:], 'err': out.stderr[-2000:]}

def scenario_snapshots():
    import openpyxl
    import other_sheets as osh
    snaps = {}
    for scen in ('Bull', 'Bear'):
        p = os.path.join(BASE, f'MA_{scen}.xlsx')
        wb = openpyxl.load_workbook(OUT); wb['Model'][f'{NOTE}2'] = scen; wb.save(p)
        res = recalc(p); print(scen, res.get('status'), res.get('total_errors'))
        v = openpyxl.load_workbook(p, data_only=True)['Bull-Base-Bear']
        snaps[scen] = {(r, c): v[f'{c}{r}'].value for r in range(9, 51) for c in 'CDEFGHIJKL' if v[f'{c}{r}'].value is not None}
        os.remove(p)
    wb = openpyxl.load_workbook(OUT)
    osh.snapshot_values(wb, snaps['Bull'], snaps['Bear'])
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    osh.inject_charts(OUT)

NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
def inject_values(src_xlsx, recalc_xlsx):
    """Store cached formula results (from a LibreOffice-recalculated copy) in the openpyxl file, so the workbook reads correctly in any
    viewer; fullCalcOnLoad stays set so Excel recalculates on open. Same approach as the LECO / ESAB finishing step."""
    import zipfile, shutil, openpyxl
    from lxml import etree
    wbv = openpyxl.load_workbook(recalc_xlsx, data_only=True)
    zin = zipfile.ZipFile(src_xlsx)
    wbxml = etree.fromstring(zin.read('xl/workbook.xml'))
    rels = etree.fromstring(zin.read('xl/_rels/workbook.xml.rels'))
    rid2t = {r.get('Id'): r.get('Target') for r in rels}
    sheets = {}
    for s in wbxml.iter('{%s}sheet' % NS):
        rid = s.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
        sheets['xl/' + rid2t[rid].lstrip('/').replace('xl/', '')] = s.get('name')
    tmp = src_xlsx + '.tmp'
    zout = zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED)
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename in sheets:
            ws = wbv[sheets[item.filename]]
            root = etree.fromstring(data)
            for c in root.iter('{%s}c' % NS):
                f = c.find('{%s}f' % NS)
                if f is None: continue
                v = ws[c.get('r')].value
                old = c.find('{%s}v' % NS)
                if old is not None: c.remove(old)
                if v is None: continue
                ve = etree.SubElement(c, '{%s}v' % NS)
                if isinstance(v, bool):
                    c.set('t', 'b'); ve.text = '1' if v else '0'
                elif isinstance(v, (int, float)):
                    if 't' in c.attrib: del c.attrib['t']
                    ve.text = repr(float(v)) if isinstance(v, float) else str(v)
                elif isinstance(v, str) and v.startswith('#'):
                    c.set('t', 'e'); ve.text = v
                else:
                    c.set('t', 'str'); ve.text = str(v)
            data = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
        zout.writestr(item, data)
    zout.close(); zin.close()
    shutil.move(tmp, src_xlsx)

def store_values():
    import shutil
    p = os.path.join(BASE, 'MA_recalc.xlsx')
    shutil.copy(OUT, p)
    res = recalc(p); print('recalc', res.get('status'), res.get('total_errors'))
    inject_values(OUT, p)
    os.remove(p)

if __name__ == '__main__':
    main()
    if '--snap' in sys.argv: scenario_snapshots()
    if '--recalc' in sys.argv: store_values()
