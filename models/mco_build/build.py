"""Build MCO_Model.xlsx from the MLM template (same layout, formatting and process as the prior models)."""
import sys, copy, os, json
from openpyxl.styles import Font
from fw import *
import fw
from spec import build_spec, scen_layout, SCEN_ROWS, SC
from spec2 import build_spec2

OUT = os.path.join(os.path.dirname(BASE), 'MCO_Model.xlsx')
G = MAN['guidance_fy26']['current']; GR = MAN['guidance_fy26']['reconciliation']

# ====================================================================== inputs (V0)
def wc_inputs():
    q = DATA['2025']['x']
    rev = q['rev']
    return {'dso': round(q['bs_ar'] / rev * 365), 'dr': round(q['bs_dr'] / rev, 3), 'ap': round(q['bs_ap'] / rev, 3)}
WC = wc_inputs()
WC['note'] = (f'Forecast inputs at fiscal-year-end (31-Dec-2025) levels on FY2025 revenue (DSO {WC["dso"]}d, deferred revenue {WC["dr"]:.1%} and payables & accrued '
              f'{WC["ap"]:.1%} of revenue); year ends are compared with year ends (Q1 receivables are seasonally high, deferred revenue peaks in Q1).')
AM = MAN['amort_sched']['schedule']
HSD, MSD = (0.09, 0.08, 0.07), (0.06, 0.05, 0.04)

V0 = {
    'out_hdr': 'FY26 guidance — Q2-26 release (22-Jul-2026)',
    'out_note': ('Revenue up high-single-digit %, adjusted operating margin 52–53%, adjusted diluted EPS $16.50–17.00 (raised from $16.40–17.00), diluted EPS $16.00–16.50, FCF $2.7–2.9bn, '
                 'share repurchases up to $3.0bn, ETR 23–25%, net interest $220–240m; MIS revenue up high-single-digit % at ~65% margin; MA revenue up mid-single-digit % at 34–35%. '
                 'Qualitative ranges are read as high-single-digit = 7–9% and mid-single-digit = 4–6% (model interpretation).'),
    'outlook': [
        ('out_mis_g', '  FY26 MIS revenue growth', HSD, 'pct', 'Guided "increase in the high-single-digit percent range" (read as 7–9%). Drives the Q3/Q4-26E MIS revenue calibration.'),
        ('out_ma_g', '  FY26 MA revenue growth', MSD, 'pct', 'Guided "increase in the mid-single-digit percent range" (read as 4–6%; incl. ~4 pts inorganic drag from the Learning and Regulatory Solutions divestitures and ~(0.5) pts FX). Drives the MA calibration.'),
        ('out_eps', '  FY26 adjusted diluted EPS ($)', (17.00, 16.75, 16.50), 'ps', 'Guided $16.50–17.00. Drives the Q3/Q4-26E segment-margin calibration.'),
        ('out_mco_g', '  FY26 MCO revenue growth — memo', HSD, 'pct', 'Guided "high-single-digit" (information row; the model total follows the segment end-points).'),
        ('out_aom', '  FY26 Adjusted Operating Margin — memo', (0.53, 0.525, 0.52), 'pct', 'Guided 52–53% (information row).'),
        ('out_misam', '  FY26 MIS Adjusted Operating Margin — memo', (0.65, 0.65, 0.65), 'pct', 'Guided ~65% (information row).'),
        ('out_maam', '  FY26 MA Adjusted Operating Margin — memo', (0.35, 0.345, 0.34), 'pct', 'Guided 34–35% (information row).'),
        ('out_geps', '  FY26 diluted EPS ($) — memo', (16.50, 16.25, 16.00), 'ps', 'Guided $16.00–16.50 (information row).'),
        ('out_fcf', '  FY26 Free Cash Flow ($m) — memo', (2900.0, 2800.0, 2700.0), 'm', 'Guided $2.7–2.9bn (operating cash flow $3.15–3.35bn less capex ~$450m) (information row).'),
        ('bb26', '  2H26E share repurchases ($m)', (835.0, 685.0, 535.0), 'm', '1H26 $2,165m (treasury shares); FY26 guidance "up to $3.0bn" → 2H up to $835m (high), $2.85bn FY (mid), $2.7bn FY (low). $1.8bn authorization remaining at 30-Jun-2026.'),
    ],
    'points': [
        ('pt_int', '  Q3/Q4-26E interest expense, net (per quarter, $m; expense −)', -53.0, 'm1', 'FY26 guidance $220–240m (mid $230m) − 1H26 $124m = $106m for 2H.'),
        ('pt_onop', '  Q3/Q4-26E other non-operating income, net (per quarter, $m)', 0.0, 'm1', 'Guided non-operating income $70–90m incl. interest and divestiture gains → other non-operating ≈ nil in 2H (1H26 +$16m).'),
        ('pt_gain', '  Q4/26E contingent gain on the Regulatory Solutions divestiture ($m)', 119.0, 'm1',
         'Up to $119m of contingent consideration payable on post-closing conditions in 2H26 (Q2-26 10-Q); the FY26 EPS bridge gain of ~$(1.25) less the Q2 $(0.72) implies ~$0.53 net ≈ $119m pre-tax. Timing assumed Q4.'),
        ('pt_amort', '  Q3/Q4-26E acquired-intangible amortization (per quarter, $m)', round((AM['2026'] - 106.0) / 2, 1), 'm1', 'FY2025 10-K schedule 2026 $207m − 1H26 $106m, split evenly.'),
        ('pt_dep', '  Q3/Q4-26E other depreciation & amortization (per quarter, $m)', 68.0, 'm1', 'D&A guided ~6% of revenue; 1H26 other D&A $142m ($71m/qtr) → $68m/qtr with amortization ≈ 6.1% of FY revenue.'),
        ('pt_restr', '  Q3/Q4-26E restructuring (per quarter, $m)', 30.0, 'm1', 'Guided ~1.5% of revenue (≈ $118m) − 1H26 $59m ≈ $59m for 2H (programme expanded Jul-2026 to $300–350m total).'),
        ('pt_othx', '  Q3/Q4-26E other items excluded from Adjusted Operating Income (per quarter, $m)', 11.5, 'm1',
         'Reserve for an international non-income tax obligation ~0.5% of revenue (1H26 $36m) and duplicate rent (NY HQ; EPS bridge ~$0.10 ≈ $22m pre-tax for FY; 1H26 $2m).'),
        ('pt_etr', '  Q3/Q4-26E tax rate (GAAP and on adjusting items)', 0.24, 'pct', 'FY26 guidance 23–25%; 1H26 24.5%.'),
        ('pt_nci', '  Q3/Q4-26E net income attributable to NCI (per quarter, $m)', 1.0, 'm1', 'ICRA minority; 1H26 $1m.'),
        ('pt_capex', '  FY26 capital additions ($m)', 450.0, 'm1', 'Guided ~$450m (incl. ~$100m of New York and London office relocations); 1H26 $186m.'),
    ],
    'lev': (1.50, 1.25, 1.00),
    'lev_note': ('Net debt / Adjusted Operating Income was ~1.2x at YE2025 (total debt $7.0bn less cash and short-term investments $2.4bn ÷ $4.0bn); FY2021–25 range 1.2–2.4x (2.4x after the RMS acquisition). '
                 'Target = leverage up to which surplus cash is returned via buybacks (assumption; Moody\'s gives no numeric net-leverage target). Bull = more buybacks.'),
    'blocks': [
        ('MIS_CFG', 'MIS_CFG (Corporate Finance y/y)', 0.03, -0.05, [0.07, 0.06, 0.06, 0.05], 'FY2015–25 CAGR 6.7%; 1H26 +19% on refinancing / M&A issuance; leveraged-finance refinancing walls 2027–28.', 'pct', False),
        ('MIS_SFG', 'MIS_SFG (Structured Finance y/y)', 0.03, -0.05, [0.05, 0.04, 0.04, 0.04], '1H26 +5%; private credit / CLO / ABS growth vs CMBS cyclicality.', 'pct', False),
        ('MIS_FIG', 'MIS_FIG (Financial Institutions y/y)', 0.02, -0.04, [0.06, 0.05, 0.05, 0.05], '1H26 +9%; bank / insurer issuance; private-credit fund ratings.', 'pct', False),
        ('MIS_PPIF', 'MIS_PPIF (Public, Project & Infrastructure Finance y/y)', 0.03, -0.05, [0.06, 0.06, 0.05, 0.05], '1H26 +23%; infrastructure / energy-transition project finance and US public finance.', 'pct', False),
        ('MIS_MOTH', 'MIS_MOTH (MIS Other y/y)', 0.02, -0.03, [0.06, 0.05, 0.05, 0.05], 'ICRA non-ratings and other; small (~$50m p.a.).', 'pct', False),
        ('MA_BANK', 'MA_BANK (Decision Solutions — Banking y/y)', 0.02, -0.04, [0.07, 0.07, 0.06, 0.06],
         '2026 depressed by the Learning (Q4-25) and Regulatory Solutions (Q2-26) divestitures; 2027E includes 1H anniversary; organic cc Banking +9% in 1H26; ARR +10%.', 'pct', False),
        ('MA_INS', 'MA_INS (Decision Solutions — Insurance y/y)', 0.02, -0.04, [0.09, 0.08, 0.08, 0.07], '1H26 +10% (CAPE Analytics, RMS cloud migration); ARR +9%.', 'pct', False),
        ('MA_KYC', 'MA_KYC (Decision Solutions — KYC y/y)', 0.03, -0.04, [0.12, 0.11, 0.10, 0.09], '1H26 +15% (cc +11%); ARR +13%; Orbis / Maxsight entity data.', 'pct', False),
        ('MA_RI', 'MA_RI (Research & Insights y/y)', 0.02, -0.03, [0.05, 0.05, 0.05, 0.05], '1H26 +5%; ARR +6%; mature research subscriptions.', 'pct', False),
        ('MA_DI', 'MA_DI (Data & Information y/y)', 0.02, -0.03, [0.08, 0.08, 0.07, 0.07], '1H26 +9%; ARR +8%; Orbis company data.', 'pct', False),
        ('MIS_M', 'MIS_M (MIS Adjusted Operating Margin, company basis)', 0.01, -0.03, [0.655, 0.66, 0.665, 0.67],
         'FY26 guided ~65% (1H26 67.5%, FY2025 63.6%, FY2021 62.2%, FY2022 51.8%); operating leverage on issuance.', 'pct', False),
        ('MA_M', 'MA_M (MA Adjusted Operating Margin, company basis)', 0.01, -0.02, [0.355, 0.365, 0.375, 0.385],
         'FY26 guided 34–35% (FY2025 33.1%, FY2021 26.0%); restructuring savings ($300–350m p.a. company-wide) and SaaS mix; Moody\'s medium-term MA ambition is not quantified on EDGAR in 2026 filings.', 'pct', False),
        ('MNA', 'M&A_spend (bolt-on acquisitions, $m)', 300.0, -200.0, [200.0, 200.0, 200.0, 200.0],
         'Bolt-ons 2022–26 (kompany, GCR, Praedicat, Numerated, CAPE Analytics $223m, ICR Chile, Fintellix, MERIS) ran ~$0–250m p.a.; large deals (BvD $3.5bn 2017, RMS $1.9bn 2021) are not modelled. Base $200m p.a. (assumption).', 'm', True),
    ],
    'm_mult': 7.0, 'm_mult_note': 'Analytics / data bolt-ons: RDC (2020) $702m, RMS (2021) $1,927m on ~$300m revenue (~6–7x); BvD (2017) $3.5bn on ~$280m (~12x). 7.0x input (assumption).',
    'm_intp': 0.40, 'm_intp_note': 'PPAs: BvD 38% of consideration in intangibles, RMS 40%, RDC 40%, Reis 37%, Cortera 41%, PassFort 15%.',
    'opr_pct': 0.235, 'sbc_pct': 0.030, 'etr': 0.240, 'nci': 4.0, 'onop': 10.0,
    'dep_pct': {2027: 0.036, 2028: 0.037, 2029: 0.038, 2030: 0.038},
    'capex_pct': {2027: 0.052, 2028: 0.050, 2029: 0.050, 2030: 0.050},
    'capex_note': '2026E: FY26 guidance ~$450m (incl. ~$100m office relocations). 2027E+: revenue × capex % (FY2022–25 4.2–5.2%).',
    'dps_q': 1.03, 'dps_g': 0.09,
    'dps_note': 'Dividends declared per share in the quarter as reported (from 2008 the 10-K shows each dividend in the quarter it was declared). Q3/Q4-26E $1.03 (declared 21-Jul-2026). 2027E+: prior year × (1 + growth); FY2021–26 +9–10% p.a.',
    'def_2h': -10.0, 'def_ae': -20.0, 'eq_2h': -20.0, 'eq_ae': -80.0, 'iss_2h': 0.25, 'iss_ae': 0.6, 'dil': 0.5,
    'ae_restr': {2027: 60.0, 2028: 0.0, 2029: 0.0, 2030: 0.0},
    'restr_note': (MAN['restructuring']['program'] + ': total $300–350m (expanded Jul-2026), $212m charged to 30-Jun-2026, substantially complete by end-2027 → ~$60m left for 2027E (mid-point $325m − ~$270m by YE2026E).'),
    'ae_othx': {2027: 20.0, 2028: 0.0, 2029: 0.0, 2030: 0.0},
    'gain_note': 'Q4/26E: contingent consideration on the Regulatory Solutions sale (point estimate); 2027E+: none (excluded from adjusted EPS by the company).',
    'amort_sched': {2027: AM['2027'], 2028: AM['2028'], 2029: AM['2029'], 2030: AM['2030']},
    'amort_note': MAN['amort_sched']['src'] + ': 2027 $195m, 2028 $182m, 2029 $150m, 2030 $126m (thereafter $1,006m); excludes 2026 acquisitions (MERIS, Fintellix; immaterial).',
    'iss_notes': {2026: 0.0, 2027: 1071.0, 2028: 1000.0, 2029: 900.0, 2030: 1357.0},
    'repay': {2026: 0.0, 2027: 571.0, 2028: 500.0, 2029: 400.0, 2030: 857.0},
    'mat': {2026: 571.0, 2027: 500.0, 2028: 400.0, 2029: 857.0, 2030: 0.0},
    'repay_note': ('Maturities (Q2-26 10-Q Note 13): Mar-2027 €500m 1.75% ($571m at 30-Jun-2026 FX); Jan-2028 $500m 3.25%; Feb-2029 $400m 4.25%; Feb-2030 €750m 0.95% ($857m); '
                   'Aug-2031 $600m 2.00%; then 2032–2061. EUR notes held at 30-Jun-2026 FX (assumption).'),
    'min_cash': 1500.0, 'min_cash_note': 'Moody\'s holds ~$1.5–2.5bn of cash and short-term investments (≈ half outside the US); $1.5bn floor (assumption).',
    'rcf': 1250.0,
    'r_notes': {2027: 0.034, 2028: 0.036, 2029: 0.037, 2030: 0.039},
    'r_notes_note': ('Weighted coupon ~3.3% on $7.1bn of notes (0.95–5.25%), raised by the receive-fixed / pay-floating swaps ($1.75bn notional) and net investment hedges; expense on borrowings 1H26 $110m '
                     '(~3.2% annualised). Refinancing at ~5% lifts the rate as low-coupon EUR notes mature (assumption).'),
    'ds_fees': 50.0, 'r_cash': 0.020, 'r_cash_note': '1H26 interest income $21m on ~$2bn average cash (~2% annualised).',
    'wc': WC,
    'px': 470.0,
    'px_note': ('VALUATION — share price input $470.00 (user-provided, 12-Oct-2026). Cross-check: latest SEC-filed price = June-2026 average repurchase price $450.64 '
                '(Q2-26 10-Q Part II Item 2). Update to market.'),
    'mna_note': 'Scenario lever M&A_spend (Base $200m p.a. of MA bolt-ons). Funded from cash / commercial paper (debt schedule).',
    'mis_guid_note': 'Company FY26 MIS revenue guidance: increase in the high-single-digit percent range (22-Jul-2026; unchanged since 18-Feb-2026).',
    'ma_guid_note': 'Company FY26 MA revenue guidance: increase in the mid-single-digit percent range (incl. ~4 pts divestiture drag) (22-Jul-2026).',
    'eps_guid_note': 'Company FY26 adjusted diluted EPS guidance $16.50–17.00 (raised 22-Jul-2026 from $16.40–17.00). Drives the segment-margin calibration.',
}

# ====================================================================== compute engine
def hist_value(row, col):
    kw = row.kw; path = kw.get('data')
    if not path: return None
    if kw.get('ppa'):
        d = MAN['ppa'].get(pkey(col))
        if not isinstance(d, dict) or d.get('consideration') is None: return None
        cash = d.get('cash_acquired') or 0.0
        f = path.split('_', 1)[1]
        v = {'cons': d['consideration'] - cash, 'ca': (d.get('current_assets') or 0.0) - cash, 'ppe': d.get('ppe'), 'intang': d.get('intangibles'),
             'gw': d.get('goodwill'), 'onca': (d.get('other_assets') or 0.0) + (d.get('rou') or 0.0), 'cl': d.get('liabilities_current'),
             'debt': d.get('debt_assumed'), 'ncl': (d.get('deferred_tax_liab') or 0.0) + (d.get('other_liab') or 0.0), 'names': d.get('name')}[f]
        return v if kw.get('textdata') else num(v)
    v = xget(pkey(col), path)
    return v if kw.get('textdata') else num(v)

def first(kw, keys):
    for k in keys:
        if k in kw: return k
    return None

def compute(row, col):
    kw = row.kw; k = col.kind
    CTX.col = col
    if kw.get('ann') and col.q: return None, False
    if k in HISTK:
        v = hist_value(row, col)
        if v is not None: return v, True
        for key in ('hist', 'all'):
            if key in kw:
                return kw[key](), False
        return None, False
    if k == 'QE':
        if row.spec2: return None, False
        if 'qe_in' in kw:
            v = kw['qe_in']; return (v[col.q] if isinstance(v, dict) else v), True
        key = first(kw, ('qe', 'fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    if k == 'E26':
        if 'e26_in' in kw: return kw['e26_in'], True
        if 'e26' in kw: return kw['e26'](), False
        if kw.get('avg') and ('qe' in kw or 'qe_in' in kw): return f'=AVERAGE({Q1}{row.r}:{QE}{row.r})', False
        if kw.get('flow') and not row.spec2 and any(x in kw for x in ('qe', 'qe_in', 'fc')): return S26(row.key), False
        key = first(kw, ('fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    if k == 'AE':
        if 'ae_in' in kw:
            v = kw['ae_in']; return (v[col.year] if isinstance(v, dict) else v), True
        key = first(kw, ('ae', 'fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    return None, False

# ====================================================================== writers
CAGR_G = {'is_rev', 'g_rev', 'mis_rev', 'ma_rev', 'mis_cfg', 'mis_sfg', 'mis_fig', 'mis_ppif', 'ma_bank', 'ma_ins', 'ma_kyc', 'ma_ri', 'ma_di', 'ma_arr',
          'o_adj', 'g_aoi', 'r_adj', 'r_ebitda', 'r_ebit', 'is_oi', 'd_dna', 'e_adjni', 'e_eps', 'cf_cfo', 'fcf', 'bs_ta', 'sh_dil', 'is_opr', 'is_sga',
          'd_amort', 'd_dep', 'mis_aoi', 'ma_aoi', 'is_ni', 'eps_dil', 'dps', 'mis_rec', 'ma_rec'}
CAGR_A = {'mis_m', 'ma_m', 'g_aoim', 'o_adj_m', 'r_adj_m', 'is_om', 'r_ebit_m', 'gm_rev', 'gm_mis', 'gm_ma', 'gm_ebitda_m', 'gm_op_m', 'gm_net_m', 'gm_opr',
          'gm_sga', 'gm_adjop_m', 'fcf_m', 'fcf_conv', 'capex_pct', 'ra_roic', 'roe', 'lv_x', 'd_opr', 'd_dep_pct', 'd_sbc_pct', 'wc_nwc_p', 'is_etr', 'e_rate',
          'fcf_eb', 'mis_mix', 'ma_mix', 'mis_rec_p', 'ma_rec_p', 'wc_dso', 'wc_drp'}
A = {y: ANN[y].c for y in ANN}

def style_cagr(ws, r, on):
    src = 16 if on else 12
    for c in CAGR:
        ws[f'{c}{r}']._style = copy.copy(TM[f'{L(CI(c) + TSHIFT)}{src}']._style)

def write_row(ws, row):
    r = row.r; arch = ARCH[row.arch]
    clone_row_style(ws, arch, r)
    ws.cell(r, 2).value = row.label or None
    if row.kw.get('note'): ws[f'{NOTE}{r}'] = row.kw['note']
    style_cagr(ws, r, row.key in CAGR_G or row.key in CAGR_A)
    if row.arch in ('section', 'block', 'blank', 'sub') and not row.kw.get('data'):
        return
    if row.key in DEFS:
        write_defs(ws, row); return
    nf = NF[row.nf]
    for col in COLS:
        v, inp = compute(row, col)
        if v is None: continue
        cell = ws[f'{col.c}{r}']
        cell.value = v
        if not row.kw.get('textdata'): cell.number_format = nf
        if inp:
            set_color(cell, BLUE, bold=True if col.kind in FCK else None)
            lk = row.kw.get('labkey')
            if lk and col.kind in HISTK and isinstance(v, (int, float)) and abs(v) > 1e-9:
                t = xget(pkey(col), lk)
                if t: comment(cell, f'Company items ({pkey(col)}): {t}', 380, 110)
        elif row.arch == 'growth' and col.kind in HISTK:
            set_color(cell, GRAY)
        else:
            set_color(cell, BLACK)
    rc = CELL_COMMENTS.get(row.key, {})
    for c, t in rc.items(): comment(ws[f'{c}{r}'], t)
    if row.key in CAGR_G:
        ws[f'{CAGR[0]}{r}'] = f'=IFERROR(({A[2030]}{r}/{A[2025]}{r})^(1/5)-1,"n/m")'
        ws[f'{CAGR[1]}{r}'] = f'=IFERROR(({A[2025]}{r}/{A[2020]}{r})^(1/5)-1,"n/m")'
        ws[f'{CAGR[2]}{r}'] = f'=IFERROR(({A[2025]}{r}/{A[2015]}{r})^(1/10)-1,"n/m")'
        ws[f'{CAGR[3]}{r}'] = f'=IFERROR(({A[2025]}{r}/{A[2006]}{r})^(1/19)-1,"n/m")'
        for c in CAGR: ws[f'{c}{r}'].number_format = '0%;\\(0%\\);"n/m"'
    elif row.key in CAGR_A:
        ws[f'{CAGR[0]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2026, 2031))}),"n/m")'
        ws[f'{CAGR[1]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2021, 2026))}),"n/m")'
        ws[f'{CAGR[2]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2016, 2026))}),"n/m")'
        ws[f'{CAGR[3]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2006, 2026))}),"n/m")'
        for c in CAGR: ws[f'{c}{r}'].number_format = NF['pct']
    t = LABEL_SOURCES.get(row.key)
    if t: comment(ws.cell(r, 2), t)

REL = "Source: Moody's earnings releases (Form 8-K Item 2.02, Ex. 99.1) — "
LABEL_SOURCES = {
    'mis_cfg': (REL + 'revenue by line of business table (external revenue, as originally reported in each release; 2006–07 on the 2008 MIS / MA basis from the FY2007 / FY2008 '
                'comparatives — the old segments were "Moody\'s Investors Service" incl. research and "Moody\'s KMV"). REIT revenue moved from SFG to CFG in the 2019 releases (2018 recast, not applied).'),
    'ma_bank': (REL + 'MA revenue by line. Decision Solutions / Research & Insights / Data & Information from 2022 (2021 from the 2022 restated comparatives); Banking / Insurance / KYC '
                'from the 2023 10-Q revenue notes and 2024+ releases (earlier quarters of the current basis from the following year\'s restated comparatives, used only where the restated '
                'lines sum to the originally reported MA revenue). Legacy structures are memo rows below.'),
    'ma_l_rda': ('Legacy MA lines as originally reported: Subscriptions / Software / Consulting (2006–08; 2006–07 from the 2008 recast); Research, Data & Analytics / Risk Management Software / '
                 'Professional Services (2009–11); RD&A / Enterprise Risk Solutions / Professional Services (2012–20; PS dropped from 2021 after the MAKS sale); Decision Solutions total (2022–23).'),
    'mis_aoi': REL + 'segment table (Adjusted Operating Income by segment; from Q3-12; 10-K segment notes for FY2011–12). Segment allocations were refined each 1 January; values as first published.',
    'is_rev': REL + 'consolidated statements of operations; annual columns per the 10-K (the 10-K wins where it differs from the Q4 release).',
    'd_amort': 'Source: XBRL company facts (AmortizationOfIntangibleAssets, 10-K / 10-Q as first filed; Q4 = FY − 9M).',
    'o_pub': REL + '"Adjusted Operating Income and Adjusted Operating Margin" table (from Q3-12; FY2012+).',
    'e_pub': REL + 'non-GAAP / adjusted diluted EPS reconciliation (published when the company had adjusting items; definition in force per period — see the definition row).',
    'cf_ni': 'Source: 10-Q / 10-K cash-flow statements (year to date; discrete quarters derived, Q4 = FY − 9M).',
    'bs_cash': 'Source: 10-Q / 10-K balance sheets (period ends); shares outstanding from the release ("actual number of shares outstanding") or the 10-Q / 10-K.',
    'ppa_cons': ("Source: 10-K acquisition notes (final allocations): Bureau van Dijk (Aug-2017), Reis (Oct-2018), RDC (Feb-2020), Catylist / ZM / Acquire Media (Q4-20), Cortera (Mar-2021), "
                 "RMS (Sep-2021), PassFort (Nov-2021). Consideration and current assets shown net of cash acquired. Other deals (CAPE Analytics, GCR, Praedicat, Numerated, ICR Chile, "
                 "Fintellix, MERIS, kompany) were described as not material; no allocation published."),
    'sh_basic': REL + 'statements of operations (weighted average shares). No stock splits since 2005.',
    'ds_total': 'Source: 10-Q / 10-K balance sheets; Q2-26 10-Q Note 13 (indebtedness; 14 senior notes, $7,128m principal) and FY2025 10-K Note 16 (CP programme, 2024 credit facility).',
    'elim': REL + 'segment table (intersegment royalty charged by MIS to MA; MA intersegment revenue from 2011).',
}
CELL_COMMENTS = {}
def comments_setup():
    CELL_COMMENTS['mis_cfg'] = {A[2006]: 'FY2006 on the 2008 MIS / MA basis (FY2007 10-K / release comparatives); see the source comment on the label.',
                                A[2007]: 'FY2007 on the 2008 MIS / MA basis (FY2008 release comparatives).'}
    CELL_COMMENTS['mis_rev'] = {QC[(2014, 1)].c: 'Q1–Q3/14 as originally reported (before the Q4-14 realignment moved KIS non-ratings revenue from MA RD&A to the new "MIS Other" line); '
                                                  'FY2014 is on the realigned basis, so Σ quarters ≠ FY (MIS −$10.1m, MA +$10.1m).'}
    CELL_COMMENTS['is_restr'] = {ANN[2016].c: 'FY2016 per the 10-K ($12.0m); the Q4-16 release kept $3.6m inside SG&A, so Σ quarters (8.4) ≠ FY.'}
    CELL_COMMENTS['is_oi'] = {QC[(2016, 4)].c: 'Q4-16 operating loss: $863.8m settlement charge with the US DOJ and 21 states (excluded from Adjusted Operating Income and adjusted EPS).'}
    CELL_COMMENTS['eps_pub'] = {QC[(2016, 4)].c: 'Q4-16 diluted EPS (2.22) as printed in the Q4-16 release; the 10-K and later filings show (2.25) (dilutive shares excluded in a loss quarter).'}
    CELL_COMMENTS['cf_fx'] = {QC[(2019, 2)].c: 'Includes −$8.0m "reclassification of cash to assets held for sale" (MAKS).', QC[(2019, 3)].c: 'Includes −$11.1m reclassification of cash to assets held for sale (MAKS).'}
    for p, x in DATA.items():
        xx = x.get('x', {})
        if xx.get('adj_oi_old') is not None and p.isdigit() and int(p) in ANN:
            CELL_COMMENTS.setdefault('o_adj', {})[ANN[int(p)].c] = (f'Company measure in force: non-GAAP operating income {xx["adj_oi_old"]:,.1f} (items: {xx.get("adj_oi_old_lab")}) — before D&A, so not '
                                                                     'comparable; the model applies the current definition (operating income + D&A + restructuring + IS-face other operating items).')
        if xx.get('adj_ni_derived') and pkey_col(p):
            CELL_COMMENTS.setdefault('e_pubni', {})[pkey_col(p)] = 'Adjusted net income not published; derived as published adjusted EPS × diluted shares.'
        if xx.get('seg_recast') and pkey_col(p):
            pass
def pkey_col(p):
    if p.isdigit(): return ANN[int(p)].c if int(p) in ANN else None
    y, q = int('20' + p[-2:]), int(p[1])
    return QC[(y, q)].c if (y, q) in QC else None

DEFS = {}
def defs_setup():
    q = lambda y, n: QC[(y, n)].c
    DEFS['o_def'] = {
        'C': ('2006–11: no "Adjusted Operating Income" (before D&A) published — model applies the current definition (OI + D&A + restructuring + IS-face other items); 2006–09 company '
              '"non-GAAP operating income" (excl. building-sale gain / SBC / restructuring) in cell comments', None),
        A[2012]: ('FY2012 (from Q3-12): Adjusted Operating Income = OI excl. D&A, restructuring and goodwill impairment', 'First published in the Q3-12 release; FY2010–11 figures appeared only later as comparatives (839.2 / 967.6).'),
        q(2013, 1): ('2013–Q2/16: Adjusted Operating Income = OI excl. D&A (and the Q4-13 goodwill impairment)', None),
        q(2016, 3): ('Q3-16+: also excludes restructuring; Q4-16 the $863.8m DOJ / states settlement charge', None),
        q(2017, 2): ('Q2-17+: acquisition-related expenses (Bureau van Dijk) excluded', None),
        q(2019, 2): ('2019: MAKS impairment / divestiture loss and a captive-insurance settlement excluded', None),
        q(2024, 2): ('Q2-24+: charges related to asset abandonment excluded', None),
        q(2026, 1): ('2026: reserve for an international non-income tax obligation and duplicate rent (NY HQ) excluded', None),
    }
    DEFS['e_def'] = {
        'C': ('2006–15: non-GAAP EPS published when items existed (building-sale gain 2006; restructuring and legacy tax 2007–12; ICRA gain 2014; litigation Q4-13) — '
              'adjusted = GAAP where none published (flag)', None),
        q(2016, 3): ('Q3-16+: adjusted EPS excludes restructuring; Q4-16 the settlement charge and an FX gain', None),
        q(2017, 1): ('2017: CCXI gain, purchase-price hedge gain and acquisition-related expenses excluded; acquired-intangible amortization excluded from Q3-17 (2016 recast later, not applied)', None),
        q(2017, 4): ('Q4-17 / 2018: US Tax Act items excluded', None),
        q(2022, 2): ('2022: Russia exit (FX translation) and the Q4-22 debt-buyback gain excluded', None),
        q(2025, 4): ('Q4-25 / 2026: gains on business divestitures (Learning / Regulatory Solutions) and the international non-income tax reserve excluded', None),
    }
def write_defs(ws, row):
    for c, (txt, cm) in DEFS[row.key].items():
        cell = ws[f'{c}{row.r}']; cell.value = txt
        cell.font = Font(name='Calibri', sz=10, b=True, i=True, color='FFC00000')
        if cm: comment(cell, cm)

def write_scen(ws):
    srcrow = {'hdr': 9, 'out_hdr': 11, 'pt_hdr': 15, 'out': 12, 'pt': 16, 'lev': 28, 'blk': 30, 'yr': 31}
    for spec_ in SCEN_ROWS:
        r, kind = spec_[0], spec_[1]
        for ci in range(CI(SL), CI(SNOTE) + 1):
            ws.cell(r, ci)._style = copy.copy(TM.cell(srcrow[kind], ci + TSHIFT)._style)
        if kind in ('hdr', 'out_hdr', 'pt_hdr'):
            _, _, a, b, c_, d_, e = spec_
            for col, v in zip([SL, SBULL, SBASE, SBEAR, SNOTE], [a, b, c_, d_, e]):
                if v is not None: ws[f'{col}{r}'] = v
        elif kind == 'out':
            _, _, lab, hi, mid, lo, note, nf = spec_
            ws[f'{SL}{r}'] = lab
            for col, v in zip([SBULL, SBASE, SBEAR], [hi, mid, lo]):
                ws[f'{col}{r}'] = v; ws[f'{col}{r}'].number_format = NF[nf]; set_color(ws[f'{col}{r}'], BLUE)
            ws[f'{SNOTE}{r}'] = note
        elif kind == 'pt':
            _, _, lab, _, val, _, note, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBASE}{r}'] = val; ws[f'{SBASE}{r}'].number_format = NF[nf]; set_color(ws[f'{SBASE}{r}'], BLUE); ws[f'{SNOTE}{r}'] = note
        elif kind == 'lev':
            _, _, lab, a, b, c_, note = spec_
            ws[f'{SL}{r}'] = lab
            for col, v in zip([SBULL, SBASE, SBEAR], [a, b, c_]):
                ws[f'{col}{r}'] = v; ws[f'{col}{r}'].number_format = NF['x']
            ws[f'{SNOTE}{r}'] = note
        elif kind == 'blk':
            _, _, lab, db, mid, de, note, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBULL}{r}'] = db; ws[f'{SBASE}{r}'] = mid; ws[f'{SBEAR}{r}'] = de; ws[f'{SNOTE}{r}'] = note
            for col in (SBULL, SBEAR):
                ws[f'{col}{r}'].number_format = '\\+#,##0;\\-#,##0' if nf == 'm' else '\\+0.0%;\\-0.0%'
        elif kind == 'yr':
            _, _, lab, hdr, base, clamp, _, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBASE}{r}'] = base
            if clamp:
                ws[f'{SBULL}{r}'] = f'=MAX(0,{SBASE}{r}+${SBULL}${hdr})'; ws[f'{SBEAR}{r}'] = f'=MAX(0,{SBASE}{r}+${SBEAR}${hdr})'
            else:
                ws[f'{SBULL}{r}'] = f'={SBASE}{r}+${SBULL}${hdr}'; ws[f'{SBEAR}{r}'] = f'={SBASE}{r}+${SBEAR}${hdr}'
            for col in (SBULL, SBASE, SBEAR): ws[f'{col}{r}'].number_format = NF[nf]

def write_header(ws):
    for r in range(1, 8):
        clone_row_style(ws, r, r, 2, CI(LASTC))
    ws[f'{NOTE}1'] = 'SCENARIO SWITCH ▼ (Bull / Base / Bear) — drives all scenario-linked forecast drivers'
    ws[f'{NOTE}2'] = 'Base'
    comment(ws[f'{NOTE}2'], 'Scenario toggle. Choose Bull / Base / Bear. Drives: FY26 MIS / MA revenue-growth and adjusted-EPS guidance end-points (high / mid / low), '
            '2H26 buybacks, 2027E–2030E revenue growth by line, segment margins, M&A spend and the leverage target for 2027E+ buybacks '
            f'(scenario table {SL}:{SNOTE}).')
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type='list', formula1='"Bull,Base,Bear"', allow_blank=False); ws.add_data_validation(dv); dv.add(f'{NOTE}2')
    ws['B2'] = "Moody's Corporation (NYSE: MCO) — 3-Statement Financial Model"
    ws['B3'] = ("US GAAP as reported · USD millions (per-share data in USD; shares in millions) · Calendar fiscal year · Source: Moody's Forms 10-K / 10-Q, Form 8-K earnings releases (Ex. 99.1), "
                "SEC EDGAR (CIK 1059556); XBRL company facts; investor-relations releases (ir.moodys.com) are the same documents as filed on EDGAR")
    ws['B4'] = ("Basis notes: as originally reported (10-K over the Q4 release for annual columns). Segments: MIS (ratings) and MA (analytics) since the Jan-2008 reorganisation (2006–07 revenue lines "
                "on the 2008 basis). MA lines on the current structure from 2021 (restated comparatives). Acquisitions: BvD Aug-17, Reis Oct-18, RDC Feb-20, RMS Sep-21; divestitures: MAKS Nov-19, "
                "Learning Solutions Q4-25, Regulatory Solutions Q2-26. Q3/Q4-26E calibrated to the FY26 guidance (22-Jul-2026). Share price $470.00 (user, 12-Oct-2026).")
    ws[f'{CAGR[0]}5'] = 'Forecast'
    for c in CAGR[1:]: ws[f'{c}5'] = 'Historical'
    ws['B6'] = '(USDm)'
    for col in COLS: ws[f'{col.c}6'] = col.label
    ws[f'{NOTE}6'] = 'Modelling Notes'
    for c, a, b in zip(CAGR, ['5Y CAGR', '5Y CAGR', '10Y CAGR', '19Y CAGR'], ["25-'30", "20-'25", "15-'25", "06-'25"]):
        ws[f'{c}6'] = a; ws[f'{c}7'] = b

def set_columns(ws):
    for k in list(ws.column_dimensions.keys()): del ws.column_dimensions[k]
    ws.column_dimensions['A'].width = TM.column_dimensions['A'].width
    ws.column_dimensions['B'].width = TM.column_dimensions['B'].width
    wid = {'A': 13.0, 'Q': 9.5703125, 'QE': 11.42578125, 'E26': 14.140625, 'AE': 13.0}
    for col in COLS:
        cd = ws.column_dimensions[col.c]; cd.width = wid[col.kind]
        if col.q:
            cd.outlineLevel = 1; cd.hidden = True
    for ci in range(CI(NOTE), CI(LASTC) + 1):
        ws.column_dimensions[L(ci)].width = TM.column_dimensions[L(ci + TSHIFT)].width or 13.0
    ws.sheet_format.outlineLevelCol = 1

def main():
    import openpyxl
    wb = openpyxl.load_workbook(TEMPLATE)
    if 'DCF_old' in wb.sheetnames: del wb['DCF_old']
    ws = wb['Model']
    blank_style = copy.copy(ws['A1']._style)
    for rowc in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for c in rowc:
            c.value = None; c.comment = None; c._style = copy.copy(blank_style)
    for k in list(ws.row_dimensions.keys()):
        rd = ws.row_dimensions[k]; rd.outlineLevel = 0; rd.hidden = False
    ws.data_validations.dataValidation = []
    scen_layout(V0)
    build_spec(V0)
    n1 = len(fw.SPEC)
    build_spec2(V0)
    for i, row in enumerate(fw.SPEC): row.spec2 = i >= n1
    last = number_rows(8)
    for r in range(1, last + 1): ws.row_dimensions[r].height = 15.0
    ws.row_dimensions[2].height = 21.0
    defs_setup(); comments_setup()
    write_header(ws)
    for row in fw.SPEC: write_row(ws, row)
    write_scen(ws)
    set_columns(ws)
    ws.freeze_panes = 'C8'
    ws.sheet_view.zoomScale = 70
    import other_sheets as osh
    osh.do_all(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    osh.inject_charts(OUT)
    print('rows', fw.SPEC[-1].r, 'saved', OUT)
    return wb

RECALC = os.environ.get('XLSX_RECALC', '/mnt/skills/public/xlsx/scripts/recalc.py')
def recalc(path, timeout=600):
    import subprocess
    out = subprocess.run(['python3', RECALC, path, str(timeout)], capture_output=True, text=True, cwd=os.path.dirname(RECALC))
    try:
        return json.loads(out.stdout)
    except Exception:
        return {'raw': out.stdout[-2000:], 'err': out.stderr[-2000:]}

def scenario_snapshots():
    import openpyxl
    import other_sheets as osh
    snaps = {}
    for scen in ('Bull', 'Bear'):
        p = os.path.join(BASE, f'MCO_{scen}.xlsx')
        wb = openpyxl.load_workbook(OUT); wb['Model'][f'{NOTE}2'] = scen; wb.save(p)
        res = recalc(p); print(scen, res.get('status'), res.get('total_errors'))
        v = openpyxl.load_workbook(p, data_only=True)['Bull-Base-Bear']
        snaps[scen] = {(r, c): v[f'{c}{r}'].value for r in range(9, 51) for c in 'CDEFGHIJKL' if v[f'{c}{r}'].value is not None}
        os.remove(p)
    wb = openpyxl.load_workbook(OUT)
    osh.snapshot_values(wb, snaps['Bull'], snaps['Bear'])
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    osh.inject_charts(OUT)

NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
def inject_values(src_xlsx, recalc_xlsx):
    """Store cached formula results (from a LibreOffice-recalculated copy) in the openpyxl file, so the workbook reads correctly in any
    viewer; fullCalcOnLoad stays set so Excel recalculates on open. Same approach as the LECO / ESAB finishing step."""
    import zipfile, shutil, openpyxl
    from lxml import etree
    wbv = openpyxl.load_workbook(recalc_xlsx, data_only=True)
    zin = zipfile.ZipFile(src_xlsx)
    wbxml = etree.fromstring(zin.read('xl/workbook.xml'))
    rels = etree.fromstring(zin.read('xl/_rels/workbook.xml.rels'))
    rid2t = {r.get('Id'): r.get('Target') for r in rels}
    sheets = {}
    for s in wbxml.iter('{%s}sheet' % NS):
        rid = s.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
        sheets['xl/' + rid2t[rid].lstrip('/').replace('xl/', '')] = s.get('name')
    tmp = src_xlsx + '.tmp'
    zout = zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED)
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename in sheets:
            ws = wbv[sheets[item.filename]]
            root = etree.fromstring(data)
            for c in root.iter('{%s}c' % NS):
                f = c.find('{%s}f' % NS)
                if f is None: continue
                v = ws[c.get('r')].value
                old = c.find('{%s}v' % NS)
                if old is not None: c.remove(old)
                if v is None: continue
                ve = etree.SubElement(c, '{%s}v' % NS)
                if isinstance(v, bool):
                    c.set('t', 'b'); ve.text = '1' if v else '0'
                elif isinstance(v, (int, float)):
                    if 't' in c.attrib: del c.attrib['t']
                    ve.text = repr(float(v)) if isinstance(v, float) else str(v)
                elif isinstance(v, str) and v.startswith('#'):
                    c.set('t', 'e'); ve.text = v
                else:
                    c.set('t', 'str'); ve.text = str(v)
            data = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
        zout.writestr(item, data)
    zout.close(); zin.close()
    shutil.move(tmp, src_xlsx)

def store_values():
    import shutil
    p = os.path.join(BASE, 'MCO_recalc.xlsx')
    shutil.copy(OUT, p)
    res = recalc(p); print('recalc', res.get('status'), res.get('total_errors'))
    inject_values(OUT, p)
    os.remove(p)

if __name__ == '__main__':
    main()
    if '--snap' in sys.argv: scenario_snapshots()
    if '--recalc' in sys.argv: store_values()
