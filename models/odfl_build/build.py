"""Build ODFL_Model.xlsx from the MLM template (same layout, formatting and process)."""
import sys, copy, os, json
from openpyxl.styles import Font
from fw import *
import fw
import normalize
import workdays
from spec import build_spec, scen_layout, SCEN_ROWS, SC, seas
from spec2 import build_spec2

OUT = os.path.join(os.path.dirname(BASE), 'ODFL_Model.xlsx')
normalize.run()
D = DATA

# ====================================================================== inputs (V0)
def wc_inputs():
    b, i = D['2025']['bs'], D['2025']['is']
    opx = i['total_opex'] - i['dda']
    return {'dso': round(b['receivables'] / i['revenue'] * 365), 'dpo': round(b['ap'] / opx * 365),
            'comp': round(b['accrued_comp'] / i['revenue'], 3), 'claims': round(b['claims_cur'] / i['revenue'], 3)}
WC = wc_inputs()
WC['note'] = (f'Forecast inputs set at FY2025 levels (DSO {WC["dso"]}d, payables {WC["dpo"]}d of operating expenses excl. D&A, '
              f'compensation & benefits {WC["comp"]:.1%} and current claims accruals {WC["claims"]:.1%} of revenue).')

V0 = {
    'out_hdr': 'FY26 calibration — Q3/26E anchored to the 3-Sep-26 mid-quarter update; seasonality 2016–25 (formulas on the historical columns)',
    'out_note': 'ODFL gives no revenue / EPS guidance. Bull / Base / Bear = best / average / worst historical sequential change, 2016–2025 (Q2→Q3 and Q3→Q4).',
    'outlook': [
        ('seq_or3', '  Q2→Q3 operating-ratio change (best / average / worst, 2016–25)',
         (seas('MIN', 'or_seq', 3), seas('AVERAGE', 'or_seq', 3), seas('MAX', 'or_seq', 3)), 'pct',
         'Applied to the Q2/26 operating ratio (70.1%). Best case = 2020 (−332 bps, COVID rebound); excluding 2020 the best is −166 bps (2023).'),
        ('seq_or4', '  Q3→Q4 operating-ratio change (best / average / worst, 2016–25)',
         (seas('MIN', 'or_seq', 4), seas('AVERAGE', 'or_seq', 4), seas('MAX', 'or_seq', 4)), 'pct',
         'Applied to the Q3/26E operating ratio. Q4 is seasonally weaker (fewer work days, winter weather, holiday freight mix).'),
        ('seq_tpd', '  Q3→Q4 LTL tons per day, sequential % (best / average / worst)',
         (seas('MAX', 'l_tpd_sq', 4), seas('AVERAGE', 'l_tpd_sq', 4), seas('MIN', 'l_tpd_sq', 4)), 'pct', 'Applied to Q3/26E tons per day.'),
        ('seq_xf', '  Q3→Q4 LTL revenue/cwt ex fuel, sequential % (best / average / worst)',
         (seas('MAX', 'y_xf_sq', 4), seas('AVERAGE', 'y_xf_sq', 4), seas('MIN', 'y_xf_sq', 4)), 'pct', 'Applied to Q3/26E ex-fuel yield.'),
        ('bb26', '  2H/26 share repurchases ($m)', (400, 300, 150), 'm',
         '1H/26 $239.7m; $1.31bn remaining under the $3.0bn 2023 programme at 30-Jun-26. FY25 $730m, FY24 $967m.'),
    ],
    'points': [
        ('q3_tpd_g', '  Q3/26E LTL tons per day y/y (Aug-26, 3-Sep-26 update)', -0.009, 'pct',
         'August 2026 LTL tons per day −0.9% y/y (shipments/day −2.4%, weight/shipment +1.7%). July not separately disclosed; Q2/26 −4.1%.'),
        ('q3_xf_g', '  Q3/26E LTL revenue/cwt ex fuel y/y (Jul–Aug QTD)', 0.048, 'pct', '3-Sep-26 update: quarter-to-date +4.8%.'),
        ('q3_cwt_g', '  Q3/26E LTL revenue/cwt incl. fuel y/y (Jul–Aug QTD)', 0.113, 'pct', '3-Sep-26 update: quarter-to-date +11.3% (fuel surcharge per cwt implied).'),
        ('q3_rpd_g', '  Memo: Aug-26 revenue per day y/y', 0.124, 'pct', '3-Sep-26 update: August revenue per day +12.4% (memo; not a calibration target).'),
        ('pt_dda', '  FY26 depreciation & amortization ($m)', 372.0, 'm', '1H/26 $184.0m + 2H at the Q2/26 run-rate ($91.7m) + new equipment; not guided on EDGAR.'),
        ('pt_int', '  FY26 interest expense (income), net ($m)', -9.0, 'm', '1H/26 −$4.3m (interest income $4.7m on cash; $20m notes).'),
        ('pt_etr', '  Tax rate', 0.25, 'pct', '1H/26 effective rate 25.0%; FY25 24.8%.'),
        ('pt_capex', '  Capital expenditures ($m)', 380.0, 'm', 'Q2-26 release: ~$380m (real estate & service centers $180m; tractors & trailers $155m; IT & other $45m). 1H/26 $139.6m.'),
    ],
    'lev': (150, 250, 400),
    'lev_note': 'ODFL runs with minimal debt and returns surplus cash (cash $120m at YE25, $284m at 30-Jun-26). Target cash: Bull $150m / Base $250m / Bear $400m.',
    'blocks': [
        ('TPD', 'TPD (LTL tons per day y/y)', 0.02, -0.03, [0.04, 0.05, 0.04, 0.035],
         'Freight-cycle recovery from the 2023–26 industrial recession (tons/day −15% from the 2022 peak) and share gains on service (99% on-time, 0.1% claims ratio).', 'pct', False),
        ('YXF', 'YXF (LTL revenue/cwt ex fuel y/y)', 0.01, -0.015, [0.05, 0.045, 0.045, 0.04],
         'Yield discipline: cost-plus pricing targets inflation + reinvestment; ex-fuel revenue/cwt +4–6% p.a. 2016–25 (weight-per-shipment mix headwind as volumes recover).', 'pct', False),
        ('OR', 'OR (operating ratio)', -0.015, 0.025, [0.705, 0.695, 0.69, 0.685],
         'FY2022 record 70.6%; FY25 74.7% (de-leverage at −8% tons/day). Density recovery on a network with ~30% excess capacity; long-term ambition low/mid-60s.', 'pct', False),
    ],
    'qe_days': {3: workdays.quarter(2026, 3), 4: workdays.quarter(2026, 4)},
    'ae_days': {y: workdays.year(y) for y in (2027, 2028, 2029, 2030)},
    'oth_g': 0.04, 'dda_pct': 0.066, 'sbc_pct': 0.0024, 'etr': 0.25, 'dps_g': 0.08, 'oth_ae': 3.5,
    'def_tax_ae': 20.0, 'ofin_ae': -10.0, 'capex_pct': 0.11, 'disp_pct': 0.006, 'wc': WC, 'min_cash': 100.0,
    'capex_note': '2027E+: input 11% of revenue (2016–25 average ~15%; 2025 7.6% after the 2021–24 network build-out; 2026 plan ~6.5%).',
    'mat': {2026: 0.0, 2027: -20.0, 2028: 0.0, 2029: 0.0, 2030: 0.0},
    'mat_note': 'Series B 2.79% senior notes: $20m paid May-2026 (in 1H/26 actuals); final $20m due May-2027. Revolver undrawn.',
    'debt_note': 'Debt at 30-Jun-26 $20.0m: Series B senior notes (PGIM Note Agreement, May-2020). $400m revolving credit agreement (Wells Fargo, to Mar-2028) undrawn; $31.8m letters of credit.',
    'px': 198.39,
    'px_note': 'VALUATION — share price input $198.39 (latest observable on EDGAR: Form 4 open-market sale 25-Aug-2026 at $198.39; update to market)',
}

# ====================================================================== compute engine
def hist_value(row, col):
    kw = row.kw; path = kw.get('data')
    if not path: return None
    if kw.get('annual_only') and col.kind != 'A': return None
    v = get(pkey(col), path)
    return v if kw.get('textdata') else num(v)

def first(kw, keys):
    for k in keys:
        if k in kw: return k
    return None

def compute(row, col):
    kw = row.kw; k = col.kind
    CTX.col = col
    if kw.get('ann') and col.q: return None, False
    if k in HISTK:
        v = hist_value(row, col)
        if v is not None: return v, True
        for key in ('hist', 'all'):
            if key in kw:
                return kw[key](), False
        return None, False
    if k == 'QE':
        if row.spec2: return None, False
        if 'qe_in' in kw:
            v = kw['qe_in']; return (v[col.q] if isinstance(v, dict) else v), True
        key = first(kw, ('qe', 'fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    if k == 'E26':
        if 'e26_in' in kw: return kw['e26_in'], True
        if 'e26' in kw: return kw['e26'](), False
        if kw.get('avg') and ('qe' in kw or 'qe_in' in kw): return f'=AVERAGE(BW{row.r},BX{row.r},BY{row.r},BZ{row.r})', False
        if kw.get('flow') and not row.spec2 and any(x in kw for x in ('qe', 'qe_in', 'fc')): return S26(row.key), False
        key = first(kw, ('fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    if k == 'AE':
        if 'ae_in' in kw:
            v = kw['ae_in']; return (v[col.year] if isinstance(v, dict) else v), True
        key = first(kw, ('ae', 'fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    return None, False

# ====================================================================== writers
CAGR_G = {'l_tpd', 'l_tons', 'l_ship', 'y_xf', 'y_cwt', 'r_ltl', 'r_oth', 'r_tot', 'is_rev', 'is_swb', 'is_ops', 'is_dda', 'is_op', 'is_ni',
          'eps_dil', 'dps', 'r_ebitda', 'r_adj', 'r_ebit', 'e_adjni', 'e_eps', 'cf_cfo', 'fcf', 'fcf_ps', 'bs_ta', 'bs_teq', 'sh_dil', 'n_sc', 'bs_ppe'}
CAGR_A = {'l_tpd_g', 'y_xf_g', 'y_cwt_g', 'y_fsc_pct', 'or', 'is_or', 'is_opm', 'is_etr', 'r_adj_m', 'r_ebit_m', 'r_ebitda_m', 'gm_rev', 'gm_tpd', 'gm_xf',
          'gm_or', 'gm_ebitda_m', 'gm_op_m', 'gm_net_m', 'gm_swb', 'fcf_m', 'fcf_conv', 'capex_pct', 'ncapex_pct', 'ra_roic', 'rn_ronta', 'roe', 'payout',
          'c_swb', 'c_ops', 'c_gen', 'c_taxlic', 'c_ins', 'c_comm', 'c_pt', 'c_dda', 'c_misc', 'lv_x', 'wc_nwc_p', 'r_rra_pct', 'r_aor'}
ANN_HIST = [ANN[y].c for y in range(2006, 2026)]

def style_cagr(ws, r, on):
    src = 16 if on else 12
    for c in ('CH', 'CI', 'CJ', 'CK'):
        ws[f'{c}{r}']._style = copy.copy(TM[f'{c}{src}']._style)

def write_row(ws, row):
    r = row.r; arch = ARCH[row.arch]
    clone_row_style(ws, arch, r)
    ws.cell(r, 2).value = row.label or None
    if row.kw.get('note'): ws[f'CF{r}'] = row.kw['note']
    if row.kw.get('group'):
        ws.row_dimensions[r].outlineLevel = 1; ws.row_dimensions[r].hidden = True
    cg = row.key in CAGR_G or row.key in CAGR_A
    style_cagr(ws, r, cg)
    if row.arch in ('section', 'block', 'blank', 'sub') and not row.kw.get('data'):
        return
    if row.key in DEFS:
        write_defs(ws, row); return
    nf = NF[row.nf]
    for col in COLS:
        v, inp = compute(row, col)
        if v is None: continue
        cell = ws[f'{col.c}{r}']
        cell.value = v
        if not row.kw.get('textdata'): cell.number_format = nf
        if inp:
            set_color(cell, BLUE, bold=True if col.kind in FCK else None)
        elif row.arch == 'growth' and col.kind in HISTK:
            set_color(cell, GRAY)
        else:
            set_color(cell, BLACK)
        add_cell_comment(row, col, cell)
    if row.key in CAGR_G:
        ws[f'CH{r}'] = f'=IFERROR((CE{r}/BV{r})^(1/5)-1,"n/m")'
        ws[f'CI{r}'] = f'=IFERROR((BV{r}/AW{r})^(1/5)-1,"n/m")'
        ws[f'CJ{r}'] = f'=IFERROR((BV{r}/X{r})^(1/10)-1,"n/m")'
        ws[f'CK{r}'] = f'=IFERROR((BV{r}/C{r})^(1/19)-1,"n/m")'
        for c in ('CH', 'CI', 'CJ', 'CK'): ws[f'{c}{r}'].number_format = '0%;\\(0%\\);"n/m"'
    elif row.key in CAGR_A:
        ws[f'CH{r}'] = f'=IFERROR(AVERAGE(CA{r},CB{r},CC{r},CD{r},CE{r}),"n/m")'
        ws[f'CI{r}'] = f'=IFERROR(AVERAGE({",".join(f"{ANN[y].c}{r}" for y in range(2021, 2026))}),"n/m")'
        ws[f'CJ{r}'] = f'=IFERROR(AVERAGE({",".join(f"{ANN[y].c}{r}" for y in range(2016, 2026))}),"n/m")'
        ws[f'CK{r}'] = f'=IFERROR(AVERAGE({",".join(f"{c}{r}" for c in ANN_HIST)}),"n/m")'
        for c in ('CH', 'CI', 'CJ', 'CK'): ws[f'{c}{r}'].number_format = NF['pct']
    t = LABEL_SOURCES.get(row.key)
    if t: comment(ws.cell(r, 2), t)

def add_cell_comment(row, col, cell):
    if col.kind not in HISTK: return
    d = D.get(pkey(col)) or {}
    x, st = d.get('x') or {}, d.get('st') or {}
    if row.key == 'r_items' and x.get('op_labels'):
        comment(cell, 'Company-identified (pre-tax): ' + x['op_labels'])
    elif row.key == 'e_disc' and x.get('disc_labels'):
        comment(cell, 'Company-identified discrete tax item: ' + x['disc_labels'])
    elif row.key == 'l_days' and st.get('days_calc'):
        comment(cell, 'Not published in this release — computed with the ODFL work-day calendar (workdays.py).', 300, 60)
    elif row.key == 'l_tpd' and (d.get('stats') or {}).get('tons_day') is None and st.get('tpd') is not None:
        comment(cell, 'Derived: tons ÷ work days (not published in this release).', 260, 50)

SRC = 'Source: ODFL earnings releases (Form 8-K Ex. 99.1) — '
LABEL_SOURCES = {
    'l_days': SRC + '"Operating Statistics" table (work days from Q4-12); earlier years computed (ODFL calendar).',
    'l_tpd': SRC + '"Operating Statistics" (LTL tons per day / LTL tonnage per day); annual = Q4 release year-to-date column.',
    'y_xf': SRC + '"Operating Statistics" — LTL revenue per hundredweight, excluding fuel surcharges (statistics basis).',
    'r_ltl': SRC + 'summary table "LTL services revenue / Other services revenue" (published from 2019).',
    'n_sc': 'Source: Form 10-K Item 1 Business / Item 2 Properties (service centers, tractors, linehaul trailers) and Human Capital (employees).',
    'is_rev': 'Source: Statements of Operations (Q-release three-month columns; annual = Q4 release twelve-month column; 10-K for FY2009).',
    'is_dda': 'Source: Statements of Operations (release); cash-flow D&A in the 10-K / 10-Q differs by <$0.05m.',
    'r_op_xb': 'Source: SEC XBRL companyfacts (CIK 0000878927), earliest-filed fact for the period (as originally reported); Q4 = FY − 9M.',
    'r_gain': 'Source: Statements of Cash Flows (10-K / 10-Q year-to-date; discrete quarters derived) — "(Gain) loss on disposal / sale of property and equipment".',
    'r_items': 'Source: ODFL earnings releases / 10-Q MD&A — only items quantified by the company (see cell comments).',
    'cf_ni': 'Source: Statements of Cash Flows (10-K / 10-Q year-to-date; discrete quarters derived; Q4 = FY − 9M).',
    'bs_cash': 'Source: Balance Sheets (10-K / 10-Q).',
    'dps': 'Source: releases / 10-Q / 10-K (split-adjusted). Quarterly dividend initiated Q1-17.',
    'eps_pub': 'Source: releases; restated ×1/6.75 (pre-Aug-2010), ×1/4.5 (to Sep-2012), ×1/3 (to Mar-2020), ×1/2 (to Mar-2024) by document filing date.',
}

DEFS = {}
def defs_setup():
    q = lambda y, n: QC[(y, n)].c
    DEFS['l_def'] = {
        'C': ('To 2013: total tons & shipments (LTL + truckload), revenue/cwt on the total basis; work days computed to 2011',
              'As originally reported. The 2014 releases restated 2013 on the LTL basis; the model keeps the original 2013 figures.'),
        q(2014, 1): ('From Q1-14: LTL tons, LTL shipments and LTL revenue/cwt (statistics exclude undelivered-freight adjustments & non-weight services)', None),
    }
    DEFS['r_def'] = {
        'C': ('Model definition: EBITDA = net income + tax + net interest + other non-op + D&A; Adjusted excludes net (gain) loss on disposal of P&E and company-named one-offs',
              'ODFL publishes no non-GAAP reconciliation. Items: FY2007 −$2.0m (Q2-07 customer pricing resolution, revenue).'),
        q(2017, 4): ('Q4-17: + $9.8m special bonus to non-executive employees after the Tax Cuts and Jobs Act (SWB)', None),
        q(2022, 3): ('Q3-22: − $15.8m one-time SWB reduction (termination of the Executive Chairman employment agreement)', None),
    }
    DEFS['e_def'] = {
        'C': ('Model adjusted EPS = (GAAP net income + adjusting items × (1 − effective tax rate excl. discrete items) + discrete tax items) ÷ diluted shares', None),
        q(2017, 4): ('Q4-17: TCJA revaluation of the net deferred tax liability (−$104.9m benefit removed)', None),
        q(2018, 1): ('From 2018: 21% federal statutory rate (TCJA)', None),
    }

def write_defs(ws, row):
    for c, (txt, cm) in DEFS[row.key].items():
        cell = ws[f'{c}{row.r}']; cell.value = txt
        cell.font = Font(name='Calibri', sz=10, b=True, i=True, color='FFC00000')
        if cm: comment(cell, cm)

def _v(x): return x() if callable(x) else x

def write_scen(ws):
    for spec_ in SCEN_ROWS:
        r, kind = spec_[0], spec_[1]
        src = {'hdr': 9, 'out_hdr': 11, 'pt_hdr': 15, 'out': 12, 'pt': 16, 'lev': 28, 'blk': 30, 'yr': 31}[kind]
        for ci in range(CI('CN'), CI('CR') + 1):
            ws.cell(r, ci)._style = copy.copy(TM.cell(src, ci)._style)
        if kind in ('hdr', 'out_hdr', 'pt_hdr'):
            _, _, a, b, c_, d_, e = spec_
            for col, v in zip(['CN', 'CO', 'CP', 'CQ', 'CR'], [a, b, c_, d_, e]):
                if v is not None: ws[f'{col}{r}'] = v
        elif kind == 'out':
            _, _, lab, hi, mid, lo, note, nf = spec_
            ws[f'CN{r}'] = lab
            for col, v in zip(['CO', 'CP', 'CQ'], [hi, mid, lo]):
                v = _v(v)
                ws[f'{col}{r}'] = v; ws[f'{col}{r}'].number_format = NF[nf]
                set_color(ws[f'{col}{r}'], BLACK if isinstance(v, str) and v.startswith('=') else BLUE)
            ws[f'CR{r}'] = note
        elif kind == 'pt':
            _, _, lab, _, val, _, note, nf = spec_
            ws[f'CN{r}'] = lab; ws[f'CP{r}'] = val; ws[f'CP{r}'].number_format = NF[nf]; set_color(ws[f'CP{r}'], BLUE); ws[f'CR{r}'] = note
        elif kind == 'lev':
            _, _, lab, a, b, c_, note = spec_
            ws[f'CN{r}'] = lab
            for col, v in zip(['CO', 'CP', 'CQ'], [a, b, c_]):
                ws[f'{col}{r}'] = v; ws[f'{col}{r}'].number_format = NF['m']; set_color(ws[f'{col}{r}'], BLUE)
            ws[f'CR{r}'] = note
        elif kind == 'blk':
            _, _, lab, db, mid, de, note, nf = spec_
            ws[f'CN{r}'] = lab; ws[f'CO{r}'] = db; ws[f'CP{r}'] = mid; ws[f'CQ{r}'] = de; ws[f'CR{r}'] = note
            for col in ('CO', 'CQ'):
                ws[f'{col}{r}'].number_format = '\\+#,##0;\\-#,##0' if nf == 'm' else '\\+0.0%;\\-0.0%'
        elif kind == 'yr':
            _, _, lab, hdr, base, clamp, _, nf = spec_
            ws[f'CN{r}'] = lab; ws[f'CP{r}'] = base
            if clamp:
                ws[f'CO{r}'] = f'=MAX(0,CP{r}+$CO${hdr})'; ws[f'CQ{r}'] = f'=MAX(0,CP{r}+$CQ${hdr})'
            else:
                ws[f'CO{r}'] = f'=CP{r}+$CO${hdr}'; ws[f'CQ{r}'] = f'=CP{r}+$CQ${hdr}'
            for col in ('CO', 'CP', 'CQ'): ws[f'{col}{r}'].number_format = NF[nf]

def write_header(ws):
    for r in range(1, 8):
        clone_row_style(ws, r, r, 2, CI('CR'))
    ws['CF1'] = 'SCENARIO SWITCH ▼ (Bull / Base / Bear) — drives all scenario-linked forecast drivers'
    ws['CF2'] = 'Base'
    comment(ws['CF2'], 'Scenario toggle. Choose Bull / Base / Bear. Drives: Q3/Q4-26E operating ratio and Q4/26E tons/day & ex-fuel yield '
            '(best / average / worst 2016–25 seasonality), 2H/26 buybacks, 2027E–2030E tons/day growth, ex-fuel yield growth and operating ratio, '
            'and the target cash balance for 2027E+ buybacks (scenario table CN:CR).')
    ws['B2'] = 'Old Dominion Freight Line, Inc. (Nasdaq: ODFL) — 3-Statement Financial Model'
    ws['B3'] = ('US GAAP as reported · USD millions (per-share data in USD; shares in millions, split-adjusted; tons, shipments & miles in thousands) · '
                'Source: ODFL Forms 10-K / 10-Q, Form 8-K earnings releases (Ex. 99.1) and mid-quarter updates, SEC XBRL companyfacts, SEC EDGAR CIK 0000878927 '
                '(investor-relations releases are the same documents as filed on EDGAR)')
    ws['B4'] = ('Basis notes: calendar fiscal year; single reportable segment (LTL); all periods as originally reported. Stock splits 3-for-2 Aug-2010, Sep-2012, Mar-2020 '
                'and 2-for-1 Mar-2024 (earlier shares / per-share data restated). Operating statistics on the total-tons basis to 2013, LTL from 2014; LTL / other services '
                'revenue split from 2019. ODFL publishes no non-GAAP measures or annual guidance: EBITDA / Adjusted EBITDA / adjusted EPS are model definitions; '
                'Q3/26E anchored to the 3-Sep-26 mid-quarter update, Q4/26E and the Q3/Q4 operating ratio on 2016–25 seasonality.')
    ws['CH5'] = 'Forecast'; ws['CI5'] = 'Historical'; ws['CJ5'] = 'Historical'; ws['CK5'] = 'Historical'
    ws['B6'] = '(USDm)'
    for col in COLS: ws[f'{col.c}6'] = col.label
    ws['CF6'] = 'Modelling Notes'
    for c, a, b in [('CH', '5Y CAGR', "25-'30"), ('CI', '5Y CAGR', "20-'25"), ('CJ', '10Y CAGR', "15-'25"), ('CK', '19Y CAGR', "06-'25")]:
        ws[f'{c}6'] = a; ws[f'{c}7'] = b

def main():
    import openpyxl
    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb['Model']
    blank_style = copy.copy(ws['A1']._style)
    for rowc in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for c in rowc:
            c.value = None; c.comment = None; c._style = copy.copy(blank_style)
    for k in list(ws.row_dimensions.keys()):
        rd = ws.row_dimensions[k]; rd.outlineLevel = 0; rd.hidden = False
    scen_layout(V0)
    build_spec(V0)
    n1 = len(fw.SPEC)
    build_spec2(V0)
    for i, row in enumerate(fw.SPEC): row.spec2 = i >= n1
    last = number_rows(8)
    for r in range(1, last + 1): ws.row_dimensions[r].height = 15.0
    ws.row_dimensions[2].height = 21.0
    defs_setup()
    write_header(ws)
    for row in fw.SPEC: write_row(ws, row)
    write_scen(ws)
    ws.freeze_panes = TM.freeze_panes
    import other_sheets as osh
    osh.do_all(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    osh.inject_charts(OUT)
    print('rows', fw.SPEC[-1].r, 'saved', OUT)
    return wb

RECALC = os.environ.get('XLSX_RECALC', '/mnt/skills/public/xlsx/scripts/recalc.py')
def recalc(path, timeout=600):
    import subprocess
    out = subprocess.run(['python3', RECALC, path, str(timeout)], capture_output=True, text=True, cwd=os.path.dirname(RECALC))
    try:
        return json.loads(out.stdout)
    except Exception:
        return {'raw': out.stdout[-2000:], 'err': out.stderr[-2000:]}

def scenario_snapshots():
    import openpyxl
    import other_sheets as osh
    snaps = {}
    for scen in ('Bull', 'Bear'):
        p = os.path.join(BASE, f'ODFL_{scen}.xlsx')
        wb = openpyxl.load_workbook(OUT); wb['Model']['CF2'] = scen; wb.save(p)
        res = recalc(p); print(scen, res.get('status'), res.get('total_errors'))
        v = openpyxl.load_workbook(p, data_only=True)['Bull-Base-Bear']
        snaps[scen] = {(r, c): v[f'{c}{r}'].value for r in range(9, 48) for c in 'CDEFGHIJKL' if v[f'{c}{r}'].value is not None}
        os.remove(p)
    wb = openpyxl.load_workbook(OUT)
    osh.snapshot_values(wb, snaps['Bull'], snaps['Bear'])
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    osh.inject_charts(OUT)

if __name__ == '__main__':
    main()
    if '--snap' in sys.argv: scenario_snapshots()
