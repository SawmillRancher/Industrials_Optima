"""Rewire the Bull-Base-Bear, DCF and Charts tabs of the MLM template (current layout: DCF with live valuation date, fade period,
implied ROIC, target-IRR entry price, reverse DCF and Mauboussin EV/NOPAT; Bull-Base-Bear with buyback rows) to the BAH Model rows /
columns; keep the template chart XML. The template's DCF_old tab is not carried.

Column mapping: template 2025A base = BAH FY2026A; template 2026E-2030E = BAH FY2027E-FY2031E. Valuation net debt is taken at the
FY2027E year end (31-Mar-2027, after the Ultra payment), so the FY2027E UFCF is excluded from the PV (no double count) and the explicit
discounted years are FY2028E-FY2031E."""
import re, zipfile, shutil
from openpyxl.utils import get_column_letter as L
from fw import R, TEMPLATE, ANN, NOTE

# template Model row -> BAH row key
MAP = {158: 'is_rev', 164: 'r_cash', 166: 'is_oi', 182: 'd_dep', 186: 'sh_dil', 211: 'r_adj', 212: 'r_adj_m', 219: 'r_ebita', 221: 'r_adj',
       229: 'e_adjni', 231: 'e_eps', 238: 'gm_ns', 245: 'gm_ebitda_m', 247: 'gm_op_m', 260: 'cf_wc', 262: 'cf_capex', 263: 'cf_acq',
       280: 'fcf', 282: 'fcf_m', 290: 'bb_px', 291: 'bb_cash', 292: 'bb_sh', 331: 'bs_nci', 408: 'ra_ic', 411: 'ra_roic', 422: 'rn_ronta',
       428: 'lv_nd', 433: 'v_px', 434: 'v_sh'}
# template Model column -> BAH column (template 2025A = BAH FY2026A; template 2026E-2030E = BAH FY2027E-FY2031E)
CMAP = {'BV': ANN[2026].c, 'CA': ANN[2027].c, 'CB': ANN[2028].c, 'CC': ANN[2029].c, 'CD': ANN[2030].c, 'CE': ANN[2031].c, 'CF': NOTE}
REF = re.compile(r"Model!(\$?)([A-Z]{1,3})(\$?)(\d+)")
HIST = [ANN[y].c for y in range(2012, 2027)]

def remap(f):
    def rep(m):
        c, r = m.group(2), int(m.group(4))
        c2 = CMAP.get(c, c)
        r2 = R[MAP[r]] if r in MAP else r
        if c == 'CF' and r == 2: r2 = 2
        return f'Model!{m.group(1)}{c2}{m.group(3)}{r2}'
    return REF.sub(rep, f)

def hist_stat(f):
    """LT historical AVERAGE / MIN / MAX over the template's annual columns -> BAH FY2012-FY2026 annual columns."""
    m = re.match(r'=IFERROR\((AVERAGE|MIN|MAX)\(Model!C(\d+),', f)
    if not m: return None
    fn, r = m.group(1), int(m.group(2))
    rr = R[MAP[r]]
    return f'=IFERROR({fn}({",".join(f"Model!{c}{rr}" for c in HIST)}),"")'

def fy_dates(f):
    f = re.sub(r'DATE\((\d{4}),12,31\)', lambda m: f'DATE({int(m.group(1)) + 1},3,31)', f)
    f = re.sub(r'DATE\((\d{4})-1,12,31\)', lambda m: f'DATE({int(m.group(1)) + 1}-1,3,31)', f)
    return f

TEXT = {
    'Bull-Base-Bear': {
        'B2': 'Booz Allen Hamilton Holding Corporation (BAH) — Bull / Base / Bear Case P&L Summary',
        'B3': ('Adjusted figures (company definitions: adjusted EBITDA, adjusted net income, adjusted diluted EPS) · USD millions · fiscal years ended 31 March · '
               'Live panel linked to Model — toggle scenario via Model!AU2. Bull and Bear panels are value snapshots generated from the Model with the switch set to each '
               'case. FY2027E = Q1 actual + Q2–Q4 calibrated to FY27 guidance (revenue $11.2–11.7bn, adjusted EBITDA $1,240–1,290m), plus Ultra Mission Solutions from 24-Aug-2026.'),
        'G26': 13, 'H26': 13, 'S26': 15, 'T26': 15, 'AD26': 10, 'AE26': 10, 'AF26': 10,
    },
    'DCF': {
        'B2': 'DCF VALUATION — BOOZ ALLEN HAMILTON HOLDING CORPORATION (BAH)',
        'C5': 0.0425, 'D5': 'Assumption (~US 10Y Treasury yield, Oct-2026; government rate sources not reachable from the build environment); update to market',
        'D6': 'Assumption: US implied ERP ~4.5–5.5%',
        'C7': 1.00, 'D7': ('Assumption: government IT / consulting services peers (CACI, SAIC, Leidos, Parsons) ~0.6–1.0; ~98% U.S. government revenue on '
                           'cost-reimbursable contracts, but Civil budget and procurement-policy risk (FY26 Civil −22%) and the 2025–26 de-rating argue for the top of the range'),
        'C9': 0.0575, 'D9': 'Aug-2026 notes at 5.375% (2030) and 5.900% (2034) (8-K 4-Aug-2026); 2033 / 2035 notes 5.950%; marginal cost ~5.75%',
        'C10': 0.225, 'D10': 'Model FY2028E+ effective tax rate (FY27 guidance 20–23%)',
        'C12': 0.30, 'D12': 'Net debt ~$3.9bn at FY2027E YE (post-Ultra) vs ~$9.0bn market cap ($75.00 × ~120m diluted shares); net leverage ~3.0x',
        'D15': ('Periods = (31-Mar of fiscal year − today) / 365.25; cash flows assumed at fiscal year-end (31 March). FY2027E UFCF is excluded from the PV '
                '(its cash flows, incl. the $720m Ultra payment, are in the FY2027E year-end net debt deducted below); explicit years FY2028E–31E.'),
        'B18': 'Adj. EBITA (before acquisition amortization) less acquisition, restructuring, legal & financing costs',
        'B21': 'plus depreciation & other amortization (acquisition amortization is outside the NOPAT basis)',
        'B22': '− Capex and acquisitions (Ultra in FY2027E; M&A lever 2028E–31E)',
        'B26': 'Fade-period growth (2032E–36E; 4.5% → 3.0% linear)', 'I26': 0.045, 'M26': 0.03, 'C27': 0.025,
        'D32': ("Booz Allen's adjusted EBITDA expenses stock-based compensation (not adjusted out since FY2016). NOPAT is taxed on adjusted EBITA (before acquisition "
                'amortization, which is non-cash and excluded) less the cash items the company adjusts out (other corporate / acquisition costs, restructuring, legal '
                'matters, financing costs); row 21 adds back depreciation & other amortization only. FY2027E (column D) is shown but not discounted: its cash flows, '
                'including the $720m Ultra acquisition, are already in the FY2027E year-end net debt. Row 22 includes bolt-on acquisitions in FY2028E–31E; the fade '
                'period and terminal value assume no new acquisitions. Operating leases are treated as operating costs (not in net debt).'),
        'D43': 'Opening IC = Model FY2026A book IC (equity + net debt), rolled forward with capex + acquisitions (incl. Ultra in FY2027E) − depreciation + ΔWC.',
        'B48': 'Less: Net debt + NCI (Model FY2027E year end, 31-Mar-2027, after the Ultra payment)',
        'D48': 'BAH: no non-controlling interests after FY2025; net debt taken at the FY2027E year end, so the FY2027E UFCF (column D) is excluded from the PV.',
        'B50': 'Diluted shares (m, Model FY2027E period end)',
        'D76': 'Rationale (Booz Allen characteristics)',
        'C78': 0.06, 'D78': 'Model adjusted EBITA CAGR FY2027–31E: low-to-mid single-digit National Security growth, Civil trough, Ultra double-digit growth and slight margin expansion → ~6% NOPAT growth',
        'C79': 0.25, 'D79': ('Asset-light services model: invested capital is mostly acquisition goodwill and receivables; organic growth needs little capital '
                             '(capex ~1% of revenue), so incremental returns are high; acquisitions (Liberty, EverWatch, Ultra) earn high-single-digit returns on price paid'),
        'C80': 12, 'D80': ('Cleared workforce, incumbent positions on multi-year IDIQ vehicles and a $39bn backlog sustain the advantage; exposure to U.S. budget '
                           'priorities, procurement reform and outcome-based pricing shortens it'),
        'B90': 'VALUATION  —  Applied to FY2029E NOPAT',
        'B91': 'FY2029E Adj. EBITA less acquisition, restructuring, legal & financing costs',
        'B93': 'FY2029E NOPAT',
        'B95': 'Implied Enterprise Value (FY2029E basis)',
        'B96': 'Less: FY2028E year-end net debt + NCI (Model)',
        'B98': 'Diluted shares FY2029E (m, Model)',
        'B99': 'Implied share price FY2029 ($)',
        'B106': 'Plus: FY2027E year-end net debt + NCI',
        'B108': 'BAH market-implied EV/NOPAT (on FY2029E NOPAT)',
    },
    'Charts': {
        'B2': 'Booz Allen Hamilton — Operating Performance (FY2012–FY2026) and Model Forecast (FY2027E–FY2031E)',
        'B3': ('USDm · fiscal years ended 31 March · Linked to Model. Bars = contribution to revenue growth (pp): National Security (Defense + Intelligence) vs '
               'Civil & Commercial, Ultra and other (customer split published from FY2017; earlier years shown as other). Adjusted EBITDA = company definition. '
               'ROIC / RONTA on adjusted EBIT (full-cost).'),
        'B8': 'National Security contribution (pp)', 'B9': 'Civil & Commercial, Ultra & other (pp)', 'B12': 'Adjusted EBITDA', 'B13': 'Operating income (GAAP)',
        'B14': 'Adj. EBITDA Margin %', 'B15': 'Operating Margin %',
    },
}
BBB_ROWS = {9: 'Revenue', 10: '   Revenue y/y %', 11: '   National Security revenue y/y %', 12: 'Adjusted EBITDA',
            15: 'Operating Income (GAAP)', 16: '   Operating income y/y %', 17: '   Operating margin (GAAP) %', 19: 'Adjusted Net Income',
            35: 'Adjusted EBITDA (leverage basis)', 36: '   Adjusted EBITDA % of revenue'}
CHART_TITLES = {
    'Aggregates Revenue Drivers: Shipment Volume vs. Price (ASP), y/y %': 'Revenue Growth: National Security vs. Civil &amp; Commercial, Ultra &amp; Other, pp',
    'Adjusted EBITDA &amp; Earnings from Operations (USDm)': 'Adjusted EBITDA &amp; Operating Income (USDm)',
}

def build_charts(ws):
    for i, y in enumerate(range(2012, 2032)):
        c = L(3 + i); mc = ANN[y].c; pc = ANN[y - 1].c if y - 1 in ANN else None
        ws[f'{c}7'] = y if y <= 2026 else f'{y}E'
        if pc:
            ws[f'{c}8'] = (f'=IF(AND(ISNUMBER(Model!{mc}{R["ns"]}),ISNUMBER(Model!{pc}{R["ns"]})),'
                           f'IFERROR((Model!{mc}{R["ns"]}-Model!{pc}{R["ns"]})/Model!{pc}{R["is_rev"]},0),0)')
            ws[f'{c}9'] = f'=IFERROR(Model!{mc}{R["gm_rev"]}-{c}8,0)'
        else:
            ws[f'{c}8'] = 0
            ws[f'{c}9'] = 0
        ws[f'{c}11'] = f'={c}7'
        ws[f'{c}12'] = f'=Model!{mc}{R["r_adj"]}'
        ws[f'{c}13'] = f'=Model!{mc}{R["is_oi"]}'
        ws[f'{c}14'] = f'=Model!{mc}{R["gm_ebitda_m"]}'
        ws[f'{c}15'] = f'=Model!{mc}{R["gm_op_m"]}'
        ws[f'{c}17'] = f'={c}7'
        ws[f'{c}18'] = f'=IFERROR(Model!{mc}{R["ra_roic"]}+0,NA())'
        ws[f'{c}19'] = f'=IFERROR(Model!{mc}{R["rn_ronta"]}+0,NA())'
        ws[f'{c}20'] = f'=IFERROR(Model!{mc}{R["gm_ebitda_m"]}+0,NA())'
        ws[f'{c}21'] = f'=IFERROR(Model!{mc}{R["fcf_m"]}+0,NA())'

def do_all(wb):
    for name in ('Bull-Base-Bear', 'DCF', 'Charts'):
        ws = wb[name]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and 'Model!' in c.value:
                    h = hist_stat(c.value) if name == 'Bull-Base-Bear' else None
                    c.value = h or remap(c.value)
                if isinstance(c.value, str) and 'DATE(' in c.value:
                    c.value = fy_dates(c.value)
        for addr, t in TEXT.get(name, {}).items():
            ws[addr] = t
    b = wb['Bull-Base-Bear']
    for r, t in BBB_ROWS.items():
        for c in ('B', 'N', 'Z'): b[f'{c}{r}'] = t
    for c0 in ('C', 'O', 'AA'):
        i = b[f'{c0}7'].column
        b.cell(7, i).value = '2026A'
        for k in range(1, 6): b.cell(7, i + k).value = f'{2026 + k}E'
        b.cell(7, i + 6).value = "26-'31 CAGR"
    b['J2'] = 'Share price ($):'
    for c in ('B', 'N', 'Z'): b[f'{c}29'] = 'IRR (to 31-Mar of Y-1)'
    for c in ('B', 'N', 'Z'):
        if b[f'{c}40'].value: b[f'{c}40'] = '   Share repurchases ($m; FY2027E = Q2–Q4)'
    d = wb['DCF']
    for i, col in enumerate('DEFGHIJKLM'):
        d[f'{col}17'] = f'{2027 + i}E'
    a26, a27 = ANN[2026].c, ANN[2027].c
    d['D22'] = f'=Model!{a27}{R["cf_capex"]}+Model!{a27}{R["cf_acq"]}'
    d['D34'] = f'=Model!{a26}{R["ra_ic"]}'
    d['D31'] = '=0*D24*D30'
    d['P24'] = '=IFERROR((H24/E24)^(1/(COLUMNS(E24:H24)-1))-1,"n/m")'
    d['Q24'] = '=IFERROR((M24/E24)^(1/(COLUMNS(E24:M24)-1))-1,"n/m")'
    from openpyxl.comments import Comment
    c3 = Comment('FY2027E UFCF excluded from the PV: valuation net debt (C48) is taken at the FY2027E year end, which already reflects the FY2027E cash flows '
                 '(including the $720m Ultra Mission Solutions payment).', 'BAH model')
    c3.width, c3.height = 380, 90; d['D31'].comment = c3
    cm = Comment('P/E on adjusted diluted EPS (company definition). BAH at $75.00 (user-provided, 12-Oct-2026) trades at ~12x FY2027E adjusted EPS. Exit multiples are '
                 'assumptions: Base 13x, Bull 15x, Bear 10x applied to FY2030E–31E (FY2029E–31E for Bear), reflecting government-services peers and the Civil '
                 'budget risk.', 'BAH model')
    cm.width, cm.height = 420, 120; b['N26'].comment = cm
    cm2 = Comment(f'Snapshot (values) of Model rows {R["bb_cash"]}, {R["bb_px"]} and {R["bb_sh"]} (Share Buyback Schedule) with Model!{NOTE}2 = Bull / Bear. '
                  'Base panel (D40:H42) is live-linked.', 'BAH model')
    cm2.width, cm2.height = 420, 90; b['P40'].comment = cm2
    d['P17'] = 'CAGR 28–31E'; d['Q17'] = 'CAGR 28–36E'
    build_charts(wb['Charts'])

def snapshot_values(wb, bull, bear):
    ws = wb['Bull-Base-Bear']
    cols = 'CDEFGHIJKL'
    for snap, off in ((bull, 'OPQRSTUVWX'), (bear, ['AA', 'AB', 'AC', 'AD', 'AE', 'AF', 'AG', 'AH', 'AI', 'AJ'])):
        for (r, c), v in snap.items():
            tgt = f'{off[cols.index(c)]}{r}'
            cur = ws[tgt].value
            if isinstance(cur, str) and cur.startswith('='): continue
            if r in (7, 26): continue
            ws[tgt] = v

def inject_charts(path):
    """Replace openpyxl-rewritten chart parts with the template chart XML (combo charts, colours, layout), retitled, ranges kept
    (Charts!C:V = FY2012-FY2031E) and cached values removed so Excel redraws from the BAH cells."""
    tz = zipfile.ZipFile(TEMPLATE)
    tmpl = {n: tz.read(n) for n in tz.namelist() if n.startswith('xl/charts/chart')}
    tmp = path + '.tmp'
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in tmpl:
                t = tmpl[item.filename].decode('utf-8')
                for a, b2 in CHART_TITLES.items(): t = t.replace(a, b2)
                t = re.sub(r'<c:numCache>.*?</c:numCache>', '', t, flags=re.S)
                t = re.sub(r'<c:strCache>.*?</c:strCache>', '', t, flags=re.S)
                data = t.encode('utf-8')
            zout.writestr(item, data)
    shutil.move(tmp, path)
