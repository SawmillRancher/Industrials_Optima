"""Bull / Base / Bear scenario input table (right of the Model notes column) — same layout as the RBA / HII models."""
from openpyxl.styles import Alignment

from .framework import BLACK, BLUE, FMT_EPS, FMT_NUM, FMT_PCT, GREY, LIGHT, MID, WHITE, fill, font, put_comment
from .rows import S_BUYB, S_CAPEX, S_EPS_FY, S_EPS_Q3, S_GAAP_FY, S_GR, S_MA, S_MG, S_TAX, SEG_NAMES


def write_scenarios(ws, cols, S):
    L, NB, NBa, NBe, NOTE = cols.SCN_LABEL, cols.SCN["Bull"], cols.SCN["Base"], cols.SCN["Bear"], cols.SCN_NOTE

    def hdr(r, text, a="Bull", b="Base", c="Bear"):
        ws[f"{L}{r}"] = text
        ws[f"{L}{r}"].font = font(12, True, color=WHITE)
        ws[f"{L}{r}"].fill = fill(MID)
        for col, v in ((NB, a), (NBa, b), (NBe, c)):
            if v is not None:
                ws[f"{col}{r}"] = v
            ws[f"{col}{r}"].font = font(12, True, color=WHITE)
            ws[f"{col}{r}"].fill = fill(MID)
            ws[f"{col}{r}"].alignment = Alignment(horizontal="center")

    ws[f"{L}9"] = "SCENARIO INPUT TABLE — driver assumptions per Bull / Base / Bear"
    ws[f"{L}9"].font = font(11, True, color=BLACK)
    for col, v in ((NB, "Bull"), (NBa, "Base"), (NBe, "Bear")):
        ws[f"{col}9"] = v
        ws[f"{col}9"].font = font(12, True, color=BLACK)
        ws[f"{col}9"].fill = fill(LIGHT)
        ws[f"{col}9"].alignment = Alignment(horizontal="center")
    ws[f"{NOTE}9"] = ("Bull / Bear = Base + Δ (Δ inputs on each block header). Outlook block = high / mid / low end "
                      "of FY26 ranges.")
    ws[f"{NOTE}9"].font = font(10)

    def val(cell, v, fmt):
        ws[cell] = v
        is_formula = isinstance(v, str) and v.startswith("=")
        ws[cell].font = font(12, color=BLACK if is_formula else BLUE)
        ws[cell].number_format = fmt

    hdr(11, S["outlook_hdr"], "High", "Mid", "Low")
    put_comment(ws[f"{L}11"], S["outlook_comment"])
    for r, lab, vals, fmt in [(S_EPS_FY, "  FY26 non-GAAP diluted EPS ($)", S["eps_fy"], FMT_EPS),
                              (S_EPS_Q3, "  Q3/26 non-GAAP diluted EPS ($)", S["eps_q3"], FMT_EPS),
                              (S_GAAP_FY, "  FY26 GAAP diluted EPS ($, memo)", S["gaap_fy"], FMT_EPS)]:
        ws[f"{L}{r}"] = lab
        ws[f"{L}{r}"].font = font(12)
        for col, v in zip((NB, NBa, NBe), vals):
            val(f"{col}{r}", v, fmt)
    hdr(15, "FY26 point estimates (all cases)", None, None, None)
    for r, lab, v, fmt in [(S_TAX, "  Tax rate (GAAP & non-GAAP)", S["tax"], FMT_PCT),
                           (S_CAPEX, "  Capital expenditures ($m)", S["capex"], FMT_NUM)]:
        ws[f"{L}{r}"] = lab
        ws[f"{L}{r}"].font = font(12)
        val(f"{NBa}{r}", v, fmt)
    if S.get("point_note"):
        ws[f"{NOTE}15"] = S["point_note"]
        ws[f"{NOTE}15"].font = font(10)

    def lever(base_row, name, base_vals, d_bull, d_bear, fmt=FMT_PCT, years=(2027, 2028, 2029, 2030), floor=False,
              dfmt='\\+0.0%;\\-0.0%', note=""):
        hdr(base_row, name, None, "Δ →", None)
        for col, d in ((NB, d_bull), (NBe, d_bear)):
            ws[f"{col}{base_row}"] = d
            ws[f"{col}{base_row}"].font = font(12, True, color=BLUE)
            ws[f"{col}{base_row}"].fill = fill(GREY)
            ws[f"{col}{base_row}"].number_format = dfmt
        if note:
            ws[f"{NOTE}{base_row}"] = note
            ws[f"{NOTE}{base_row}"].font = font(10)
        for i, (y, bv) in enumerate(zip(years, base_vals)):
            r = base_row + 1 + i
            ws[f"{L}{r}"] = f"  {y}E" if isinstance(y, int) else f"  {y}"
            ws[f"{L}{r}"].font = font(12)
            val(f"{NBa}{r}", bv, fmt)
            for col in (NB, NBe):
                f = f"={NBa}{r}+${col}${base_row}"
                if floor:
                    f = f"=MAX(0,{NBa}{r}+${col}${base_row})"
                val(f"{col}{r}", f, fmt)

    for s in ("di", "inst", "ade", "es"):
        g = S["gr"][s]
        lever(S_GR[s], f"{s.upper()}_organicGrowth ({SEG_NAMES[s]} sales y/y)", g["base"], *g["d"], note=g["note"])
    for s in ("di", "inst", "ade", "es"):
        m = S["mg"][s]
        lever(S_MG[s], f"{s.upper()}_margin ({SEG_NAMES[s]} non-GAAP OI margin)", m["base"], *m["d"], note=m["note"])
    lever(S_MA, "M&A_spend (acquisition spend, $m)", S["ma"]["base"], *S["ma"]["d"], fmt=FMT_NUM, floor=True,
          dfmt='\\+#,##0;\\-#,##0', note=S["ma"]["note"])
    lever(S_BUYB, "Buybacks ($m)", S["buyb"]["base"], *S["buyb"]["d"], fmt=FMT_NUM,
          years=("2H/26E", 2027, 2028, 2029, 2030), floor=True, dfmt='\\+#,##0;\\-#,##0', note=S["buyb"]["note"])
