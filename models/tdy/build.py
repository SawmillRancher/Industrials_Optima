"""Build models/TDY_Model.xlsx — Teledyne Technologies (NYSE: TDY) 3-statement model on the RBA template.

Usage:  RECALC_SCRIPT=<path to recalc.py> python -m models.tdy.build   (from the repo root)
"""
from __future__ import annotations

import json
import os
import subprocess
import sys

import openpyxl

from .assumptions import ASSUMP, DEF_NI, DEF_OI, SCENARIOS
from .data import CATS, load
from .framework import Cols, put_comment
from .other_tabs import write_bbb, write_charts, write_dcf
from .rows import build_rows
from .scenario import write_scenarios
from .writer import Data, set_widths, write_model

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
OUT = os.path.join(ROOT, "models", "TDY_Model.xlsx")
RECALC = os.environ.get("RECALC_SCRIPT", "")

TITLE = "Teledyne Technologies Incorporated (NYSE: TDY) — 3-Statement Financial Model"
SUB = ("US GAAP as reported · USD millions (per-share data in USD) · Source: Teledyne Forms 10-K / 10-Q and Form 8-K "
       "earnings releases (Ex. 99.1), SEC EDGAR CIK 0001094285 (investor-relations releases are the same documents as "
       "filed on EDGAR)")
BASIS = ("Basis notes: 52/53-week fiscal year (FY2009, FY2015, FY2020 = 53 weeks). FY2006–07 include Teledyne Continental "
         "Motors (piston engines, sold Apr-2011; discontinued from the FY2010 10-K). Segments on the current four-segment "
         "basis from FY2008 (FY2010 10-K recast); FY2006–07 legacy segments shown as memo. ASU 2017-07 (2018) recast 2017; "
         "ASC 606 from 2018 (2017 not restated). FLIR from 14-May-2021. Q3/26E–Q4/26E = FY26 non-GAAP EPS outlook less "
         "1H/26 actuals.")

FC_CATS = {"AMORT", "ACQ", "STEPUP"}          # always present in the bridges (forecast drivers)


def used_cats(P):
    used_o, used_n = set(FC_CATS), set(FC_CATS)
    for v in P.values():
        for k, val in v.items():
            if not isinstance(val, (int, float)) or abs(val) < 1e-9:
                continue
            if k.startswith("adjO."):
                used_o.add(k[5:])
            elif k.startswith("adjN."):
                used_n.add(k[5:])
    cats_o = [(k, lab) for k, lab, o, n in CATS if o and k in used_o]
    cats_n = [(k, lab) for k, lab, o, n in CATS if n and k in used_n]
    return cats_o, cats_n


def add_item_comments(ws, cols, rowmap, labels):
    for c in cols.list:
        for key, labs in labels.get(c.key, {}).items():
            pre, cat = key.split(".", 1)
            if pre == "segitems":
                rk = "corp_it" if cat == "corp" else f"it_{cat}"
            elif pre in ("adjO", "adjN"):
                if cat == "_per_share":
                    rk = "b_adj" if pre == "adjN" else "o_adj"
                else:
                    rk = ("o_" if pre == "adjO" else "n_") + cat
            else:
                continue
            if rk in rowmap and labs:
                cell = ws.cell(rowmap[rk], c.idx)
                if cell.value is not None:
                    put_comment(cell, "Company wording: " + " | ".join(labs))


def current_def_keys(cols, P):
    """Columns where the company's own non-GAAP definition is the current one (amortization-based, from the
    Q2-21 release, recast to Q1/20)."""
    out = set()
    for c in cols.list:
        if c.is_forecast or c.year >= 2021:
            out.add(c.key)
        elif c.year == 2020 and P.get(c.key, {}).get("adjN.adj_ni_published") is not None:
            out.add(c.key)
    return out


def build(path=OUT, scenario="Base"):
    P, labels, notes, supp, outlook, conflicts = load()
    if conflicts:
        print("NOTE data overlaps (later file wins):", conflicts[:10])
    cats_o, cats_n = used_cats(P)
    cols = Cols()
    info = {"assump": ASSUMP, "def_oi": DEF_OI, "def_ni": DEF_NI, "cats_o": cats_o, "cats_n": cats_n,
            "o_keys": {f"o_{k}" for k, _ in cats_o}, "cur_from": current_def_keys(cols, P)}
    rows = build_rows(info)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model"
    rowmap = write_model(ws, cols, rows, Data(P), TITLE, SUB, BASIS)
    ws[f"{cols.NOTES}2"] = scenario
    add_item_comments(ws, cols, rowmap, labels)
    set_widths(ws, cols)
    write_scenarios(ws, cols, SCENARIOS)
    write_bbb(wb, cols, rowmap)
    write_dcf(wb, cols, rowmap)
    write_charts(wb, cols, rowmap)
    wb.create_sheet("Sheet1")
    wb.save(path)
    return cols, rowmap


def recalc(path):
    if not RECALC:
        print("RECALC_SCRIPT not set — skipping recalculation")
        return None
    out = subprocess.run([sys.executable, RECALC, path, os.environ.get("RECALC_TIMEOUT", "900")], capture_output=True,
                         text=True)
    try:
        return json.loads(out.stdout)
    except Exception:
        print(out.stdout, out.stderr)
        return None


def snapshot(path, cols, rowmap):
    """Run the model in Bull and Bear and paste values into the Bull-Base-Bear snapshot panels."""
    vals = {}
    for scn in ("Bull", "Bear"):
        tmp = path.replace(".xlsx", f"_{scn}.xlsx")
        wb = openpyxl.load_workbook(path)
        wb["Model"][f"{cols.NOTES}2"] = scn
        wb.save(tmp)
        recalc(tmp)
        wbv = openpyxl.load_workbook(tmp, data_only=True)
        b = wbv["Bull-Base-Bear"]
        vals[scn] = {(r, i): b.cell(r, 3 + i).value for r in range(9, 48) for i in range(0, 11)}
        os.remove(tmp)
    wb = openpyxl.load_workbook(path)
    b = wb["Bull-Base-Bear"]
    for scn, start in (("Bull", 15), ("Bear", 27)):   # O and AA columns
        for (r, i), v in vals[scn].items():
            if r in (26, 27, 28, 29, 34, 44, 45, 46, 47):
                continue      # keep the live P/E-driven valuation rows as formulas
            cell = b.cell(r, start + i)
            if cell.value is not None and isinstance(cell.value, str) and cell.value.startswith("="):
                cell.value = v
    wb.save(path)


if __name__ == "__main__":
    cols, rowmap = build()
    res = recalc(OUT)
    print(json.dumps(res, indent=1)[:3000] if res else "no recalc")
    if res and res.get("status") in ("success", "errors_found"):
        snapshot(OUT, cols, rowmap)
        res = recalc(OUT)
        print(json.dumps(res, indent=1)[:2000])
