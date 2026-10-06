"""Build LOAR_Model.xlsx from the MLM template (same layout, formatting and process as the prior models)."""
import sys, copy, os, json
from openpyxl.styles import Font
from fw import *
import fw
from spec import build_spec, scen_layout, SCEN_ROWS, SC
from spec2 import build_spec2

OUT = os.path.join(os.path.dirname(BASE), 'LOAR_Model.xlsx')

# ====================================================================== inputs (V0)
def wc_inputs():
    x = DATA['2025']['x']; q = DATA['Q2-26']['x']
    rev, cogs = (DATA['Q1-26']['x']['rev'] + q['rev']) * 2, (DATA['Q1-26']['x']['cogs'] + q['cogs']) * 2
    return {'dso': round(q['bs_ar'] / rev * 365), 'dio': round(q['bs_inv'] / cogs * 365), 'dpo': round(q['bs_ap'] / cogs * 365),
            'ocl': round(q['bs_ocl'] / rev, 3)}
WC = wc_inputs()
WC['note'] = (f'Forecast inputs at 30-Jun-2026 levels on annualised 1H/26 sales / cost of sales (DSO {WC["dso"]}d, inventory {WC["dio"]}d, '
              f'payables {WC["dpo"]}d, accrued & other current liabilities {WC["ocl"]:.1%} of sales).')

V0 = {
    'out_hdr': 'FY26 outlook — Q2-26 release (6-Aug-26)',
    'out_note': 'Raised in the Q2-26 release (from $645–655m sales / $257–262m Adjusted EBITDA / $1.26–1.30 Adjusted EPS in Q1-26). Includes LMB (Dec-25) and Harper Engineering (Jan-26).',
    'outlook': [
        ('out_rev', '  FY26 net sales ($m)', (675.0, 670.0, 665.0), 'm1', 'Guided $665–675m (+35% y/y at mid). Drives the Q3/Q4-26E revenue calibration.'),
        ('out_ebitda', '  FY26 Adjusted EBITDA ($m)', (270.0, 267.5, 265.0), 'm1', 'Guided $265–270m (~40% margin). Drives the Q3/Q4-26E margin calibration.'),
        ('out_ni', '  FY26 net income ($m) — memo', (60.0, 58.0, 56.0), 'm1', 'Guided $56–60m (memo; compared with the model, not forced).'),
        ('out_eps', '  FY26 Adjusted EPS ($, current definition) — memo', (1.36, 1.34, 1.32), 'ps', 'Guided $1.32–1.36 (diluted EPS $0.57–0.62).'),
        ('bb26', '  2H/26 share repurchases ($m)', (0.0, 0.0, 0.0), 'm', 'No repurchase authorization; capital deployed to M&A / deleveraging.'),
    ],
    'points': [
        ('pt_int', '  Interest expense, net ($m)', 80.0, 'm1', 'FY26 guidance: ~$80m (1H/26 $38.7m).'),
        ('pt_dep', '  Depreciation ($m)', 15.0, 'm1', 'FY26 guidance: ~$15m (1H/26 $6.6m).'),
        ('pt_amort', '  Amortization ($m)', 65.0, 'm1', 'FY26 guidance: ~$65m (1H/26 $32.3m).'),
        ('pt_etr', '  Tax rate (GAAP & adjusted)', 0.24, 'pct', '1H/26 effective rate 23.9% (Q2-26 23.3%); not guided. Net income guidance mid-point implies ~24%.'),
        ('pt_capex', '  Capital expenditures ($m)', 16.0, 'm1', 'Not guided; 1H/26 $7.0m, FY25 $13.0m; ~2.4% of FY26E sales.'),
        ('pt_trans', '  Transaction expenses ($m)', 4.0, 'm1', '1H/26 $2.8m (Harper Engineering); analyst point estimate for ongoing deal activity.'),
    ],
    'lev': (3.0, 2.5, 2.0),
    'lev_note': ('Net debt $830m at 30-Jun-26 = ~3.6x LTM Adjusted EBITDA (~3.1x FY26E guidance mid-point) after LMB / Harper; < 5.5x keeps the SOFR + 4.25% margin. '
                 'Floor = leverage at which surplus cash is returned via buybacks; until then free cash flow prepays term loans. Bull keeps more firepower deployed (higher floor).'),
    'blocks': [
        ('COM_OEM', 'COM_OEM (commercial aerospace OEM sales y/y)', 0.03, -0.06, [0.14, 0.12, 0.10, 0.09],
         'Boeing 737 / 787 and Airbus A320 / A350 rate ramps; Harper interiors latches; FY26 commercial + BJ OEM outlook +17–20%.', 'pct', False),
        ('COM_AM', 'COM_AM (commercial aerospace aftermarket y/y)', 0.02, -0.04, [0.10, 0.09, 0.08, 0.08],
         'Installed-base growth, aging fleets and price (proprietary, sole-source parts); FY26 outlook low-double digits.', 'pct', False),
        ('BJ_OEM', 'BJ_OEM (business jet & GA OEM y/y)', 0.02, -0.05, [0.07, 0.06, 0.06, 0.05],
         'Business-jet deliveries (Gulfstream G700/G800, Bombardier Global 8000, Textron); auto-throttle content.', 'pct', False),
        ('BJ_AM', 'BJ_AM (business jet & GA aftermarket y/y)', 0.02, -0.03, [0.08, 0.07, 0.07, 0.06],
         'Flight hours of the business-jet fleet; brakes and replacement parts.', 'pct', False),
        ('DEF_OEM', 'DEF_OEM (defense OEM y/y)', 0.02, -0.04, [0.08, 0.07, 0.07, 0.06],
         'Military aircraft and systems build (LMB fans & motors from Dec-25); FY26 defense outlook mid-single digits.', 'pct', False),
        ('DEF_AM', 'DEF_AM (defense aftermarket y/y)', 0.02, -0.03, [0.06, 0.06, 0.05, 0.05], 'Readiness / sustainment spares.', 'pct', False),
        ('OTH_OEM', 'OTH_OEM (other / non-aerospace OEM y/y)', 0.02, -0.04, [0.04, 0.04, 0.04, 0.04], 'Industrial applications (~6% of sales).', 'pct', False),
        ('OTH_AM', 'OTH_AM (other / non-aerospace aftermarket y/y)', 0.02, -0.04, [0.04, 0.04, 0.04, 0.04], 'Industrial aftermarket.', 'pct', False),
        ('EBM', 'EBM (Adjusted EBITDA margin, current portfolio)', 0.01, -0.025, [0.405, 0.41, 0.415, 0.42],
         'FY25 38.1%; 1H/26 40.5%; FY26 guidance ~40%. Strategic value drivers (pricing, cost, productivity) + mix towards aftermarket; ~50 bps p.a. expansion.', 'pct', False),
        ('MNA', 'M&A_spend (bolt-on acquisitions, $m)', 200.0, -300.0, [300.0, 300.0, 300.0, 300.0],
         'Loar deployed $60m (2023), $384m (2024), $508m (2025) and $250m (1H/26) on acquisitions; ~$750m pipeline. Base ~$300m p.a.', 'm', True),
    ],
    'qe_acq': {3: 31.0, 4: 30.0},
    'qe_acq_note': ('Q3/Q4-26E inputs: LMB (~$15m a quarter; ~$60m FY26 per the company) + Harper Engineering (~$16m a quarter) not in the prior-year quarter. '
                    'Q2-26 net acquisition sales $33.3m incl. Beadlight (anniversaries 28-Jul-26). Q4: LMB consolidated from 23-Dec-25 (≈ one week in Q4-25).'),
    'dep_pct': 0.022, 'sga_pct': 0.27, 'sbc_pct': 0.026, 'etr': 0.24, 'def_tax_ae': -3.0, 'capex_pct': 0.024, 'opt_ae': 2.0,
    'dcost_q': 1.0,
    'amort_sched': {2027: 65.2, 2028: 65.1, 2029: 65.1, 2030: 62.5},
    'amort_note': (MAN['intang_sched_src'] + ' 10-K schedule 2027 $50.6m, 2028–29 $50.6m, 2030 $48.0m; + Harper ~$10.2m p.a. ($153.5m over ~15 yrs) '
                   '+ ~$4.4m other long-term-asset amortization / LMB measurement-period step-up (FY26 guide $65m vs ~$60.6m schedule).'),
    'wc': WC, 'min_cash': 75.0,
    'r_fac': {2027: 0.0825, 2028: 0.08, 2029: 0.08, 2030: 0.08},
    'r_fac_note': 'Q2-26 interest $20.0m on ~$963m term loans ≈ 8.3% (SOFR ~3.9% + 4.25% + debt-cost amortization ~0.4%, less interest income). 2028+: lower SOFR / repricing as leverage falls.',
    'px': 75.00,
    'px_note': 'VALUATION — share price input $75.00 (NYSE: LOAR, early Oct-2026 quote; ~$7.1bn market cap); update to market',
    'mna_note': 'Scenario lever M&A_spend (Base ~$300m p.a.; Loar acquired $508m in 2025 incl. LMB and $250m in 1H/26). Spend funded from cash / term-loan draws (debt schedule).',
    'rev_guid_note': 'Company FY26 net sales guidance $665–675m (raised in the Q2-26 release, 6-Aug-2026). Drives the revenue-outlook calibration.',
    'eb_guid_note': 'Company FY26 Adjusted EBITDA guidance $265–270m (raised in the Q2-26 release). Drives the margin calibration.',
}

# ====================================================================== compute engine
def hist_value(row, col):
    kw = row.kw; path = kw.get('data')
    if not path: return None
    if kw.get('ppa'):
        v = MAN['ppa'].get(pkey(col), {}).get(path.split('_', 1)[1])
        return v if kw.get('textdata') else num(v)
    v = xget(pkey(col), path)
    return v if kw.get('textdata') else num(v)

def first(kw, keys):
    for k in keys:
        if k in kw: return k
    return None

def compute(row, col):
    kw = row.kw; k = col.kind
    CTX.col = col
    if kw.get('ann') and col.q: return None, False
    if k in HISTK:
        v = hist_value(row, col)
        if v is not None: return v, True
        for key in ('hist', 'all'):
            if key in kw:
                return kw[key](), False
        return None, False
    if k == 'QE':
        if row.spec2: return None, False
        if 'qe_in' in kw:
            v = kw['qe_in']; return (v[col.q] if isinstance(v, dict) else v), True
        key = first(kw, ('qe', 'fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    if k == 'E26':
        if 'e26_in' in kw: return kw['e26_in'], True
        if 'e26' in kw: return kw['e26'](), False
        if kw.get('avg') and ('qe' in kw or 'qe_in' in kw): return f'=AVERAGE({Q1}{row.r},{Q2}{row.r},{Q3}{row.r},{Q4}{row.r})', False
        if kw.get('flow') and not row.spec2 and any(x in kw for x in ('qe', 'qe_in', 'fc')): return S26(row.key), False
        key = first(kw, ('fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    if k == 'AE':
        if 'ae_in' in kw:
            v = kw['ae_in']; return (v[col.year] if isinstance(v, dict) else v), True
        key = first(kw, ('ae', 'fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    return None, False

# ====================================================================== writers
CAGR_G = {'is_rev', 'g_rev', 'com_tot', 'bj_tot', 'def_tot', 'oth_tot', 'mx_oem', 'mx_am', 'r_adj', 'g_adj', 'r_ebitda', 'is_op', 'is_dda',
          'o_adj', 'e_adjni', 'e_eps', 'cf_cfo', 'fcf', 'bs_ta', 'bs_teq', 'sh_dil', 'is_gp', 'is_sga', 'd_amort', 'd_dep', 'a_org'}
CAGR_A = {'g_adjm', 'r_adj_m', 'is_gm', 'is_opm', 'o_adj_m', 'r_ebit_m', 'gm_rev', 'gm_org', 'gm_gm', 'gm_ebitda_m', 'gm_op_m', 'gm_net_m', 'gm_sga',
          'gm_adjop_m', 'fcf_m', 'fcf_conv', 'capex_pct', 'ra_roic', 'rn_ronta', 'roe', 'lv_x', 'a_org_g', 'a_acq_g', 'd_sga', 'd_dep_pct',
          'd_sbc_pct', 'wc_nwc_p', 'mx_am_pct', 'is_etr', 'fcf_eb'}

def style_cagr(ws, r, on):
    src = 16 if on else 12
    for c in CAGR:
        ws[f'{c}{r}']._style = copy.copy(TM[f'{L(CI(c) + TSHIFT)}{src}']._style)

def write_row(ws, row):
    r = row.r; arch = ARCH[row.arch]
    clone_row_style(ws, arch, r)
    ws.cell(r, 2).value = row.label or None
    if row.kw.get('note'): ws[f'{NOTE}{r}'] = row.kw['note']
    if row.kw.get('group'):
        ws.row_dimensions[r].outlineLevel = 1; ws.row_dimensions[r].hidden = True
    style_cagr(ws, r, row.key in CAGR_G or row.key in CAGR_A)
    if row.arch in ('section', 'block', 'blank', 'sub') and not row.kw.get('data'):
        return
    if row.key in DEFS:
        write_defs(ws, row); return
    nf = NF[row.nf]
    for col in COLS:
        v, inp = compute(row, col)
        if v is None: continue
        cell = ws[f'{col.c}{r}']
        cell.value = v
        if not row.kw.get('textdata'): cell.number_format = nf
        if inp:
            set_color(cell, BLUE, bold=True if col.kind in FCK else None)
        elif row.arch == 'growth' and col.kind in HISTK:
            set_color(cell, GRAY)
        else:
            set_color(cell, BLACK)
    a = {y: ANN[y].c for y in ANN}
    if row.key in CAGR_G:
        ws[f'AR{r}'] = f'=IFERROR(({a[2030]}{r}/{a[2025]}{r})^(1/5)-1,"n/m")'
        ws[f'AS{r}'] = f'=IFERROR(({a[2025]}{r}/{a[2020]}{r})^(1/5)-1,"n/m")'
        ws[f'AT{r}'] = f'=IFERROR(({a[2025]}{r}/{a[2015]}{r})^(1/10)-1,"n/m")'
        ws[f'AU{r}'] = f'=IFERROR(({a[2025]}{r}/{a[2012]}{r})^(1/13)-1,"n/m")'
        for c in CAGR: ws[f'{c}{r}'].number_format = '0%;\\(0%\\);"n/m"'
    elif row.key in CAGR_A:
        ws[f'AR{r}'] = f'=IFERROR(AVERAGE({",".join(f"{a[y]}{r}" for y in range(2026, 2031))}),"n/m")'
        ws[f'AS{r}'] = f'=IFERROR(AVERAGE({",".join(f"{a[y]}{r}" for y in range(2021, 2026))}),"n/m")'
        ws[f'AT{r}'] = f'=IFERROR(AVERAGE({",".join(f"{a[y]}{r}" for y in range(2016, 2026))}),"n/m")'
        ws[f'AU{r}'] = f'=IFERROR(AVERAGE({",".join(f"{a[y]}{r}" for y in range(2012, 2026))}),"n/m")'
        for c in CAGR: ws[f'{c}{r}'].number_format = NF['pct']
    t = LABEL_SOURCES.get(row.key)
    if t: comment(ws.cell(r, 2), t)

SRC = 'Source: Loar earnings releases (Form 8-K Ex. 99.1) — '
LABEL_SOURCES = {
    'com_oem': SRC + 'Table 5 "Sales by End-Market" (from Q2-24; Q1-24 / 2023 quarters from comparatives); FY2022–23 from the IPO prospectus (424B4) revenue note.',
    'a_org': SRC + 'narrative "Organically, net sales increased …" (quarter and year-to-date); FY2023 from the prospectus MD&A.',
    'is_rev': ('Source: release Table 2 (Q1-24+); FY2022–23 audited statements, Q1/22–Q4/23 quarterly results and FY2012–21 EBITDA reconciliation '
               'from the IPO prospectus (424B4, 26-Apr-2024). FY2017 = predecessor (1-Jan–1-Oct-2017) + successor (2-Oct–31-Dec-2017) as presented.'),
    'd_dep': SRC + 'Table 4 EBITDA reconciliation; prospectus reconciliations before 2024.',
    'r_adj_pub': SRC + 'Table 4 "Reconciliation of Net income to EBITDA and Adjusted EBITDA"; prospectus annual (FY2012–23) and quarterly (Q1/22–Q4/23) reconciliations.',
    'e_pub': SRC + 'Table 6 (current definition, from the Q1-26 release: Q1-25 / Q2-25 restated as comparatives).',
    'e_pub_old': SRC + 'Table 6 as originally reported (Q2-24 – Q4-25; FY2024, FY2025).',
    'e_tax': SRC + 'Table 6 tax adjustment ("tax effect of the adjustments at the applicable effective tax rate", excluding the effect of transaction expenses and SBC).',
    'cf_ni': 'Source: release Table 3 / prospectus cash-flow statements (year-to-date; discrete quarters derived; Q4 = FY − 9M). 2022 quarters: prospectus "Other data".',
    'bs_cash': 'Source: release Table 1 (quarter ends from Q1-24); prospectus balance sheets at 31-Dec-2022 / 2023.',
    'ppa_cons': 'Source: prospectus Note 2 (SCHROTH); FY2025 10-K Note 2 (DAC / CAV, AAI, Beadlight); Q2-26 10-Q Note 2 (LMB, Harper).',
    'sh_basic': SRC + 'Table 6 / Table 2. Pre-IPO (to Q1-24) Loar Holdings, LLC reported 204 common units — per-share data not meaningful.',
}

DEFS = {}
def defs_setup():
    q = lambda y, n: QC[(y, n)].c
    DEFS['r_def'] = {
        'C': ('FY2012–21: prospectus reconciliation (adds inventory step-up, other (income) loss, transaction, SBC, integration, COVID 2020–21, management fees to 2017)',
              'Loss on extinguishment of debt (2016–17), FX gain (2017) and insurance gain (2015) sit between net income and operating income.'),
        q(2022, 1): ('2022–23: prospectus quarterly / annual reconciliations (COVID add-back to Q4-22)', None),
        q(2024, 1): ('Release Table 4: + refinancing costs below operating income (Q2-24, Q4-24); other income / contingent consideration deducted', None),
        q(2026, 2): ('Q2-26: + other expense (increase in Harper Engineering contingent consideration)', None),
    }
    DEFS['e_def'] = {
        'C': ('No Adjusted EPS before the IPO (Loar Holdings, LLC common units)', None),
        q(2024, 2): ('Q2-24 – Q4-25 as published: prior definition (refinancing + EBITDA adjustments, tax-effected; amortization not added back)',
                     'Shown here on the current definition (model); the prior-definition figures and their checks are in the memo rows below.'),
        q(2025, 1): ('Q1-25 / Q2-25 restated on the current definition in the Q1-26 / Q2-26 releases (checks = 0)', None),
        q(2026, 1): ('Q1-26+: current definition — + amortization of acquired intangible assets (tax-effected); FY26 guidance on this basis', None),
    }

def write_defs(ws, row):
    for c, (txt, cm) in DEFS[row.key].items():
        cell = ws[f'{c}{row.r}']; cell.value = txt
        cell.font = Font(name='Calibri', sz=10, b=True, i=True, color='FFC00000')
        if cm: comment(cell, cm)

def write_scen(ws):
    srcrow = {'hdr': 9, 'out_hdr': 11, 'pt_hdr': 15, 'out': 12, 'pt': 16, 'lev': 28, 'blk': 30, 'yr': 31}
    for spec_ in SCEN_ROWS:
        r, kind = spec_[0], spec_[1]
        for ci in range(CI(SL), CI(SNOTE) + 1):
            ws.cell(r, ci)._style = copy.copy(TM.cell(srcrow[kind], ci + TSHIFT)._style)
        if kind in ('hdr', 'out_hdr', 'pt_hdr'):
            _, _, a, b, c_, d_, e = spec_
            for col, v in zip([SL, SBULL, SBASE, SBEAR, SNOTE], [a, b, c_, d_, e]):
                if v is not None: ws[f'{col}{r}'] = v
        elif kind == 'out':
            _, _, lab, hi, mid, lo, note, nf = spec_
            ws[f'{SL}{r}'] = lab
            for col, v in zip([SBULL, SBASE, SBEAR], [hi, mid, lo]):
                ws[f'{col}{r}'] = v; ws[f'{col}{r}'].number_format = NF[nf]; set_color(ws[f'{col}{r}'], BLUE)
            ws[f'{SNOTE}{r}'] = note
        elif kind == 'pt':
            _, _, lab, _, val, _, note, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBASE}{r}'] = val; ws[f'{SBASE}{r}'].number_format = NF[nf]; set_color(ws[f'{SBASE}{r}'], BLUE); ws[f'{SNOTE}{r}'] = note
        elif kind == 'lev':
            _, _, lab, a, b, c_, note = spec_
            ws[f'{SL}{r}'] = lab
            for col, v in zip([SBULL, SBASE, SBEAR], [a, b, c_]):
                ws[f'{col}{r}'] = v; ws[f'{col}{r}'].number_format = NF['x']
            ws[f'{SNOTE}{r}'] = note
        elif kind == 'blk':
            _, _, lab, db, mid, de, note, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBULL}{r}'] = db; ws[f'{SBASE}{r}'] = mid; ws[f'{SBEAR}{r}'] = de; ws[f'{SNOTE}{r}'] = note
            for col in (SBULL, SBEAR):
                ws[f'{col}{r}'].number_format = '\\+#,##0;\\-#,##0' if nf == 'm' else '\\+0.0%;\\-0.0%'
        elif kind == 'yr':
            _, _, lab, hdr, base, clamp, _, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBASE}{r}'] = base
            if clamp:
                ws[f'{SBULL}{r}'] = f'=MAX(0,{SBASE}{r}+${SBULL}${hdr})'; ws[f'{SBEAR}{r}'] = f'=MAX(0,{SBASE}{r}+${SBEAR}${hdr})'
            else:
                ws[f'{SBULL}{r}'] = f'={SBASE}{r}+${SBULL}${hdr}'; ws[f'{SBEAR}{r}'] = f'={SBASE}{r}+${SBEAR}${hdr}'
            for col in (SBULL, SBASE, SBEAR): ws[f'{col}{r}'].number_format = NF[nf]

def write_header(ws):
    for r in range(1, 8):
        clone_row_style(ws, r, r, 2, CI(LASTC))
    ws[f'{NOTE}1'] = 'SCENARIO SWITCH ▼ (Bull / Base / Bear) — drives all scenario-linked forecast drivers'
    ws[f'{NOTE}2'] = 'Base'
    comment(ws[f'{NOTE}2'], 'Scenario toggle. Choose Bull / Base / Bear. Drives: FY26 sales and Adjusted EBITDA outlook end-points (high / mid / low), '
            '2027E–2030E end-market OEM / aftermarket growth, Adjusted EBITDA margin, M&A spend and the leverage floor for 2027E+ buybacks '
            f'(scenario table {SL}:{SNOTE}).')
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type='list', formula1='"Bull,Base,Bear"', allow_blank=False); ws.add_data_validation(dv); dv.add(f'{NOTE}2')
    ws['B2'] = 'Loar Holdings Inc. (NYSE: LOAR) — 3-Statement Financial Model'
    ws['B3'] = ('US GAAP as reported · USD millions (per-share data in USD; shares in millions) · Source: Loar Forms 10-K / 10-Q, Form 8-K earnings releases (Ex. 99.1), '
                'deal 8-Ks and the IPO prospectus (424B4, 26-Apr-2024), SEC EDGAR CIK 0002000178 (investor-relations releases at ir.loargroup.com are the same documents as filed on EDGAR)')
    ws['B4'] = ('Basis notes: calendar fiscal year; single reportable segment; all periods as originally reported. FY2012–21: prospectus net income → Adjusted EBITDA '
                'reconciliation only (FY2017 = predecessor Jan–Oct + successor Oct–Dec: acquisition of Loar Group Inc. by Loar Holdings, LLC, 2-Oct-2017). Loar Holdings, LLC to the IPO (29-Apr-2024; follow-on Dec-2024) — '
                'per-share data from Q2-24. Acquisitions: SCHROTH Jul-22, DAC Jul-23, CAV Sep-23, Applied Avionics Aug-24, Beadlight Jul-25, LMB Dec-25, Harper Engineering Jan-26. '
                'Adjusted EPS on the current definition (adds back amortization of acquired intangibles, from Q1-26) for all periods. Q3/26E–Q4/26E calibrated to FY26 sales and Adjusted EBITDA guidance.')
    ws[f'{CAGR[0]}5'] = 'Forecast'
    for c in CAGR[1:]: ws[f'{c}5'] = 'Historical'
    ws['B6'] = '(USDm)'
    for col in COLS: ws[f'{col.c}6'] = col.label
    ws[f'{NOTE}6'] = 'Modelling Notes'
    for c, a, b in zip(CAGR, ['5Y CAGR', '5Y CAGR', '10Y CAGR', '13Y CAGR'], ["25-'30", "20-'25", "15-'25", "12-'25"]):
        ws[f'{c}6'] = a; ws[f'{c}7'] = b

def set_columns(ws):
    from openpyxl.worksheet.dimensions import ColumnDimension
    for k in list(ws.column_dimensions.keys()): del ws.column_dimensions[k]
    ws.column_dimensions['A'].width = TM.column_dimensions['A'].width
    ws.column_dimensions['B'].width = TM.column_dimensions['B'].width
    wid = {'A': 13.0, 'Q': 9.5703125, 'QE': 11.42578125, 'E26': 14.140625, 'AE': 13.0}
    for col in COLS:
        cd = ws.column_dimensions[col.c]; cd.width = wid[col.kind]
        if col.q:
            cd.outlineLevel = 1; cd.hidden = True
    for ci in range(CI(NOTE), CI(LASTC) + 1):
        ws.column_dimensions[L(ci)].width = TM.column_dimensions[L(ci + TSHIFT)].width or 13.0
    ws.sheet_format.outlineLevelCol = 1

def main():
    import openpyxl
    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb['Model']
    blank_style = copy.copy(ws['A1']._style)
    for rowc in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for c in rowc:
            c.value = None; c.comment = None; c._style = copy.copy(blank_style)
    for k in list(ws.row_dimensions.keys()):
        rd = ws.row_dimensions[k]; rd.outlineLevel = 0; rd.hidden = False
    ws.data_validations.dataValidation = []
    scen_layout(V0)
    build_spec(V0)
    n1 = len(fw.SPEC)
    build_spec2(V0)
    for i, row in enumerate(fw.SPEC): row.spec2 = i >= n1
    last = number_rows(8)
    for r in range(1, last + 1): ws.row_dimensions[r].height = 15.0
    ws.row_dimensions[2].height = 21.0
    defs_setup()
    write_header(ws)
    for row in fw.SPEC: write_row(ws, row)
    write_scen(ws)
    set_columns(ws)
    ws.freeze_panes = 'C8'
    import other_sheets as osh
    osh.do_all(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    osh.inject_charts(OUT)
    print('rows', fw.SPEC[-1].r, 'saved', OUT)
    return wb

RECALC = os.environ.get('XLSX_RECALC', '/mnt/skills/public/xlsx/scripts/recalc.py')
def recalc(path, timeout=600):
    import subprocess
    out = subprocess.run(['python3', RECALC, path, str(timeout)], capture_output=True, text=True, cwd=os.path.dirname(RECALC))
    try:
        return json.loads(out.stdout)
    except Exception:
        return {'raw': out.stdout[-2000:], 'err': out.stderr[-2000:]}

def scenario_snapshots():
    import openpyxl
    import other_sheets as osh
    snaps = {}
    for scen in ('Bull', 'Bear'):
        p = os.path.join(BASE, f'LOAR_{scen}.xlsx')
        wb = openpyxl.load_workbook(OUT); wb['Model'][f'{NOTE}2'] = scen; wb.save(p)
        res = recalc(p); print(scen, res.get('status'), res.get('total_errors'))
        v = openpyxl.load_workbook(p, data_only=True)['Bull-Base-Bear']
        snaps[scen] = {(r, c): v[f'{c}{r}'].value for r in range(9, 48) for c in 'CDEFGHIJKL' if v[f'{c}{r}'].value is not None}
        os.remove(p)
    wb = openpyxl.load_workbook(OUT)
    osh.snapshot_values(wb, snaps['Bull'], snaps['Bear'])
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    osh.inject_charts(OUT)

if __name__ == '__main__':
    main()
    if '--snap' in sys.argv: scenario_snapshots()
