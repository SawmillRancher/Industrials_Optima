"""Framework: MLM-template layout re-based on LOAR's columns, style cloning, row registry and per-column formula context.

Columns (Model sheet):  C-L FY2012-2021 (annual; prospectus P&L / EBITDA-bridge lines only) | M-AF Q1/22-Q4/25 with FY
subtotals | AG-AH Q1/26-Q2/26 | AI-AJ Q3/26E-Q4/26E | AK 2026E | AL-AO 2027E-2030E | AP modelling notes | AR-AU CAGRs |
AX-BB scenario input table. Template columns right of the forecast block keep their order (template letter − 42).
"""
import copy, json, os
import openpyxl
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI
from openpyxl.styles import Font
from openpyxl.comments import Comment
import normalize

BASE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(BASE, 'MLM_Model.xlsx')

class Col:
    def __init__(self, c, label, kind, year, q=None):
        self.c, self.label, self.kind, self.year, self.q = c, label, kind, year, q
        self.prior = None; self.prevq = None
    def __repr__(self): return f'{self.c}:{self.label}'

COLS, ANN, QC = [], {}, {}
_i = CI('C')
for y in range(2012, 2022):
    col = Col(L(_i), y, 'A', y); COLS.append(col); ANN[y] = col; _i += 1
for y in range(2022, 2027):
    yy = str(y)[2:]
    for q in (1, 2, 3, 4, None):
        if q:
            kind, lab = ('QE', f'Q{q}/{yy}E') if (y == 2026 and q >= 3) else ('Q', f'Q{q}/{yy}')
            col = Col(L(_i), lab, kind, y, q); QC[(y, q)] = col
        else:
            kind, lab = ('E26', '2026E') if y == 2026 else ('A', y)
            col = Col(L(_i), lab, kind, y); ANN[y] = col
        COLS.append(col); _i += 1
for y in range(2027, 2031):
    col = Col(L(_i), f'{y}E', 'AE', y); COLS.append(col); ANN[y] = col; _i += 1
assert ANN[2025].c == 'AF' and QC[(2026, 2)].c == 'AH' and ANN[2026].c == 'AK' and ANN[2030].c == 'AO', (ANN[2025], ANN[2026], ANN[2030])
for col in COLS:
    if col.q is None:
        if col.year - 1 in ANN: col.prior = ANN[col.year - 1].c
    else:
        if (col.year - 1, col.q) in QC: col.prior = QC[(col.year - 1, col.q)].c
        if col.q > 1: col.prevq = QC[(col.year, col.q - 1)].c
        elif (col.year - 1, 4) in QC: col.prevq = QC[(col.year - 1, 4)].c
COLMAP = {c.c: c for c in COLS}
HISTK = ('A', 'Q')
FCK = ('QE', 'E26', 'AE')

# key column letters (used by the row specs)
Q1, Q2, Q3, Q4, E26 = 'AG', 'AH', 'AI', 'AJ', 'AK'          # 2026 quarters / 2026E
A25, A24 = 'AF', 'AA'
NOTE = 'AP'
SW = '$AP$2'                                                # scenario switch
SL, SBULL, SBASE, SBEAR, SNOTE = 'AX', 'AY', 'AZ', 'BA', 'BB'  # scenario table columns
CAGR = ('AR', 'AS', 'AT', 'AU')
LASTC = 'BB'
TSHIFT = 42                                                 # template column index − TSHIFT = LOAR column index (right block)

def pkey(col):
    if col.kind == 'A': return str(col.year)
    if col.kind == 'Q': return f'Q{col.q}-{str(col.year)[2:]}'
    return None

# ---------------------------------------------------------------- data
DATA = normalize.run()
MAN = normalize.MAN
def num(x):
    if isinstance(x, bool) or x is None: return None
    return x if isinstance(x, (int, float)) else None
def xget(per, key):
    return (DATA.get(per, {}).get('x') or {}).get(key)

# ---------------------------------------------------------------- styles (cloned from the MLM template)
_tmpl_wb = openpyxl.load_workbook(TEMPLATE)
TM = _tmpl_wb['Model']
ARCH = {'section': 8, 'blank': 9, 'block': 10, 'line_b': 162, 'line': 161, 'growth': 237, 'pct': 212, 'memo': 180, 'check': 179,
        'driver': 153, 'total': 166, 'text': 215, 'sub': 217, 'eps': 187, 'eps_total': 188, 'eps_memo': 191, 'sched': 378, 'memo_calc': 181}
NF = {'m': '#,##0;\\(#,##0\\);\\-', 'm1': '#,##0.0;\\(#,##0.0\\);\\-', 'pct': '0.0%;\\(0.0%\\);\\-', 'ps': '0.00;\\(0.00\\);\\-',
      'x': '0.0\\x', 'x2': '0.00\\x;\\(0.00"x)";\\-', 'd': '0\\d;\\(0")d";\\-', 'gen': 'General', 'yrs': '0" yrs"',
      'm2': '#,##0.00;\\(#,##0.00\\);\\-', 'pct2': '0.00%;\\(0.00%\\);\\-'}
BLUE, BLACK, GRAY = 'FF0000FF', 'FF000000', 'FFA6A6A6'
SRC_KIND = {'A': 'BV', 'Q': 'BX', 'QE': 'BY', 'E26': 'CA', 'AE': 'CB'}

def tcol(ci):
    """template column letter for LOAR column index ci"""
    letter = L(ci)
    if letter == 'B': return 'B'
    if letter in COLMAP: return SRC_KIND[COLMAP[letter].kind]
    return L(ci + TSHIFT)

def clone_row_style(ws, src_row, dst_row, c0=2, c1=CI(LASTC)):
    for ci in range(c0, c1 + 1):
        ws.cell(dst_row, ci)._style = copy.copy(TM[f'{tcol(ci)}{src_row}']._style)

def set_color(cell, rgb, bold=None):
    f = copy.copy(cell.font)
    cell.font = Font(name=f.name, sz=f.sz, b=f.b if bold is None else bold, i=f.i, u=f.u, color=rgb)

def comment(cell, text, w=420, h=160):
    c = Comment(text, 'LOAR model'); c.width, c.height = w, h
    cell.comment = c

# ---------------------------------------------------------------- registry
class Row:
    def __init__(self, key, label, arch, nf, kw):
        self.key, self.label, self.arch, self.nf, self.kw = key, label, arch, nf, kw
        self.r = None; self.spec2 = False
SPEC, R = [], {}
def add(key, label='', arch='line', nf='m', **kw):
    if key in R: raise KeyError('dup ' + key)
    SPEC.append(Row(key, label, arch, nf, kw)); R[key] = None
    return key
_bn = [0]
def blank():
    _bn[0] += 1
    return add(f'_b{_bn[0]}', '', 'blank', 'gen')
def number_rows(start):
    r = start
    for row in SPEC:
        R[row.key] = r; row.r = r; r += 1
    return r

class _Ctx: col = None
CTX = _Ctx()
def V(key, c=None): return f'{c or CTX.col.c}{R[key]}'
def P(key): return f'{CTX.col.prior}{R[key]}'
def PQ(key): return f'{CTX.col.prevq}{R[key]}'
def Q(key, y, q): return f'{QC[(y, q)].c}{R[key]}'
def Y(key, y): return f'{ANN[y].c}{R[key]}'
def S26(key): return f'=SUM({Q1}{R[key]},{Q2}{R[key]},{Q3}{R[key]},{Q4}{R[key]})'
def H1(key, y=2026): return f'({Q(key, y, 1)}+{Q(key, y, 2)})'
def SUMQ(key):
    """annual column = sum of its four quarters (historical annual columns 2022+)"""
    y = CTX.col.year
    return f'=SUM({QC[(y, 1)].c}{R[key]}:{QC[(y, 4)].c}{R[key]})'
