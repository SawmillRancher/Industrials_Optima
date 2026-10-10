import openpyxl, re, sys
from openpyxl.utils import column_index_from_string as CI, get_column_letter as L
sys.setrecursionlimit(100000)
wb = openpyxl.load_workbook(sys.argv[1] if len(sys.argv) > 1 else '../LOAR_Model.xlsx'); ws = wb['Model']
G = {}
ref = re.compile(r"(?<![A-Za-z!])\$?([A-Z]{1,3})\$?(\d+)(?::\$?([A-Z]{1,3})\$?(\d+))?")
for row in ws.iter_rows():
    for c in row:
        v = c.value
        if isinstance(v, str) and v.startswith('='):
            deps = set()
            s = re.sub(r'"[^"]*"', '', v)
            for m in ref.finditer(s):
                a, r1, b, r2 = m.groups()
                if b:
                    for ci in range(CI(a), CI(b) + 1):
                        for r in range(int(r1), int(r2) + 1): deps.add(f'{L(ci)}{r}')
                else: deps.add(f'{a}{r1}')
            G[c.coordinate] = deps
color = {}
cyc = []
def dfs(u, stack):
    color[u] = 1; stack.append(u)
    for w in G.get(u, ()):
        if color.get(w) == 1:
            cyc.append(stack[stack.index(w):] + [w]); 
            if len(cyc) > 5: return
        elif color.get(w) is None and w in G:
            dfs(w, stack)
    stack.pop(); color[u] = 2
for u in list(G):
    if color.get(u) is None: dfs(u, [])
    if len(cyc) > 5: break
lab = lambda a: ws.cell(int(re.sub('[A-Z]', '', a)), 2).value
for c in cyc[:5]: print([f'{a}:{str(lab(a)).strip()[:30]}' for a in c][:12])
print('cycles', len(cyc))
