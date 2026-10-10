"""List formula errors and non-zero check rows in a recalculated LOAR model."""
import sys, openpyxl
from collections import defaultdict
from openpyxl.utils import get_column_letter as L
p = sys.argv[1]
wb = openpyxl.load_workbook(p, data_only=True)
ws = wb['Model']
hdr = {c.column_letter: c.value for c in ws[6]}
errs = defaultdict(list)
for name in wb.sheetnames:
    for row in wb[name].iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith('#'):
                errs[name].append(c.coordinate)
for k, v in errs.items(): print('ERR', k, len(v), v[:40])
for r in range(8, ws.max_row + 1):
    lab = ws.cell(r, 2).value
    if not (isinstance(lab, str) and ('Check' in lab or 'tie-out' in lab)): continue
    bad = []
    for c in range(3, 42):
        v = ws.cell(r, c).value
        if isinstance(v, (int, float)) and abs(v) > 1e-9:
            bad.append(f'{hdr.get(L(c))}:{round(v, 3)}')
    if bad: print(r, lab.strip()[:70], len(bad), bad[:25])
