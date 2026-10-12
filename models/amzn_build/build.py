"""Build AMZN_Model.xlsx from the MLM template (same layout, formatting and process as the prior models)."""
import sys, copy, os, json
from openpyxl.styles import Font
from fw import *
import fw
from spec import build_spec, scen_layout, SCEN_ROWS, SC
from spec2 import build_spec2

OUT = os.path.join(os.path.dirname(BASE), 'AMZN_Model.xlsx')
G = MAN['guidance_q3_26']
DEBT = MAN['debt']

# ====================================================================== inputs (V0)
def wc_inputs():
    q = DATA['2025']['x']
    rev, cogs = q['rev'], q['cogs']
    return {'dso': round(q['bs_ar'] / rev * 365), 'dio': round(q['bs_inv'] / cogs * 365), 'dpo': round(q['bs_ap'] / cogs * 365),
            'accr': round(q['bs_accr_x'] / rev, 3), 'unearned': round(q['bs_unearned'] / rev, 3)}
WC = wc_inputs()
WC['note'] = (f'Forecast inputs at fiscal-year-end (31-Dec-2025) levels on FY2025 net sales / cost of sales (receivables {WC["dso"]}d, inventory {WC["dio"]}d, '
              f'payables {WC["dpo"]}d, accrued excl. debt {WC["accr"]:.1%} and unearned revenue {WC["unearned"]:.1%} of net sales); year ends are compared with year ends '
              '(Q4 holiday payables make mid-year days unrepresentative).')
MAT = DEBT['maturities_existing_dec25']; NEW = DEBT['new_2026_due_2028_2030']
POST = DEBT['post_q2_issues']
ISS26 = POST['usd_jul'] + POST['gbp_sep_gbp'] * POST['gbp_fx_assumed']
MATY = {2027: MAT['2027'], 2028: MAT['2028'] + NEW['2028'], 2029: MAT['2029'] + NEW['2029'], 2030: MAT['2030'] + NEW['2030'], 2031: 10000.0}
X25 = DATA['Q2-26']['x']

V0 = {
    'out_hdr': 'Q3-26 guidance — Q2-26 release (8-K, 30-Jul-2026)',
    'out_note': ('Net sales $197.0–202.0bn (+9–12% y/y; ~80 bp FX headwind; Prime Day in Q2 in 2026 vs Q3 in 2025 — ex Prime Day growth ~400 bp higher); '
                 'operating income $22.5–26.5bn (assumes no energy-derivative remeasurement, acquisitions, restructurings or legal settlements). Amazon guides one quarter ahead.'),
    'outlook': [
        ('out_rev', '  Q3-26 net sales ($m)', tuple(G['net_sales']), 'm', 'Guided $197.0–202.0bn. Drives the Q3/26E net-sales calibration (retail product groups).'),
        ('out_oi', '  Q3-26 operating income ($m)', tuple(G['operating_income']), 'm', 'Guided $22.5–26.5bn (Q3-25 $17.4bn incl. $4.3bn special charges). Drives the segment-margin calibration.'),
    ],
    'points': [
        ('pt_prime', '  Prime Day timing shift on Q3-26 total growth (pp)', -0.04, 'pct', 'Guidance: "nearly 400 basis points" (8-K 30-Jul-2026).'),
        ('pt_capex_q', '  Gross capex per quarter, Q3/Q4-26E ($m)', 54208.0, 'm', 'Q2-26 purchases of P&E $54.2bn (Q1 $44.2bn); held at the Q2 run-rate (assumption).'),
        ('pt_proc', '  Proceeds from P&E sales and incentives, % of gross capex', 0.02, 'pct', '1H-26 2.1%; FY2025 2.7%.'),
        ('pt_othop', '  Other operating expense (income), net per quarter ($m)', 300.0, 'm', '1H-26 $537m.'),
        ('pt_yield', '  Interest income yield on cash & marketable securities (annual)', 0.036, 'pct', 'Q2-26 $1,295m on $143.1bn opening cash and securities (3.6%).'),
        ('pt_int_q3', '  Q3/26E interest expense ($m)', 1750.0, 'm', 'Q2-26 $1,314m on $122bn opening debt; + $25.0bn July notes (4.60–6.25%) from 9-Jul (assumption).'),
        ('pt_int_q4', '  Q4/26E interest expense ($m)', 1950.0, 'm', 'Full quarter of the July notes + £4.25bn September notes (5.20–6.65%) (assumption).'),
        ('pt_etr', '  Q3/Q4-26E effective tax rate', 0.19, 'pct', 'FY2025 19.6%; FY2024 13.5% (assumption; no investment marks forecast).'),
        ('pt_iss_q', '  Shares issued under equity plans per quarter (m)', 30.0, 'm1', 'Shares outstanding +52m in 1H-26; +138m in 2025.'),
        ('pt_def', '  2H-26E deferred income taxes ($m)', 4000.0, 'm', '1H-26 $30.5bn incl. deferred tax on Anthropic gains; underlying ~$2bn a quarter (assumption).'),
        ('pt_lease', '  2H-26E principal repayments of finance leases & financing obligations ($m)', -1000.0, 'm', '1H-26 −$1,037m.'),
        ('pt_fladd', '  2H-26E P&E acquired under finance leases ($m; non-cash)', 1500.0, 'm', '1H-26 $2,128m; FY2025 $2,911m.'),
    ],
    'lev': (0.25, 0.0, -0.25),
    'lev_note': ('Amazon has not repurchased shares since 2022 ($10bn authorization of Mar-2022, $6bn used) and has financed the AI build-out with $66bn of notes in 1H-26. '
                 'Target = net debt / adjusted EBITDA up to which surplus cash is returned via buybacks (model assumption; no company target): Base 0.0x (free cash flow beyond '
                 'net-debt neutrality returned), Bull 0.25x, Bear −0.25x (keeps a net-cash cushion). Net debt was −$54bn at YE2025 and +$9bn at 30-Jun-2026.'),
    'blocks': [
        ('ONLINE', 'ONLINE (Online stores y/y)', 0.02, -0.03, [0.085, 0.08, 0.075, 0.07],
         'FY2025 +9.0%; 1H-26 +13.3% (Q1 +11.9%, Q2 incl. Prime Day). Faster delivery, everyday essentials; maturing category.', 'pct', False),
        ('PHYS', 'PHYS (Physical stores y/y)', 0.02, -0.03, [0.05, 0.05, 0.045, 0.04], 'FY2025 +6.3%; 1H-26 +4.6% (Whole Foods, Amazon Fresh).', 'pct', False),
        ('TP', 'TP (Third-party seller services y/y)', 0.02, -0.03, [0.11, 0.10, 0.095, 0.09], 'FY2025 +10.3%; 1H-26 +16.3% (fulfilment fees, mix shift to 3P).', 'pct', False),
        ('SUBS', 'SUBS (Subscription services y/y)', 0.02, -0.03, [0.11, 0.10, 0.095, 0.09], 'FY2025 +11.8%; 1H-26 +14.6% (Prime, video, music).', 'pct', False),
        ('ADS', 'ADS (Advertising services y/y)', 0.03, -0.05, [0.19, 0.17, 0.15, 0.13], 'FY2025 +22.1%; Q2-26 +26%; Prime Video ads, sponsored products.', 'pct', False),
        ('AWS_G', 'AWS_G (AWS y/y)', 0.05, -0.08, [0.28, 0.24, 0.20, 0.17],
         'FY2025 +19.7%; Q2-26 +36.8% (fastest in 18 quarters; AI and chips run-rates >$25bn each). Capacity-constrained growth as capex lands.', 'pct', False),
        ('OTHER', 'OTHER (Other net sales y/y)', 0.02, -0.03, [0.08, 0.07, 0.06, 0.06], 'FY2025 +9.4% (co-branded credit cards, licensing, shipping services).', 'pct', False),
        ('NA_M', 'NA_M (North America segment margin)', 0.01, -0.015, [0.085, 0.09, 0.095, 0.10],
         'FY2025 6.9% (7.5% ex the $2.5bn FTC charge); 1H-26 7.9%. Regionalised fulfilment, advertising mix vs quick-commerce and Leo investment.', 'pct', False),
        ('INTL_M', 'INTL_M (International segment margin)', 0.01, -0.015, [0.045, 0.05, 0.055, 0.06], 'FY2025 2.9%; 1H-26 3.9%; scale in emerging markets offsets price investment.', 'pct', False),
        ('AWS_M', 'AWS_M (AWS segment margin)', 0.02, -0.03, [0.36, 0.355, 0.35, 0.35],
         'FY2025 35.4%; 1H-26 38.6% (incl. $0.6bn energy-contract gains). Rising depreciation from the AI build-out caps the margin.', 'pct', False),
        ('CAPEX', 'CAPEX (gross capex % of net sales)', -0.01, 0.015, [0.23, 0.21, 0.19, 0.17],
         'FY2025 18.4%; 1H-26 25.8%. AI / data-centre build-out peaking in 2026–27, then normalising (assumption).', 'pct', False),
        ('MNA', 'M&A_spend (acquisitions & strategic investments, $m)', 3000.0, -3000.0, [5000.0, 5000.0, 5000.0, 5000.0],
         'Cash-flow line incl. non-marketable investments (Anthropic): FY2024 $7.1bn, FY2025 and 1H-26 per the cash-flow statement. Base $5bn p.a. (assumption).', 'm', True),
    ],
    'dep_r': {2027: 0.135, 2028: 0.135, 2029: 0.135, 2030: 0.135},
    'dep_note': 'Q2-26 annualised 14.0% of opening PP&E; FY2025 16.6% of opening (13.7% of average). 2027E+ 13.5% on opening + ½ net additions (servers ~6-year lives from 2025).',
    'oda_pct': 0.032, 'sbc_pct': 0.026, 'etr': 0.19, 'othop': 1200.0,
    'cost': {'ful': {y: 0.145 for y in range(2027, 2031)}, 'tech': {2027: 0.170, 2028: 0.172, 2029: 0.172, 2030: 0.172},
             'mkt': {y: 0.058 for y in range(2027, 2031)}, 'ga': {y: 0.014 for y in range(2027, 2031)}},
    'def_ae': 6000.0, 'lease_ae': -2000.0, 'proc_pct': 0.02, 'iss_ae': 120.0, 'dil': 135.0, 'min_cash': 50000.0, 'fladd_ae': 3000.0, 'flint': 600.0,
    'iss_notes': {2026: ISS26, 2027: MATY[2027], 2028: MATY[2028], 2029: MATY[2029], 2030: MATY[2030]},
    'repay': {2026: 0.0, 2027: MATY[2027], 2028: MATY[2028], 2029: MATY[2029], 2030: MATY[2030]},
    'mat': {2026: MATY[2027], 2027: MATY[2028], 2028: MATY[2029], 2029: MATY[2030], 2030: MATY[2031]},
    'repay_note': (f'Maturities: existing notes per the FY2025 10-K schedule (XBRL; 2026 ${MAT["2026"]:,.0f}m — repaid in 1H-26; 2027 ${MAT["2027"]:,.0f}m; 2028 ${MAT["2028"]:,.0f}m; '
                   f'2029 ${MAT["2029"]:,.0f}m; 2030 ${MAT["2030"]:,.0f}m) + 2026 issues due 2028–30 (${NEW["2028"]:,.0f}m / ${NEW["2029"]:,.0f}m / ${NEW["2030"]:,.0f}m; '
                   + NEW['src'] + '). 2031 maturities $10.0bn assumed. 2027E+ maturities refinanced with new notes (input row above).'),
    'r_notes': {2027: 0.045, 2028: 0.046, 2029: 0.047, 2030: 0.047},
    'r_notes_note': ('Q2-26 interest expense $1,314m ≈ 4.3% annualised on opening debt (incl. lease interest); 2026 issues priced at 3.85–6.25% (USD), 2.80–4.85% (EUR), '
                     '0.84–2.08% (CHF), 3.40–5.00% (CAD), 5.20–6.65% (GBP). Weighted rate drifting up (assumption).'),
    'r_cash': 0.033, 'r_cash_note': 'Q2-26 3.6% annualised; drifting down with policy rates (assumption).',
    'wc': WC,
    'px': MAN['price']['px'],
    'px_note': ('VALUATION — share price input $260.00 (provided by the user, 12-Oct-2026); update to market'),
    'mna_note': 'Scenario lever M&A_spend (Base $5bn p.a.; half assumed to buy operating businesses at 3.0x sales, half strategic investments with no revenue).',
    'rev_guid_note': 'Company Q3-26 net sales guidance $197.0–202.0bn (8-K 30-Jul-2026). Drives the net-sales calibration.',
    'oi_guid_note': 'Company Q3-26 operating income guidance $22.5–26.5bn (8-K 30-Jul-2026). Drives the segment-margin calibration.',
}

# ====================================================================== compute engine
def hist_value(row, col):
    kw = row.kw; path = kw.get('data')
    if not path: return None
    v = xget(pkey(col), path)
    return v if kw.get('textdata') else num(v)

def first(kw, keys):
    for k in keys:
        if k in kw: return k
    return None

def compute(row, col):
    kw = row.kw; k = col.kind
    CTX.col = col
    if kw.get('ann') and col.q: return None, False
    if k in HISTK:
        v = hist_value(row, col)
        if v is not None: return v, True
        if 'hist_in' in kw: return kw['hist_in'](), True
        for key in ('hist', 'all'):
            if key in kw:
                return kw[key](), False
        return None, False
    if k == 'QE':
        if row.spec2: return None, False
        if 'qe_in' in kw:
            v = kw['qe_in']; return (v[col.q] if isinstance(v, dict) else v), True
        if 'fc_in' in kw: return kw['fc_in'], True
        key = first(kw, ('qe', 'fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    if k == 'E26':
        if 'e26_in' in kw: return kw['e26_in'], True
        if 'e26' in kw: return kw['e26'](), False
        if 'fc_in' in kw: return kw['fc_in'], True
        if kw.get('avg') and ('qe' in kw or 'qe_in' in kw): return f'=AVERAGE({Q1}{row.r}:{Q4E}{row.r})', False
        if kw.get('flow') and not row.spec2 and any(x in kw for x in ('qe', 'qe_in', 'fc')): return S26(row.key), False
        key = first(kw, ('fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    if k == 'AE':
        if 'ae_in' in kw:
            v = kw['ae_in']; return (v[col.year] if isinstance(v, dict) else v), True
        if 'fc_in' in kw: return kw['fc_in'], True
        key = first(kw, ('ae', 'fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    return None, False

# ====================================================================== writers
CAGR_G = {'is_rev', 'g_rev', 'pg_online', 'pg_physical', 'pg_tp', 'pg_subs', 'pg_ads', 'pg_awsg', 'pg_other', 'seg_na_rev', 'seg_intl_rev', 'seg_aws_rev',
          'seg_na_oi', 'seg_intl_oi', 'seg_aws_oi', 'o_adj', 'g_adj', 'r_adj', 'r_ebitda', 'is_oi', 'd_dda', 'd_dep', 'e_adjni', 'e_eps', 'cf_cfo', 'fcf',
          'bs_ta', 'bs_eq', 'sh_dil', 'is_ni', 'eps_dil', 'cf_capex', 'd_sbc', 'is_ful', 'is_tech'}
CAGR_A = {'seg_na_m', 'seg_intl_m', 'seg_aws_m', 'g_adjm', 'o_adj_m', 'r_adj_m', 'is_gm', 'is_oi_m', 'r_ebit_m', 'gm_rev', 'gm_aws', 'gm_ads', 'gm_gm', 'gm_ebitda_m',
          'gm_op_m', 'gm_net_m', 'gm_adjop_m', 'fcf_m', 'fcf_conv', 'capex_pct', 'ra_roic', 'rn_ronta', 'roe', 'lv_x', 'd_ful', 'd_tech', 'd_mkt', 'd_ga',
          'd_sbc_pct', 'd_oda_pct', 'wc_nwc_p', 'is_etr', 'e_rate', 'fcf_eb', 'pg_ads_mix', 'pg_aws_mix', 'pg_svc_mix', 'seg_na_sh', 'gm_ful', 'gm_tech'}
AY = {y: ANN[y].c for y in ANN}

def style_cagr(ws, r, on):
    src = 16 if on else 12
    for c in CAGR:
        ws[f'{c}{r}']._style = copy.copy(TM[f'{c}{src}']._style)

def write_row(ws, row):
    r = row.r; arch = ARCH[row.arch]
    clone_row_style(ws, arch, r)
    ws.cell(r, 2).value = row.label or None
    if row.kw.get('note'): ws[f'{NOTE}{r}'] = row.kw['note']
    style_cagr(ws, r, row.key in CAGR_G or row.key in CAGR_A)
    if row.arch in ('section', 'block', 'blank', 'sub') and not row.kw.get('data'):
        return
    if row.key in DEFS:
        write_defs(ws, row); return
    nf = NF[row.nf]
    for col in COLS:
        v, inp = compute(row, col)
        if v is None: continue
        cell = ws[f'{col.c}{r}']
        cell.value = v
        if not row.kw.get('textdata'): cell.number_format = nf
        if inp:
            set_color(cell, BLUE, bold=True if col.kind in FCK else None)
            lk = row.kw.get('labkey')
            if lk and col.kind in HISTK and isinstance(v, (int, float)) and abs(v) > 1e-9:
                t = xget(pkey(col), lk)
                if t: comment(cell, f'{pkey(col)}: {t}', 420, 130)
        elif row.arch == 'growth' and col.kind in HISTK:
            set_color(cell, GRAY)
        else:
            set_color(cell, BLACK)
    rc = CELL_COMMENTS.get(row.key, {})
    for c, t in rc.items(): comment(ws[f'{c}{r}'], t)
    if row.key in CAGR_G:
        ws[f'{CAGR[0]}{r}'] = f'=IFERROR(({AY[2030]}{r}/{AY[2025]}{r})^(1/5)-1,"n/m")'
        ws[f'{CAGR[1]}{r}'] = f'=IFERROR(({AY[2025]}{r}/{AY[2020]}{r})^(1/5)-1,"n/m")'
        ws[f'{CAGR[2]}{r}'] = f'=IFERROR(({AY[2025]}{r}/{AY[2015]}{r})^(1/10)-1,"n/m")'
        ws[f'{CAGR[3]}{r}'] = f'=IFERROR(({AY[2025]}{r}/{AY[2006]}{r})^(1/19)-1,"n/m")'
        for c in CAGR: ws[f'{c}{r}'].number_format = '0%;\\(0%\\);"n/m"'
    elif row.key in CAGR_A:
        ws[f'{CAGR[0]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{AY[y]}{r}" for y in range(2026, 2031))}),"n/m")'
        ws[f'{CAGR[1]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{AY[y]}{r}" for y in range(2021, 2026))}),"n/m")'
        ws[f'{CAGR[2]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{AY[y]}{r}" for y in range(2016, 2026))}),"n/m")'
        ws[f'{CAGR[3]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{AY[y]}{r}" for y in range(2006, 2026))}),"n/m")'
        for c in CAGR: ws[f'{c}{r}'].number_format = NF['pct']
    t = LABEL_SOURCES.get(row.key)
    if t: comment(ws.cell(r, 2), t)

REL = 'Source: Amazon.com earnings releases (Form 8-K Item 2.02, Ex. 99.1) — '
LABEL_SOURCES = {
    'pg_online': REL + 'supplemental "Net Sales" by groups of similar products and services (quarterly from Q2-16); each quarter from the latest release that presents it '
                 '(Amazon re-presents the groups: advertising services split out of "Other" in the Q4-21 release, back to Q3-20). FY = Σ quarters (computed).',
    'seg_na_rev': REL + 'segment information. AWS segment from Q1-15; FY2013 and Q1–Q4/FY2014 from the AWS recast (8-K 23-Apr-2015, Ex. 99.3). 2015 segment operating income from the '
                  'restated comparatives in the 2016 releases (stock-based compensation and other operating items allocated to segments from 2016). '
                  'FY2006–12 and 2013 quarters: legacy NA (incl. AWS) / International basis — memo rows.',
    'leg_na_rev': REL + 'segment information as originally reported (North America incl. AWS; segment OI before stock-based compensation and other operating expense) to Q4-14.',
    'orig_na_oi': REL + '2015 quarterly / annual segment results as originally reported in the 2015 releases.',
    'mx_media': REL + 'supplemental net sales by segment: Media, Electronics and other general merchandise, Other (2006–2016; NA + International summed).',
    'is_rev': REL + 'consolidated statements of operations (three-month columns; FY from the twelve-month columns of the Q4 release), as originally reported.',
    'd_dda': REL + 'consolidated statements of cash flows ("Depreciation and amortization of property and equipment and capitalized content costs, operating lease assets, and other").',
    'd_dep': 'Source: XBRL company facts, us-gaap:Depreciation (10-K FY2014+; 10-Q quarters from Q2-24; Q4 = FY − 9M computed).',
    'd_sc': 'Source: release text (Form 8-K Ex. 99.1) of Q4-22, Q1-23, Q3-25 and Q4-25 — amounts quoted verbatim in manual_items.json.',
    'd_en': 'Source: Q2-26 Form 10-Q Note 1 (energy contracts subject to derivative accounting): net unrealized gains $551m in Q2-26 and $599m in 1H-26 (Q1-26 $48m computed).',
    'o_pub': 'Source: Q3-25 release ("Without these charges, operating income would have been $21.7 billion") and Q4-25 release ("… $27.4 billion").',
    'cf_ni': REL + 'consolidated statements of cash flows (three-month columns as published; FY twelve-month columns).',
    'fcf_pub': REL + 'supplemental financial information / free-cash-flow reconciliations (TTM, as published in each quarter\'s release).',
    'bs_cash': REL + 'consolidated balance sheets (period ends); common shares outstanding from the supplemental table (split-adjusted ×20 before 6-Jun-2022).',
    'bs_std': 'Source: XBRL company facts LongTermDebtCurrent (2011+) and ShortTermBorrowings (2021+), as first filed.',
    'sh_basic': REL + 'statements of operations (weighted-average shares; ×20 for documents filed before the 20-for-1 split of 6-Jun-2022).',
    'ds_total': 'Source: release balance sheets; XBRL current portion / short-term borrowings; Q2-26 10-Q Note 5; 8-Ks of 13-Mar, 16-Mar, 12-Jun, 9-Jul and 14-Sep-2026 (notes) and 10-Jun-2026 (DDTL).',
    'hr_fl': 'Source: XBRL FinanceLeaseLiabilityCurrent / Noncurrent (2019+); CapitalLeaseObligationsCurrent / Noncurrent (2010–2018; absolute values — the 2010 current fact is tagged negative).',
    'fx_ww': REL + 'supplemental financial information ("WW net sales — Y/Y growth, excluding F/X").',
}
CELL_COMMENTS = {}
def comments_setup():
    for p in ('2013', 'Q1-14', 'Q2-14', 'Q3-14', 'Q4-14', '2014'):
        c = ANN[int(p)].c if p.isdigit() else QC[(2014, int(p[1]))].c
        CELL_COMMENTS.setdefault('seg_na_rev', {})[c] = f'{p}: AWS recast (8-K 23-Apr-2015, Ex. 99.3); originally reported NA incl. AWS in the legacy memo rows. Segment OI before SBC (unallocated row).'
    for p in ('Q1-15', 'Q2-15', 'Q3-15', 'Q4-15', '2015'):
        c = ANN[2015].c if p.isdigit() else QC[(2015, int(p[1]))].c
        x = DATA[p]['x']
        CELL_COMMENTS.setdefault('seg_na_oi', {})[c] = (f'{p}: restated comparative from the 2016 release (SBC and other operating items allocated). As originally reported: '
                                                        f'NA {x["orig_na_oi"]:,.0f}, International {x["orig_intl_oi"]:,.0f}, AWS {x["orig_aws_oi"]:,.0f}, unallocated {x["orig_unalloc"]:,.0f} (memo rows).')
    CELL_COMMENTS.setdefault('d_dda', {})[ANN[2022].c] = ('FY2022 as presented in the Q4-22 release includes the 9M-22 "other operating expense (income), net" ($460m) shown separately in the '
                                                          'Q1–Q3/22 releases, so Σ quarters is $460m lower (documented check_data exception).')
    CELL_COMMENTS.setdefault('pg_other', {})[ANN[2020].c] = 'FY2020: advertising services are split out only from Q3-20, so the FY "Other" includes Q1–Q2/20 advertising (FY advertising blank).'
    CELL_COMMENTS.setdefault('cf_capex', {})[ANN[2017].c] = ('FY2017 (Q4-17 release) shows gross purchases and "proceeds from property and equipment incentives" separately; '
                                                              'Q1–Q3/17 releases showed purchases net of proceeds and incentives.')
    CELL_COMMENTS.setdefault('cf_debt', {})[ANN[2014].c] = 'FY2014 (Q4-14 release) splits capital / finance lease repayments out; Q1–Q3/14 releases show one combined repayment line (in this row).'
    for p in ('Q1-17', 'Q2-17', 'Q3-17', 'Q1-18', 'Q2-18', 'Q3-18'):
        x = DATA[p]['x']
        CELL_COMMENTS.setdefault('fcf_basis', {})[QC[(2000 + int(p[-2:]), int(p[1]))].c] = f'{p}: {x["fcf_basis_note"]}'
    CELL_COMMENTS.setdefault('is_rev', {})[QC[(2026, 2)].c] = 'Q2-26 includes Prime Day (held in Q3 in 2025).'
    CELL_COMMENTS.setdefault('is_othinc', {})[QC[(2026, 2)].c] = 'Q2-26: "non-operating pre-tax other income of $53.4 billion, primarily from our investments in Anthropic" (release text).'

DEFS = {}
def defs_setup():
    q = lambda y, n: QC[(y, n)].c
    DEFS['o_def'] = {
        'C': ('No adjusted operating income published — model adjusted OI = GAAP OI + items Amazon quantified (none before Q4-22)', None),
        q(2022, 4): ('Q4-22: ~$2.7bn charges (self-insurance estimates, impairments, severance) added back', MAN['special_charges']['Q4-22']['text']),
        q(2023, 1): ('Q1-23: ~$0.5bn severance added back', MAN['special_charges']['Q1-23']['text']),
        q(2025, 3): ('Q3-25: FTC settlement $2.5bn + severance $1.8bn; company-published OI ex charges $21.7bn', MAN['special_charges']['Q3-25']['text']),
        q(2025, 4): ('Q4-25: Italy tax disputes & lawsuit $1.1bn, severance $0.73bn, impairments $0.61bn; ex charges $27.4bn', MAN['special_charges']['Q4-25']['text']),
        q(2026, 1): ('2026: energy-contract derivative remeasurements (gains) removed — excluded from Q3-26 guidance', MAN['energy_derivatives']['source']),
    }
    DEFS['e_def'] = {
        'C': ('Model definition (Amazon publishes no adjusted EPS): ex other income (expense), net and equity-method activity, + special items; statutory tax 35%', None),
        AY[2018]: ('From 2018: adjustments tax-effected at the 21% US statutory rate', None),
        q(2021, 4): ('Q4-21–2024: excludes Rivian marks (in other income); 2025–26: excludes Anthropic gains', MAN['nonop_highlights']['Q2-26']),
    }
    DEFS['fcf_def'] = {
        'C': ('FCF = operating cash flow − purchases of fixed assets (TTM, as published)', None),
        q(2014, 4): ('Q4-14+: lease-adjusted FCF variants published (FCF less lease principal; less finance-lease principal & assets under capital leases)', None),
        q(2015, 4): ('Q4-15+: purchases of P&E net of proceeds from sales and incentives', None),
        q(2019, 1): ('2019+: variants renamed "less principal repayments of finance leases and financing obligations" / "less equipment finance leases …" (ASC 842)', None),
        q(2025, 1): ('2025+: only FCF published (lease-adjusted variants discontinued)', None),
    }
def write_defs(ws, row):
    for c, (txt, cm) in DEFS[row.key].items():
        cell = ws[f'{c}{row.r}']; cell.value = txt
        cell.font = Font(name='Calibri', sz=10, b=True, i=True, color='FFC00000')
        if cm: comment(cell, cm)

def write_scen(ws):
    srcrow = {'hdr': 9, 'out_hdr': 11, 'pt_hdr': 15, 'out': 12, 'pt': 16, 'lev': 28, 'blk': 30, 'yr': 31}
    for spec_ in SCEN_ROWS:
        r, kind = spec_[0], spec_[1]
        for ci in range(CI(SL), CI(SNOTE) + 1):
            ws.cell(r, ci)._style = copy.copy(TM.cell(srcrow[kind], ci)._style)
        if kind in ('hdr', 'out_hdr', 'pt_hdr'):
            _, _, a, b, c_, d_, e = spec_
            for col, v in zip([SL, SBULL, SBASE, SBEAR, SNOTE], [a, b, c_, d_, e]):
                if v is not None: ws[f'{col}{r}'] = v
        elif kind == 'out':
            _, _, lab, hi, mid, lo, note, nf = spec_
            ws[f'{SL}{r}'] = lab
            for col, v in zip([SBULL, SBASE, SBEAR], [hi, mid, lo]):
                ws[f'{col}{r}'] = v; ws[f'{col}{r}'].number_format = NF[nf]; set_color(ws[f'{col}{r}'], BLUE)
            ws[f'{SNOTE}{r}'] = note
        elif kind == 'pt':
            _, _, lab, _, val, _, note, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBASE}{r}'] = val; ws[f'{SBASE}{r}'].number_format = NF[nf]; set_color(ws[f'{SBASE}{r}'], BLUE); ws[f'{SNOTE}{r}'] = note
        elif kind == 'lev':
            _, _, lab, a, b, c_, note = spec_
            ws[f'{SL}{r}'] = lab
            for col, v in zip([SBULL, SBASE, SBEAR], [a, b, c_]):
                ws[f'{col}{r}'] = v; ws[f'{col}{r}'].number_format = NF['x']; set_color(ws[f'{col}{r}'], BLUE)
            ws[f'{SNOTE}{r}'] = note
        elif kind == 'blk':
            _, _, lab, db, mid, de, note, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBULL}{r}'] = db; ws[f'{SBASE}{r}'] = mid; ws[f'{SBEAR}{r}'] = de; ws[f'{SNOTE}{r}'] = note
            for col in (SBULL, SBEAR):
                ws[f'{col}{r}'].number_format = '\\+#,##0;\\-#,##0' if nf == 'm' else '\\+0.0%;\\-0.0%'
                set_color(ws[f'{col}{r}'], BLUE)
        elif kind == 'yr':
            _, _, lab, hdr, base, clamp, _, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBASE}{r}'] = base; set_color(ws[f'{SBASE}{r}'], BLUE)
            if clamp:
                ws[f'{SBULL}{r}'] = f'=MAX(0,{SBASE}{r}+${SBULL}${hdr})'; ws[f'{SBEAR}{r}'] = f'=MAX(0,{SBASE}{r}+${SBEAR}${hdr})'
            else:
                ws[f'{SBULL}{r}'] = f'={SBASE}{r}+${SBULL}${hdr}'; ws[f'{SBEAR}{r}'] = f'={SBASE}{r}+${SBEAR}${hdr}'
            for col in (SBULL, SBASE, SBEAR): ws[f'{col}{r}'].number_format = NF[nf]

def write_header(ws):
    for r in range(1, 8):
        clone_row_style(ws, r, r, 2, CI(LASTC))
    ws[f'{NOTE}1'] = 'SCENARIO SWITCH ▼ (Bull / Base / Bear) — drives all scenario-linked forecast drivers'
    ws[f'{NOTE}2'] = 'Base'
    comment(ws[f'{NOTE}2'], 'Scenario toggle. Choose Bull / Base / Bear. Drives: Q3-26 net-sales and operating-income guidance end-points (high / mid / low), '
            '2027E–2030E product-group growth, segment margins, capex intensity, M&A spend and the leverage target for 2027E+ buybacks '
            f'(scenario table {SL}:{SNOTE}).')
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type='list', formula1='"Bull,Base,Bear"', allow_blank=False); ws.add_data_validation(dv); dv.add(f'{NOTE}2')
    ws['B2'] = 'Amazon.com, Inc. (NASDAQ: AMZN) — 3-Statement Financial Model'
    ws['B3'] = ('US GAAP as reported · USD millions (per-share data in USD; shares in millions, split-adjusted) · Fiscal years ended 31 December · '
                'Source: Amazon.com Forms 10-K / 10-Q, Form 8-K earnings releases (Ex. 99.1, incl. the 23-Apr-2015 AWS recast Ex. 99.3) and financing 8-Ks, SEC EDGAR '
                '(CIK 0001018724); XBRL company facts; the investor-relations releases are the same documents as filed on EDGAR. Guidance: Q3-26 only (8-K 30-Jul-2026); '
                'the ~$200bn 2026 capex figure is from the call and UNVERIFIED (information row only).')
    ws['B4'] = ('Basis notes: as originally reported. Segments NA / International to 2014 (AWS inside NA), AWS from 2015 (FY2013–14 recast); 2015 segment OI restated for SBC allocation '
                '(2016 releases). 20-for-1 split 6-Jun-2022. ASC 606 (2018, modified retrospective), ASC 842 (2019). Acquisitions: Whole Foods 28-Aug-17 ($13.2bn), MGM 17-Mar-22 ($6.1bn + $2.5bn debt), One Medical 22-Feb-23 ($3.5bn). '
                'Q3/26E calibrated to Q3-26 guidance (net sales, operating income); Q4/26E carries the Q3 Δs (no guidance).')
    ws[f'{CAGR[0]}5'] = 'Forecast'
    for c in CAGR[1:]: ws[f'{c}5'] = 'Historical'
    ws['B6'] = '(USDm)'
    for col in COLS: ws[f'{col.c}6'] = col.label
    ws[f'{NOTE}6'] = 'Modelling Notes'
    for c, a, b in zip(CAGR, ['5Y CAGR', '5Y CAGR', '10Y CAGR', '19Y CAGR'], ["25-'30", "20-'25", "15-'25", "06-'25"]):
        ws[f'{c}6'] = a; ws[f'{c}7'] = b

def set_columns(ws):
    for k in list(ws.column_dimensions.keys()): del ws.column_dimensions[k]
    for ci in range(1, CI(LASTC) + 1):
        L_ = L(ci)
        ws.column_dimensions[L_].width = TM.column_dimensions[L_].width or 13.0
    for col in COLS:
        if col.q:
            cd = ws.column_dimensions[col.c]; cd.outlineLevel = 1; cd.hidden = True
    ws.sheet_format.outlineLevelCol = 1

def main():
    import openpyxl
    wb = openpyxl.load_workbook(TEMPLATE)
    if 'DCF_old' in wb.sheetnames: del wb['DCF_old']
    ws = wb['Model']
    blank_style = copy.copy(ws['A1']._style)
    for rowc in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for c in rowc:
            c.value = None; c.comment = None; c._style = copy.copy(blank_style)
    for k in list(ws.row_dimensions.keys()):
        rd = ws.row_dimensions[k]; rd.outlineLevel = 0; rd.hidden = False
    ws.data_validations.dataValidation = []
    ws.conditional_formatting = type(ws.conditional_formatting)()
    scen_layout(V0)
    build_spec(V0)
    n1 = len(fw.SPEC)
    build_spec2(V0)
    for i, row in enumerate(fw.SPEC): row.spec2 = i >= n1
    last = number_rows(8)
    for r in range(1, last + 1): ws.row_dimensions[r].height = 15.0
    ws.row_dimensions[2].height = 21.0
    defs_setup(); comments_setup()
    write_header(ws)
    for row in fw.SPEC: write_row(ws, row)
    write_scen(ws)
    set_columns(ws)
    ws.freeze_panes = 'C8'
    ws.sheet_view.pane.topLeftCell = f'{ANN[2019].c}8'         # opens on 2019 (quarterly columns collapsed)
    ws.sheet_view.zoomScale = 70
    import other_sheets as osh
    osh.do_all(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    osh.inject_charts(OUT)
    print('rows', fw.SPEC[-1].r, 'saved', OUT)
    return wb

RECALC = os.environ.get('XLSX_RECALC', '/mnt/skills/public/xlsx/scripts/recalc.py')
def recalc(path, timeout=900):
    import subprocess
    out = subprocess.run(['python3', RECALC, path, str(timeout)], capture_output=True, text=True, cwd=os.path.dirname(RECALC))
    try:
        return json.loads(out.stdout)
    except Exception:
        return {'raw': out.stdout[-2000:], 'err': out.stderr[-2000:]}

def scenario_snapshots():
    import openpyxl
    import other_sheets as osh
    snaps = {}
    for scen in ('Bull', 'Bear'):
        p = os.path.join(BASE, f'AMZN_{scen}.xlsx')
        wb = openpyxl.load_workbook(OUT); wb['Model'][f'{NOTE}2'] = scen; wb.save(p)
        res = recalc(p); print(scen, res.get('status'), res.get('total_errors'))
        v = openpyxl.load_workbook(p, data_only=True)['Bull-Base-Bear']
        snaps[scen] = {(r, c): v[f'{c}{r}'].value for r in range(9, 51) for c in 'CDEFGHIJKL' if v[f'{c}{r}'].value is not None}
        os.remove(p)
    wb = openpyxl.load_workbook(OUT)
    osh.snapshot_values(wb, snaps['Bull'], snaps['Bear'])
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    osh.inject_charts(OUT)

NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
def inject_values(src_xlsx, recalc_xlsx):
    """Store cached formula results (from a LibreOffice-recalculated copy) in the openpyxl file, so the workbook reads correctly in any
    viewer; fullCalcOnLoad stays set so Excel recalculates on open. Same approach as the reference kits' finishing step."""
    import zipfile, shutil, openpyxl
    from lxml import etree
    wbv = openpyxl.load_workbook(recalc_xlsx, data_only=True)
    zin = zipfile.ZipFile(src_xlsx)
    wbxml = etree.fromstring(zin.read('xl/workbook.xml'))
    rels = etree.fromstring(zin.read('xl/_rels/workbook.xml.rels'))
    rid2t = {r.get('Id'): r.get('Target') for r in rels}
    sheets = {}
    for s in wbxml.iter('{%s}sheet' % NS):
        rid = s.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
        sheets['xl/' + rid2t[rid].lstrip('/').replace('xl/', '')] = s.get('name')
    tmp = src_xlsx + '.tmp'
    zout = zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED)
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename in sheets:
            ws = wbv[sheets[item.filename]]
            root = etree.fromstring(data)
            for c in root.iter('{%s}c' % NS):
                f = c.find('{%s}f' % NS)
                if f is None: continue
                v = ws[c.get('r')].value
                old = c.find('{%s}v' % NS)
                if old is not None: c.remove(old)
                if v is None: continue
                ve = etree.SubElement(c, '{%s}v' % NS)
                if isinstance(v, bool):
                    c.set('t', 'b'); ve.text = '1' if v else '0'
                elif isinstance(v, (int, float)):
                    if 't' in c.attrib: del c.attrib['t']
                    ve.text = repr(float(v)) if isinstance(v, float) else str(v)
                elif isinstance(v, str) and v.startswith('#'):
                    c.set('t', 'e'); ve.text = v
                else:
                    c.set('t', 'str'); ve.text = str(v)
            data = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
        zout.writestr(item, data)
    zout.close(); zin.close()
    shutil.move(tmp, src_xlsx)

def store_values():
    import shutil
    p = os.path.join(BASE, 'AMZN_recalc.xlsx')
    shutil.copy(OUT, p)
    res = recalc(p); print('recalc', res.get('status'), res.get('total_errors'))
    inject_values(OUT, p)
    os.remove(p)

if __name__ == '__main__':
    main()
    if '--snap' in sys.argv: scenario_snapshots()
    if '--recalc' in sys.argv: store_values()
