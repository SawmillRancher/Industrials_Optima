"""Rewire the Bull-Base-Bear, DCF (current template layout) and Charts tabs of the MLM template to the AMZN Model rows; drop the
template's DCF_old tab; keep the template chart XML (retitled). The Model sheet keeps the template's column layout, so only the
row references are remapped (template row -> AMZN row key)."""
import re, zipfile, shutil
from openpyxl.worksheet.formula import ArrayFormula
from fw import R, TEMPLATE, ANN, comment

MAP = {158: 'is_rev', 166: 'is_oi', 164: 'r_cash', 182: 'd_dep', 186: 'sh_dil', 211: 'r_adj', 212: 'r_adj_m', 219: 'r_ebit',
       221: 'r_adj', 229: 'e_adjni', 231: 'e_eps', 238: 'gm_aws', 245: 'gm_ebitda_m', 247: 'gm_op_m', 260: 'cf_wc', 262: 'cf_capex',
       263: 'cf_acq', 280: 'fcf', 282: 'fcf_m', 290: 'bb_px', 291: 'bb_cash', 292: 'bb_sh', 331: 'bs_nci', 408: 'ra_ic', 411: 'ra_roic',
       422: 'rn_ronta', 428: 'lv_nd', 433: 'v_px', 434: 'v_sh', 12: 'gm_rev', 30: 'gm_rev', 14: 'gm_aws', 31: 'gm_aws', 353: 'ra_ic', 355: 'ra_ic'}
REF = re.compile(r"Model!(\$?)([A-Z]{1,3})(\$?)(\d+)")

def remap(f):
    def rep(m):
        r = int(m.group(4))
        if m.group(2) == 'CF' and r == 2: return m.group(0)          # scenario switch
        if r in MAP: return f'Model!{m.group(1)}{m.group(2)}{m.group(3)}{R[MAP[r]]}'
        raise KeyError(f'unmapped template Model row {r} in {f}')
    return REF.sub(rep, f)

BBB = {
    'B2': 'Amazon.com, Inc. (AMZN) — Bull / Base / Bear Case P&L Summary',
    'B3': ('Model definitions (Amazon publishes no adjusted measures besides free cash flow): adjusted operating income / EBITDA add back only the special charges Amazon '
           'quantified and energy-contract remeasurements; adjusted EPS also excludes other income (expense) — investment marks (Anthropic, Rivian) — and equity-method activity. '
           'USD millions · Live panel linked to Model — toggle the scenario with the switch on the Model sheet (cell CF2). Bull and Bear panels are value snapshots generated from the Model with the switch set '
           'to each case. 2026E = 1H actual + Q3 calibrated to the Q3-26 guidance + Q4 carrying the Q3 Δs.'),
    'B9': 'Net sales', 'B10': '   Net sales y/y %', 'B11': '   AWS segment net sales y/y %',
    'B12': 'Adjusted EBITDA (model definition)', 'B15': 'Operating Income (GAAP)', 'B16': '   Operating income y/y %',
    'B19': 'Adjusted Net Income (model definition)', 'B35': 'Adjusted EBITDA (leverage basis)', 'B36': '   Adjusted EBITDA % of net sales',
    'B29': 'IRR (to Dec-Y-1)',
}
for k in list(BBB):
    if k[0] == 'B' and k not in ('B2', 'B3'):
        BBB['N' + k[1:]] = BBB[k]; BBB['Z' + k[1:]] = BBB[k]

DCF = {
    'B2': 'DCF VALUATION — AMAZON.COM, INC. (AMZN)',
    'D5': 'Assumption (~US 10Y Treasury yield, Oct-2026; government rate sources not reachable from the build environment); update to market',
    'D6': 'Assumption: US implied ERP ~4.5–5.5%',
    'D7': ('Assumption: mega-cap internet / cloud peers (MSFT, GOOGL, META, ORCL) ~1.0–1.3; Amazon\'s retail cyclicality offset by AWS / advertising contract revenue; '
           'heavy AI capex raises operating leverage'),
    'D9': 'AA / A1 issuer; 2026 issues priced at 3.85–6.25% (USD, 8-Ks 13-Mar and 9-Jul-2026); marginal long-dated cost ~5.3%',
    'D10': 'Model 2027E+ effective tax rate (FY2025 19.6%)',
    'D12': 'Gross debt ~$165bn after the Jul/Sep-2026 notes vs ~$2.8tn market cap ($260 × ~10.9bn shares) → ~5% D/V; net debt near zero',
    'D15': 'Periods = (Dec-31-Year − today) / 365.25; cash flows assumed at year-end. 2026E UFCF is excluded (row 24) because valuation net debt is taken at YE2026E',
    'B18': 'Adjusted operating income (full-cost: after SBC and all D&A) less cash special charges',
    'B21': 'plus depreciation of property & equipment (operating-lease, content and other amortization is reinvested — see Model cash-flow offset line)',
    'B22': '− Capex net of proceeds, acquisitions & investments (M&A lever 2027E–30E) and finance-lease principal',
    'B26': 'Fade-period growth (2031E–35E; 9% → 5% linear)',
    'D32': ('Amazon\'s operating income expenses stock-based compensation (not adjusted out) and all D&A, so NOPAT is taxed on full-cost adjusted operating income less cash '
            'special charges (settlements, severance). Only P&E depreciation is added back: operating-lease and content amortization is matched by lease payments / content spend '
            '(held flat in the Model). Finance-lease and financing-obligation principal is deducted with capex because leases are excluded from net debt. Non-marketable equity '
            'investments (Anthropic) are neither in UFCF nor added to EV (conservative; their marks sit in other income). 2026E UFCF is set to zero (cash flows already in YE2026E net debt).'),
    'D43': ('Opening IC = Model 2025A book invested capital (equity + net debt; includes the non-marketable investments in other assets), rolled forward with DCF net investment '
            '(net capex + M&A + lease principal − P&E depreciation − ΔWC).'),
    'B48': 'Less: Net debt + NCI (Model 2026E year end)',
    'D48': 'AMZN: no noncontrolling interests; net debt = debt (incl. Jul/Sep-2026 notes) − cash − marketable securities; finance-lease liabilities excluded (principal deducted in UFCF).',
    'B50': 'Diluted shares (m, Model 2026E period end)',
    'B70': ('To solve for the market-implied path: Data ▸ What-If Analysis ▸ Goal Seek — set C51 to the current price (C52) by changing C27 (terminal growth) '
            'or the fade-period growth inputs (I26:M26).'),
    'D76': 'Rationale (Amazon characteristics)',
    'D78': ('Model adjusted operating income CAGR 2025–30E (~15%): AWS growth on the AI build-out, advertising mix and retail margin expansion, partly offset by rising depreciation'),
    'D79': ('Advertising, AWS and third-party services earn high incremental returns on a shared logistics / data-centre base; the 2025–27 AI capex wave lowers near-term '
            'returns until utilisation catches up'),
    'D80': ('Scale in fulfilment, Prime ecosystem, AWS installed base and switching costs sustain a long competitive-advantage period; capital intensity and regulatory risk cap it'),
    'B90': 'VALUATION  —  Applied to 2028E NOPAT',
    'B91': '2028E adjusted operating income less cash special charges',
    'B96': 'Less: YE2027 net debt + NCI (Model)',
    'B108': 'AMZN market-implied EV/NOPAT (on 2028E NOPAT)',
    'B36': 'less P&E depreciation',
}
CHARTS = {
    'B2': 'Amazon.com, Inc. — Historical Operating Performance (2006–2025)',
    'B3': ('All historicals · USDm · Linked to Model. Bars = contribution to net-sales growth of North America + International vs AWS (pp; stacked = total growth; AWS a segment from 2015, recast 2013–14, so 2006–13 bars show total growth). '
           'Adjusted EBITDA = model definition (operating income + total D&A + Amazon-quantified special items). ROIC / RONTA on adjusted operating income (full-cost); blank (plotted at zero) where invested capital is negative (2006–11, n/m).'),
    'B8': 'North America + International contribution (pp; total growth before 2014)', 'B9': 'AWS contribution (pp)', 'B12': 'Adjusted EBITDA (model)', 'B13': 'Operating Income (GAAP)',
    'B14': 'Adj. EBITDA Margin %', 'B15': 'GAAP Op. Margin %',
}
CHART_TITLES = {
    'Aggregates Revenue Drivers: Shipment Volume vs. Price (ASP), y/y %': 'Net Sales Growth: Retail vs. AWS Contribution, pp',
    'Adjusted EBITDA &amp; Earnings from Operations (USDm)': 'Adjusted EBITDA &amp; Operating Income (USDm)',
}

def _rewire(ws):
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if isinstance(v, ArrayFormula):
                if 'Model!' in (v.text or ''): c.value = ArrayFormula(v.ref, remap(v.text))
            elif isinstance(v, str) and 'Model!' in v:
                c.value = remap(v)

def build_charts(ws):
    for i, y in enumerate(range(2006, 2026)):
        c = chr(ord('C') + i); mc = ANN[y].c
        ws[f'{c}7'] = y
        ws[f'{c}8'] = f'=IFERROR(Model!{mc}{R["gc_na"]}+Model!{mc}{R["gc_intl"]},IFERROR(Model!{mc}{R["gm_rev"]}+0,0))'
        ws[f'{c}9'] = f'=IFERROR(Model!{mc}{R["gc_aws"]}+0,0)'
        ws[f'{c}11'] = f'={c}7'
        ws[f'{c}12'] = f'=Model!{mc}{R["r_adj"]}'
        ws[f'{c}13'] = f'=Model!{mc}{R["is_oi"]}'
        ws[f'{c}14'] = f'=Model!{mc}{R["gm_ebitda_m"]}'
        ws[f'{c}15'] = f'=Model!{mc}{R["gm_op_m"]}'
        ws[f'{c}17'] = f'={c}7'
        ws[f'{c}18'] = f'=IFERROR(Model!{mc}{R["ra_roic"]}+0,"")'
        ws[f'{c}19'] = f'=IFERROR(Model!{mc}{R["rn_ronta"]}+0,"")'
        ws[f'{c}20'] = f'=IFERROR(Model!{mc}{R["gm_ebitda_m"]}+0,NA())'
        ws[f'{c}21'] = f'=IFERROR(Model!{mc}{R["fcf_m"]}+0,NA())'

def do_all(wb):
    if 'DCF_old' in wb.sheetnames: del wb['DCF_old']
    for name, text in (('Bull-Base-Bear', BBB), ('DCF', DCF), ('Charts', CHARTS)):
        ws = wb[name]
        _rewire(ws)
        for addr, t in text.items(): ws[addr] = t
    d = wb['DCF']
    for col, mc in zip('DEFGH', [ANN[y].c for y in range(2026, 2031)]):
        d[f'{col}21'] = f'=Model!{mc}{R["d_dep"]}'
        d[f'{col}22'] = f'=Model!{mc}{R["cf_capex"]}+Model!{mc}{R["cf_capex_proc"]}+Model!{mc}{R["cf_acq"]}+Model!{mc}{R["cf_lease"]}'
    d['I22'] = '=H22*(1+I$26)'
    d['D24'] = '=0*(D20+D21+D22+D23)'
    d['D34'] = f'=Model!BV{R["ra_ic"]}'
    d['P24'] = '=IFERROR((H24/E24)^(1/(COLUMNS(E24:H24)-1))-1,"n/m")'; d['Q24'] = '=IFERROR((M24/E24)^(1/(COLUMNS(E24:M24)-1))-1,"n/m")'
    comment(d['P24'], 'UFCF CAGRs start from 2027E because 2026E UFCF is excluded (zero) in this layout.')
    d['C5'] = 0.0425; d['C6'] = 0.05; d['C7'] = 1.10; d['C9'] = 0.053; d['C10'] = 0.19; d['C12'] = 0.05
    d['I26'] = 0.09; d['M26'] = 0.05; d['C27'] = 0.035
    d['C78'] = 0.12; d['C79'] = 0.20; d['C80'] = 15
    comment(d['D24'], '2026E unlevered FCF excluded (×0): valuation net debt is taken at YE2026E, which already reflects the 2026 cash flows (playbook rule 5.7: do not double-count the current year).')
    b = wb['Bull-Base-Bear']
    b['H26'] = 28; b['G26'] = 28; b['T26'] = 32; b['S26'] = 32; b['AF26'] = 22; b['AE26'] = 22; b['AD26'] = None
    b['N26'].comment = None; b['P40'].comment = None
    comment(b['B26'], 'P/E on adjusted EPS (model definition: ex investment marks and equity-method activity, + Amazon-quantified special items, tax-effected at 21%). '
            'At $260.00 the stock trades on ~33x 2026E model adjusted EPS. Exit multiples are assumptions: Base 28x, Bull 32x, Bear 22x applied to 2029E–30E.')
    comment(b['P40'], f'Snapshot (values) of Model rows {R["bb_cash"]}, {R["bb_px"]} and {R["bb_sh"]} (Share Buyback Schedule) with the Model scenario switch (cell CF2) = Bull; '
            'Bear panel (AB40:AF42) likewise. Base panel (D40:H42) is live-linked.')
    for c in ('B', 'N', 'Z'):
        if b[f'{c}40'].value: b[f'{c}40'] = '   Share repurchases ($m; 2026E = 2H)'
    build_charts(wb['Charts'])

def snapshot_values(wb, bull, bear):
    ws = wb['Bull-Base-Bear']
    cols = 'CDEFGHIJKL'
    for snap, off in ((bull, 'OPQRSTUVWX'), (bear, ['AA', 'AB', 'AC', 'AD', 'AE', 'AF', 'AG', 'AH', 'AI', 'AJ'])):
        for (r, c), v in snap.items():
            tgt = f'{off[cols.index(c)]}{r}'
            cur = ws[tgt].value
            if isinstance(cur, str) and cur.startswith('='): continue
            if r in (26,): continue
            ws[tgt] = v

def inject_charts(path):
    """Replace openpyxl-rewritten chart parts with the template chart XML (keeps the combo chart, colours and layout), retitled, with
    cached values removed so Excel redraws from the AMZN cells."""
    tz = zipfile.ZipFile(TEMPLATE)
    tmpl = {n: tz.read(n) for n in tz.namelist() if n.startswith('xl/charts/chart')}
    tmp = path + '.tmp'
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in tmpl:
                t = tmpl[item.filename].decode('utf-8')
                for a, b2 in CHART_TITLES.items(): t = t.replace(a, b2)
                t = re.sub(r'<c:numCache>.*?</c:numCache>', '', t, flags=re.S)
                t = re.sub(r'<c:strCache>.*?</c:strCache>', '', t, flags=re.S)
                data = t.encode('utf-8')
            zout.writestr(item, data)
    shutil.move(tmp, path)
