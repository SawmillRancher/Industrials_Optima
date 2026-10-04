"""Build models/RMS_Model.xlsx — Hermès International (Euronext Paris: RMS) 3-statement model on the MLM template.

Usage (from the repo root):
    RECALC_SCRIPT=<path to xlsx skill scripts/recalc.py> python3 -m models.rms.build
"""
from __future__ import annotations

import json
import os
import subprocess
import sys

import openpyxl

from .assumptions import ASSUMP
from .data import load
from .framework import Cols
from .other_tabs import write_bbb, write_charts, write_dcf
from .rows import build_rows
from .scenario import SCENARIOS, write_scenarios
from .writer import Data, set_widths, write_model

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
OUT = os.path.join(ROOT, "models", "RMS_Model.xlsx")
RECALC = os.environ.get("RECALC_SCRIPT", "")

TITLE = "Hermès International S.C.A. (Euronext Paris: RMS) — 3-Statement Financial Model"
SUB = ("IFRS as reported · EUR millions (per-share data in EUR) · Source: Hermès registration documents / Universal "
       "Registration Documents (annual financial reports FY2010–FY2025; FY2009 comparatives and FY2006–08 key "
       "figures from the 2010 document), half-year financial reports (H1 2013–H1 2026), half-year / full-year results "
       "and quarterly revenue releases — finance.hermes.com (documents filed with the AMF). Hermès is not an SEC "
       "registrant: EDGAR holds only ADR (Form F-6) registrations, no 10-K / 10-Q / 8-K.")
BASIS = ("Basis notes: calendar fiscal year; interim columns are half-years (Hermès publishes full statements for H1 "
         "and FY only): H1 as reported, H2 = FY − H1 (balances = 31-Dec). Each period as originally reported (2018 "
         "pre-IFRS 16; IFRS 16 from 2019, 2018 not restated). Tableware folded into Other Hermès sectors (separate to "
         "2013). Free-share plan expense reclassified from S&A to other income & expenses from 2015 (2014 as "
         "reported). Non-recurring items: 2018 HK property gain €52.7m; 2020 Shang Xia deconsolidation €91.1m. French "
         "exceptional contribution on large companies' profits in 2025 (€331m) and renewed for 2026. H2/26E–2030E = "
         "forecast (no company numeric guidance; scenario-driven).")
SWITCH_COMMENT = ("Scenario toggle. Choose Bull / Base / Bear. Drives: H2/26E currency effect, constant-currency "
                  "growth by métier (H2/26E, 2027E–30E), gross margin, recurring operating margin, buybacks, "
                  "exceptional dividends and payout (scenario input table to the right). Bull / Bear = Base + Δ.")


def build(path=OUT, scenario="Base"):
    P, notes, conflicts, texts = load()
    if conflicts:
        print("NOTE data overlaps:", conflicts[:10])
    cols = Cols()
    rows = build_rows({"assump": ASSUMP, "P": P})
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model"
    rowmap = write_model(ws, cols, rows, Data(P), TITLE, SUB, BASIS, SWITCH_COMMENT)
    ws[f"{cols.NOTES}2"] = scenario
    set_widths(ws, cols)
    write_scenarios(ws, cols, SCENARIOS)
    write_bbb(wb, cols, rowmap)
    write_dcf(wb, cols, rowmap)
    write_charts(wb, cols, rowmap)
    wb.save(path)
    return cols, rowmap


def recalc(path):
    if not RECALC:
        print("RECALC_SCRIPT not set — skipping recalculation")
        return None
    out = subprocess.run([sys.executable, RECALC, path, os.environ.get("RECALC_TIMEOUT", "300")], capture_output=True,
                         text=True)
    try:
        return json.loads(out.stdout)
    except Exception:
        print(out.stdout, out.stderr)
        return None


VAL_ROWS = (26, 27, 28, 29, 34, 44, 45, 46, 47)   # live P/E-driven rows stay as formulas in the snapshot panels


def snapshot(path, cols):
    """Run the model in Bull and Bear and paste values into the Bull-Base-Bear snapshot panels."""
    vals = {}
    for scn in ("Bull", "Bear"):
        tmp = path.replace(".xlsx", f"_{scn}.xlsx")
        wb = openpyxl.load_workbook(path)
        wb["Model"][f"{cols.NOTES}2"] = scn
        wb.save(tmp)
        recalc(tmp)
        wbv = openpyxl.load_workbook(tmp, data_only=True)
        b = wbv["Bull-Base-Bear"]
        vals[scn] = {(r, i): b.cell(r, 3 + i).value for r in range(9, 48) for i in range(0, 11)}
        os.remove(tmp)
    wb = openpyxl.load_workbook(path)
    b = wb["Bull-Base-Bear"]
    for scn, start in (("Bull", 15), ("Bear", 27)):   # O and AA columns
        for (r, i), v in vals[scn].items():
            if r in VAL_ROWS:
                continue
            cell = b.cell(r, start + i)
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.value = v
    wb.save(path)


if __name__ == "__main__":
    cols, rowmap = build()
    res = recalc(OUT)
    print(json.dumps(res, indent=1)[:3000] if res else "no recalc")
    if res and res.get("status") in ("success", "errors_found"):
        snapshot(OUT, cols)
        res = recalc(OUT)
        print(json.dumps(res, indent=1)[:2000])
