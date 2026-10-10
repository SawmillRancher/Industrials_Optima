"""Bull / Base / Bear scenario input table (right of the Model notes column) — layout of the MLM / TDY models.

Every lever block is 6 rows: header (Bull / Bear Δ inputs), H2/26E (or 2026E), 2027E, 2028E, 2029E, 2030E.
"""
from openpyxl.styles import Alignment

from .framework import BLACK, BLUE, FMT_NUM, FMT_PCT, GREY, LIGHT, MID, WHITE, fill, font, put_comment

# --- absolute rows in the scenario columns
S_FX26, S_TAX26, S_EXC26, S_CAPEX26 = 12, 13, 14, 15
S_CC = {"leather": 18, "rtw": 24, "silk": 30, "other_hermes": 36, "perfume": 42, "watches": 48, "other_products": 54}
S_GM, S_ROI, S_BUYB, S_XDIV, S_PAYOUT = 60, 66, 72, 78, 84

SECTOR_NAMES = {
    "leather": "Leather Goods & Saddlery",
    "rtw": "Ready-to-wear and Accessories",
    "silk": "Silk and Textiles",
    "other_hermes": "Other Hermès sectors",
    "perfume": "Perfume and Beauty",
    "watches": "Watches",
    "other_products": "Other products",
}


def srow(base, c):
    """Scenario value row for column c: H2/26E (and 2026E for annual levers) → base+1; year y → base+1+(y−2026)."""
    if c.kind == "HE":
        return base + 1
    return base + 1 + (c.year - 2026)


SCENARIOS = {
    "point_hdr": "2026E point estimates — H2/26E drivers (company has no numeric FY26 guidance)",
    "point_comment": ("Hermès does not give revenue or margin guidance: H1/26 release (29-Jul-26) 'confirms an ambitious "
                      "goal for revenue growth at constant exchange rates' in the medium term. H1/26 report: 2026 ETR "
                      "expected at ~33% incl. the renewed French exceptional contribution (~28% excluding it); currency "
                      "impact > €360m on H1 revenue."),
    "fx26": (-0.02, -0.03, -0.05),
    "fx26_note": ("H2/26E currency effect on reported revenue growth (H1/26: −4.5 pts: +1.6% reported vs +6.1% at "
                  "constant rates). USD, JPY and CNY weaker vs the euro than in H2/25."),
    "tax26": 0.33, "exc26": 0.05, "capex26": 1150.0,
    "point_note": ("ETR 33% and contribution ≈ 5 pts of pre-tax income: H1/26 report outlook. Capex: 2025 €1,161m, "
                   "H1/26 €344m; 2024 URD guided ~€1bn for 2025."),
    # constant-currency growth levers: base for H2/26E, 2027E..2030E; (Bull Δ, Bear Δ)
    "cc": {
        "leather": dict(base=(0.09, 0.09, 0.09, 0.085, 0.08), d=(0.02, -0.03),
                        note="H1/26 +9.8% (Q2 slightly above Q1); FY25 +13.1%, FY24 +18.3%. Capacity: 25th workshop "
                             "(Loupes, 2026), three more by 2030."),
        "rtw": dict(base=(0.04, 0.06, 0.06, 0.06, 0.055), d=(0.02, -0.03),
                    note="H1/26 +2.0% with a Q2 acceleration (women's AW26 collection); FY25 +6.1%, FY24 +15.4%."),
        "silk": dict(base=(0.08, 0.06, 0.06, 0.055, 0.05), d=(0.02, -0.03),
                     note="H1/26 +9.7% (acceleration in Q2); FY25 +4.7%, FY24 +3.8%."),
        "other_hermes": dict(base=(0.06, 0.08, 0.08, 0.075, 0.07), d=(0.02, -0.03),
                             note="Jewellery and Home: H1/26 +5.4%; FY25 +11.2%, FY24 +17.1%."),
        "perfume": dict(base=(-0.02, 0.02, 0.03, 0.03, 0.03), d=(0.02, -0.03),
                        note="H1/26 −4.5%; FY25 −7.6%, FY24 +9.3%."),
        "watches": dict(base=(0.03, 0.05, 0.05, 0.05, 0.05), d=(0.02, -0.03),
                        note="H1/26 +0.2% with a solid Q2; FY25 −1.5%, FY24 −4.2%."),
        "other_products": dict(base=(0.03, 0.04, 0.04, 0.04, 0.04), d=(0.02, -0.03),
                               note="John Lobb, Saint-Louis, Puiforcat, textiles (Holding Textile Hermès): H1/26 +2.8%; FY25 +5.5%."),
    },
    "gm": dict(base=(0.713, 0.712, 0.713, 0.714, 0.714), d=(0.005, -0.010),
               note="Gross margin: 71.1% FY25 (H2/25 71.5%), 71.1% H1/26 (FX headwind offset by price increases)."),
    "roi": dict(base=(0.402, 0.405, 0.408, 0.410, 0.410), d=(0.015, -0.025),
                note="Recurring operating margin (Hermès KPI): 41.0% FY25 (H2/25 40.7%), 41.0% H1/26 (41.4% H1/25). "
                     "H2/26E: H2/25 40.7% less the H1 y/y drift (−0.4 pt) and FX."),
    "buyb": dict(base=(50.0, 250.0, 250.0, 250.0, 250.0), d=(300.0, -200.0),
                 note="Share buybacks (€m; H2/26E then 2027E+): H1/26 €160m (94,846 shares) for employee free-share "
                      "plans; 2025 €8m net. Floor at 0."),
    "xdiv": dict(base=(0.0, 0.0, 0.0, 0.0, 0.0), d=(1000.0, 0.0),
                 note="Exceptional dividends paid (€m; 2026E then 2027E+): €10/share (~€1.05bn) paid in 2024 and 2025; "
                      "none for FY25. Floor at 0."),
    "payout": dict(base=(0.40, 0.40, 0.40, 0.40, 0.40), d=(0.05, 0.0),
                   note="Ordinary DPS / recurring diluted EPS (2026E then 2027E+): FY25 €18.00 / €46.2 = 39%; FY24 "
                        "€16.00 / €43.9 = 36%."),
}


def write_scenarios(ws, cols, S):
    L, NB, NBa, NBe, NOTE = cols.SCN_LABEL, cols.SCN["Bull"], cols.SCN["Base"], cols.SCN["Bear"], cols.SCN_NOTE

    def hdr(r, text, a="Bull", b="Base", c="Bear"):
        ws[f"{L}{r}"] = text
        ws[f"{L}{r}"].font = font(12, True, color=WHITE)
        ws[f"{L}{r}"].fill = fill(MID)
        for col, v in ((NB, a), (NBa, b), (NBe, c)):
            if v is not None:
                ws[f"{col}{r}"] = v
            ws[f"{col}{r}"].font = font(12, True, color=WHITE)
            ws[f"{col}{r}"].fill = fill(MID)
            ws[f"{col}{r}"].alignment = Alignment(horizontal="center")

    ws[f"{L}9"] = "SCENARIO INPUT TABLE — driver assumptions per Bull / Base / Bear"
    ws[f"{L}9"].font = font(11, True, color=BLACK)
    for col, v in ((NB, "Bull"), (NBa, "Base"), (NBe, "Bear")):
        ws[f"{col}9"] = v
        ws[f"{col}9"].font = font(12, True, color=BLACK)
        ws[f"{col}9"].fill = fill(LIGHT)
        ws[f"{col}9"].alignment = Alignment(horizontal="center")
    ws[f"{NOTE}9"] = ("Bull / Bear = Base + Δ (Δ inputs on each block header). Point-estimate block applies to all "
                      "cases except the H2/26E currency effect.")
    ws[f"{NOTE}9"].font = font(10)

    def val(cell, v, fmt):
        ws[cell] = v
        is_formula = isinstance(v, str) and v.startswith("=")
        ws[cell].font = font(12, color=BLACK if is_formula else BLUE)
        ws[cell].number_format = fmt

    def note(r, text):
        ws[f"{NOTE}{r}"] = text
        ws[f"{NOTE}{r}"].font = font(10)

    hdr(11, S["point_hdr"], None, None, None)
    put_comment(ws[f"{L}11"], S["point_comment"])
    ws[f"{L}{S_FX26}"] = "  H2/26E currency effect on revenue growth (pts)"
    for col, v in zip((NB, NBa, NBe), S["fx26"]):
        val(f"{col}{S_FX26}", v, FMT_PCT)
    note(S_FX26, S["fx26_note"])
    for r, lab, v, fmt in [(S_TAX26, "  2026E effective tax rate (incl. French exceptional contribution)", S["tax26"],
                            FMT_PCT),
                           (S_EXC26, "  2026E exceptional contribution (pts of pre-tax income)", S["exc26"], FMT_PCT),
                           (S_CAPEX26, "  2026E operating investments (€m)", S["capex26"], FMT_NUM)]:
        ws[f"{L}{r}"] = lab
        val(f"{NBa}{r}", v, fmt)
    for r in (S_FX26, S_TAX26, S_EXC26, S_CAPEX26):
        ws[f"{L}{r}"].font = font(12)
    note(S_TAX26, S["point_note"])

    def lever(base_row, name, base_vals, d_bull, d_bear, fmt=FMT_PCT, first="H2/26E", floor=False,
              dfmt='\\+0.0%;\\-0.0%', note_txt=""):
        hdr(base_row, name, None, "Δ →", None)
        for col, d in ((NB, d_bull), (NBe, d_bear)):
            ws[f"{col}{base_row}"] = d
            ws[f"{col}{base_row}"].font = font(12, True, color=BLUE)
            ws[f"{col}{base_row}"].fill = fill(GREY)
            ws[f"{col}{base_row}"].number_format = dfmt
        if note_txt:
            note(base_row, note_txt)
        for i, bv in enumerate(base_vals):
            r = base_row + 1 + i
            ws[f"{L}{r}"] = f"  {first}" if i == 0 else f"  {2026 + i}E"
            ws[f"{L}{r}"].font = font(12)
            val(f"{NBa}{r}", bv, fmt)
            for col in (NB, NBe):
                f = f"={NBa}{r}+${col}${base_row}"
                if floor:
                    f = f"=MAX(0,{NBa}{r}+${col}${base_row})"
                val(f"{col}{r}", f, fmt)

    for s, b in S_CC.items():
        g = S["cc"][s]
        lever(b, f"{s.upper()}_ccGrowth ({SECTOR_NAMES[s]} revenue y/y, constant FX)", g["base"], *g["d"],
              note_txt=g["note"])
    lever(S_GM, "GM (gross margin %)", S["gm"]["base"], *S["gm"]["d"], note_txt=S["gm"]["note"])
    lever(S_ROI, "ROI_margin (recurring operating margin %)", S["roi"]["base"], *S["roi"]["d"],
          note_txt=S["roi"]["note"])
    lever(S_BUYB, "Buybacks (€m)", S["buyb"]["base"], *S["buyb"]["d"], fmt=FMT_NUM, floor=True,
          dfmt='\\+#,##0;\\-#,##0', note_txt=S["buyb"]["note"])
    lever(S_XDIV, "Exceptional dividends paid (€m)", S["xdiv"]["base"], *S["xdiv"]["d"], fmt=FMT_NUM, first="2026E",
          floor=True, dfmt='\\+#,##0;\\-#,##0', note_txt=S["xdiv"]["note"])
    lever(S_PAYOUT, "Payout (ordinary DPS / recurring EPS)", S["payout"]["base"], *S["payout"]["d"], first="2026E",
          note_txt=S["payout"]["note"])
