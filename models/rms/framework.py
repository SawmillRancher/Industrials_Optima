"""Sheet-building framework for the Hermès model (layout / styling of the MLM model, itself on the TDY template).

The Model tab is an ordered list of Row objects. Each row knows how to fill every period column: historical
hard-codes (blue) from the extracted filings, formulas (black) for derived lines, and forecast formulas / inputs
for H2/26E, 2026E and 2027E–2030E.

Hermès publishes full financial statements twice a year (H1 and FY), so the interim columns are half-years:
H1 as reported, H2 = FY − H1 (flows) or the year-end balance (stocks).
"""
from __future__ import annotations

from dataclasses import dataclass
from typing import Callable, Optional

from openpyxl.comments import Comment
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------------------
# Palette / formats (MLM model)
# --------------------------------------------------------------------------------------
DARK = "2F3E46"
MID = "52796F"
SAGE = "84A98C"
LIGHT = "CAD2C5"
CREAM = "FFFDE7"
GREY = "E9ECEC"
YELLOW = "FFFF00"
BLUE = "0000FF"
BLACK = "000000"
GREEN = "008000"
RED = "C00000"
NAVY = "1F3864"
WHITE = "FFFFFF"

FMT_NUM = '#,##0;\\(#,##0\\);\\-'
FMT_NUM1 = '#,##0.0;\\(#,##0.0\\);\\-'
FMT_PCT = '0.0%;\\(0.0%\\);\\-'
FMT_PCT0 = '0%;\\(0%\\);\\-'
FMT_EPS = '0.00;\\(0.00\\);\\-'
FMT_SHR = '#,##0.00;\\(#,##0.00\\);\\-'
FMT_X = '0.00\\x;\\(0.00"x)";\\-'
FMT_DAYS = '0\\d;\\(0")d";\\-'
FMT_CAGR = '0%;\\(0%\\);"n/m"'

FONT = "Calibri"


def font(sz=12, b=False, i=False, color=None):
    return Font(name=FONT, sz=sz, b=b, i=i, color=color)


def fill(rgb):
    return PatternFill("solid", fgColor=rgb)


THIN_TOP = Border(top=Side(style="thin"))

# --------------------------------------------------------------------------------------
# Columns
# --------------------------------------------------------------------------------------
ANNUAL_ONLY_YEARS = list(range(2006, 2013))   # 2006–2012: annual only
HALF_YEARS = list(range(2013, 2026))          # 2013–2025: H1, H2, FY
FC_YEAR = 2026                                # H1/26 (actual), H2/26E, 2026E
OUT_YEARS = [2027, 2028, 2029, 2030]


@dataclass
class Col:
    key: str          # "2006", "H1/13", "H2/13", "2013", "H1/26", "H2/26E", "2026E", "2027E"
    label: object
    kind: str         # A (annual-only hist) | H (hist H1) | H2 (hist H2 = FY − H1) | Y (hist FY)
                      # HE (forecast H2/26E) | E26 (2026E) | E (2027E+)
    year: int
    idx: int = 0
    hidden: bool = False

    @property
    def L(self):
        return get_column_letter(self.idx)

    @property
    def is_forecast(self):
        return self.kind in ("HE", "E26", "E")

    @property
    def is_annual(self):
        return self.kind in ("A", "Y", "E26", "E")

    @property
    def is_half(self):
        return self.kind in ("H", "H2", "HE")


def build_columns(first_idx=3):
    cols = []
    for y in ANNUAL_ONLY_YEARS:
        cols.append(Col(str(y), y, "A", y))
    for y in HALF_YEARS + [FC_YEAR]:
        yy = f"{y % 100:02d}"
        fc = y == FC_YEAR
        cols.append(Col(f"H1/{yy}", f"H1/{yy}", "H", y, hidden=not fc))
        if fc:
            cols.append(Col(f"H2/{yy}E", f"H2/{yy}E", "HE", y))
            cols.append(Col(f"{y}E", f"{y}E", "E26", y))
        else:
            cols.append(Col(f"H2/{yy}", f"H2/{yy}", "H2", y, hidden=True))
            cols.append(Col(str(y), y, "Y", y))
    for y in OUT_YEARS:
        cols.append(Col(f"{y}E", f"{y}E", "E", y))
    for i, c in enumerate(cols):
        c.idx = first_idx + i
    return cols


class Cols:
    def __init__(self):
        self.list = build_columns()
        self.by_key = {c.key: c for c in self.list}
        self.last = self.list[-1]
        n = self.last.idx
        self.NOTES = get_column_letter(n + 1)
        self.GAP1 = get_column_letter(n + 2)
        self.CAGR = [get_column_letter(n + 3 + i) for i in range(4)]   # 25-30, 20-25, 15-25, 06-25
        self.GAP2 = get_column_letter(n + 7)
        self.SPACER = get_column_letter(n + 8)
        self.SCN_LABEL = get_column_letter(n + 9)
        self.SCN = {s: get_column_letter(n + 10 + i) for i, s in enumerate(["Bull", "Base", "Bear"])}
        self.SCN_NOTE = get_column_letter(n + 13)
        self.SWITCH = f"${self.NOTES}$2"

    def get(self, key):
        return self.by_key[key]

    def annual(self, year):
        for k in (str(year), f"{year}E"):
            if k in self.by_key:
                return self.by_key[k]
        return None

    def h1(self, year):
        return self.by_key.get(f"H1/{year % 100:02d}")

    def h2(self, year):
        yy = f"{year % 100:02d}"
        return self.by_key.get(f"H2/{yy}") or self.by_key.get(f"H2/{yy}E")

    def prior_same(self, c: Col):
        """Same period one year earlier."""
        if c.is_annual:
            return self.annual(c.year - 1)
        if c.kind == "H":
            return self.h1(c.year - 1)
        if c.kind in ("H2", "HE"):
            return self.h2(c.year - 1)
        return None

    def prev_period(self, c: Col):
        """Immediately preceding period: year → prior year; H1 → prior year end; H2 → H1."""
        if c.is_annual or c.kind == "H":
            return self.annual(c.year - 1)
        if c.kind in ("H2", "HE"):
            return self.h1(c.year)
        return None


# --------------------------------------------------------------------------------------
# Row definitions
# --------------------------------------------------------------------------------------
@dataclass
class Row:
    key: Optional[str]
    label: str = ""
    style: str = "line"         # section | block | total | sub | line | pct | growth | check | memo | sched | blank | note
    fmt: str = FMT_NUM
    src: Optional[str] = None   # data key for historical hard-codes, e.g. "is.revenue"
    agg: str = "sum"            # how H2 derives from FY and H1: sum (H2 = FY − H1) | last (H2 = FY) | none
    hist: Optional[Callable] = None       # formula for all historical columns (overrides src)
    he: Optional[Callable] = None         # H2/26E
    e26: Optional[Callable] = None        # 2026E
    e: Optional[Callable] = None          # 2027E–2030E
    fc: Optional[Callable] = None         # shortcut: same formula for HE, E26 and E (unless given)
    annual_fc: bool = False     # forecast built annually (2026E from FY25 YE); H2/26E derived from 2026E − H1/26
    halves: bool = True         # False → leave half-year cells empty (annual-only data)
    note: str = ""
    comment: str = ""
    cagr: Optional[str] = None  # "cagr" | "avg" | None
    input_fc: bool = False      # forecast cells are inputs (blue bold)
    texts: Optional[dict] = None   # deftext rows: colkey -> text


class Ctx:
    """Formula context for one column."""

    def __init__(self, cols: Cols, rows: dict, c: Col):
        self.cols, self.rows, self.c = cols, rows, c

    def r(self, key):
        return self.rows[key]

    def a(self, key, col: Col = None):
        col = col or self.c
        return f"{col.L}{self.rows[key]}"

    def py(self, key):
        p = self.cols.prior_same(self.c)
        return None if p is None else f"{p.L}{self.rows[key]}"

    def pp(self, key):
        p = self.cols.prev_period(self.c)
        return None if p is None else f"{p.L}{self.rows[key]}"

    def at(self, key, colkey):
        return f"{self.cols.get(colkey).L}{self.rows[key]}"

    def abs_at(self, key, colkey):
        return f"${self.cols.get(colkey).L}${self.rows[key]}"

    def fy(self, key, year=None):
        return f"{self.cols.annual(year or self.c.year).L}{self.rows[key]}"

    def h1(self, key, year=None):
        return f"{self.cols.h1(year or self.c.year).L}{self.rows[key]}"

    def scn(self, scn_row):
        s = self.cols.SCN
        return (f'CHOOSE(MATCH({self.cols.SWITCH},{{"Bull","Base","Bear"}},0),'
                f'${s["Bull"]}${scn_row},${s["Base"]}${scn_row},${s["Bear"]}${scn_row})')

    def base(self, scn_row):
        return f'${self.cols.SCN["Base"]}${scn_row}'

    @property
    def half_mult(self):
        """Annualisation factor for stock/flow ratios in half-year columns."""
        return 2 if self.c.is_half else 1


def yoy(key):
    return lambda x: (f'=IF(AND(ISNUMBER({x.a(key)}),ISNUMBER({x.py(key)})),IFERROR({x.a(key)}/{x.py(key)}-1,""),"")'
                      if x.py(key) else None)


def ratio(num, den):
    return lambda x: f'=IF(AND(ISNUMBER({x.a(num)}),ISNUMBER({x.a(den)})),IFERROR({x.a(num)}/{x.a(den)},""),"")'


def put_comment(cell, text):
    if text:
        cm = Comment(text, "Model")
        cm.width, cm.height = 420, 160
        cell.comment = cm
