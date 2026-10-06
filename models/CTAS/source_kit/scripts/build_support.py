"""Adapt Bull-Base-Bear, DCF and Charts sheets of the MLM template to the CTAS Model sheet."""
import copy, re
from openpyxl.utils import get_column_letter as L
import mdl_core as M

def mref(R, key, period, absolute=False):
    p = M.BYNAME[period]
    return f"Model!${p.L}${R[key]}" if absolute else f"Model!{p.L}{R[key]}"

ANN_H = [f'FY{y}' for y in range(2006, 2027)]

def build_bbb(wb, R):
    ws = wb['Bull-Base-Bear']
    sw = f"Model!${L(M.NOTE_COL)}$2"
    ws['B2'] = 'Cintas Corporation (CTAS) — Bull / Base / Bear Case P&L Summary'
    ws['B3'] = ('Adjusted figures (company definition: adjusted diluted EPS excl. acquisition transaction & integration costs; model Adjusted EBITDA) · USD millions · '
                'Fiscal years end 31-May · Live panel linked to Model — toggle scenario via Model!' + L(M.NOTE_COL) + '2. Bull and Bear panels are value snapshots generated from the Model '
                'with the switch set to each case. UniFirst consolidated from an assumed 31-Dec-2026 close (toggle in the Model scenario table).')
    ws['K2'] = '=' + mref(R, 'px', 'FY2027', True)
    ws['C5'] = '=' + sw
    cols = ['C', 'D', 'E', 'F', 'G', 'H']
    pers = ['FY2026', 'FY2027', 'FY2028', 'FY2029', 'FY2030', 'FY2031']
    heads = ['FY2026A', 'FY2027E', 'FY2028E', 'FY2029E', 'FY2030E', 'FY2031E']
    for c, h in zip(cols, heads):
        ws[f'{c}7'] = h
    ws['I7'] = "26-'31 CAGR"
    for off in (12, 24):   # snapshot panels headers (O.. and AA..)
        for i, h in enumerate(heads):
            ws.cell(7, 15 + i if off == 12 else 27 + i).value = h
        ws.cell(7, 21 if off == 12 else 33).value = "26-'31 CAGR"
    links = {9: 'is_rev', 11: 'gm_org', 12: 'adj_ebitda', 15: 'is_oi', 19: 'a_ni', 21: 'sh_d', 23: 'a_eps', 31: 'fcf', 35: 'adj_ebitda_pf', 38: 'nd', 44: 'roic', 45: 'ronta'}
    for r, k in links.items():
        for c, p in zip(cols, pers):
            ws[f'{c}{r}'] = '=' + mref(R, k, p)
    for c, p in zip(cols[1:], pers[1:]):
        ws[f'{c}40'] = '=' + mref(R, 'bb_cash', p)
        ws[f'{c}41'] = '=' + mref(R, 'bb_px', p)
        ws[f'{c}42'] = '=' + mref(R, 'bb_sh', p)
    ws['B11'] = '   Organic revenue growth % (legacy Cintas; FY2027E ex-workday)'
    ws['B12'] = 'Adjusted EBITDA (model; excl. transaction & integration costs)'
    ws['B15'] = 'Operating Income (GAAP)'
    ws['B16'] = '   Operating income y/y %'
    ws['B19'] = 'Adjusted Income from Continuing Operations (company definition)'
    ws['B35'] = 'Adjusted EBITDA — pro forma (FY2027E incl. UniFirst full year)'
    ws['B29'] = 'IRR (to May-Y-1)'
    from copy import copy as _cp
    for r in (26, 27, 28, 29):
        ws[f'F{r}']._style = _cp(ws[f'G{r}']._style)
    ws['F26'] = 35; ws['G26'] = 35; ws['H26'] = 35
    ws['F27'] = '=F26*F23'; ws['F28'] = '=IFERROR(F27/$K$2-1,"")'
    ws['F29'] = '=IFERROR((F27/$K$2)^(1/((DATE(2029-1,5,31)-$K$3)/365.25))-1,"")'
    ws['B26'] = 'P / E (x) — assumed (adjusted EPS; CTAS 5-yr range ~30–50x)'
    ws['G29'] = '=IFERROR((G27/$K$2)^(1/((DATE(2030-1,5,31)-$K$3)/365.25))-1,"")'
    ws['H29'] = '=IFERROR((H27/$K$2)^(1/((DATE(2031-1,5,31)-$K$3)/365.25))-1,"")'
    for c in ('F', 'G', 'H'):
        ws[f'{c}48'] = f'={c}38'
    ws['B48'] = '(+) Net Debt'
    # historical stats J/K/L over FY2006-FY2026 annual columns
    stat_rows = {13: 'adj_ebitda_m', 17: 'gm_oim', 32: 'fcf_m', 36: 'gm_ebm', 44: 'roic', 45: 'ronta'}
    for r, k in stat_rows.items():
        refs = ','.join(mref(R, k, p) for p in ANN_H)
        ws[f'J{r}'] = f'=IFERROR(AVERAGE({refs}),"")'
        ws[f'K{r}'] = f'=IFERROR(MIN({refs}),"")'
        ws[f'L{r}'] = f'=IFERROR(MAX({refs}),"")'
    return ws

def build_dcf(wb, R):
    ws = wb['DCF']
    ws['B2'] = 'DCF VALUATION — CINTAS CORPORATION (CTAS)'
    ws['C5'] = 0.045; ws['D5'] = 'Assumption (~US 10Y Treasury yield); update to market'
    ws['C6'] = 0.05; ws['D6'] = 'US implied ERP ~4.5–5.5% (Damodaran)'
    ws['C7'] = 0.95; ws['D7'] = 'Business-services / uniform peers (UniFirst, Aramark, Vestis, Rollins, ABM) ~0.8–1.1; recurring route-based revenue and >90% retention dampen cyclicality'
    ws['C9'] = 0.05; ws['D9'] = 'A3/A- rated; UniFirst acquisition debt assumed at 5.0% (424B3); 4.0% 2032 and 6.15% 2036 notes outstanding'
    ws['C10'] = 0.21; ws['D10'] = 'Model FY2028E+ effective tax rate'
    ws['C12'] = 0.06; ws['D12'] = 'Net debt ~$4.3bn at FY2027E YE (post-UniFirst) vs ~$82bn market cap ($198.95 × ~413m shares) — ~5% D/V; 6% target'
    ws['D15'] = 'Periods = (31-May-FY − today) / 365.25; cash flows assumed at fiscal year-end'
    heads = ['FY2027E', 'FY2028E', 'FY2029E', 'FY2030E', 'FY2031E', 'FY2032E', 'FY2033E', 'FY2034E', 'FY2035E', 'FY2036E']
    dc = ['D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L', 'M']
    for c, h in zip(dc, heads): ws[f'{c}17'] = h
    ws['P17'] = 'CAGR 27–31E'; ws['Q17'] = 'CAGR 27–36E'
    pers = ['FY2027', 'FY2028', 'FY2029', 'FY2030', 'FY2031']
    for c, p in zip(dc[:5], pers):
        ws[f'{c}18'] = f"={mref(R, 'adj_ebit', p)}-{mref(R, 'is_trans', p)}"
        ws[f'{c}21'] = '=' + mref(R, 'is_da', p)
        ws[f'{c}23'] = '=' + mref(R, 'cf_wc', p)
        ws[f'{c}22'] = ('=' + mref(R, 'cf_capex', p)) if p == 'FY2027' else f"={mref(R, 'cf_capex', p)}+{mref(R, 'cf_acq', p)}"
    ws['I22'] = f"={mref(R, 'cf_capex', 'FY2031')}*(1+I$26)"
    ws['B18'] = 'Adj. EBIT less transaction & integration costs'
    ws['B21'] = 'plus D&A (incl. UniFirst purchase accounting)'
    ws['B22'] = '− Capex and bolt-on acquisitions (M&A FY2028E–31E)'
    ws['B26'] = 'Fade-period growth (FY2032E–36E; 7% → 5% linear)'
    ws['I26'] = 0.07; ws['M26'] = 0.05; ws['C27'] = 0.04
    for i, c in enumerate(dc):
        ws[f'{c}29'] = f'=(DATE({2027 + i},5,31)-$C$15)/365.25'
    ws['D32'] = ("Cintas' adjusted EBIT expenses stock-based compensation and deducts all D&A incl. ~$98m p.a. of UniFirst PPA amortization; "
                 "transaction & integration costs deducted as cash costs. UniFirst consideration is captured through FY2027E net debt and the 14.07m shares issued.")
    ws['D34'] = f"={mref(R, 'r_ic', 'FY2026')}+{mref(R, 'ppa_cash', 'FY2027')}+{mref(R, 'ppa_eq', 'FY2027')}"
    ws['D43'] = 'Opening IC = Model FY2026A book IC (equity + net debt) + UniFirst total consideration ($5.35bn cash & shares), rolled forward with capex + M&A − D&A + ΔNWC.'
    ws['B48'] = 'Less: Net debt (Model FY2027E YE, post-UniFirst)'
    ws['C48'] = '=' + mref(R, 'nd', 'FY2027')
    ws['D48'] = 'FY2027E year-end net debt includes the UniFirst cash consideration and $2.8bn acquisition debt; no non-controlling interests.'
    ws['B50'] = 'Diluted shares (m, Model FY2027E incl. UniFirst shares)'
    ws['C50'] = '=' + mref(R, 'v_sh', 'FY2027')
    ws['C52'] = '=' + mref(R, 'v_px', 'FY2027')
    ws['D76'] = 'Rationale (Cintas characteristics)'
    ws['C78'] = 0.10; ws['D78'] = 'Model adjusted EBIT CAGR FY2027–31E: high-single-digit organic growth, margin expansion and UniFirst synergies → ~10% NOPAT growth'
    ws['C79'] = 0.25; ws['D79'] = 'Route density and garment reuse earn ~25–30% ROIC on organic investment (FY2026 ROIC ~29%); acquired growth earns less (goodwill)'
    ws['C80'] = 15; ws['D80'] = 'Largest North American route network (~1.5m customers combined), high retention, cross-sell runway (first aid, fire, hygiene); regulatory moats in fire inspection'
    ws['B90'] = 'VALUATION  —  Applied to FY2029E NOPAT'
    ws['B91'] = 'FY2029E Adj. EBIT less transaction & integration costs'
    ws['C91'] = f"={mref(R, 'adj_ebit', 'FY2029')}-{mref(R, 'is_trans', 'FY2029')}"
    ws['B93'] = 'FY2029E NOPAT'
    ws['B95'] = 'Implied Enterprise Value (FY2029E basis)'
    ws['B96'] = 'Less: FY2028E YE net debt (Model)'
    ws['C96'] = '=' + mref(R, 'nd', 'FY2028')
    ws['B98'] = 'Diluted shares FY2029E (m, Model)'
    ws['C98'] = '=' + mref(R, 'v_sh', 'FY2029')
    ws['B99'] = 'Implied share price FY2029 ($)'
    ws['B106'] = 'Plus: FY2027E YE net debt'
    ws['B108'] = 'CTAS market-implied EV/NOPAT (on FY2029E NOPAT)'
    return ws

def build_charts(wb, R):
    ws = wb['Charts']
    ws['B2'] = 'Cintas Corporation — Historical Operating Performance (FY2006–FY2026)'
    ws['B3'] = ('All historicals · USDm · fiscal years ending 31-May · Linked to Model. Organic growth = company-published organic revenue growth '
                '(FY2007–09: reported growth — organic not disclosed in releases); acquisitions/FX/workdays = reported growth − organic.')
    # extend from 20 columns (C..V) to 21 (C..W): copy column V styles into W
    for r in range(1, ws.max_row + 1):
        src = ws.cell(r, 22); dst = ws.cell(r, 23)
        if src.has_style:
            dst._style = copy.copy(src._style)
    ws.column_dimensions['W'].width = ws.column_dimensions['V'].width or 9
    for i, p in enumerate(ANN_H):
        c = L(3 + i)
        ws[f'{c}7'] = p
        ws[f'{c}11'] = f'={c}7'; ws[f'{c}17'] = f'={c}7'
        ws[f'{c}8'] = f"=IF(ISNUMBER({mref(R, 'og_org', p)}),{mref(R, 'og_org', p)},IF(ISNUMBER({mref(R, 'og_rep', p)}),{mref(R, 'og_rep', p)},0))"
        ws[f'{c}9'] = f"=IF(ISNUMBER({mref(R, 'og_adj', p)}),{mref(R, 'og_adj', p)},0)"
        ws[f'{c}12'] = '=' + mref(R, 'adj_ebitda', p)
        ws[f'{c}13'] = '=' + mref(R, 'is_oi', p)
        ws[f'{c}14'] = '=' + mref(R, 'gm_ebm', p)
        ws[f'{c}15'] = '=' + mref(R, 'gm_oim', p)
        ws[f'{c}18'] = f"=IFERROR({mref(R, 'roic', p)}+0,NA())"
        ws[f'{c}19'] = f"=IFERROR({mref(R, 'ronta', p)}+0,NA())"
        ws[f'{c}20'] = f"=IFERROR({mref(R, 'gm_ebm', p)}+0,NA())"
        ws[f'{c}21'] = f"=IFERROR({mref(R, 'fcf_m', p)}+0,NA())"
    ws['B8'] = 'Organic revenue growth %'
    ws['B9'] = 'Acquisitions, divestitures, FX & workdays %'
    ws['B13'] = 'Operating Income (GAAP)'
    ws['B15'] = 'GAAP Op. Margin %'
    titles = ['Revenue Drivers: Organic Growth vs. Acquisitions / Divestitures / FX / Workdays, y/y %',
              'Adjusted EBITDA & Operating Income (USDm)',
              'Returns Profile: ROIC, RONTA, Adjusted EBITDA Margin & FCF Margin']
    for ch, t in zip(ws._charts, titles):
        try:
            ch.title.tx.rich.p[0].r[0].t = t
        except Exception:
            ch.title = t
        for s in ch.series:
            for ref in (s.val.numRef if s.val is not None else None, s.cat.numRef if (s.cat is not None and s.cat.numRef is not None) else None,
                        s.cat.strRef if (s.cat is not None and s.cat.strRef is not None) else None):
                if ref is not None and ref.f:
                    ref.f = ref.f.replace('$V$', '$W$')
    return ws
