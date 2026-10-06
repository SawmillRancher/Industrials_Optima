import openpyxl,json,sys
from openpyxl.utils import get_column_letter as L
f=sys.argv[1]
wb=openpyxl.load_workbook(f,data_only=True)
ws=wb['Model']
R=json.load(open('rowmap.json'))
inv={v:k for k,v in R.items()}
errs={}
for row in ws.iter_rows(min_row=1,max_row=ws.max_row):
    for c in row:
        if isinstance(c.value,str) and c.value.startswith('#') or (isinstance(c.value,str) and c.value in ('Err:502','Err:504','Err:508','Err:522','Err:523')):
            errs.setdefault(inv.get(c.row,c.row),[]).append(c.coordinate)
print('ERRORS:',{k:(len(v),v[:6]) for k,v in errs.items()})
hdr={c:ws.cell(6,c).value for c in range(3,90)}
for k in [k for k in R if k.startswith('chk') or k.endswith('_chk') or k in ('e_chk','a_chk','fcf_chk','r_chk','ppa_chk')]:
    r=R[k]; bad=[]
    for c in range(3,89):
        v=ws.cell(r,c).value
        if isinstance(v,(int,float)) and abs(v)>(0.005 if k in('a_chk','r_chk') else 0.05): bad.append((hdr[c],round(v,3)))
    if bad: print('CHECK FAIL',k,bad[:40])
