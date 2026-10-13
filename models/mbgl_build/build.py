"""Build MBGL_Model.xlsx (Mobility Global) from the MLM template (same layout, formatting and process as the prior models; kit adapted from wwd_build)."""
import sys, copy, os, json
from openpyxl.styles import Font
from fw import *
import fw
from spec import build_spec, scen_layout, SCEN_ROWS, SC
from spec2 import build_spec2

OUT = os.path.join(os.path.dirname(BASE), 'MBGL_Model.xlsx')
G26 = MAN['guidance_fy26']
AM = MAN['amort_sched']

# ====================================================================== inputs (V0)
def wc_inputs():
    x = DATA['2025']['x']
    rev = x['rev']
    return {'dso': round(x['bs_ar'] / rev * 365), 'ap': round(x['bs_ap'] / rev, 3), 'ue': round(x['bs_unearned'] / rev, 3),
            'ocl': round((x['bs_ocl'] - DATA['2025']['bs']['due_to_c']) / rev, 3)}
WC = wc_inputs()
WC['note'] = (f'Forecast inputs at fiscal-year-end (31-Dec-2025) levels on FY2025 revenue: DSO {WC["dso"]}d, payables {WC["ap"]:.1%}, unearned revenue {WC["ue"]:.1%}, '
              f'accrued & other current liabilities {WC["ocl"]:.1%} (excluding the $19m due to S&P Global, settled at the separation). Year ends compared with year ends.')
YRS = list(range(2027, 2032))

V0 = {
    'out_hdr': 'FY26 outlook — Q2-26 release (7-Aug-2026)',
    'out_note': ('First guidance as a public company: revenue $1,870–1,885m (+6.9–7.7%), Adjusted EBITDA $745–760m (~40% margin); 1H carve-out + 2H consolidated. '
                 'No EPS / interest / tax / capex / FCF guidance. "Medium-term financial targets" referenced in the release are not filed on EDGAR (unverified; not used).'),
    'outlook': [
        ('out_rev', '  FY26 revenue ($m)', (1885.0, 1877.5, 1870.0), 'm', 'Guided $1,870–1,885m. Drives the Q3/Q4-26E revenue calibration.'),
        ('out_ebitda', '  FY26 Adjusted EBITDA ($m)', (760.0, 752.5, 745.0), 'm', 'Guided $745–760m. Drives the Q3/Q4-26E segment-margin calibration.'),
        ('bb26', '  2H/26 share repurchases ($m)', (0.0, 0.0, 0.0), 'm', 'No repurchase authorization announced (10-Q, 8-Ks). Dividend $0.06 a quarter.'),
    ],
    'points': [
        ('pt_int', '  2H/26 interest expense, net ($m)', 54.0, 'm', ('Assumption: coupons on $2.0bn notes $27.7m a quarter + ~$0.8m accretion/fees − interest on ~$250m cash. '
                                                               'Form 10 pro forma: $114m p.a. gross (5.74% incl. amortization).')),
        ('pt_etr', '  2H/26 effective tax rate', 0.29, 'pct', 'Assumption: 1H/26 29.4% (Q2 29.3%); not guided.'),
        ('pt_capex', '  FY26 capital expenditures ($m)', 26.0, 'm', 'Assumption: 1H/26 $12m; FY2025 $24m.'),
    ],
    'lev': (2.5, 2.0, 1.5),
    'lev_note': ('Assumption — no leverage target published. Net debt ~$1.8bn at 30-Jun-26 ≈ 2.5x LTM Adjusted EBITDA ($740m); covenant ≤ 3.50x; notes rated investment grade '
                 '(change-of-control put on a downgrade below IG). Target = leverage at which surplus cash is returned via buybacks; Bull = more buybacks (higher target).'),
    'blocks': [
        ('CFX_G', 'CFX_G (CARFAX revenue y/y)', 0.02, -0.03, [0.080, 0.078, 0.075, 0.072, 0.070],
         'FY2024 +12%, FY2025 +10%, 1H/26 +8%: dealer subscriptions (Advantage, Listings, CARFAX for Life), BIG data customers, Showroom / Homegrown launches, CARFAX Germany.', 'pct', False),
        ('B2B_G', 'B2B_G (B2B revenue y/y)', 0.02, -0.03, [0.050, 0.050, 0.045, 0.045, 0.040],
         'FY2024 +3%, FY2025 +6%, 1H/26 +6% (Q2 +4%): OEM / dealer marketing (automotiveMastermind), Polk data, Market Scan; non-subscription (recalls, marketing) cyclical.', 'pct', False),
        ('CFX_M', 'CFX_M (CARFAX Segment Adjusted EBITDA margin)', 0.010, -0.025, [0.480, 0.485, 0.490, 0.495, 0.500],
         'FY2023 45.5%, FY2024 45.8%, FY2025 46.9%, 1H/26 48.0%; operating leverage on the data estate vs advertising / strategic investment.', 'pct', False),
        ('B2B_M', 'B2B_M (B2B Segment Adjusted EBITDA margin)', 0.015, -0.030, [0.335, 0.340, 0.345, 0.350, 0.355],
         'FY2023 34.3%, FY2024 35.4%, FY2025 32.4% (2025 allocation refinement), 1H/26 32.6%; "One Mobility Global" integration and AI workflow savings.', 'pct', False),
        ('MNA', 'M&A_spend (tuck-in acquisitions, $m)', 100.0, 0.0, [0.0, 0.0, 0.0, 0.0, 0.0],
         'None announced; the Form 10 cites inorganic opportunities in a fragmented industry (Market Scan Feb-2023, $214m). Base nil; Bull $100m p.a. (assumption).', 'm', True),
    ],
    'corp_q': {3: -12.0, 4: -12.0},
    'corp_note': ('History: corporate unallocated (S&P Global allocations; FY2025 −$22m, 1H/26 −$9m). Q3/Q4-26E input (assumption): −$12m a quarter = 1H/26 run-rate '
                  '(−$4.5m) + standalone public-company costs. The Form 10 pro forma quantifies only $12m p.a. (TSA $3m + new corporate leases $9m); the rest '
                  '(board, audit, investor relations, treasury, insurance) is an assumption. 2027E+: % of revenue.'),
    'corp_pct': {2027: -0.024, 2028: -0.023, 2029: -0.022, 2030: -0.021, 2031: -0.020},
    'opex_pct': 0.295, 'dep_pct': 0.008, 'etr': 0.285, 'capex_pct': 0.014,
    'amort_ae': {2027: AM['2027'], 2028: AM['2028'], 2029: AM['2029'], 2030: AM['2030'], 2031: 290.0},
    'amort_note': AM['src'] + ' 2026–29 $296m, 2030 $293m; 2031E $290m assumed (databases & software largely amortized by ~2032).',
    'sbc_q': {3: 6.0, 4: 6.0}, 'sbc_pct': 0.013,
    'sbc_note': 'History: S&P Global awards allocated (FY2025 $22m). Q3/Q4-26E $6m a quarter (assumption: spin-date replacement and new MBGL awards, Form 4 grants Jul / Sep-2026). 2027E+: % of revenue.',
    'trans_q': {3: 10.0, 4: 10.0}, 'trans_ae': {2027: 10.0, 2028: 0.0, 2029: 0.0, 2030: 0.0, 2031: 0.0},
    'trans_note': ('Stand-up of the spin business (company adjusting item): FY2025 $21m, 1H/26 $57m (carve-out). Costs after the separation are borne by Mobility Global (Form 10). '
                   'Q3/Q4-26E $10m a quarter and 2027E $10m are assumptions; nil thereafter.'),
    'oth_q': {3: 2.0, 4: 2.0}, 'oth_ae': 8.0,
    'dps_q': {3: 0.06, 4: 0.06}, 'dps_g': {2027: 1.0, 2028: 0.05, 2029: 0.05, 2030: 0.05, 2031: 0.05},
    'dil_q': 1.0, 'dil_ae': 1.5, 'iss_2h': 0.2, 'iss_ae': 0.8,
    'def_2h': -38.0, 'def_ae': -80.0,
    'def_note': ('Deferred-tax benefit from book amortization of acquired intangibles (FY2025 −$90m, 1H/26 −$38m). 2H/26E −$38m and 2027E+ −$80m p.a. are assumptions '
                 '(~27% × $296m amortization).'),
    'wc': WC, 'min_cash': 200.0,
    'iss_notes': {2027: 0.0, 2028: 0.0, 2029: 650.0, 2030: 0.0, 2031: 650.0},
    'repay': {2027: 0.0, 2028: 0.0, 2029: 650.0, 2030: 0.0, 2031: 650.0},
    'repay_note': ('5.050% notes due 2029 ($650m) and 5.450% notes due 2031 ($650m) assumed refinanced at maturity (assumption; modelled at the year end; '
                   'maturity months are not given in the 10-Q).'),
    'accr_2h': 1.5, 'accr_ae': 3.0,
    'r_notes': {2027: 0.0553, 2028: 0.0553, 2029: 0.0553, 2030: 0.0568, 2031: 0.0568},
    'r_notes_note': ('Weighted coupon on $2.0bn: (650 × 5.050% + 650 × 5.450% + 700 × 6.050%) / 2,000 = 5.53%; from 2030E the 2029 notes refinanced at ~5.5% '
                     '(assumption) → 5.68%. Accretion of OID / issuance costs added separately.'),
    'px': MAN['price']['px'],
    'px_note': ('VALUATION — share price input $18.00 (provided by the user, 13-Oct-2026; latest SEC-filed price: Form 4 withholding $19.83 on 1-Sep-2026). '
                'Update to market.'),
    'mna_note': 'Scenario lever M&A_spend (Base nil; no acquisitions announced). Funded from cash / revolver (debt schedule).',
    'rev_guid_note': 'Company FY26 revenue guidance $1,870–1,885m (Q2-26 release, 7-Aug-2026). Drives the revenue calibration.',
    'eb_guid_note': 'Company FY26 Adjusted EBITDA guidance $745–760m (Q2-26 release). Drives the segment-margin calibration.',
}

# ====================================================================== compute engine
def hist_value(row, col):
    kw = row.kw; path = kw.get('data')
    if not path: return None
    if kw.get('ppa'):
        v = MAN['ppa'].get(pkey(col), {}).get(path.split('_', 1)[1])
        return v if kw.get('textdata') else num(v)
    v = xget(pkey(col), path)
    return v if kw.get('textdata') else num(v)

def first(kw, keys):
    for k in keys:
        if k in kw: return k
    return None

def compute(row, col):
    kw = row.kw; k = col.kind
    CTX.col = col
    if kw.get('ann') and col.q: return None, False
    if k in HISTK:
        v = hist_value(row, col)
        if v is not None: return v, True
        for key in ('hist', 'all'):
            if key in kw:
                return kw[key](), False
        return None, False
    if k == 'QE':
        if row.spec2: return None, False
        if 'qe_in' in kw:
            v = kw['qe_in']; return (v[col.q] if isinstance(v, dict) else v), True
        key = first(kw, ('qe', 'fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    if k == 'E26':
        if 'e26_in' in kw: return kw['e26_in'], True
        if 'e26' in kw: return kw['e26'](), False
        if kw.get('avg') and ('qe' in kw or 'qe_in' in kw): return f'=AVERAGE({Q1}{row.r}:{QE}{row.r})', False
        if kw.get('flow') and not row.spec2 and any(x in kw for x in ('qe', 'qe_in', 'fc')): return S26(row.key), False
        key = first(kw, ('fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    if k == 'AE':
        if 'ae_in' in kw:
            v = kw['ae_in']; return (v[col.year] if isinstance(v, dict) else v), True
        key = first(kw, ('ae', 'fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    return None, False

# ====================================================================== writers
CAGR_G = {'is_rev', 'g_rev', 'cfx_rev', 'b2b_rev', 'cfx_adj', 'b2b_adj', 'g_adj', 'r_adj', 'r_ebitda', 'is_op', 'is_ni', 'o_adj', 'e_adjni', 'e_eps',
          'cf_cfo', 'fcf', 'bs_eq', 'sh_dil', 'eps_dil', 'dps', 'rv_splice', 'sp_rev', 'sp_adjop', 'd_amort', 'd_dep', 'is_sga', 'is_opex'}
CAGR_A = {'cfx_m', 'b2b_m', 'g_adjm', 'o_adj_m', 'r_adj_m', 'is_opm', 'r_ebit_m', 'gm_rev', 'gm_cfx', 'gm_b2b', 'gm_ebitda_m', 'gm_op_m', 'gm_net_m', 'gm_sga',
          'gm_adjop_m', 'fcf_m', 'fcf_conv', 'capex_pct', 'ra_roic', 'rn_ronta', 'roe', 'lv_x', 'd_opex', 'd_dep_pct', 'd_sbc_pct', 'wc_nwc_p', 'is_etr',
          'fcf_eb', 'corp_pct', 'cfx_subp', 'b2b_subp', 'sp_adjop_m', 'rv_splice_g', 'sp_rev_g'}
A = {y: ANN[y].c for y in ANN}

def style_cagr(ws, r, on):
    src = 16 if on else 12
    for c in CAGR:
        ws[f'{c}{r}']._style = copy.copy(TM[f'{L(CI(c) + TSHIFT)}{src}']._style)

def write_row(ws, row):
    r = row.r; arch = ARCH[row.arch]
    clone_row_style(ws, arch, r)
    ws.cell(r, 2).value = row.label or None
    if row.kw.get('note'): ws[f'{NOTE}{r}'] = row.kw['note']
    style_cagr(ws, r, row.key in CAGR_G or row.key in CAGR_A)
    if row.arch in ('section', 'block', 'blank', 'sub') and not row.kw.get('data'):
        return
    if row.key in DEFS:
        write_defs(ws, row); return
    nf = NF[row.nf]
    for col in COLS:
        v, inp = compute(row, col)
        if v is None: continue
        cell = ws[f'{col.c}{r}']
        cell.value = v
        if not row.kw.get('textdata'): cell.number_format = nf
        if inp:
            set_color(cell, BLUE, bold=True if col.kind in FCK else None)
            lk = row.kw.get('labkey')
            if lk and col.kind in HISTK and isinstance(v, (int, float)) and abs(v) > 1e-9:
                t = xget(pkey(col), lk)
                if t: comment(cell, f'Company items ({pkey(col)}): {t}', 380, 110)
        elif row.arch == 'growth' and col.kind in HISTK:
            set_color(cell, GRAY)
        else:
            set_color(cell, BLACK)
    rc = CELL_COMMENTS.get(row.key, {})
    for c, t in rc.items(): comment(ws[f'{c}{r}'], t)
    if row.key in CAGR_G:
        ws[f'{CAGR[0]}{r}'] = f'=IFERROR(({A[2031]}{r}/{A[2026]}{r})^(1/5)-1,"n/m")'
        ws[f'{CAGR[1]}{r}'] = f'=IFERROR(({A[2025]}{r}/{A[2023]}{r})^(1/2)-1,"n/m")'
        ws[f'{CAGR[2]}{r}'] = f'=IFERROR({A[2025]}{r}/{A[2024]}{r}-1,"n/m")'
        ws[f'{CAGR[3]}{r}'] = f'=IFERROR(({A[2025]}{r}/{A[2022]}{r})^(1/3)-1,"n/m")'
        for c in CAGR: ws[f'{c}{r}'].number_format = '0%;\\(0%\\);"n/m"'
    elif row.key in CAGR_A:
        ws[f'{CAGR[0]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2027, 2032))}),"n/m")'
        ws[f'{CAGR[1]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2023, 2026))}),"n/m")'
        ws[f'{CAGR[2]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2024, 2026))}),"n/m")'
        ws[f'{CAGR[3]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2022, 2026))}),"n/m")'
        for c in CAGR: ws[f'{c}{r}'].number_format = NF['pct']
    t = LABEL_SOURCES.get(row.key)
    if t: comment(ws.cell(r, 2), t)

F10 = 'Source: Form 10 information statement (10-12B/A Ex. 99.1, 27-May-2026) — '
LABEL_SOURCES = {
    'cfx_rev': (F10 + 'MD&A segment results (FY2023–25, Q1-25, Q1-26); Q2-25 / Q2-26: 10-Q MD&A and release Exhibit 4; Q3-25 / Q4-25: FY2025 quarterly '
                'financial data (8-K Ex. 99.2, 7-Aug-2026, "unaudited and preliminary"). Two segments since the 2025 realignment (all periods recast by the company).'),
    'cfx_adj': F10 + 'reconciliation of segment operating profit to Segment Adjusted EBITDA; release Exhibit 5 (Q2); Ex. 99.2 (Q3-25 / Q4-25).',
    'corp_adj': F10 + 'Corporate column of the Segment Adjusted EBITDA reconciliation (corporate unallocated expense).',
    'sp_rev': ('Source: S&P Global Inc. (CIK 64040) earnings releases, 8-K Item 2.02 Ex. 99.1 — "Mobility" segment revenue / expenses / operating profit tables, '
               'adjusted segment operating profit reconciliations and revenue-by-type tables (Q1-23 – Q2-26). FY2022: S&P Global non-GAAP pro forma full year '
               '(IHS Markit merger 28-Feb-2022; reported Mar–Dec revenue $1,142m).'),
    'is_rev': (F10 + 'audited combined statements of income (FY2023–25) and unaudited Q1-26 / Q1-25; 10-Q (Q2-26 / Q2-25); 8-K Ex. 99.2 (Q3-25 / Q4-25). '
               'Carve-out (combined) basis: includes S&P Global corporate allocations; whole $ millions as published.'),
    'd_dep': F10 + 'statements of income (depreciation and amortization of intangibles shown separately; Q3-25 / Q4-25 only combined in Ex. 99.2).',
    'r_adj_pub': F10 + 'MD&A "reconciliation of our net income to Adjusted EBITDA"; release Exhibit 5 (Q2-26 / Q2-25); Ex. 99.2 (2025 quarters).',
    'o_pub': F10 + 'MD&A reconciliation of operating profit to Adjusted operating profit (published for FY2023–25, Q1-25 and Q1-26 only).',
    'cf_ni': (F10 + 'combined statements of cash flows (FY2023–25; Q1 three months); 10-Q six months (Q2 = 1H − Q1). Q3-25 / Q4-25 not published.'),
    'bs_cash': ('Source: combined balance sheets — 31-Dec-2023: first Form 10 draft (DRS, 23-Oct-2025); 31-Dec-2024 / 2025 and 31-Mar-2026: Form 10; '
                '30-Jun-2026: 10-Q. Pre-spin equity = parent company investment.'),
    'sh_basic': 'Source: 10-Q statements of operations and Note 3 (EPS) — 294,821,320 shares basic and diluted for all periods presented; cover page (1-Jul-2026).',
    'ds_face': 'Source: 10-Q Note 4 (Debt); 8-K 2-Jul-2026 (separation); Form 10 pro forma note (b).',
    'fcf_pub': F10 + 'Free cash flow reconciliation; release "Computation of Free Cash Flow" (Q2).',
    'dps': 'Source: Q2-26 release (dividend declared, payable 10-Sep-2026) and 10-Q MD&A ("Common Stockholder Dividends": $0.06 per share quarterly).',
}
CELL_COMMENTS = {}
def comments_setup():
    for c in (QC[(2025, 3)].c, QC[(2025, 4)].c):
        CELL_COMMENTS.setdefault('is_rev', {})[c] = ('8-K Ex. 99.2 (7-Aug-2026): "Pre-separation quarterly information, including Q3 and Q4 2025, is unaudited and preliminary, '
                                                     'and reflects the unaudited carve-out financials ... as reported by S&P Global Inc." Ties to FY2025 (Form 10) in total.')
    CELL_COMMENTS['sp_rev'] = {A[2022]: 'S&P Global non-GAAP pro forma revenue for the full year 2022 (as if IHS Markit had been owned from 1-Jan-2022); reported revenue (Mar–Dec) $1,142m.'}
    CELL_COMMENTS['sp_op'] = {A[2022]: 'Pro forma operating profit (S&P Global Q4-22 release).'}
    CELL_COMMENTS['bs_cash'] = {A[2023]: 'From the first Form 10 draft (DRS, 23-Oct-2025): FY2024 / FY2023 audited combined balance sheets.'}
    CELL_COMMENTS['cf_par'] = {QC[(2026, 2)].c: '1H/26 net transfers to Parent −$2,011m incl. the $2.0bn dividend paid to S&P Global from the senior-note proceeds (10-Q Note 4).'}
    CELL_COMMENTS['o_pub'] = {A[2024]: 'FY2024 total as published ($617m); the segment columns add to $618m (company rounding).'}

DEFS = {}
def defs_setup():
    q = lambda y, n: QC[(y, n)].c
    DEFS['r_def'] = {
        A[2023]: ('Form 10 definition: NI + interest + taxes + D&A + SBC + separation transaction costs + severance + acquisition integration + acquisition / disposition costs '
                  '+ legal recovery + other (other income kept in EBITDA)', 'Form 10 MD&A, "Adjusted EBITDA and Adjusted EBITDA Margin".'),
        q(2026, 2): ('Q2-26 release: same items; severance and other combined ("employee severance charges and other")', None),
    }
    DEFS['o_def'] = {
        A[2023]: ('Company Adjusted operating profit (Form 10 only): OP + amortization + severance + transaction + integration + ELT transition + acq./disposition + legal recovery + other',
                  'Not presented in the Q2-26 release; the model bridge continues the definition.'),
    }
    DEFS['e_def'] = {
        A[2023]: ('Model definition — Mobility Global publishes no adjusted net income / EPS; shown only where a share count exists (Q2-25, Q2-26, forecast)', None),
    }
    DEFS['rv_basis'] = {
        A[2022]: ('S&P Global pro forma (full year)', None),
        q(2023, 1): ('S&P Global segment (2023–24 quarters)', None),
        A[2023]: ('Carve-out (Form 10)', None),
        q(2025, 1): ('Carve-out (Form 10 / 10-Q / Ex. 99.2)', None),
    }
def write_defs(ws, row):
    for c, (txt, cm) in DEFS[row.key].items():
        cell = ws[f'{c}{row.r}']; cell.value = txt
        cell.font = Font(name='Calibri', sz=10, b=True, i=True, color='FFC00000')
        if cm: comment(cell, cm)

def write_scen(ws):
    srcrow = {'hdr': 9, 'out_hdr': 11, 'pt_hdr': 15, 'out': 12, 'pt': 16, 'lev': 28, 'blk': 30, 'yr': 31}
    for spec_ in SCEN_ROWS:
        r, kind = spec_[0], spec_[1]
        for ci in range(CI(SL), CI(SNOTE) + 1):
            ws.cell(r, ci)._style = copy.copy(TM.cell(srcrow[kind], ci + TSHIFT)._style)
        if kind in ('hdr', 'out_hdr', 'pt_hdr'):
            _, _, a, b, c_, d_, e = spec_
            for col, v in zip([SL, SBULL, SBASE, SBEAR, SNOTE], [a, b, c_, d_, e]):
                if v is not None: ws[f'{col}{r}'] = v
        elif kind == 'out':
            _, _, lab, hi, mid, lo, note, nf = spec_
            ws[f'{SL}{r}'] = lab
            for col, v in zip([SBULL, SBASE, SBEAR], [hi, mid, lo]):
                ws[f'{col}{r}'] = v; ws[f'{col}{r}'].number_format = NF[nf]; set_color(ws[f'{col}{r}'], BLUE)
            ws[f'{SNOTE}{r}'] = note
        elif kind == 'pt':
            _, _, lab, _, val, _, note, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBASE}{r}'] = val; ws[f'{SBASE}{r}'].number_format = NF[nf]; set_color(ws[f'{SBASE}{r}'], BLUE); ws[f'{SNOTE}{r}'] = note
        elif kind == 'lev':
            _, _, lab, a, b, c_, note = spec_
            ws[f'{SL}{r}'] = lab
            for col, v in zip([SBULL, SBASE, SBEAR], [a, b, c_]):
                ws[f'{col}{r}'] = v; ws[f'{col}{r}'].number_format = NF['x']
            ws[f'{SNOTE}{r}'] = note
        elif kind == 'blk':
            _, _, lab, db, mid, de, note, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBULL}{r}'] = db; ws[f'{SBASE}{r}'] = mid; ws[f'{SBEAR}{r}'] = de; ws[f'{SNOTE}{r}'] = note
            for col in (SBULL, SBEAR):
                ws[f'{col}{r}'].number_format = '\\+#,##0;\\-#,##0' if nf == 'm' else '\\+0.0%;\\-0.0%'
        elif kind == 'yr':
            _, _, lab, hdr, base, clamp, _, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBASE}{r}'] = base
            if clamp:
                ws[f'{SBULL}{r}'] = f'=MAX(0,{SBASE}{r}+${SBULL}${hdr})'; ws[f'{SBEAR}{r}'] = f'=MAX(0,{SBASE}{r}+${SBEAR}${hdr})'
            else:
                ws[f'{SBULL}{r}'] = f'={SBASE}{r}+${SBULL}${hdr}'; ws[f'{SBEAR}{r}'] = f'={SBASE}{r}+${SBEAR}${hdr}'
            for col in (SBULL, SBASE, SBEAR): ws[f'{col}{r}'].number_format = NF[nf]

def write_header(ws):
    for r in range(1, 8):
        clone_row_style(ws, r, r, 2, CI(LASTC))
    ws[f'{NOTE}1'] = 'SCENARIO SWITCH ▼ (Bull / Base / Bear) — drives all scenario-linked forecast drivers'
    ws[f'{NOTE}2'] = 'Base'
    comment(ws[f'{NOTE}2'], 'Scenario toggle. Choose Bull / Base / Bear. Drives: FY26 revenue and Adjusted EBITDA outlook end-points (high / mid / low), '
            '2H/26 buybacks, 2027E–2031E segment revenue growth and Segment Adjusted EBITDA margins, M&A spend and the leverage target for 2027E+ buybacks '
            f'(scenario table {SL}:{SNOTE}).')
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type='list', formula1='"Bull,Base,Bear"', allow_blank=False); ws.add_data_validation(dv); dv.add(f'{NOTE}2')
    ws['B2'] = 'Mobility Global Inc. (NYSE: MBGL) — 3-Statement Financial Model'
    ws['B3'] = ('US GAAP · USD millions (whole millions as published; per-share data in USD; shares in millions) · calendar fiscal year · Source: Mobility Global Form 10 '
                '(information statement, 10-12B/A, 27-May-2026; first draft DRS 23-Oct-2025), Form 10-Q (Q2-26), Form 8-K (earnings release Ex. 99.1 and FY2025 quarterly '
                'data Ex. 99.2, 7-Aug-2026; separation 2-Jul-2026), SEC EDGAR CIK 0002090312 ("S&P Global Mobility Holding Co" to Jan-2026); S&P Global Inc. '
                'earnings releases (CIK 0000064040) for the Mobility segment memo; investor-relations sites are blocked — the releases are the documents filed on EDGAR')
    ws['B4'] = ('Basis notes: spun off from S&P Global on 1-Jul-2026 (one MBGL share per SPGI share; 294.8m shares). History to 30-Jun-2026 = carve-out (combined) statements '
                'incl. S&P Global corporate allocations: FY2023–25 annual, 2025 quarters (Q3/Q4 "unaudited and preliminary"), Q1/Q2-26. Pre-2025 quarters and FY2022 = '
                'S&P Global "Mobility" segment memo (revenue comparable within $0–4m; profit not comparable). Segments: CARFAX, B2B (2025 realignment; recast by the company). '
                'Acquisition: Market Scan Feb-2023. $2.0bn senior notes (May-2026) funded a $2.0bn dividend to S&P Global. Q3/Q4-26E calibrated to FY26 revenue and '
                'Adjusted EBITDA guidance (7-Aug-2026).')
    ws[f'{CAGR[0]}5'] = 'Forecast'
    for c in CAGR[1:]: ws[f'{c}5'] = 'Historical'
    ws['B6'] = '(USDm)'
    for col in COLS: ws[f'{col.c}6'] = col.label
    ws[f'{NOTE}6'] = 'Modelling Notes'
    for c, a, b in zip(CAGR, ['5Y CAGR', '2Y CAGR', '1Y', '3Y CAGR'], ["26-'31", "23-'25", "24-'25", "22-'25"]):
        ws[f'{c}6'] = a; ws[f'{c}7'] = b

def set_columns(ws):
    for k in list(ws.column_dimensions.keys()): del ws.column_dimensions[k]
    ws.column_dimensions['A'].width = TM.column_dimensions['A'].width
    ws.column_dimensions['B'].width = TM.column_dimensions['B'].width
    wid = {'A': 13.0, 'Q': 9.5703125, 'QE': 11.42578125, 'E26': 14.140625, 'AE': 13.0}
    for col in COLS:
        cd = ws.column_dimensions[col.c]; cd.width = wid[col.kind]
        if col.q:
            cd.outlineLevel = 1; cd.hidden = True
    for ci in range(CI(NOTE), CI(LASTC) + 1):
        ws.column_dimensions[L(ci)].width = TM.column_dimensions[L(ci + TSHIFT)].width or 13.0
    ws.sheet_format.outlineLevelCol = 1

def main():
    import openpyxl
    wb = openpyxl.load_workbook(TEMPLATE)
    if 'DCF_old' in wb.sheetnames: del wb['DCF_old']
    ws = wb['Model']
    blank_style = copy.copy(ws['A1']._style)
    for rowc in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for c in rowc:
            c.value = None; c.comment = None; c._style = copy.copy(blank_style)
    for k in list(ws.row_dimensions.keys()):
        rd = ws.row_dimensions[k]; rd.outlineLevel = 0; rd.hidden = False
    ws.data_validations.dataValidation = []
    scen_layout(V0)
    build_spec(V0)
    n1 = len(fw.SPEC)
    build_spec2(V0)
    for i, row in enumerate(fw.SPEC): row.spec2 = i >= n1
    last = number_rows(8)
    for r in range(1, last + 1): ws.row_dimensions[r].height = 15.0
    ws.row_dimensions[2].height = 21.0
    defs_setup(); comments_setup()
    write_header(ws)
    for row in fw.SPEC: write_row(ws, row)
    write_scen(ws)
    set_columns(ws)
    ws.freeze_panes = 'C8'
    ws.sheet_view.zoomScale = 70
    import other_sheets as osh
    osh.do_all(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    osh.inject_charts(OUT)
    print('rows', fw.SPEC[-1].r, 'saved', OUT)
    return wb

RECALC = os.environ.get('XLSX_RECALC', '/mnt/skills/public/xlsx/scripts/recalc.py')
def recalc(path, timeout=600):
    import subprocess
    out = subprocess.run(['python3', RECALC, path, str(timeout)], capture_output=True, text=True, cwd=os.path.dirname(RECALC))
    try:
        return json.loads(out.stdout)
    except Exception:
        return {'raw': out.stdout[-2000:], 'err': out.stderr[-2000:]}

def scenario_snapshots():
    import openpyxl
    import other_sheets as osh
    snaps = {}
    for scen in ('Bull', 'Bear'):
        p = os.path.join(BASE, f'MBGL_{scen}.xlsx')
        wb = openpyxl.load_workbook(OUT); wb['Model'][f'{NOTE}2'] = scen; wb.save(p)
        res = recalc(p); print(scen, res.get('status'), res.get('total_errors'))
        v = openpyxl.load_workbook(p, data_only=True)['Bull-Base-Bear']
        snaps[scen] = {(r, c): v[f'{c}{r}'].value for r in range(9, 51) for c in 'CDEFGHIJKL' if v[f'{c}{r}'].value is not None}
        os.remove(p)
    wb = openpyxl.load_workbook(OUT)
    osh.snapshot_values(wb, snaps['Bull'], snaps['Bear'])
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    osh.inject_charts(OUT)

NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
def inject_values(src_xlsx, recalc_xlsx):
    """Store cached formula results (from a LibreOffice-recalculated copy) in the openpyxl file, so the workbook reads correctly in any
    viewer; fullCalcOnLoad stays set so Excel recalculates on open. Same approach as the LECO / ESAB finishing step."""
    import zipfile, shutil, openpyxl
    from lxml import etree
    wbv = openpyxl.load_workbook(recalc_xlsx, data_only=True)
    zin = zipfile.ZipFile(src_xlsx)
    wbxml = etree.fromstring(zin.read('xl/workbook.xml'))
    rels = etree.fromstring(zin.read('xl/_rels/workbook.xml.rels'))
    rid2t = {r.get('Id'): r.get('Target') for r in rels}
    sheets = {}
    for s in wbxml.iter('{%s}sheet' % NS):
        rid = s.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
        sheets['xl/' + rid2t[rid].lstrip('/').replace('xl/', '')] = s.get('name')
    tmp = src_xlsx + '.tmp'
    zout = zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED)
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename in sheets:
            ws = wbv[sheets[item.filename]]
            root = etree.fromstring(data)
            for c in root.iter('{%s}c' % NS):
                f = c.find('{%s}f' % NS)
                if f is None: continue
                v = ws[c.get('r')].value
                old = c.find('{%s}v' % NS)
                if old is not None: c.remove(old)
                if v is None: continue
                ve = etree.SubElement(c, '{%s}v' % NS)
                if isinstance(v, bool):
                    c.set('t', 'b'); ve.text = '1' if v else '0'
                elif isinstance(v, (int, float)):
                    if 't' in c.attrib: del c.attrib['t']
                    ve.text = repr(float(v)) if isinstance(v, float) else str(v)
                elif isinstance(v, str) and v.startswith('#'):
                    c.set('t', 'e'); ve.text = v
                else:
                    c.set('t', 'str'); ve.text = str(v)
            data = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
        zout.writestr(item, data)
    zout.close(); zin.close()
    shutil.move(tmp, src_xlsx)

def store_values():
    import shutil
    p = os.path.join(BASE, 'MBGL_recalc.xlsx')
    shutil.copy(OUT, p)
    res = recalc(p); print('recalc', res.get('status'), res.get('total_errors'))
    inject_values(OUT, p)
    os.remove(p)

if __name__ == '__main__':
    main()
    if '--snap' in sys.argv: scenario_snapshots()
    if '--recalc' in sys.argv: store_values()
