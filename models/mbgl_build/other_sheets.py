"""Rewire the Bull-Base-Bear, DCF and Charts tabs of the MLM template (current layout: DCF with live valuation date, fade period,
implied ROIC, target-IRR entry price, reverse DCF and Mauboussin EV/NOPAT; Bull-Base-Bear with buyback rows) to the MBGL Model rows /
columns; keep the template chart XML. The template's DCF_old tab is not carried.

Column mapping (rolled one year, as for WWD): the template's 2025A base column becomes MBGL 2026E (Q1-Q2 actual, Q3-Q4 calibrated to
guidance; the valuation date of Oct-2026 is within this year, so its cash flow sits in YE2026E net debt) and the template's 2026E-2030E
become 2027E-2031E, so the DCF's explicit years start in 2027E and net debt is taken at 31-Dec-2026."""
import re, zipfile, shutil
from openpyxl.utils import get_column_letter as L
from fw import R, TEMPLATE, ANN, NOTE

# template Model row -> MBGL row key
MAP = {158: 'is_rev', 164: 'r_cash', 166: 'is_op', 182: 'is_dda', 186: 'sh_dil', 211: 'r_adj', 212: 'r_adj_m', 219: 'r_ebit_adj', 221: 'r_adj',
       229: 'e_adjni', 231: 'e_eps', 238: 'gm_cfx', 245: 'gm_ebitda_m', 247: 'gm_op_m', 260: 'cf_wc', 262: 'cf_capex', 263: 'cf_acq',
       280: 'fcf', 282: 'fcf_m', 290: 'bb_px', 291: 'bb_cash', 292: 'bb_sh', 331: 'v_nci', 408: 'ra_ic', 411: 'ra_roic', 422: 'rn_ronta',
       428: 'lv_nd', 433: 'v_px', 434: 'v_sh', 12: 'gm_cfx', 14: 'gm_b2b', 30: 'gm_cfx', 31: 'gm_b2b'}
CMAP = {'BV': ANN[2026].c, 'CA': ANN[2027].c, 'CB': ANN[2028].c, 'CC': ANN[2029].c, 'CD': ANN[2030].c, 'CE': ANN[2031].c, 'CF': NOTE}
REF = re.compile(r"Model!(\$?)([A-Z]{1,3})(\$?)(\d+)")
HIST = [ANN[y].c for y in range(2023, 2026)]

def remap(f):
    def rep(m):
        c, r = m.group(2), int(m.group(4))
        c2 = CMAP.get(c, c)
        r2 = R[MAP[r]] if r in MAP else r
        if c == 'CF' and r == 2: r2 = 2
        return f'Model!{m.group(1)}{c2}{m.group(3)}{r2}'
    return REF.sub(rep, f)

def hist_stat(f):
    """LT historical AVERAGE / MIN / MAX over the template's annual columns -> MBGL 2023-2025 annual columns."""
    m = re.match(r'=IFERROR\((AVERAGE|MIN|MAX)\(Model!C(\d+),', f)
    if not m: return None
    fn, r = m.group(1), int(m.group(2))
    rr = R[MAP[r]]
    return f'=IFERROR({fn}({",".join(f"Model!{c}{rr}" for c in HIST)}),"")'

def fy_dates(f):
    """explicit DCF years roll one year (base year 2026E); year ends stay 31-December"""
    f = re.sub(r'DATE\((\d{4}),12,31\)', lambda m: f'DATE({int(m.group(1)) + 1},12,31)', f)
    f = re.sub(r'DATE\((\d{4})-1,12,31\)', lambda m: f'DATE({int(m.group(1)) + 1}-1,12,31)', f)
    return f

TEXT = {
    'Bull-Base-Bear': {
        'B2': 'Mobility Global Inc. (MBGL) — Bull / Base / Bear Case P&L Summary',
        'B3': ('Company definitions: Adjusted EBITDA (company), Adjusted EPS = model definition (Mobility Global publishes none: + amortization of intangibles, '
               'transaction and other items, tax-effected) · USD millions · Live panel linked to Model — toggle scenario via Model!AC2. Bull and Bear panels are value '
               'snapshots generated from the Model with the switch set to each case. 2026E = 1H carve-out actual + 2H calibrated to the FY26 outlook '
               '(revenue $1,870–1,885m, Adjusted EBITDA $745–760m).'),
        'G26': 13, 'H26': 13, 'S26': 15, 'T26': 15, 'AD26': 10, 'AE26': 10, 'AF26': 10,
    },
    'DCF': {
        'B2': 'DCF VALUATION — MOBILITY GLOBAL INC. (MBGL)',
        'C5': 0.0425, 'D5': 'Assumption (~US 10Y Treasury yield, Oct-2026; rate sources not reachable from the build environment); update to market',
        'D6': 'Assumption: US implied ERP ~4.5–5.5%',
        'C7': 1.15, 'D7': ('Assumption: no trading history before Jul-2026; data / information-services peers ~0.9–1.2, raised for a newly levered spin (~2.2x) '
                           'with auto-dealer and OEM end-market exposure; ~81% subscription revenue dampens cyclicality'),
        'C9': 0.0575, 'D9': 'Senior notes 5.050% 2029 / 5.450% 2031 / 6.050% 2036 (May-2026; 5.74% all-in per the Form 10 pro forma); revolver SOFR + 1.0–1.625%',
        'C10': 0.27, 'D10': 'Assumption: normalised tax rate on adjusted EBIT (model 2027E+ GAAP effective rate 28.5%; amortization is largely non-deductible)',
        'C12': 0.25, 'D12': 'Net debt ~$1.8bn at 30-Jun-2026 vs ~$5.3bn market cap ($18.00 × ~296m diluted shares); D/V ~25%',
        'D15': 'Periods = (31-Dec of the year − today) / 365.25; cash flows assumed at year-end; 2026E (valuation date in Q4) is the base year and its cash flow sits in YE2026E net debt',
        'B18': 'Adj. EBITA (Adjusted EBITDA − depreciation) less SBC and cash transaction / severance / other costs',
        'B21': 'plus depreciation (amortization of acquired intangibles is excluded above, not added back)',
        'B22': '− Capex and acquisitions (M&A lever 2027E–31E)',
        'B26': 'Fade-period growth (2032E–36E; 6.0% → 4.0% linear)', 'I26': 0.06, 'M26': 0.04, 'C27': 0.03,
        'D32': ("NOPAT basis (standard): taxed on pre-amortization Adjusted EBITA with a depreciation-only add-back. The ~$296m p.a. of acquired-intangible "
                'amortization (IHS Markit purchase accounting) is non-cash and largely non-deductible (deferred-tax liability $967m at 30-Jun-2026), so taxing EBIT after '
                'amortization would overstate the cash tax shield. Mobility Global\'s Adjusted EBITDA adds back stock-based compensation, a real cost: row 18 deducts SBC and '
                'the cash transaction (separation stand-up), severance and other items the company excludes. Base year 2026E: valuation date (Oct-2026) inside the year; '
                'YE2026E net debt captures 2H/26 cash flow.'),
        'D43': 'Opening IC = Model 2026E book IC (equity + net debt at 31-Dec-2026; equity carries the IHS Markit purchase goodwill and intangibles), rolled forward with capex + M&A − D&A + ΔWC.',
        'B48': 'Less: Net debt + NCI (Model 2026E year end, 31-Dec-2026)',
        'D48': 'MBGL: no noncontrolling interests; net debt = senior notes (carrying value) + revolver − cash at the base-year end.',
        'B50': 'Diluted shares (m, Model 2026E period end)',
        'D76': 'Rationale (Mobility Global characteristics)',
        'C78': 0.09, 'D78': 'Model Adjusted EBIT CAGR 2026–31E: ~7% revenue growth (CARFAX ~8%, B2B ~5%) with modest margin expansion → high-single-digit NOPAT growth',
        'C79': 0.30, 'D79': 'Asset-light data business (capex ~1.4% of revenue): incremental returns on organic capital are very high; book IC is inflated by IHS Markit purchase accounting',
        'C80': 12, 'D80': 'CARFAX consumer brand and vehicle-history data estate (38bn records, 832m VINs) and dealer network effects; exposure to dealer / OEM budgets and AI-driven disruption',
        'B90': 'VALUATION  —  Applied to 2029E NOPAT',
        'B91': '2029E Adj. EBITA less SBC and cash adjusting items',
        'B93': '2029E NOPAT',
        'B95': 'Implied Enterprise Value (2029E basis)',
        'B96': 'Less: YE2028E net debt + NCI (Model)',
        'B98': 'Diluted shares 2029E (m, Model)',
        'B99': 'Implied share price 2029 ($)',
        'B106': 'Plus: YE2026E net debt + NCI',
        'B108': 'MBGL market-implied EV/NOPAT (on 2029E NOPAT)',
    },
    'Charts': {
        'B2': 'Mobility Global — Operating Performance (2023–2025) and Model Forecast (2026E–2031E)',
        'B3': ('USDm · Linked to Model. Bars = CARFAX / B2B revenue growth (carve-out; 2023 growth n/a on the carve-out basis — see the S&P Global segment memo). Adjusted EBITDA '
               '= company definition. ROIC / RONTA on Adjusted EBIT (full-cost); returns are depressed by the IHS Markit purchase-accounting goodwill and intangibles in equity. '
               'Carve-out history starts in FY2023.'),
        'B8': 'CARFAX revenue y/y %', 'B9': 'B2B revenue y/y %', 'B12': 'Adjusted EBITDA', 'B13': 'Operating profit (GAAP)',
        'B14': 'Adj. EBITDA Margin %', 'B15': 'GAAP Op. Margin %',
    },
}
BBB_ROWS = {9: 'Revenue', 10: '   Revenue y/y %', 11: '   CARFAX revenue y/y %', 12: 'Adjusted EBITDA', 13: '   Adjusted EBITDA margin %',
            15: 'Operating Profit (GAAP)', 16: '   Operating profit y/y %', 17: '   Operating margin (GAAP) %', 19: 'Adjusted Net Income (model definition)',
            35: 'Adjusted EBITDA (leverage basis)', 36: '   Adjusted EBITDA % of revenue'}
CHART_TITLES = {
    'Aggregates Revenue Drivers: Shipment Volume vs. Price (ASP), y/y %': 'Revenue Growth by Segment: CARFAX vs. B2B, y/y %',
    'Adjusted EBITDA &amp; Earnings from Operations (USDm)': 'Adjusted EBITDA &amp; Operating Profit (USDm)',
}
CYEARS = list(range(2023, 2032))

def build_charts(ws):
    for i, y in enumerate(CYEARS):
        c = L(3 + i); mc = ANN[y].c
        ws[f'{c}7'] = y if y <= 2025 else f'{y}E'
        ws[f'{c}8'] = f'=IFERROR(Model!{mc}{R["gm_cfx"]}+0,0)'
        ws[f'{c}9'] = f'=IFERROR(Model!{mc}{R["gm_b2b"]}+0,0)'
        ws[f'{c}11'] = f'={c}7'
        ws[f'{c}12'] = f'=IFERROR(Model!{mc}{R["r_adj"]}+0,NA())'
        ws[f'{c}13'] = f'=IFERROR(Model!{mc}{R["is_op"]}+0,NA())'
        ws[f'{c}14'] = f'=IFERROR(Model!{mc}{R["gm_ebitda_m"]}+0,NA())'
        ws[f'{c}15'] = f'=IFERROR(Model!{mc}{R["gm_op_m"]}+0,NA())'
        ws[f'{c}17'] = f'={c}7'
        ws[f'{c}18'] = f'=IFERROR(Model!{mc}{R["ra_roic"]}+0,NA())'
        ws[f'{c}19'] = f'=IFERROR(Model!{mc}{R["rn_ronta"]}+0,NA())'
        ws[f'{c}20'] = f'=IFERROR(Model!{mc}{R["gm_ebitda_m"]}+0,NA())'
        ws[f'{c}21'] = f'=IFERROR(Model!{mc}{R["fcf_m"]}+0,NA())'
    last = 3 + len(CYEARS)
    for ci in range(last, 23):
        for r in (7, 8, 9, 11, 12, 13, 14, 15, 17, 18, 19, 20, 21):
            ws.cell(r, ci).value = None

def do_all(wb):
    for name in ('Bull-Base-Bear', 'DCF', 'Charts'):
        ws = wb[name]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and 'Model!' in c.value:
                    h = hist_stat(c.value) if name == 'Bull-Base-Bear' else None
                    c.value = h or remap(c.value)
                if isinstance(c.value, str) and 'DATE(' in c.value:
                    c.value = fy_dates(c.value)
        for addr, t in TEXT.get(name, {}).items():
            ws[addr] = t
    b = wb['Bull-Base-Bear']
    for r, t in BBB_ROWS.items():
        for c in ('B', 'N', 'Z'): b[f'{c}{r}'] = t
    for c0 in ('C', 'O', 'AA'):
        i = b[f'{c0}7'].column
        b.cell(7, i).value = '2026E'
        for k in range(1, 6): b.cell(7, i + k).value = f'{2026 + k}E'
        b.cell(7, i + 6).value = "26-'31 CAGR"
    from openpyxl.comments import Comment
    cm = Comment('P/E on adjusted EPS (model definition: + amortization of intangibles and adjusting items, tax-effected). MBGL at $18.00 trades at ~13x 2027E adjusted EPS. '
                 'Exit multiples are assumptions: Base 13x, Bull 15x, Bear 10x (no trading history before Jul-2026; leveraged spin with mid-to-high-single-digit growth).', 'MBGL model')
    cm.width, cm.height = 420, 120; b['N26'].comment = cm
    cm2 = Comment(f'Snapshot (values) of Model rows {R["bb_cash"]}, {R["bb_px"]} and {R["bb_sh"]} (Share Buyback Schedule) with Model!{NOTE}2 = Bull / Bear. '
                  'Base panel (D40:H42) is live-linked.', 'MBGL model')
    cm2.width, cm2.height = 420, 90; b['P40'].comment = cm2
    b['J2'] = 'Share price ($):'
    for c in ('B', 'N', 'Z'): b[f'{c}29'] = 'IRR (to 31-Dec of Y-1)'
    for c in ('B', 'N', 'Z'):
        if b[f'{c}40'].value: b[f'{c}40'] = '   Share repurchases ($m; 2026E = 2H)'
    d = wb['DCF']
    for i, col in enumerate('DEFGHIJKLM'):
        d[f'{col}17'] = f'{2027 + i}E'
    a26, a27 = ANN[2026].c, ANN[2027].c
    d['D22'] = f'=Model!{a27}{R["cf_capex"]}+Model!{a27}{R["cf_acq"]}'
    d['D34'] = f'=Model!{a26}{R["ra_ic"]}'
    d['C48'] = f'=Model!{a26}{R["lv_nd"]}+Model!{a26}{R["v_nci"]}'
    d['C50'] = f'=Model!{a26}{R["v_sh"]}'
    d['C52'] = f'=Model!{a26}{R["v_px"]}'
    d['P17'] = 'CAGR 27–31E'; d['Q17'] = 'CAGR 27–36E'
    d['C91'] = f'=Model!{ANN[2029].c}{R["r_ebita"]}-Model!{ANN[2029].c}{R["r_cash"]}-Model!{ANN[2029].c}{R["d_sbc"]}'
    for col, y in zip('DEFGH', range(2027, 2032)):     # standard NOPAT basis (playbook 5.7): pre-amortization, after SBC and cash items
        d[f'{col}18'] = f'=Model!{ANN[y].c}{R["r_ebita"]}-Model!{ANN[y].c}{R["r_cash"]}-Model!{ANN[y].c}{R["d_sbc"]}'
        d[f'{col}21'] = f'=Model!{ANN[y].c}{R["d_dep"]}'
    build_charts(wb['Charts'])

def snapshot_values(wb, bull, bear):
    ws = wb['Bull-Base-Bear']
    cols = 'CDEFGHIJKL'
    for snap, off in ((bull, 'OPQRSTUVWX'), (bear, ['AA', 'AB', 'AC', 'AD', 'AE', 'AF', 'AG', 'AH', 'AI', 'AJ'])):
        for (r, c), v in snap.items():
            tgt = f'{off[cols.index(c)]}{r}'
            cur = ws[tgt].value
            if isinstance(cur, str) and cur.startswith('='): continue
            if r in (7, 26): continue
            ws[tgt] = v

def inject_charts(path):
    """Replace openpyxl-rewritten chart parts with the template chart XML (combo charts, colours, layout), retitled, re-ranged
    to Charts!C:K (2023-2031E) and cached values removed so Excel redraws from the MBGL cells."""
    tz = zipfile.ZipFile(TEMPLATE)
    tmpl = {n: tz.read(n) for n in tz.namelist() if n.startswith('xl/charts/chart')}
    tmp = path + '.tmp'
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in tmpl:
                t = tmpl[item.filename].decode('utf-8')
                for a, b2 in CHART_TITLES.items(): t = t.replace(a, b2)
                t = re.sub(r'(Charts!\$C\$\d+:\$)V(\$\d+)', r'\1K\2', t)
                t = re.sub(r'<c:numCache>.*?</c:numCache>', '', t, flags=re.S)
                t = re.sub(r'<c:strCache>.*?</c:strCache>', '', t, flags=re.S)
                data = t.encode('utf-8')
            zout.writestr(item, data)
    shutil.move(tmp, path)
