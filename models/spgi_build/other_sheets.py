"""Rewire the Bull-Base-Bear, DCF and Charts tabs of the MLM template (current layout: DCF with live valuation date, fade period,
implied ROIC, target-IRR entry price, reverse DCF and Mauboussin EV/NOPAT; Bull-Base-Bear with buyback rows) to the SPGI Model rows /
columns; keep the template chart XML. The template's DCF_old tab is not carried.

Column mapping: the template's 2025A base column becomes SPGI FY2026E (1H actual + 2H calibrated to guidance; the valuation date,
Oct-2026, falls in Q4) and the template's 2026E-2030E become SPGI 2027E-2031E. Net debt is taken at 31-Dec-2026E, so the 2026E UFCF
is excluded from the DCF (no double counting) and the explicit years start in 2027E."""
import re, zipfile, shutil, copy
from openpyxl.utils import get_column_letter as L
from fw import R, TEMPLATE, ANN, NOTE

# template Model row -> SPGI row key
MAP = {158: 'is_rev', 164: 'r_cash', 166: 'is_op', 182: 'd_dep', 186: 'sh_dil', 211: 'r_adj', 212: 'r_adj_m', 219: 'r_ebit', 221: 'r_adj',
       229: 'e_adjni', 231: 'e_eps', 238: 'gm_rat', 245: 'gm_ebitda_m', 247: 'gm_op_m', 260: 'cf_wc', 262: 'cf_capex', 263: 'cf_acq',
       280: 'fcf', 282: 'fcf_m', 290: 'bb_px', 291: 'bb_cash', 292: 'bb_sh', 331: 'v_nci', 408: 'ra_ic', 411: 'ra_roic', 422: 'roe',
       428: 'lv_nd', 433: 'v_px', 434: 'v_sh'}
CMAP = {'BV': ANN[2026].c, 'CA': ANN[2027].c, 'CB': ANN[2028].c, 'CC': ANN[2029].c, 'CD': ANN[2030].c, 'CE': ANN[2031].c, 'CF': NOTE}
REF = re.compile(r"Model!(\$?)([A-Z]{1,3})(\$?)(\d+)")
HIST = [ANN[y].c for y in range(2010, 2026)]

def remap(f):
    def rep(m):
        c, r = m.group(2), int(m.group(4))
        c2 = CMAP.get(c, c)
        r2 = R[MAP[r]] if r in MAP else r
        if c == 'CF' and r == 2: r2 = 2
        return f'Model!{m.group(1)}{c2}{m.group(3)}{r2}'
    return REF.sub(rep, f)

def hist_stat(f):
    """LT historical AVERAGE / MIN / MAX over the template's annual columns -> SPGI FY2010-FY2025 annual columns."""
    m = re.match(r'=IFERROR\((AVERAGE|MIN|MAX)\(Model!C(\d+),', f)
    if not m: return None
    fn, r = m.group(1), int(m.group(2))
    rr = R[MAP[r]]
    return f'=IFERROR({fn}({",".join(f"Model!{c}{rr}" for c in HIST)}),"")'

def fy_dates(f):
    f = re.sub(r'DATE\((\d{4}),12,31\)', lambda m: f'DATE({int(m.group(1)) + 1},12,31)', f)
    f = re.sub(r'DATE\((\d{4})-1,12,31\)', lambda m: f'DATE({int(m.group(1)) + 1}-1,12,31)', f)
    return f

TEXT = {
    'Bull-Base-Bear': {
        'B2': 'S&P Global Inc. (SPGI) — Bull / Base / Bear Case P&L Summary',
        'B3': ('Adjusted figures (company definitions: adjusted operating profit, adjusted diluted EPS; adjusted EBITDA = model definition) · USD millions · continuing operations excluding '
               'Mobility · Live panel linked to Model — toggle scenario via Model!CC2. Bull and Bear panels are value snapshots generated from the Model with the switch set to each case. '
               'FY2026E = 1H actual (pro forma) + 2H calibrated to the FY26 guidance (revenue +5.9–7.9%, adjusted EPS $17.50–17.75).'),
        'F26': 26, 'G26': 26, 'H26': 26, 'R26': 30, 'S26': 30, 'T26': 30, 'AD26': 20, 'AE26': 20, 'AF26': 20,
    },
    'DCF': {
        'B2': 'DCF VALUATION — S&P GLOBAL INC. (SPGI)',
        'C5': 0.0425, 'D5': 'Assumption (~US 10Y Treasury yield, Oct-2026; government rate sources not reachable from the build environment); update to market',
        'D6': 'Assumption: US implied ERP ~4.5–5.5%',
        'C7': 1.00, 'D7': ('Assumption: information-services / benchmark peers (MCO, MSCI, ICE, CME, VRSK, FDS) ~0.9–1.1; Ratings issuance and Indices asset-linked fees are market-sensitive, '
                           'subscriptions defensive'),
        'C9': 0.05, 'D9': 'Recent issues: $600m 4.25% 2031 and $400m 4.80% 2035 notes (Dec-2025); A-rated issuer; pre-tax cost of new debt ~5% (assumption)',
        'C10': 0.225, 'D10': 'Model 2027E+ effective tax rate (FY26 guidance 22–23%)',
        'C12': 0.10, 'D12': 'Net debt ~$12bn at YE2026E (after >$7bn of 2026 buybacks) vs ~$118bn market cap ($415 × ~283m diluted shares); ~1.6x adjusted EBITDA; plus ~$5bn redeemable NCI',
        'D15': 'Periods = (31-Dec of the year − today) / 365.25; cash flows assumed at year end; FY2026E (Q4 under way at the valuation date) is the base year — its cash flow is in YE2026E net debt',
        'B18': 'Adj. EBITA (before deal amortization) less severance, acquisition, disposition & legal cash costs',
        'B21': 'plus depreciation (deal-related amortization is not deducted above, so it is not added back)',
        'B22': '− Capex and acquisitions (bolt-ons; M&A lever 2027E–31E)',
        'B26': 'Fade-period growth (2032E–36E; 7.0% → 5.0% linear)', 'I26': 0.07, 'M26': 0.05, 'C27': 0.04,
        'D32': ("S&P Global's adjusted operating profit excludes deal-related amortization (IHS Markit and other acquired intangibles) but is after stock-based compensation, which the company does not "
                'adjust out. NOPAT therefore taxes adjusted EBITA (before acquired-intangible amortization) less the cash severance, acquisition, disposition and legal costs the company adds back '
                '(row 18), and adds back depreciation only (row 21). Row 22 includes bolt-on acquisitions in 2027E–31E; the fade period and terminal value assume no new acquisitions. '
                'Redeemable NCI (CME Group 27% of S&P DJI) and other NCI are deducted in the EV bridge. Mobility (spun off 1-Jul-2026) is excluded throughout.'),
        'D43': 'Opening IC = Model FY2026E book IC (equity incl. redeemable NCI + net debt at 31-Dec-2026E, after the Mobility spin-off), rolled forward with capex + M&A − depreciation + ΔWC.',
        'B48': 'Less: Net debt + redeemable & other NCI (Model FY2026E year end)',
        'D48': 'Net debt at 31-Dec-2026E (after the >$7bn 2026 buybacks) plus redeemable NCI (CME 27% of S&P DJI, carrying value) and other NCI; FY2026E cash flow is not double counted.',
        'B50': 'Diluted shares (m, Model FY2026E period end)',
        'D76': 'Rationale (S&P Global characteristics)',
        'C78': 0.09, 'D78': 'Model adjusted EBITA CAGR 2026–31E: mid / high-single-digit revenue growth with modest margin expansion and the M&A lever → ~8–9% NOPAT growth',
        'C79': 0.30, 'D79': ('Benchmarks (S&P 500 index licensing, Platts price assessments) and ratings are franchise assets with very high incremental returns; reported ROIC is depressed by the '
                             '$31bn IHS Markit goodwill, so incremental (not average) returns matter'),
        'C80': 20, 'D80': ('Regulatory-licensed ratings duopoly, embedded benchmarks and data networks give a long competitive-advantage period; risks are issuance cyclicality, passive-fee '
                           'compression and AI-driven disruption of data / workflow products'),
        'B90': 'VALUATION  —  Applied to FY2029E NOPAT',
        'B91': 'FY2029E Adj. EBITA less severance / acquisition / disposition / legal costs',
        'B93': 'FY2029E NOPAT',
        'B95': 'Implied Enterprise Value (FY2029E basis)',
        'B96': 'Less: FY2028E year-end net debt + NCI (Model)',
        'B98': 'Diluted shares FY2029E (m, Model)',
        'B99': 'Implied share price FY2029 ($)',
        'B106': 'Plus: FY2026E year-end net debt + NCI',
        'B108': 'SPGI market-implied EV/NOPAT (on FY2029E NOPAT)',
    },
    'Charts': {
        'B2': 'S&P Global Inc. — Operating Performance (FY2012–FY2025) and Model Forecast (FY2026E–FY2031E)',
        'B3': ('USDm · fiscal years ended 31 December · Linked to Model. Bars = contribution of Ratings and of the other divisions (Indices, Energy, Market Intelligence) to total revenue growth (pp); '
               'FY2022 includes the IHS Markit merger; FY2025+ excludes Mobility. Adjusted EBITDA = model definition (adjusted OP + depreciation); ROIC on adjusted EBITA.'),
        'B8': 'Ratings contribution to growth (pp)', 'B9': 'Indices, Energy & MI contribution (pp)', 'B12': 'Adjusted EBITDA', 'B13': 'Operating profit (GAAP)',
        'B14': 'Adj. EBITDA Margin %', 'B15': 'Operating Margin %', 'B19': 'ROE (RONTA n/m: negative net tangible assets)',
    },
}
BBB_ROWS = {9: 'Revenue', 10: '   Revenue y/y %', 11: '   Ratings revenue y/y %', 12: 'Adjusted EBITDA (model)',
            15: 'Operating profit (GAAP)', 16: '   Operating profit y/y %', 17: '   Operating margin (GAAP) %', 19: 'Adjusted Net Income',
            35: 'Adjusted EBITDA (leverage basis)', 36: '   Adjusted EBITDA % of revenue', 45: 'ROE'}
CHART_TITLES = {
    'Aggregates Revenue Drivers: Shipment Volume vs. Price (ASP), y/y %': 'Revenue Growth: Ratings vs. Other Divisions Contribution, pp',
    'Adjusted EBITDA &amp; Earnings from Operations (USDm)': 'Adjusted EBITDA &amp; Operating Profit (USDm)',
}

def build_charts(ws):
    for i, y in enumerate(range(2012, 2032)):
        c = L(3 + i); mc = ANN[y].c
        ws[f'{c}7'] = y if y <= 2025 else f'{y}E'
        ws[f'{c}8'] = f'=IFERROR(Model!{mc}{R["gc_rat"]}+0,0)'
        ws[f'{c}9'] = f'=IFERROR(Model!{mc}{R["gc_idx"]}+0,0)+IFERROR(Model!{mc}{R["gc_en"]}+0,0)+IFERROR(Model!{mc}{R["gc_mi"]}+0,0)'
        ws[f'{c}11'] = f'={c}7'
        ws[f'{c}12'] = f'=Model!{mc}{R["r_adj"]}'
        ws[f'{c}13'] = f'=Model!{mc}{R["is_op"]}'
        ws[f'{c}14'] = f'=Model!{mc}{R["gm_ebitda_m"]}'
        ws[f'{c}15'] = f'=Model!{mc}{R["gm_op_m"]}'
        ws[f'{c}17'] = f'={c}7'
        ws[f'{c}18'] = f'=IFERROR(Model!{mc}{R["ra_roic"]}+0,NA())'
        ws[f'{c}19'] = f'=IFERROR(Model!{mc}{R["roe"]}+0,NA())'
        ws[f'{c}20'] = f'=IFERROR(Model!{mc}{R["gm_ebitda_m"]}+0,NA())'
        ws[f'{c}21'] = f'=IFERROR(Model!{mc}{R["fcf_m"]}+0,NA())'

def do_all(wb):
    for name in ('Bull-Base-Bear', 'DCF', 'Charts'):
        ws = wb[name]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and 'Model!' in c.value:
                    h = hist_stat(c.value) if name == 'Bull-Base-Bear' else None
                    c.value = h or remap(c.value)
                if isinstance(c.value, str) and 'DATE(' in c.value:
                    c.value = fy_dates(c.value)
        for addr, t in TEXT.get(name, {}).items():
            ws[addr] = t
    b = wb['Bull-Base-Bear']
    for r, t in BBB_ROWS.items():
        for c in ('B', 'N', 'Z'): b[f'{c}{r}'] = t
    for c0 in ('C', 'O', 'AA'):
        i = b[f'{c0}7'].column
        b.cell(7, i).value = '2026E'
        for k in range(1, 6): b.cell(7, i + k).value = f'{2026 + k}E'
        b.cell(7, i + 6).value = "26-'31 CAGR"
    from openpyxl.comments import Comment
    cm = Comment('P/E on adjusted diluted EPS (company definition, continuing operations). SPGI at $415 trades at ~23.5x FY2026E adjusted EPS (~$17.6 guidance mid-point) and '
                 '~21x 2027E. Exit multiples are assumptions applied to 2029E–31E: Base 26x, Bull 30x (re-rating toward the ~28–33x the stock carried in 2019–24), '
                 'Bear 20x (AI-disruption / issuance-downturn de-rating).', 'SPGI model')
    cm.width, cm.height = 420, 120; b['N26'].comment = cm
    cm2 = Comment(f'Snapshot (values) of Model rows {R["bb_cash"]}, {R["bb_px"]} and {R["bb_sh"]} (Share Buyback Schedule) with Model!{NOTE}2 = Bull / Bear. '
                  'Base panel (D40:H42) is live-linked.', 'SPGI model')
    cm2.width, cm2.height = 420, 90; b['P40'].comment = cm2
    b['J2'] = 'Share price ($):'
    for c in ('F', 'R'):            # price target / IRR for the first exit year too (template filled two years for Base / Bull)
        b[f'{c}27'] = f'={c}26*{c}23'; b[f'{c}28'] = f'=IFERROR({c}27/$K$2-1,"")'
        b[f'{c}29'] = f'=IFERROR(({c}27/$K$2)^(1/((DATE(2029-1,12,31)-$K$3)/365.25))-1,"")'
        for r in (27, 28, 29): b[f'{c}{r}']._style = copy.copy(b[f'G{r}']._style)
    for c in ('B', 'N', 'Z'): b[f'{c}29'] = 'IRR (to 31-Dec of Y-1)'
    for c in ('B', 'N', 'Z'):
        if b[f'{c}40'].value: b[f'{c}40'] = '   Share repurchases ($m; FY2026E = 2H)'
    d = wb['DCF']
    for i, col in enumerate('DEFGHIJKLM'):
        d[f'{col}17'] = f'{2027 + i}E'
    a26, a27 = ANN[2026].c, ANN[2027].c
    d['D22'] = f'=Model!{a27}{R["cf_capex"]}+Model!{a27}{R["cf_acq"]}'
    d['D34'] = f'=Model!{a26}{R["ra_ic"]}'
    d['C48'] = f'=Model!{a26}{R["lv_nd"]}+Model!{a26}{R["v_nci"]}'
    cm3 = Comment('S&P Global has negative operating working capital (unearned revenue billed in advance), so revenue growth releases cash: net reinvestment in the fade period is '
                  'negative and the implied terminal RONIC (g / reinvestment) is not meaningful. Value creation rests on the high returns of the existing franchise.', 'SPGI model')
    cm3.width, cm3.height = 420, 100; d['B42'].comment = cm3
    d['C50'] = f'=Model!{a26}{R["v_sh"]}'
    d['C52'] = f'=Model!{a26}{R["v_px"]}'
    d['P17'] = 'CAGR 27–31E'; d['Q17'] = 'CAGR 27–36E'
    build_charts(wb['Charts'])

def snapshot_values(wb, bull, bear):
    ws = wb['Bull-Base-Bear']
    cols = 'CDEFGHIJKL'
    for snap, off in ((bull, 'OPQRSTUVWX'), (bear, ['AA', 'AB', 'AC', 'AD', 'AE', 'AF', 'AG', 'AH', 'AI', 'AJ'])):
        for (r, c), v in snap.items():
            tgt = f'{off[cols.index(c)]}{r}'
            cur = ws[tgt].value
            if isinstance(cur, str) and cur.startswith('='): continue
            if r in (7, 26): continue
            ws[tgt] = v

def inject_charts(path):
    """Replace openpyxl-rewritten chart parts with the template chart XML (combo charts, colours, layout), retitled, ranges kept
    (Charts!C:V = FY2012-FY2031E) and cached values removed so Excel redraws from the SPGI cells."""
    tz = zipfile.ZipFile(TEMPLATE)
    tmpl = {n: tz.read(n) for n in tz.namelist() if n.startswith('xl/charts/chart')}
    tmp = path + '.tmp'
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in tmpl:
                t = tmpl[item.filename].decode('utf-8')
                for a, b2 in CHART_TITLES.items(): t = t.replace(a, b2)
                t = re.sub(r'<c:numCache>.*?</c:numCache>', '', t, flags=re.S)
                t = re.sub(r'<c:strCache>.*?</c:strCache>', '', t, flags=re.S)
                data = t.encode('utf-8')
            zout.writestr(item, data)
    shutil.move(tmp, path)
