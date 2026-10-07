"""Build STE_Model.xlsx from the MLM template (same layout, formatting and process as the prior models)."""
import sys, copy, os, json
from openpyxl.styles import Font
from fw import *
import fw
from spec import build_spec, scen_layout, SCEN_ROWS, SC
from spec2 import build_spec2

OUT = os.path.join(os.path.dirname(BASE), 'STE_Model.xlsx')
G27 = MAN['guidance_fy27']

# ====================================================================== inputs (V0)
def wc_inputs():
    q = DATA['2026']['x']
    rev, cogs = q['rev'], q['cogs']
    return {'dso': round(q['bs_ar'] / rev * 365), 'dio': round(q['bs_inv'] / cogs * 365), 'dpo': round(q['bs_ap'] / cogs * 365),
            'ocl': round(q['bs_ocl'] / rev, 3)}
WC = wc_inputs()
WC['note'] = (f'Forecast inputs at fiscal-year-end (31-Mar-2026) levels on FY2026 revenue / cost of revenues (DSO {WC["dso"]}d, inventory {WC["dio"]}d, '
              f'payables {WC["dpo"]}d, other current liabilities {WC["ocl"]:.1%} of revenue); year ends are compared with year ends (Q1 is seasonally low in receivables).')
AM = MAN['amort_sched_10k']

V0 = {
    'out_hdr': 'FY27 outlook — Q4 FY26 release (11-May-26), reiterated 5-Aug-26',
    'out_note': ('As-reported revenue growth 7–8% (incl. Healthcare tuck-ins, slightly favourable FX); constant-currency organic 6–7%; adjusted EPS $11.10–11.30 '
                 '(+9–11% vs $10.17); ~25% tax rate; capex ~$450m (raised from $375m in Aug-26); free cash flow ~$800m (from $850m).'),
    'outlook': [
        ('out_g', '  FY27 revenue growth, as reported', tuple(G27['rev_growth']), 'pct', 'Guided 7–8%. Drives the Q2–Q4/27E revenue calibration.'),
        ('out_cc', '  FY27 constant-currency organic growth — memo', tuple(G27['cc_organic']), 'pct', 'Guided 6–7% (memo; compare the growth-bridge row).'),
        ('out_eps', '  FY27 adjusted EPS ($)', tuple(G27['adj_eps']), 'ps', 'Guided $11.10–11.30. Drives the Q2–Q4/27E segment-margin calibration.'),
        ('out_geps', '  FY27 GAAP diluted EPS, continuing ($) — memo', tuple(G27['gaap_eps']), 'ps',
         'Guided $8.92–9.12 (amortization $2.08, acquisition $0.01, restructuring $0.07, other $0.02 per share); memo, not forced.'),
        ('bb27', '  Q2–Q4/27E share repurchases ($m)', (400.0, 300.0, 200.0), 'm', 'Q1/27 $115.5m; $900m left on the $1.0bn May-2026 authorization.'),
    ],
    'points': [
        ('pt_int', '  Interest expense ($m)', 63.0, 'm1', 'Not guided; FY26 $60.7m, Q1/27 $15.8m; Feb-2027 private-placement maturity refinanced on the revolver.'),
        ('pt_onop', '  Other non-operating (income) expense, net ($m)', -10.0, 'm1', 'Interest & miscellaneous income (Q1/27 −$3.1m) net of other items.'),
        ('pt_amort', '  Amortization of acquired intangibles ($m)', 262.0, 'm1', 'Q1/27 $65.3m; 10-K schedule $253.1m + FY27 tuck-ins; FY27 outlook $2.08 per share after tax.'),
        ('pt_dep', '  Depreciation & other amortization ($m)', 238.0, 'm1', 'Q1/27 $58.5m (D&A $123.8m − amortization $65.3m); rising with capex.'),
        ('pt_etr', '  Tax rate (GAAP & adjusted)', G27['etr'], 'pct', 'FY27 outlook ~25% (Q1/27 GAAP 26.5%, adjusted 25.9%).'),
        ('pt_capex', '  Capital expenditures ($m)', G27['capex'], 'm1', 'FY27 outlook ~$450m (Q1/27 release), incl. the North Carolina chemistries and Mentor, OH plants.'),
        ('pt_fcf', '  Free cash flow outlook ($m) — memo', G27['fcf'], 'm1', 'FY27 outlook ~$800m (CFO ~$1,250m − capex ~$450m).'),
    ],
    'lev': (1.0, 0.75, 0.5),
    'lev_note': ('Net debt $1.41bn at 30-Jun-26 ≈ 0.9x adjusted EBITDA (model definition). Target = leverage up to which surplus cash is returned via buybacks; '
                 'below it free cash flow is held / repays the revolver. Bull = more buybacks (higher target).'),
    'blocks': [
        ('HC_CAP', 'HC_CAP (Healthcare capital equipment y/y)', 0.02, -0.04, [0.045, 0.04, 0.04, 0.04],
         'Hospital capex cycle (sterilizers, washers, OR tables / lights, integration); backlog $444m at 30-Jun-26 (+10% y/y).', 'pct', False),
        ('HC_CONS', 'HC_CONS (Healthcare consumables y/y)', 0.015, -0.03, [0.065, 0.065, 0.06, 0.06],
         'Procedure volumes (US surgical procedures +MSD), endoscopy consumables (Cantel), share gains, price.', 'pct', False),
        ('HC_SERV', 'HC_SERV (Healthcare service y/y)', 0.015, -0.03, [0.085, 0.08, 0.075, 0.07],
         'Equipment service contracts, instrument repair (BD instruments, Aug-23), outsourced reprocessing; FY26 +12%.', 'pct', False),
        ('AST_SERV', 'AST_SERV (AST contract sterilization & lab y/y)', 0.02, -0.04, [0.08, 0.08, 0.075, 0.075],
         'Med-device volumes (bioprocessing recovery), price; EO / gamma / X-ray capacity adds. FY26 +11%.', 'pct', False),
        ('AST_CAP', 'AST_CAP (AST capital equipment y/y)', 0.0, -0.10, [0.0, 0.0, 0.0, 0.0], 'Lumpy X-ray / E-beam system sales (~$20–30m p.a.).', 'pct', False),
        ('LS_CAP', 'LS_CAP (Life Sciences capital equipment y/y)', 0.03, -0.05, [0.04, 0.04, 0.04, 0.04], 'Pharma / biopharma capital (backlog $110m).', 'pct', False),
        ('LS_CONS', 'LS_CONS (Life Sciences consumables y/y)', 0.015, -0.03, [0.06, 0.06, 0.06, 0.06], 'Barrier & cleaning chemistries for pharma manufacturing.', 'pct', False),
        ('LS_SERV', 'LS_SERV (Life Sciences service y/y)', 0.01, -0.03, [0.04, 0.04, 0.04, 0.04], 'Installed-base service (CECS divested Q4 FY24).', 'pct', False),
        ('HC_M', 'HC_M (Healthcare segment operating margin)', 0.01, -0.02, [0.251, 0.254, 0.257, 0.26],
         'FY26 24.6%; FY27E ~24.8% (calibrated to the EPS outlook). Volume / price / productivity vs tariffs and inflation (~+30 bps p.a.).', 'pct', False),
        ('AST_M', 'AST_M (AST segment operating margin)', 0.01, -0.02, [0.453, 0.456, 0.459, 0.462], 'FY26 46.1%; high fixed-cost network — operating leverage on volume.', 'pct', False),
        ('LS_M', 'LS_M (Life Sciences segment operating margin)', 0.01, -0.02, [0.413, 0.416, 0.419, 0.422], 'FY26 42.6%; FY27E calibrated ~40.8%.', 'pct', False),
        ('MNA', 'M&A_spend (tuck-in acquisitions, $m)', 150.0, -150.0, [150.0, 150.0, 150.0, 150.0],
         'STERIS spent $16–55m p.a. on tuck-ins (FY25–Q1/27) besides BD instruments ($540m, FY24) and Cantel ($4.6bn incl. debt, FY22). Base $150m p.a.', 'm', True),
    ],
    'qe_acq': {2: 7.0, 3: 7.0, 4: 7.0},
    'qe_acq_note': 'Q2–Q4/27E input: Healthcare tuck-ins not in the prior-year quarter (Q1/27 impact of acquisitions $6.9m).',
    'qe_fx': {2: 7.0, 3: 7.0, 4: 6.0},
    'qe_acq_cash': 0.0,
    'corp_pct': -0.071,
    'dep_pct': {2028: 0.038, 2029: 0.039, 2030: 0.040, 2031: 0.040}, 'sga_pct': {2028: 0.232, 2029: 0.229, 2030: 0.226, 2031: 0.222},
    'rd_pct': 0.019, 'sbc_pct': 0.0105, 'etr': 0.25, 'capex_pct': {2028: 0.068, 2029: 0.065, 2030: 0.063, 2031: 0.062},
    'dps_q': 0.63, 'def_q': -4.0, 'def_ae': -15.0, 'opt_q': 3.0, 'opt_ae': 15.0,
    'qe_restr': {2: 3.0, 3: 3.0, 4: 3.0}, 'ae_restr': {2028: 20.0, 2029: 20.0, 2030: 12.0, 2031: 0.0},
    'restr_note': MAN['restructuring_2027'] + ' Model: FY27E ~$9m (outlook $0.07 per share), FY28E $20m, FY29E $20m, FY30E $12m (total ~$61m).',
    'amort_sched': {int(y) + 0: AM[y] + 8.0 for y in ('2028', '2029', '2030', '2031')},
    'amort_note': MAN['amort_sched_src'] + ' FY28 $248.2m, FY29 $246.3m, FY30 $241.3m, FY31 $209.1m; + ~$8m p.a. from FY27 tuck-ins (Q1/27 run-rate $261m vs $253.1m scheduled).',
    'repay': {2027: 118.9, 2028: 150.0, 2029: 127.5, 2030: 0.0, 2031: 100.0},
    'mat': {2027: 150.0, 2028: 127.5, 2029: 0.0, 2030: 100.0, 2031: 61.5},
    'min_cash': 400.0,
    'r_notes': {2028: 0.033, 2029: 0.033, 2030: 0.033, 2031: 0.034},
    'r_notes_note': 'Weighted coupon ~3.2% (2.70% / 3.75% public notes; 1.86–4.03% private placements); FY31: 2031 notes refinanced at ~5% from Mar-2031.',
    'wc': WC,
    'px': 211.07,
    'px_note': 'VALUATION — share price input $211.07 (NYSE: STE, close 5-Oct-2026; ~$20.6bn market cap); update to market',
    'mna_note': 'Scenario lever M&A_spend (Base $150m p.a. of tuck-ins). Funded from cash / revolver (debt schedule).',
    'rev_guid_note': 'Company FY27 as-reported revenue growth outlook 7–8% (11-May-2026, reiterated 5-Aug-2026). Drives the revenue calibration.',
    'eps_guid_note': 'Company FY27 adjusted EPS outlook $11.10–11.30 (reiterated 5-Aug-2026). Drives the segment-margin calibration.',
}
# ====================================================================== compute engine
def hist_value(row, col):
    kw = row.kw; path = kw.get('data')
    if not path: return None
    if kw.get('ppa'):
        v = MAN['ppa'].get(pkey(col), {}).get(path.split('_', 1)[1])
        return v if kw.get('textdata') else num(v)
    v = xget(pkey(col), path)
    return v if kw.get('textdata') else num(v)

def first(kw, keys):
    for k in keys:
        if k in kw: return k
    return None

def compute(row, col):
    kw = row.kw; k = col.kind
    CTX.col = col
    if kw.get('ann') and col.q: return None, False
    if k in HISTK:
        v = hist_value(row, col)
        if v is not None: return v, True
        for key in ('hist', 'all'):
            if key in kw:
                return kw[key](), False
        return None, False
    if k == 'QE':
        if row.spec2: return None, False
        if 'qe_in' in kw:
            v = kw['qe_in']; return (v[col.q] if isinstance(v, dict) else v), True
        key = first(kw, ('qe', 'fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    if k == 'E26':
        if 'e26_in' in kw: return kw['e26_in'], True
        if 'e26' in kw: return kw['e26'](), False
        if kw.get('avg') and ('qe' in kw or 'qe_in' in kw): return f'=AVERAGE({Q1}{row.r},{Q2}{row.r},{Q3}{row.r},{Q4}{row.r})', False
        if kw.get('flow') and not row.spec2 and any(x in kw for x in ('qe', 'qe_in', 'fc')): return S27(row.key), False
        key = first(kw, ('fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    if k == 'AE':
        if 'ae_in' in kw:
            v = kw['ae_in']; return (v[col.year] if isinstance(v, dict) else v), True
        key = first(kw, ('ae', 'fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    return None, False

# ====================================================================== writers
CAGR_G = {'is_rev', 'g_rev', 'hc_rev', 'ast_rev', 'ls_rev', 'mx_cap', 'mx_cons', 'mx_serv', 'o_adj', 'g_adj', 'r_adj', 'r_ebitda', 'is_op', 'is_dda',
          'e_adjni', 'e_eps', 'cf_cfo', 'fcf', 'bs_ta', 'bs_teq', 'sh_dil', 'is_gp', 'is_sga', 'd_amort', 'd_dep', 'hc_oi', 'ast_oi', 'ls_oi',
          'is_nica', 'eps_dilc', 'dps', 'hc_cap', 'hc_cons', 'hc_serv', 'ast_serv', 'ls_cap', 'ls_cons', 'ls_serv'}
CAGR_A = {'hc_m', 'ast_m', 'ls_m', 'g_adjm', 'o_adj_m', 'r_adj_m', 'is_gm', 'is_opm', 'r_ebit_m', 'gm_rev', 'gm_org', 'gm_gm', 'gm_agm', 'gm_ebitda_m',
          'gm_op_m', 'gm_net_m', 'gm_sga', 'gm_rd', 'gm_adjop_m', 'fcf_m', 'fcf_conv', 'capex_pct', 'ra_roic', 'rn_ronta', 'roe', 'lv_x', 'og_cc', 'og_og',
          'd_sga', 'd_dep_pct', 'd_sbc_pct', 'wc_nwc_p', 'mx_rec', 'is_etr', 'e_rate', 'fcf_eb', 'corp_pct'}
A = {y: ANN[y].c for y in ANN}

def style_cagr(ws, r, on):
    src = 16 if on else 12
    for c in CAGR:
        ws[f'{c}{r}']._style = copy.copy(TM[f'{L(CI(c) + TSHIFT)}{src}']._style)

def write_row(ws, row):
    r = row.r; arch = ARCH[row.arch]
    clone_row_style(ws, arch, r)
    ws.cell(r, 2).value = row.label or None
    if row.kw.get('note'): ws[f'{NOTE}{r}'] = row.kw['note']
    style_cagr(ws, r, row.key in CAGR_G or row.key in CAGR_A)
    if row.arch in ('section', 'block', 'blank', 'sub') and not row.kw.get('data'):
        return
    if row.key in DEFS:
        write_defs(ws, row); return
    nf = NF[row.nf]
    for col in COLS:
        v, inp = compute(row, col)
        if v is None: continue
        cell = ws[f'{col.c}{r}']
        cell.value = v
        if not row.kw.get('textdata'): cell.number_format = nf
        if inp:
            set_color(cell, BLUE, bold=True if col.kind in FCK else None)
        elif row.arch == 'growth' and col.kind in HISTK:
            set_color(cell, GRAY)
        else:
            set_color(cell, BLACK)
    if row.key in CAGR_G:
        ws[f'AW{r}'] = f'=IFERROR(({A[2031]}{r}/{A[2026]}{r})^(1/5)-1,"n/m")'
        ws[f'AX{r}'] = f'=IFERROR(({A[2026]}{r}/{A[2021]}{r})^(1/5)-1,"n/m")'
        ws[f'AY{r}'] = f'=IFERROR(({A[2026]}{r}/{A[2016]}{r})^(1/10)-1,"n/m")'
        ws[f'AZ{r}'] = f'=IFERROR(({A[2026]}{r}/{A[2012]}{r})^(1/14)-1,"n/m")'
        for c in CAGR: ws[f'{c}{r}'].number_format = '0%;\\(0%\\);"n/m"'
    elif row.key in CAGR_A:
        ws[f'AW{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2027, 2032))}),"n/m")'
        ws[f'AX{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2022, 2027))}),"n/m")'
        ws[f'AY{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2017, 2027))}),"n/m")'
        ws[f'AZ{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2012, 2027))}),"n/m")'
        for c in CAGR: ws[f'{c}{r}'].number_format = NF['pct']
    t = LABEL_SOURCES.get(row.key)
    if t: comment(ws.cell(r, 2), t)

SRC = 'Source: STERIS earnings releases (Form 8-K Ex. 99.1) — '
LABEL_SOURCES = {
    'hc_cap': SRC + '"Unaudited Supplemental Financial Data" revenue by type per segment (Healthcare Products + Healthcare Specialty Services (all service) FY2016–20; FY2017 from the FY2018 release comparative, scaled to the original segment revenue).',
    'hc_oi': SRC + 'Segment Data ("income from operations before adjustments"); FY2012–15: adjusted segment operating income from the Non-GAAP reconciliation tables.',
    'og_cc': SRC + 'Non-GAAP table "Impact of Acquisitions / Divestitures / Foreign Currency Movements" (published from FY2017).',
    'is_rev': ('Source: release consolidated condensed statements of operations (FY2012–FY2022 and FY2025+ as reported). FY2023 quarters: Dental discontinued-operations recast '
               'in the Q4 FY24 release (8-May-2024; opex lines derived — R&D / restructuring / other scaled to the FY2023 restated totals, SG&A plug); Q4 FY23 / FY2023 and '
               'FY2024: Q4 FY24 release; Q1–Q3 FY24: comparatives in the Q1–Q3 FY25 releases.'),
    'd_dda': 'Source: 10-K / 10-Q cash-flow statements via SEC XBRL (DepreciationDepletionAndAmortization, year to date as first filed; discrete quarters derived).',
    'o_pub': SRC + 'Non-GAAP Financial Measures ("Income from Operations" adjusted column; FY2012–15 vertical reconciliation "Adjusted operating income").',
    'e_pub': SRC + 'Non-GAAP Financial Measures (adjusted diluted EPS; from Q4 FY24: continuing operations).',
    'e_pub_orig': SRC + 'adjusted EPS as originally reported, incl. Dental (Q1 FY23 – Q3 FY24 releases); Q4 FY24 / FY2024: total company adjusted EPS in the Q4 FY24 release.',
    'e_disc': 'Source: Q4 FY24 release "Adjusted Quarterly Results (Non-GAAP)" recast tables (income from discontinued operations, net of tax, adjusted) and Non-GAAP tables.',
    'cf_ni': 'Source: release cash-flow statements (year to date; discrete quarters derived, Q4 = FY − 9M); D&A, share-based compensation and deferred taxes from the 10-K / 10-Q (XBRL).',
    'bs_cash': 'Source: release condensed balance sheets (quarter ends); goodwill / intangibles split FY2012–15 and noncontrolling interests / shares outstanding from XBRL.',
    'ppa_cons': 'Source: FY2016 10-K Note 3 (Synergy, preliminary); FY2023 10-K Note 2 (Cantel, final); FY2024 10-K Note 2 (BD surgical instrumentation, final).',
    'sh_basic': SRC + 'statements of operations (weighted average shares). STERIS Corp common shares to Nov-2015; STERIS plc (UK) to Mar-2019; STERIS plc (Ireland) since.',
    'ds_total': 'Source: release balance sheets; FY2026 10-K Note 8 (Debt).',
}

DEFS = {}
def defs_setup():
    q = lambda y, n: QC[(y, n)].c
    DEFS['o_def'] = {
        'C': ('FY2012–15 STERIS Corp: as reported; SYSTEM 1 rebate & class action adjusted out (FY2011–13); segments Healthcare / Life Sciences / Isomedix',
              'Segment operating income published on a GAAP basis; the adjusted segment figures from the Non-GAAP tables are used. Healthcare revenue net of the SYSTEM 1 rebate program.'),
        A[2016]: ('FY2016: Synergy combination (Nov-2015); segments Healthcare Products / Healthcare Specialty Services / Life Sciences / AST; segment OI before adjustments', None),
        A[2019]: ('FY2019+: corporate costs reported separately (previously largely allocated to segments)', None),
        A[2021]: ('FY2021+: Healthcare Products and Healthcare Specialty Services combined into Healthcare', None),
        q(2022, 1): ('Q1/22: Cantel (Jun-2021) — Dental segment; STERIS plc presentation as reported', None),
        q(2023, 1): ('FY2023+: continuing operations (Dental discontinued; recast in the Q4 FY24 release and FY25 releases)',
                     'FY2023 Q1–Q3: GAAP and adjusted totals from the recast; segment lines from the original releases (Healthcare, AST and Life Sciences unaffected); corporate = adjusted OI − segments.'),
        q(2026, 1): ('Q1/26+: releases presented in USD millions (one decimal) — checks use ±0.15 tolerance', None),
    }
    DEFS['e_def'] = {
        'C': ('Adjusted EPS = GAAP diluted EPS + adjusting items, net of tax (company definition throughout); FY2012–15 items shown net of tax', None),
        q(2022, 1): ('Q1/22: GAAP net loss — adjusted EPS on diluted shares (input, implied from the release)', None),
        q(2023, 1): ('FY2023–24: continuing operations (restated); as-originally-reported adjusted EPS incl. Dental rebuilt in the memo rows below', None),
        q(2025, 1): ('FY2025+: continuing operations as reported', None),
    }
def write_defs(ws, row):
    for c, (txt, cm) in DEFS[row.key].items():
        cell = ws[f'{c}{row.r}']; cell.value = txt
        cell.font = Font(name='Calibri', sz=10, b=True, i=True, color='FFC00000')
        if cm: comment(cell, cm)

def write_scen(ws):
    srcrow = {'hdr': 9, 'out_hdr': 11, 'pt_hdr': 15, 'out': 12, 'pt': 16, 'lev': 28, 'blk': 30, 'yr': 31}
    for spec_ in SCEN_ROWS:
        r, kind = spec_[0], spec_[1]
        for ci in range(CI(SL), CI(SNOTE) + 1):
            ws.cell(r, ci)._style = copy.copy(TM.cell(srcrow[kind], ci + TSHIFT)._style)
        if kind in ('hdr', 'out_hdr', 'pt_hdr'):
            _, _, a, b, c_, d_, e = spec_
            for col, v in zip([SL, SBULL, SBASE, SBEAR, SNOTE], [a, b, c_, d_, e]):
                if v is not None: ws[f'{col}{r}'] = v
        elif kind == 'out':
            _, _, lab, hi, mid, lo, note, nf = spec_
            ws[f'{SL}{r}'] = lab
            for col, v in zip([SBULL, SBASE, SBEAR], [hi, mid, lo]):
                ws[f'{col}{r}'] = v; ws[f'{col}{r}'].number_format = NF[nf]; set_color(ws[f'{col}{r}'], BLUE)
            ws[f'{SNOTE}{r}'] = note
        elif kind == 'pt':
            _, _, lab, _, val, _, note, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBASE}{r}'] = val; ws[f'{SBASE}{r}'].number_format = NF[nf]; set_color(ws[f'{SBASE}{r}'], BLUE); ws[f'{SNOTE}{r}'] = note
        elif kind == 'lev':
            _, _, lab, a, b, c_, note = spec_
            ws[f'{SL}{r}'] = lab
            for col, v in zip([SBULL, SBASE, SBEAR], [a, b, c_]):
                ws[f'{col}{r}'] = v; ws[f'{col}{r}'].number_format = NF['x']
            ws[f'{SNOTE}{r}'] = note
        elif kind == 'blk':
            _, _, lab, db, mid, de, note, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBULL}{r}'] = db; ws[f'{SBASE}{r}'] = mid; ws[f'{SBEAR}{r}'] = de; ws[f'{SNOTE}{r}'] = note
            for col in (SBULL, SBEAR):
                ws[f'{col}{r}'].number_format = '\\+#,##0;\\-#,##0' if nf == 'm' else '\\+0.0%;\\-0.0%'
        elif kind == 'yr':
            _, _, lab, hdr, base, clamp, _, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBASE}{r}'] = base
            if clamp:
                ws[f'{SBULL}{r}'] = f'=MAX(0,{SBASE}{r}+${SBULL}${hdr})'; ws[f'{SBEAR}{r}'] = f'=MAX(0,{SBASE}{r}+${SBEAR}${hdr})'
            else:
                ws[f'{SBULL}{r}'] = f'={SBASE}{r}+${SBULL}${hdr}'; ws[f'{SBEAR}{r}'] = f'={SBASE}{r}+${SBEAR}${hdr}'
            for col in (SBULL, SBASE, SBEAR): ws[f'{col}{r}'].number_format = NF[nf]


def write_header(ws):
    for r in range(1, 8):
        clone_row_style(ws, r, r, 2, CI(LASTC))
    ws[f'{NOTE}1'] = 'SCENARIO SWITCH ▼ (Bull / Base / Bear) — drives all scenario-linked forecast drivers'
    ws[f'{NOTE}2'] = 'Base'
    comment(ws[f'{NOTE}2'], 'Scenario toggle. Choose Bull / Base / Bear. Drives: FY27 revenue-growth and adjusted-EPS outlook end-points (high / mid / low), '
            'Q2–Q4/27E buybacks, 2028E–2031E segment revenue-by-type growth, segment margins, M&A spend and the leverage target for 2028E+ buybacks '
            f'(scenario table {SL}:{SNOTE}).')
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type='list', formula1='"Bull,Base,Bear"', allow_blank=False); ws.add_data_validation(dv); dv.add(f'{NOTE}2')
    ws['B2'] = 'STERIS plc (NYSE: STE) — 3-Statement Financial Model'
    ws['B3'] = ('US GAAP as reported · USD millions (per-share data in USD; shares in millions) · Fiscal years ended 31 March (2027 = Apr-2026 – Mar-2027) · '
                'Source: STERIS Forms 10-K / 10-Q and Form 8-K earnings releases (Ex. 99.1), SEC EDGAR (STERIS Corp CIK 815065 to FY2015; STERIS plc (UK) CIK 1624899 FY2016–19; '
                'STERIS plc (Ireland) CIK 1757898 from FY2019); investor-relations releases at steris-ir.com are the same documents as filed on EDGAR')
    ws['B4'] = ('Basis notes: FY2012–FY2022 as originally reported; FY2023 onward continuing operations (Dental, acquired with Cantel Jun-2021, sold Sep-2024, recast in the Q4 FY24 '
                'and FY25 releases) — originally reported adjusted EPS rebuilt and checked in memo rows. Segments: Healthcare, Applied Sterilization Technologies (AST), Life Sciences '
                '(+ Dental FY2022). Acquisitions: Synergy Health Nov-15, Cantel Jun-21, BD surgical instrumentation Aug-23. Q2/27E–Q4/27E calibrated to FY27 revenue-growth and adjusted-EPS guidance.')
    ws[f'{CAGR[0]}5'] = 'Forecast'
    for c in CAGR[1:]: ws[f'{c}5'] = 'Historical'
    ws['B6'] = '(USDm)'
    for col in COLS: ws[f'{col.c}6'] = col.label
    ws[f'{NOTE}6'] = 'Modelling Notes'
    for c, a, b in zip(CAGR, ['5Y CAGR', '5Y CAGR', '10Y CAGR', '14Y CAGR'], ["26-'31", "21-'26", "16-'26", "12-'26"]):
        ws[f'{c}6'] = a; ws[f'{c}7'] = b

def set_columns(ws):
    for k in list(ws.column_dimensions.keys()): del ws.column_dimensions[k]
    ws.column_dimensions['A'].width = TM.column_dimensions['A'].width
    ws.column_dimensions['B'].width = TM.column_dimensions['B'].width
    wid = {'A': 13.0, 'Q': 9.5703125, 'QE': 11.42578125, 'E26': 14.140625, 'AE': 13.0}
    for col in COLS:
        cd = ws.column_dimensions[col.c]; cd.width = wid[col.kind]
        if col.q:
            cd.outlineLevel = 1; cd.hidden = True
    for ci in range(CI(NOTE), CI(LASTC) + 1):
        ws.column_dimensions[L(ci)].width = TM.column_dimensions[L(ci + TSHIFT)].width or 13.0
    ws.sheet_format.outlineLevelCol = 1

def main():
    import openpyxl
    wb = openpyxl.load_workbook(TEMPLATE)
    if 'DCF_old' in wb.sheetnames: del wb['DCF_old']
    ws = wb['Model']
    blank_style = copy.copy(ws['A1']._style)
    for rowc in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for c in rowc:
            c.value = None; c.comment = None; c._style = copy.copy(blank_style)
    for k in list(ws.row_dimensions.keys()):
        rd = ws.row_dimensions[k]; rd.outlineLevel = 0; rd.hidden = False
    ws.data_validations.dataValidation = []
    scen_layout(V0)
    build_spec(V0)
    n1 = len(fw.SPEC)
    build_spec2(V0)
    for i, row in enumerate(fw.SPEC): row.spec2 = i >= n1
    last = number_rows(8)
    for r in range(1, last + 1): ws.row_dimensions[r].height = 15.0
    ws.row_dimensions[2].height = 21.0
    defs_setup()
    write_header(ws)
    for row in fw.SPEC: write_row(ws, row)
    write_scen(ws)
    set_columns(ws)
    ws.freeze_panes = 'C8'
    import other_sheets as osh
    osh.do_all(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    osh.inject_charts(OUT)
    print('rows', fw.SPEC[-1].r, 'saved', OUT)
    return wb

RECALC = os.environ.get('XLSX_RECALC', '/mnt/skills/public/xlsx/scripts/recalc.py')
def recalc(path, timeout=600):
    import subprocess
    out = subprocess.run(['python3', RECALC, path, str(timeout)], capture_output=True, text=True, cwd=os.path.dirname(RECALC))
    try:
        return json.loads(out.stdout)
    except Exception:
        return {'raw': out.stdout[-2000:], 'err': out.stderr[-2000:]}

def scenario_snapshots():
    import openpyxl
    import other_sheets as osh
    snaps = {}
    for scen in ('Bull', 'Bear'):
        p = os.path.join(BASE, f'STE_{scen}.xlsx')
        wb = openpyxl.load_workbook(OUT); wb['Model'][f'{NOTE}2'] = scen; wb.save(p)
        res = recalc(p); print(scen, res.get('status'), res.get('total_errors'))
        v = openpyxl.load_workbook(p, data_only=True)['Bull-Base-Bear']
        snaps[scen] = {(r, c): v[f'{c}{r}'].value for r in range(9, 51) for c in 'CDEFGHIJKL' if v[f'{c}{r}'].value is not None}
        os.remove(p)
    wb = openpyxl.load_workbook(OUT)
    osh.snapshot_values(wb, snaps['Bull'], snaps['Bear'])
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    osh.inject_charts(OUT)

if __name__ == '__main__':
    main()
    if '--snap' in sys.argv: scenario_snapshots()
