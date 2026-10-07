"""Rewire the Bull-Base-Bear, DCF and Charts tabs of the MLM template (current layout: DCF with live valuation date, fade period,
implied ROIC, target-IRR entry price, reverse DCF and Mauboussin EV/NOPAT; Bull-Base-Bear with buyback rows) to the STE Model rows /
columns; keep the template chart XML. The template's DCF_old tab is not carried."""
import re, zipfile, shutil
from openpyxl.utils import get_column_letter as L
from fw import R, TEMPLATE, ANN, NOTE

# template Model row -> STE row key
MAP = {158: 'is_rev', 164: 'r_cash', 166: 'is_op', 182: 'is_dda', 186: 'sh_dil', 211: 'r_adj', 212: 'r_adj_m', 219: 'r_ebit', 221: 'r_adj',
       229: 'e_adjni', 231: 'e_eps', 238: 'gm_org', 245: 'gm_ebitda_m', 247: 'gm_op_m', 260: 'cf_wc', 262: 'cf_capex', 263: 'cf_acq',
       280: 'fcf', 282: 'fcf_m', 290: 'bb_px', 291: 'bb_cash', 292: 'bb_sh', 331: 'bs_nci', 408: 'ra_ic', 411: 'ra_roic', 422: 'rn_ronta',
       428: 'lv_nd', 433: 'v_px', 434: 'v_sh'}
# template Model column -> STE column (template 2025A = STE FY2026A; template 2026E-2030E = STE FY2027E-FY2031E)
CMAP = {'BV': ANN[2026].c, 'CA': ANN[2027].c, 'CB': ANN[2028].c, 'CC': ANN[2029].c, 'CD': ANN[2030].c, 'CE': ANN[2031].c, 'CF': NOTE}
REF = re.compile(r"Model!(\$?)([A-Z]{1,3})(\$?)(\d+)")
HIST = [ANN[y].c for y in range(2012, 2027)]

def remap(f):
    def rep(m):
        c, r = m.group(2), int(m.group(4))
        c2 = CMAP.get(c, c)
        r2 = R[MAP[r]] if r in MAP else r
        if c == 'CF' and r == 2: r2 = 2
        return f'Model!{m.group(1)}{c2}{m.group(3)}{r2}'
    return REF.sub(rep, f)

def hist_stat(f):
    """LT historical AVERAGE / MIN / MAX over the template's annual columns -> STE FY2012-FY2026 annual columns."""
    m = re.match(r'=IFERROR\((AVERAGE|MIN|MAX)\(Model!C(\d+),', f)
    if not m: return None
    fn, r = m.group(1), int(m.group(2))
    rr = R[MAP[r]]
    return f'=IFERROR({fn}({",".join(f"Model!{c}{rr}" for c in HIST)}),"")'

def fy_dates(f):
    f = re.sub(r'DATE\((\d{4}),12,31\)', lambda m: f'DATE({int(m.group(1)) + 1},3,31)', f)
    f = re.sub(r'DATE\((\d{4})-1,12,31\)', lambda m: f'DATE({int(m.group(1)) + 1}-1,3,31)', f)
    return f

TEXT = {
    'Bull-Base-Bear': {
        'B2': 'STERIS plc (STE) — Bull / Base / Bear Case P&L Summary',
        'B3': ('Adjusted figures (company definitions: adjusted income from operations, adjusted EPS from continuing operations; adjusted EBITDA = model definition) · '
               'USD millions · fiscal years ended 31 March · Live panel linked to Model — toggle scenario via Model!AU2. Bull and Bear panels are value snapshots generated '
               'from the Model with the switch set to each case. FY2027E calibrated to the FY27 outlook (revenue growth 7–8%, adjusted EPS $11.10–11.30).'),
        'G26': 22, 'H26': 22, 'S26': 25, 'T26': 25, 'AD26': 17, 'AE26': 17, 'AF26': 17,
    },
    'DCF': {
        'B2': 'DCF VALUATION — STERIS PLC (STE)',
        'C5': 0.045, 'D5': 'Assumption (~US 10Y Treasury yield, Oct-2026); update to market',
        'C7': 0.85, 'D7': 'Medtech / infection-prevention peers (Ecolab, Getinge, Cantel pre-deal, Sotera Health) ~0.7–1.0; ~75% recurring consumables & service revenue dampens cyclicality',
        'C9': 0.050, 'D9': 'Investment grade (BBB+ / Baa1); 2.70% 2031 and 3.75% 2051 public notes; marginal cost ~5%',
        'C10': 0.25, 'D10': 'Model FY2027E+ effective tax rate (FY27 outlook ~25%)',
        'C12': 0.08, 'D12': 'Net debt ~$1.4bn at 30-Jun-2026 vs ~$20.6bn market cap ($211.07 × ~98m diluted shares); ~0.9x adjusted EBITDA',
        'D15': 'Periods = (31-Mar of fiscal year − today) / 365.25; cash flows assumed at fiscal year-end (31 March)',
        'B18': 'Adj. EBIT less acquisition, integration & restructuring costs',
        'B21': 'plus DD&A (total, incl. acquired-intangible amortization)',
        'B22': '− Capex (PP&E & intangibles) and acquisitions (tuck-ins; M&A lever 2028E–31E)',
        'B26': 'Fade-period growth (2032E–36E; 6.0% → 4.5% linear)', 'I26': 0.06, 'M26': 0.045, 'C27': 0.035,
        'D32': ("STERIS's adjusted operating income excludes acquired-intangible amortization (~$262m in FY27E) but expenses share-based compensation; Adjusted EBIT "
                'here deducts all D&A (non-cash, added back in row 21). Acquisition & integration and restructuring charges (cash; adjusted out by the company, incl. the '
                'FY27–30 North Carolina Center-of-Excellence plan) are deducted inside row 18. Row 22 includes tuck-in acquisitions in FY2027E–31E; the fade period and '
                'terminal value assume no new acquisitions. Dental (sold Sep-2024) is excluded throughout (continuing operations).'),
        'D43': 'Opening IC = Model FY2026A book IC (total equity incl. NCI + net debt), rolled forward with capex + M&A − DD&A + ΔWC.',
        'B48': 'Less: Net debt + NCI (Model FY2027E year end, 31-Mar-2027)',
        'D48': 'STE: noncontrolling interests (~$14m) deducted alongside net debt.',
        'B50': 'Diluted shares (m, Model FY2027E period end)',
        'D76': 'Rationale (STERIS characteristics)',
        'C78': 0.08, 'D78': 'Model adjusted EBIT CAGR FY2027–31E: mid-to-high single-digit revenue growth (procedures, service, AST capacity), ~50 bps p.a. margin expansion and buybacks → ~8% NOPAT growth',
        'C79': 0.18, 'D79': 'Razor / razor-blade installed base (sterilizers, washers, endoscope reprocessors) drives consumables & service at high incremental returns; AST sterilization network earns ~45% segment margins on capacity adds; acquisitions (Synergy, Cantel) earn ~7–9% after tax',
        'C80': 15, 'D80': 'Regulated, validated sterilization processes and hospital standardization create high switching costs; ~75% recurring revenue; EO regulatory and litigation exposure is the main risk to the CAP',
        'B90': 'VALUATION  —  Applied to FY2029E NOPAT',
        'B91': 'FY2029E Adj. EBIT less acquisition, integration & restructuring costs',
        'B93': 'FY2029E NOPAT',
        'B95': 'Implied Enterprise Value (FY2029E basis)',
        'B96': 'Less: FY2028E year-end net debt + NCI (Model)',
        'B98': 'Diluted shares FY2029E (m, Model)',
        'B99': 'Implied share price FY2029 ($)',
        'B106': 'Plus: FY2027E year-end net debt + NCI',
        'B108': 'STE market-implied EV/NOPAT (on FY2029E NOPAT)',
    },
    'Charts': {
        'B2': 'STERIS plc — Operating Performance (FY2012–FY2026) and Model Forecast (FY2027E–FY2031E)',
        'B3': ('USDm · fiscal years ended 31 March · Linked to Model. Bars = constant-currency organic growth vs acquisitions, divestitures & FX (published bridge from FY2017; '
               'earlier years: total growth shown as organic). Adjusted EBITDA = model definition (adjusted operating income + depreciation & other amortization). '
               'ROIC / RONTA on adjusted EBIT (full-cost).'),
        'B8': 'CC organic growth %', 'B9': 'Acquisitions, divestitures & FX %', 'B12': 'Adjusted EBITDA', 'B13': 'Income from Operations (GAAP)',
    },
}
BBB_ROWS = {9: 'Revenues', 10: '   Revenues y/y %', 11: '   Constant-currency organic growth %', 12: 'Adjusted EBITDA (model)',
            15: 'Income from Operations (GAAP)', 16: '   Income from operations y/y %', 19: 'Adjusted Net Income (continuing operations)',
            35: 'Adjusted EBITDA (leverage basis)', 36: '   Adjusted EBITDA % of revenues'}
CHART_TITLES = {
    'Aggregates Revenue Drivers: Shipment Volume vs. Price (ASP), y/y %': 'Revenue Growth: Constant-Currency Organic vs. Acquisitions, Divestitures &amp; FX, y/y %',
    'Adjusted EBITDA &amp; Earnings from Operations (USDm)': 'Adjusted EBITDA &amp; Income from Operations (USDm)',
}

def build_charts(ws):
    for i, y in enumerate(range(2012, 2032)):
        c = L(3 + i); mc = ANN[y].c
        ws[f'{c}7'] = y if y <= 2026 else f'{y}E'
        ws[f'{c}8'] = f'=IFERROR(Model!{mc}{R["og_cc"]}+0,IFERROR(Model!{mc}{R["gm_rev"]}+0,0))'
        ws[f'{c}9'] = f'=IFERROR(Model!{mc}{R["gm_rev"]}-Model!{mc}{R["og_cc"]},0)'
        ws[f'{c}11'] = f'={c}7'
        ws[f'{c}12'] = f'=Model!{mc}{R["r_adj"]}'
        ws[f'{c}13'] = f'=Model!{mc}{R["is_op"]}'
        ws[f'{c}14'] = f'=Model!{mc}{R["gm_ebitda_m"]}'
        ws[f'{c}15'] = f'=Model!{mc}{R["gm_op_m"]}'
        ws[f'{c}17'] = f'={c}7'
        ws[f'{c}18'] = f'=IFERROR(Model!{mc}{R["ra_roic"]}+0,NA())'
        ws[f'{c}19'] = f'=IFERROR(Model!{mc}{R["rn_ronta"]}+0,NA())'
        ws[f'{c}20'] = f'=IFERROR(Model!{mc}{R["gm_ebitda_m"]}+0,NA())'
        ws[f'{c}21'] = f'=IFERROR(Model!{mc}{R["fcf_m"]}+0,NA())'

def do_all(wb):
    for name in ('Bull-Base-Bear', 'DCF', 'Charts'):
        ws = wb[name]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and 'Model!' in c.value:
                    h = hist_stat(c.value) if name == 'Bull-Base-Bear' else None
                    c.value = h or remap(c.value)
                if isinstance(c.value, str) and 'DATE(' in c.value:
                    c.value = fy_dates(c.value)
        for addr, t in TEXT.get(name, {}).items():
            ws[addr] = t
    b = wb['Bull-Base-Bear']
    for r, t in BBB_ROWS.items():
        for c in ('B', 'N', 'Z'): b[f'{c}{r}'] = t
    for c0 in ('C', 'O', 'AA'):
        i = b[f'{c0}7'].column
        b.cell(7, i).value = '2026A'
        for k in range(1, 6): b.cell(7, i + k).value = f'{2026 + k}E'
        b.cell(7, i + 6).value = "26-'31 CAGR"
    b['J2'] = 'Share price ($):'
    for c in ('B', 'N', 'Z'): b[f'{c}29'] = 'IRR (to 31-Mar of Y-1)'
    for c in ('B', 'N', 'Z'):
        if b[f'{c}40'].value: b[f'{c}40'] = '   Share repurchases ($m; FY2027E = Q2–Q4)'
    d = wb['DCF']
    for i, col in enumerate('DEFGHIJKLM'):
        d[f'{col}17'] = f'{2027 + i}E'
    a26, a27 = ANN[2026].c, ANN[2027].c
    d['D22'] = f'=Model!{a27}{R["cf_capex"]}+Model!{a27}{R["cf_acq"]}'
    d['D34'] = f'=Model!{a26}{R["ra_ic"]}'
    build_charts(wb['Charts'])

def snapshot_values(wb, bull, bear):
    ws = wb['Bull-Base-Bear']
    cols = 'CDEFGHIJKL'
    for snap, off in ((bull, 'OPQRSTUVWX'), (bear, ['AA', 'AB', 'AC', 'AD', 'AE', 'AF', 'AG', 'AH', 'AI', 'AJ'])):
        for (r, c), v in snap.items():
            tgt = f'{off[cols.index(c)]}{r}'
            cur = ws[tgt].value
            if isinstance(cur, str) and cur.startswith('='): continue
            if r in (7, 26): continue
            ws[tgt] = v

def inject_charts(path):
    """Replace openpyxl-rewritten chart parts with the template chart XML (combo charts, colours, layout), retitled, ranges kept
    (Charts!C:V = FY2012-FY2031E) and cached values removed so Excel redraws from the STE cells."""
    tz = zipfile.ZipFile(TEMPLATE)
    tmpl = {n: tz.read(n) for n in tz.namelist() if n.startswith('xl/charts/chart')}
    tmp = path + '.tmp'
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in tmpl:
                t = tmpl[item.filename].decode('utf-8')
                for a, b2 in CHART_TITLES.items(): t = t.replace(a, b2)
                t = re.sub(r'<c:numCache>.*?</c:numCache>', '', t, flags=re.S)
                t = re.sub(r'<c:strCache>.*?</c:strCache>', '', t, flags=re.S)
                data = t.encode('utf-8')
            zout.writestr(item, data)
    shutil.move(tmp, path)
