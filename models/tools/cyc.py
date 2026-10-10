"""Scan the Model sheet of a workbook for circular references (prints up to five cycles with their row labels).

Usage: python3 cyc.py <model.xlsx>

The MLM template is built to have none: interest runs on opening balances and buybacks / revolvers plug on pre-financing cash.
"""
import re
import sys

import openpyxl
from openpyxl.utils import column_index_from_string as CI, get_column_letter as L

sys.setrecursionlimit(100000)
REF = re.compile(r"(?<![A-Za-z!])\$?([A-Z]{1,3})\$?(\d+)(?::\$?([A-Z]{1,3})\$?(\d+))?")


def main(path):
    ws = openpyxl.load_workbook(path)['Model']
    graph = {}
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if isinstance(v, str) and v.startswith('='):
                deps = set()
                for m in REF.finditer(re.sub(r'"[^"]*"', '', v)):
                    a, r1, b, r2 = m.groups()
                    if b:
                        for ci in range(CI(a), CI(b) + 1):
                            for r in range(int(r1), int(r2) + 1):
                                deps.add(f'{L(ci)}{r}')
                    else:
                        deps.add(f'{a}{r1}')
                graph[c.coordinate] = deps

    color, cycles = {}, []

    def dfs(u, stack):
        color[u] = 1
        stack.append(u)
        for w in graph.get(u, ()):
            if color.get(w) == 1:
                cycles.append(stack[stack.index(w):] + [w])
                if len(cycles) > 5:
                    return
            elif color.get(w) is None and w in graph:
                dfs(w, stack)
        stack.pop()
        color[u] = 2

    for u in list(graph):
        if color.get(u) is None:
            dfs(u, [])
        if len(cycles) > 5:
            break

    def label(addr):
        return str(ws.cell(int(re.sub('[A-Z]', '', addr)), 2).value).strip()[:30]

    for cyc in cycles[:5]:
        print([f'{a}:{label(a)}' for a in cyc][:12])
    print('cycles', len(cycles))
    return 1 if cycles else 0


if __name__ == '__main__':
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    sys.exit(main(sys.argv[1]))
