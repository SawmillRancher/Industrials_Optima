"""Framework: column model, style cloning from the TDY template, row registry and writer for the VMC model."""
import copy, json, os, glob
import openpyxl
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI
from openpyxl.styles import Font
from openpyxl.comments import Comment

BASE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(BASE, 'TDY_Model.xlsx')

# ---------------------------------------------------------------- columns
class Col:
    def __init__(self, c, label, kind, year, q=None):
        self.c, self.label, self.kind, self.year, self.q = c, label, kind, year, q
        self.prior = None   # prior-year comparable column
        self.prev = None    # previous annual column (annual / forecast annual only)
    def __repr__(self): return f'{self.c}:{self.label}'

COLS = []
ANN, QC, HC = {}, {}, {}
_i = CI('C')
for y in range(2006, 2013):
    col = Col(L(_i), y, 'A', y); COLS.append(col); ANN[y] = col; _i += 1
for y in range(2013, 2027):
    yy = str(y)[2:]
    for lab, kind, q in [(f'Q1/{yy}', 'Q', 1), (f'Q2/{yy}', 'Q', 2), (f'1H/{yy}', 'H', None),
                         (f'Q3/{yy}', 'Q', 3), (f'Q4/{yy}', 'Q', 4), (y, 'A', None)]:
        if y == 2026:
            if q in (3, 4): kind, lab = 'QE', f'Q{q}/{yy}E'
            if kind == 'A': kind, lab = 'E26', '2026E'
        col = Col(L(_i), lab, kind, y, q); COLS.append(col); _i += 1
        if q: QC[(y, q)] = col
        elif kind == 'H': HC[y] = col
        else: ANN[y] = col
for y in range(2027, 2031):
    col = Col(L(_i), f'{y}E', 'AE', y); COLS.append(col); ANN[y] = col; _i += 1
assert ANN[2030].c == 'CS' and ANN[2026].c == 'CO' and HC[2026].c == 'CL', (ANN[2030], ANN[2026], HC[2026])
for col in COLS:
    if col.kind in ('A', 'E26', 'AE'):
        if col.year - 1 in ANN: col.prior = col.prev = ANN[col.year - 1].c
    elif col.kind in ('Q', 'QE'):
        if (col.year - 1, col.q) in QC: col.prior = QC[(col.year - 1, col.q)].c
    elif col.kind == 'H':
        if col.year - 1 in HC: col.prior = HC[col.year - 1].c
COLMAP = {c.c: c for c in COLS}
HIST = [c for c in COLS if c.kind in ('A', 'Q', 'H')]
FC = [c for c in COLS if c.kind in ('QE', 'E26', 'AE')]
AE = [c for c in COLS if c.kind == 'AE']

def pkey(col):
    """data-file key for a historical column."""
    if col.kind == 'A': return str(col.year)
    if col.kind == 'Q': return f'Q{col.q}-{str(col.year)[2:]}'
    return None

# ---------------------------------------------------------------- data
DATA = {}
def load_data():
    for f in glob.glob(os.path.join(BASE, 'data', '*.json')):
        k = os.path.basename(f)[:-5]
        try:
            DATA[k] = json.load(open(f))
        except Exception as e:
            print('BAD JSON', f, e)
    return DATA

def num(x):
    if isinstance(x, bool) or x is None: return None
    if isinstance(x, (int, float)): return x
    try:
        return float(str(x).replace(',', '').replace('$', ''))
    except Exception:
        return None

def get(k, path):
    d = DATA.get(k)
    if d is None: return None
    for p in path.split('.'):
        if not isinstance(d, dict): return None
        d = d.get(p)
        if d is None: return None
    return d

# ---------------------------------------------------------------- styles
_tmpl_wb = openpyxl.load_workbook(TEMPLATE)
TM = _tmpl_wb['Model']
ARCH = {  # row archetypes in the TDY template
    'section': 8, 'blank': 9, 'block': 10, 'line': 13, 'line_b': 99, 'total': 16, 'growth': 12,
    'pct': 20, 'memo': 17, 'check': 18, 'driver': 107, 'text': 197, 'sub': 198, 'eps_total': 181,
    'eps_memo': 182, 'ratio_total': 406, 'cf_line': 263, 'cf_total': 271, 'bs_line': 315,
    'bs_total': 326, 'header6': 6, 'legacy': 86,
}
NF = {
    'm': '#,##0;\\(#,##0\\);\\-', 'm1': '#,##0.0;\\(#,##0.0\\);\\-', 'pct': '0.0%;\\(0.0%\\);\\-',
    'ps': '0.00;\\(0.00\\);\\-', 'x': '0.0\\x', 'x2': '0.00\\x;\\(0.00"x)";\\-', 'd': '0\\d;\\(0")d";\\-',
    'gen': 'General', 'yrs': '0" yrs"',
}
BLUE, BLACK, GRAY = 'FF0000FF', 'FF000000', 'FFA6A6A6'
LASTCOL = CI('DF')

def clone_row_style(ws, src_row, dst_row, cols=range(2, CI('CY') + 1)):
    for ci in cols:
        s = TM.cell(src_row, ci); d = ws.cell(dst_row, ci)
        if s.has_style:
            d._style = copy.copy(s._style)

def set_color(cell, rgb, bold=None):
    f = copy.copy(cell.font)
    cell.font = Font(name=f.name, sz=f.sz, b=f.b if bold is None else bold, i=f.i, u=f.u, color=rgb)

# ---------------------------------------------------------------- registry
class Row:
    def __init__(self, key, label, arch='line', nf='m', **kw):
        self.key, self.label, self.arch, self.nf = key, label, arch, nf
        self.kw = kw
SPEC = []
R = {}
def add(key, label='', arch='line', nf='m', **kw):
    if key in R: raise KeyError('dup ' + key)
    SPEC.append(Row(key, label, arch, nf, **kw))
    R[key] = None
    return key
_blank_n = [0]
def blank():
    _blank_n[0] += 1
    return add(f'_blank{_blank_n[0]}', '', 'blank', 'gen')

def number_rows(start):
    r = start
    for row in SPEC:
        R[row.key] = r; row.r = r; r += 1
    return r

# current column context used by formula helpers
class _Ctx: col = None
CTX = _Ctx()
def V(key, c=None):
    return f'{c or CTX.col.c}{R[key]}'
def P(key):   # prior-year comparable
    return f'{CTX.col.prior}{R[key]}'
def PV(key):  # previous annual (forecast years)
    return f'{CTX.col.prev}{R[key]}'
def A(key, c):  # absolute
    return f'${c}${R[key]}'
