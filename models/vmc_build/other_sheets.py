"""Rewire Bull-Base-Bear, DCF and Charts sheets to the VMC Model rows."""
from fw import *
from openpyxl.comments import Comment

FC_COLS = ['CI', 'CO', 'CP', 'CQ', 'CR', 'CS']           # 2025A, 2026E..2030E
HIST_ANN = [ANN[y].c for y in range(2006, 2026)]

def hist_stat(fn, key):
    refs = ','.join(f'Model!{c}{R[key]}' for c in HIST_ANN)
    return f'=IFERROR({fn}({refs}),"")'

BBB_MAP = {  # row: (label, model key or None, stat key for J/K/L)
    9: ('Total Revenues', 'is_rev', None),
    11: ('   Organic growth %', 'gm_org', None),
    12: ('Adjusted EBITDA', 'r_cur', None),
    13: ('   Adjusted EBITDA margin %', None, 'gm_ebitda_m'),
    15: ('Operating Earnings (GAAP)', 'is_op', None),
    17: ('   Operating margin (GAAP) %', None, 'gm_op_m'),
    19: ('Adjusted Net Earnings (attributable to Vulcan)', 'e_cur', None),
    21: ('FDSO (m)', 'sh_dil', None),
    23: ('Adjusted EPS — Diluted ($)', 'e_cur_eps', None),
    31: ('Free Cash Flow', 'fcf', None),
    32: ('   FCF %', None, 'fcf_m'),
    35: ('Adjusted EBIT (model)', 'o_ebit', None),
    36: ('   Adjusted EBIT margin %', None, 'gm_ebit_m'),
    38: ('Net Debt', 'lev_nd', None),
    41: ('ROIC', 'ra_roic', 'ra_roic'),
    42: ('RONTA', 'rn_ronta', 'rn_ronta'),
}

def do_bbb(wb):
    ws = wb['Bull-Base-Bear']
    ws['B2'] = 'Vulcan Materials (VMC) — Bull / Base / Bear Case P&L Summary'
    ws['B3'] = 'Adjusted figures (company definitions) · USD millions · Live panel linked'
    ws['K2'] = f'=Model!$CO${R["val_px"]}'
    for r, (lab, key, stat) in BBB_MAP.items():
        for off in (0, 12, 24):  # live panel B, Bull snapshot N, Bear snapshot Z
            ws.cell(r, 2 + off).value = lab
        if key:
            for i, c in enumerate(FC_COLS):
                ws.cell(r, 3 + i).value = f'=Model!{c}{R[key]}'
        if stat:
            ws[f'J{r}'] = hist_stat('AVERAGE', stat); ws[f'K{r}'] = hist_stat('MIN', stat); ws[f'L{r}'] = hist_stat('MAX', stat)
    # margins / ratios on the live panel
    for i, c in enumerate('CDEFGH'):
        ws[f'{c}13'] = f'=IFERROR({c}12/{c}9,"")'
        ws[f'{c}36'] = f'=IFERROR({c}35/{c}9,"")'
        ws[f'{c}39'] = f'=IFERROR({c}38/{c}12,"")'
    for p in ('B', 'N', 'Z'):
        ws[f'{p}39'] = '   Net Debt / Adjusted EBITDA (x)'
        ws[f'{p}10'] = '   Revenues y/y %'
        ws[f'{p}16'] = '   Operating earnings y/y %'
    for c, mc in zip('EFGH', ['CP', 'CQ', 'CR', 'CS']):
        ws[f'{c}45'] = f'={c}38+Model!{mc}{R["val_nci"]}'
        ws[f'{c}47'] = f'=IFERROR({c}46/{c}12,"")'
    for c, mc in zip('QRST', ['CP', 'CQ', 'CR', 'CS']):
        ws[f'{c}45'] = f'={c}38+Model!{mc}{R["val_nci"]}'
        ws[f'{c}47'] = f'=IFERROR({c}46/{c}12,"")'
    for c, mc in zip(['AC', 'AD', 'AE', 'AF'], ['CP', 'CQ', 'CR', 'CS']):
        ws[f'{c}45'] = f'={c}38+Model!{mc}{R["val_nci"]}'
        ws[f'{c}47'] = f'=IFERROR({c}46/{c}12,"")'
    for p in ('B', 'N', 'Z'):
        ws[f'{p}26'] = 'P / E (x) — assumed (Adjusted EPS)'
        ws[f'{p}47'] = 'Implied EV / Adjusted EBITDA (x)'
    ws['G26'] = 28; ws['H26'] = 28
    ws['S26'] = 32; ws['T26'] = 32
    ws['AD26'] = 22; ws['AE26'] = 22; ws['AF26'] = 22
    c = Comment('Bull P/E on Adjusted EPS (company definition). Vulcan traded ~25–35x NTM Adjusted EPS 2021–26 '
                '(≈26x at $243 on the model 2026E Adjusted EPS). Base 28x, Bear 22x (aggregates down-cycle trough).', 'VMC model')
    c.width, c.height = 380, 120
    ws['N26'].comment = c

def snapshot_values(wb, values_bull, values_bear):
    """values_*: dict {(row, col_letter_live): value} read from recalculated scenario runs (live panel C..L)."""
    ws = wb['Bull-Base-Bear']
    live_cols = list('CDEFGHIJKL')
    for vals, cols in ((values_bull, list('OPQRSTUVWX')), (values_bear, ['AA', 'AB', 'AC', 'AD', 'AE', 'AF', 'AG', 'AH', 'AI', 'AJ'])):
        for (r, lc), v in vals.items():
            if r in (26, 27, 28, 29, 34, 44, 45, 46, 47): continue   # formula rows stay live in the snapshot
            tc = cols[live_cols.index(lc)]
            ws[f'{tc}{r}'] = v

def do_dcf(wb):
    ws = wb['DCF']
    M = lambda key, c: f'Model!{c}{R[key]}'
    ws['B2'] = 'DCF VALUATION — VULCAN MATERIALS (VMC)'
    ws['C7'] = 0.95
    ws['D7'] = 'Aggregates peers (MLM, CRH, KNF, EXP) ~0.9–1.1; infrastructure-weighted demand and pricing power dampen cyclicality'
    ws['C9'] = 0.0525
    ws['D9'] = 'BBB+ / Baa2 (investment grade); notes 3.50%–7.15% coupons, weighted-average effective rate 5.04% (30-Jun-26); marginal cost ~5.0–5.5%'
    ws['C10'] = 0.22
    ws['D10'] = 'Model 2027E+ effective tax rate (FY26 guide 22–23%)'
    ws['C12'] = 0.12
    ws['D12'] = 'Net debt ~$4.1bn vs ~$31bn market cap ($243 × ~129.6m shares); net debt / TTM Adjusted EBITDA 1.7x at 30-Jun-26'
    ws['B18'] = 'Adjusted EBIT (Model; Adjusted EBITDA − DDA&A)'
    ws['B20'] = '(−) Adjusting items (cash charges excluded from Adjusted EBITDA)'
    ws['B22'] = 'NOPAT  =  (Adjusted EBIT − adjusting items) × (1 − t)'
    ws['B23'] = 'plus depreciation, depletion, accretion & amortization'
    for col, mc in zip('DEFGH', ['CO', 'CP', 'CQ', 'CR', 'CS']):
        ws[f'{col}18'] = f'={M("o_ebit", mc)}'
        ws[f'{col}20'] = f'=-{M("o_adj", mc)}'
        ws[f'{col}23'] = f'={M("is_dda", mc)}'
        ws[f'{col}25'] = f'={M("cf_wc", mc)}'
        ws[f'{col}29'] = f'={M("cf_capex", mc)}'
        if col != 'D':
            ws[f'{col}24'] = f'={col}29+{M("cf_acq", mc)}'
    ws['I31'] = 0.07; ws['M31'] = 0.05; ws['C32'] = 0.04
    ws['D37'] = ("Vulcan's Adjusted EBITDA (and the model Adjusted EBIT) already expense share-based compensation; only gains on sale, divested-operation "
                 "charges, acquisition, restructuring and similar items are excluded (deducted again in row 20 when forecast). Acquisition spend from the M&A "
                 "lever is deducted in 2027E–30E so that acquired revenues / EBITDA in the Model are paid for; 2026E excludes the 1H/26 California ready-mix "
                 "sale proceeds and bolt-ons. The fade period and terminal value assume no new acquisitions. Operating-lease costs sit inside Adjusted EBIT.")
    ws['D42'] = f'={M("ra_ic", "CI")}'
    ws['D48'] = (f'Opening IC = Model 2025A book IC (equity + net debt, Model!CI{R["ra_ic"]}), rolled forward with DCF net investment (capex + M&A − DDA&A + ΔWC). '
                 'Terminal incremental ROIC = g ÷ reinvestment rate. Cumulative = (NOPAT 2035E − 2026E) / Σ net investment 2026E–34E.')
    ws['C53'] = f'={M("lev_nd", "CO")}'
    ws['C54'] = f'={M("val_nci", "CO")}'
    ws['C56'] = f'={M("sh_dil", "CO")}'
    ws['C58'] = f'={M("val_px", "CO")}'
    ws['B74'] = ('To solve for the market-implied path: Data ▸ What-If Analysis ▸ Goal Seek — set C57 (implied price) to the current price (C58) by changing '
                 'C32 (terminal growth) or the Model scenario levers (aggregates volume / price / unit cash cost, downstream growth and margins, M&A spend).')
    ws['D80'] = 'Rationale (Vulcan characteristics)'
    ws['C82'] = 0.08
    ws['D82'] = 'Model Adjusted EBIT CAGR 2025–30E: low-single-digit volume growth + mid-single-digit pricing with unit-profit expansion + bolt-on M&A'
    ws['C83'] = 0.30
    ws['D83'] = ('Pricing-led unit-profit growth on an existing quarry base is high-return; volume growth and greenfield / bolt-on reserves need capital '
                 '(capex ~9–10% of revenues, deals at ~13–15x EBITDA). Blended ROIIC ~25–35% vs book ROIC ~10%.')
    ws['C84'] = 15
    ws['D84'] = ('Largest US aggregates producer; local-monopoly quarries (permitting barriers, freight-limited markets) and ~70 years of reserves '
                 'support a long competitive-advantage period')
    ws['B95'] = '2028E Adjusted EBIT (Model)'
    ws['B96'] = '(−) 2028E adjusting items'
    ws['B98'] = '2028E NOPAT  =  (Adjusted EBIT − adjusting items) × (1 − t)'
    ws['C95'] = f'={M("o_ebit", "CQ")}'
    ws['C96'] = f'=-{M("o_adj", "CQ")}'
    ws['C101'] = f'={M("lev_nd", "CP")}+{M("val_nci", "CP")}'
    ws['C103'] = f'={M("sh_dil", "CQ")}'
    ws['B113'] = "VMC's market-implied EV/NOPAT (on 2028E NOPAT)"

CH_MAP = {8: ('Organic %', 'gm_org', 'IFERROR({}+0,0)'), 9: ('M&A %', 'gm_acq', 'IFERROR({}+0,0)'),
          12: ('Adjusted EBITDA', 'r_cur', '{}'), 13: ('Operating Earnings (GAAP)', 'is_op', '{}'),
          14: ('Adj. EBITDA Margin %', 'gm_ebitda_m', '{}'), 15: ('GAAP Op. Margin %', 'gm_op_m', '{}'),
          18: ('ROIC', 'ra_roic', 'IFERROR({}+0,NA())'), 19: ('RONTA', 'rn_ronta', 'IFERROR({}+0,NA())'),
          20: ('Adj. EBITDA margin %', 'gm_ebitda_m', 'IFERROR({}+0,NA())'), 21: ('FCF margin %', 'fcf_m', 'IFERROR({}+0,NA())')}
CHART_TITLES = {
    'Net Sales Growth Decomposition: Organic vs M&amp;A': 'Revenue Growth Decomposition: Organic vs M&amp;A',
    'Non-GAAP &amp; GAAP Operating Income (USDm) with Margins (%)': 'Adjusted EBITDA &amp; GAAP Operating Earnings (USDm) with Margins (%)',
    'Returns Profile: ROIC, RONTA, Non-GAAP Operating Margin &amp; FCF Margin': 'Returns Profile: ROIC, RONTA, Adjusted EBITDA Margin &amp; FCF Margin',
}

def do_charts(wb):
    ws = wb['Charts']
    ws['B2'] = 'Vulcan Materials — Historical Operating Performance (2006–2025)'
    ws['B3'] = ('All historicals · USDm · Linked to Model. M&A % = disclosed incremental revenues from acquisitions (e.g. U.S. Concrete 2021, '
                'Wake Stone / Superior Ready Mix 2024) where Vulcan quantified them. Adjusted EBITDA on the current definition (model memo series) for all years.')
    for r, (lab, key, tmpl) in CH_MAP.items():
        ws[f'B{r}'] = lab
        for i, y in enumerate(range(2006, 2026)):
            c = L(CI('C') + i)
            ws[f'{c}{r}'] = '=' + tmpl.format(f'Model!{ANN[y].c}{R[key]}')

def inject_charts(path):
    """Replace openpyxl-rewritten chart parts with the template's chart XML (keeps the bar+line combo), retitled, caches stripped."""
    import zipfile, re, shutil
    src = zipfile.ZipFile(TEMPLATE)
    tmpl = {n: src.read(n).decode('utf8') for n in src.namelist() if re.match(r'xl/charts/chart\d+\.xml$', n)}
    def clean(x):
        for a, b in CHART_TITLES.items(): x = x.replace(a, b)
        x = re.sub(r'<c:numCache>.*?</c:numCache>', '', x, flags=re.S)
        x = re.sub(r'<c:strCache>.*?</c:strCache>', '', x, flags=re.S)
        return x
    by_title = {}
    for n, x in tmpl.items():
        t = re.findall(r'<a:t>([^<]+)</a:t>', x)[0]
        by_title[t] = clean(x)
    zin = zipfile.ZipFile(path)
    tmp = path + '.tmp'
    zout = zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED)
    for item in zin.infolist():
        data = zin.read(item.filename)
        if re.match(r'xl/charts/chart\d+\.xml$', item.filename):
            s = data.decode('utf8')
            m = re.findall(r'<a:t>([^<]+)</a:t>', s)
            t = m[0] if m else None
            if t in by_title: data = by_title[t].encode('utf8')
            elif t in CHART_TITLES.values():
                inv = {v: k for k, v in CHART_TITLES.items()}
                data = by_title[inv[t]].encode('utf8')
        zout.writestr(item, data)
    zout.close(); zin.close()
    shutil.move(tmp, path)
