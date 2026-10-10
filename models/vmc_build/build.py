"""Build VMC_Model.xlsx from the TDY template."""
import sys, copy
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from fw import *
import fw
from spec import *
import spec
from spec2 import build_spec2
from normalize import normalize

OUT = os.path.join(BASE, 'VMC_Model.xlsx')
normalize()

# ---------------------------------------------------------------- scenario table layout (DB:DF)
SCEN_ROWS = []   # (row, kind, cells...)
def scen_layout():
    r = 9
    SCEN_ROWS.append((r, 'hdr', 'SCENARIO INPUT TABLE — driver assumptions per Bull / Base / Bear', 'Bull', 'Base', 'Bear',
                      'Bull / Bear = Base + Δ (Δ inputs on each block header). Outlook block = high / mid / low end of FY26 ranges.'))
    r = 11
    SCEN_ROWS.append((r, 'out_hdr', 'FY26 outlook — Q2-26 release (29-Jul-26)', 'High', 'Mid', 'Low', None)); r += 1
    SC['out_ebitda'] = r; SCEN_ROWS.append((r, 'out', '  FY26 Adjusted EBITDA ($m)', 2600, 2500, 2400, 'Guided $2.4–2.6bn (Q4-25 release, reaffirmed Q1-26 and Q2-26).', 'm')); r += 1
    SC['out_ni'] = r; SCEN_ROWS.append((r, 'out', '  FY26 net earnings attributable to Vulcan ($m, memo)', 1300, 1215, 1100,
                                        'Guided $1.1–1.3bn (Q4-25); mid = Q2-26 projected Adjusted EBITDA bridge ($1,215m).', 'm')); r += 1
    SC['out_vol'] = r; SCEN_ROWS.append((r, 'out', '  FY26 aggregates shipments growth (memo)', 0.03, 0.02, 0.01, 'Guided +1–3% (226.8m tons in 2025).', 'pct')); r += 1
    SC['out_px'] = r; SCEN_ROWS.append((r, 'out', '  FY26 freight-adjusted price growth (memo)', 0.06, 0.05, 0.04, 'Guided +4–6% ($21.98/ton in 2025).', 'pct')); r += 1
    r += 1
    SCEN_ROWS.append((r, 'pt_hdr', 'FY26 point estimates (all cases)', None, None, None,
                      'Q2-26 projected Adjusted EBITDA bridge (DDA&A $700m, interest $215m, tax $340m) and Q4-25 guidance (SAG $580–590m, capex $750–800m, ETR 22–23%).')); r += 1
    for key, lab, val, nf, note in [
        ('pt_dda', '  DDA&A ($m)', 700, 'm', 'Projected Adjusted EBITDA bridge (Q2-26 release).'),
        ('pt_int', '  Interest expense, net ($m)', 215, 'm', 'Projected bridge Q2-26 ($215m; Q4-25 guide ~$225m).'),
        ('pt_etr', '  Tax rate (GAAP & adjusted)', 0.22, 'pct', 'Guided 22–23%; bridge implies ~21.8%.'),
        ('pt_sag', '  SAG expense ($m)', 585, 'm', 'Guided $580–590m.'),
        ('pt_capex', '  Capital expenditures ($m)', 775, 'm', 'Guided $750–800m.'),
        ('pt_cost_g', '  Aggregates unit cash cost growth, FY26 (%)', 0.03, 'pct', 'Guided "low-single-digit increase" in freight-adjusted unit cash cost ($10.65 in 2025); 3% analyst point.'),
        ('pt_down_g', '  Asphalt & concrete underlying revenue growth, 2H/26 (%)', 0.0, 'pct',
         'Analyst estimate; divested Houston asphalt / California ready-mix removed via the divested-share inputs (Model rows).'),
        ('pt_ddat_g', '  DDA&A per ton / corporate DDA&A growth, 2027E+ (%)', 0.03, 'pct', 'Analyst estimate (capex ~1.1x DDA&A).'),
        ('pt_dps_g', '  Dividend per share growth, 2027E+ (%)', 0.07, 'pct', 'Analyst estimate (2026 quarterly dividend raised to $0.52).'),
    ]:
        SC[key] = r; SCEN_ROWS.append((r, 'pt', lab, None, val, None, note, nf)); r += 1
    r += 1
    def block(name, lab, d_bull, d_bear, base, note, nf='pct', clamp=False, years=(2027, 2028, 2029, 2030)):
        nonlocal r
        SCEN_ROWS.append((r, 'blk', lab, d_bull, 'Δ →', d_bear, note, 'dpct' if nf == 'pct' else 'dm'))
        hdr = r; r += 1
        SC[name] = {}
        for y, b in zip(years, base):
            SC[name][y] = r
            SCEN_ROWS.append((r, 'yr', f'  {y}E' if isinstance(y, int) else f'  {y}', hdr, b, clamp, None, nf)); r += 1
        r += 1
    block('AGG_vol', 'AGG_volumeGrowth (aggregates shipments y/y)', 0.01, -0.025, [0.02, 0.025, 0.02, 0.02],
          'IIJA-funded public work, data-centre / manufacturing megaprojects vs soft single-family housing; long-run volumes ~1–2% p.a.')
    block('AGG_price', 'AGG_priceGrowth (freight-adjusted price y/y)', 0.01, -0.02, [0.05, 0.05, 0.045, 0.045],
          'Vulcan has compounded freight-adjusted price ~7% p.a. 2019–25; FY26 guide +4–6%.')
    block('AGG_cost', 'AGG_unitCashCostGrowth (freight-adjusted unit cash cost y/y)', -0.01, 0.015, [0.035, 0.035, 0.03, 0.03],
          'FY26 guide: low-single-digit increase (diesel, parts, labour); operating-efficiency programme.')
    block('ASPH_g', 'ASPH_growth (asphalt segment revenues y/y)', 0.02, -0.03, [0.03, 0.03, 0.03, 0.03],
          'Arizona / California / Texas / Tennessee paving; liquid-asphalt pass-through.')
    block('ASPH_m', 'ASPH_cashMargin (asphalt cash gross profit % of revenues)', 0.01, -0.02, [0.19, 0.19, 0.195, 0.195],
          'FY25 ~18–19%; FY26 guide implies ~$245m asphalt cash gross profit.')
    block('CONC_g', 'CONC_growth (concrete segment revenues y/y)', 0.02, -0.03, [0.02, 0.02, 0.02, 0.02],
          'Remaining East-coast ready-mix after the California sale.')
    block('CONC_m', 'CONC_cashMargin (concrete cash gross profit % of revenues)', 0.01, -0.02, [0.11, 0.11, 0.115, 0.115],
          'FY26 guide implies ~$44m concrete cash gross profit excl. California.')
    block('MNA', 'M&A_spend (acquisition spend, $m)', 750, -500, [500, 750, 750, 750],
          'Aggregates-led bolt-ons; Vulcan deployed ~$2.3bn on Wake Stone / Superior Ready Mix etc. in 2024 and ~$0.1bn in 1H/26.', nf='m', clamp=True)
    SCEN_ROWS.append((r, 'blk', 'Buybacks 2H/26E ($m)', 250, 'Δ →', -250, 'Repurchases $400m in 1H/26 ($217m returned in Q1 incl. dividends); 3.9m shares left on the 2017 authorization.', 'dm'))
    hdr = r; r += 1
    SC['BB26'] = r; SCEN_ROWS.append((r, 'yr', '  2H/26E', hdr, 250, True, None, 'm')); r += 2
    SC['LEV'] = r
    SCEN_ROWS.append((r, 'lev', 'Min. net debt / Adjusted EBITDA (x) — buyback plug', 2.0, 1.5, 1.25,
                      'Surplus cash above this leverage floor goes to buybacks (Model leverage-buyback schedule). Vulcan targets total debt / Adj. EBITDA of 2.0–2.5x.'))
    return r

# ---------------------------------------------------------------- helpers
def comment(cell, text, w=420, h=160):
    c = Comment(text, 'VMC model'); c.width, c.height = w, h
    cell.comment = c

def write_scen(ws):
    T = TM
    for spec_ in SCEN_ROWS:
        r, kind = spec_[0], spec_[1]
        if kind == 'hdr':
            src = 9
        elif kind in ('out_hdr', 'pt_hdr'):
            src = 11 if kind == 'out_hdr' else 15
        elif kind in ('out',):
            src = 12
        elif kind == 'pt':
            src = 16
        elif kind == 'blk':
            src = 20
        elif kind in ('yr',):
            src = 21
        else:
            src = 81
        for ci in range(CI('DA'), CI('DF') + 1):
            s = T.cell(src, ci); d = ws.cell(r, ci)
            if s.has_style: d._style = copy.copy(s._style)
        if kind in ('hdr', 'out_hdr', 'pt_hdr'):
            _, _, a, b, c_, d_, e = spec_
            for col, v in zip(['DB', 'DC', 'DD', 'DE', 'DF'], [a, b, c_, d_, e]):
                if v is not None: ws[f'{col}{r}'] = v
        elif kind == 'out':
            _, _, lab, hi, mid, lo, note, nf = spec_
            ws[f'DB{r}'] = lab
            for col, v in zip(['DC', 'DD', 'DE'], [hi, mid, lo]):
                ws[f'{col}{r}'] = v; ws[f'{col}{r}'].number_format = NF[nf]
            ws[f'DF{r}'] = note
        elif kind == 'pt':
            _, _, lab, _, val, _, note, nf = spec_
            ws[f'DB{r}'] = lab; ws[f'DD{r}'] = val; ws[f'DD{r}'].number_format = NF[nf]; ws[f'DF{r}'] = note
        elif kind == 'blk':
            _, _, lab, db, mid, de, note, nf = spec_
            ws[f'DB{r}'] = lab; ws[f'DC{r}'] = db; ws[f'DD{r}'] = mid; ws[f'DE{r}'] = de; ws[f'DF{r}'] = note
            if nf == 'dm':
                for col in ('DC', 'DE'): ws[f'{col}{r}'].number_format = '\\+#,##0;\\-#,##0'
        elif kind == 'yr':
            _, _, lab, hdr, base, clamp, _, nf = spec_
            ws[f'DB{r}'] = lab; ws[f'DD{r}'] = base
            if clamp:
                ws[f'DC{r}'] = f'=MAX(0,DD{r}+$DC${hdr})'; ws[f'DE{r}'] = f'=MAX(0,DD{r}+$DE${hdr})'
            else:
                ws[f'DC{r}'] = f'=DD{r}+$DC${hdr}'; ws[f'DE{r}'] = f'=DD{r}+$DE${hdr}'
            for col in ('DC', 'DD', 'DE'): ws[f'{col}{r}'].number_format = NF[nf]
        elif kind == 'lev':
            _, _, lab, a, b, c_, note = spec_
            ws[f'DB{r}'] = lab
            for col, v in zip(['DC', 'DD', 'DE'], [a, b, c_]):
                ws[f'{col}{r}'] = v; ws[f'{col}{r}'].number_format = NF['x']
            ws[f'DF{r}'] = note

def is_formula(v): return isinstance(v, str) and v.startswith('=')

def hist_value(row, col):
    kw = row.kw
    path = kw.get('data')
    if not path: return None
    if kw.get('cfbs'):
        if col.kind == 'A': return num(get(str(col.year), path))
        if col.kind == 'H' and col.year == 2026: return num(get('Q2-26', path))
        return None
    if kw.get('annual_only'):
        return num(get(str(col.year), path)) if col.kind == 'A' else None
    if col.kind in ('A', 'Q'):
        return num(get(pkey(col), path))
    return None

def h1_formula(row, col):
    kw = row.kw
    q1, q2 = QC[(col.year, 1)].c, QC[(col.year, 2)].c
    v1, v2 = hist_value(row, QC[(col.year, 1)]), hist_value(row, QC[(col.year, 2)])
    if kw.get('h1_partial') and kw.get('flow') and (v1 is not None or v2 is not None):
        return f'=N({q1}{row.r})+N({q2}{row.r})'
    if v1 is None or v2 is None: return None
    if kw.get('flow'): return f'={q1}{row.r}+{q2}{row.r}'
    if kw.get('avg'): return f'=AVERAGE({q1}{row.r},{q2}{row.r})'
    return None

def compute(row, col):
    """Return (value, is_input)."""
    kw = row.kw; k = col.kind
    CTX.col = col; CTX.key_base = kw.get('growth')
    fcol = k in ('QE', 'E26', 'AE')
    if not fcol:
        v = hist_value(row, col)
        if v is not None: return v, True
        if k == 'H' and kw.get('data') and not kw.get('cfbs') and not kw.get('annual_only'):
            f = h1_formula(row, col)
            if f: return f, False
            if ('h1' in kw and hist_value(row, QC[(col.year, 1)]) is not None
                    and hist_value(row, QC[(col.year, 2)]) is not None):
                return kw['h1'](), False
            if 'hist' not in kw and 'all' not in kw: return None, False
        if kw.get('cfbs_f'):
            if k == 'A' or (k == 'H' and col.year == 2026):
                if kw.get('ann_only_f') and k == 'H': return None, False
                return kw['cfbs_f'](), False
            return None, False
        if kw.get('cfbs') or kw.get('annual_only'): return None, False
        if kw.get('hist_only') and False: pass
        for key in ('hist', 'all'):
            if key in kw:
                if kw.get('need_prior') and col.prior is None: return None, False
                f = kw[key]()
                return f, False
        if kw.get('growth'):
            return (growth_formula(), False) if col.prior else (None, False)
        if kw.get('prior_ratio') and col.prior:
            n, d = kw['prior_ratio']
            return f'=IFERROR(({V(n)}-{P(n)})/({V(d)}-{P(d)}),"")', False
        return None, False
    # ---------------- forecast columns
    if k == 'QE':
        if 'qe' in kw: return kw['qe'](), False
        if 'qe_input' in kw: return kw['qe_input'][col.q], True
        if 'fc' in kw: return kw['fc'](), False
    if k == 'E26':
        if 'e26' in kw: return kw['e26'](), False
        if 'e26_input' in kw: return kw['e26_input'], True
        if 'fc_ann' in kw: return kw['fc_ann'](), False
        if (('qe' in kw or 'qe_input' in kw or 'fc' in kw) and not kw.get('cfbs') and not kw.get('avg')
                and (kw.get('flow') or row.nf in ('m', 'm1'))):
            return f'=SUM({V(row.key,"CJ")},{V(row.key,"CK")},{V(row.key,"CM")},{V(row.key,"CN")})', False
        if 'fc' in kw: return kw['fc'](), False
    if k == 'AE':
        if 'ae' in kw: return kw['ae'](), False
        if 'ae_vals' in kw: return kw['ae_vals'][col.year], True
        if 'ae_input' in kw:
            return kw['ae_input'], True
        if 'fc_ann' in kw: return kw['fc_ann'](), False
        if 'fc' in kw: return kw['fc'](), False
    if k in ('E26', 'AE') and kw.get('cfbs_f'):
        return kw['cfbs_f'](), False
    if kw.get('hist_only'): return None, False
    if 'all' in kw:
        if kw.get('need_prior') and col.prior is None: return None, False
        return kw['all'](), False
    if kw.get('growth'):
        g = kw['growth']
        c = col
        if c.prior:
            return f'=IF(AND(ISNUMBER({V(g)}),ISNUMBER({P(g)})),IFERROR({V(g)}/{P(g)}-1,""),"")', False
    if kw.get('prior_ratio') and col.prior:
        n, d = kw['prior_ratio']
        return f'=IFERROR(({V(n)}-{P(n)})/({V(d)}-{P(d)}),"")', False
    return None, False

def h1_formula_fallback_ok(row, col):
    return True

# ---------------------------------------------------------------- main
def main():
    import openpyxl
    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb['Model']
    # clear Model
    for rowc in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for c in rowc:
            c.value = None; c.comment = None
            c._style = copy.copy(ws.cell(1, 1)._style)
    for k in list(ws.row_dimensions.keys()):
        rd = ws.row_dimensions[k]; rd.outlineLevel = 0; rd.hidden = False
    scen_layout()
    build_spec(); build_spec2()
    number_rows(8)
    write_header(ws)
    for row in SPEC:
        write_row(ws, row)
    write_scen(ws)
    ws.freeze_panes = 'C8'
    import other_sheets as osh
    osh.do_bbb(wb); osh.do_dcf(wb); osh.do_charts(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    osh.inject_charts(OUT)
    print('rows', SPEC[-1].r, 'saved', OUT)
    return wb

RECALC = os.environ.get('XLSX_RECALC', 'recalc.py')  # LibreOffice recalculation script (needs libreoffice-calc)
def recalc(path):
    import subprocess, json as _j
    out = subprocess.run(['python3', RECALC, path, '300'], capture_output=True, text=True, cwd=os.path.dirname(RECALC))
    return _j.loads(out.stdout)

def scenario_snapshots():
    import openpyxl, shutil
    import other_sheets as osh
    snaps = {}
    for scen in ('Bull', 'Bear'):
        p = os.path.join(BASE, f'VMC_{scen}.xlsx')
        wb = openpyxl.load_workbook(OUT); wb['Model']['CT2'] = scen; wb.save(p)
        res = recalc(p); print(scen, res.get('status'), res.get('total_errors'))
        v = openpyxl.load_workbook(p, data_only=True)['Bull-Base-Bear']
        snaps[scen] = {(r, c): v[f'{c}{r}'].value for r in range(9, 48) for c in 'CDEFGHIJKL' if v[f'{c}{r}'].value is not None}
    wb = openpyxl.load_workbook(OUT)
    osh.snapshot_values(wb, snaps['Bull'], snaps['Bear'])
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    osh.inject_charts(OUT)

def write_header(ws):
    for r in range(1, 8):
        clone_row_style(ws, r, r, range(2, CI('DF') + 1))
    ws['CT1'] = 'SCENARIO SWITCH ▼ (Bull / Base / Bear) — drives all scenario-linked forecast drivers'
    ws['CT2'] = 'Base'
    comment(ws['CT2'], 'Scenario toggle. Choose Bull / Base / Bear. Drives: FY26 Adjusted EBITDA outlook end-point (high / mid / low), '
            '2027E–2030E aggregates volume / price / unit cash cost, asphalt and concrete growth and margins, M&A spend, 2H/26 buybacks '
            'and the leverage floor for 2027E+ buybacks (scenario table DB:DF).')
    ws['B2'] = 'Vulcan Materials Company (NYSE: VMC) — 3-Statement Financial Model'
    ws['B3'] = ('US GAAP as reported · USD millions (per-share data in USD; volumes in millions of tons / cubic yards) · Source: Vulcan Forms 10-K / 10-Q '
                'and Form 8-K earnings releases (Ex. 99.1), SEC EDGAR CIK 0001396009 (legacy Vulcan CIK 0000103973 before Nov-2007; '
                'investor-relations releases are the same documents as filed on EDGAR)')
    ws['B4'] = ('Basis notes: calendar fiscal year. Florida Rock acquired Nov-2007 (Cement segment to Mar-2014, sold to Cementos Argos with Florida '
                'ready-mix). Pre-2018 total revenues = net sales + delivery revenues; ASC 606 (2018) moved freight into segment revenues. ASU 2017-07 '
                '(2018) non-service pension to nonoperating; ASU 2016-18 restricted cash in the cash-flow statement from FY2017. U.S. Concrete from Aug-2021; '
                'Wake Stone / Superior Ready Mix 2024; Houston asphalt sold Q4/25; California ready-mix sold Jun-26. All periods as originally reported. '
                'Q3/26E–Q4/26E = FY26 Adjusted EBITDA outlook less 1H/26 actuals.')
    for c in ('CV',): pass
    ws['CV5'] = 'Forecast'; ws['CW5'] = 'Historical'; ws['CX5'] = 'Historical'; ws['CY5'] = 'Historical'
    ws['B6'] = '(USDm)'
    for col in COLS:
        ws[f'{col.c}6'] = col.label
    ws['CT6'] = 'Modelling Notes'
    for c, a, b in [('CV', '5Y CAGR', "25-'30"), ('CW', '5Y CAGR', "20-'25"), ('CX', '10Y CAGR', "15-'25"), ('CY', '19Y CAGR', "06-'25")]:
        ws[f'{c}6'] = a; ws[f'{c}7'] = b

CAGR_ROWS_GROWTH = {'agg_tons', 'agg_price', 'agg_fa_rev', 'agg_rev', 'agg_cgp', 'agg_cgp_t', 'agg_cost', 'asph_rev', 'asph_cgp', 'conc_rev',
                    'conc_cgp', 'g_rev', 'g_cgp', 'g_ebitda', 'g_ebit', 'g_op', 'is_rev', 'is_gp', 'is_op', 'is_ni', 'eps_dil', 'r_adj', 'r_cur',
                    'o_ebit', 'e_adjni', 'e_eps', 'e_cur', 'e_cur_eps', 'cfo', 'fcf', 'dps', 'sag', 'is_dda', 'bs_ta', 'bs_veq', 'sh_dil', 'fcf_ps'}
CAGR_ROWS_AVG = {'agg_cgp_m', 'agg_gp_m', 'asph_cgp_m', 'conc_cgp_m', 'g_ebitda_m', 'g_cgp_m', 'r_adj_m', 'r_cur_m', 'o_ebit_m', 'gm_gm',
                 'gm_ebitda_m', 'gm_ebit_m', 'gm_op_m', 'gm_net_m', 'gm_sag', 'fcf_m', 'capex_pct', 'ra_roic', 'rn_ronta', 'roe', 'is_etr',
                 'sag_pct', 'drv_cogs', 'drv_dda', 'payout', 'gm_rev', 'gm_org', 'gm_vol', 'gm_price'}

def write_row(ws, row):
    r = row.r
    arch = ARCH[row.arch]
    clone_row_style(ws, arch, r)
    lab = ws.cell(r, 2); lab.value = row.label or None
    if row.kw.get('note'):
        ws[f'CT{r}'] = row.kw['note']
    if row.arch in ('section', 'block', 'blank', 'sub', 'text') and row.key not in ('r_def', 'e_def'):
        if row.kw.get('group'):
            ws.row_dimensions[r].outlineLevel = 1; ws.row_dimensions[r].hidden = True
        return
    if row.kw.get('group'):
        ws.row_dimensions[r].outlineLevel = 1; ws.row_dimensions[r].hidden = True
    if row.key in ('r_def', 'e_def'):
        write_defs(ws, row); return
    nf = NF[row.nf]
    for col in COLS:
        v, inp = compute(row, col)
        if v is None: continue
        cell = ws[f'{col.c}{r}']
        cell.value = v
        cell.number_format = nf
        tc = TM[f'{col.c}{arch}']
        tcol = tc.font.color.rgb if (tc.font.color is not None and isinstance(tc.font.color.rgb, str)) else None
        if inp:
            set_color(cell, BLUE, bold=True if col.kind in ('QE', 'E26', 'AE') else None)
        elif tcol == GRAY and col.kind in ('A', 'Q', 'H'):
            pass
        else:
            set_color(cell, BLACK)
        add_cell_comment(ws, row, col, cell)
    # CAGR / averages
    if row.key in CAGR_ROWS_GROWTH:
        ws[f'CV{r}'] = f'=IFERROR(({V(row.key,"CS")}/{V(row.key,"CI")})^(1/5)-1,"n/m")'
        ws[f'CW{r}'] = f'=IFERROR(({V(row.key,"CI")}/{V(row.key,"BE")})^(1/5)-1,"n/m")'
        ws[f'CX{r}'] = f'=IFERROR(({V(row.key,"CI")}/{V(row.key,"AA")})^(1/10)-1,"n/m")'
        ws[f'CY{r}'] = f'=IFERROR(({V(row.key,"CI")}/{V(row.key,"C")})^(1/19)-1,"n/m")'
        for c in ('CV', 'CW', 'CX', 'CY'): ws[f'{c}{r}'].number_format = '0%;\\(0%\\);"n/m"'
    elif row.key in CAGR_ROWS_AVG:
        fc = ','.join(f'{c}{r}' for c in ['CO', 'CP', 'CQ', 'CR', 'CS'])
        h5 = ','.join(f'{ANN[y].c}{r}' for y in range(2021, 2026))
        h10 = ','.join(f'{ANN[y].c}{r}' for y in range(2016, 2026))
        h20 = ','.join(f'{ANN[y].c}{r}' for y in range(2006, 2026))
        ws[f'CV{r}'] = f'=IFERROR(AVERAGE({fc}),"n/m")'; ws[f'CW{r}'] = f'=IFERROR(AVERAGE({h5}),"n/m")'
        ws[f'CX{r}'] = f'=IFERROR(AVERAGE({h10}),"n/m")'; ws[f'CY{r}'] = f'=IFERROR(AVERAGE({h20}),"n/m")'
        for c in ('CV', 'CW', 'CX', 'CY'): ws[f'{c}{r}'].number_format = NF['pct']
    add_label_comment(ws, row)

SRC_RELEASE = 'Source: Vulcan earnings releases (Form 8-K Ex. 99.1) — '
LABEL_SOURCES = {
    'agg_tons': SRC_RELEASE + 'Table D "Average Unit Sales Price and Unit Shipments"; FY2006–12 from the 10-K MD&A. Quarterly from Q1/13.',
    'agg_price': SRC_RELEASE + 'freight-adjusted sales price per ton (Table D / Appendix 1). 2007 price covers legacy Vulcan only (excl. Florida Rock).',
    'agg_fa_rev': SRC_RELEASE + 'Appendix 1 "Aggregates Segment Freight-Adjusted Revenues" (published from 2014).',
    'agg_rev': SRC_RELEASE + 'Table D "Segment Financial Data" total revenues; 10-K segment note for FY2006–12.',
    'agg_gp': SRC_RELEASE + 'Table D segment gross profit; 10-K segment note.',
    'agg_dda': SRC_RELEASE + 'Table D "Depreciation, Depletion, Accretion and Amortization" by segment; 10-K segment note.',
    'agg_cgp_pub': SRC_RELEASE + 'Appendix "Cash Gross Profit" (aggregates cash gross profit).',
    'asph_rev': SRC_RELEASE + 'Table D. 2007–09: asphalt mix and concrete reported as one segment (shown in the legacy "other segment" rows).',
    'conc_rev': SRC_RELEASE + 'Table D. 2007–09: see legacy rows (combined Asphalt mix & Concrete segment).',
    'cem_rev': '10-K segment note / releases: Cement segment (Florida Rock) 2007–Q1/14; Q2/14–Q3/14 releases label calcium products "Cement".',
    'calc_rev': 'Releases: Calcium segment (Brooksville, FL) from Q2/14.',
    'oseg_rev': '2006: "other" product line (single Construction Materials segment). 2007–09: combined Asphalt mix & Concrete segment (FY2008 / FY2009 10-K recast).',
    'interseg': SRC_RELEASE + 'Table D intersegment sales (total eliminations).',
    'deliv': 'Income statement "Delivery revenues" (pre-Q3/14 presentation); from Q3/14 to 2017 the releases\' "freight and delivery revenues" reconciliation (shown on the income statement memo line only).',
    'is_rev': 'Source: Consolidated Statements of Earnings (earnings release Table A; 10-K for annual).',
    'is_othop': 'Combination of all remaining operating lines (e.g. "Other operating expense, net", "Recovery from legal settlement", "Gain on sale of real estate", "Loss on impairments", "Restructuring charges"). Cell comments give each period\'s lines.',
    'r_tax': 'Source: EBITDA reconciliation (Appendix) — "Income tax expense, including discontinued operations".',
    'r_discat': '2016–Q1/23 (and earlier) Vulcan began its EBITDA build from net earnings and removed discontinued operations after tax before EBIT; from Q2/23 the pre-tax discontinued-ops (gain) loss sits in the Adjusted EBITDA bridge instead.',
    'r_adj_pub': SRC_RELEASE + 'Appendix "EBITDA and Adjusted EBITDA". Introduced FY2011; FY2011 as originally published.',
    'e_pub': SRC_RELEASE + 'Appendix "Adjusted Diluted EPS". 2006–08: "EPS from continuing operations, as adjusted"; 2014–15 from the following year\'s release comparatives where the release gave narrative only.',
    'e_items_ps': 'Company per-share bridge. Where the company bridge starts from EPS from continuing operations, the discontinued-operations per-share amount is included here so the model bridge starts from diluted EPS on net earnings.',
    'cf_ne': 'Source: Consolidated Statements of Cash Flows (10-K; 1H/26 from the Q2-26 release Table C).',
    'bs_cash': 'Source: Consolidated Balance Sheets (10-K; 30-Jun-26 from the Q2-26 release Table B).',
    'dps': 'Source: releases / 10-Q / 10-K (XBRL); Q4 = full year less nine months where the Q4 release omits it.',
    'reserves': 'Source: 10-K Item 2 (Properties) — total aggregates reserves.',
    'sch_roic_pub': 'Source: release appendix "Return on Invested Capital" (published from FY2019).',
}
def add_label_comment(ws, row):
    t = LABEL_SOURCES.get(row.key)
    if t: comment(ws.cell(row.r, 2), t)

def add_cell_comment(ws, row, col, cell):
    if col.kind not in ('A', 'Q'): return
    k = pkey(col)
    d = DATA.get(k)
    if not d: return
    if row.key == 'is_othop':
        lines = (d.get('is') or {}).get('other_op_lines')
        if lines and len(lines) > 1:
            comment(cell, 'Company wording: ' + ' | '.join(f'{a}: {b}' for a, b in lines if b not in (None, 0, 0.0)))
    elif row.kw.get('adj'):
        lines = (d.get('ngaap') or {}).get('ebitda_bridge_lines')
        if lines and num(cell.value) not in (None, 0):
            comment(cell, 'Company wording (Adjusted EBITDA bridge): ' + ' | '.join(f'{a}: {b}' for a, b in lines if num(b) not in (None, 0)))
    elif row.key == 'e_items_ps':
        lines = (d.get('ngaap') or {}).get('eps_bridge_lines')
        if lines:
            comment(cell, 'Company wording (Adjusted EPS bridge, $/share): ' + ' | '.join(f'{a}: {b}' for a, b in lines))

DEFS_R = {
    'C': ('Not published (2006): no EBITDA measure', None),
    'D': ('EBITDA (continuing operations) 2007–2010; no adjusted measure', 'Vulcan EBITDA excluded discontinued operations (Chemicals). FY2008 release EBITDA preceded the $252.7m goodwill impairment booked in the 10-K (shown in the impairment row).'),
    'H': ('Adjusted EBITDA introduced (FY2011): EBITDA excl. restructuring & exchange-offer costs', 'FY2011 as originally published ($440.4m). The FY2012 release restated FY2011 to $351.8m by also excluding the $46.4m legal recovery and $42.1m gains on sale.'),
    'I': ('FY2012+: excludes gains on sale, legal recoveries, restructuring, exchange-offer and business-development costs', None),
    'Z': ('Q4/15: deferred-revenue amortization no longer deducted from Adjusted EBITDA', 'Q1–Q3/15 as originally reported (deducted), so 2015 quarters do not sum to FY2015 ($836.3m).'),
    'BV': ('Q4/22–Q1/23: discontinued operations removed after tax in the EBITDA build', None),
    'BS': ('Q2/23+: tax incl. discontinued ops; pre-tax disc. ops (gain) loss in the Adjusted EBITDA bridge', 'FY2023 10-K restated FY2022 EBITDA to $1,517.9m on this basis (Adjusted EBITDA unchanged at $1,625.6m).'),
    'CJ': ('2026: CEO transition & reorganization charges added to the bridge', None),
}
DEFS_E = {
    'C': ('2006–08: "EPS from continuing operations, as adjusted" (gains on contractual rights / ECU earn-out, real-estate gains, Florida Rock)', None),
    'F': ('2009–2013: no adjusted EPS published', None),
    'P': ('2014+: Adjusted diluted EPS from continuing operations (items in Adjusted EBITDA net of tax + discrete tax items)', '2014 narrative only; figures from the 2015 releases\' comparatives.'),
    'AM': ('2017: tax reform remeasurement and debt-purchase interest excluded', 'Tax Cuts and Jobs Act deferred-tax remeasurement benefit ($1.99/share) and Alabama NOL valuation allowance release ($0.21) treated as discrete tax items; debt-purchase interest $0.73/share after tax.'),
    'BL': ('2022+: NOL carryforward valuation allowance (Mexico) as discrete tax item', None),
}
def write_defs(ws, row):
    defs = DEFS_R if row.key == 'r_def' else DEFS_E
    for c, (txt, cm) in defs.items():
        cell = ws[f'{c}{row.r}']; cell.value = txt
        cell.font = Font(name='Calibri', sz=10, b=True, i=True, color='FFC00000')
        if cm: comment(cell, cm)

if __name__ == '__main__':
    main()
    if '--snap' in sys.argv: scenario_snapshots()
