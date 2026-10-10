"""Build WAB_Model.xlsx from the MLM template (same layout, formatting and process)."""
import sys, copy, os, json
from openpyxl.styles import Font
from fw import *
import fw
import normalize
from spec import build_spec, scen_layout, SCEN_ROWS, SC
from spec2 import build_spec2

OUT = os.path.join(BASE, 'WAB_Model.xlsx')
DEALS = normalize.run()
D = DATA

# ====================================================================== inputs (V0)
def wc_inputs():
    b, i = D['2025']['bs'], D['2025']['is']
    ar = num(b['receivables']) + (num(b.get('unbilled')) or 0)
    return {'dso': round(ar / i['net_sales'] * 365), 'dio': round(b['inventories'] / i['cost_of_sales'] * 365),
            'dpo': round(b['ap'] / i['cost_of_sales'] * 365), 'dep': round(b['customer_deposits'] / i['net_sales'], 3),
            'acc': round(b['accrued_comp'] / i['net_sales'], 3)}
WC = wc_inputs()
WC['note'] = (f'Forecast inputs set at FY2025 levels (DSO {WC["dso"]}d incl. unbilled, inventory {WC["dio"]}d, payables {WC["dpo"]}d, '
              f'customer deposits {WC["dep"]:.1%} and accrued compensation {WC["acc"]:.1%} of sales).')

V0 = {
    'out_hdr': 'FY26 outlook — Q2-26 release (22-Jul-26)',
    'out_note': 'Raised in the Q2-26 release (from $12.19–12.49bn / $10.25–10.65 in Q1-26; $10.05–10.45 in Feb-26). Includes Dellner (closed 10-Feb-26).',
    'outlook': [
        ('out_rev', '  FY26 net sales ($m)', (12600, 12450, 12300), 'm', 'Guided $12.30–12.60bn (raised $110m at mid-point in Q2-26; +11.5% y/y).'),
        ('out_eps', '  FY26 adjusted diluted EPS ($)', (10.90, 10.75, 10.60), 'ps', 'Guided $10.60–10.90 (+19.9% y/y at mid). Drives the Q3/Q4-26E adjusted-margin calibration.'),
        ('bb26', '  2H/26 share repurchases ($m)', (500, 350, 200), 'm', 'Authorization $760m remaining at 30-Jun-26 (1H/26: $457m, 1.80m shares).'),
    ],
    'points': [
        ('pt_int', '  Interest expense, net ($m)', 311, 'm', '1H/26 $151m + Q2/26 run-rate ($80m) × 2 — Dellner financing on the revolver / receivables program; not guided.'),
        ('pt_oth', '  Other income (expense), net ($m)', 21, 'm', '1H/26 actual $21m; 2H/26 assumed nil (Q2/26 −$2m).'),
        ('pt_etr', '  Tax rate (GAAP & adjusted)', 0.245, 'pct', 'FY26 guidance slide: tax rate ~24.5%.'),
        ('pt_capex', '  Capital expenditures ($m)', 250, 'm', 'Guidance: capex ~2% of sales (≈$250m on the mid-point); 1H/26 $108m.'),
        ('pt_restr', '  Restructuring & portfolio optimization costs ($m)', 15, 'm', '1H/26 $5m; Integration 3.0 / Portfolio Optimization; analyst point estimate.'),
    ],
    'lev': (1.75, 1.5, 1.25),
    'lev_note': 'Wabtec targets net leverage of ~2.0–2.5x through the cycle; ~1.7x Adjusted EBITDA expected at YE26 (2.0x at 30-Jun-26 after Dellner). Floors leave headroom for bolt-on M&A; buybacks absorb free cash flow once leverage reaches the scenario floor.',
    'blocks': [
        ('F_SVC', 'FRT_SVC (Freight Services sales y/y)', 0.02, -0.02, [0.04, 0.05, 0.05, 0.05],
         'Modernizations (lumpy; N.A. mods down in 2026), overhauls, parts & long-term service agreements on the ~25,000 GE-built fleet.', 'pct', False),
        ('F_EQ', 'FRT_EQ (Freight Equipment sales y/y)', 0.03, -0.04, [0.06, 0.05, 0.04, 0.04],
         'New locomotives (international orders: Kazakhstan, India, Brazil, Africa) & mining; $25.3bn Freight backlog at 30-Jun-26.', 'pct', False),
        ('F_COMP', 'FRT_COMP (Freight Components sales y/y)', 0.02, -0.03, [0.03, 0.04, 0.04, 0.04],
         'N.A. railcar build (soft 2025–26) and industrial (L&M, cooling); North American freight car cycle recovery from 2027.', 'pct', False),
        ('F_DIG', 'FRT_DIG (Freight Digital Intelligence sales y/y)', 0.03, -0.03, [0.08, 0.08, 0.07, 0.07],
         'PTC / train control, signalling (Frauscher axle counters), inspection (Inspection Technologies), software; organic high-single digit.', 'pct', False),
        ('T_OE', 'TRN_OE (Transit original equipment sales y/y)', 0.025, -0.03, [0.05, 0.05, 0.04, 0.04],
         'Car-builder backlogs (Europe, India); Dellner couplers anniversary in Feb-27.', 'pct', False),
        ('T_AM', 'TRN_AM (Transit aftermarket sales y/y)', 0.02, -0.02, [0.06, 0.05, 0.05, 0.05],
         'Installed-base aftermarket (~60% of Transit); pricing.', 'pct', False),
        ('F_OPM', 'FRT_OPM (Freight adjusted operating margin)', 0.01, -0.015, [0.265, 0.270, 0.275, 0.280],
         'FY25 24.3%; 1H/26 25.9%. Five-year outlook (2025–29): 350+ bps consolidated adjusted margin expansion.', 'pct', False),
        ('T_OPM', 'TRN_OPM (Transit adjusted operating margin)', 0.01, -0.015, [0.180, 0.185, 0.190, 0.195],
         'FY25 14.8%; 1H/26 17.2% (Dellner mix, Integration 3.0, low-margin pruning).', 'pct', False),
        ('MNA', 'M&A_spend (bolt-on acquisitions, $m)', 500, -500, [750, 750, 750, 750],
         'Wabtec deployed ~$2.5bn on acquisitions in 2025 and ~$1.1bn in 1H/26 (Dellner); Base assumes ~$0.75bn p.a. of bolt-ons.', 'm', True),
    ],
    'qe_growth': {'f_svc': {3: -0.04, 4: -0.04}, 'f_eq': {3: 0.12, 4: 0.12}, 'f_comp': {3: 0.0, 4: 0.0},
                  'f_dig': {3: round(0.06 + 60 / 297, 3), 4: round(0.06 + 40 / 361, 3)},
                  't_oe': {3: round(0.05 + 35 / 367, 3), 4: round(0.05 + 35 / 351, 3)},
                  't_am': {3: round(0.06 + 35 / 426, 3), 4: round(0.06 + 35 / 491, 3)}},
    'qe_growth_note': {
        'f_svc': 'Input −4%: 1H/26 −11% y/y (lower N.A. modernization deliveries, as guided), Q2/26 −4%.',
        'f_eq': 'Input +12%: 1H/26 +43% (international locomotive deliveries); tougher comparisons in 2H (Q3-25 $677m, Q4-25 $666m).',
        'f_comp': 'Input 0%: 1H/26 −3% (N.A. railcar build down; industrial up).',
        'f_dig': 'Input = 6% organic + Frauscher (closed 1-Dec-25) not yet in the prior-year quarter (~$60m in Q3, ~$40m in Q4; analyst estimate). Inspection Technologies anniversaries 1-Jul-26.',
        't_oe': 'Input = 5% organic + ~half of Dellner (closed 10-Feb-26; $69m of sales in Q2/26) ≈ $35m per quarter.',
        't_am': 'Input = 6% organic + ~half of Dellner ≈ $35m per quarter.',
    },
    'qe_acq': {3: 130.0, 4: 110.0},
    'qe_acq_note': 'Q3/Q4-26E inputs: Frauscher (~$60m / ~$40m) + Dellner (~$70m per quarter) not in the prior-year quarter (analyst estimate; Q2/26 published $232m incl. Inspection Technologies).',
    'qe_inv': {3: 0.0, 4: 0.0},
    'inv_note': 'Inventory step-ups: Inspection Technologies $24m (Q3-25), Frauscher / Dellner $29m (Q4-25 – 1H/26). Fully expensed by Q2/26 — none forecast.',
    'c_pct': -0.017, 'dep_pct': 0.0165, 'sga_pct': 0.126, 'eng_pct': 0.021, 'sbc_pct': 0.008,
    'amort_sched': {2027: 356.0, 2028: 354.0, 2029: 352.0, 2030: 339.0},
    'amort_note': 'Q2-26 10-Q note: amortization of existing intangibles remainder-2026 $179m (Q3/Q4-26E at the Q2/26 run-rate of $91m), 2027 $356m, 2028 $354m, 2029 $352m, 2030 $339m.',
    'oth_ae': 10.0, 'nci_ae': 5.0, 'etr': 0.245, 'dps_g': 0.12, 'restr_ae': 20.0,
    'def_tax_ae': -50.0, 'ofin_ae': -40.0, 'capex_pct': 0.02, 'wc': WC, 'min_cash': 600.0,
    'mat': {2026: -1250.0, 2027: -568.0, 2028: -1250.0, 2029: 0.0, 2030: -1225.0},
    'mat_note': '2026E: 3.45% notes ($750m, Nov-26) + 364-day term loan ($500m, 27-Nov-26); 2027: €500m 1.25% notes ($568m book); 2028: 4.70% notes $1.25bn; 2030: 4.90% notes $500m + term loan $725m (Apr-30).',
    'refi_pct': {2026: 1.0, 2027: 1.0, 2028: 1.0, 2029: 1.0, 2030: 1.0},
    'debt_note': ('Debt at 30-Jun-26 $6,571m: revolver $638m + receivables program $400m (prepayable facilities) and notes & term loans $5,533m '
                  '(3.45% 2026 $750m, €500m 1.25% 2027, 4.70% 2028 $1.25bn, 4.90% 2030 $500m, 5.611% 2034 $500m, 5.50% 2035 $750m, term loans $725m 2030 / $500m Nov-26). '
                  '$2bn commercial paper program established 1-Oct-26 (nothing issued).'),
    'r_notes': {2027: 0.047, 2028: 0.048, 2029: 0.050, 2030: 0.051},
    'r_notes_note': 'Blended coupon ~4.5% on the 30-Jun-26 notes & term loans; drifts up as the 3.45% 2026, 1.25% EUR 2027 and 4.70% 2028 notes refinance at ~5.0–5.5%.',
    'px': 280.0,
    'px_note': 'VALUATION — share price input $280.00 (latest observable on EDGAR: Form 4 open-market sales 1-Sep-2026 at $278.93–281.74; update to market)',
    'mna_note': 'Scenario lever M&A_spend. Wabtec deployed ~$2.5bn (2025: Inspection Technologies, Frauscher) and ~$1.1bn (1H/26: Dellner); Base assumes ~$0.75bn p.a. of bolt-ons from 2027.',
    'rev_guid_note': 'Company FY26 sales guidance $12.30–12.60bn (raised in the Q2-26 release, 22-Jul-2026; includes Dellner). Drives the revenue-outlook calibration.',
    'eps_guid_note': 'Company FY26 adjusted EPS guidance $10.60–10.90 (raised in the Q2-26 release). Drives the adjusted-margin calibration.',
    'legacy_pl': [('spe', 'Specialty Products & Electronics (Freight Electronics & Specialty Products to 2010)'), ('brake', 'Brake Products'),
                  ('reman', 'Remanufacturing, Overhaul & Build'), ('transit', 'Transit Products (Other Transit Products from 2008)'), ('other', 'Other')],
}

# ====================================================================== PPA by closing quarter
PPA = {}
for d in DEALS or []:
    p = d.get('ppa_latest')
    cd = str(d.get('closing_date') or '')
    if not p or len(cd) < 7 or not cd[:4].isdigit(): continue
    y, m = int(cd[:4]), int(cd[5:7]); q = (m - 1) // 3 + 1
    key = f'Q{q}-{str(y)[2:]}'
    PPA[key] = {'cons': d.get('cash_consideration'), 'cashacq': p.get('cash_acquired'), 'ar': p.get('receivables'), 'inv': p.get('inventories'),
                'oca': p.get('other_current_assets'), 'ppe': p.get('ppe'), 'intang': p.get('intangibles_total'), 'gw': p.get('goodwill'),
                'onca': p.get('other_noncurrent_assets'), 'cl': p.get('current_liabilities'), 'ncl': p.get('noncurrent_liabilities'),
                'dtl': p.get('deferred_taxes'), 'nci': None,
                'names': f'{d["name"].split("(")[0].strip()} — closed {cd}; {d.get("segment", "")}'}

# ====================================================================== compute engine
SPEC2_START = None
def hist_value(row, col):
    kw = row.kw; path = kw.get('data')
    if not path: return None
    if kw.get('annual_only') and col.kind != 'A': return None
    if kw.get('ppa'):
        v = PPA.get(pkey(col), {}).get(path.split('.')[1])
        return v if kw.get('textdata') else num(v)
    v = get(pkey(col), path)
    return v if kw.get('textdata') else num(v)

def first(kw, keys):
    for k in keys:
        if k in kw: return k
    return None

def compute(row, col):
    kw = row.kw; k = col.kind
    CTX.col = col
    if kw.get('ann') and col.q: return None, False
    if k in HISTK:
        v = hist_value(row, col)
        if v is not None: return v, True
        for key in ('hist', 'all'):
            if key in kw:
                f = kw[key]()
                return f, False
        return None, False
    if k == 'QE':
        if row.spec2: return None, False
        if 'qe_in' in kw:
            v = kw['qe_in']; return (v[col.q] if isinstance(v, dict) else v), True
        key = first(kw, ('qe', 'fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    if k == 'E26':
        if 'e26_in' in kw: return kw['e26_in'], True
        if 'e26' in kw: return kw['e26'](), False
        if kw.get('avg') and ('qe' in kw or 'qe_in' in kw): return f'=AVERAGE(BW{row.r},BX{row.r},BY{row.r},BZ{row.r})', False
        if kw.get('flow') and not row.spec2 and any(x in kw for x in ('qe', 'qe_in', 'fc')): return S26(row.key), False
        key = first(kw, ('fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    if k == 'AE':
        if 'ae_in' in kw:
            v = kw['ae_in']; return (v[col.year] if isinstance(v, dict) else v), True
        key = first(kw, ('ae', 'fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    return None, False

# ====================================================================== writers
CAGR_G = {'f_rev', 't_rev', 'f_svc', 'f_eq', 'f_comp', 'f_dig', 't_oe', 't_am', 'f_adjop', 't_adjop', 'g_rev', 'g_adjop', 'b_tot',
          'is_rev', 'is_gp', 'is_sga', 'is_eng', 'is_amort', 'is_op', 'is_ni', 'is_dda', 'eps_dil', 'dps', 'r_ebitda', 'r_adj', 'r_ebit',
          'o_adj', 'e_adjni', 'e_eps', 'cf_cfo', 'fcf', 'fcf_ps', 'bs_ta', 'bs_weq', 'sh_dil'}
CAGR_A = {'f_adjm', 't_adjm', 'f_gm', 't_gm', 'g_adjm', 'is_gm', 'is_opm', 'is_etr', 'r_adj_m', 'r_ebit_m', 'o_adj_m', 'gm_rev', 'gm_org',
          'gm_frt', 'gm_trn', 'gm_gm', 'gm_adjm', 'gm_ebitda_m', 'gm_op_m', 'gm_net_m', 'gm_sga', 'gm_eng', 'fcf_m', 'fcf_conv', 'ocf_conv',
          'capex_pct', 'ra_roic', 'rn_ronta', 'roe', 'payout', 'd_dep_pct', 'd_sga', 'd_eng', 'd_sbc', 'a_org_g', 'a_acq_g', 'lv_x', 'wc_nwc_p'}
ANN_HIST = [ANN[y].c for y in range(2006, 2026)]

def style_cagr(ws, r, on):
    src = 16 if on else 12
    for c in ('CH', 'CI', 'CJ', 'CK'):
        ws[f'{c}{r}']._style = copy.copy(TM[f'{c}{src}']._style)

def write_row(ws, row):
    r = row.r; arch = ARCH[row.arch]
    clone_row_style(ws, arch, r)
    ws.cell(r, 2).value = row.label or None
    if row.kw.get('note'): ws[f'CF{r}'] = row.kw['note']
    if row.kw.get('group'):
        ws.row_dimensions[r].outlineLevel = 1; ws.row_dimensions[r].hidden = True
    cg = row.key in CAGR_G or row.key in CAGR_A
    style_cagr(ws, r, cg)
    if row.arch in ('section', 'block', 'blank', 'sub') and not row.kw.get('data'):
        return
    if row.key in DEFS:
        write_defs(ws, row); return
    nf = NF[row.nf]
    for col in COLS:
        v, inp = compute(row, col)
        if v is None: continue
        cell = ws[f'{col.c}{r}']
        cell.value = v
        if not row.kw.get('textdata'): cell.number_format = nf
        if inp:
            set_color(cell, BLUE, bold=True if col.kind in FCK else None)
        elif row.arch == 'growth' and col.kind in HISTK:
            set_color(cell, GRAY)
        else:
            set_color(cell, BLACK)
        add_cell_comment(row, col, cell)
    if row.key in CAGR_G:
        ws[f'CH{r}'] = f'=IFERROR((CE{r}/BV{r})^(1/5)-1,"n/m")'
        ws[f'CI{r}'] = f'=IFERROR((BV{r}/AW{r})^(1/5)-1,"n/m")'
        ws[f'CJ{r}'] = f'=IFERROR((BV{r}/X{r})^(1/10)-1,"n/m")'
        ws[f'CK{r}'] = f'=IFERROR((BV{r}/C{r})^(1/19)-1,"n/m")'
        for c in ('CH', 'CI', 'CJ', 'CK'): ws[f'{c}{r}'].number_format = '0%;\\(0%\\);"n/m"'
    elif row.key in CAGR_A:
        ws[f'CH{r}'] = f'=IFERROR(AVERAGE(CA{r},CB{r},CC{r},CD{r},CE{r}),"n/m")'
        ws[f'CI{r}'] = f'=IFERROR(AVERAGE({",".join(f"{ANN[y].c}{r}" for y in range(2021, 2026))}),"n/m")'
        ws[f'CJ{r}'] = f'=IFERROR(AVERAGE({",".join(f"{ANN[y].c}{r}" for y in range(2016, 2026))}),"n/m")'
        ws[f'CK{r}'] = f'=IFERROR(AVERAGE({",".join(f"{c}{r}" for c in ANN_HIST)}),"n/m")'
        for c in ('CH', 'CI', 'CJ', 'CK'): ws[f'{c}{r}'].number_format = NF['pct']
    t = LABEL_SOURCES.get(row.key)
    if t: comment(ws.cell(r, 2), t)

def add_cell_comment(row, col, cell):
    if col.kind not in HISTK: return
    d = D.get(pkey(col)) or {}
    ng = d.get('ngaap') or {}
    if row.key == 'o_tot' and ng.get('adj_table'):
        rows = [f'{a}: {b.get("op")}' for a, b in ng['adj_table'][1:-1] if isinstance(b, dict) and num(b.get('op'))]
        if rows: comment(cell, 'Company wording (Reported → Adjusted, income from operations column): ' + ' | '.join(rows))
    elif row.key == 'e_disc' and ng.get('adj_table'):
        rows = [f'{a}: tax {b.get("tax")}, NCI {b.get("nci")}' for a, b in ng['adj_table'][1:-1] if isinstance(b, dict)
                and (num(b.get('op')) in (None, 0)) and (num(b.get('tax')) or num(b.get('nci')))]
        if rows: comment(cell, 'Tax / NCI-only rows in the company table: ' + ' | '.join(rows))
    elif row.key == 'e_ps' and ng.get('eps_bridge_lines'):
        comment(cell, 'Company per-share bridge ($/share, split-adjusted): ' + ' | '.join(f'{a}: {b}' for a, b in ng['eps_bridge_lines']))
    elif row.key == 'r_oth' and num(cell.value):
        comment(cell, 'Company Adjusted EBITDA add-back less restructuring / inventory step-up / transaction items in the Reported→Adjusted table '
                '(e.g. 2019 one-time PPA & policy harmonization; 2022–24 EBITDA add-back differs from the adjusted-results restructuring row).')

SRC = 'Source: Wabtec earnings releases (Form 8-K Ex. 99.1) — '
LABEL_SOURCES = {
    'f_svc': SRC + 'Appendix "Sales by Product Line" (from 2019; 2019 Q1–Q3 on the Q4-19 recast basis).',
    'f_rev': SRC + 'segment table; FY2006–12 from the 10-K segment note. Freight Group / Transit Group to 2018.',
    'f_adjop': SRC + 'Appendix "Reconciliation of Reported Results to Adjusted Results — by Segment" (from 2020; 2019 margins only).',
    'f_bl12': SRC + 'backlog table (12-month / total by segment); 10-K Item 1 backlog for 2006–2016 year ends.',
    'a_acq': SRC + 'Appendix "Reconciliation of Changes in Net Sales" (2019+); 10-Q / 10-K MD&A sales-change tables before.',
    'is_rev': 'Source: Condensed Consolidated Statements of Income (release Appendix A; 10-K for FY2006–12).',
    'is_dda': 'Source: EBITDA reconciliation (release appendix) where published, otherwise the cash-flow statement.',
    'r_ebitda_pub': SRC + 'Appendix "EBITDA Reconciliation" (published from 2019).',
    'o_adj_pub': SRC + 'Appendix "Reconciliation of Reported Results to Adjusted Results" (income from operations column); Q3-16 narrative.',
    'e_pub': SRC + 'adjusted diluted EPS as published (table bridges from Q4-17; narrative per-share bridges: FY2006, Q4-09, FY2011, Q1–Q2-16, Q3-16).',
    'cf_ni': 'Source: Consolidated Statements of Cash Flows (10-K / 10-Q year-to-date; discrete quarters derived; Q4 = FY − 9M).',
    'bs_cash': 'Source: Consolidated Balance Sheets (10-K / 10-Q).',
    'ppa_cons': 'Source: 10-K 2025 / Q1-26 & Q2-26 10-Q Note 3 Acquisitions (latest purchase-price allocation; Frauscher & Dellner preliminary).',
    'lp_spe': 'Source: 10-K note "Sales by product" (pre-2019 categories), as originally reported.',
    'dps': 'Source: releases / 10-Q / 10-K (split-adjusted: 2-for-1 split June 2013).',
}

DEFS = {}
def defs_setup():
    q = lambda y, n: QC[(y, n)].c
    DEFS['r_def'] = {
        'C': ('Not published 2006–2018 (EBITDA introduced with the GE Transportation merger)', None),
        q(2019, 1): ('2019: EBITDA = income from operations + D&A; Adjusted EBITDA adds restructuring/transaction, one-time PPA, policy harmonization',
                     'Q1–Q4 2019 as originally reported (merger basis). Later releases restated 2019 on the 2020 definition.'),
        q(2020, 1): ('2020+: EBITDA = income from operations + other income + D&A; Adjusted EBITDA + restructuring & transaction costs', None),
        q(2025, 3): ('Q3-25+: add-back also includes inventory purchase-accounting charges (Inspection Technologies, Frauscher, Dellner)', None),
    }
    DEFS['o_def'] = {
        'C': ('2006–2016: no adjusted operating income (narrative adjusted EPS only: FY2006, Q4-09, FY2011, 2016 quarters)', None),
        q(2017, 4): ('Q4-17–2018: Reported→Adjusted tables (restructuring / integration, contract adjustments, Faiveley & GE merger costs)', None),
        q(2019, 1): ('2019: GE Transportation merger items (inventory step-up, restructuring / transaction, policy harmonization); amortization NOT added back',
                     'As originally reported. The 2020 releases restated FY2019 adjusted EPS to $4.86 (vs $4.17) on the new basis.'),
        q(2020, 1): ('Q1-20+: all non-cash intangible amortization added back (FX losses also excluded through Q4-21)', None),
        q(2023, 4): ('Q4-23+: "Restructuring and Portfolio Optimization" (Integration 2.0 / 3.0)', None),
        q(2025, 3): ('Q3-25+: inventory purchase-accounting and transaction costs shown separately (2025–26 acquisitions)', None),
    }
    DEFS['e_def'] = {
        'C': ('FY2006: EPS ex Q3 restructuring & Q4 tax benefit (narrative); none published 2007–08, 2010, 2012–15 (FY2009: Q4 only; FY2011 narrative)', None),
        q(2016, 1): ('2016 quarters: adjusted EPS ex Faiveley transaction costs (narrative per-share bridges)', None),
        q(2017, 4): ('Q4-17: + U.S. tax reform (net $7.9m) as discrete tax item; 2018: transition-tax true-ups, GE merger financing', None),
        q(2019, 1): ('2019: merger basis (amortization not added back); tax on non-deductible transaction costs', None),
        q(2020, 1): ('Q1-20+: + non-cash amortization (after tax); Q4-21 amended-return tax benefit excluded', None),
    }

def write_defs(ws, row):
    for c, (txt, cm) in DEFS[row.key].items():
        cell = ws[f'{c}{row.r}']; cell.value = txt
        cell.font = Font(name='Calibri', sz=10, b=True, i=True, color='FFC00000')
        if cm: comment(cell, cm)

def write_scen(ws):
    for spec_ in SCEN_ROWS:
        r, kind = spec_[0], spec_[1]
        src = {'hdr': 9, 'out_hdr': 11, 'pt_hdr': 15, 'out': 12, 'pt': 16, 'lev': 28, 'blk': 30, 'yr': 31}[kind]
        for ci in range(CI('CN'), CI('CR') + 1):
            ws.cell(r, ci)._style = copy.copy(TM.cell(src, ci)._style)
        if kind in ('hdr', 'out_hdr', 'pt_hdr'):
            _, _, a, b, c_, d_, e = spec_
            for col, v in zip(['CN', 'CO', 'CP', 'CQ', 'CR'], [a, b, c_, d_, e]):
                if v is not None: ws[f'{col}{r}'] = v
        elif kind == 'out':
            _, _, lab, hi, mid, lo, note, nf = spec_
            ws[f'CN{r}'] = lab
            for col, v in zip(['CO', 'CP', 'CQ'], [hi, mid, lo]):
                ws[f'{col}{r}'] = v; ws[f'{col}{r}'].number_format = NF[nf]; set_color(ws[f'{col}{r}'], BLUE)
            ws[f'CR{r}'] = note
        elif kind == 'pt':
            _, _, lab, _, val, _, note, nf = spec_
            ws[f'CN{r}'] = lab; ws[f'CP{r}'] = val; ws[f'CP{r}'].number_format = NF[nf]; set_color(ws[f'CP{r}'], BLUE); ws[f'CR{r}'] = note
        elif kind == 'lev':
            _, _, lab, a, b, c_, note = spec_
            ws[f'CN{r}'] = lab
            for col, v in zip(['CO', 'CP', 'CQ'], [a, b, c_]):
                ws[f'{col}{r}'] = v; ws[f'{col}{r}'].number_format = NF['x']
            ws[f'CR{r}'] = note
        elif kind == 'blk':
            _, _, lab, db, mid, de, note, nf = spec_
            ws[f'CN{r}'] = lab; ws[f'CO{r}'] = db; ws[f'CP{r}'] = mid; ws[f'CQ{r}'] = de; ws[f'CR{r}'] = note
            for col in ('CO', 'CQ'):
                ws[f'{col}{r}'].number_format = '\\+#,##0;\\-#,##0' if nf == 'm' else '\\+0.0%;\\-0.0%'
        elif kind == 'yr':
            _, _, lab, hdr, base, clamp, _, nf = spec_
            ws[f'CN{r}'] = lab; ws[f'CP{r}'] = base
            if clamp:
                ws[f'CO{r}'] = f'=MAX(0,CP{r}+$CO${hdr})'; ws[f'CQ{r}'] = f'=MAX(0,CP{r}+$CQ${hdr})'
            else:
                ws[f'CO{r}'] = f'=CP{r}+$CO${hdr}'; ws[f'CQ{r}'] = f'=CP{r}+$CQ${hdr}'
            for col in ('CO', 'CP', 'CQ'): ws[f'{col}{r}'].number_format = NF[nf]

def write_header(ws):
    for r in range(1, 8):
        clone_row_style(ws, r, r, 2, CI('CR'))
    ws['CF1'] = 'SCENARIO SWITCH ▼ (Bull / Base / Bear) — drives all scenario-linked forecast drivers'
    ws['CF2'] = 'Base'
    comment(ws['CF2'], 'Scenario toggle. Choose Bull / Base / Bear. Drives: FY26 sales and adjusted-EPS outlook end-points (high / mid / low), '
            '2H/26 buybacks, 2027E–2030E product-line growth and segment adjusted margins, M&A spend and the leverage floor for 2027E+ buybacks '
            '(scenario table CN:CR).')
    ws['B2'] = 'Westinghouse Air Brake Technologies Corporation (Wabtec, NYSE: WAB) — 3-Statement Financial Model'
    ws['B3'] = ('US GAAP as reported · USD millions (per-share data in USD; shares in millions, split-adjusted) · Source: Wabtec Forms 10-K / 10-Q and Form 8-K '
                'earnings releases (Ex. 99.1) and earnings decks, deal 8-Ks, SEC EDGAR CIK 0000943452 (investor-relations releases are the same documents as filed on EDGAR)')
    ws['B4'] = ('Basis notes: calendar fiscal year; all periods as originally reported. 2-for-1 stock split Jun-2013 (earlier shares / per-share data restated). '
                'Faiveley Transport from 30-Nov-2016; GE Transportation merger 25-Feb-2019 (segments recast Q4-19; product-line sales from 2019); '
                'adjusted results add back non-cash amortization from Q1-20 (2019 merger basis as reported). Inspection Technologies 1-Jul-2025; Frauscher 1-Dec-2025; '
                'Dellner Couplers 10-Feb-2026. Cash flow incl. restricted cash from 2018 (ASU 2016-18). Q3/26E–Q4/26E calibrated to FY26 sales and adjusted-EPS guidance.')
    ws['CH5'] = 'Forecast'; ws['CI5'] = 'Historical'; ws['CJ5'] = 'Historical'; ws['CK5'] = 'Historical'
    ws['B6'] = '(USDm)'
    for col in COLS: ws[f'{col.c}6'] = col.label
    ws['CF6'] = 'Modelling Notes'
    for c, a, b in [('CH', '5Y CAGR', "25-'30"), ('CI', '5Y CAGR', "20-'25"), ('CJ', '10Y CAGR', "15-'25"), ('CK', '19Y CAGR', "06-'25")]:
        ws[f'{c}6'] = a; ws[f'{c}7'] = b

def main():
    import openpyxl
    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb['Model']
    blank_style = copy.copy(ws['A1']._style)
    for rowc in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for c in rowc:
            c.value = None; c.comment = None; c._style = copy.copy(blank_style)
    for k in list(ws.row_dimensions.keys()):
        rd = ws.row_dimensions[k]; rd.outlineLevel = 0; rd.hidden = False
    scen_layout(V0)
    build_spec(V0)
    n1 = len(fw.SPEC)
    build_spec2(V0)
    for i, row in enumerate(fw.SPEC): row.spec2 = i >= n1
    last = number_rows(8)
    for r in range(1, last + 1): ws.row_dimensions[r].height = 15.0
    ws.row_dimensions[2].height = 21.0
    defs_setup()
    write_header(ws)
    for row in fw.SPEC: write_row(ws, row)
    write_scen(ws)
    ws.freeze_panes = TM.freeze_panes
    import other_sheets as osh
    osh.do_all(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    osh.inject_charts(OUT)
    print('rows', fw.SPEC[-1].r, 'saved', OUT)
    return wb

RECALC = os.environ.get('XLSX_RECALC', '/mnt/skills/public/xlsx/scripts/recalc.py')
def recalc(path, timeout=600):
    import subprocess
    out = subprocess.run(['python3', RECALC, path, str(timeout)], capture_output=True, text=True, cwd=os.path.dirname(RECALC))
    try:
        return json.loads(out.stdout)
    except Exception:
        return {'raw': out.stdout[-2000:], 'err': out.stderr[-2000:]}

def scenario_snapshots():
    import openpyxl, shutil
    import other_sheets as osh
    snaps = {}
    for scen in ('Bull', 'Bear'):
        p = os.path.join(BASE, f'WAB_{scen}.xlsx')
        wb = openpyxl.load_workbook(OUT); wb['Model']['CF2'] = scen; wb.save(p)
        res = recalc(p); print(scen, res.get('status'), res.get('total_errors'))
        v = openpyxl.load_workbook(p, data_only=True)['Bull-Base-Bear']
        snaps[scen] = {(r, c): v[f'{c}{r}'].value for r in range(9, 48) for c in 'CDEFGHIJKL' if v[f'{c}{r}'].value is not None}
    wb = openpyxl.load_workbook(OUT)
    osh.snapshot_values(wb, snaps['Bull'], snaps['Bear'])
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    osh.inject_charts(OUT)

if __name__ == '__main__':
    main()
    if '--snap' in sys.argv: scenario_snapshots()
