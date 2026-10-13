"""Rewire the Bull-Base-Bear, DCF and Charts tabs of the MLM template (current layout: DCF with live valuation date, fade period,
implied ROIC, target-IRR entry price, reverse DCF and Mauboussin EV/NOPAT; Bull-Base-Bear with buyback rows) to the GOOG Model rows;
keep the template chart XML. The template's DCF_old tab is not carried.

Columns: GOOG uses the template's columns unchanged (2025A base column BV, 2026E CA, 2027E-2030E CB-CE), so only rows are remapped.
DCF: net debt is taken at YE2026E (Model CA), so the 2026E UFCF is excluded from the PV sum (its cash flows are already in YE2026E net
debt; BWXT / WWD lesson); explicit years 2027E-2030E, fade 2031E-2035E."""
import re, zipfile, shutil
from openpyxl.utils import get_column_letter as L
from fw import R, TEMPLATE, ANN, NOTE

# template Model row -> GOOG row key
MAP = {158: 'is_rev', 164: 'r_cash', 166: 'is_oi', 182: 'is_dda', 186: 'sh_dil', 211: 'r_adj', 212: 'r_adj_m', 219: 'r_ebit', 221: 'r_adj',
       229: 'e_adjni', 231: 'e_eps', 238: 'gm_cloud', 245: 'gm_ebitda_m', 247: 'gm_op_m', 260: 'cf_wc', 262: 'cf_capex', 263: 'cf_acq',
       280: 'fcf', 282: 'fcf_m', 290: 'bb_px', 291: 'bb_cash', 292: 'bb_sh', 331: 'v_stake', 408: 'ra_ic', 411: 'ra_roic', 422: 'rn_ronta',
       428: 'lv_nd', 433: 'v_px', 434: 'v_sh'}
REF = re.compile(r"Model!(\$?)([A-Z]{1,3})(\$?)(\d+)")
HIST = [ANN[y].c for y in range(2006, 2026)]

def remap(f):
    def rep(m):
        c, r = m.group(2), int(m.group(4))
        r2 = R[MAP[r]] if r in MAP else r
        if c == 'CF' and r == 2: r2 = 2
        return f'Model!{m.group(1)}{c}{m.group(3)}{r2}'
    return REF.sub(rep, f)

def hist_stat(f):
    """LT historical AVERAGE / MIN / MAX over the template's annual columns -> GOOG FY2006-FY2025 annual columns."""
    m = re.match(r'=IFERROR\((AVERAGE|MIN|MAX)\(Model!C(\d+),', f)
    if not m: return None
    fn, r = m.group(1), int(m.group(2))
    rr = R[MAP[r]]
    return f'=IFERROR({fn}({",".join(f"Model!{c}{rr}" for c in HIST)}),"")'

TEXT = {
    'Bull-Base-Bear': {
        'B2': 'Alphabet Inc. (GOOG) — Bull / Base / Bear Case P&L Summary',
        'B3': ('Model-defined adjusted figures (adjusted EBITDA / operating income / EPS exclude company-reconciled EC fines and severance charges, equity-security gains and the Tax Act; '
               'Alphabet publishes no EBITDA or adjusted EPS) · USD millions · Live panel linked to Model — toggle scenario via Model!CF2. Bull and Bear panels are value snapshots '
               'generated from the Model with the switch set to each case. 2026E = Q1–Q2 actual + Q3–Q4 estimate (no company guidance).'),
        'G26': 26, 'H26': 26, 'S26': 30, 'T26': 30, 'AD26': 20, 'AE26': 20, 'AF26': 20,
    },
    'DCF': {
        'B2': 'DCF VALUATION — ALPHABET INC. (GOOG)',
        'C5': 0.0425, 'D5': 'Assumption (~US 10Y Treasury yield, Oct-2026; government rate sources not reachable from the build environment); update to market',
        'D6': 'Assumption: US implied ERP ~4.5–5.5%',
        'C7': 1.05, 'D7': 'Assumption: mega-cap internet / cloud peers (MSFT, AMZN, META) ~1.0–1.2; advertising cyclicality offset by subscription and cloud backlog ($519.5bn RPO)',
        'C9': 0.047, 'D9': '2026 senior notes: USD w.a. coupon 4.80% (15 years), GBP 5.31%, EUR 3.90%, CAD 4.35% (Q2-26 10-Q Note 6); AA+ / Aa2 issuer',
        'C10': 0.175, 'D10': 'Model 2027E+ effective tax rate (FY2025 16.8%)',
        'C12': 0.05, 'D12': 'Net cash ~$0.1tn at YE2026E vs ~$4.3tn market cap ($350 × ~12.3bn diluted shares); debt ~$100bn, preferred converts in 2029',
        'D15': 'Periods = (31-Dec of the year − today) / 365.25; cash flows at year-end. 2026E UFCF is excluded from the sum: YE2026E net debt already reflects 2026 cash flows',
        'B18': 'Adj. EBIT less cash adjusting items (EC fines, severance & office space)',
        'B21': 'plus D&A (depreciation + amortization of intangibles)',
        'B22': '− Capex (property & equipment) and acquisitions (M&A lever 2027E–30E; Wiz / Intersect in 2026E)',
        'B26': 'Fade-period growth (2031E–35E; 9.0% → 6.0% linear)', 'I26': 0.09, 'M26': 0.06, 'C27': 0.035,
        'D32': ("Alphabet's adjusted EBIT (model) is after stock-based compensation (~6.5% of revenues) and amortization; only company-reconciled cash items (EC fines, severance and "
                'office-space charges) are added back, and they are deducted again in row 18. D&A (row 21) is added back in full because adjusted EBIT is full-cost. Row 22 includes '
                'capex at the call outlook / scenario levels ($200bn 2026E, $250–300bn 2027E–30E; unverified) and acquisitions. Equity-security gains are excluded from NOPAT; '
                'non-marketable equity stakes are added back at carrying value in the equity bridge (row 48). 2026E UFCF is not summed (net debt at YE2026E).'),
        'D43': 'Opening IC = Model 2025A invested capital (equity + net debt − non-marketable equity securities), rolled forward with capex + M&A − D&A + ΔWC (2026E includes Wiz / Intersect).',
        'B48': 'Less: Net debt − non-marketable equity securities (Model 2026E YE)',
        'D48': 'GOOG: net debt (debt − cash − marketable securities; negative = net cash) less non-marketable equity stakes at carrying value (outside invested capital). No NCI.',
        'B50': 'Diluted shares (m, Model 2026E period end; incl. RSUs and preferred if-converted)',
        'D76': 'Rationale (Alphabet characteristics)',
        'C78': 0.10, 'D78': 'Model adjusted EBIT CAGR 2025–30E: Search / YouTube / subscriptions high-single to low-double digit, Cloud 20–45% with rising margins, against rising depreciation → ~10% NOPAT growth',
        'C79': 0.20, 'D79': 'Search and YouTube earn very high returns on little capital; incremental returns now hinge on AI infrastructure (capex ~45% of revenues in 2026E–27E) monetised through Cloud backlog and AI products',
        'C80': 15, 'D80': 'Search distribution, data and TPU / infrastructure scale sustain a long competitive-advantage period; risks: AI-assistant substitution of search, antitrust remedies (US search and ad-tech cases, EC fines)',
        'B90': 'VALUATION  —  Applied to 2028E NOPAT',
        'B91': '2028E Adj. EBIT less cash adjusting items',
        'B96': 'Less: YE2027 net debt − non-marketable equity securities (Model)',
        'B106': 'Plus: YE2026 net debt − non-marketable equity securities',
        'B108': 'GOOG market-implied EV/NOPAT (on 2028E NOPAT)',
        'C45': '=SUM(E31:M31)',
    },
    'Charts': {
        'B2': 'Alphabet Inc. — Operating Performance (2011–2025) and Model Forecast (2026E–2030E)',
        'B3': ('USDm · Linked to Model. Bars = contribution of Google Services and Google Cloud to total revenue growth (pp; current revenue taxonomy from 2017–18, earlier years n/a). '
               'Adjusted EBITDA and operating income = model definitions. ROIC / RONTA on adjusted EBIT (invested capital excludes equity stakes).'),
        'B8': 'Google Services contribution to growth (pp)', 'B9': 'Google Cloud contribution to growth (pp)', 'B12': 'Adjusted EBITDA (model)', 'B13': 'Income from operations (GAAP)',
        'B14': 'Adj. EBITDA Margin %', 'B15': 'Operating Margin %',
    },
}
BBB_ROWS = {9: 'Revenues', 10: '   Revenues y/y %', 11: '   Google Cloud revenues y/y %', 12: 'Adjusted EBITDA (model)',
            15: 'Income from operations (GAAP)', 16: '   Operating income y/y %', 17: '   Operating margin (GAAP) %', 19: 'Adjusted Net Income (model)',
            35: 'Adjusted EBITDA (leverage basis)', 36: '   Adjusted EBITDA % of revenues', 38: 'Net Debt (negative = net cash)',
            48: '(+) Net debt − non-marketable equity securities'}
CHART_TITLES = {
    'Aggregates Revenue Drivers: Shipment Volume vs. Price (ASP), y/y %': 'Revenue Growth: Google Services vs. Google Cloud Contribution, pp',
    'Adjusted EBITDA &amp; Earnings from Operations (USDm)': 'Adjusted EBITDA &amp; Income from Operations (USDm)',
}

def build_charts(ws):
    for i, y in enumerate(range(2011, 2031)):
        c = L(3 + i); mc = ANN[y].c
        ws[f'{c}7'] = y if y <= 2025 else f'{y}E'
        ws[f'{c}8'] = f'=IFERROR(Model!{mc}{R["gc_svc"]}+0,0)'
        ws[f'{c}9'] = f'=IFERROR(Model!{mc}{R["gc_cloud"]}+0,0)'
        ws[f'{c}11'] = f'={c}7'
        ws[f'{c}12'] = f'=Model!{mc}{R["r_adj"]}'
        ws[f'{c}13'] = f'=Model!{mc}{R["is_oi"]}'
        ws[f'{c}14'] = f'=Model!{mc}{R["gm_ebitda_m"]}'
        ws[f'{c}15'] = f'=Model!{mc}{R["gm_op_m"]}'
        ws[f'{c}17'] = f'={c}7'
        ws[f'{c}18'] = f'=IFERROR(Model!{mc}{R["ra_roic"]}+0,NA())'
        ws[f'{c}19'] = f'=IFERROR(Model!{mc}{R["rn_ronta"]}+0,NA())'
        ws[f'{c}20'] = f'=IFERROR(Model!{mc}{R["gm_ebitda_m"]}+0,NA())'
        ws[f'{c}21'] = f'=IFERROR(Model!{mc}{R["fcf_m"]}+0,NA())'

def do_all(wb):
    for name in ('Bull-Base-Bear', 'DCF', 'Charts'):
        ws = wb[name]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and 'Model!' in c.value:
                    h = hist_stat(c.value) if name == 'Bull-Base-Bear' else None
                    c.value = h or remap(c.value)
        for addr, t in TEXT.get(name, {}).items():
            ws[addr] = t
    b = wb['Bull-Base-Bear']
    for r, t in BBB_ROWS.items():
        for c in ('B', 'N', 'Z'): b[f'{c}{r}'] = t
    from openpyxl.comments import Comment
    cm = Comment('P/E on adjusted EPS (model definition: excludes equity-security gains, EC fines and the Tax Act). GOOG at $350 trades at ~29x 2026E model adjusted EPS. '
                 'Exit multiples are assumptions: Base 26x, Bull 30x, Bear 20x applied to 2029E–30E (2028E–30E for Bear), vs the FY2021–25 range of ~18–30x.', 'GOOG model')
    cm.width, cm.height = 420, 120; b['N26'].comment = cm
    cm2 = Comment(f'Snapshot (values) of Model rows {R["bb_cash"]}, {R["bb_px"]} and {R["bb_sh"]} (Share Schedule) with Model!{NOTE}2 = Bull / Bear. '
                  'Base panel (D40:H42) is live-linked.', 'GOOG model')
    cm2.width, cm2.height = 420, 90; b['P40'].comment = cm2
    b['J2'] = 'Share price ($):'
    for c in ('B', 'N', 'Z'):
        if b[f'{c}40'].value: b[f'{c}40'] = '   Share repurchases ($m; 2026E = 2H)'
    d = wb['DCF']
    a25, a26, a27 = ANN[2025].c, ANN[2026].c, ANN[2027].c
    d['D22'] = f'=Model!{a26}{R["cf_capex"]}+Model!{a26}{R["cf_acq"]}'
    d['D34'] = f'=Model!{a25}{R["ra_ic"]}'
    d['C48'] = f'=Model!{a26}{R["lv_nd"]}+Model!{a26}{R["v_stake"]}'
    d['C50'] = f'=Model!{a26}{R["v_sh"]}'
    d['C52'] = f'=Model!{a26}{R["v_px"]}'
    d['C96'] = f'=Model!{a27}{R["lv_nd"]}+Model!{a27}{R["v_stake"]}'
    # fade-period capex converges linearly from the 2030E level to D&A + NOPAT x g / terminal RONIC in 2035E (a perpetuity at the
    # 2030E capex intensity would imply ~88% reinvestment at ~7% returns); the terminal RONIC is an input
    d['B28'] = 'Terminal RONIC on new investment (fade capex converges to D&A + NOPAT × g ÷ RONIC by 2035E)'
    d['C28'] = 0.20
    d['C28'].number_format = '0.0%'
    d['D28'] = ('Assumption. 2030E capex ($300bn, 36% of revenues, 2.6x D&A) is a build-out level; holding that intensity in perpetuity would imply ~88% reinvestment at ~7% '
                'returns. Model ROIC 2030E ~17%; 20% terminal RONIC = reinvestment of 17.5% of NOPAT at 3.5% growth.')
    tgt = '(-($M$21+$M$20*$C$27/$C$28))'
    for k, col in enumerate('IJKLM', start=1):
        d[f'{col}22'] = f'=$H$22+({tgt}-$H$22)*{k}/5'
    d['B22'] = '− Capex (property & equipment) and acquisitions (M&A lever 2027E–30E; Wiz / Intersect in 2026E); fade: converges to the terminal reinvestment level'
    for r in range(17, 32):                      # DCF CAGR columns: n/m when a base is negative (2026E UFCF after $200bn capex)
        for c in ('P', 'Q'):
            v = d[f'{c}{r}'].value
            if isinstance(v, str) and v.startswith('=') and not v.startswith('=IFERROR('):
                d[f'{c}{r}'] = f'=IFERROR({v[1:]},"n/m")'
    build_charts(wb['Charts'])

def snapshot_values(wb, bull, bear):
    ws = wb['Bull-Base-Bear']
    cols = 'CDEFGHIJKL'
    for snap, off in ((bull, 'OPQRSTUVWX'), (bear, ['AA', 'AB', 'AC', 'AD', 'AE', 'AF', 'AG', 'AH', 'AI', 'AJ'])):
        for (r, c), v in snap.items():
            tgt = f'{off[cols.index(c)]}{r}'
            cur = ws[tgt].value
            if isinstance(cur, str) and cur.startswith('='): continue
            if r in (7, 26): continue
            ws[tgt] = v

def inject_charts(path):
    """Replace openpyxl-rewritten chart parts with the template chart XML (combo charts, colours, layout), retitled, ranges kept
    (Charts!C:V = 2011-2030E) and cached values removed so Excel redraws from the GOOG cells."""
    tz = zipfile.ZipFile(TEMPLATE)
    tmpl = {n: tz.read(n) for n in tz.namelist() if n.startswith('xl/charts/chart')}
    tmp = path + '.tmp'
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in tmpl:
                t = tmpl[item.filename].decode('utf-8')
                for a, b2 in CHART_TITLES.items(): t = t.replace(a, b2)
                t = re.sub(r'<c:numCache>.*?</c:numCache>', '', t, flags=re.S)
                t = re.sub(r'<c:strCache>.*?</c:strCache>', '', t, flags=re.S)
                data = t.encode('utf-8')
            zout.writestr(item, data)
    shutil.move(tmp, path)
