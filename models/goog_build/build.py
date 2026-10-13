"""Build GOOG_Model.xlsx (Alphabet Inc.) from the MLM template (same layout, formatting and process as the prior models)."""
import sys, copy, os, json
from openpyxl.styles import Font
from fw import *
import fw
from spec import build_spec, scen_layout, SCEN_ROWS, SC
from spec2 import build_spec2

OUT = os.path.join(os.path.dirname(BASE), 'GOOG_Model.xlsx')
GD = MAN['guidance']

# ====================================================================== inputs (V0)
def wc_inputs():
    q = DATA['2025']['x']
    return {'dso': round(q['bs_ar'] / q['rev'] * 365), 'apc': 0.12, 'ocl': round(q['bs_ocl'] / q['rev'], 3)}
WC = wc_inputs()
WC['note'] = (f'Forecast inputs at FY2025 year-end levels: DSO {WC["dso"]}d on revenues; payables 12% of annual capex (FY2025 13.3%, Q2-26 11.3% annualised); '
              f'other current liabilities {WC["ocl"]:.1%} of revenues (incl. accrued legal & regulatory fines and settlements of $17.4bn at 30-Jun-2026).')

V0 = {
    'out_hdr': 'Q3–Q4/26E outlook — Alphabet gives no revenue, margin or EPS guidance (scenario shifts vs the 1H26 run-rate)',
    'out_note': ('Capex outlook 2026 $195–205bn (raised from $180–190bn) is call-only (22-Jul-2026) and UNVERIFIED; the Q2-26 10-Q says only that 2026 '
                 'technical-infrastructure investment will "significantly increase" vs 2025 (1H26 $80.6bn).'),
    'outlook': [
        ('out_dg', '  Q3–Q4/26E revenue growth shift vs 1H26 y/y (pp; all revenue types)', (0.02, 0.0, -0.03), 'pct',
         'Base = 1H26 y/y run-rate (Search +17.9%, YouTube +11.8%, Network −2.3%, subscriptions +17.2%, Cloud +73% incl. Wiz from 11-Mar-2026). Assumption.'),
        ('out_dm', '  Q3–Q4/26E Services / Cloud operating-margin shift (pp)', (0.01, 0.0, -0.015), 'pct',
         "Base = Q2-26 margins (Services with last year's sequential seasonality). Bear reflects leased third-party data-centre capacity in Q3 (call, unverified)."),
        ('bb26', '  2H26E share repurchases ($m)', (5000.0, 0.0, 0.0), 'm', 'No repurchases in 1H26 after the Jun-2026 equity raise; $69.5bn authorization remaining.'),
    ],
    'points': [
        ('out_rev', '  Optional FY26 revenue target ($m; blank = driver-based)', None, 'm', 'Blank by default (no company guidance). Enter a value (e.g. consensus) to solve the 2H revenue shift to it.'),
        ('pt_capex', '  FY26 capital expenditures ($m) — UNVERIFIED call guidance mid-point', GD['capex_2026']['low'] / 2 + GD['capex_2026']['high'] / 2, 'm', GD['capex_2026']['status']),
        ('pt_depg', '  Q3–Q4/26E depreciation, sequential growth', 0.096, 'pct', 'Q2/Q1-26 +9.6% ($7,104m vs $6,482m); capex ramp.'),
        ('pt_amort', '  Q3–Q4/26E amortization of intangibles ($m per quarter)', 370.0, 'm', 'Q2-26 $367m (Wiz intangibles from Mar-2026).'),
        ('pt_etr', '  Q3–Q4/26E effective tax rate', 0.18, 'pct', '1H26 19.1% incl. tax on equity-security gains; FY2025 16.8%.'),
        ('pt_yield', '  Q3–Q4/26E interest yield on cash + marketable securities (annual)', 0.033, 'pct', 'Q2-26 $1,430m on ~$185bn average cash and securities (~3.1%); $242.5bn at 30-Jun-2026.'),
        ('pt_intexp', '  Q3–Q4/26E interest expense ($m per quarter)', 1150.0, 'm', 'Q2-26 $1,278m (net of $363m capitalized interest); ~$100bn of debt at ~4.0% weighted coupon.'),
        ('pt_prefq', '  Q3–Q4/26E preferred dividends accrued ($m per quarter)', 296.875, 'm', '6.25% × $19.0bn liquidation preference ÷ 4.'),
        ('pt_def', '  2H26E deferred income taxes ($m)', 0.0, 'm', 'Not forecast (assumption).'),
    ],
    'lev': (-0.10, -0.25, -0.40),
    'lev_note': ('Alphabet is net cash ($142bn at 30-Jun-2026 after the $49.6bn equity raise; $100bn of debt vs $242bn of cash and marketable securities). '
                 'Negative target = net-cash floor: buybacks return liquidity above |x| × adjusted EBITDA. Bull −0.10x (more buybacks), Base −0.25x, Bear −0.40x (assumption).'),
    'blocks': [
        ('SEARCH', 'SEARCH (Google Search & other y/y)', 0.02, -0.04, [0.11, 0.09, 0.08, 0.07],
         '1H26 +17.9% (FY2025 +13.4%, FY2024 +13.2%); AI Overviews / AI Mode monetisation vs competition from AI assistants. Assumption.', 'pct', False),
        ('YOUTUBE', 'YOUTUBE (YouTube ads y/y)', 0.02, -0.04, [0.11, 0.10, 0.09, 0.08], '1H26 +11.8% (FY2025 +11.7%). Assumption.', 'pct', False),
        ('NETWORK', 'NETWORK (Google Network y/y)', 0.02, -0.03, [-0.03, -0.03, -0.03, -0.03], '1H26 −2.3% (FY2025 −1.9%, FY2024 −3.0%): structural decline. Assumption.', 'pct', False),
        ('SUBS', 'SUBS (subscriptions, platforms & devices y/y)', 0.02, -0.04, [0.14, 0.12, 0.11, 0.10],
         '1H26 +17.2% (FY2025 +19.1%): YouTube subscriptions, Google One / AI plans, Play. Assumption.', 'pct', False),
        ('CLOUD', 'CLOUD (Google Cloud y/y)', 0.08, -0.12, [0.45, 0.32, 0.25, 0.20],
         '1H26 +73% (Q2-26 +82%; Wiz from 11-Mar-2026; TPU system sales); backlog $513.9bn at 30-Jun-2026 (about half within 24 months) vs FY2025 revenue $58.7bn. Capacity-constrained. Assumption.', 'pct', False),
        ('OB', 'OB (Other Bets revenues y/y)', 0.10, -0.10, [-0.30, 0.25, 0.25, 0.25],
         'GFiber (internet services; revenue not disclosed) deconsolidated at its late-2026 closing → 2027E −30% (assumption); Waymo growth thereafter.', 'pct', False),
        ('SVC_M', 'SVC_M (Google Services operating margin)', 0.015, -0.025, [0.420, 0.405, 0.395, 0.385],
         '1H26 43.5% (FY2025 40.7%, FY2024 39.8%). Steps down as allocated AI-infrastructure depreciation rises (model D&A 6% → 14% of revenues 2025–30E on $250–300bn p.a. capex), '
         'keeping group adjusted EBITDA margin ~42–44% (1H26 41%). Assumption.', 'pct', False),
        ('CLOUD_M', 'CLOUD_M (Google Cloud operating margin)', 0.02, -0.05, [0.33, 0.32, 0.32, 0.32],
         '1H26 34.4%, Q2-26 35.6% (FY2025 23.7%, FY2024 14.1%); depreciation of new capacity, leased third-party data centres and Wiz amortization offset scale. Assumption.', 'pct', False),
        ('OB_OI', 'OB_OI (Other Bets operating income, $m)', 1000.0, -1500.0, [-8000.0, -8500.0, -9000.0, -9000.0],
         'FY2025 $(7.5)bn, 1H26 $(3.9)bn: Waymo expansion. Assumption.', 'm', False),
        ('ALPHA_P', 'ALPHA_P (Alphabet-level costs % of revenues)', -0.005, 0.007, [0.050, 0.050, 0.048, 0.046],
         'FY2025 4.2%, 1H26 4.9%: frontier AI model development (Google DeepMind), corporate costs, fines & settlements. Assumption.', 'pct', False),
        ('CAPEX', 'CAPEX (purchases of property & equipment, $m)', -10000.0, 15000.0, [250000.0, 270000.0, 285000.0, 300000.0],
         '2026E $200bn (UNVERIFIED call outlook); "increase significantly in 2027" (call, unverified; no figure). Assumption.', 'm', True),
        ('MNA', 'M&A_spend (acquisitions, $m)', 5000.0, -5000.0, [5000.0, 5000.0, 5000.0, 5000.0],
         'Wiz $29.5bn and Intersect $5.9bn in 2026; FY2021–25 $0.5–6.3bn p.a. Base $5bn p.a. (assumption).', 'm', True),
    ],
    'rv_notes': {
        'search': 'Revenue types per the 10-Q / 10-K revenue note (XBRL; latest filing presenting the period). YouTube ads and Cloud are disclosed from FY2017 (annual) / Q1-19 (quarterly).',
        'subs': '"Google other" to Q4-23, renamed "Google subscriptions, platforms, and devices" from Q1-24 (no recast).',
        'cloud': 'Google Cloud segment revenues (from Q4-20; earlier the Google Cloud revenue type, incl. in the Google segment). Includes Wiz from 11-Mar-2026.',
        'ob': 'Other Bets segment revenues (from FY2015; recast basis where presented).',
    },
    'mna_note': 'Scenario lever M&A_spend (Base $5bn p.a.). Funded from cash / commercial paper (debt schedule).',
    'gf_cash': 1500.0, 'gf_hfs': 7100.0,
    'dep_r': 0.115, 'dep_note': 'FY2025 12.4% of opening net property & equipment; 1H26 ~11% annualised (longer-lived data-centre buildings and land in the mix). 2027E+ input 11.5% (assumption).',
    'amort_sched': {2027: 1450.0, 2028: 1400.0, 2029: 1350.0, 2030: 1300.0},
    'amort_note': 'Wiz intangibles $8.3bn (developed technology $3.6bn / 7 years, customer relationships $4.5bn / 10 years, trade names $0.2bn / 7 years ≈ $1.0bn p.a.) + other (~$0.4bn). Assumption (no post-Wiz schedule filed).',
    'sbc_p': 0.065, 'rd_p': 0.152, 'sm_p': 0.068, 'ga_p': 0.048, 'etr': 0.175,
    'dps_q': 0.22, 'dps_g': 0.05, 'def_ae': 0.0, 'iss_2h': 60.0, 'iss_ae': 120.0, 'dil': 145.0, 'sbcnet_r': 0.60, 'atm': 0.0, 'pref_cash_2h': 527.725,
    'repay': {2026: 2000.0, 2027: 1000.0, 2028: 2676.0, 2029: 1764.0, 2030: 5500.0},
    'mat': {2026: 1000.0, 2027: 2676.0, 2028: 1764.0, 2029: 5500.0, 2030: 5000.0},
    'repay_note': ('FY2025 10-K maturity ladder (face): 2026 $2,000m (2.00% notes, Aug-2026), 2027 $1,000m, 2028 $2,676m, 2029 $1,764m, 2030 $5,500m. '
                   'The 2026 USD / GBP / CHF / EUR / CAD / JPY issues mature 2029–2126 and are not split by year in the 10-Q (assumed beyond 2030). Maturities refinanced (assumption).'),
    'min_cash': 25000.0,
    'r_notes': 0.042, 'r_notes_note': 'Weighted coupon ~4.0% on $101bn of notes at 30-Jun-2026 (2026 issues 1.06–5.31%, 2025 issues 2.38–5.70%, 2020 notes 0.80–2.25%) + credit-facility / project debt. Assumption 4.2%.',
    'yield': 0.032, 'yield_note': 'Interest income ÷ average cash and marketable securities: FY2025 ~4.0%, Q2-26 ~3.1% annualised. Assumption 3.2%.',
    'wc': WC,
    'px': MAN['price']['px'],
    'px_note': ('VALUATION — share price input $350.00 (GOOG Class C, user-provided 13-Oct-2026; latest SEC-filed Class C price $340.45, Form 4 sale 30-Sep-2026); '
                'applied to all share classes (GOOGL trades within ~1%); update to market'),
}

# ====================================================================== compute engine
def hist_value(row, col):
    kw = row.kw; path = kw.get('data')
    if not path: return None
    if kw.get('ppa'):
        v = MAN['ppa'].get(pkey(col), {}).get(path.split('_', 1)[1])
        return v if kw.get('textdata') else num(v)
    v = xget(pkey(col), path)
    return v if kw.get('textdata') else num(v)

def first(kw, keys):
    for k in keys:
        if k in kw: return k
    return None

def compute(row, col):
    kw = row.kw; k = col.kind
    CTX.col = col
    if kw.get('ann') and col.q: return None, False
    if k in HISTK:
        v = hist_value(row, col)
        if v is not None: return v, True
        for key in ('hist', 'all'):
            if key in kw:
                return kw[key](), False
        return None, False
    if k == 'QE':
        if row.spec2: return None, False
        if 'qe_in' in kw:
            v = kw['qe_in']; return (v[col.q] if isinstance(v, dict) else v), True
        key = first(kw, ('qe', 'fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    if k == 'E26':
        if 'e26_in' in kw: return kw['e26_in'], True
        if 'e26' in kw: return kw['e26'](), False
        if kw.get('avg') and ('qe' in kw or 'qe_in' in kw): return f'=AVERAGE({Q1}{row.r}:{QE}{row.r})', False
        if kw.get('flow') and not row.spec2 and any(x in kw for x in ('qe', 'qe_in', 'fc')): return S26(row.key), False
        key = first(kw, ('fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    if k == 'AE':
        if 'ae_in' in kw:
            v = kw['ae_in']; return (v[col.year] if isinstance(v, dict) else v), True
        key = first(kw, ('ae', 'fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    return None, False

# ====================================================================== writers
CAGR_G = {'is_rev', 'g_rev', 'rv_tot', 'rv_search', 'rv_youtube', 'rv_network', 'rv_subs', 'rv_cloud', 'rv_svc', 'rv_adv', 'so_svc', 'g_oi', 'is_oi',
          'o_adj', 'r_adj', 'r_ebitda', 'd_dda', 'd_dep', 'd_sbc', 'e_adjni', 'e_eps', 'cf_cfo', 'fcf', 'bs_ta', 'bs_teq', 'sh_dil', 'is_ni', 'eps_dil',
          'is_cor', 'is_rd', 'is_sm', 'is_ga', 'kpi_tac', 'bs_ppe', 'kpi_emp'}
CAGR_A = {'so_svc_m', 'so_cloud_m', 'so_alpha_p', 'g_oim', 'o_adj_m', 'r_adj_m', 'is_gm', 'is_oim', 'r_ebit_m', 'gm_rev', 'gm_svc', 'gm_cloud', 'gm_gm',
          'gm_ebitda_m', 'gm_op_m', 'gm_net_m', 'gm_rd', 'gm_sm', 'gm_ga', 'gm_adjop_m', 'fcf_m', 'fcf_conv', 'capex_pct', 'ra_roic', 'rn_ronta', 'roe', 'lv_x',
          'd_rd', 'd_sm', 'd_ga', 'd_sbc_p', 'd_dep_r', 'wc_nwc_p', 'is_etr', 'e_rate', 'rv_cloud_mix', 'kpi_tac_p', 'rv_tot_g', 'rv_search_g', 'rv_youtube_g',
          'rv_network_g', 'rv_subs_g', 'rv_cloud_g'}
A = {y: ANN[y].c for y in ANN}

def style_cagr(ws, r, on):
    src = 16 if on else 12
    for c in CAGR:
        ws[f'{c}{r}']._style = copy.copy(TM[f'{L(CI(c) + TSHIFT)}{src}']._style)

def write_row(ws, row):
    r = row.r; arch = ARCH[row.arch]
    clone_row_style(ws, arch, r)
    ws.cell(r, 2).value = row.label or None
    if row.kw.get('note'): ws[f'{NOTE}{r}'] = row.kw['note']
    style_cagr(ws, r, row.key in CAGR_G or row.key in CAGR_A)
    if row.arch in ('section', 'block', 'blank', 'sub') and not row.kw.get('data'):
        return
    if row.key in DEFS:
        write_defs(ws, row); return
    nf = NF[row.nf]
    for col in COLS:
        v, inp = compute(row, col)
        if v is None: continue
        cell = ws[f'{col.c}{r}']
        cell.value = v
        if not row.kw.get('textdata'): cell.number_format = nf
        if inp:
            set_color(cell, BLUE, bold=True if col.kind in FCK else None)
            lk = row.kw.get('labkey')
            if lk and col.kind in HISTK and isinstance(v, (int, float)) and abs(v) > 1e-9:
                t = xget(pkey(col), lk)
                if t: comment(cell, f'Company items ({pkey(col)}): {t}', 380, 110)
        elif row.arch == 'growth' and col.kind in HISTK:
            set_color(cell, GRAY)
        else:
            set_color(cell, BLACK)
    rc = CELL_COMMENTS.get(row.key, {})
    for c, t in rc.items(): comment(ws[f'{c}{r}'], t)
    if row.key in CAGR_G:
        ws[f'{CAGR[0]}{r}'] = f'=IFERROR(({A[2030]}{r}/{A[2025]}{r})^(1/5)-1,"n/m")'
        ws[f'{CAGR[1]}{r}'] = f'=IFERROR(({A[2025]}{r}/{A[2020]}{r})^(1/5)-1,"n/m")'
        ws[f'{CAGR[2]}{r}'] = f'=IFERROR(({A[2025]}{r}/{A[2015]}{r})^(1/10)-1,"n/m")'
        ws[f'{CAGR[3]}{r}'] = f'=IFERROR(({A[2025]}{r}/{A[2006]}{r})^(1/19)-1,"n/m")'
        for c in CAGR: ws[f'{c}{r}'].number_format = '0%;\\(0%\\);"n/m"'
    elif row.key in CAGR_A:
        ws[f'{CAGR[0]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2026, 2031))}),"n/m")'
        ws[f'{CAGR[1]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2021, 2026))}),"n/m")'
        ws[f'{CAGR[2]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2016, 2026))}),"n/m")'
        ws[f'{CAGR[3]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2006, 2026))}),"n/m")'
        for c in CAGR: ws[f'{c}{r}'].number_format = NF['pct']
    t = LABEL_SOURCES.get(row.key)
    if t: comment(ws.cell(r, 2), t)

REL = 'Source: Google / Alphabet earnings releases (Form 8-K Item 2.02, Ex. 99.1) — '
XB = 'Source: Forms 10-K / 10-Q (XBRL instance documents of each filing; FY2006–08 from the 10-K text) — '
LABEL_SOURCES = {
    'rv_search': XB + 'revenue note "Disaggregation of revenue" (dimensional facts), latest filing presenting the period (FY2017–18 annual from the FY2019 10-K; 2019 quarters from the 2020 10-Qs). Q4 = FY − (Q1 + Q2 + Q3).',
    'rv_cloud': XB + 'segment note (Google Cloud segment from Q4-20, recast to 2018); revenue type "Google Cloud" for FY2017 and 2019 quarters.',
    'so_svc': XB + 'segment note "Information about Segments" (operating income by segment, latest filing presenting the period: Google Services / Google Cloud / Other Bets / Alphabet-level from FY2018 annual and Q1-20 quarters).',
    'lso_google': XB + 'segment notes as first presented (Google / Other Bets FY2015–Q4-20; Google / Motorola Mobile FY2012–13).',
    'lg_web': XB + "MD&A revenue tables of each period's own 10-Q / 10-K (three-month columns; Q4 = FY − 9M); FY2006–08 from the 10-K text.",
    'is_rev': XB + 'consolidated statements of income, as originally reported (Q1–Q3 = 10-Q three months; Q4 = 10-K fiscal year − Q3 10-Q nine months). FY2006–08 from the 10-K text.',
    'is_oie': XB + 'OI&E note (interest income, interest expense, gains (losses) on equity securities — XBRL EquitySecuritiesFvNiGainLoss from 2018).',
    'sh_basic': XB + 'EPS note (Class A + B + C weighted shares; Class A diluted assumes conversion of Class B, + Class C diluted). Split-adjusted by filing date (x40 before 3-Apr-2014; x20 before 15-Jul-2022). Q4 = 4 × FY − Σ Q1–Q3 (derived).',
    'eps_pub': XB + 'diluted EPS per Class A share (= Class C except Q2-15 / FY2015, when the Class C adjustment payment raised Class C EPS); Q4 from the Q4 release three-month column.',
    'kpi_tac': REL + '"Revenues, Traffic Acquisition Costs (TAC) and number of employees" table.',
    'o_pub': REL + 'non-GAAP reconciliations "excluding the EC fine" (Q2-17, Q2-18, Q1-19, Q3-25) and "excluding severance and related and office space charges" (Q1-24 release, Q1-23 and Q1-24).',
    'ng_ni': REL + '"Reconciliations of non-GAAP results of operations" (Google Inc. / Alphabet, to the Q4-16 release): excludes stock-based compensation net of tax and discontinued operations.',
    'd_fine': 'Source: manual_items.json (items / company_adjusted), each with its Ex. 99.1 URL.',
    'cf_ni': XB + 'consolidated statements of cash flows (year to date; discrete quarters derived, Q4 = FY − 9M). Cash includes restricted cash from 2018.',
    'bs_cash': XB + 'consolidated balance sheets (period ends); shares outstanding = Class A + B + C (equity note / cover).',
    'fcf_pub': REL + '"Reconciliation from net cash provided by operating activities to free cash flow" (quarterly; three-month column).',
    'ppa_cons': 'Source: Q2-26 10-Q Note 8 (Wiz, Intersect).',
    'ds_total': 'Source: Q2-26 10-Q Note 6 (debt); FY2025 10-K debt note (maturities).',
    'mc_h': "Source: Q2-26 10-Q Note 11 (Stockholders' equity: common stock issuance, mandatory convertible preferred stock, capped calls, ATM program).",
    'dv_hdr': 'Source: Q2-26 10-Q Note 8 (pending GFiber divestiture).',
}
CELL_COMMENTS = {}
def comments_setup():
    q = lambda y, n: QC[(y, n)].c
    CELL_COMMENTS['is_ni'] = {q(2015, 1): 'Q1-15 as originally reported ($3,586m). The Q2-15 10-Q revised income tax for 2008–Q1-15 (Q1-15 $3,515m); quarters are kept as first reported, so Σ 2015 quarters exceeds FY2015 by $71m.',
                              A[2017]: 'FY2017 includes the $9.9bn Tax Act charge (Q4-17) and the $2.7bn EC fine (Q2-17).'}
    CELL_COMMENTS['so_basis'] = {A[2018]: 'Performance fees reclassified from operating expenses to other income (expense), net in 2019; 2018 segments recast ($27,524m vs $26,321m operating income as originally reported).'}
    CELL_COMMENTS['bs_restr'] = {q(2014, 1): 'Q1–Q3 2014 cash-flow statements include Motorola Mobile cash held for sale ($160m).'}
    CELL_COMMENTS['is_cor'] = {A[2012]: 'FY2012–13 cost of revenues tagged by segment only (Google / Motorola) — derived as total costs and expenses − R&D − S&M − G&A.'}
    CELL_COMMENTS['eps_pub'] = {q(2015, 2): 'Class A EPS $4.93 (pre-split; $0.2465 split-adjusted). Class C EPS was $6.43, including the $522m Class C adjustment payment.'}
    CELL_COMMENTS['is_oie'] = {A[2009]: 'FY2009 interest income and other, net not tagged in XBRL — derived as income before taxes − income from operations.'}
    CELL_COMMENTS['rv_cloud'] = {q(2026, 1): 'Includes Wiz from 11-Mar-2026 (revenue not disclosed; not stripped from the 1H26 growth used for Q3–Q4/26E).'}
    CELL_COMMENTS['fcf_chk'] = {q(2016, 3): 'Known difference: the Q3-16 release FCF ($7,291m) predates the 10-Q cash-flow revision; the 10-Q year-to-date statements give $7,262m. Check reads 0 within that documented $29m.'}

DEFS = {}
def defs_setup():
    q = lambda y, n: QC[(y, n)].c
    DEFS['o_def'] = {
        'C': ('2006–Q4-16: company non-GAAP operating income excluded stock-based compensation (model adjusted OI = GAAP; SBC treated as a cost)',
              'Google Inc. / Alphabet published non-GAAP operating income excluding SBC (and in 2008 equity impairments below OI) to the Q4-2016 release.'),
        q(2017, 2): ('Q2-17: operating income excluding the EC fine ($2,736m) — published; the model adds back company-reconciled items only', None),
        q(2018, 2): ('Q2-18: excluding the EC fines ($5,071m)', None),
        q(2019, 1): ('Q1-19: excluding the EC fine ($1,697m)', None),
        q(2023, 1): ('Q1-23 / Q1-24: excluding severance & office-space charges (published in the Q1-24 release)', None),
        q(2025, 3): ('Q3-25: excluding the EC fine ($3,457m; in G&A)', None),
    }
    DEFS['e_def'] = {
        'C': ('Model adjusted EPS (all periods): excl. discontinued operations, company-reconciled operating items (EC fines; severance & office space), equity-security gains / losses '
              '(2018+; 2008 impairment) and the Tax Act; company non-GAAP EPS (excl. SBC net of tax) published to Q4-16 — memo with check', None),
        q(2017, 2): ('2017–19: company EPS excluding the EC fine (Q2-17, Q2-18, Q1-19) / the Tax Act (Q4-17) — memo with check', None),
        A[2018]: ('2018+: ASU 2016-01 — equity-security fair-value changes through OI&E (excluded by the model)', None),
    }
def write_defs(ws, row):
    for c, (txt, cm) in DEFS[row.key].items():
        cell = ws[f'{c}{row.r}']; cell.value = txt
        cell.font = Font(name='Calibri', sz=10, b=True, i=True, color='FFC00000')
        if cm: comment(cell, cm)

def write_scen(ws):
    srcrow = {'hdr': 9, 'out_hdr': 11, 'pt_hdr': 15, 'out': 12, 'pt': 16, 'lev': 28, 'blk': 30, 'yr': 31}
    for spec_ in SCEN_ROWS:
        r, kind = spec_[0], spec_[1]
        for ci in range(CI(SL), CI(SNOTE) + 1):
            ws.cell(r, ci)._style = copy.copy(TM.cell(srcrow[kind], ci + TSHIFT)._style)
        if kind in ('hdr', 'out_hdr', 'pt_hdr'):
            _, _, a, b, c_, d_, e = spec_
            for col, v in zip([SL, SBULL, SBASE, SBEAR, SNOTE], [a, b, c_, d_, e]):
                if v is not None: ws[f'{col}{r}'] = v
        elif kind == 'out':
            _, _, lab, hi, mid, lo, note, nf = spec_
            ws[f'{SL}{r}'] = lab
            for col, v in zip([SBULL, SBASE, SBEAR], [hi, mid, lo]):
                ws[f'{col}{r}'] = v; ws[f'{col}{r}'].number_format = NF[nf]; set_color(ws[f'{col}{r}'], BLUE)
            ws[f'{SNOTE}{r}'] = note
        elif kind == 'pt':
            _, _, lab, _, val, _, note, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBASE}{r}'] = val; ws[f'{SBASE}{r}'].number_format = NF[nf]; set_color(ws[f'{SBASE}{r}'], BLUE); ws[f'{SNOTE}{r}'] = note
        elif kind == 'lev':
            _, _, lab, a, b, c_, note = spec_
            ws[f'{SL}{r}'] = lab
            for col, v in zip([SBULL, SBASE, SBEAR], [a, b, c_]):
                ws[f'{col}{r}'] = v; ws[f'{col}{r}'].number_format = NF['x']
            ws[f'{SNOTE}{r}'] = note
        elif kind == 'blk':
            _, _, lab, db, mid, de, note, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBULL}{r}'] = db; ws[f'{SBASE}{r}'] = mid; ws[f'{SBEAR}{r}'] = de; ws[f'{SNOTE}{r}'] = note
            for col in (SBULL, SBEAR):
                ws[f'{col}{r}'].number_format = '\\+#,##0;\\-#,##0' if nf == 'm' else '\\+0.0%;\\-0.0%'
        elif kind == 'yr':
            _, _, lab, hdr, base, clamp, _, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBASE}{r}'] = base
            if clamp:
                ws[f'{SBULL}{r}'] = f'=MAX(0,{SBASE}{r}+${SBULL}${hdr})'; ws[f'{SBEAR}{r}'] = f'=MAX(0,{SBASE}{r}+${SBEAR}${hdr})'
            else:
                ws[f'{SBULL}{r}'] = f'={SBASE}{r}+${SBULL}${hdr}'; ws[f'{SBEAR}{r}'] = f'={SBASE}{r}+${SBEAR}${hdr}'
            for col in (SBULL, SBASE, SBEAR): ws[f'{col}{r}'].number_format = NF[nf]

def write_header(ws):
    for r in range(1, 8):
        clone_row_style(ws, r, r, 2, CI(LASTC))
    ws[f'{NOTE}1'] = 'SCENARIO SWITCH ▼ (Bull / Base / Bear) — drives all scenario-linked forecast drivers'
    ws[f'{NOTE}2'] = 'Base'
    comment(ws[f'{NOTE}2'], 'Scenario toggle. Choose Bull / Base / Bear. Drives: Q3–Q4/26E revenue-growth and margin shifts and 2H26 buybacks, '
            '2027E–2030E revenue growth by type, Services / Cloud margins, Other Bets losses, Alphabet-level costs, capex, M&A spend and the net-debt target for 2027E+ buybacks '
            f'(scenario table {SL}:{SNOTE}).')
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type='list', formula1='"Bull,Base,Bear"', allow_blank=False); ws.add_data_validation(dv); dv.add(f'{NOTE}2')
    ws['B2'] = 'Alphabet Inc. (NASDAQ: GOOG) — 3-Statement Financial Model'
    ws['B3'] = ('US GAAP as originally reported · USD millions (per-share data in USD; shares in millions, split-adjusted) · Calendar fiscal year · '
                'Source: SEC EDGAR only — Google Inc. (CIK 1288776, to Q3-2015) and Alphabet Inc. (CIK 1652044) Forms 10-K / 10-Q (XBRL instance documents), '
                'Form 8-K earnings releases (Ex. 99.1); investor-relations releases are the same documents as filed on EDGAR. Capex outlook from the earnings call is UNVERIFIED.')
    ws['B4'] = ('Basis notes: statements as originally reported; revenue types and segments from the latest filing presenting the period (recasts: Cloud split Q4-20, performance fees to OI&E 2019, '
                'Google DeepMind to Alphabet-level 2023). Splits: 2-for-1 Class C dividend Apr-2014, 20-for-1 Jul-2022. Disc ops: Motorola Home 2012–13, Motorola Mobile 2014. '
                'Deals: Wiz (11-Mar-2026, $29.5bn), Intersect (10-Mar-2026, $5.9bn); GFiber divestiture pending (late 2026). Jun-2026: $30.5bn common + $19.0bn 6.25% mandatory convertible preferred. '
                'No company guidance: Q3–Q4/26E from the 1H26 run-rate with scenario shifts.')
    ws[f'{CAGR[0]}5'] = 'Forecast'
    for c in CAGR[1:]: ws[f'{c}5'] = 'Historical'
    ws['B6'] = '(USDm)'
    for col in COLS: ws[f'{col.c}6'] = col.label
    ws[f'{NOTE}6'] = 'Modelling Notes'
    for c, a, b in zip(CAGR, ['5Y CAGR', '5Y CAGR', '10Y CAGR', '19Y CAGR'], ["25-'30", "20-'25", "15-'25", "06-'25"]):
        ws[f'{c}6'] = a; ws[f'{c}7'] = b

def set_columns(ws):
    for k in list(ws.column_dimensions.keys()): del ws.column_dimensions[k]
    ws.column_dimensions['A'].width = TM.column_dimensions['A'].width
    ws.column_dimensions['B'].width = TM.column_dimensions['B'].width
    wid = {'A': 13.0, 'Q': 9.5703125, 'QE': 11.42578125, 'E26': 14.140625, 'AE': 13.0}
    for col in COLS:
        cd = ws.column_dimensions[col.c]; cd.width = wid[col.kind]
        if col.q:
            cd.outlineLevel = 1; cd.hidden = True
    for ci in range(CI(NOTE), CI(LASTC) + 1):
        ws.column_dimensions[L(ci)].width = TM.column_dimensions[L(ci + TSHIFT)].width or 13.0
    ws.sheet_format.outlineLevelCol = 1

def main():
    import openpyxl
    wb = openpyxl.load_workbook(TEMPLATE)
    if 'DCF_old' in wb.sheetnames: del wb['DCF_old']
    ws = wb['Model']
    blank_style = copy.copy(ws['A1']._style)
    for rowc in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for c in rowc:
            c.value = None; c.comment = None; c._style = copy.copy(blank_style)
    for k in list(ws.row_dimensions.keys()):
        rd = ws.row_dimensions[k]; rd.outlineLevel = 0; rd.hidden = False
    ws.data_validations.dataValidation = []
    scen_layout(V0)
    build_spec(V0)
    n1 = len(fw.SPEC)
    build_spec2(V0)
    for i, row in enumerate(fw.SPEC): row.spec2 = i >= n1
    last = number_rows(8)
    for r in range(1, last + 1): ws.row_dimensions[r].height = 15.0
    ws.row_dimensions[2].height = 21.0
    defs_setup(); comments_setup()
    write_header(ws)
    for row in fw.SPEC: write_row(ws, row)
    write_scen(ws)
    set_columns(ws)
    ws.freeze_panes = 'C8'
    ws.sheet_view.zoomScale = 70
    import other_sheets as osh
    osh.do_all(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    osh.inject_charts(OUT)
    print('rows', fw.SPEC[-1].r, 'saved', OUT)
    return wb

RECALC = os.environ.get('XLSX_RECALC', '/mnt/skills/public/xlsx/scripts/recalc.py')
def recalc(path, timeout=600):
    import subprocess
    out = subprocess.run(['python3', RECALC, path, str(timeout)], capture_output=True, text=True, cwd=os.path.dirname(RECALC))
    try:
        return json.loads(out.stdout)
    except Exception:
        return {'raw': out.stdout[-2000:], 'err': out.stderr[-2000:]}

def scenario_snapshots():
    import openpyxl
    import other_sheets as osh
    snaps = {}
    for scen in ('Bull', 'Bear'):
        p = os.path.join(BASE, f'GOOG_{scen}.xlsx')
        wb = openpyxl.load_workbook(OUT); wb['Model'][f'{NOTE}2'] = scen; wb.save(p)
        res = recalc(p); print(scen, res.get('status'), res.get('total_errors'))
        v = openpyxl.load_workbook(p, data_only=True)['Bull-Base-Bear']
        snaps[scen] = {(r, c): v[f'{c}{r}'].value for r in range(9, 51) for c in 'CDEFGHIJKL' if v[f'{c}{r}'].value is not None}
        os.remove(p)
    wb = openpyxl.load_workbook(OUT)
    osh.snapshot_values(wb, snaps['Bull'], snaps['Bear'])
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    osh.inject_charts(OUT)

NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
def inject_values(src_xlsx, recalc_xlsx):
    """Store cached formula results (from a LibreOffice-recalculated copy) in the openpyxl file, so the workbook reads correctly in any
    viewer; fullCalcOnLoad stays set so Excel recalculates on open. Same approach as the LECO / ESAB finishing step."""
    import zipfile, shutil, openpyxl
    from lxml import etree
    wbv = openpyxl.load_workbook(recalc_xlsx, data_only=True)
    zin = zipfile.ZipFile(src_xlsx)
    wbxml = etree.fromstring(zin.read('xl/workbook.xml'))
    rels = etree.fromstring(zin.read('xl/_rels/workbook.xml.rels'))
    rid2t = {r.get('Id'): r.get('Target') for r in rels}
    sheets = {}
    for s in wbxml.iter('{%s}sheet' % NS):
        rid = s.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
        sheets['xl/' + rid2t[rid].lstrip('/').replace('xl/', '')] = s.get('name')
    tmp = src_xlsx + '.tmp'
    zout = zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED)
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename in sheets:
            ws = wbv[sheets[item.filename]]
            root = etree.fromstring(data)
            for c in root.iter('{%s}c' % NS):
                f = c.find('{%s}f' % NS)
                if f is None: continue
                v = ws[c.get('r')].value
                old = c.find('{%s}v' % NS)
                if old is not None: c.remove(old)
                if v is None: continue
                ve = etree.SubElement(c, '{%s}v' % NS)
                if isinstance(v, bool):
                    c.set('t', 'b'); ve.text = '1' if v else '0'
                elif isinstance(v, (int, float)):
                    if 't' in c.attrib: del c.attrib['t']
                    ve.text = repr(float(v)) if isinstance(v, float) else str(v)
                elif isinstance(v, str) and v.startswith('#'):
                    c.set('t', 'e'); ve.text = v
                else:
                    c.set('t', 'str'); ve.text = str(v)
            data = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
        zout.writestr(item, data)
    zout.close(); zin.close()
    shutil.move(tmp, src_xlsx)

def store_values():
    import shutil
    p = os.path.join(BASE, 'GOOG_recalc.xlsx')
    shutil.copy(OUT, p)
    res = recalc(p); print('recalc', res.get('status'), res.get('total_errors'))
    inject_values(OUT, p)
    os.remove(p)

if __name__ == '__main__':
    main()
    if '--snap' in sys.argv: scenario_snapshots()
    if '--recalc' in sys.argv: store_values()
