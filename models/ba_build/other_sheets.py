"""Rewire the Bull-Base-Bear, DCF and Charts tabs of the MLM template (current layout: DCF with live valuation date, fade period,
implied ROIC, target-IRR entry price, reverse DCF and Mauboussin EV/NOPAT; Bull-Base-Bear with buyback rows) to the BA Model rows /
columns; keep the template chart XML. The template's DCF_old tab is not carried.

Column mapping: the template's 2025A base column becomes BA 2026E (1H actual + Q3/Q4 estimates; the valuation date, 10-Oct-2026, is in
Q4) and the template's 2026E-2030E become BA 2027E-2031E, so the DCF's explicit years start in 2027E and net debt is taken at
31-Dec-2026 (2026E cash flows are inside year-end net debt, not double counted)."""
import re, zipfile, shutil
from openpyxl.utils import get_column_letter as L
from fw import R, TEMPLATE, ANN, NOTE

# template Model row -> BA row key
MAP = {158: 'is_rev', 164: 'r_gain', 166: 'is_efo', 182: 'is_dda', 186: 'sh_dil', 211: 'r_adj', 212: 'r_adj_m', 219: 'r_ebit', 221: 'r_adj',
       229: 'e_coreni', 231: 'e_eps', 238: 'gm_del', 245: 'gm_ebitda_m', 247: 'gm_op_m', 260: 'cf_wc', 262: 'cf_capex', 263: 'cf_acq',
       280: 'fcf', 282: 'fcf_m', 290: 'bb_px', 291: 'bb_cash', 292: 'bb_sh', 331: 'v_dl', 408: 'ra_ic', 411: 'ra_roic', 422: 'rn_ronta',
       428: 'lv_nd', 433: 'v_px', 434: 'v_sh'}
CMAP = {'BV': ANN[2026].c, 'CA': ANN[2027].c, 'CB': ANN[2028].c, 'CC': ANN[2029].c, 'CD': ANN[2030].c, 'CE': ANN[2031].c, 'CF': NOTE}
REF = re.compile(r"Model!(\$?)([A-Z]{1,3})(\$?)(\d+)")
HIST = [ANN[y].c for y in range(2006, 2026)]

def remap(f):
    def rep(m):
        c, r = m.group(2), int(m.group(4))
        c2 = CMAP.get(c, c)
        r2 = R[MAP[r]] if r in MAP else r
        if c == 'CF' and r == 2: r2 = 2
        return f'Model!{m.group(1)}{c2}{m.group(3)}{r2}'
    return REF.sub(rep, f)

def hist_stat(f):
    """LT historical AVERAGE / MIN / MAX over the template's annual columns -> BA FY2006-FY2025 annual columns."""
    m = re.match(r'=IFERROR\((AVERAGE|MIN|MAX)\(Model!C(\d+),', f)
    if not m: return None
    fn, r = m.group(1), int(m.group(2))
    rr = R[MAP[r]]
    return f'=IFERROR({fn}({",".join(f"Model!{c}{rr}" for c in HIST)}),"")'

def fy_dates(f):
    """Template explicit years 2026E-2030E -> BA 2027E-2031E (calendar year ends)."""
    f = re.sub(r'DATE\((\d{4}),12,31\)', lambda m: f'DATE({int(m.group(1)) + 1},12,31)', f)
    f = re.sub(r'DATE\((\d{4})-1,12,31\)', lambda m: f'DATE({int(m.group(1)) + 1}-1,12,31)', f)
    return f

TEXT = {
    'Bull-Base-Bear': {
        'B2': 'The Boeing Company (BA) — Bull / Base / Bear Case P&L Summary',
        'B3': ('Core figures (company non-GAAP definitions: core operating earnings, core EPS) and model adjusted EBITDA (core operating earnings + D&A) · USD millions · '
               'Live panel linked to Model — toggle scenario via Model!CG2. Bull and Bear panels are value snapshots generated from the Model with the switch set to each case. '
               '2026E = 1H26 actual + Q3/Q4 calibrated to 2026 delivery, margin and FCF end-points (UNVERIFIED call statements).'),
        'G26': 30, 'H26': 30, 'S26': 34, 'T26': 34, 'AD26': 22, 'AE26': 22, 'AF26': 22,
    },
    'DCF': {
        'B2': 'DCF VALUATION — THE BOEING COMPANY (BA)',
        'C5': 0.0425, 'D5': 'Assumption (~US 10Y Treasury yield, Oct-2026; government rate sources not reachable from the build environment); update to market',
        'D6': 'Assumption: US implied ERP ~4.5–5.5%',
        'C7': 1.20, 'D7': 'Assumption: commercial aerospace OEM with high operating leverage, 2019–24 losses and BBB- / Baa3-area credit; A&D primes ~0.8–1.0, commercial aero suppliers ~1.2–1.4',
        'C9': 0.060, 'D9': 'Interest and debt expense ≈5.2% of debt (Q2-26); 2024 notes issued at 6.3–7.1%; marginal pre-tax cost ~6%',
        'C10': 0.21, 'D10': 'US statutory rate (cash taxes lower near term: valuation allowance $9.8bn and NOLs at YE25)',
        'C12': 0.20, 'D12': 'Net debt ~$26bn + pension / retiree health ~$6bn vs ~$150bn market cap ($185 × ~833m diluted incl. preferred conversion shares)',
        'D15': 'Periods = (31-Dec of the year − today) / 365.25; cash flows assumed at year end; 2026E (Q4 in progress at the valuation date) is the base year',
        'B18': 'Adj. EBIT (core operating earnings) less gains on dispositions',
        'B21': 'plus D&A (depreciation + amortization of acquired intangibles)',
        'B23': '± Δ Working capital (cash-flow statement; + = release)',
        'B22': '− Capex (PP&E) and acquisitions (M&A lever, zero in all cases)',
        'B26': 'Fade-period growth (2032E–36E; 6.0% → 4.0% linear)', 'I26': 0.06, 'M26': 0.04, 'C27': 0.03,
        'D32': ("Boeing's core operating earnings are full-cost: after D&A, share-based plans, reach-forward losses and abnormal production costs; only the FAS/CAS service cost "
                'adjustment is removed (non-cash pension recovery; also excluded from UFCF). D&A is added back in row 21. Row 18 deducts gains on dispositions (2025 Digital Aviation '
                'Solutions $9.6bn) so NOPAT is operating only. 401(k) treasury-share contributions (~$1.6bn a year) are a real compensation cost: they sit inside core earnings and '
                'are funded with shares, not cash, so UFCF is not adjusted for them (their dilution is in the share count). Working capital (row 23) carries the inventory / advances '
                'unwind that drives Boeing\'s cash recovery.'),
        'D43': 'Opening IC = Model 2026E book IC (equity + net debt at 31-Dec-2026), rolled forward with capex + M&A − D&A + ΔWC.',
        'B48': 'Less: Net debt + pension & retiree health + NCI (Model 2026E year end)',
        'D48': 'Net debt = debt − cash − short-term investments; unfunded pension and retiree health liabilities are debt-like (pre-tax). The mandatory convertible preferred is in the diluted share count (converts Oct-2027), not in net debt.',
        'B50': 'Diluted shares (m, Model 2026E period end, incl. preferred conversion shares)',
        'D76': 'Rationale (Boeing characteristics)',
        'C78': 0.10, 'D78': 'Applied from 2029E (after the recovery years): model NOPAT growth 2029E–31E ~15% a year, then the 6% → 4% fade; ~10% a year over the stage-1 period (assumption)',
        'C79': 0.25, 'D79': 'Duopoly in large commercial jets with ~$600bn commercial backlog; returns on incremental capital are high at normal production rates (2015–18 ROIC well above WACC) but development programs and quality failures can destroy value for years',
        'C80': 12, 'D80': 'Very long product cycles (737 / 777 / 787 franchises, 20–30-year services annuities) and certification barriers sustain a long competitive-advantage period; risks: regulator oversight, execution, a new-airplane program in the 2030s',
        'B90': 'VALUATION  —  Applied to 2029E NOPAT',
        'B91': '2029E Adj. EBIT (core operating earnings) less gains on dispositions',
        'B93': '2029E NOPAT',
        'B95': 'Implied Enterprise Value (2029E basis)',
        'B96': 'Less: 2028E year-end net debt + pension + NCI (Model)',
        'B98': 'Diluted shares 2029E (m, Model)',
        'B99': 'Implied share price 2029 ($)',
        'B106': 'Plus: 2026E year-end net debt + pension + NCI',
        'B108': 'BA market-implied EV/NOPAT (on 2029E NOPAT)',
    },
    'Charts': {
        'B2': 'The Boeing Company — Operating Performance (FY2012–FY2025) and Model Forecast (2026E–2031E)',
        'B3': ('USDm · fiscal years ended 31 December · Linked to Model. Bars = commercial deliveries y/y and BCA revenue per delivery y/y. Adjusted EBITDA = model definition '
               '(core operating earnings + D&A); core operating earnings = company definition (published from FY2011; earlier years rebuilt / = GAAP, flagged in the Model). '
               'ROIC / RONTA on core operating earnings excl. divestiture gains; n/m when invested capital or equity is negative.'),
        'B8': 'Commercial deliveries y/y', 'B9': 'BCA revenue per delivery y/y', 'B12': 'Adjusted EBITDA (model)', 'B13': 'Earnings from operations (GAAP)',
        'B14': 'Adj. EBITDA Margin %', 'B15': 'Operating Margin (GAAP) %',
    },
}
BBB_ROWS = {9: 'Total revenues', 10: '   Revenues y/y %', 11: '   Commercial deliveries y/y %', 12: 'Adjusted EBITDA (model)',
            15: 'Earnings from operations (GAAP)', 16: '   Earnings from operations y/y %', 17: '   Operating margin (GAAP) %', 19: 'Core earnings (for core EPS)',
            21: 'Diluted shares (m)', 23: 'Core EPS ($)', 26: 'P / E (x) — assumed (core EPS)', 35: 'Adjusted EBITDA (leverage basis)',
            36: '   Adjusted EBITDA % of revenues', 48: '(+) Net debt, pension & retiree health, NCI'}
CHART_TITLES = {
    'Aggregates Revenue Drivers: Shipment Volume vs. Price (ASP), y/y %': 'Commercial Airplanes Drivers: Deliveries vs. Revenue per Delivery, y/y %',
    'Adjusted EBITDA &amp; Earnings from Operations (USDm)': 'Adjusted EBITDA (model) &amp; Earnings from Operations (USDm)',
}

def build_charts(ws):
    for i, y in enumerate(range(2012, 2032)):
        c = L(3 + i); mc = ANN[y].c
        ws[f'{c}7'] = y if y <= 2025 else f'{y}E'
        ws[f'{c}8'] = f'=IFERROR(Model!{mc}{R["del_g"]}+0,0)'
        ws[f'{c}9'] = f'=IFERROR(Model!{mc}{R["bca_rpd_g"]}+0,0)'
        ws[f'{c}11'] = f'={c}7'
        ws[f'{c}12'] = f'=Model!{mc}{R["r_adj"]}'
        ws[f'{c}13'] = f'=Model!{mc}{R["is_efo"]}'
        ws[f'{c}14'] = f'=Model!{mc}{R["gm_ebitda_m"]}'
        ws[f'{c}15'] = f'=Model!{mc}{R["gm_op_m"]}'
        ws[f'{c}17'] = f'={c}7'
        ws[f'{c}18'] = f'=IFERROR(Model!{mc}{R["ra_roic"]}+0,NA())'
        ws[f'{c}19'] = f'=IFERROR(Model!{mc}{R["rn_ronta"]}+0,NA())'
        ws[f'{c}20'] = f'=IFERROR(Model!{mc}{R["gm_ebitda_m"]}+0,NA())'
        ws[f'{c}21'] = f'=IFERROR(Model!{mc}{R["fcf_m"]}+0,NA())'

def do_all(wb):
    for name in ('Bull-Base-Bear', 'DCF', 'Charts'):
        ws = wb[name]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and 'Model!' in c.value:
                    h = hist_stat(c.value) if name == 'Bull-Base-Bear' else None
                    c.value = h or remap(c.value)
                if isinstance(c.value, str) and 'DATE(' in c.value:
                    c.value = fy_dates(c.value)
        for addr, t in TEXT.get(name, {}).items():
            ws[addr] = t
    b = wb['Bull-Base-Bear']
    for r, t in BBB_ROWS.items():
        for c in ('B', 'N', 'Z'): b[f'{c}{r}'] = t
    for c0 in ('C', 'O', 'AA'):
        i = b[f'{c0}7'].column
        b.cell(7, i).value = '2026E'
        for k in range(1, 6): b.cell(7, i + k).value = f'{2026 + k}E'
        b.cell(7, i + 6).value = "26-'31 CAGR"
    from openpyxl.comments import Comment
    cm = Comment('P/E on core EPS (company definition). At $185 Boeing trades on depressed 2026E–27E core EPS (loss / break-even), so the P/E is applied to 2029E–31E '
                 'earnings only. Exit multiples are assumptions: Base 30x, Bull 34x, Bear 22x (2015–18 core P/E ~15–25x at peak margins; recovery-stage premium). ', 'BA model')
    cm.width, cm.height = 420, 120; b['N26'].comment = cm
    cm2 = Comment(f'Snapshot (values) of Model rows {R["bb_cash"]}, {R["bb_px"]} and {R["bb_sh"]} (Share Buyback Schedule) with Model!{NOTE}2 = Bull / Bear. '
                  'Base panel (D40:H42) is live-linked.', 'BA model')
    cm2.width, cm2.height = 420, 90; b['P40'].comment = cm2
    b['J2'] = 'Share price ($):'
    for c in ('B', 'N', 'Z'): b[f'{c}29'] = 'IRR (to 31-Dec of Y-1)'
    for c in ('B', 'N', 'Z'):
        if b[f'{c}40'].value: b[f'{c}40'] = '   Share repurchases ($m; 2026E = 2H)'
    d = wb['DCF']
    for i, col in enumerate('DEFGHIJKLM'):
        d[f'{col}17'] = f'{2027 + i}E'
    a26, a27 = ANN[2026].c, ANN[2027].c
    d['D22'] = f'=Model!{a27}{R["cf_capex"]}+Model!{a27}{R["cf_acq"]}'
    d['D34'] = f'=Model!{a26}{R["ra_ic"]}'
    d['C48'] = f'=Model!{a26}{R["lv_nd"]}+Model!{a26}{R["v_dl"]}'
    d['C50'] = f'=Model!{a26}{R["v_sh"]}'
    d['C52'] = f'=Model!{a26}{R["v_px"]}'
    d['P17'] = 'CAGR 27–31E'; d['Q17'] = 'CAGR 27–36E'
    build_charts(wb['Charts'])

def snapshot_values(wb, bull, bear):
    ws = wb['Bull-Base-Bear']
    cols = 'CDEFGHIJKL'
    for snap, off in ((bull, 'OPQRSTUVWX'), (bear, ['AA', 'AB', 'AC', 'AD', 'AE', 'AF', 'AG', 'AH', 'AI', 'AJ'])):
        for (r, c), v in snap.items():
            tgt = f'{off[cols.index(c)]}{r}'
            cur = ws[tgt].value
            if isinstance(cur, str) and cur.startswith('='): continue
            if r in (7, 26): continue
            ws[tgt] = v

def inject_charts(path):
    """Replace openpyxl-rewritten chart parts with the template chart XML (combo charts, colours, layout), retitled, ranges kept
    (Charts!C:V = FY2012-FY2031E) and cached values removed so Excel redraws from the BA cells."""
    tz = zipfile.ZipFile(TEMPLATE)
    tmpl = {n: tz.read(n) for n in tz.namelist() if n.startswith('xl/charts/chart')}
    tmp = path + '.tmp'
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in tmpl:
                t = tmpl[item.filename].decode('utf-8')
                for a, b2 in CHART_TITLES.items(): t = t.replace(a, b2)
                t = re.sub(r'<c:numCache>.*?</c:numCache>', '', t, flags=re.S)
                t = re.sub(r'<c:strCache>.*?</c:strCache>', '', t, flags=re.S)
                data = t.encode('utf-8')
            zout.writestr(item, data)
    shutil.move(tmp, path)
