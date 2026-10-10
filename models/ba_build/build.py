"""Build BA_Model.xlsx (The Boeing Company) from the MLM template (same layout, formatting and process as the prior models)."""
import sys, copy, os, json
from openpyxl.styles import Font
from fw import *
import fw
from spec import build_spec, scen_layout, SCEN_ROWS, SC
from spec2 import build_spec2

OUT = os.path.join(os.path.dirname(BASE), 'BA_Model.xlsx')
G26 = MAN['guidance_2026']
EF = MAN['edgar_facts']

# ====================================================================== inputs (V0)
def wc_inputs():
    q = DATA['2025']['x']
    rev, cogs = q['rev'], q['cogs']
    return {'dso': round(q['bs_ar'] / rev * 365), 'dpo': round(q['bs_ap'] / cogs * 365), 'adv': round(q['bs_adv'] / rev, 3), 'ocl': round(q['bs_ocl'] / rev, 3)}
WC = wc_inputs()
WC['note'] = (f'Forecast inputs at year-end 2025 levels on FY2025 revenues / cost of sales (receivable days {WC["dso"]}d, payable days {WC["dpo"]}d, advances '
              f'{WC["adv"]:.1%} of revenues in 2026E, accrued & other current liabilities {WC["ocl"]:.1%}); year ends are compared with year ends.')
UNV = 'UNVERIFIED (earnings-call statement, not in an EDGAR filing)'

V0 = {
    'out_hdr': '2026 outlook — Q4-25 call (27-Jan-2026), reaffirmed Q2-26 call (28-Jul-2026) — ' + UNV,
    'out_note': ('Boeing publishes no numeric outlook in its 8-K releases (since 2019). Call statements: free cash flow $1–3bn (incl. ~$1bn Spirit drag and a ~$700m DOJ payment in Q3), '
                 '~500 737 and 90–100 787 deliveries, BDS margin ~2.5% (after the $280m VC-25B charge), BGS ~18%, BCA margins negative in 2026. Sources in manual_items.json.'),
    'outlook': [
        ('out_d737', '  2026 737 deliveries (units)', (510.0, 500.0, 480.0), 'm', 'Target ~500 (' + UNV + '); Bull / Bear ±. 1H26: 243. Drives Q3/Q4-26E 737 deliveries.'),
        ('out_d787', '  2026 787 deliveries (units)', (100.0, 95.0, 90.0), 'm', 'Target 90–100 (' + UNV + '). 1H26: 40.'),
        ('out_bdsm', '  2026 BDS operating margin', (0.030, 0.025, 0.020), 'pct', '~2.5% incl. the Q2-26 VC-25B charge (' + UNV + '). Drives the Q3/Q4-26E BDS margin calibration.'),
        ('out_bgsm', '  2026 BGS operating margin', (0.185, 0.180, 0.175), 'pct', '~18% (' + UNV + '). Drives the Q3/Q4-26E BGS margin calibration.'),
        ('out_fcf', '  2026 free cash flow ($m)', (3000.0, 2000.0, 1000.0), 'm', '$1–3bn (' + UNV + '); 1H26 −$823m. Solved through 2026E year-end inventories.'),
        ('out_bcam', '  2H26 BCA operating margin (assumption, not guidance)', (-0.005, -0.020, -0.040), 'pct',
         'Assumption: 1H26 −4.2% (Q2 −2.7%); company indicated negative BCA margins in 2026 (' + UNV + ').'),
        ('bb26', '  2H26 share repurchases ($m)', (0.0, 0.0, 0.0), 'm', 'No repurchases since Q1-2020; deleveraging priority.'),
    ],
    'points': [
        ('pt_unal', '  Q3/Q4-26E unallocated items, core basis ($m per quarter)', -480.0, 'm1', '1H26 −$978m (Q1 −$348m, Q2 −$630m). Assumption.'),
        ('pt_fascas', '  Q3/Q4-26E FAS/CAS service cost adjustment ($m per quarter)', 155.0, 'm1', 'Q1 and Q2-26 $155m each.'),
        ('pt_dda', '  Q3/Q4-26E depreciation & amortization ($m per quarter)', 585.0, 'm1', '1H26 $1,169m.'),
        ('pt_nonop', '  Q3/Q4-26E non-operating pension & postretirement income ($m per quarter)', -64.0, 'm1', 'Q1 −$65m, Q2 −$64m (expense).'),
        ('pt_intinc', '  Q3/Q4-26E interest & other income ($m per quarter)', 150.0, 'm1', 'Q2-26 $143m on ~$20bn cash & investments. Assumption.'),
        ('pt_int', '  Q3/Q4-26E interest and debt expense ($m per quarter)', 590.0, 'm1', 'Q2-26 $600m; debt $45.9bn at 30-Jun-2026. Assumption.'),
        ('pt_etr', '  Q3/Q4-26E effective tax rate', 0.10, 'pct', 'Valuation allowance; 2025 GAAP rate 15.1%. Assumption.'),
        ('pt_capex', '  2026 capital expenditures ($m)', 4200.0, 'm1', '1H26 $2,008m; not guided in EDGAR filings. Assumption.'),
        ('pt_pref', '  Q3/Q4-26E preferred dividends ($m per quarter)', 86.25, 'm1', '$5,750m × 6.00% ÷ 4.'),
    ],
    'lev': (1.0, 0.5, 0.0),
    'lev_note': ('Buybacks resume only once net debt / adjusted EBITDA (model) is below the target; Boeing\'s stated priority is debt reduction to protect its investment-grade rating '
                 '(10-K). Net debt ~$26bn at 30-Jun-2026 (≈7x 2026E adjusted EBITDA). Targets are assumptions: Bull 1.0x / Base 0.5x / Bear 0.0x (higher target = earlier, larger buybacks).'),
    'blocks': [
        ('D737', 'D737 (737 deliveries, units)', 30.0, -60.0, [560.0, 610.0, 640.0, 650.0, 660.0],
         'Rate 47/month from mid-2026, 52/month planned (Everett North Line); 737-7/-10 first deliveries 2027 (Q2-26 release). 2018 peak 580.', 'm', True),
        ('D787', 'D787 (787 deliveries, units)', 10.0, -15.0, [110.0, 125.0, 135.0, 140.0, 140.0], 'Rate 8/month (Q4-25 release) rising toward 10; 2019 peak 158.', 'm', True),
        ('D777', 'D777 (777 / 777X deliveries, units)', 5.0, -15.0, [35.0, 55.0, 65.0, 70.0, 70.0], '777X certification / first delivery assumed 2027; 2025: 35. Assumption.', 'm', True),
        ('D767', 'D767 (767 commercial deliveries, units)', 0.0, 0.0, [25.0, 5.0, 0.0, 0.0, 0.0], '767F production ends 2027 (assumption); 2025: 30.', 'm', True),
        ('BCA_RPD', 'BCA_RPD (BCA revenue per delivery y/y)', 0.01, -0.01, [0.04, 0.03, 0.03, 0.03, 0.03],
         'Price escalation ~2–3% + widebody / 777X mix; 1H26 $66.7m per delivery (2025 $69.2m; 2018 $75.3m).', 'pct', False),
        ('BCA_M', 'BCA_M (BCA operating margin)', 0.02, -0.03, [0.03, 0.06, 0.09, 0.11, 0.12],
         'Positive margins from 2027 (company indication, UNVERIFIED); 2017–18 margins 9.4% / 13.0% at 760–806 deliveries; Spirit integration and abnormal costs fade.', 'pct', False),
        ('BDS_G', 'BDS_G (BDS revenue growth)', 0.02, -0.02, [0.05, 0.05, 0.045, 0.04, 0.04], '1H26 +17% y/y; backlog $85bn (Q2-26); US / allied defense budgets.', 'pct', False),
        ('BDS_M', 'BDS_M (BDS operating margin)', 0.01, -0.02, [0.06, 0.075, 0.085, 0.09, 0.09],
         'Fixed-price development programs maturing (KC-46A, T-7A, MQ-25 Milestone C, VC-25B); 2013–17 margins 9–11%; high-single-digit long-term (UNVERIFIED).', 'pct', False),
        ('BGS_G', 'BGS_G (BGS revenue growth)', 0.01, -0.02, [0.05, 0.06, 0.06, 0.05, 0.05], 'Ex-Digital Aviation Solutions base; fleet utilisation, parts and government services; backlog $33bn.', 'pct', False),
        ('BGS_M', 'BGS_M (BGS operating margin)', 0.005, -0.01, [0.180, 0.180, 0.180, 0.180, 0.180], '2023–24 18–19%; 1H26 18.1%.', 'pct', False),
        ('MNA', 'M&A_spend (acquisitions, $m)', 0.0, 0.0, [0.0, 0.0, 0.0, 0.0, 0.0], 'None assumed (deleveraging after Spirit).', 'm', True),
    ],
    'un_rev_pct': -0.002,
    'un_pct': {2027: -0.020, 2028: -0.019, 2029: -0.018, 2030: -0.017, 2031: -0.016},
    'fascas': {2027: 600.0, 2028: 550.0, 2029: 500.0, 2030: 450.0, 2031: 400.0},
    'dda_pct': {2027: 0.024, 2028: 0.024, 2029: 0.023, 2030: 0.023, 2031: 0.023},
    'amort_sched': {2027: 182.0, 2028: 155.0, 2029: 150.0, 2030: 144.0, 2031: 140.0},
    'amort_note': 'FY2025 10-K Note 4: 2026 $197m, 2027 $182m, 2028 $155m, 2029 $150m, 2030 $144m; 2031 $140m assumed. Spirit intangibles ($109m, ~5-year life) included.',
    'amort_2h': 98.5,
    'ga_pct': {2027: 0.058, 2028: 0.057, 2029: 0.056, 2030: 0.055, 2031: 0.055},
    'rd_pct': {2027: 0.039, 2028: 0.038, 2029: 0.037, 2030: 0.037, 2031: 0.037},
    'sbc_pct': 0.005, 'k401_g': 0.03, 'nonop': -250.0, 'opinv': 0.0,
    'etr': {2027: 0.10, 2028: 0.12, 2029: 0.15, 2030: 0.18, 2031: 0.20},
    'pref_div': {2027: 272.2, 2028: 0.0, 2029: 0.0, 2030: 0.0, 2031: 0.0}, 'mcps_frac': round(288 / 365, 4),
    'dps': 0.0, 'pens_2h': -100.0, 'pens_cf': -300.0, 'tax_old': 0.35,
    'capex_pct': {2027: 0.042, 2028: 0.040, 2029: 0.037, 2030: 0.035, 2031: 0.033},
    'dio': {2027: 370, 2028: 355, 2029: 340, 2030: 325, 2031: 315},
    'advp': {2027: 0.64, 2028: 0.62, 2029: 0.60, 2030: 0.58, 2031: 0.57},
    'iss_2h': 0.3, 'iss_ae': 1.0, 'dil': 6.0, 'dil_e26': 5.0,
    'repay': {2026: 0.0, 2027: 4403.0, 2028: 2739.0, 2029: 2508.0, 2030: 5274.0, 2031: 3000.0},
    'mat': {2026: 4403.0, 2027: 2739.0, 2028: 2508.0, 2029: 5274.0, 2030: 3000.0, 2031: 3000.0},
    'refi': {2027: 0.0, 2028: 0.5, 2029: 0.75, 2030: 1.0, 2031: 1.0},
    'repay_note': ('FY2025 10-K: scheduled principal 2026 $8,351m (repaid in 1H26: debt repayments $8,376m, so 2H26 = 0), 2027 $4,403m, 2028 $2,739m, 2029 $2,508m, '
                   '2030 $5,274m; 2031 $3,000m assumed (≈$30.8bn due after 2030). Refinancing % is an assumption (deleveraging first).'),
    'debt_note': EF['debt'],
    'min_cash': 8000.0,
    'r_notes': {2027: 0.052, 2028: 0.052, 2029: 0.052, 2030: 0.053, 2031: 0.053},
    'r_notes_note': 'Q2-26 interest and debt expense $600m on $45.9bn (≈5.2% annualised); 2024 notes issued at 6.3–7.1%. Assumption.',
    'yield': 0.035, 'yield_note': '2025 interest & other income ≈$0.9bn on ~$30bn average cash & investments (~3%). Assumption.',
    'wc': WC,
    'px': MAN['price']['px'],
    'px_note': 'VALUATION — share price input $185.00 (user-provided, 10-Oct-2026); update to market',
    'mna_note': 'Scenario lever M&A_spend (zero in all cases).',
}

# ====================================================================== compute engine
def hist_value(row, col):
    kw = row.kw; path = kw.get('data')
    if not path: return None
    if kw.get('ppa'):
        v = MAN['ppa'].get(pkey(col), {}).get(path.split('_', 1)[1])
        return v if kw.get('textdata') else num(v)
    v = xget(pkey(col), path)
    return v if kw.get('textdata') else num(v)

def first(kw, keys):
    for k in keys:
        if k in kw: return k
    return None

def compute(row, col):
    kw = row.kw; k = col.kind
    CTX.col = col
    if kw.get('ann') and col.q: return None, False
    if k in HISTK:
        v = hist_value(row, col)
        if v is not None: return v, True
        for key in ('hist', 'all'):
            if key in kw:
                return kw[key](), False
        return None, False
    if k == 'QE':
        if row.spec2: return None, False
        if 'qe_in' in kw:
            v = kw['qe_in']; return (v[col.q] if isinstance(v, dict) else v), True
        key = first(kw, ('qe', 'fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    if k == 'E26':
        if 'e26_in' in kw: return kw['e26_in'], True
        if 'e26' in kw: return kw['e26'](), False
        if kw.get('avg') and ('qe' in kw or 'qe_in' in kw): return f'=AVERAGE({Q1}{row.r}:{QE}{row.r})', False
        if kw.get('flow') and not row.spec2 and any(x in kw for x in ('qe', 'qe_in', 'fc')): return S26(row.key), False
        key = first(kw, ('fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    if k == 'AE':
        if 'ae_in' in kw:
            v = kw['ae_in']; return (v[col.year] if isinstance(v, dict) else v), True
        key = first(kw, ('ae', 'fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    return None, False

# ====================================================================== writers
CAGR_G = {'is_rev', 'g_rev', 'bca_rev', 'bds_rev', 'bgs_rev', 'del_tot', 'del_737', 'del_787', 'bca_rpd', 'o_core', 'g_core', 'r_adj', 'r_ebitda', 'is_efo',
          'd_dda', 'e_coreni', 'cf_cfo', 'bs_ta', 'sh_dil', 'is_gp', 'is_ga', 'is_rd', 'bca_oi', 'bds_oi', 'bgs_oi'}
CAGR_A = {'bca_m', 'bds_m', 'bgs_m', 'g_core_m', 'o_core_m', 'r_adj_m', 'is_gm', 'is_efo_m', 'r_ebit_m', 'gm_rev', 'gm_bca', 'gm_bds', 'gm_bgs', 'gm_gm', 'gm_ebitda_m',
          'gm_op_m', 'gm_net_m', 'gm_ga', 'gm_rd', 'gm_core_m', 'fcf_m', 'capex_pct', 'ra_roic', 'rn_ronta', 'd_ga', 'd_rd', 'd_dda_pct', 'd_sbc_pct', 'wc_nwc_p',
          'un_pct', 'del_g', 'bca_rpd_g', 'bds_g', 'bgs_g'}
A = {y: ANN[y].c for y in ANN}

def style_cagr(ws, r, on):
    src = 16 if on else 12
    for c in CAGR:
        ws[f'{c}{r}']._style = copy.copy(TM[f'{L(CI(c) + TSHIFT)}{src}']._style)

def write_row(ws, row):
    r = row.r; arch = ARCH[row.arch]
    clone_row_style(ws, arch, r)
    ws.cell(r, 2).value = row.label or None
    if row.kw.get('note'): ws[f'{NOTE}{r}'] = row.kw['note']
    style_cagr(ws, r, row.key in CAGR_G or row.key in CAGR_A)
    if row.arch in ('section', 'block', 'blank', 'sub') and not row.kw.get('data'):
        return
    if row.key in DEFS:
        write_defs(ws, row); return
    nf = NF[row.nf]
    for col in COLS:
        v, inp = compute(row, col)
        if v is None: continue
        cell = ws[f'{col.c}{r}']
        cell.value = v
        if not row.kw.get('textdata'): cell.number_format = nf
        if inp:
            set_color(cell, BLUE, bold=True if col.kind in FCK else None)
            lk = row.kw.get('labkey')
            if lk and col.kind in HISTK and isinstance(v, (int, float)) and abs(v) > 1e-9:
                t = xget(pkey(col), lk)
                if t: comment(cell, f'Company items ({pkey(col)}): {t}', 380, 110)
        elif row.arch == 'growth' and col.kind in HISTK:
            set_color(cell, GRAY)
        else:
            set_color(cell, BLACK)
    rc = CELL_COMMENTS.get(row.key, {})
    for c, t in rc.items(): comment(ws[f'{c}{r}'], t)
    if row.key in CAGR_G:
        def cg(a, b, n): return f'=IFERROR(IF(AND({A[a]}{r}>0,{A[b]}{r}>0),({A[a]}{r}/{A[b]}{r})^(1/{n})-1,"n/m"),"n/m")'
        ws[f'{CAGR[0]}{r}'] = cg(2031, 2026, 5)
        ws[f'{CAGR[1]}{r}'] = cg(2025, 2020, 5)
        ws[f'{CAGR[2]}{r}'] = cg(2025, 2015, 10)
        ws[f'{CAGR[3]}{r}'] = cg(2025, 2006, 19)
        for c in CAGR: ws[f'{c}{r}'].number_format = '0%;\\(0%\\);"n/m"'
    elif row.key in CAGR_A:
        ws[f'{CAGR[0]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2027, 2032))}),"n/m")'
        ws[f'{CAGR[1]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2021, 2026))}),"n/m")'
        ws[f'{CAGR[2]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2016, 2026))}),"n/m")'
        ws[f'{CAGR[3]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2006, 2026))}),"n/m")'
        for c in CAGR: ws[f'{c}{r}'].number_format = NF['pct']
    t = LABEL_SOURCES.get(row.key)
    if t: comment(ws.cell(r, 2), t)

REL = 'Source: Boeing earnings releases (Form 8-K Item 2.02, Ex. 99.1) — '
LABEL_SOURCES = {
    'del_737': REL + 'deliveries table (commercial airplanes by program; footnote markers ignored). 737 = 737 Next-Generation to 2017, 737 NG + MAX from 2017.',
    'bca_rev': REL + 'segment revenues and earnings from operations (as originally reported; Global Services formed Q3-17 with services moved out of BCA / BDS; 2019 realignment not recast).',
    'bds_l_bma': REL + 'segment table sub-segments (Integrated Defense Systems / Defense, Space & Security).',
    'un_core': REL + 'unallocated items, eliminations and other detail table (share-based plans, deferred compensation, capitalized interest, pension, postretirement, eliminations).',
    'pens_item': REL + 'segment table (FAS/CAS service cost adjustment, 2018+) and the unallocated items detail (pension / postretirement, to 2017).',
    'is_rev': REL + 'consolidated statements of operations, as originally reported (2016–17 ASC 606 / ASU 2017-07 restatements in the 2018 releases are not recast).',
    'o_pub': REL + 'Table 1 summary and "Reconciliation of Non-GAAP Measures" (core operating earnings; FY2011 from the Q4-2012 release comparative).',
    'e_pub': REL + 'core EPS reconciliation (published from the Q4-2012 release, FY2011 comparative).',
    'cf_ni': REL + 'consolidated statements of cash flows (year to date; discrete quarters derived: Q2 = H1 − Q1, Q3 = 9M − H1, Q4 = FY − 9M).',
    'bs_cash': REL + 'consolidated statements of financial position; shares outstanding = 1,012,261,159 issued − treasury shares (caption; XBRL TreasuryStockCommonShares for 2017–19 year ends).',
    'ppa_cons': 'Source: FY2025 10-K Note 2 (Spirit Acquisition, preliminary purchase-price allocation).',
    'sh_dil': REL + 'statements of operations / core reconciliation (diluted weighted average shares; loss periods: diluted = basic).',
    'ds_total': 'Source: release balance sheets; FY2025 10-K debt note (scheduled principal payments); Q2-26 10-Q liquidity section.',
    'fcf_pub': REL + 'Table 2 "Cash Flow" (free cash flow = operating cash flow − additions to PP&E).',
    'cal_h': 'Guidance end-points are not in EDGAR filings. Sources (third-party call transcripts / summaries) are listed in manual_items.json → guidance_2026. Treat as UNVERIFIED.',
}
CELL_COMMENTS = {}
def comments_setup():
    A = {y: ANN[y].c for y in ANN}
    CELL_COMMENTS['eps_chk'] = {A[2006]: 'n/c: the FY2006 release prints diluted EPS $2.85, while net earnings $2,215m ÷ the 787.6m diluted shares shown = $2.81 (the presented share count is not the EPS denominator). Carried as published; documented exception.'}
    CELL_COMMENTS['bca_rev'] = {QC[(2017, 3)].c: 'Q3-17: Global Services formed — commercial services moved out of BCA (and government services out of BDS). Q1/Q2-17 columns are as originally reported on the prior structure.'}
    CELL_COMMENTS['bgs_oi'] = {QC[(2025, 4)].c: 'Q4-25 includes the $9.6bn pre-tax gain on the Digital Aviation Solutions divestiture (closed 31-Oct-2025).'}
    CELL_COMMENTS['e_chk'] = {QC[(2020, 4)].c: 'Q4-20: published core EPS $(15.25) vs model $(15.23) (per-share items rounded in the release); within the ±0.025 tolerance stated in the label.'}
    CELL_COMMENTS['eps_num'] = {QC[(2025, 4)].c: 'Q4-25 diluted EPS ($10.23) is on the if-converted basis for the mandatory convertible preferred: net earnings attributable ($8,220m) ÷ 803.8m shares (release).'}
    CELL_COMMENTS['bds_rev'] = {A[2010]: '2010–11 releases print BMA / N&SS / GS&S without a BDS total; the total is the sum (no estimate).', A[2011]: 'Sum of sub-segments (no total printed).'}

DEFS = {}
def defs_setup():
    A = {y: ANN[y].c for y in ANN}
    q = lambda y, n: QC[(y, n)].c
    DEFS['o_def'] = {
        'C': ('2006–07: core operating earnings not published; unallocated pension detail not disclosed → core = GAAP (flag)',
              'The 2006–07 releases published "adjusted EPS" on a different basis (e.g. FY2006 $3.62, excluding tax benefits and other items); not carried.'),
        A[2008]: ('2008–10: not published — model rebuild = earnings from operations + unallocated pension & postretirement expense (unallocated detail)', None),
        A[2011]: ('2011–17: company core = earnings from operations + unallocated pension / postretirement expense (FY2011 from the Q4-12 release comparative)',
                  'Core operating earnings, core operating margin and core EPS were first published in the Q4-2012 release (30-Jan-2013) with FY2011 comparatives.'),
        A[2018]: ('2018+: core = earnings from operations − FAS/CAS service cost adjustment (ASU 2017-07; non-service pension moved below operating earnings)',
                  'Q1-2018 release: 2017 restated on the new definition (FY2017 core EPS $12.33 vs $12.04 as originally published); the model keeps the originally published 2017 figures.'),
    }
    DEFS['e_def'] = {
        'C': ('2006–10: core EPS not published (model rebuild at 35% from 2008; = GAAP 2006–07)', None),
        A[2011]: ('2011–17: core EPS = GAAP diluted EPS + unallocated pension / postretirement expense, net of tax (35% statutory; 2016–17 three-line per-share bridge)', None),
        A[2018]: ('2018+: core EPS = GAAP diluted EPS − FAS/CAS service cost adjustment − non-operating pension & postretirement income + deferred taxes on the adjustments (21%)', None),
        q(2024, 4): ('Q4-24+: mandatory convertible preferred dividends deducted (if-converted in Q4-25)', None),
    }
def write_defs(ws, row):
    for c, (txt, cm) in DEFS[row.key].items():
        cell = ws[f'{c}{row.r}']; cell.value = txt
        cell.font = Font(name='Calibri', sz=10, b=True, i=True, color='FFC00000')
        if cm: comment(cell, cm)

def write_scen(ws):
    srcrow = {'hdr': 9, 'out_hdr': 11, 'pt_hdr': 15, 'out': 12, 'pt': 16, 'lev': 28, 'blk': 30, 'yr': 31}
    for spec_ in SCEN_ROWS:
        r, kind = spec_[0], spec_[1]
        for ci in range(CI(SL), CI(SNOTE) + 1):
            ws.cell(r, ci)._style = copy.copy(TM.cell(srcrow[kind], ci + TSHIFT)._style)
        if kind in ('hdr', 'out_hdr', 'pt_hdr'):
            _, _, a, b, c_, d_, e = spec_
            for col, v in zip([SL, SBULL, SBASE, SBEAR, SNOTE], [a, b, c_, d_, e]):
                if v is not None: ws[f'{col}{r}'] = v
        elif kind == 'out':
            _, _, lab, hi, mid, lo, note, nf = spec_
            ws[f'{SL}{r}'] = lab
            for col, v in zip([SBULL, SBASE, SBEAR], [hi, mid, lo]):
                ws[f'{col}{r}'] = v; ws[f'{col}{r}'].number_format = NF[nf]; set_color(ws[f'{col}{r}'], BLUE)
            ws[f'{SNOTE}{r}'] = note
        elif kind == 'pt':
            _, _, lab, _, val, _, note, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBASE}{r}'] = val; ws[f'{SBASE}{r}'].number_format = NF[nf]; set_color(ws[f'{SBASE}{r}'], BLUE); ws[f'{SNOTE}{r}'] = note
        elif kind == 'lev':
            _, _, lab, a, b, c_, note = spec_
            ws[f'{SL}{r}'] = lab
            for col, v in zip([SBULL, SBASE, SBEAR], [a, b, c_]):
                ws[f'{col}{r}'] = v; ws[f'{col}{r}'].number_format = NF['x']
            ws[f'{SNOTE}{r}'] = note
        elif kind == 'blk':
            _, _, lab, db, mid, de, note, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBULL}{r}'] = db; ws[f'{SBASE}{r}'] = mid; ws[f'{SBEAR}{r}'] = de; ws[f'{SNOTE}{r}'] = note
            for col in (SBULL, SBEAR):
                ws[f'{col}{r}'].number_format = '\\+#,##0;\\-#,##0' if nf == 'm' else '\\+0.0%;\\-0.0%'
        elif kind == 'yr':
            _, _, lab, hdr, base, clamp, _, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBASE}{r}'] = base
            if clamp:
                ws[f'{SBULL}{r}'] = f'=MAX(0,{SBASE}{r}+${SBULL}${hdr})'; ws[f'{SBEAR}{r}'] = f'=MAX(0,{SBASE}{r}+${SBEAR}${hdr})'
            else:
                ws[f'{SBULL}{r}'] = f'={SBASE}{r}+${SBULL}${hdr}'; ws[f'{SBEAR}{r}'] = f'={SBASE}{r}+${SBEAR}${hdr}'
            for col in (SBULL, SBASE, SBEAR): ws[f'{col}{r}'].number_format = NF[nf]

def write_header(ws):
    for r in range(1, 8):
        clone_row_style(ws, r, r, 2, CI(LASTC))
    ws[f'{NOTE}1'] = 'SCENARIO SWITCH ▼ (Bull / Base / Bear) — drives all scenario-linked forecast drivers'
    ws[f'{NOTE}2'] = 'Base'
    comment(ws[f'{NOTE}2'], 'Scenario toggle. Choose Bull / Base / Bear. Drives: 2026 delivery, BDS / BGS margin and free-cash-flow end-points (high / mid / low; UNVERIFIED call statements), '
            '2H26 BCA margin, 2027E–2031E deliveries by program, revenue per delivery, segment growth and margins, M&A and the leverage target for 2027E+ buybacks '
            f'(scenario table {SL}:{SNOTE}).')
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type='list', formula1='"Bull,Base,Bear"', allow_blank=False); ws.add_data_validation(dv); dv.add(f'{NOTE}2')
    ws['B2'] = 'The Boeing Company (NYSE: BA) — 3-Statement Financial Model'
    ws['B3'] = ('US GAAP as reported · USD millions (per-share data in USD; shares and deliveries in millions / units) · Fiscal years ended 31 December · '
                'Source: Boeing Forms 10-K / 10-Q, Form 8-K earnings releases (Ex. 99.1) and deal / financing filings, SEC EDGAR (CIK 0000012927); XBRL company facts for cross-checks; '
                'investor-relations releases at boeing.com are the same documents as filed on EDGAR. 2026 outlook figures are UNVERIFIED call statements (not in EDGAR filings)')
    ws['B4'] = ('Basis notes: as originally reported in each period\'s release. Segments: BCA, IDS→BDS (2010), BCC (to 2022), Other (to 2013); Global Services formed Q3-17; '
                '2018 adoption of ASC 606 (2016–17 restated, not recast here) and ASU 2017-07 (FAS/CAS line; core definition change). Deals: Spirit AeroSystems acquired 8-Dec-25; '
                'Digital Aviation Solutions sold 31-Oct-25; Oct-24 equity raise ($18.2bn common + $5.75bn 6% mandatory convertible preferred, converts Oct-27). '
                'Q3/Q4-26E calibrated to 2026 delivery, BDS / BGS margin and FCF end-points (UNVERIFIED).')
    ws[f'{CAGR[0]}5'] = 'Forecast'
    for c in CAGR[1:]: ws[f'{c}5'] = 'Historical'
    ws['B6'] = '(USDm)'
    for col in COLS: ws[f'{col.c}6'] = col.label
    ws[f'{NOTE}6'] = 'Modelling Notes'
    for c, a, b in zip(CAGR, ['5Y CAGR', '5Y CAGR', '10Y CAGR', '19Y CAGR'], ["26-'31", "20-'25", "15-'25", "06-'25"]):
        ws[f'{c}6'] = a; ws[f'{c}7'] = b

def set_columns(ws):
    for k in list(ws.column_dimensions.keys()): del ws.column_dimensions[k]
    ws.column_dimensions['A'].width = TM.column_dimensions['A'].width
    ws.column_dimensions['B'].width = TM.column_dimensions['B'].width
    wid = {'A': 13.0, 'Q': 9.5703125, 'QE': 11.42578125, 'E26': 14.140625, 'AE': 13.0}
    for col in COLS:
        cd = ws.column_dimensions[col.c]; cd.width = wid[col.kind]
        if col.q:
            cd.outlineLevel = 1; cd.hidden = True
    for ci in range(CI(NOTE), CI(LASTC) + 1):
        ws.column_dimensions[L(ci)].width = TM.column_dimensions[L(ci + TSHIFT)].width or 13.0
    ws.sheet_format.outlineLevelCol = 1

def main():
    import openpyxl
    wb = openpyxl.load_workbook(TEMPLATE)
    if 'DCF_old' in wb.sheetnames: del wb['DCF_old']
    ws = wb['Model']
    blank_style = copy.copy(ws['A1']._style)
    for rowc in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for c in rowc:
            c.value = None; c.comment = None; c._style = copy.copy(blank_style)
    for k in list(ws.row_dimensions.keys()):
        rd = ws.row_dimensions[k]; rd.outlineLevel = 0; rd.hidden = False
    ws.data_validations.dataValidation = []
    scen_layout(V0)
    build_spec(V0)
    n1 = len(fw.SPEC)
    build_spec2(V0)
    for i, row in enumerate(fw.SPEC): row.spec2 = i >= n1
    last = number_rows(8)
    for r in range(1, last + 1): ws.row_dimensions[r].height = 15.0
    ws.row_dimensions[2].height = 21.0
    defs_setup(); comments_setup()
    write_header(ws)
    for row in fw.SPEC: write_row(ws, row)
    write_scen(ws)
    set_columns(ws)
    ws.freeze_panes = 'C8'
    ws.sheet_view.zoomScale = 70
    import other_sheets as osh
    osh.do_all(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    osh.inject_charts(OUT)
    print('rows', fw.SPEC[-1].r, 'saved', OUT)
    return wb

RECALC = os.environ.get('XLSX_RECALC', '/mnt/skills/public/xlsx/scripts/recalc.py')
def recalc(path, timeout=600):
    import subprocess
    out = subprocess.run(['python3', RECALC, path, str(timeout)], capture_output=True, text=True, cwd=os.path.dirname(RECALC))
    try:
        return json.loads(out.stdout)
    except Exception:
        return {'raw': out.stdout[-2000:], 'err': out.stderr[-2000:]}

def scenario_snapshots():
    import openpyxl
    import other_sheets as osh
    snaps = {}
    for scen in ('Bull', 'Bear'):
        p = os.path.join(BASE, f'BA_{scen}.xlsx')
        wb = openpyxl.load_workbook(OUT); wb['Model'][f'{NOTE}2'] = scen; wb.save(p)
        res = recalc(p); print(scen, res.get('status'), res.get('total_errors'))
        v = openpyxl.load_workbook(p, data_only=True)['Bull-Base-Bear']
        snaps[scen] = {(r, c): v[f'{c}{r}'].value for r in range(9, 51) for c in 'CDEFGHIJKL' if v[f'{c}{r}'].value is not None}
        os.remove(p)
    wb = openpyxl.load_workbook(OUT)
    osh.snapshot_values(wb, snaps['Bull'], snaps['Bear'])
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    osh.inject_charts(OUT)

NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
def inject_values(src_xlsx, recalc_xlsx):
    """Store cached formula results (from a LibreOffice-recalculated copy) in the openpyxl file, so the workbook reads correctly in any
    viewer; fullCalcOnLoad stays set so Excel recalculates on open. Same approach as the LECO / ESAB finishing step."""
    import zipfile, shutil, openpyxl
    from lxml import etree
    wbv = openpyxl.load_workbook(recalc_xlsx, data_only=True)
    zin = zipfile.ZipFile(src_xlsx)
    wbxml = etree.fromstring(zin.read('xl/workbook.xml'))
    rels = etree.fromstring(zin.read('xl/_rels/workbook.xml.rels'))
    rid2t = {r.get('Id'): r.get('Target') for r in rels}
    sheets = {}
    for s in wbxml.iter('{%s}sheet' % NS):
        rid = s.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
        sheets['xl/' + rid2t[rid].lstrip('/').replace('xl/', '')] = s.get('name')
    tmp = src_xlsx + '.tmp'
    zout = zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED)
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename in sheets:
            ws = wbv[sheets[item.filename]]
            root = etree.fromstring(data)
            for c in root.iter('{%s}c' % NS):
                f = c.find('{%s}f' % NS)
                if f is None: continue
                v = ws[c.get('r')].value
                old = c.find('{%s}v' % NS)
                if old is not None: c.remove(old)
                if v is None: continue
                ve = etree.SubElement(c, '{%s}v' % NS)
                if isinstance(v, bool):
                    c.set('t', 'b'); ve.text = '1' if v else '0'
                elif isinstance(v, (int, float)):
                    if 't' in c.attrib: del c.attrib['t']
                    ve.text = repr(float(v)) if isinstance(v, float) else str(v)
                elif isinstance(v, str) and v.startswith('#'):
                    c.set('t', 'e'); ve.text = v
                else:
                    c.set('t', 'str'); ve.text = str(v)
            data = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
        zout.writestr(item, data)
    zout.close(); zin.close()
    shutil.move(tmp, src_xlsx)

def store_values():
    import shutil
    p = os.path.join(BASE, 'BA_recalc.xlsx')
    shutil.copy(OUT, p)
    res = recalc(p); print('recalc', res.get('status'), res.get('total_errors'))
    inject_values(OUT, p)
    os.remove(p)

if __name__ == '__main__':
    main()
    if '--snap' in sys.argv: scenario_snapshots()
    if '--recalc' in sys.argv: store_values()
