"""Build FICO_Model.xlsx (Fair Isaac Corporation) from the MLM template (same layout, formatting and process as the prior models)."""
import sys, copy, os, json
from openpyxl.styles import Font
from fw import *
import fw
from spec import build_spec, scen_layout, SCEN_ROWS, SC
from spec2 import build_spec2

OUT = os.path.join(os.path.dirname(BASE), 'FICO_Model.xlsx')

# ====================================================================== inputs (V0)
def wc_inputs():
    q = DATA['2025']['x']
    rev = q['rev']
    return {'dso': round(q['bs_ar'] / rev * 365), 'ap': round(q['bs_ap'] / rev, 3), 'acc': round(q['bs_accomp'] / rev, 3), 'dr': round(q['bs_defrev'] / rev, 3)}
WC = wc_inputs()
WC['note'] = (f'Forecast inputs at fiscal-year-end (30-Sep-2025) levels on FY2025 revenue (DSO {WC["dso"]}d — Scores are billed in arrears; payables & accrued {WC["ap"]:.1%}, '
              f'accrued compensation {WC["acc"]:.1%}, deferred revenue {WC["dr"]:.1%} of revenue); year ends are compared with year ends (quarter ends are seasonal).')
G = MAN['guidance_fy26']['current']
NM = G['ng_eps']; G9 = {k: sum(DATA[f'Q{q}-26']['x'][k] for q in (1, 2, 3)) for k in ('adj_sbc', 'ng_co_etb')}

V0 = {
    'out_hdr': 'FY26 guidance — Q3 FY26 release (29-Jul-2026; point guidance)',
    'out_note': ('Revenue $2.53bn (raised from $2.45bn), GAAP net income $850m / EPS $36.86, non-GAAP net income $979m / EPS $42.43 (raised from $40.45); reconciliation: SBC $188m, '
                 'income tax adjustments −$46m, excess tax benefit −$13m. No segment, margin, tax, interest or cash-flow guidance. Point guidance, so high = mid = low.'),
    'outlook': [
        ('out_rev', '  FY26 revenue ($m)', (G['revenue'],) * 3, 'm', 'Guided $2.53bn. Drives the Q4/26E revenue calibration.'),
        ('out_eps', '  FY26 non-GAAP EPS ($)', (NM,) * 3, 'ps', 'Guided $42.43. Drives the Q4/26E segment-margin calibration.'),
        ('out_gni', '  FY26 GAAP net income ($m) — memo', (G['gaap_ni'],) * 3, 'm', 'Guided $850m (information row; excludes the $27m Q4 restructuring announced 6-Oct-2026).'),
        ('out_geps', '  FY26 GAAP EPS ($) — memo', (G['gaap_eps'],) * 3, 'ps', 'Guided $36.86 (information row).'),
        ('bb26', '  Q4/26E share repurchases ($m; cash)', (500.0, 300.0, 100.0), 'm',
         'Assumption. $800m left on the June 2026 program at 30-Jun-2026, of which $300m is the ASR prepayment (settles in Q4 with no further cash) → $500m available; 9M FY26 $3,046m.'),
    ],
    'points': [
        ('pt_int', '  Q4/26E interest expense, net ($m)', 78.0, 'm1', 'Not guided. Full quarter on ~$5.6bn debt: notes $3.4bn at ~5.5%, term loan $1.5bn at 5.86%, revolver ~$0.7bn at 5.64% (Q3 $59.9m incl. 25 days of the term loan).'),
        ('pt_othinc', '  Q4/26E other income (expense), net ($m)', 3.0, 'm1', '9M FY26 $9.9m (Q3 $11.9m incl. investment gains).'),
        ('pt_dda', '  Q4/26E depreciation & amortization ($m)', 4.3, 'm1', 'Q3 $4.2m; rising with capitalized software.'),
        ('pt_sbc', '  Q4/26E share-based compensation ($m)', round(G['sbc'] - G9['adj_sbc'], 1), 'm1', f'FY26 guidance $188m − 9M actual ${G9["adj_sbc"]:.1f}m.'),
        ('pt_restr', '  Q4/26E restructuring charges ($m)', MAN['restructuring_oct2026']['q4_fy26_charge'], 'm1', '8-K Item 2.05 (1-Oct-2026): ~$27.0m pre-tax in Q4 FY26 (15% workforce reduction).'),
        ('pt_capex', '  Q4/26E purchases of property and equipment ($m)', 0.6, 'm1', '9M FY26 $1.4m.'),
        ('pt_capsw', '  Q4/26E capitalized internal-use software ($m)', 9.5, 'm1', '9M FY26 $26.5m (Q3 $9.2m).'),
        ('pt_ngtax', '  Q4/26E non-GAAP effective tax rate', 0.255, 'pct', '9M FY26 25.4% (non-GAAP tax = GAAP tax − income tax adjustments − excess tax benefit).'),
        ('pt_trate', '  Q4/26E tax rate on adjusting items', 0.245, 'pct', 'FY26 guidance: income tax adjustments −$46m on SBC $188m = 24.5%.'),
        ('pt_etb', '  Q4/26E excess tax benefit line ($m; + = shortfall added back)', round(G['etb'] - G9['ng_co_etb'], 2), 'm1', f'FY26 guidance −$13m − 9M actual ${G9["ng_co_etb"]:.1f}m.'),
    ],
    'lev': (3.25, 3.0, 2.5),
    'lev_note': ('Net debt $5.28bn at 30-Jun-2026 ≈ 3.5x LTM adjusted EBITDA (model); credit-agreement covenant 4.5x to Dec-2026, 4.0x to Dec-2027, 3.5x thereafter. Target (assumption) = leverage up to which '
                 'free cash flow and debt capacity are returned via buybacks; FICO has bought back stock every year since FY2013. Bull = more buybacks (higher target).'),
    'blocks': [
        ('SC_B2B', 'SC_B2B (B2B scores revenue y/y)', 0.05, -0.12, [0.15, 0.12, 0.10, 0.09, 0.08],
         '9M FY26 +54%, FY25 +33%, FY24 +27% — mortgage origination score unit price (10-Q MD&A). Base assumes pricing gains moderate. Bear: FHFA approval of other score models / VantageScore competition (10-K risk factors).', 'pct', False),
        ('SC_B2C', 'SC_B2C (B2C scores revenue y/y)', 0.02, -0.04, [0.04, 0.04, 0.04, 0.03, 0.03], '9M FY26 +5%, FY25 +6%, FY24 −2% (myFICO subscriptions; royalties on scores sold through the bureaus).', 'pct', False),
        ('SW_PLAT', 'SW_PLAT (platform software revenue y/y)', 0.08, -0.12, [0.30, 0.25, 0.20, 0.17, 0.15],
         '9M FY26 +53%; platform ARR +62% and net retention 148% at 30-Jun-2026 (Q3 FY26 release); FY25 +19%.', 'pct', False),
        ('SW_NONPLAT', 'SW_NONPLAT (non-platform software revenue y/y)', 0.05, -0.05, [-0.15, -0.10, -0.07, -0.05, -0.04],
         '9M FY26 −17%; non-platform ARR −17% and net retention 82% at 30-Jun-2026 (migration to the platform, de-emphasis).', 'pct', False),
        ('SW_PS', 'SW_PS (professional services revenue y/y)', 0.03, -0.03, [-0.05, -0.03, 0.0, 0.0, 0.0], '9M FY26 −9%; FY25 −5%; FY22–25 declining as implementations move to partners.', 'pct', False),
        ('SC_M', 'SC_M (Scores segment operating margin)', 0.01, -0.04, [0.905, 0.905, 0.905, 0.905, 0.905],
         '9M FY26 90.2% (Q3 91%), FY25 87.8%, FY22 88.1%: very high incremental margins on price; held flat (Bear: price concessions).', 'pct', False),
        ('SW_M', 'SW_M (Software segment operating margin)', 0.03, -0.04, [0.27, 0.29, 0.31, 0.32, 0.33],
         '9M FY26 27.6%, FY25 30.1%, FY23 32.6%; platform scale and the Oct-2026 15% workforce reduction (savings not quantified — assumption).', 'pct', False),
        ('MNA', 'M&A_spend (software bolt-ons, $m)', 0.0, 0.0, [0.0, 0.0, 0.0, 0.0, 0.0],
         'FICO has made no material acquisition since FY2014 and returns capital through buybacks; Base $0 (lever available).', 'm', True),
    ],
    'corp_pct': -0.075,
    'corp_note': 'Input −7.5% of revenue from FY2027 (9M FY26 −8.1%, FY25 −9.1%, FY24 −10.9%; operating leverage and the Oct-2026 workforce reduction — assumption).',
    'dda_pct': 0.008, 'cogs_pct': 0.135, 'rd_pct': 0.08, 'sbc_pct': 0.072,
    'ae_restr': {2027: 5.0, 2028: 0.0, 2029: 0.0, 2030: 0.0, 2031: 0.0},
    'othinc': 12.0, 'trate': 0.245, 'etb': -10.0, 'ngtax': 0.25,
    'capex_pct': 0.002, 'capsw_pct': 0.014, 'taxw_pct': 0.5,
    'stock_q4': 8.0, 'stock_ae': 25.0, 'taxw_q4': -5.0,
    'asr_final': round(MAN['buyback']['asr']['initial_shares_m'] / MAN['buyback']['asr']['initial_pct'] - MAN['buyback']['asr']['initial_shares_m'], 3),
    'iss_q4': 0.02, 'iss_ae': 0.15, 'dil': 0.10, 'pxg': 0.08,
    'iss_notes': {2028: 1762.5},
    'repay_notes': {2028: 900.0},
    'tl_rep': {2026: 75.0, 2027: 337.5, 2028: 1087.5},
    'mat': {2026: 337.5, 2027: 1987.5, 2028: 0.0, 2029: 0.0, 2030: 0.0, 2031: 0.0},
    'min_cash': 150.0, 'rev_cap': 750.0,
    'r_notes': {2027: 0.0554, 2028: 0.0554, 2029: 0.0615, 2030: 0.0615, 2031: 0.0615},
    'r_notes_note': ('Weighted coupon 5.54% on the $3.4bn notes (4.00% 2028, 6.00% 2033, 6.25% 2034); from FY2029 6.15% after the assumed FY2028 refinancing of the 2028 notes and the term-loan balance at ~6.25%.'),
    'r_tl': 0.057, 'r_rev': 0.055, 'fees': 4.0, 'r_cash': 0.02,
    'wc': WC,
    'px': MAN['price']['px'],
    'px_note': ('VALUATION — share price input $740.00 (user-provided, 13-Oct-2026; not an SEC-sourced price — update to market. The Q3 FY26 10-Q average repurchase price was $1,148.84)'),
    'mna_note': 'Scenario lever M&A_spend (Base $0). Funded from cash / revolver (debt schedule).',
    'rev_guid_note': 'Company FY26 revenue guidance $2.53bn (raised 29-Jul-2026 from $2.45bn; $2.35bn initial, 5-Nov-2025). Drives the revenue calibration.',
    'eps_guid_note': 'Company FY26 non-GAAP EPS guidance $42.43 (raised 29-Jul-2026 from $40.45; $38.17 initial). Drives the segment-margin calibration.',
}

# ====================================================================== compute engine
def hist_value(row, col):
    kw = row.kw; path = kw.get('data')
    if not path: return None
    if kw.get('ppa'):
        v = MAN['ppa'].get(pkey(col), {}).get(path.split('_', 1)[1])
        return v if kw.get('textdata') else num(v)
    v = xget(pkey(col), path)
    return v if kw.get('textdata') else num(v)

def first(kw, keys):
    for k in keys:
        if k in kw: return k
    return None

def compute(row, col):
    kw = row.kw; k = col.kind
    CTX.col = col
    if kw.get('ann') and col.q: return None, False
    if k in HISTK:
        v = hist_value(row, col)
        if v is not None: return v, True
        for key in ('hist', 'all'):
            if key in kw:
                return kw[key](), False
        return None, False
    if k == 'QE':
        if row.spec2: return None, False
        if 'qe_in' in kw:
            v = kw['qe_in']; return (v[col.q] if isinstance(v, dict) else v), True
        key = first(kw, ('qe', 'fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    if k == 'E26':
        if 'e26_in' in kw: return kw['e26_in'], True
        if 'e26' in kw: return kw['e26'](), False
        if kw.get('avg') and ('qe' in kw or 'qe_in' in kw): return f'=AVERAGE({Q1}{row.r}:{QE}{row.r})', False
        if kw.get('flow') and not row.spec2 and any(x in kw for x in ('qe', 'qe_in', 'fc')): return S26(row.key), False
        key = first(kw, ('fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    if k == 'AE':
        if 'ae_in' in kw:
            v = kw['ae_in']; return (v[col.year] if isinstance(v, dict) else v), True
        key = first(kw, ('ae', 'fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    return None, False

# ====================================================================== writers
CAGR_G = {'is_rev', 'g_rev', 'sc_rev', 'sw_rev', 'sc_b2b', 'sc_b2c', 'sw_plat', 'sw_nonplat', 'sw_ps', 'sc_oi', 'sw_oi', 'o_adj', 'g_adj', 'r_adj',
          'r_ebitda', 'r_ebit', 'is_op', 'd_dda', 'e_adjni', 'e_eps', 'cf_cfo', 'fcf', 'bs_ta', 'sh_dil', 'is_sga', 'is_rd', 'is_cogs', 'is_ni', 'eps_dil', 'd_sbc'}
CAGR_A = {'sc_m', 'sw_m', 'g_adjm', 'o_adj_m', 'r_adj_m', 'is_op_m', 'r_ebit_m', 'gm_rev', 'gm_sc', 'gm_sw', 'gm_ebitda_m', 'gm_op_m', 'gm_net_m', 'gm_sga',
          'gm_rd', 'gm_adjop_m', 'fcf_m', 'fcf_conv', 'capex_pct', 'capsw_pct', 'ra_roic', 'rn_ronta', 'lv_x', 'd_cogs', 'd_rd', 'd_sga', 'd_sbc_pct',
          'wc_nwc_p', 'is_etr', 'e_rate', 'fcf_eb', 'corp_pct', 'sc_mix', 'sc_b2b_p', 'sw_plat_p'}
A = {y: ANN[y].c for y in ANN}

def style_cagr(ws, r, on):
    src = 16 if on else 12
    for c in CAGR:
        ws[f'{c}{r}']._style = copy.copy(TM[f'{L(CI(c) + TSHIFT)}{src}']._style)

def write_row(ws, row):
    r = row.r; arch = ARCH[row.arch]
    clone_row_style(ws, arch, r)
    ws.cell(r, 2).value = row.label or None
    if row.kw.get('note'): ws[f'{NOTE}{r}'] = row.kw['note']
    style_cagr(ws, r, row.key in CAGR_G or row.key in CAGR_A)
    if row.arch in ('section', 'block', 'blank', 'sub') and not row.kw.get('data'):
        return
    if row.key in DEFS:
        write_defs(ws, row); return
    nf = NF[row.nf]
    for col in COLS:
        v, inp = compute(row, col)
        if v is None: continue
        cell = ws[f'{col.c}{r}']
        cell.value = v
        if not row.kw.get('textdata'): cell.number_format = nf
        if inp:
            set_color(cell, BLUE, bold=True if col.kind in FCK else None)
            lk = row.kw.get('labkey')
            if lk and col.kind in HISTK and isinstance(v, (int, float)) and abs(v) > 1e-9:
                t = xget(pkey(col), lk)
                if t: comment(cell, f'Company items ({pkey(col)}): {t}', 380, 110)
        elif row.arch == 'growth' and col.kind in HISTK:
            set_color(cell, GRAY)
        else:
            set_color(cell, BLACK)
    rc = CELL_COMMENTS.get(row.key, {})
    for c, t in rc.items(): comment(ws[f'{c}{r}'], t)
    if row.key in CAGR_G:
        ws[f'{CAGR[0]}{r}'] = f'=IFERROR(({A[2031]}{r}/{A[2026]}{r})^(1/5)-1,"n/m")'
        ws[f'{CAGR[1]}{r}'] = f'=IFERROR(({A[2025]}{r}/{A[2020]}{r})^(1/5)-1,"n/m")'
        ws[f'{CAGR[2]}{r}'] = f'=IFERROR(({A[2025]}{r}/{A[2015]}{r})^(1/10)-1,"n/m")'
        ws[f'{CAGR[3]}{r}'] = f'=IFERROR(({A[2025]}{r}/{A[2006]}{r})^(1/19)-1,"n/m")'
        for c in CAGR: ws[f'{c}{r}'].number_format = '0%;\\(0%\\);"n/m"'
    elif row.key in CAGR_A:
        ws[f'{CAGR[0]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2027, 2032))}),"n/m")'
        ws[f'{CAGR[1]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2021, 2026))}),"n/m")'
        ws[f'{CAGR[2]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2016, 2026))}),"n/m")'
        ws[f'{CAGR[3]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2006, 2026))}),"n/m")'
        for c in CAGR: ws[f'{c}{r}'].number_format = NF['pct']
    t = LABEL_SOURCES.get(row.key)
    if t: comment(ws.cell(r, 2), t)

REL = 'Source: Fair Isaac earnings releases (Form 8-K Item 2.02, Ex. 99.1) — '
FIL = 'Source: Fair Isaac Forms 10-Q / 10-K — '
LABEL_SOURCES = {
    'sc_b2b': (FIL + 'MD&A tables "business-to-business / business-to-consumer Scores" and "platform / non-platform software" revenue (first published in the FY2021 10-K '
               'for FY2019–21; FY2021 quarters from the FY2022 10-Q prior-year columns; Q4 = fiscal year − nine months).'),
    'sc_pub': FIL + 'segment note and MD&A segment summaries (each period as first reported; Q4 = fiscal year − nine months; FY2021 Q4 nine months from the FY2022 Q3 10-Q re-presented column).',
    'sw_pub': 'Consolidated revenue (release) − Scores segment revenue (10-Q / 10-K).',
    'sc_oi': FIL + 'MD&A "segment operating income" summary (segments, unallocated corporate expenses, share-based compensation, amortization, restructuring, gains → operating income).',
    'leg_sms_rev': FIL + 'segment note / MD&A; product-line segments to FY2009 (Strategy Machine Solutions, Scoring Solutions, Professional Services, Analytic Software Tools), '
                   'Applications / Scores / Tools FY2010–15, Applications / Scores / Decision Management Software FY2016–Q3/21, Scores / Software from the FY2021 10-K.',
    'corp': FIL + 'MD&A segment summary ("unallocated corporate expenses"; none before FY2010 — corporate costs were allocated to the segments).',
    'is_rev': REL + 'condensed consolidated statements of income, as originally reported (three-month columns; FY2006–12 from the 10-K, FY2013+ from the Q4 release twelve-month column).',
    'is_int': FIL + 'segment note ("unallocated interest expense, net"; FY2006–13 interest expense − interest income); the releases show interest inside "other expense, net".',
    'd_dda': REL + 'condensed statements of cash flows (year to date; discrete quarters derived).',
    'e_pubni': REL + '"Non-GAAP Results" table (published from Q1 FY2013).',
    'fcf_pub': REL + '"Free cash flow" table (published from Q1 FY2013; FY2013–Q2/17 the company also deducted dividends paid).',
    'cf_ni': REL + 'condensed statements of cash flows (year to date; discrete quarters derived, Q4 = FY − 9M); FY2006–12 from the 10-K.',
    'bs_cash': REL + 'condensed balance sheets (period ends; FY2006–12 from the 10-K); shares outstanding from XBRL (CommonStockSharesOutstanding, as first filed).',
    'sh_basic': REL + 'statements of income (weighted-average shares, thousands). No stock splits in the history.',
    'ds_total': 'Source: release balance sheets; Q3 FY26 10-Q Note 6 (debt); 8-Ks of 11-Mar-2026 (2026 Senior Notes) and 5-Jun-2026 (term loan).',
    'cal_guid': 'Source: Q3 FY26 release (8-K Ex. 99.1, 29-Jul-2026), outlook and "Reconciliation of Non-GAAP Guidance".',
    'd_restr': 'Source: 8-K Item 2.05 dated 1-Oct-2026 (filed 6-Oct-2026).',
    'bb_asr': 'Source: Q3 FY26 10-Q Note 7 and Part II Item 2 (ASR agreement with Wells Fargo Securities, 5-Jun-2026).',
}
CELL_COMMENTS = {}
def comments_setup():
    for y in (2006, 2007, 2008, 2009):
        CELL_COMMENTS.setdefault('sc_pub', {})[A[y]] = f'FY{y}: "Scoring Solutions" product-line segment (as originally reported; not the later Scores segment definition).'
    CELL_COMMENTS['sc_oi'] = {A[2021]: 'FY2021 as reported in the FY2021 10-K on the new Scores / Software basis (FY2021 quarters on the Applications / Scores / DMS basis; Q4/21 = FY − 9M re-presented).'}
    CELL_COMMENTS['is_disc'] = {A[2007]: 'Discontinued operations (business units sold in FY2008), as reported in the FY2007 10-K.'}
    CELL_COMMENTS['e_tax'] = {QC[(2013, 1)].c: 'FY2013–Q2/16: the company showed items net of tax; the tax effect is derived (published non-GAAP NI − GAAP NI − pre-tax items).'}
    CELL_COMMENTS['fcf_div'] = {QC[(2013, 1)].c: 'FY2013–Q2/17 the published free cash flow = CFO − capex − dividends paid; the check adds the dividends back.'}

DEFS = {}
def defs_setup():
    q = lambda y, n: QC[(y, n)].c
    DEFS['o_def'] = {
        'C': ('FY2006–09: product-line segments; segment results after amortization and corporate costs; unallocated = SBC, restructuring, (gains) / losses on sales',
              'FY2006–09 10-K MD&A: segment operating income for Strategy Machine Solutions, Scoring Solutions, Professional Services and Analytic Software Tools; only share-based '
              'compensation, restructuring / acquisition-related and product-line sale results were unallocated.'),
        A[2010]: ('FY2010+: Applications / Scores / Tools; unallocated corporate expenses, SBC, amortization, restructuring', 'FY2010 10-K. Earlier years are not recast in the model.'),
        A[2016]: ('FY2016+: Decision Management Software replaces Tools', 'FY2016 10-K (FY2016 annual column); quarters from Q1/17.'),
        A[2021]: ('FY2021 10-K: Scores / Software (quarters from Q1/22)', 'Software = Applications + Decision Management Software, with re-allocated costs.'),
        A[2025]: ('FY2025 10-K: "total segment operating income" shown before unallocated corporate expenses — model measure unchanged', None),
    }
    DEFS['e_def'] = {
        'C': ('FY2006–12: no non-GAAP measures published — non-GAAP = GAAP (flag)', None),
        q(2013, 1): ('Q1 FY13+: non-GAAP NI / EPS excl. amortization, restructuring & acquisition, SBC (items shown net of tax); FCF published', None),
        q(2016, 3): ('Q3 FY16+: items pre-tax with a separate "income tax adjustments" line', None),
        q(2017, 1): ('Q1 FY17+: excess tax benefit on share-based awards excluded (ASU 2016-09)', None),
        q(2018, 1): ('FY2018–19: Tax Cuts and Jobs Act effects and a realized gain on a cost-method investment excluded', None),
        q(2021, 1): ('FY2021–24: gains on product-line asset sales and the C&R divestiture excluded', None),
    }
def write_defs(ws, row):
    for c, (txt, cm) in DEFS[row.key].items():
        cell = ws[f'{c}{row.r}']; cell.value = txt
        cell.font = Font(name='Calibri', sz=10, b=True, i=True, color='FFC00000')
        if cm: comment(cell, cm)

def write_scen(ws):
    srcrow = {'hdr': 9, 'out_hdr': 11, 'pt_hdr': 15, 'out': 12, 'pt': 16, 'lev': 28, 'blk': 30, 'yr': 31}
    for spec_ in SCEN_ROWS:
        r, kind = spec_[0], spec_[1]
        for ci in range(CI(SL), CI(SNOTE) + 1):
            ws.cell(r, ci)._style = copy.copy(TM.cell(srcrow[kind], ci + TSHIFT)._style)
        if kind in ('hdr', 'out_hdr', 'pt_hdr'):
            _, _, a, b, c_, d_, e = spec_
            for col, v in zip([SL, SBULL, SBASE, SBEAR, SNOTE], [a, b, c_, d_, e]):
                if v is not None: ws[f'{col}{r}'] = v
        elif kind == 'out':
            _, _, lab, hi, mid, lo, note, nf = spec_
            ws[f'{SL}{r}'] = lab
            for col, v in zip([SBULL, SBASE, SBEAR], [hi, mid, lo]):
                ws[f'{col}{r}'] = v; ws[f'{col}{r}'].number_format = NF[nf]; set_color(ws[f'{col}{r}'], BLUE)
            ws[f'{SNOTE}{r}'] = note
        elif kind == 'pt':
            _, _, lab, _, val, _, note, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBASE}{r}'] = val; ws[f'{SBASE}{r}'].number_format = NF[nf]; set_color(ws[f'{SBASE}{r}'], BLUE); ws[f'{SNOTE}{r}'] = note
        elif kind == 'lev':
            _, _, lab, a, b, c_, note = spec_
            ws[f'{SL}{r}'] = lab
            for col, v in zip([SBULL, SBASE, SBEAR], [a, b, c_]):
                ws[f'{col}{r}'] = v; ws[f'{col}{r}'].number_format = NF['x']
            ws[f'{SNOTE}{r}'] = note
        elif kind == 'blk':
            _, _, lab, db, mid, de, note, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBULL}{r}'] = db; ws[f'{SBASE}{r}'] = mid; ws[f'{SBEAR}{r}'] = de; ws[f'{SNOTE}{r}'] = note
            for col in (SBULL, SBEAR):
                ws[f'{col}{r}'].number_format = '\\+#,##0;\\-#,##0' if nf == 'm' else '\\+0.0%;\\-0.0%'
        elif kind == 'yr':
            _, _, lab, hdr, base, clamp, _, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBASE}{r}'] = base
            if clamp:
                ws[f'{SBULL}{r}'] = f'=MAX(0,{SBASE}{r}+${SBULL}${hdr})'; ws[f'{SBEAR}{r}'] = f'=MAX(0,{SBASE}{r}+${SBEAR}${hdr})'
            else:
                ws[f'{SBULL}{r}'] = f'={SBASE}{r}+${SBULL}${hdr}'; ws[f'{SBEAR}{r}'] = f'={SBASE}{r}+${SBEAR}${hdr}'
            for col in (SBULL, SBASE, SBEAR): ws[f'{col}{r}'].number_format = NF[nf]

def write_header(ws):
    for r in range(1, 8):
        clone_row_style(ws, r, r, 2, CI(LASTC))
    ws[f'{NOTE}1'] = 'SCENARIO SWITCH ▼ (Bull / Base / Bear) — drives all scenario-linked forecast drivers'
    ws[f'{NOTE}2'] = 'Base'
    comment(ws[f'{NOTE}2'], 'Scenario toggle. Choose Bull / Base / Bear. Drives: Q4/26E buybacks, 2027E–2031E revenue-stream growth (B2B / B2C Scores, platform / non-platform software, '
            f'professional services), segment margins, M&A spend and the leverage target for 2027E+ buybacks (scenario table {SL}:{SNOTE}). FY26 guidance is a point estimate in all cases.')
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type='list', formula1='"Bull,Base,Bear"', allow_blank=False); ws.add_data_validation(dv); dv.add(f'{NOTE}2')
    ws['B2'] = 'Fair Isaac Corporation (NYSE: FICO) — 3-Statement Financial Model'
    ws['B3'] = ('US GAAP as reported · USD millions (per-share data in USD; shares in millions) · Fiscal years ended 30 September (2026 = Oct-2025 – Sep-2026) · '
                'Source: Fair Isaac Forms 10-K / 10-Q, Form 8-K earnings releases (Ex. 99.1) and financing / restructuring 8-Ks, SEC EDGAR (CIK 814547); XBRL company facts; '
                'investor-relations releases at fico.com are the same documents as filed on EDGAR')
    ws['B4'] = ('Basis notes: as originally reported. Segments: Strategy Machine Solutions / Scoring Solutions / Professional Services / Analytic Software Tools (FY2006–09), Applications / '
                'Scores / Tools (FY2010–15), Applications / Scores / DMS (FY2016–Q3/21), Scores / Software (FY2021 10-K on). Disc ops FY2007–09. Non-GAAP published from Q1 FY13. '
                'Q4/26E calibrated to FY26 revenue and non-GAAP EPS guidance (29-Jul-2026) plus the Oct-2026 restructuring; $1.5bn ASR (Jun-2026) settles in Q4 FY26.')
    ws[f'{CAGR[0]}5'] = 'Forecast'
    for c in CAGR[1:]: ws[f'{c}5'] = 'Historical'
    ws['B6'] = '(USDm)'
    for col in COLS: ws[f'{col.c}6'] = col.label
    ws[f'{NOTE}6'] = 'Modelling Notes'
    for c, a_, b_ in zip(CAGR, ['5Y CAGR', '5Y CAGR', '10Y CAGR', '19Y CAGR'], ["26-'31", "20-'25", "15-'25", "06-'25"]):
        ws[f'{c}6'] = a_; ws[f'{c}7'] = b_

def set_columns(ws):
    for k in list(ws.column_dimensions.keys()): del ws.column_dimensions[k]
    ws.column_dimensions['A'].width = TM.column_dimensions['A'].width
    ws.column_dimensions['B'].width = TM.column_dimensions['B'].width
    wid = {'A': 13.0, 'Q': 9.5703125, 'QE': 11.42578125, 'E26': 14.140625, 'AE': 13.0}
    for col in COLS:
        cd = ws.column_dimensions[col.c]; cd.width = wid[col.kind]
        if col.q:
            cd.outlineLevel = 1; cd.hidden = True
    for ci in range(CI(NOTE), CI(LASTC) + 1):
        ws.column_dimensions[L(ci)].width = TM.column_dimensions[L(ci + TSHIFT)].width or 13.0
    ws.sheet_format.outlineLevelCol = 1

def main():
    import openpyxl
    wb = openpyxl.load_workbook(TEMPLATE)
    if 'DCF_old' in wb.sheetnames: del wb['DCF_old']
    ws = wb['Model']
    blank_style = copy.copy(ws['A1']._style)
    for rowc in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for c in rowc:
            c.value = None; c.comment = None; c._style = copy.copy(blank_style)
    for k in list(ws.row_dimensions.keys()):
        rd = ws.row_dimensions[k]; rd.outlineLevel = 0; rd.hidden = False
    ws.data_validations.dataValidation = []
    scen_layout(V0)
    build_spec(V0)
    n1 = len(fw.SPEC)
    build_spec2(V0)
    for i, row in enumerate(fw.SPEC): row.spec2 = i >= n1
    last = number_rows(8)
    for r in range(1, last + 1): ws.row_dimensions[r].height = 15.0
    ws.row_dimensions[2].height = 21.0
    defs_setup(); comments_setup()
    write_header(ws)
    for row in fw.SPEC: write_row(ws, row)
    write_scen(ws)
    set_columns(ws)
    ws.freeze_panes = 'C8'
    ws.sheet_view.zoomScale = 70
    import other_sheets as osh
    osh.do_all(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    osh.inject_charts(OUT)
    print('rows', fw.SPEC[-1].r, 'saved', OUT)
    return wb

RECALC = os.environ.get('XLSX_RECALC', '/mnt/skills/public/xlsx/scripts/recalc.py')
def recalc(path, timeout=600):
    import subprocess
    out = subprocess.run(['python3', RECALC, path, str(timeout)], capture_output=True, text=True, cwd=os.path.dirname(RECALC))
    try:
        return json.loads(out.stdout)
    except Exception:
        return {'raw': out.stdout[-2000:], 'err': out.stderr[-2000:]}

def scenario_snapshots():
    import openpyxl
    import other_sheets as osh
    snaps = {}
    for scen in ('Bull', 'Bear'):
        p = os.path.join(BASE, f'FICO_{scen}.xlsx')
        wb = openpyxl.load_workbook(OUT); wb['Model'][f'{NOTE}2'] = scen; wb.save(p)
        res = recalc(p); print(scen, res.get('status'), res.get('total_errors'))
        v = openpyxl.load_workbook(p, data_only=True)['Bull-Base-Bear']
        snaps[scen] = {(r, c): v[f'{c}{r}'].value for r in range(9, 51) for c in 'CDEFGHIJKL' if v[f'{c}{r}'].value is not None}
        os.remove(p)
    wb = openpyxl.load_workbook(OUT)
    osh.snapshot_values(wb, snaps['Bull'], snaps['Bear'])
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    osh.inject_charts(OUT)

NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
def inject_values(src_xlsx, recalc_xlsx):
    """Store cached formula results (from a LibreOffice-recalculated copy) in the openpyxl file, so the workbook reads correctly in any
    viewer; fullCalcOnLoad stays set so Excel recalculates on open. Same approach as the LECO / ESAB finishing step."""
    import zipfile, shutil, openpyxl
    from lxml import etree
    wbv = openpyxl.load_workbook(recalc_xlsx, data_only=True)
    zin = zipfile.ZipFile(src_xlsx)
    wbxml = etree.fromstring(zin.read('xl/workbook.xml'))
    rels = etree.fromstring(zin.read('xl/_rels/workbook.xml.rels'))
    rid2t = {r.get('Id'): r.get('Target') for r in rels}
    sheets = {}
    for s in wbxml.iter('{%s}sheet' % NS):
        rid = s.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
        sheets['xl/' + rid2t[rid].lstrip('/').replace('xl/', '')] = s.get('name')
    tmp = src_xlsx + '.tmp'
    zout = zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED)
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename in sheets:
            ws = wbv[sheets[item.filename]]
            root = etree.fromstring(data)
            for c in root.iter('{%s}c' % NS):
                f = c.find('{%s}f' % NS)
                if f is None: continue
                v = ws[c.get('r')].value
                old = c.find('{%s}v' % NS)
                if old is not None: c.remove(old)
                if v is None: continue
                ve = etree.SubElement(c, '{%s}v' % NS)
                if isinstance(v, bool):
                    c.set('t', 'b'); ve.text = '1' if v else '0'
                elif isinstance(v, (int, float)):
                    if 't' in c.attrib: del c.attrib['t']
                    ve.text = repr(float(v)) if isinstance(v, float) else str(v)
                elif isinstance(v, str) and v.startswith('#'):
                    c.set('t', 'e'); ve.text = v
                else:
                    c.set('t', 'str'); ve.text = str(v)
            data = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
        zout.writestr(item, data)
    zout.close(); zin.close()
    shutil.move(tmp, src_xlsx)

def store_values():
    import shutil
    p = os.path.join(BASE, 'FICO_recalc.xlsx')
    shutil.copy(OUT, p)
    res = recalc(p); print('recalc', res.get('status'), res.get('total_errors'))
    inject_values(OUT, p)
    os.remove(p)

if __name__ == '__main__':
    main()
    if '--snap' in sys.argv: scenario_snapshots()
    if '--recalc' in sys.argv: store_values()
