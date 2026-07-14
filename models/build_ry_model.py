"""
Build a full three-statement financial model for R. & Y. Tool and Die Co. Limited.

Source data: Annual Compilation Financial Statements (unaudited) prepared by
GD Barkwell Chartered Professional Accountant, extracted from the two provided PDFs:
  - Summary of Statement of Income and Retained Earnings (+ Cost of Sales schedule)
  - Balance Sheet (5-year summary)

Design/template: the "Model" tab of Model_PL_v1.xlsx (Moog Inc. 3-statement model).
This script reproduces that template's visual language (dark-slate section bands,
sage sub-totals, blue hard-coded inputs vs. black formulas, banker number formats)
applied to R&Y's actual line items, annuals only (FY2021–FY2025).

The cash flow statement is fully IMPLIED — derived from the change in each balance
sheet account plus net income and the non-cash amortization from the income
statement. Every year's net change in cash reconciles exactly to the reported cash
balance (a tie-out check row proves this in the workbook).

Units: Canadian dollars, whole dollars (as reported).
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ----------------------------------------------------------------------------
# Style tokens (lifted from the Moog "Model" template)
# ----------------------------------------------------------------------------
NAVY      = "FF1F3864"   # title
GRAY      = "FF595959"   # subtitle
GRAY_LBL  = "FF404040"   # sub-metric label
SLATE     = "FF2F3E46"   # section band fill
GREEN     = "FF52796F"   # sub-section band fill
SAGE      = "FFCAD2C5"   # key-total highlight fill
WHITE     = "FFFFFFFF"
BLACK     = "FF000000"
BLUE      = "FF0000FF"   # hard-coded input

MONEY = r'#,##0;(#,##0);\-'
PCT   = r'0.0%;(0.0%);\-'
RATIO = r'0.00"x";(0.00"x");\-'
DAYS  = r'0" days";(0);\-'
YRS   = r'0.0" yrs";(0.0);\-'

thin = Side(style="thin", color=BLACK)
dbl  = Side(style="double", color=BLACK)
TOP        = Border(top=thin)
TOP_DBL    = Border(top=thin, bottom=dbl)

YEARS = [2021, 2022, 2023, 2024, 2025]
YCOLS = ["C", "D", "E", "F", "G"]          # data columns
PRIOR = {"D": "C", "E": "D", "F": "E", "G": "F"}
CAGR_COL = "H"
NOTE_COL = "I"

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Model"
ws.sheet_view.showGridLines = False

# column widths
ws.column_dimensions["A"].width = 1.4
ws.column_dimensions["B"].width = 44
for col in YCOLS:
    ws.column_dimensions[col].width = 13.5
ws.column_dimensions[CAGR_COL].width = 12
ws.column_dimensions[NOTE_COL].width = 58

# ----------------------------------------------------------------------------
# Two-pass build: first lay out rows to assign numbers, then fill formulas.
# Each entry: (kind, key, label, opts)
# kinds: title, subtitle, header(col band), band, subband, spacer,
#        data (a numbered line)
# ----------------------------------------------------------------------------
R = {}          # key -> row number
layout = []     # ordered list of entries

def add(kind, key=None, label="", **opts):
    layout.append({"kind": kind, "key": key, "label": label, "opts": opts})

# ---- Header block ----
add("title",    label="R. & Y. Tool and Die Co. Limited — Three-Statement Financial Model")
add("subtitle", label="Unaudited · extracted from Annual Compilation Financial Statements prepared by GD Barkwell Chartered Professional Accountant · Canadian dollars (whole $) · fiscal year ends December 31")
add("subtitle", label="Income Statement & Balance Sheet are as reported. The Cash Flow Statement is IMPLIED — derived from the year-over-year change in every balance sheet account plus net income and non-cash amortization. FY2021 cash flow is not computable (no FY2020 balance sheet).")
add("spacer")
add("colhdr")   # the ($) | 2021.. | CAGR | Notes row
add("spacer")

# ---- INCOME STATEMENT ----
add("band", label="INCOME STATEMENT")
add("data", "revenue",     "Revenue", style="input", cagr=True,
    note="Reported net sales.")
add("data", "cogs",        "Cost of sales (schedule)", style="formula_link",
    note="Links to Cost of Sales schedule total below.")
add("data", "gross_margin","Gross margin", style="subtotal",
    formula="{c}%REV-{c}%COGS", bold=True, top=True, cagr=True)
add("data", "gm_pct",      "  Gross margin %", style="pct",
    formula="IFERROR({c}%GROSS/{c}%REV,0)")
add("data", "other_income","Other income", style="input",
    note="2021 includes $60,000 COVID grants.")
add("data", "gm_plus_oi",  "Gross margin plus other income", style="subtotal",
    formula="{c}%GROSS+{c}%OI", bold=True, top=True)
add("spacer")
add("subband", label="General and administrative expenses")
add("data", "ga_mgmt",   "  Management salaries and bonuses", style="input")
add("data", "ga_bank",   "  Bank charges and interest", style="input")
add("data", "ga_amort",  "  Amortization — office equipment", style="input",
    note="Non-cash; added back in cash flow.")
add("data", "ga_adv",    "  Advertising, promotion and sponsorship", style="input")
add("data", "ga_meals",  "  Meals and entertainment", style="input")
add("data", "ga_tele",   "  Telecommunication", style="input")
add("data", "ga_comp",   "  Computer technology", style="input")
add("data", "ga_office", "  Office and miscellaneous", style="input")
add("data", "ga_leases", "  Office equipment leases", style="input")
add("data", "ga_consult","  Consulting fees", style="input")
add("data", "ga_iso",    "  ISO Quality Management System certification", style="input")
add("data", "ga_acct",   "  Accounting and bookkeeping fees", style="input")
add("data", "ga_total",  "Total general and administrative expenses", style="subtotal",
    formula="SUM({c}%GA_FIRST:{c}%GA_LAST)", bold=True, top=True)
add("spacer")
add("data", "nibt", "Net income before income taxes", style="keytotal",
    formula="{c}%GMOI-{c}%GATOT", bold=True, top=True)
add("data", "tax_current",  "  Income taxes — current", style="input")
add("data", "tax_deferred", "  Income taxes — deferred (recovery)", style="input",
    note="Non-cash; equals the change in the future income tax liability.")
add("data", "tax_total",    "Total income taxes", style="subtotal",
    formula="{c}%TAXC+{c}%TAXD", bold=True, top=True)
add("data", "tax_rate",     "  Effective tax rate", style="pct",
    formula="IFERROR({c}%TAXTOT/{c}%NIBT,0)")
add("data", "net_income", "Net income for the year", style="keytotal",
    formula="{c}%NIBT-{c}%TAXTOT", bold=True, top=True, cagr=True)
add("spacer")
add("subband", label="Net income → EBITDA − capex bridge (bottom-up)")
add("data", "br_ni",   "  Net income for the year", style="formula",
    formula="{c}%NI")
add("data", "br_tax",  "  Add: income taxes", style="formula",
    formula="{c}%TAXTOT")
add("data", "br_int",  "  Add: bank charges and interest", style="formula",
    formula="{c}%GABANK", note="Interest is bundled within bank charges & interest.")
add("data", "br_ebit", "  = EBIT", style="subtotal",
    formula="{c}%NI+{c}%TAXTOT+{c}%GABANK", bold=True, top=True)
add("data", "br_amoff","  Add: amortization — office equipment", style="formula",
    formula="{c}%GAAMORT")
add("data", "br_ambld","  Add: amortization — building and shop equipment", style="formula",
    formula="{c}%CSAMORT")
add("data", "br_ebitda","  = EBITDA", style="subtotal",
    formula="{c}%BREBIT+{c}%GAAMORT+{c}%CSAMORT", bold=True, top=True)
add("data", "br_ebitdapct","  EBITDA margin %", style="pct",
    formula="IFERROR({c}%BREBITDA/{c}%REV,0)")
add("data", "br_ebitdainc","  Incremental EBITDA % (ΔEBITDA / ΔRevenue)", style="pct",
    formula="IFERROR(({c}%BREBITDA-{p}%BREBITDA)/({c}%REV-{p}%REV),\"\")", cfonly=True,
    note="Drop-through of incremental revenue to EBITDA; needs a prior year, so FY2021 is n/a.")
add("data", "br_capex","  Less: capital expenditures (implied)", style="formula",
    formula="{c}%CFCAPEX", cfonly=True,
    note="From the implied cash flow; requires prior-year PP&E, so FY2021 is n/a.")
add("data", "br_final","  = EBITDA − capex", style="keytotal",
    formula="{c}%BREBITDA+{c}%CFCAPEX", bold=True, top=True, cfonly=True)
add("spacer")
add("subband", label="Retained earnings reconciliation")
add("data", "re_beg", "  Retained earnings, beginning of year", style="re_beg")
add("data", "re_ni",  "  Net income for the year", style="formula",
    formula="{c}%NI")
add("data", "re_div", "  Dividends", style="input")
add("data", "re_end", "Retained earnings, end of year", style="subtotal",
    formula="{c}%REBEG+{c}%RENI-{c}%REDIV", bold=True, top=True)
add("spacer")
add("subband", label="Growth & margins")
add("data", "g_rev",  "  Revenue y/y %", style="pct",
    formula="IFERROR({c}%REV/{p}%REV-1,\"\")", yoy=True)
add("data", "g_gm",   "  Gross margin y/y %", style="pct",
    formula="IFERROR({c}%GROSS/{p}%GROSS-1,\"\")", yoy=True)
add("data", "g_ni",   "  Net income y/y %", style="pct",
    formula="IFERROR({c}%NI/{p}%NI-1,\"\")", yoy=True)
add("data", "m_ga",   "  G&A % of revenue", style="pct",
    formula="IFERROR({c}%GATOT/{c}%REV,0)")
add("data", "m_net",  "  Net income margin %", style="pct",
    formula="IFERROR({c}%NI/{c}%REV,0)")
add("spacer")

# ---- COST OF SALES SCHEDULE ----
add("band", label="COST OF SALES SCHEDULE")
add("subband", label="Direct costs")
add("data", "cs_wages",   "  Direct wages", style="input")
add("data", "cs_mat",     "  Materials, supplies and miscellaneous", style="input")
add("data", "cs_sub",     "  Subcontracts", style="input")
add("data", "cs_direct",  "Total direct costs", style="subtotal",
    formula="SUM({c}%CSD_FIRST:{c}%CSD_LAST)", bold=True, top=True)
add("subband", label="Direct costs as % of revenue")
add("data", "csp_wages", "  Direct wages % of revenue", style="pct",
    formula="IFERROR({c}%CSWAGES/{c}%REV,0)")
add("data", "csp_mat",   "  Materials, supplies and miscellaneous % of revenue", style="pct",
    formula="IFERROR({c}%CSMAT/{c}%REV,0)")
add("data", "csp_sub",   "  Subcontracts % of revenue", style="pct",
    formula="IFERROR({c}%CSSUB/{c}%REV,0)")
add("data", "csp_total", "  Total direct costs % of revenue", style="pct",
    formula="IFERROR({c}%CSDIR/{c}%REV,0)", bold=True)
add("subband", label="Indirect costs")
add("data", "cs_benefits","  Employee group benefits", style="input")
add("data", "cs_wsib",    "  WSIB", style="input")
add("data", "cs_ins",     "  Insurance and property tax", style="input")
add("data", "cs_util",    "  Utilities", style="input")
add("data", "cs_rm_eq",   "  Repairs and maintenance — equipment", style="input")
add("data", "cs_rm_bld",  "  Repairs and maintenance — buildings", style="input")
add("data", "cs_freight", "  Freight and duty", style="input")
add("data", "cs_truck",   "  Truck and travel", style="input")
add("data", "cs_amort",   "  Amortization — building and shop equipment", style="input",
    note="Non-cash; added back in cash flow.")
add("data", "cs_indirect","Total indirect costs", style="subtotal",
    formula="SUM({c}%CSI_FIRST:{c}%CSI_LAST)", bold=True, top=True)
add("data", "cs_total",   "Total cost of sales", style="keytotal",
    formula="{c}%CSDIR+{c}%CSIND", bold=True, top=True)
add("spacer")

# ---- BALANCE SHEET ----
add("band", label="BALANCE SHEET")
add("subband", label="Assets — current")
add("data", "bs_cash",   "  Cash in bank", style="input")
add("data", "bs_ar",     "  Accounts receivable", style="input")
add("data", "bs_gst",    "  Government sales tax recoverable", style="input")
add("data", "bs_dfd",    "  Due from director", style="input")
add("data", "bs_dfkel",  "  Due from KEL Tooling, non-interest bearing", style="input")
add("data", "bs_ca",     "Total current assets", style="subtotal",
    formula="SUM({c}%CA_FIRST:{c}%CA_LAST)", bold=True, top=True)
add("data", "bs_dfrg",   "Due from Related Group of Companies", style="input")
add("data", "bs_ppe",    "Land, buildings and equipment, net", style="input",
    note="Net of accumulated amortization.")
add("data", "bs_ta",     "TOTAL ASSETS", style="keytotal",
    formula="{c}%CA+{c}%DFRG+{c}%PPE", bold=True, top=True, dbl=True)
add("spacer")
add("subband", label="Liabilities — current")
add("data", "bs_loan",   "  Bank loan", style="input")
add("data", "bs_ap",     "  Accounts payable and accrued liabilities", style="input")
add("data", "bs_payroll","  Payroll deductions payable", style="input")
add("data", "bs_taxpay", "  Income tax payable", style="input")
add("data", "bs_dtkel",  "  Due to KEL Tooling, non-interest bearing", style="input")
add("data", "bs_dtd",    "  Due to director", style="input")
add("data", "bs_cl",     "Total current liabilities", style="subtotal",
    formula="SUM({c}%CL_FIRST:{c}%CL_LAST)", bold=True, top=True)
add("data", "bs_fitl",   "Future income tax liability", style="input")
add("data", "bs_tl",     "Total liabilities", style="subtotal",
    formula="{c}%CL+{c}%FITL", bold=True, top=True)
add("subband", label="Shareholders' equity")
add("data", "bs_sc",     "  Share capital", style="input")
add("data", "bs_re",     "  Retained earnings", style="formula_link",
    note="Links to retained earnings, end of year (Income Statement).")
add("data", "bs_te",     "Total shareholders' equity", style="subtotal",
    formula="{c}%SC+{c}%BSRE", bold=True, top=True)
add("data", "bs_tle",    "TOTAL LIABILITIES AND EQUITY", style="keytotal",
    formula="{c}%TL+{c}%TE", bold=True, top=True, dbl=True)
add("data", "bs_check",  "Balance check (Total assets − Total liabilities & equity)", style="check",
    formula="{c}%TA-{c}%TLE")
add("spacer")

# ---- IMPLIED CASH FLOW STATEMENT ----
add("band", label="CASH FLOW STATEMENT (implied from balance sheet & income statement)")
add("subband", label="Operating activities")
add("data", "cf_ni",    "  Net income for the year", style="cf",
    formula="{c}%NI")
add("data", "cf_amoff", "  Add: amortization — office equipment", style="cf",
    formula="{c}%GAAMORT")
add("data", "cf_ambld", "  Add: amortization — building and shop equipment", style="cf",
    formula="{c}%CSAMORT")
add("data", "cf_fitl",  "  Increase / (decrease) in future income tax liability", style="cf",
    formula="{c}%FITL-{p}%FITL")
add("data", "cf_ar",    "  (Increase) / decrease in accounts receivable", style="cf",
    formula="-({c}%AR-{p}%AR)")
add("data", "cf_gst",   "  (Increase) / decrease in government sales tax recoverable", style="cf",
    formula="-({c}%GST-{p}%GST)")
add("data", "cf_ap",    "  Increase / (decrease) in accounts payable and accrued liabilities", style="cf",
    formula="{c}%AP-{p}%AP")
add("data", "cf_payroll","  Increase / (decrease) in payroll deductions payable", style="cf",
    formula="{c}%PAYROLL-{p}%PAYROLL")
add("data", "cf_taxpay","  Increase / (decrease) in income tax payable", style="cf",
    formula="{c}%TAXPAY-{p}%TAXPAY")
add("data", "cf_cfo",   "Net cash from operating activities", style="cf_subtotal",
    formula="SUM({c}%CFO_FIRST:{c}%CFO_LAST)", bold=True, top=True)
add("subband", label="Investing activities")
add("data", "cf_capex", "  Purchase of land, buildings and equipment (implied)", style="cf",
    formula="-(({c}%PPE-{p}%PPE)+{c}%GAAMORT+{c}%CSAMORT)",
    note="Implied capex = Δ net PP&E + total amortization. A positive figure here is a cash inflow (net disposals > additions), as in FY2023.")
add("data", "cf_cfi",   "Net cash from investing activities", style="cf_subtotal",
    formula="SUM({c}%CFI_FIRST:{c}%CFI_LAST)", bold=True, top=True)
add("subband", label="Financing activities")
add("data", "cf_loan",  "  Increase / (decrease) in bank loan", style="cf",
    formula="{c}%LOAN-{p}%LOAN")
add("data", "cf_dtd",   "  Increase / (decrease) in due to director", style="cf",
    formula="{c}%DTD-{p}%DTD")
add("data", "cf_dtkel", "  Increase / (decrease) in due to KEL Tooling", style="cf",
    formula="{c}%DTKEL-{p}%DTKEL")
add("data", "cf_dfrg",  "  (Increase) / decrease in due from Related Group of Companies", style="cf",
    formula="-({c}%DFRG-{p}%DFRG)", note="Related-party advances — reclassified to financing.")
add("data", "cf_dfd",   "  (Increase) / decrease in due from director", style="cf",
    formula="-({c}%DFD-{p}%DFD)")
add("data", "cf_dfkel", "  (Increase) / decrease in due from KEL Tooling", style="cf",
    formula="-({c}%DFKEL-{p}%DFKEL)")
add("data", "cf_div",   "  Dividends paid", style="cf",
    formula="-{c}%REDIV")
add("data", "cf_sc",    "  Issuance of share capital", style="cf",
    formula="{c}%SC-{p}%SC")
add("data", "cf_cff",   "Net cash from financing activities", style="cf_subtotal",
    formula="SUM({c}%CFF_FIRST:{c}%CFF_LAST)", bold=True, top=True)
add("data", "cf_netchg","Net change in cash", style="keytotal",
    formula="{c}%CFO+{c}%CFI+{c}%CFF", bold=True, top=True)
add("data", "cf_beg",   "Cash, beginning of year", style="cf",
    formula="{p}%CASH")
add("data", "cf_end",   "Cash, end of year", style="cf_subtotal",
    formula="{c}%CFBEG+{c}%CFNET", bold=True, top=True)
add("data", "cf_tie",   "Tie-out check (implied ending cash − balance sheet cash)", style="check",
    formula="{c}%CFEND-{c}%CASH")
add("spacer")

# ---- AMORTIZATION & PP&E SCHEDULE ----
add("band", label="AMORTIZATION & PP&E SCHEDULE")
add("subband", label="Amortization (depreciation) by category")
add("data", "am_off",   "  Amortization — office equipment (G&A)", style="formula",
    formula="{c}%GAAMORT")
add("data", "am_bld",   "  Amortization — building and shop equipment (COGS)", style="formula",
    formula="{c}%CSAMORT")
add("data", "am_total", "Total amortization (D&A)", style="subtotal",
    formula="{c}%GAAMORT+{c}%CSAMORT", bold=True, top=True)
add("subband", label="PP&E rollforward (net book value)")
add("data", "pp_beg",   "  PP&E, beginning of year", style="formula",
    formula="{p}%PPE", cfonly=True)
add("data", "pp_capex", "  Add: capital expenditures (implied)", style="formula",
    formula="-{c}%CFCAPEX", cfonly=True)
add("data", "pp_amort", "  Less: total amortization", style="formula",
    formula="-({c}%GAAMORT+{c}%CSAMORT)", cfonly=True)
add("data", "pp_end",   "  PP&E, end of year", style="subtotal",
    formula="{c}%PPBEG+{c}%PPCAPEX+{c}%PPAMORT", bold=True, top=True, cfonly=True)
add("data", "pp_check", "  Check (rollforward − balance sheet PP&E)", style="check",
    formula="{c}%PPEND-{c}%PPE", cfonly=True)
add("subband", label="Implied useful-life analysis")
add("data", "ul_net",   "  Net PP&E, end of year", style="formula",
    formula="{c}%PPE")
add("data", "ul_da",    "  Total amortization (D&A)", style="formula",
    formula="{c}%AMTOT")
add("data", "ul_rate",  "  Implied depreciation rate (D&A / beginning net PP&E)", style="pct",
    formula="IFERROR({c}%AMTOT/{p}%PPE,\"\")", cfonly=True,
    note="Annual amortization as a % of opening net book value.")
add("data", "ul_life",  "  Implied remaining useful life (net PP&E / D&A)", style="years",
    formula="IFERROR({c}%PPE/{c}%AMTOT,\"\")",
    note="Years of amortization remaining at the current run-rate on the net book value.")
add("data", "ul_capint","  Capex intensity (implied capex / D&A)", style="ratio",
    formula="IFERROR(-{c}%CFCAPEX/{c}%AMTOT,\"\")", cfonly=True,
    note="Reinvestment ratio: >1x grows the asset base, <1x shrinks it.")
add("data", "ul_caprev","  Capex % of revenue", style="pct",
    formula="IFERROR(-{c}%CFCAPEX/{c}%REV,\"\")", cfonly=True)
add("spacer")

# ---- FREE CASH FLOW & RATIO ANALYSIS ----
add("band", label="RATIO ANALYSIS & FREE CASH FLOW")
add("subband", label="Cash generation")
add("data", "fcf",      "  Free cash flow (CFO − capex)", style="formula",
    formula="{c}%CFO+{c}%CFCAPEX")
add("data", "fcf_conv", "  FCF conversion (FCF / net income)", style="ratio",
    formula="IFERROR(({c}%CFO+{c}%CFCAPEX)/{c}%NI,\"\")", cfonly=True)
add("data", "capex_rev","  Capex % of revenue", style="pct",
    formula="IFERROR(-{c}%CFCAPEX/{c}%REV,\"\")", cfonly=True)
# --- Returns (mirrors the template's ROIC decomposition) ---
add("subband", label="Returns")
add("data", "ret_ebit",   "Operating earnings (EBIT)", style="formula",
    formula="{c}%NIBT+{c}%GABANK")
add("data", "ret_taxrate","Normalised tax rate", style="pct",
    formula="{c}%TAXRATE")
add("data", "ret_nopat",  "NOPAT = EBIT × (1 − tax rate)", style="formula",
    formula="{c}%RETEBIT*(1-{c}%RETTAX)")
add("data", "ret_ic",     "Invested capital (Equity + Net debt)", style="formula",
    formula="{c}%RETEQ+{c}%RETND")
add("data", "ret_eq",     "  Total equity", style="formula",
    formula="{c}%TE")
add("data", "ret_nd",     "  Net debt (bank loan − cash)", style="formula",
    formula="{c}%LOAN-{c}%CASH")
add("data", "ret_roic",   "ROIC = NOPAT / Invested capital", style="pct",
    formula="IFERROR({c}%RETNOPAT/{c}%RETIC,0)", bold=True, top=True)
add("data", "ret_opm",    "  Operating margin", style="pct",
    formula="IFERROR({c}%RETEBIT/{c}%REV,0)")
add("data", "ret_turn",   "  Capital turnover (Sales / IC)", style="ratio",
    formula="IFERROR({c}%REV/{c}%RETIC,0)")
add("data", "ret_burden", "  Tax burden (1 − tax rate)", style="pct",
    formula="1-{c}%RETTAX")
add("data", "ret_check",  "  Check: Margin × Turnover × Tax burden", style="pct",
    formula="{c}%RETOPM*{c}%RETTURN*{c}%RETBURDEN")
# --- RONTA decomposition ---
add("subband", label="RONTA decomposition")
add("data", "ron_ebit",  "EBIT before special items", style="formula",
    formula="{c}%RETEBIT")
add("data", "ron_nopat", "NOPAT = EBIT × (1 − tax rate)", style="formula",
    formula="{c}%RONEBIT*(1-{c}%RETTAX)")
add("data", "ron_lbl",   "Net tangible assets (NTA):", style="label")
add("data", "ron_ta",    "  Total assets", style="formula",
    formula="{c}%TA")
add("data", "ron_intang","  − Intangible assets (none)", style="formula",
    formula="0")
add("data", "ron_nibcl", "  − Non-interest-bearing current liabilities", style="formula",
    formula="-({c}%CL-{c}%LOAN)", note="Current liabilities excluding the interest-bearing bank loan.")
add("data", "ron_cash",  "  − Cash", style="formula",
    formula="-{c}%CASH")
add("data", "ron_nta",   "  = Net tangible assets", style="subtotal",
    formula="{c}%RONTATA+{c}%RONINTANG+{c}%RONNIBCL+{c}%RONCASH", bold=True, top=True)
add("data", "ron_ronta", "RONTA = NOPAT / NTA", style="pct",
    formula="IFERROR({c}%RONNOPAT/{c}%RONNTA,0)", bold=True)
# --- Return on Equity ---
add("subband", label="Return on Equity (ROE)")
add("data", "roe_ni", "Net income (profit for the year)", style="formula",
    formula="{c}%NI")
add("data", "roe_eq", "Total equity", style="formula",
    formula="{c}%TE")
add("data", "roe",    "ROE = Net income / Equity", style="pct",
    formula="IFERROR({c}%ROENI/{c}%ROEEQ,0)", bold=True)
add("data", "roa",    "Return on assets (NI / total assets)", style="pct",
    formula="IFERROR({c}%NI/{c}%TA,0)")
# --- Leverage ---
add("subband", label="Leverage")
add("data", "lev_nd",     "Net debt (bank loan − cash)", style="formula",
    formula="{c}%LOAN-{c}%CASH")
add("data", "lev_ebitda", "EBITDA (Operating earnings + D&A)", style="formula",
    formula="{c}%RETEBIT+{c}%GAAMORT+{c}%CSAMORT")
add("data", "lev_nde",    "Net debt / EBITDA  (x)", style="ratio",
    formula="IFERROR({c}%LEVND/{c}%LEVEBITDA,\"\")")
add("data", "lev_cov",    "Interest coverage (EBIT / interest expense)", style="ratio",
    formula="IFERROR({c}%RETEBIT/{c}%GABANK,\"\")")
add("data", "lev_de",     "Debt / equity (bank loan / equity)", style="ratio",
    formula="IFERROR({c}%LOAN/{c}%TE,0)")
# --- Liquidity & working capital ---
add("subband", label="Liquidity & working capital")
add("data", "r_cr",   "Current ratio (CA / CL)", style="ratio",
    formula="IFERROR({c}%CA/{c}%CL,\"\")")
add("data", "r_qr",   "Quick ratio ((cash + A/R) / CL)", style="ratio",
    formula="IFERROR(({c}%CASH+{c}%AR)/{c}%CL,\"\")")
add("data", "r_dso",  "DSO (A/R / revenue × 365)", style="days",
    formula="IFERROR({c}%AR/{c}%REV*365,\"\")")
add("data", "r_dpo",  "DPO (A/P / cost of sales × 365)", style="days",
    formula="IFERROR({c}%AP/{c}%COGS*365,\"\")")

# ----------------------------------------------------------------------------
# Pass 1: assign row numbers
# ----------------------------------------------------------------------------
r = 2
for e in layout:
    e["row"] = r
    if e["kind"] == "data":
        R[e["key"]] = r
    r += 1

# named ranges for SUM blocks (first/last inclusive)
R["GA_FIRST"], R["GA_LAST"] = R["ga_mgmt"], R["ga_acct"]
R["CSD_FIRST"], R["CSD_LAST"] = R["cs_wages"], R["cs_sub"]
R["CSI_FIRST"], R["CSI_LAST"] = R["cs_benefits"], R["cs_amort"]
R["CA_FIRST"], R["CA_LAST"] = R["bs_cash"], R["bs_dfkel"]
R["CL_FIRST"], R["CL_LAST"] = R["bs_loan"], R["bs_dtd"]
R["CFO_FIRST"], R["CFO_LAST"] = R["cf_ni"], R["cf_taxpay"]
R["CFI_FIRST"], R["CFI_LAST"] = R["cf_capex"], R["cf_capex"]
R["CFF_FIRST"], R["CFF_LAST"] = R["cf_loan"], R["cf_sc"]

# token -> row key mapping used inside formula templates ({c}%TOKEN)
TOK = {
    "REV": "revenue", "COGS": "cogs", "GROSS": "gross_margin", "OI": "other_income",
    "GMOI": "gm_plus_oi", "GATOT": "ga_total", "NIBT": "nibt", "TAXC": "tax_current",
    "TAXD": "tax_deferred", "TAXTOT": "tax_total", "NI": "net_income",
    "REBEG": "re_beg", "RENI": "re_ni", "REDIV": "re_div",
    "CSDIR": "cs_direct", "CSIND": "cs_indirect", "GAAMORT": "ga_amort",
    "CSAMORT": "cs_amort", "GABANK": "ga_bank",
    "CA": "bs_ca", "DFRG": "bs_dfrg", "PPE": "bs_ppe", "TA": "bs_ta",
    "CL": "bs_cl", "FITL": "bs_fitl", "TL": "bs_tl", "SC": "bs_sc",
    "BSRE": "bs_re", "TE": "bs_te", "TLE": "bs_tle",
    "AR": "bs_ar", "GST": "bs_gst", "AP": "bs_ap", "PAYROLL": "bs_payroll",
    "TAXPAY": "bs_taxpay", "DFD": "bs_dfd", "DFKEL": "bs_dfkel",
    "DTD": "bs_dtd", "DTKEL": "bs_dtkel", "LOAN": "bs_loan", "CASH": "bs_cash",
    "CFO": "cf_cfo", "CFI": "cf_cfi", "CFF": "cf_cff", "CFNET": "cf_netchg",
    "CFBEG": "cf_beg", "CFEND": "cf_end", "CFCAPEX": "cf_capex",
    # income-statement bridge
    "BREBIT": "br_ebit", "BREBITDA": "br_ebitda",
    # cost-of-sales % of revenue
    "CSWAGES": "cs_wages", "CSMAT": "cs_mat", "CSSUB": "cs_sub",
    # amortization & PP&E schedule
    "AMTOT": "am_total", "PPBEG": "pp_beg", "PPCAPEX": "pp_capex",
    "PPAMORT": "pp_amort", "PPEND": "pp_end",
    # returns / RONTA / ROE / leverage
    "TAXRATE": "tax_rate",
    "RETEBIT": "ret_ebit", "RETTAX": "ret_taxrate", "RETNOPAT": "ret_nopat",
    "RETIC": "ret_ic", "RETEQ": "ret_eq", "RETND": "ret_nd",
    "RETOPM": "ret_opm", "RETTURN": "ret_turn", "RETBURDEN": "ret_burden",
    "RONEBIT": "ron_ebit", "RONNOPAT": "ron_nopat", "RONTATA": "ron_ta",
    "RONINTANG": "ron_intang", "RONNIBCL": "ron_nibcl", "RONCASH": "ron_cash",
    "RONNTA": "ron_nta", "ROENI": "roe_ni", "ROEEQ": "roe_eq",
    "LEVND": "lev_nd", "LEVEBITDA": "lev_ebitda",
    # SUM block anchors
    "GA_FIRST": "GA_FIRST", "GA_LAST": "GA_LAST",
    "CSD_FIRST": "CSD_FIRST", "CSD_LAST": "CSD_LAST",
    "CSI_FIRST": "CSI_FIRST", "CSI_LAST": "CSI_LAST",
    "CA_FIRST": "CA_FIRST", "CA_LAST": "CA_LAST",
    "CL_FIRST": "CL_FIRST", "CL_LAST": "CL_LAST",
    "CFO_FIRST": "CFO_FIRST", "CFO_LAST": "CFO_LAST",
    "CFI_FIRST": "CFI_FIRST", "CFI_LAST": "CFI_LAST",
    "CFF_FIRST": "CFF_FIRST", "CFF_LAST": "CFF_LAST",
}

def rownum(token):
    key = TOK[token]
    return R[key]

def build_formula(template, col):
    """Expand a formula template for a given column letter.
    {c}=current col, {p}=prior col, %TOKEN -> <col><rownum>."""
    p = PRIOR.get(col, col)
    out = []
    i = 0
    s = template
    # replace %TOKEN occurrences (they are preceded by {c} or {p})
    # We process by scanning for '{c}%' and '{p}%'
    s = s.replace("{c}%", "\x00c\x00").replace("{p}%", "\x00p\x00")
    # now tokens look like \x00c\x00TOKEN ; split
    result = ""
    j = 0
    while j < len(s):
        if s[j] == "\x00":
            which = s[j+1]
            j += 3  # skip \x00 c \x00
            # read token chars
            tok = ""
            while j < len(s) and (s[j].isalnum() or s[j] == "_"):
                tok += s[j]; j += 1
            colletter = col if which == "c" else p
            result += f"{colletter}{rownum(tok)}"
        else:
            result += s[j]; j += 1
    return "=" + result

# ----------------------------------------------------------------------------
# Input data arrays (FY2021..FY2025)
# ----------------------------------------------------------------------------
DATA = {
    "revenue":     [1125192, 1129492, 1294937, 1481483, 1669107],
    "other_income":[60379, 587, 477, 1956, 605],
    "ga_mgmt":     [160000, 204000, 195352, 166000, 186000],
    "ga_bank":     [1781, 1403, 2222, 1532, 8832],
    "ga_amort":    [11240, 10505, 8563, 7845, 7176],
    "ga_adv":      [553, 1560, 929, 1400, 8035],
    "ga_meals":    [9714, 15156, 19809, 21560, 19662],
    "ga_tele":     [8407, 7686, 8182, 9193, 10163],
    "ga_comp":     [21431, 31221, 18723, 11399, 14170],
    "ga_office":   [12767, 16609, 18360, 13520, 14490],
    "ga_leases":   [684, -210, 750, 547, 1519],
    "ga_consult":  [0, 0, 0, 31645, 37254],
    "ga_iso":      [0, 0, 0, 4308, 5592],
    "ga_acct":     [11697, 14792, 18060, 9252, 8963],
    "tax_current": [28643, 10179, 16792, 30672, 34327],
    "tax_deferred":[0, -5949, -5879, -3362, -849],
    "re_div":      [0, 0, 170000, 0, 0],
    # Cost of sales schedule
    "cs_wages":    [155912, 207400, 124776, 246407, 331037],
    "cs_mat":      [196164, 234466, 423028, 320248, 233181],
    "cs_sub":      [123993, 124289, 156114, 208375, 242821],
    "cs_benefits": [6052, 5039, 9316, 7971, 18149],
    "cs_wsib":     [1840, 1078, 3724, 2996, -47],
    "cs_ins":      [28591, 33771, 29558, 33572, 29258],
    "cs_util":     [21844, 24879, 25382, 25785, 25115],
    "cs_rm_eq":    [29368, 47451, 39527, 13007, 21655],
    "cs_rm_bld":   [35600, 22714, 34750, 50649, 141141],
    "cs_freight":  [23583, 14673, 18134, 23910, 22757],
    "cs_truck":    [4307, 10175, 7805, 11969, 11486],
    "cs_amort":    [89055, 74522, 53208, 45953, 42284],
    # Balance sheet
    "bs_cash":     [301695, 331894, 323927, 451664, 700989],
    "bs_ar":       [174456, 102521, 172686, 283263, 249660],
    "bs_gst":      [10746, 8957, 11269, 4041, 10200],
    "bs_dfd":      [121461, 243508, 0, 0, 6193],
    "bs_dfkel":    [0, 0, 23308, 16108, 8908],
    "bs_dfrg":     [24277, 24277, 12780, 16132, 18139],
    "bs_ppe":      [570987, 513774, 441755, 403824, 502314],
    "bs_loan":     [25000, 11500, 0, 0, 155242],
    "bs_ap":       [246579, 269806, 157001, 148733, 152022],
    "bs_payroll":  [2687, 26638, 6158, 6157, 10936],
    "bs_taxpay":   [37030, 10179, 16792, 30644, 629],
    "bs_dtkel":    [2239, 0, 0, 0, 0],
    "bs_dtd":      [0, 0, 6616, 6616, 0],
    "bs_fitl":     [16039, 10090, 4211, 849, 0],
    "bs_sc":       [1000, 1000, 1000, 1000, 1000],
}
RE_BEG_2021 = 670703   # retained earnings beginning of FY2021 (input; rolls thereafter)

# ----------------------------------------------------------------------------
# Style helpers
# ----------------------------------------------------------------------------
def set_cell(coord, value, *, font_color=BLACK, bold=False, italic=False, sz=11,
             fill=None, numfmt=None, align=None, border=None, wrap=False):
    c = ws[coord]
    c.value = value
    c.font = Font(name="Calibri", size=sz, bold=bold, italic=italic, color=font_color)
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    if numfmt:
        c.number_format = numfmt
    if align or wrap:
        c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    if border:
        c.border = border

# ----------------------------------------------------------------------------
# Pass 2: write everything
# ----------------------------------------------------------------------------
for e in layout:
    row = e["row"]
    kind = e["kind"]
    opts = e["opts"]

    if kind == "title":
        set_cell(f"B{row}", e["label"], font_color=NAVY, bold=True, sz=16)
        continue
    if kind == "subtitle":
        set_cell(f"B{row}", e["label"], font_color=GRAY, sz=10)
        ws.row_dimensions[row].height = 13
        continue
    if kind == "spacer":
        continue
    if kind == "colhdr":
        set_cell(f"B{row}", "(C$)", font_color=WHITE, bold=True, sz=12, fill=SLATE)
        for col, yr in zip(YCOLS, YEARS):
            set_cell(f"{col}{row}", f"FY{yr}", font_color=WHITE, bold=True, sz=12,
                     fill=SLATE, align="center")
        set_cell(f"{CAGR_COL}{row}", "'21–'25 CAGR", font_color=WHITE, bold=True,
                 sz=11, fill=SLATE, align="center")
        set_cell(f"{NOTE_COL}{row}", "Notes & assumptions", font_color=WHITE,
                 bold=True, sz=11, fill=SLATE, align="left")
        # also paint the spacer col A cell for a clean band edge
        set_cell(f"A{row}", None, fill=SLATE)
        continue
    if kind == "band":
        for col in ["A", "B"] + YCOLS + [CAGR_COL, NOTE_COL]:
            set_cell(f"{col}{row}", e["label"] if col == "B" else None,
                     font_color=WHITE, bold=True, sz=12, fill=SLATE, align="left")
        ws.row_dimensions[row].height = 18
        continue
    if kind == "subband":
        for col in ["A", "B"] + YCOLS + [CAGR_COL, NOTE_COL]:
            set_cell(f"{col}{row}", e["label"] if col == "B" else None,
                     font_color=WHITE, bold=True, sz=11, fill=GREEN, align="left")
        continue

    # ---- data row ----
    style = opts.get("style")
    key = e["key"]
    bold = opts.get("bold", False)
    top = opts.get("top", False)
    dbl = opts.get("dbl", False)
    border = TOP_DBL if dbl else (TOP if top else None)

    # plain label-only row (e.g. "Net tangible assets (NTA):")
    if style == "label":
        set_cell(f"B{row}", e["label"], font_color=BLACK, bold=True, sz=12)
        continue

    is_pct = style in ("pct",)
    is_ratio = style == "ratio"
    is_days = style == "days"
    is_years = style == "years"
    numfmt = (PCT if is_pct else RATIO if is_ratio else DAYS if is_days
              else YRS if is_years else MONEY)

    # label styling
    lbl = e["label"]
    lbl_indented = lbl.startswith("  ")
    lbl_color = GRAY_LBL if (lbl_indented and style in ("pct", "ratio", "days", "years")) else BLACK
    highlight = SAGE if style == "keytotal" else None
    if style == "keytotal":
        # label cell shares the sage fill & money format like template EBIT/Total rows
        set_cell(f"B{row}", lbl, font_color=BLACK, bold=True, sz=12, fill=SAGE,
                 numfmt=MONEY, align="left", border=border)
    else:
        set_cell(f"B{row}", lbl, font_color=lbl_color, bold=bold, sz=12, border=border)

    # decide per-column content
    for col in YCOLS:
        coord = f"{col}{row}"
        cf_row = style in ("cf", "cf_subtotal") or opts.get("cfonly")
        # FY2021 has no cash-flow (no prior BS)
        if cf_row and col == "C":
            set_cell(coord, None, numfmt=numfmt, align="right", fill=highlight,
                     bold=bold, border=border, sz=12)
            continue

        if style == "input":
            val = DATA[key][YCOLS.index(col)]
            set_cell(coord, val, font_color=BLUE, numfmt=MONEY, align="right",
                     bold=bold, border=border, sz=12)
        elif style == "re_beg":
            if col == "C":
                set_cell(coord, RE_BEG_2021, font_color=BLUE, numfmt=MONEY,
                         align="right", sz=12)
            else:
                set_cell(coord, f"={PRIOR[col]}{R['re_end']}", font_color=BLACK,
                         numfmt=MONEY, align="right", sz=12)
        elif style == "formula_link":
            if key == "cogs":
                set_cell(coord, f"={col}{R['cs_total']}", font_color=BLACK,
                         numfmt=MONEY, align="right", border=border, sz=12)
            elif key == "bs_re":
                set_cell(coord, f"={col}{R['re_end']}", font_color=BLACK,
                         numfmt=MONEY, align="right", border=border, sz=12)
        else:
            tmpl = opts.get("formula")
            fcolor = BLACK
            fnt_bold = bold
            fill = highlight
            f = build_formula(tmpl, col)
            set_cell(coord, f, font_color=fcolor, bold=fnt_bold, numfmt=numfmt,
                     align="right", fill=fill, border=border, sz=12)

    # CAGR column
    if opts.get("cagr"):
        set_cell(f"{CAGR_COL}{row}",
                 f"=IFERROR((G{row}/C{row})^(1/4)-1,\"\")",
                 font_color=GRAY_LBL, numfmt=PCT, align="center", sz=11)

    # Notes column
    if opts.get("note"):
        set_cell(f"{NOTE_COL}{row}", opts["note"], font_color=GRAY, sz=10,
                 align="left", wrap=True)

# freeze panes below header, right of labels
ws.freeze_panes = "C7"

out = "/home/user/Industrials_Optima/models/RY_Tool_3_Statement_Model.xlsx"
wb.save(out)
print("saved", out)
print("rows used:", r)
