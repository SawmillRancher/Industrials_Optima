"""Build DHR_Model.xlsx from the MLM template (same layout, formatting and process as the WAB / ODFL / LOAR / WCN builds)."""
import sys, copy, os, json
from openpyxl.styles import Font
from fw import *
import fw
import normalize
from spec import build_spec, scen_layout, SCEN_ROWS, SC, LEGACY
from spec2 import build_spec2

OUT = os.path.join(os.path.dirname(BASE), 'DHR_Model.xlsx')
D = normalize.run()
fw.DATA = D
DEALS = json.load(open(os.path.join(BASE, 'data', 'deals.json')))

# ====================================================================== inputs (V0)
def wc_inputs():
    x = D['2025']
    return {'dso': round(x['bs_ar'] / x['is_sales'] * 365), 'dio': round(x['bs_inventories'] / x['is_cogs'] * 365),
            'dpo': round(x['bs_ap'] / x['is_cogs'] * 365), 'acc': round(x['bs_accrued'] / x['is_sales'], 3)}
WC = wc_inputs()
WC['note'] = (f'Forecast inputs at FY2025 levels (DSO {WC["dso"]}d, inventory {WC["dio"]}d, payables {WC["dpo"]}d, accrued expenses '
              f'{WC["acc"]:.1%} of sales). Masimo working capital enters through the Q2/26 balance sheet.')
MAT = DEALS['debt_q2_26']['maturities_by_year_derived_carrying_usd']['value']
AM = {int(a): b for a, b in json.load(open(os.path.join(BASE, 'data', '2025.json')))['extra']['amort_schedule']}

V0 = {
    'out_hdr': 'FY26 outlook — Q2-26 release (21-Jul-26)',
    'out_note': 'Raised adjusted EPS in the Q2-26 release ($8.45–8.60 vs $8.35–8.55 in Q1-26; $8.35–8.50 in Jan-26). Includes Masimo (closed 10-Jun-2026); excludes pending StatLab.',
    'outlook': [
        ('out_core', '  FY26 core sales growth (non-GAAP)', (0.04, 0.035, 0.03), 'pct', 'Guided +3.0% to +4.0% (Q2-26 release). Drives the Q3/Q4-26E core-growth calibration.'),
        ('out_eps', '  FY26 adjusted diluted EPS ($)', (8.60, 8.525, 8.45), 'ps', 'Guided $8.45–8.60 (raised in Q2-26). Drives the Q3/Q4-26E adjusted-margin calibration.'),
        ('bb26', '  2H/26 share repurchases ($m)', (1000, 500, 0), 'm', 'Q2/26: 5.0m shares for $903m; ~32m shares of authorization left. Net debt ~2.6x Adjusted EBITDA after Masimo — Base assumes modest repurchases.'),
    ],
    'points': [
        ('pt_int', '  Interest expense, net ($m)', 310, 'm', 'Q2-26 release: FY26 interest expense, net ~$(310)m (Q3 ~$(115)m) after the Masimo financing.'),
        ('pt_etr', '  Adjusted effective tax rate', 0.17, 'pct', 'Q2-26 release: effective tax rate ~17.0% (non-GAAP supplemental).'),
        ('pt_capex', '  Capital expenditures ($m)', 1150, 'm', '1H/26 $506m; FY25 $1,156m (not guided).'),
        ('pt_q3m', '  Q3/26 adjusted operating profit margin', 0.265, 'pct', 'Q2-26 release: Q3/26 adjusted operating profit margin ~26.5% (Q1/26 guide was ~28.5%; actual 30.2%). Drives the Q3 margin calibration.'),
        ('pt_amq3', '  Q3/26 amortization of acquired intangibles ($m)', 500, 'm', 'Q2-26 release: Q3 ~$(500)m incl. Masimo.'),
        ('pt_am', '  FY26 amortization of acquired intangibles ($m)', 1900, 'm', 'Q2-26 release: FY26 ~$(1,900)m incl. Masimo (10-K schedule pre-Masimo ~$1.7bn).'),
    ],
    'lev': (2.25, 2.0, 1.75),
    'lev_note': 'Danaher de-levers to ~2x after large deals (Cytiva 2020, Aldevron 2021) before resuming buybacks / M&A. Net debt ~2.6x at 26-Jun-26 (model Adjusted EBITDA). Buybacks absorb surplus cash once leverage reaches the scenario floor.',
    'blocks': [
        ('BIO_CORE', 'BIO_CORE (Biotechnology core growth)', 0.02, -0.025, [0.07, 0.075, 0.075, 0.075],
         'Bioprocessing (Cytiva / Pall): consumables recovery after the 2023–24 destock; FY25 +6.5%, 1H/26 +4.7%; long-term HSD.', 'pct', False),
        ('LS_CORE', 'LS_CORE (Life Sciences core growth)', 0.015, -0.02, [0.04, 0.045, 0.045, 0.045],
         'Instruments (pharma / academic capex), reagents (Abcam, IDT), genomic medicines (Aldevron); FY25 −1.5%, 1H/26 +3.0%; guided +3–4% FY26.', 'pct', False),
        ('DX_CORE', 'DX_CORE (Diagnostics core growth, legacy ex Masimo / StatLab)', 0.015, -0.02, [0.04, 0.045, 0.05, 0.05],
         'Clinical diagnostics (Beckman, Radiometer, Leica) MSD; Cepheid respiratory normalising (~$1.6bn FY26 vs ~$1.9bn FY25).', 'pct', False),
        ('BIO_OPM', 'BIO_OPM (Biotechnology adjusted operating margin)', 0.01, -0.015, [0.395, 0.40, 0.405, 0.41],
         'FY25 39.3% (adjusted operating profit $2,867m / sales $7,293m); bioprocessing volume leverage.', 'pct', False),
        ('LS_OPM', 'LS_OPM (Life Sciences adjusted operating margin)', 0.01, -0.015, [0.225, 0.23, 0.235, 0.24],
         'FY25 21.4% ($1,570m / $7,334m); FY24 23.2%; cost actions and Abcam integration.', 'pct', False),
        ('DX_OPM', 'DX_OPM (Diagnostics adjusted operating margin, legacy)', 0.01, -0.015, [0.29, 0.295, 0.30, 0.305],
         'FY25 28.6% ($2,846m / $9,941m); respiratory mix headwind fading.', 'pct', False),
        ('MS_G', 'MS_G (Masimo sales growth)', 0.02, -0.03, [0.08, 0.08, 0.075, 0.075],
         'Company: Masimo high-single-digit core growth long term; 1Q26 standalone +8.5% y/y.', 'pct', False),
        ('ST_G', 'ST_G (StatLab sales growth)', 0.02, -0.03, [0.08, 0.08, 0.08, 0.08], 'Company: StatLab high-single-digit core growth long term.', 'pct', False),
        ('MNA', 'M&A_spend (bolt-on acquisitions, $m)', 1500, -1000, [1000, 2000, 2500, 2500],
         'Danaher deployed ~$10bn on Masimo in 2026 (plus StatLab). Base resumes bolt-ons as leverage falls toward ~2x.', 'm', True),
    ],
    'qe_core': {'bio': {3: 0.05, 4: 0.07}, 'ls': {3: 0.035, 4: 0.04}, 'dx': {3: 0.0, 4: 0.045}},
    'qe_core_note': {
        'bio': 'Q2-26 release segment guidance: Q3 +MSD; FY +MSD. Q4 input +7% (bioprocessing; Q4-25 comparison +6%).',
        'ls': 'Q2-26 release: Q3 +3.0–4.0%; FY +3.0–4.0%. Q4 input +4%.',
        'dx': 'Q2-26 release: Q3 flat (respiratory +2.5 pt headwind); FY up slightly; Q4 respiratory flat → input +4.5%.',
    },
    'qe_fx': {3: -0.01, 4: -0.007},
    'ms_g26': {3: 0.08, 4: 0.08}, 'ms_m26': {3: 0.23, 4: 0.25}, 'ms_m': {2027: 0.26, 2028: 0.265, 2029: 0.27, 2030: 0.275},
    'ms_m_note': ('Masimo standalone FY25 GAAP operating margin 20.3% (MASI 10-K, continuing operations); Danaher target EBITDA >$530m in 2027 under its '
                  'ownership (Feb-2026 deal release) → ~26% adjusted operating margin + depreciation. Synergies added separately.'),
    'ms_csyn': {2027: 30.0, 2028: 60.0, 2029: 90.0, 2030: 115.0}, 'ms_rsyn': {2027: 10.0, 2028: 25.0, 2029: 40.0, 2030: 50.0},
    'ms_inv': {3: 30.0, 4: 0.0},
    'ms_inv_note': 'Q2/26: $108m (inventory step-up $46m, transaction costs $34m, share-based / change-in-control payments). Q3/26E input: remaining inventory step-up (estimate; total step-up not disclosed).',
    'ms_dep': 0.02,
    'st_rev26': 270.0, 'st_rev_note': 'Q2-26 release: StatLab ~$250m revenue in FY2025; 2026E pro forma at +8%.',
    'st_m': 0.30, 'st_m_note': 'Assumption: histology consumables (>85% recurring revenue); ~30% adjusted operating margin before synergies.',
    'st_mult': 7.5, 'st_mult_note': 'ASSUMPTION — purchase price not disclosed in any SEC filing. 7.5x sales (~$2.0bn) for a recurring-revenue consumables business; update when disclosed.',
    'st_price_note': 'Purchase price = 2026E pro forma sales × EV/sales input × closing switch. Paid in 2026E (funded with commercial paper).',
    'mna_note': 'Scenario lever M&A_spend. Danaher acquisitions: Cytiva $20.7bn (2020), Aldevron $9.6bn (2021), Abcam $5.6bn (2023), Masimo $9.8bn (2026).',
    'corp_q': -90.0, 'c_pct': -0.013,
    'core_guid_note': 'Company FY26 core sales growth guidance +3.0% to +4.0% (Q2-26 release, 21-Jul-2026).',
    'eps_guid_note': 'Company FY26 adjusted diluted EPS guidance $8.45–8.60 (raised in the Q2-26 release). Drives the adjusted-margin calibration (StatLab excluded).',
    'dep_pct': 0.031, 'sga_pct': 0.312, 'rnd_pct': 0.065, 'sbc_pct': 0.012,
    'amort_sched': {y: AM[y] for y in (2027, 2028, 2029, 2030)},
    'amort_note': ('Q3/Q4-26E: amortization guidance (Q3 ~$500m, FY ~$1.9bn) less Masimo. 2027E+: 2025 10-K schedule of existing intangibles '
                   f'(2027 ${AM[2027]:,.0f}m, 2028 ${AM[2028]:,.0f}m, 2029 ${AM[2029]:,.0f}m, 2030 ${AM[2030]:,.0f}m; published in $bn to one decimal).'),
    'intinc_q': 15.0, 'etr': 0.17, 'item_tax': {3: 0.18, 4: 0.18, 2027: 0.18, 2028: 0.18, 2029: 0.18, 2030: 0.18},
    'dps_g': 0.10, 'def_tax_2h': -200.0, 'def_tax_ae': -400.0, 'ofin_ae': -130.0, 'stock_ae': 100.0, 'oinv_ae': -100.0,
    'capex_pct': 0.045, 'wc': WC, 'min_cash': 2000.0, 'dil': 2.3,
    'mat': {2026: -float(MAT['2026']), 2027: -float(MAT['2027']), 2028: -float(MAT['2028']), 2029: -float(MAT['2029']), 2030: -float(MAT['2030'])},
    'mat_note': ('2H/26: €800m 2.1% notes (30-Sep-26, $912m; repaid). 2027: CHF / yen / €600m notes ($1,181m); 2028: €1.3bn 0.45% Biopharma, €500m floating, CHF ($2,249m); '
                 '2029: CHF / $800m 2.6% ($1,198m); 2030: €750m 2.5% + €750m 3.25% ($1,760m). Carrying amounts at 26-Jun-26 (Q2-26 10-Q Note 10).'),
    'refi_pct': {2026: 0.0, 2027: 1.0, 2028: 1.0, 2029: 1.0, 2030: 1.0},
    'debt_note': ('Debt at 26-Jun-26 $26,558m: euro commercial paper $4,703m (2.5%) + senior notes $21,855m (EUR / CHF / JPY / USD, coupons 0.2–4.4%). '
                  'Masimo financing: €3.0bn notes (Apr-26), CHF 2.38bn private placement (Jun-26), euro CP. $5bn revolver (2028) + $5bn 364-day facility undrawn.'),
    'r_notes': {2027: 0.021, 2028: 0.022, 2029: 0.024, 2030: 0.025},
    'r_notes_note': 'Blended coupon ~1.9% on the 26-Jun-26 notes (low-coupon EUR / CHF paper + 2026 Masimo issues at 1.65–4.0%); drifts up as 0.2–1.2% notes refinance at ~3%.',
    'r_fac': {2027: 0.025, 2028: 0.025, 2029: 0.025, 2030: 0.025},
    'px': 221.21,
    'px_note': 'VALUATION — share price input $221.21 (NYSE close 5-Oct-2026); update to market.',
}

# ====================================================================== compute engine
def hist_value(row, col):
    kw = row.kw; path = kw.get('data')
    if not path: return None
    if kw.get('annual_only') and col.kind != 'A': return None
    v = D.get(pkey(col), {}).get(path)
    return v if kw.get('textdata') else num(v)

def first(kw, keys):
    for k in keys:
        if k in kw: return k
    return None

def compute(row, col):
    kw = row.kw; k = col.kind
    CTX.col = col
    if kw.get('ann') and col.q: return None, False
    if k in HISTK:
        v = hist_value(row, col)
        if v is not None: return v, True
        for key in ('hist', 'all'):
            if key in kw:
                return kw[key](), False
        return None, False
    if k == 'QE':
        if row.spec2: return None, False
        if 'qe_in' in kw:
            v = kw['qe_in']; return (v[col.q] if isinstance(v, dict) else v), True
        key = first(kw, ('qe', 'fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    if k == 'E26':
        if 'e26_in' in kw: return kw['e26_in'], True
        if 'e26' in kw: return kw['e26'](), False
        if kw.get('avg') and ('qe' in kw or 'qe_in' in kw): return f'=AVERAGE(BW{row.r},BX{row.r},BY{row.r},BZ{row.r})', False
        if kw.get('flow') and not row.spec2 and any(x in kw for x in ('qe', 'qe_in', 'fc')): return S26(row.key), False
        key = first(kw, ('fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    if k == 'AE':
        if 'ae_in' in kw:
            v = kw['ae_in']; return (v[col.year] if isinstance(v, dict) else v), True
        key = first(kw, ('ae', 'fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    return None, False

# ====================================================================== writers
CAGR_G = {'bio_rev', 'ls_rev', 'dx_rev', 'g_rev', 'g_adjop', 'is_rev', 'is_gp', 'is_sga', 'is_rnd', 'is_op', 'is_nicont', 'is_dda', 'is_amort',
          'eps_cont', 'dps', 'r_ebitda', 'r_adj', 'r_ebit', 'o_adj', 'e_adjni', 'e_eps', 'cf_cfo', 'fcf', 'fcf_ps', 'bs_ta', 'bs_weq', 'sh_dil',
          'bio_adjop', 'ls_adjop', 'dx_adjop'}
CAGR_A = {'bio_adjm', 'ls_adjm', 'dx_adjm', 'g_adjm', 'is_gm', 'is_opm', 'is_etr', 'r_adj_m', 'r_ebit_m', 'o_adj_m', 'gm_rev', 'gm_core', 'gm_gm',
          'gm_adjm', 'gm_ebitda_m', 'gm_op_m', 'gm_net_m', 'gm_sga', 'gm_rnd', 'fcf_m', 'fcf_conv', 'capex_pct', 'ra_roic', 'rn_ronta', 'roe',
          'payout', 'd_dep_pct', 'd_sga', 'd_rnd', 'd_sbc', 'g_core', 'bio_core', 'ls_core', 'dx_core', 'lv_x', 'wc_nwc_p', 'e_etr'}
ANN_HIST = [ANN[y].c for y in range(2006, 2026)]

def style_cagr(ws, r, on):
    src = 16 if on else 12
    for c in ('CH', 'CI', 'CJ', 'CK'):
        ws[f'{c}{r}']._style = copy.copy(TM[f'{c}{src}']._style)

def write_row(ws, row):
    r = row.r; arch = ARCH[row.arch]
    clone_row_style(ws, arch, r)
    ws.cell(r, 2).value = row.label or None
    if row.kw.get('note'): ws[f'CF{r}'] = row.kw['note']
    if row.kw.get('group'):
        ws.row_dimensions[r].outlineLevel = 1; ws.row_dimensions[r].hidden = True
    style_cagr(ws, r, row.key in CAGR_G or row.key in CAGR_A)
    if row.arch in ('section', 'block', 'blank', 'sub') and not row.kw.get('data'):
        return
    if row.key in DEFS:
        write_defs(ws, row); return
    nf = NF[row.nf]
    for col in COLS:
        v, inp = compute(row, col)
        if v is None: continue
        cell = ws[f'{col.c}{r}']
        cell.value = v
        if not row.kw.get('textdata'): cell.number_format = nf
        if inp:
            set_color(cell, BLUE, bold=True if col.kind in FCK else None)
        elif row.arch == 'growth' and col.kind in HISTK:
            set_color(cell, GRAY)
        else:
            set_color(cell, BLACK)
        add_cell_comment(row, col, cell)
    if row.key in CAGR_G:
        ws[f'CH{r}'] = f'=IFERROR((CE{r}/BV{r})^(1/5)-1,"n/m")'
        ws[f'CI{r}'] = f'=IFERROR((BV{r}/AW{r})^(1/5)-1,"n/m")'
        ws[f'CJ{r}'] = f'=IFERROR((BV{r}/X{r})^(1/10)-1,"n/m")'
        ws[f'CK{r}'] = f'=IFERROR((BV{r}/C{r})^(1/19)-1,"n/m")'
        for c in ('CH', 'CI', 'CJ', 'CK'): ws[f'{c}{r}'].number_format = '0%;\\(0%\\);"n/m"'
    elif row.key in CAGR_A:
        ws[f'CH{r}'] = f'=IFERROR(AVERAGE(CA{r},CB{r},CC{r},CD{r},CE{r}),"n/m")'
        ws[f'CI{r}'] = f'=IFERROR(AVERAGE({",".join(f"{ANN[y].c}{r}" for y in range(2021, 2026))}),"n/m")'
        ws[f'CJ{r}'] = f'=IFERROR(AVERAGE({",".join(f"{ANN[y].c}{r}" for y in range(2016, 2026))}),"n/m")'
        ws[f'CK{r}'] = f'=IFERROR(AVERAGE({",".join(f"{c}{r}" for c in ANN_HIST)}),"n/m")'
        for c in ('CH', 'CI', 'CJ', 'CK'): ws[f'{c}{r}'].number_format = NF['pct']
    t = LABEL_SOURCES.get(row.key)
    if t: comment(ws.cell(r, 2), t)

def raw(col): return normalize.RAW.get(pkey(col)) or {}
def add_cell_comment(row, col, cell):
    if col.kind not in HISTK: return
    k = pkey(col); x = D.get(k) or {}
    if row.key == 'is_rev' and k in normalize.BASIS:
        comment(cell, 'Hybrid recast basis — ' + normalize.BASIS[k] + ' Balance sheet and cash flow as originally reported.')
    elif row.key in ('o_oth', 'e_oth', 'e_nonop') and isinstance(cell.value, (int, float)) and abs(cell.value) > 0.05:
        items = (raw(col).get('ngaap') or {}).get('item_amounts') or []
        want = 'nonop' if row.key == 'e_nonop' else 'op'
        rows = [f'{a[0]}: {a[1]} pretax / {a[2] if len(a) > 2 else "n/a"} after tax' for a in items
                if isinstance(a, list) and len(a) > 3 and str(a[3]) == want and 'amortization' not in str(a[0]).lower()]
        if rows: comment(cell, 'Company footnotes (release): ' + ' | '.join(rows)[:900])
    elif row.key == 'e_oth2' and isinstance(cell.value, (int, float)) and abs(cell.value) > 0.05:
        if x.get('_resid_recast'):
            comment(cell, 'Recast period: adjusted EPS shown on the recast continuing-operations basis (prior-year comparative in the following year\'s release); '
                    'the item detail is on the original basis, so the difference (adjusting items of the discontinued business) is shown here.')
        else:
            comment(cell, 'Per-share rounding / other per-share-only lines in the company bridge × diluted shares.')
    elif row.key == 'e_pref' and isinstance(cell.value, (int, float)) and abs(cell.value) > 0.05:
        comment(cell, 'Mandatory convertible preferred stock: preferred dividends added back where the company EPS uses the if-converted method (MCPS "as if converted").')
    elif row.key == 'e_pub' and x.get('adj_eps_orig') is not None:
        comment(cell, f'Recast comparative (following year\'s release). As originally published: ${x["adj_eps_orig"]:.2f}.')
    elif row.key in ('bio_oadj', 'ls_oadj', 'dx_oadj') and k in (normalize.MAN.get('period_overrides') or {}):
        src = normalize.MAN['period_overrides'][k].get('_src')
        if src: comment(cell, 'Allocated from the release footnote: ' + src)

SRC = 'Source: Danaher earnings releases (Form 8-K Ex. 99.1) — '
LABEL_SOURCES = {
    'bio_rev': 'Source: 10-K / 10-Q segment note and release segment tables (Biotechnology from Q4-22; 2022 quarters on the Q1-23 recast basis).',
    'bio_core': SRC + '"Sales Growth by Segment, Core Sales Growth by Segment" tables.',
    'bio_amort': 'Source: segment note (amortization, depreciation by segment from 2024; Q4 = FY − 9M where not published).',
    'bio_oadj': SRC + 'Q4-25 segment table (Other Operating Profit Adjustments) and adjusted-EPS footnotes (segment named in each footnote).',
    'ms_sa_rev': 'Source: Masimo Corp. (CIK 0000937556) Forms 10-K / 10-Q via SEC XBRL companyfacts; continuing operations as last reported.',
    'ms_rev': 'Source: Q2-26 10-Q Note 2 (Masimo, closed 10-Jun-2026) and Q2-26 release (acquisition impact on growth: Diagnostics 4.0%).',
    'st_pf': 'Source: Q2-26 release "Pending SLMP LLC (StatLab) Acquisition" (~$250m 2025 revenue; close expected by end-2026).',
    'g_core': SRC + 'core sales growth (non-GAAP) — total company, as published (hybrid-recast periods use the restated comparative).',
    'is_rev': 'Source: Consolidated Statements of Earnings (release / 10-Q / 10-K). Hybrid recast: see the comments on recast periods.',
    'o_adj_pub': SRC + 'Q4-25 release segment table: Adjusted Operating Profit (non-GAAP) for Q4-25, Q4-24, FY25, FY24.',
    'e_pub': SRC + 'Reconciliation of Diluted Net EPS to Adjusted Diluted Net EPS (narrative bridges 2006–2013; tables from 2014).',
    'cf_ni': 'Source: Consolidated Statements of Cash Flows (10-K / 10-Q year-to-date; discrete quarters derived; Q4 = FY − 9M).',
    'bs_cash': 'Source: Consolidated Balance Sheets (10-K / 10-Q).',
    'fcf_pub': SRC + 'free cash flow tables (published from 2018; continuing operations from 2023).',
    'ppa_cons': 'Source: Q2-26 10-Q Note 2 Acquisitions (Masimo preliminary purchase-price allocation).',
    'dps': 'Source: releases / 10-Q / 10-K equity note (split-adjusted: 2-for-1 split June 2010).',
    'lg_rev': 'Source: 10-K / 10-Q segment notes as reported in each period (hybrid recast where restated).',
}

DEFS = {}
def defs_setup():
    q = lambda y, n: QC[(y, n)].c
    DEFS['r_def'] = {'C': ('Model definition (not published by Danaher): continuing net earnings + taxes + net interest + D&A + other operating profit adjustments + nonoperating items = adjusted operating profit + depreciation', None)}
    DEFS['o_def'] = {
        'C': ('Company definition (introduced in the Q4-25 release): operating profit + amortization of intangible assets + other operating profit adjustments; applied by the model to all periods',
              'Danaher published adjusted operating profit for FY24, Q4-24, FY25 and Q4-25 (checked). Before 2015 the other-adjustment detail comes from the adjusted-EPS narrative bridges.'),
    }
    DEFS['e_def'] = {
        'C': ('2006–2014: adjusted EPS excludes discrete items only (restructuring above plan, acquisition fair-value charges, gains, discrete tax); amortization NOT added back', None),
        q(2015, 1): ('Q1-15+: amortization of acquisition-related intangibles added back (after tax)', 'Prior-year comparatives in 2015 releases were restated to the new definition.'),
        q(2016, 2): ('Q2-16+: single tax-effect line; discrete tax adjustments excluded; investment fair-value gains / losses excluded (from 2018)', None),
        q(2019, 1): ('2019–Q2/23: MCPS — preferred dividends / if-converted shares where dilutive', None),
        q(2023, 4): ('Q4-23+: continuing operations after the Veralto spin-off (30-Sep-2023)', None),
    }

def write_defs(ws, row):
    for c, (txt, cm) in DEFS[row.key].items():
        cell = ws[f'{c}{row.r}']; cell.value = txt
        cell.font = Font(name='Calibri', sz=10, b=True, i=True, color='FFC00000')
        if cm: comment(cell, cm)

def write_scen(ws):
    for spec_ in SCEN_ROWS:
        r, kind = spec_[0], spec_[1]
        src = {'hdr': 9, 'out_hdr': 11, 'pt_hdr': 15, 'out': 12, 'pt': 16, 'lev': 28, 'blk': 30, 'yr': 31}[kind]
        for ci in range(CI('CN'), CI('CR') + 1):
            ws.cell(r, ci)._style = copy.copy(TM.cell(src, ci)._style)
        if kind in ('hdr', 'out_hdr', 'pt_hdr'):
            _, _, a, b, c_, d_, e = spec_
            for col, v in zip(['CN', 'CO', 'CP', 'CQ', 'CR'], [a, b, c_, d_, e]):
                if v is not None: ws[f'{col}{r}'] = v
        elif kind == 'out':
            _, _, lab, hi, mid, lo, note, nf = spec_
            ws[f'CN{r}'] = lab
            for col, v in zip(['CO', 'CP', 'CQ'], [hi, mid, lo]):
                ws[f'{col}{r}'] = v; ws[f'{col}{r}'].number_format = NF[nf]; set_color(ws[f'{col}{r}'], BLUE)
            ws[f'CR{r}'] = note
        elif kind == 'pt':
            _, _, lab, _, val, _, note, nf = spec_
            ws[f'CN{r}'] = lab; ws[f'CP{r}'] = val; ws[f'CP{r}'].number_format = NF[nf]; set_color(ws[f'CP{r}'], BLUE); ws[f'CR{r}'] = note
        elif kind == 'lev':
            _, _, lab, a, b, c_, note = spec_
            ws[f'CN{r}'] = lab
            for col, v in zip(['CO', 'CP', 'CQ'], [a, b, c_]):
                ws[f'{col}{r}'] = v; ws[f'{col}{r}'].number_format = NF['x']
            ws[f'CR{r}'] = note
        elif kind == 'blk':
            _, _, lab, db, mid, de, note, nf = spec_
            ws[f'CN{r}'] = lab; ws[f'CO{r}'] = db; ws[f'CP{r}'] = mid; ws[f'CQ{r}'] = de; ws[f'CR{r}'] = note
            for col in ('CO', 'CQ'):
                ws[f'{col}{r}'].number_format = '\\+#,##0;\\-#,##0' if nf == 'm' else '\\+0.0%;\\-0.0%'
        elif kind == 'yr':
            _, _, lab, hdr, base, clamp, _, nf = spec_
            ws[f'CN{r}'] = lab; ws[f'CP{r}'] = base
            if clamp:
                ws[f'CO{r}'] = f'=MAX(0,CP{r}+$CO${hdr})'; ws[f'CQ{r}'] = f'=MAX(0,CP{r}+$CQ${hdr})'
            else:
                ws[f'CO{r}'] = f'=CP{r}+$CO${hdr}'; ws[f'CQ{r}'] = f'=CP{r}+$CQ${hdr}'
            for col in ('CO', 'CP', 'CQ'): ws[f'{col}{r}'].number_format = NF[nf]

def write_header(ws):
    for r in range(1, 8):
        clone_row_style(ws, r, r, 2, CI('CR'))
    ws['CF1'] = 'SCENARIO SWITCH ▼ (Bull / Base / Bear) — drives all scenario-linked forecast drivers'
    ws['CF2'] = 'Base'
    comment(ws['CF2'], 'Scenario toggle. Choose Bull / Base / Bear. Drives: FY26 core-growth and adjusted-EPS outlook end-points (high / mid / low), '
            '2H/26 buybacks, 2027E–2030E segment core growth and adjusted margins, Masimo / StatLab growth, M&A spend and the leverage floor for 2027E+ buybacks '
            '(scenario table CN:CR).')
    ws['B2'] = 'Danaher Corporation (NYSE: DHR) — 3-Statement Financial Model'
    ws['B3'] = ('US GAAP as reported · USD millions (per-share data in USD; shares in millions, split-adjusted) · Source: Danaher Forms 10-K / 10-Q and Form 8-K '
                'earnings releases (Ex. 99.1), deal 8-Ks (Masimo Feb-2026; financing Apr–Jun 2026), SEC EDGAR CIK 0000313616; Masimo standalone history from Masimo Corp. filings '
                '(CIK 0000937556) (investor-relations releases are the same documents as filed on EDGAR)')
    ws['B4'] = ('Basis notes: calendar fiscal year (quarters end on the Friday nearest the calendar quarter end). Hybrid recast: income statement, segments and adjusted EPS of the '
                'periods before each separation use the restated continuing-operations comparatives of the following year\'s filings (Apex JV 2010, Communications 2015, '
                'Fortive 2-Jul-2016, ASU 2017-07 pension recast, Envista Dec-2019, Veralto 30-Sep-2023); balance sheet and cash flow as originally reported. 2-for-1 split Jun-2010. '
                'Segments: Biotechnology / Life Sciences / Diagnostics from Q4-22 (legacy segments in memo rows). Cytiva 31-Mar-2020; Aldevron Aug-2021; Abcam Dec-2023; '
                'Masimo 10-Jun-2026 (in Diagnostics; modelled separately); StatLab pending (assumed closed end-2026). Q3/26E–Q4/26E calibrated to FY26 core-growth and adjusted-EPS guidance.')
    ws['CH5'] = 'Forecast'; ws['CI5'] = 'Historical'; ws['CJ5'] = 'Historical'; ws['CK5'] = 'Historical'
    ws['B6'] = '(USDm)'
    for col in COLS: ws[f'{col.c}6'] = col.label
    ws['CF6'] = 'Modelling Notes'
    for c, a, b in [('CH', '5Y CAGR', "25-'30"), ('CI', '5Y CAGR', "20-'25"), ('CJ', '10Y CAGR', "15-'25"), ('CK', '19Y CAGR', "06-'25")]:
        ws[f'{c}6'] = a; ws[f'{c}7'] = b

def main():
    import openpyxl
    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb['Model']
    blank_style = copy.copy(ws['A1']._style)
    for rowc in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for c in rowc:
            c.value = None; c.comment = None; c._style = copy.copy(blank_style)
    for k in list(ws.row_dimensions.keys()):
        rd = ws.row_dimensions[k]; rd.outlineLevel = 0; rd.hidden = False
    scen_layout(V0)
    build_spec(V0)
    n1 = len(fw.SPEC)
    build_spec2(V0)
    for i, row in enumerate(fw.SPEC): row.spec2 = i >= n1
    last = number_rows(8)
    for r in range(1, last + 1): ws.row_dimensions[r].height = 15.0
    ws.row_dimensions[2].height = 21.0
    defs_setup()
    write_header(ws)
    for row in fw.SPEC: write_row(ws, row)
    write_scen(ws)
    # view: same as the template (frozen at C8, opened scrolled to 2019, 70% zoom)
    ws.freeze_panes = 'C8'
    ws.sheet_view.pane.topLeftCell = 'AR8'
    ws.sheet_view.zoomScale = 70; ws.sheet_view.zoomScaleNormal = 70
    import other_sheets as osh
    osh.do_all(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    osh.inject_charts(OUT)
    print('rows', fw.SPEC[-1].r, 'saved', OUT)
    return wb

RECALC = os.environ.get('XLSX_RECALC', '/root/.claude/skills/synced/0a79f692-77b1-47a2-a1d3-b4d57a310bf9_ec54663f-c361-4f22-b606-54eff84e01bd/xlsx/scripts/recalc.py')
def recalc(path, timeout=900):
    import subprocess
    out = subprocess.run(['python3', RECALC, path, str(timeout)], capture_output=True, text=True, cwd=os.path.dirname(RECALC))
    try:
        return json.loads(out.stdout)
    except Exception:
        return {'raw': out.stdout[-2000:], 'err': out.stderr[-2000:]}

def scenario_snapshots():
    import openpyxl
    import other_sheets as osh
    snaps = {}
    for scen in ('Bull', 'Bear'):
        p = os.path.join(BASE, f'DHR_{scen}.xlsx')
        wb = openpyxl.load_workbook(OUT); wb['Model']['CF2'] = scen; wb.save(p)
        res = recalc(p); print(scen, res.get('status'), res.get('total_errors'))
        v = openpyxl.load_workbook(p, data_only=True)['Bull-Base-Bear']
        snaps[scen] = {(r, c): v[f'{c}{r}'].value for r in range(9, 51) for c in 'CDEFGHIJKL' if v[f'{c}{r}'].value is not None}
        os.remove(p)
    wb = openpyxl.load_workbook(OUT)
    osh.snapshot_values(wb, snaps['Bull'], snaps['Bear'])
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    osh.inject_charts(OUT)

if __name__ == '__main__':
    main()
    if '--snap' in sys.argv: scenario_snapshots()
    if '--recalc' in sys.argv:          # store calculated values (readable in any viewer), as in the WCN build
        print('recalc', recalc(OUT))
