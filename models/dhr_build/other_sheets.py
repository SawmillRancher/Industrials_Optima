"""Rewire the Bull-Base-Bear, DCF (new template layout) and Charts tabs of the MLM template to the DHR Model rows;
drop the template's DCF_old tab; keep the template chart XML (retitled)."""
import re, zipfile, shutil
from openpyxl.worksheet.formula import ArrayFormula
from fw import R, TEMPLATE

MAP = {158: 'is_rev', 166: 'is_op', 164: 'o_oth', 182: 'd_dep', 186: 'sh_dil', 211: 'r_adj', 212: 'r_adj_m', 219: 'o_adj',
       221: 'r_adj_pf', 229: 'e_adjni', 231: 'e_eps', 238: 'gm_core', 245: 'gm_ebitda_m', 247: 'gm_op_m', 260: 'cf_wc', 262: 'cf_capex',
       263: 'cf_acq', 280: 'fcf', 282: 'fcf_m', 290: 'bb_px', 291: 'bb_cash', 292: 'bb_sh', 331: 'bs_nci', 408: 'ra_ic', 411: 'ra_roic',
       422: 'rn_ronta', 428: 'lv_nd', 433: 'v_px', 434: 'v_sh', 12: 'gm_core', 30: 'gm_core', 14: 'gm_rev', 31: 'gm_rev'}
REF = re.compile(r"Model!(\$?)([A-Z]{1,3})(\$?)(\d+)")

def remap(f):
    def rep(m):
        r = int(m.group(4))
        if m.group(2) == 'CF' and r == 2: return m.group(0)          # scenario switch
        if r in MAP: return f'Model!{m.group(1)}{m.group(2)}{m.group(3)}{R[MAP[r]]}'
        raise KeyError(f'unmapped template Model row {r} in {f}')
    return REF.sub(rep, f)

BBB = {
    'B2': 'Danaher Corporation (DHR) — Bull / Base / Bear Case P&L Summary',
    'B3': ('Adjusted figures (company definitions: core sales growth, adjusted operating profit, adjusted diluted EPS incl. add-back of acquired-intangible amortization; '
           'Adjusted EBITDA = model definition) · USD millions · Live panel linked to Model — toggle scenario via Model!CF2. Bull and Bear panels are value snapshots '
           'generated from the Model with the switch set to each case. 2026E includes Masimo from 10-Jun-2026; StatLab assumed closed end-2026.'),
    'B9': 'Sales', 'B10': '   Sales y/y %', 'B11': '   Core sales growth % (non-GAAP)',
    'B12': 'Adjusted EBITDA (model definition)', 'B15': 'Operating Profit (GAAP)', 'B16': '   Operating profit y/y %',
    'B19': 'Adjusted Net Earnings (continuing operations, attributable to common)',
    'B35': 'Adjusted EBITDA — pro forma (2026E incl. Masimo full year)', 'B36': '   Adjusted EBITDA (pro forma) % of sales',
}
for k in list(BBB):
    if k[0] == 'B' and k not in ('B2', 'B3'):
        BBB['N' + k[1:]] = BBB[k]; BBB['Z' + k[1:]] = BBB[k]

DCF = {
    'B2': 'DCF VALUATION — DANAHER CORPORATION (DHR)',
    'D5': 'Assumption (~US 10Y Treasury yield); update to market',
    'D7': 'Life-science tools & diagnostics peers (TMO, A, RVTY, WAT, ABT diagnostics) ~0.8–1.1; ~80% recurring revenue (consumables, service) dampens cyclicality',
    'D9': 'Investment grade (A- / A3); 2026 Masimo financing: €3.0bn notes 3.25–4.0% (floating 2028), CHF 1.65–2.51% private placement; USD marginal cost ~4.5–5.0%',
    'D10': 'Model 2027E+ adjusted effective tax rate (FY26 guidance ~17%)',
    'D12': 'Net debt ~$22bn at 26-Jun-26 (post Masimo) vs ~$156bn market cap ($221 × ~705m shares) → ~12–15% D/V',
    'B18': 'Adjusted operating profit (before acquired-intangible amortization) less other operating profit adjustments',
    'B21': 'plus depreciation (acquired-intangible amortization is excluded from EBIT above, as in adjusted EPS)',
    'B22': '− Capex (PP&E) and acquisitions (M&A lever 2027E–30E; Masimo / StatLab in opening invested capital)',
    'B26': 'Fade-period growth (2031E–35E; 6% → 4.5% linear)',
    'D32': ("Danaher's adjusted operating profit expenses stock-based compensation and excludes ~$1.9bn p.a. of acquired-intangible amortization (largely not tax-deductible, "
            'finite-lived): NOPAT is taxed on adjusted operating profit and only depreciation is added back, so the terminal value carries no perpetual amortization add-back. '
            'Masimo ($9.8bn, closed Jun-26) and StatLab (assumed ~$2.0bn) consideration is excluded from 2026E UFCF and added to opening invested capital.'),
    'D43': ('Opening IC = Model 2025A book invested capital (equity + net debt) + Masimo consideration (cash $9,843m + replacement awards $45m) + StatLab assumed price, '
            'rolled forward with DCF net investment (capex + M&A − depreciation + ΔWC). Terminal incremental ROIC = g / reinvestment rate.'),
    'B48': 'Less: Net debt + NCI (Model 2026E YE, post Masimo & StatLab)',
    'D48': 'DHR-specific: NCI immaterial (~$11m); net debt includes the Masimo financing (euro CP, €/CHF notes) and the assumed StatLab funding.',
    'B50': 'Diluted shares (m, Model 2026E period end)',
    'B70': ('To solve for the market-implied path: Data ▸ What-If Analysis ▸ Goal Seek — set C51 to the current price (C52) by changing C27 (terminal growth) '
            'or the fade-period growth inputs (I26:M26).'),
    'D76': 'Rationale (Danaher characteristics)',
    'D78': ('Model NOPAT CAGR 2026–30E: mid-single-digit core growth (bioprocessing HSD), 50–100 bps p.a. adjusted-margin expansion, Masimo synergies, '
            'amortization roll-off → ~8% NOPAT growth'),
    'D79': ('Consumables-heavy portfolio (~80% recurring) and DBS earn high returns on organic investment (capex ~4–5% of sales); acquired growth (Cytiva, Aldevron, Abcam, Masimo) '
            'earns ~5–7% after tax on purchase price, diluting incremental returns'),
    'D80': 'Installed-base razor / razor-blade franchises (bioprocessing, clinical diagnostics, Cepheid), regulatory switching costs, multi-decade DBS compounding record',
    'B91': '2028E adjusted operating profit less other operating profit adjustments',
    'B96': 'Less: YE2027 net debt + NCI (Model)',
    'B108': 'DHR market-implied EV/NOPAT (on 2028E NOPAT)',
    'B36': 'less depreciation',
}
CHARTS = {
    'B2': 'Danaher Corporation — Historical Operating Performance (2006–2025)',
    'B3': ('All historicals · USDm · Linked to Model. Bars = core sales growth (non-GAAP, as published) vs total sales growth (y/y; hybrid recast basis around the Fortive, '
           'Envista and Veralto separations). Adjusted EBITDA = model definition (adjusted operating profit + depreciation). ROIC / RONTA on Adjusted EBIT (after amortization).'),
    'B8': 'Core sales growth %', 'B9': 'Total sales growth %', 'B12': 'Adjusted EBITDA (model)', 'B13': 'Operating Profit (GAAP)',
}
CHART_TITLES = {
    'Aggregates Revenue Drivers: Shipment Volume vs. Price (ASP), y/y %': 'Sales Growth: Core vs. Total, y/y %',
    'Adjusted EBITDA &amp; Earnings from Operations (USDm)': 'Adjusted EBITDA &amp; Operating Profit (USDm)',
}

def _rewire(ws):
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if isinstance(v, ArrayFormula):
                if 'Model!' in (v.text or ''): c.value = ArrayFormula(v.ref, remap(v.text))
            elif isinstance(v, str) and 'Model!' in v:
                c.value = remap(v)

def do_all(wb):
    if 'DCF_old' in wb.sheetnames: del wb['DCF_old']
    wb['DCF']['D34'] = None                       # MLM-specific opening IC (LNA consideration) — rebuilt below
    for name, text in (('Bull-Base-Bear', BBB), ('DCF', DCF), ('Charts', CHARTS)):
        ws = wb[name]
        _rewire(ws)
        for addr, t in text.items(): ws[addr] = t
    d = wb['DCF']
    d['D34'] = f'=Model!BV{R["ra_ic"]}+Model!BX{R["ppa_cons"]}+Model!BX{R["ppa_stock"]}+Model!CA{R["st_price"]}'
    d['C5'] = 0.0425; d['C6'] = 0.05; d['C7'] = 0.95; d['C9'] = 0.045; d['C10'] = 0.17; d['C12'] = 0.15
    d['I26'] = 0.06; d['M26'] = 0.045; d['C27'] = 0.035
    d['C78'] = 0.08; d['C79'] = 0.12; d['C80'] = 15
    b = wb['Bull-Base-Bear']
    b['H26'] = 25; b['T26'] = 28; b['AF26'] = 20
    from fw import comment
    b['N26'].comment = None
    comment(b['B26'], 'P/E on adjusted EPS (continuing operations, company definition incl. amortization add-back). Danaher has traded at ~22–35x NTM adjusted EPS '
            'over 2016–26 (~26x 2026E at $221.21). Base 25x, Bull 28x, Bear 20x applied to 2030E.')
    comment(b['P40'], f'Snapshot (values) of Model rows {R["bb_px"]}–{R["bb_sh"]} (Share Buyback Schedule) with Model!CF2 = Bull; Bear panel AB40:AF42 likewise. '
            'Base panel (D40:H42) is live-linked.')

def snapshot_values(wb, bull, bear):
    ws = wb['Bull-Base-Bear']
    cols = 'CDEFGHIJKL'
    for snap, off in ((bull, 'OPQRSTUVWX'), (bear, ['AA', 'AB', 'AC', 'AD', 'AE', 'AF', 'AG', 'AH', 'AI', 'AJ'])):
        for (r, c), v in snap.items():
            tgt = f'{off[cols.index(c)]}{r}'
            cur = ws[tgt].value
            if isinstance(cur, str) and cur.startswith('='): continue
            if r in (26,): continue
            ws[tgt] = v

def inject_charts(path):
    """Replace openpyxl-rewritten chart parts with the template chart XML (keeps the combo chart, colours and layout), retitled."""
    tz = zipfile.ZipFile(TEMPLATE)
    tmpl = {n: tz.read(n) for n in tz.namelist() if n.startswith('xl/charts/chart')}
    tmp = path + '.tmp'
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in tmpl:
                t = tmpl[item.filename].decode('utf-8')
                for a, b2 in CHART_TITLES.items(): t = t.replace(a, b2)
                data = t.encode('utf-8')
            zout.writestr(item, data)
    shutil.move(tmp, path)
