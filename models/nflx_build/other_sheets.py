"""Rewire the Bull-Base-Bear, DCF and Charts tabs of the MLM template (current layout: DCF with live valuation date, fade period,
implied ROIC, target-IRR entry price, reverse DCF and Mauboussin EV/NOPAT; Bull-Base-Bear with buyback rows) to the NFLX Model rows /
columns; keep the template chart XML. The template's DCF_old tab is not carried.

Column mapping: the template's 2025A base column becomes NFLX FY2026E (1H actual, 2H calibrated to the guidance; the year ends
2.5 months after the valuation date) and the template's 2026E-2030E become NFLX 2027E-2031E, so the DCF's explicit years start in
2027E and net debt is taken at the 2026E year end (its cash flows, incl. the 2H buybacks, are in that balance)."""
import re, zipfile, shutil
from openpyxl.utils import get_column_letter as L
from fw import R, TEMPLATE, ANN, NOTE

# template Model row -> NFLX row key
MAP = {158: 'is_rev', 164: 'r_cash', 166: 'is_oi', 182: 'is_dda', 186: 'sh_dil', 211: 'r_adj', 212: 'r_adj_m', 219: 'r_ebit', 221: 'r_adj',
       229: 'e_adjni', 231: 'e_eps', 238: 'gm_ucan', 245: 'gm_ebitda_m', 247: 'gm_op_m', 260: 'wc_cf', 262: 'cf_capex', 263: 'cf_acq',
       280: 'fcf', 282: 'fcf_m', 290: 'bb_px', 291: 'bb_cash', 292: 'bb_sh', 331: 'bs_nci', 408: 'ra_ic', 411: 'ra_roic', 422: 'rn_ronta',
       428: 'lv_nd', 433: 'v_px', 434: 'v_sh'}
CMAP = {'BV': ANN[2026].c, 'CA': ANN[2027].c, 'CB': ANN[2028].c, 'CC': ANN[2029].c, 'CD': ANN[2030].c, 'CE': ANN[2031].c, 'CF': NOTE}
REF = re.compile(r"Model!(\$?)([A-Z]{1,3})(\$?)(\d+)")
HIST = [ANN[y].c for y in range(2006, 2026)]

def remap(f):
    def rep(m):
        c, r = m.group(2), int(m.group(4))
        c2 = CMAP.get(c, c)
        r2 = R[MAP[r]] if r in MAP else r
        if c == 'CF' and r == 2: r2 = 2
        return f'Model!{m.group(1)}{c2}{m.group(3)}{r2}'
    return REF.sub(rep, f)

def hist_stat(f):
    """LT historical AVERAGE / MIN / MAX over the template's annual columns -> NFLX FY2006-FY2025 annual columns."""
    m = re.match(r'=IFERROR\((AVERAGE|MIN|MAX)\(Model!C(\d+),', f)
    if not m: return None
    fn, r = m.group(1), int(m.group(2))
    rr = R[MAP[r]]
    return f'=IFERROR({fn}({",".join(f"Model!{c}{rr}" for c in HIST)}),"")'

def fy_dates(f):
    """template fiscal-year-end dates move one year with the column re-mapping (2026E -> 2027E)."""
    f = re.sub(r'DATE\((\d{4}),12,31\)', lambda m: f'DATE({int(m.group(1)) + 1},12,31)', f)
    f = re.sub(r'DATE\((\d{4})-1,12,31\)', lambda m: f'DATE({int(m.group(1)) + 1}-1,12,31)', f)
    return f

TEXT = {
    'Bull-Base-Bear': {
        'B2': 'Netflix, Inc. (NFLX) — Bull / Base / Bear Case P&L Summary',
        'B3': ('Adjusted EBITDA = company definition (NI + interest & other + tax + non-content D&A + SBC); adjusted net income / EPS = GAAP + company-quantified items (model) · USD millions · '
               'Live panel linked to Model — toggle scenario via Model!CG2. Bull and Bear panels are value snapshots generated from the Model with the switch set to each case. '
               '2026E = 1H actual + 2H calibrated to the Q3-26 and FY26 guidance (revenue $51.0–51.4bn, operating margin 31.5%, FCF ~$12.5bn).'),
        'G26': 22, 'H26': 22, 'S26': 28, 'T26': 28, 'AD26': 16, 'AE26': 16, 'AF26': 16,
    },
    'DCF': {
        'B2': 'DCF VALUATION — NETFLIX, INC. (NFLX)',
        'C5': 0.0425, 'D5': 'Assumption (~US 10Y Treasury yield, Oct-2026; government rate sources not reachable from the build environment); update to market',
        'D6': 'Assumption: US implied ERP ~4.5–5.5%',
        'C7': 1.15, 'D7': 'Assumption: large-cap media / streaming and internet platforms ~1.0–1.3 (DIS, WBD, SPOT, ROKU higher; GOOGL, META ~1.1); high-growth subscription franchise',
        'C9': 0.0525, 'D9': 'Jul-2026 $1bn 5.250% senior notes due 2036 (8-K 22-Jul-2026); A-range investment-grade issuer',
        'C10': 0.17, 'D10': 'Model 2027E+ effective tax rate (1H-26 ~16–17% excl. the termination fee; FY25 13.7%)',
        'C12': 0.05, 'D12': 'Net debt ~$5.2bn at 30-Jun-2026 vs ~$290bn market cap ($70 × ~4.1bn diluted shares at YE2026E); <0.5x adjusted EBITDA',
        'D15': 'Periods = (31-Dec of year − today) / 365.25; cash flows assumed at year end; 2026E is the base year (net debt at its year end), explicit years 2027E–2031E',
        'B18': 'Adj. EBIT (= operating income after content amortization, D&A and SBC) less company-quantified one-offs',
        'B21': 'plus non-content D&A (content amortization stays a cost)',
        'B22': '− Capex (PP&E) and acquisitions (M&A lever 2027E–31E)',
        'B23': '− Cash content spend above amortization and other working capital (cash-flow statement)',
        'B26': 'Fade-period growth (2032E–36E; 8.0% → 5.0% linear)', 'I26': 0.08, 'M26': 0.05, 'C27': 0.04,
        'D32': ("Netflix's operating income already deducts content amortization and stock-based compensation; only non-content D&A is added back (row 21). "
                'Row 23 deducts cash content spend in excess of content amortization (guided ~1.1x for 2026; scenario lever 2027E+) plus other operating working capital. '
                'Company-quantified one-offs (Q2-22 restructuring, Q3-25 Brazil tax) are treated as cash costs; none are forecast. Leases are operating (not in net debt). '
                'The $2.8bn WBD termination fee (Q1-26) sits in the 2026E year-end net debt, not in UFCF.'),
        'D43': 'Opening IC = Model 2026E book IC (equity + net debt at 31-Dec-2026E), rolled forward with capex + M&A − D&A + content & working-capital investment.',
        'B48': 'Less: Net debt + NCI (Model 2026E year end)',
        'D48': 'NFLX: no noncontrolling interests; net debt (senior notes − cash − short-term investments) at the base-year end, so 2027E+ cash flows are not double counted.',
        'B50': 'Diluted shares (m, Model 2026E period end)',
        'D76': 'Rationale (Netflix characteristics)',
        'C78': 0.12, 'D78': 'Model operating income CAGR 2026–31E ~13%: ~9% revenue growth (pricing, ads, membership) with ~7pp of margin expansion → low-teens NOPAT growth',
        'C79': 0.30, 'D79': 'Global scale spreads a fixed content budget over 300m+ paid memberships (301.6m at Q4-24, the last reported); incremental revenue (price, ads, paid sharing) carries high incremental margins; capital-light ex content',
        'C80': 15, 'D80': 'Scale advantage in content spend and engagement data, global distribution and brand; risks: competition for viewing time, content-cost inflation, saturation in UCAN',
        'B90': 'VALUATION  —  Applied to 2029E NOPAT',
        'B91': '2029E Adj. EBIT less company-quantified one-offs',
        'B93': '2029E NOPAT',
        'B95': 'Implied Enterprise Value (2029E basis)',
        'B96': 'Less: 2028E year-end net debt + NCI (Model)',
        'B98': 'Diluted shares 2029E (m, Model)',
        'B99': 'Implied share price 2029 ($)',
        'B106': 'Plus: 2026E year-end net debt + NCI',
        'B108': 'NFLX market-implied EV/NOPAT (on 2029E NOPAT)',
    },
    'Charts': {
        'B2': 'Netflix, Inc. — Operating Performance (2012–2025) and Model Forecast (2026E–2031E)',
        'B3': ('USDm · fiscal years ended 31 December · Linked to Model. Bars = contribution to revenue growth (pp): UCAN vs international (EMEA + LATAM + APAC) from 2018; '
               'legacy domestic (streaming + DVD) vs international streaming before. Adjusted EBITDA = company definition. ROIC / RONTA on adjusted EBIT (full-cost).'),
        'B8': 'UCAN / domestic contribution to growth (pp)', 'B9': 'International contribution to growth (pp)', 'B12': 'Adjusted EBITDA', 'B13': 'Operating income (GAAP)',
        'B14': 'Adj. EBITDA Margin %', 'B15': 'Operating Margin %',
    },
}
BBB_ROWS = {9: 'Revenues', 10: '   Revenues y/y %', 11: '   UCAN revenue y/y %', 12: 'Adjusted EBITDA (company definition)',
            15: 'Operating income (GAAP)', 16: '   Operating income y/y %', 17: '   Operating margin (GAAP) %', 19: 'Adjusted Net Income (model)',
            35: 'Adjusted EBITDA (leverage basis)', 36: '   Adjusted EBITDA % of revenues'}
CHART_TITLES = {
    'Aggregates Revenue Drivers: Shipment Volume vs. Price (ASP), y/y %': 'Revenue Growth: UCAN vs. International Contribution, pp',
    'Adjusted EBITDA &amp; Earnings from Operations (USDm)': 'Adjusted EBITDA &amp; Operating Income (USDm)',
}

def build_charts(ws):
    for i, y in enumerate(range(2012, 2032)):
        c = L(3 + i); mc = ANN[y].c
        ws[f'{c}7'] = y if y <= 2025 else f'{y}E'
        ws[f'{c}8'] = f'=IFERROR(Model!{mc}{R["gc_dom"]}+0,0)'
        ws[f'{c}9'] = f'=IFERROR(Model!{mc}{R["gc_int"]}+0,0)'
        ws[f'{c}11'] = f'={c}7'
        ws[f'{c}12'] = f'=Model!{mc}{R["r_adj"]}'
        ws[f'{c}13'] = f'=Model!{mc}{R["is_oi"]}'
        ws[f'{c}14'] = f'=Model!{mc}{R["gm_ebitda_m"]}'
        ws[f'{c}15'] = f'=Model!{mc}{R["gm_op_m"]}'
        ws[f'{c}17'] = f'={c}7'
        ws[f'{c}18'] = f'=IFERROR(Model!{mc}{R["ra_roic"]}+0,NA())'
        ws[f'{c}19'] = f'=IFERROR(Model!{mc}{R["rn_ronta"]}+0,NA())'
        ws[f'{c}20'] = f'=IFERROR(Model!{mc}{R["gm_ebitda_m"]}+0,NA())'
        ws[f'{c}21'] = f'=IFERROR(Model!{mc}{R["fcf_m"]}+0,NA())'

def do_all(wb):
    for name in ('Bull-Base-Bear', 'DCF', 'Charts'):
        ws = wb[name]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and 'Model!' in c.value:
                    h = hist_stat(c.value) if name == 'Bull-Base-Bear' else None
                    c.value = h or remap(c.value)
                if isinstance(c.value, str) and 'DATE(' in c.value:
                    c.value = fy_dates(c.value)
        for addr, t in TEXT.get(name, {}).items():
            ws[addr] = t
    b = wb['Bull-Base-Bear']
    for r, t in BBB_ROWS.items():
        for c in ('B', 'N', 'Z'): b[f'{c}{r}'] = t
    for c0 in ('C', 'O', 'AA'):
        i = b[f'{c0}7'].column
        b.cell(7, i).value = '2026E'
        for k in range(1, 6): b.cell(7, i + k).value = f'{2026 + k}E'
        b.cell(7, i + 6).value = "26-'31 CAGR"
    from openpyxl.comments import Comment
    cm = Comment('P/E on adjusted EPS (model: GAAP + company-quantified items; = GAAP in the forecast). NFLX at $70.00 trades at ~23x 2026E adjusted EPS ($3.03, excl. the '
                 'WBD termination fee; ~20x GAAP) and ~19x 2027E. Exit multiples are assumptions: Base 22x, Bull 28x, Bear 16x applied to 2030E–31E (2029E–31E for Bear) — a modest premium to the '
                 'market for high-teens EPS growth in Base, compressing as growth matures.', 'NFLX model')
    cm.width, cm.height = 420, 130; b['N26'].comment = cm
    cm2 = Comment(f'Snapshot (values) of Model rows {R["bb_cash"]}, {R["bb_px"]} and {R["bb_sh"]} (Share Buyback Schedule) with Model!{NOTE}2 = Bull / Bear. '
                  'Base panel (D40:H42) is live-linked.', 'NFLX model')
    cm2.width, cm2.height = 420, 90; b['P40'].comment = cm2
    b['J2'] = 'Share price ($):'
    for c in ('B', 'N', 'Z'): b[f'{c}29'] = 'IRR (to 31-Dec of Y-1)'
    for c in ('B', 'N', 'Z'):
        if b[f'{c}40'].value: b[f'{c}40'] = '   Share repurchases ($m)'
    d = wb['DCF']
    for i, col in enumerate('DEFGHIJKLM'):
        d[f'{col}17'] = f'{2027 + i}E'
    a26, a27 = ANN[2026].c, ANN[2027].c
    d['D22'] = f'=Model!{a27}{R["cf_capex"]}+Model!{a27}{R["cf_acq"]}'
    d['D34'] = f'=Model!{a26}{R["ra_ic"]}'
    d['C48'] = f'=Model!{a26}{R["lv_nd"]}+Model!{a26}{R["bs_nci"]}'
    d['C50'] = f'=Model!{a26}{R["v_sh"]}'
    d['C52'] = f'=Model!{a26}{R["v_px"]}'
    d['P17'] = 'CAGR 27–31E'; d['Q17'] = 'CAGR 27–36E'
    build_charts(wb['Charts'])

def snapshot_values(wb, bull, bear):
    ws = wb['Bull-Base-Bear']
    cols = 'CDEFGHIJKL'
    for snap, off in ((bull, 'OPQRSTUVWX'), (bear, ['AA', 'AB', 'AC', 'AD', 'AE', 'AF', 'AG', 'AH', 'AI', 'AJ'])):
        for (r, c), v in snap.items():
            tgt = f'{off[cols.index(c)]}{r}'
            cur = ws[tgt].value
            if isinstance(cur, str) and cur.startswith('='): continue
            if r in (7, 26): continue
            ws[tgt] = v

def inject_charts(path):
    """Replace openpyxl-rewritten chart parts with the template chart XML (combo charts, colours, layout), retitled, ranges kept
    (Charts!C:V = 2012-2031E) and cached values removed so Excel redraws from the NFLX cells."""
    tz = zipfile.ZipFile(TEMPLATE)
    tmpl = {n: tz.read(n) for n in tz.namelist() if n.startswith('xl/charts/chart')}
    tmp = path + '.tmp'
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in tmpl:
                t = tmpl[item.filename].decode('utf-8')
                for a, b2 in CHART_TITLES.items(): t = t.replace(a, b2)
                t = re.sub(r'<c:numCache>.*?</c:numCache>', '', t, flags=re.S)
                t = re.sub(r'<c:strCache>.*?</c:strCache>', '', t, flags=re.S)
                data = t.encode('utf-8')
            zout.writestr(item, data)
    shutil.move(tmp, path)
