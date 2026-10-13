"""Build NFLX_Model.xlsx from the MLM template (same layout, formatting and process as the prior models)."""
import sys, copy, os, json
from openpyxl.styles import Font
from fw import *
import fw
from spec import build_spec, scen_layout, SCEN_ROWS, SC
from spec2 import build_spec2

OUT = os.path.join(os.path.dirname(BASE), 'NFLX_Model.xlsx')
G26 = MAN['guidance_fy26']
DEBT = MAN['debt']

# ====================================================================== inputs (V0)
def wc_inputs():
    q = DATA['2025']['x']
    rev = q['rev']
    return {'oca': round(q['bs_oca'] / rev, 4), 'ap': round(q['bs_ap'] / rev, 4), 'accr': round(q['bs_accr'] / rev, 4), 'def': round(q['bs_defrev'] / rev, 4)}
WC = wc_inputs()
WC['note'] = (f'Forecast inputs at fiscal-year-end 2025 levels on FY2025 revenue (other current assets {WC["oca"]:.1%}, payables {WC["ap"]:.1%}, '
              f'accrued {WC["accr"]:.1%}, deferred revenue {WC["def"]:.1%}); year ends are compared with year ends.')

V0 = {
    'out_hdr': 'FY26 outlook — Q2-26 shareholder letter (16-Jul-2026)',
    'out_note': ('Q3-26 forecast: revenue $12,860m, operating margin 33.2%, net income $3,452m, diluted EPS $0.82. FY26: revenue $51.0–51.4bn (narrowed from $50.7–51.7bn), '
                 'operating margin 31.5%, content amortization ~+10%, cash content spend ~1.1x amortization, FCF ~$12.5bn (raised from ~$11bn in April), ads revenue ~$3bn. '
                 'All from the 8-K Ex. 99.1 letters on EDGAR; point guidance applies to all cases.'),
    'outlook': [
        ('out_rev', '  FY26 revenue ($m)', tuple(G26['revenue']), 'm', G26['revenue_note'] + '. Drives the Q4-26E revenue calibration.'),
        ('out_q3rev', '  Q3-26 revenue ($m)', (G26['q3']['revenue'],) * 3, 'm', 'Q3-26 forecast $12,860m (+11.7% y/y). Drives the Q3-26E revenue calibration (point guidance).'),
        ('out_om', '  FY26 operating margin', (G26['op_margin'],) * 3, 'pct', G26['op_margin_note'] + '. Drives the Q4-26E cost calibration.'),
        ('out_q3om', '  Q3-26 operating margin', (G26['q3']['op_margin'],) * 3, 'pct', 'Q3-26 forecast 33.2% (operating income $4,268m). Drives the Q3-26E cost calibration.'),
        ('out_fcf', '  FY26 free cash flow ($m)', (G26['fcf'],) * 3, 'm', G26['fcf_note'] + '. Solved through 2H working capital.'),
    ],
    'points': [
        ('pt_camg', '  FY26 content amortization growth', G26['content_amort_growth'], 'pct', G26['content_amort_note'] + '.'),
        ('pt_ccr', '  FY26 cash content spend / content amortization (x)', G26['cash_content_ratio'], 'x2', G26['cash_content_note'] + '.'),
        ('pt_int', '  Q3/Q4-26E interest expense ($m, each)', 185.0, 'm1', 'Not guided; Q2-26 $175.7m plus the $1bn 5.250% notes issued 22-Jul-2026, less the $1bn 4.375% notes due Nov-2026 (assumption).'),
        ('pt_oth', '  Q3/Q4-26E interest and other income ($m, each)', 50.0, 'm1', 'Not guided; Q2-26 $51.7m (interest income $82m less F/X and hedging); Q3/Q4-25 $36m / $45m (assumption).'),
        ('pt_etr', '  Q3/Q4-26E effective tax rate', 0.165, 'pct', 'Implied by the Q3-26 guide: net income $3,452m on operating income $4,268m less ~$135m net interest → ~16.5% (1H-26 excl. the termination fee ~16–17%).'),
        ('pt_capex', '  2H-26 capex % of revenue', 0.017, 'pct', '1H-26 1.7% of revenue ($414.8m); not guided (assumption).'),
        ('pt_q3ni', '  Q3-26 net income guidance ($m) — information', G26['q3']['net_income'], 'm1', 'Information row: model Q3-26E net income is compared, not forced.'),
        ('pt_q3eps', '  Q3-26 diluted EPS guidance ($) — information', G26['q3']['eps_diluted'], 'm2', 'Information row ($0.82).'),
    ],
    'lev': (1.75, 2.0, 2.5),
    'lev_note': ('Netflix dropped its former anchors ("$10–15bn gross debt and minimum cash equivalent to two months of revenue", Q1-24 letter) for policies consistent with a solid '
                 'investment-grade rating and returns excess cash via buybacks. Year-end cash was 2.4 months of revenue at YE2025 and 30-Jun-2026 cash $9.1bn. '
                 'Base keeps two months; Bull runs leaner (more buybacks); Bear holds more cash.'),
    'blocks': [
        ('UCAN_G', 'UCAN_G (UCAN revenue y/y)', 0.02, -0.03, [0.090, 0.080, 0.070, 0.065, 0.060],
         'FY25 +15%, 1H-26 +12% (price change partly in Q2-26, ads, paid sharing anniversaried); mature penetration → price- and ads-led high-single-digit growth.', 'pct', False),
        ('EMEA_G', 'EMEA_G (EMEA revenue y/y)', 0.02, -0.03, [0.110, 0.100, 0.090, 0.080, 0.070],
         'FY25 +17%, 1H-26 +16% reported (~11–12% F/X neutral per the letters); membership growth plus pricing and ads.', 'pct', False),
        ('LATAM_G', 'LATAM_G (LATAM revenue y/y)', 0.02, -0.04, [0.120, 0.110, 0.100, 0.090, 0.080],
         'FY25 +11% reported (~20%+ F/X neutral per the letters), 1H-26 +20% reported; high-inflation pricing and currency make reported growth volatile.', 'pct', False),
        ('APAC_G', 'APAC_G (APAC revenue y/y)', 0.02, -0.04, [0.130, 0.120, 0.110, 0.100, 0.090],
         'FY25 +21%, 1H-26 +18%; lowest-penetration region, membership-led growth.', 'pct', False),
        ('OM', 'OM (GAAP operating margin)', 0.01, -0.025, [0.335, 0.350, 0.365, 0.375, 0.385],
         'FY23 20.6% → FY24 26.7% → FY25 29.5% (30.9% excl. the $619m Brazil expense) → FY26 guided 31.5%; margin has expanded ~2–3pp a year as content amortization grows slower than revenue.', 'pct', False),
        ('CAMORT_G', 'CAMORT_G (content amortization y/y)', -0.01, 0.01, [0.080, 0.070, 0.065, 0.060, 0.055],
         'FY24 +8%, FY25 +7%, FY26 guided ~+10% (higher in 1H on title timing); content growing slower than revenue drives margin. Bull = slower growth (more FCF).', 'pct', False),
        ('CCR', 'CCR (cash content spend / content amortization, x)', -0.03, 0.05, [1.10, 1.08, 1.07, 1.06, 1.05],
         'FY23 0.93x (strikes), FY24 1.11x, FY25 1.08x, FY26 guided ~1.1x; originals and live sport keep cash spend ahead of amortization.', 'x2', False),
        ('MNA', 'M&A_spend (capability acquisitions, $m)', -250.0, 500.0, [500.0, 500.0, 500.0, 500.0, 500.0],
         'Mar-2026 acquisition ~$587m; 2021–23 studio / games deals. Base $500m p.a. (assumption); no revenue attached.', 'm', True),
    ],
    'mkt_pct': {2027: 0.070, 2028: 0.069, 2029: 0.068, 2030: 0.068, 2031: 0.067},
    'td_pct': {2027: 0.077, 2028: 0.075, 2029: 0.074, 2030: 0.073, 2031: 0.072},
    'ga_pct': {2027: 0.042, 2028: 0.041, 2029: 0.040, 2030: 0.040, 2031: 0.039},
    'dep_pct': 0.008, 'sbc_pct': 0.011, 'etr': 0.17, 'r_cash': 0.025,
    'r_cash_note': 'FY25 interest and other income $172m on ~$8bn average cash and investments (incl. F/X and hedging); 1H-26 interest income $152m (10-Q). 2.5% (assumption).',
    'capex_pct': {2027: 0.016, 2028: 0.015, 2029: 0.014, 2030: 0.013, 2031: 0.013},
    'def_ae': 0.0, 'stock_2h': 100.0, 'stock_ae': 250.0, 'iss_2h': 2.0, 'iss_ae': 5.0, 'dil': 72.0,
    'repay': {2026: 1000.0, 2027: 1486.0, 2028: 3500.0, 2029: 4329.0, 2030: 2257.0, 2031: 0.0},
    'mat': {2026: 1486.0, 2027: 3500.0, 2028: 4329.0, 2029: 2257.0, 2030: 0.0, 2031: 1000.0},
    'iss_2026': 1000.0,
    'repay_note': ('Par maturities (Q2-26 10-Q Note 7): Nov-2026 $1,000m 4.375%; May-2027 EUR 3.625% $1,486m; 2028 $3,500m (4.875% Apr, 5.875% Nov); 2029 $4,329m '
                   '(EUR 4.625% $1,257m, 6.375% $800m, EUR 3.875% $1,372m, 5.375% $900m); Jun-2030 $2,257m (EUR 3.625% $1,257m, 4.875% $1,000m); 2034 $1,000m; 2036 $1,000m '
                   '(new); 2054 $800m. EUR notes at 30-Jun-2026 FX.'),
    'min_cash': 3000.0,
    'r_notes': {2027: 0.049, 2028: 0.050, 2029: 0.051, 2030: 0.052, 2031: 0.052},
    'r_notes_note': (DEBT['coupon_note'] + ' Refinancing at ~5.25% (Jul-2026 notes) lifts the average coupon as low-coupon EUR notes mature (assumption).'),
    'wc': WC,
    'px': MAN['price']['px'],
    'px_note': ('VALUATION — share price input $70.00 (provided by the user, 13-Oct-2026); update to market'),
    'mna_note': 'Scenario lever M&A_spend (Base $500m a year of capability acquisitions). Cash only; funded from cash (reduces buybacks).',
    'rev_guid_note': G26['revenue_note'] + '. Drives the Q4-26E revenue calibration.',
    'om_guid_note': 'Q3-26 operating margin 33.2% and FY26 31.5% (Q2-26 letter). Drive the Q3-26E and Q4-26E cost calibrations.',
}

# ====================================================================== compute engine
def hist_value(row, col):
    kw = row.kw; path = kw.get('data')
    if not path: return None
    if kw.get('ppa'):
        v = MAN['ppa'].get(pkey(col), {}).get(path.split('_', 1)[1])
        return v if kw.get('textdata') else num(v)
    v = xget(pkey(col), path)
    return v if kw.get('textdata') else num(v)

def first(kw, keys):
    for k in keys:
        if k in kw: return k
    return None

def compute(row, col):
    kw = row.kw; k = col.kind
    CTX.col = col
    if kw.get('ann') and col.q: return None, False
    if k in HISTK:
        v = hist_value(row, col)
        if v is not None: return v, True
        for key in ('hist', 'all'):
            if key in kw:
                return kw[key](), False
        return None, False
    if k == 'QE':
        if row.spec2: return None, False
        if 'qe_in' in kw:
            v = kw['qe_in']; return (v[col.q] if isinstance(v, dict) else v), True
        key = first(kw, ('qe', 'fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    if k == 'E26':
        if 'e26_in' in kw: return kw['e26_in'], True
        if 'e26' in kw: return kw['e26'](), False
        if kw.get('avg') and ('qe' in kw or 'qe_in' in kw): return f'=AVERAGE({Q1}{row.r}:{QE}{row.r})', False
        if kw.get('flow') and not row.spec2 and any(x in kw for x in ('qe', 'qe_in', 'fc')): return S26(row.key), False
        key = first(kw, ('fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    if k == 'AE':
        if 'ae_in' in kw:
            v = kw['ae_in']; return (v[col.year] if isinstance(v, dict) else v), True
        key = first(kw, ('ae', 'fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    return None, False

# ====================================================================== writers
CAGR_G = {'is_rev', 'g_rev', 'ucan_rev', 'emea_rev', 'latam_rev', 'apac_rev', 'rg_paid', 'd_camort', 'is_oi', 'o_adj', 'r_adj', 'r_ebitda', 'e_adjni', 'e_eps',
          'cf_cfo', 'fcf', 'bs_ta', 'bs_eq', 'sh_dil', 'is_gp', 'is_mkt', 'is_td', 'is_ga', 'd_dep', 'd_sbc', 'is_ni', 'eps_dil', 'ccs', 'bs_cont'}
CAGR_A = {'is_oi_m', 'o_adj_m', 'r_adj_m', 'is_gm', 'r_ebit_m', 'gm_rev', 'gm_ucan', 'gm_intl', 'gm_gm', 'gm_ebitda_m', 'gm_op_m', 'gm_net_m', 'gm_mkt', 'gm_td',
          'gm_ga', 'gm_camort', 'fcf_m', 'fcf_conv', 'capex_pct', 'ra_roic', 'rn_ronta', 'roe', 'lv_x', 'd_camort_p', 'd_dep_pct', 'd_sbc_pct', 'wc_nwc_p', 'is_etr',
          'fcf_eb', 'ccr', 'ucan_mix', 'emea_mix', 'latam_mix', 'apac_mix', 'd_camort_g'}
A = {y: ANN[y].c for y in ANN}

def style_cagr(ws, r, on):
    src = 16 if on else 12
    for c in CAGR:
        ws[f'{c}{r}']._style = copy.copy(TM[f'{L(CI(c) + TSHIFT)}{src}']._style)

def write_row(ws, row):
    r = row.r; arch = ARCH[row.arch]
    clone_row_style(ws, arch, r)
    ws.cell(r, 2).value = row.label or None
    if row.kw.get('note'): ws[f'{NOTE}{r}'] = row.kw['note']
    style_cagr(ws, r, row.key in CAGR_G or row.key in CAGR_A)
    if row.arch in ('section', 'block', 'blank', 'sub') and not row.kw.get('data'):
        return
    if row.key in DEFS:
        write_defs(ws, row); return
    nf = NF[row.nf]
    for col in COLS:
        v, inp = compute(row, col)
        if v is None: continue
        cell = ws[f'{col.c}{r}']
        cell.value = v
        if not row.kw.get('textdata'): cell.number_format = nf
        if inp:
            set_color(cell, BLUE, bold=True if col.kind in FCK else None)
            lk = row.kw.get('labkey')
            if lk and col.kind in HISTK and isinstance(v, (int, float)) and abs(v) > 1e-9:
                t = xget(pkey(col), lk)
                if t: comment(cell, f'Company items ({pkey(col)}): {t}', 380, 110)
        elif row.arch == 'growth' and col.kind in HISTK:
            set_color(cell, GRAY)
        else:
            set_color(cell, BLACK)
    rc = CELL_COMMENTS.get(row.key, {})
    for c, t in rc.items(): comment(ws[f'{c}{r}'], t)
    if row.key in CAGR_G:
        ws[f'{CAGR[0]}{r}'] = f'=IFERROR(({A[2031]}{r}/{A[2026]}{r})^(1/5)-1,"n/m")'
        ws[f'{CAGR[1]}{r}'] = f'=IFERROR(({A[2025]}{r}/{A[2020]}{r})^(1/5)-1,"n/m")'
        ws[f'{CAGR[2]}{r}'] = f'=IFERROR(({A[2025]}{r}/{A[2015]}{r})^(1/10)-1,"n/m")'
        ws[f'{CAGR[3]}{r}'] = f'=IFERROR(({A[2025]}{r}/{A[2006]}{r})^(1/19)-1,"n/m")'
        for c in CAGR: ws[f'{c}{r}'].number_format = '0%;\\(0%\\);"n/m"'
    elif row.key in CAGR_A:
        ws[f'{CAGR[0]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2027, 2032))}),"n/m")'
        ws[f'{CAGR[1]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2021, 2026))}),"n/m")'
        ws[f'{CAGR[2]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2016, 2026))}),"n/m")'
        ws[f'{CAGR[3]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2006, 2026))}),"n/m")'
        for c in CAGR: ws[f'{c}{r}'].number_format = NF['pct']
    t = LABEL_SOURCES.get(row.key)
    if t: comment(ws.cell(r, 2), t)

REL = 'Source: Netflix Forms 10-Q / 10-K (SEC EDGAR, CIK 1065280) — '
LTR = 'Source: Netflix shareholder letters / earnings releases (Form 8-K Item 2.02, Ex. 99.1) — '
LABEL_SOURCES = {
    'ucan_rev': (REL + 'revenue by region (revenue note / MD&A): FY2017–19 from the FY2019 10-K; Q1–Q3 2019 from the 2020 10-Q prior-year columns (latest filing presenting the period); '
                 'Q4 = fiscal year (10-K) − nine months (Q3 10-Q). Paid memberships at period end; ARM from the MD&A (Q4 from the Q4 letter regional table).'),
    'sg_dom_rev': REL + 'segment note 2011–2019 (Domestic streaming, International streaming, Domestic DVD; revenue, cost of revenues, marketing, contribution profit, memberships). Q4 = FY − 9M.',
    'is_rev': (REL + 'consolidated statements of operations, as originally reported (Q1–Q3 = the 10-Q three-month column; Q4 = FY − nine months, tied to the Q4 letter). '
               '2018 quarters from the 2019 10-Q comparatives (Q4-2018 cost reclassification).'),
    'd_camort': REL + 'cash-flow statements (amortization of streaming content assets + amortization of DVD content assets; year to date, discrete quarters derived).',
    'r_adj_pub': LTR + '"Non-GAAP Adjusted EBITDA reconciliation" (letters Q2-15 – Q1-23; each quarter takes the first published value, matched on GAAP net income).',
    'e_pubni': LTR + '"Non-GAAP net income reconciliation" (2006 – Q2-10 excl. stock-based compensation; Q1-13 debt extinguishment; Q4-14 / Q4-15 tax-accrual releases; Q1-15 F/X impact).',
    'fcf_pub': LTR + '"Non-GAAP free cash flow reconciliation" (column matched on net cash from operating activities).',
    'cf_ni': REL + 'cash-flow statements (year to date; discrete quarters derived, Q4 = FY − 9M). "Working capital & other" = CFO − the named lines.',
    'bs_cash': REL + 'consolidated balance sheets (period ends); shares outstanding from XBRL (CommonStockSharesOutstanding, as first filed, split-adjusted).',
    'sh_basic': REL + 'weighted-average shares; split-adjusted by the filing date of each source (7-for-1 effective 14-Jul-2015; 10-for-1 effective 14-Nov-2025).',
    'ds_total': 'Source: Q2-26 10-Q Note 7 (debt); 8-K 22-Jul-2026 (Item 8.01: $1bn 5.250% notes due 2036); $3bn revolving credit agreement (Apr-2024).',
    'o_restr': MAN['one_offs']['restructuring_q222']['src'],
    'n_fee': MAN['deals']['wbd'],
    'm_spend': MAN['deals']['acq_2026'] + ' ' + MAN['deals']['acq_hist'],
}
CELL_COMMENTS = {}
def comments_setup():
    import normalize
    for q in (1, 2, 3):
        c = QC[(2018, q)].c
        CELL_COMMENTS.setdefault('is_cor', {})[c] = ('Recast: ' + normalize.RECAST[2018] + f' Originally reported cost of revenues ${xget(f"Q{q}-18", "orig_cor"):,.1f}m, '
                                                     f'marketing ${xget(f"Q{q}-18", "orig_mkt"):,.1f}m, T&D ${xget(f"Q{q}-18", "orig_td"):,.1f}m, G&A ${xget(f"Q{q}-18", "orig_ga"):,.1f}m.')
    CELL_COMMENTS.setdefault('sg_dom_rev', {})[A[2011]] = ('FY2011: segment reporting began in Q4-2011 (Qwikster separation of streaming and DVD); the FY2011 10-K presents segments for '
                                                          'the fourth quarter only, so the annual column is left blank. Memberships are in the 10-K segment tables for FY2015–17 only.')
    CELL_COMMENTS.setdefault('is_rev', {})[A[2006]] = 'FY2006 10-K: columns presented oldest first (2004, 2005, 2006); the 2006 column is used.'
    CELL_COMMENTS.setdefault('fcf_pub', {})[A[2007]] = MAN['corrections'][2]
    CELL_COMMENTS.setdefault('e_pubni', {})[QC[(2014, 4)].c] = MAN['corrections'][0]
    CELL_COMMENTS.setdefault('o_brazil', {})[QC[(2025, 3)].c] = MAN['one_offs']['brazil_q325']['src']
    CELL_COMMENTS.setdefault('o_restr', {})[QC[(2022, 2)].c] = MAN['one_offs']['restructuring_q222']['src']
    CELL_COMMENTS.setdefault('n_fee', {})[QC[(2026, 1)].c] = MAN['one_offs']['wbd_fee_q126']['src']
    for q in (1, 2, 3):
        CELL_COMMENTS.setdefault('ucan_rev', {})[QC[(2019, q)].c] = 'From the 2020 Form 10-Q prior-year column (regional disclosure began with the FY2019 10-K).'
    CELL_COMMENTS.setdefault('ucan_rev', {})[A[2017]] = 'FY2017–18 from the FY2019 10-K regional tables (first regional disclosure).'
    CELL_COMMENTS.setdefault('fcf_pub', {})[QC[(2013, 2)].c] = 'Q2-13 letter shows Q2 CFO $33.9m after restating Q1-13; the 10-Q year-to-date figures give $33.9m (H1 − Q1 as first filed, $0.1m different). Published FCF matched within $0.1m.'

DEFS = {}
def defs_setup():
    q = lambda y, n: QC[(y, n)].c
    DEFS['o_def'] = {
        'C': ('Model definition: Netflix publishes no adjusted operating income; only one-offs the company itself quantified in its letters are added back', None),
        q(2022, 2): ('Q2-22: restructuring ~$150m (severance ~$70m + real-estate lease impairment ~$80m) quantified in the letter', MAN['one_offs']['restructuring_q222']['src']),
        q(2025, 3): ('Q3-25: Brazilian non-income tax expense ~$619m quantified in the letter', MAN['one_offs']['brazil_q325']['src']),
    }
    DEFS['r_def'] = {
        'C': ('2006 – Q1-14: not published — rebuilt on the company definition (NI + interest & other + tax + D&A + SBC)', None),
        q(2014, 2): ('Q2-14 – Q1-23: company-published adjusted EBITDA (letters Q2-15 – Q1-23; checked)', 'Each quarter takes its first published value; the definition is unchanged across the letters.'),
        q(2023, 2): ('Q2-23+: no longer published — model rebuild on the same definition', None),
    }
    DEFS['e_def'] = {
        'C': ('FY2006–09: company non-GAAP net income excl. stock-based compensation, net of tax (checked)', 'Published in the Q4 releases 2006–2009 (and quarterly to Q2-10).'),
        A[2010]: ('FY2010–12: not published (quarters Q1–Q2-10 only) — model: GAAP + quantified items (none)', None),
        q(2013, 1): ('Q1-13: excl. loss on extinguishment of debt (company)', None),
        q(2014, 4): ('Q4-14: excl. release of tax accrual (company)', MAN['corrections'][0]),
        q(2015, 1): ('Q1-15: excl. foreign-exchange impact (company)', None),
        q(2015, 4): ('Q4-15: excl. release of tax accrual (company)', None),
        q(2016, 1): ('2016+: not published — model: GAAP + F/X remeasurement on euro debt (2017+), restructuring (Q2-22), Brazil expense (Q3-25), WBD fee (Q1-26), tax-effected at the bounded ETR', None),
    }
def write_defs(ws, row):
    for c, (txt, cm) in DEFS[row.key].items():
        cell = ws[f'{c}{row.r}']; cell.value = txt
        cell.font = Font(name='Calibri', sz=10, b=True, i=True, color='FFC00000')
        if cm: comment(cell, cm)

def write_scen(ws):
    srcrow = {'hdr': 9, 'out_hdr': 11, 'pt_hdr': 15, 'out': 12, 'pt': 16, 'lev': 28, 'blk': 30, 'yr': 31}
    for spec_ in SCEN_ROWS:
        r, kind = spec_[0], spec_[1]
        for ci in range(CI(SL), CI(SNOTE) + 1):
            ws.cell(r, ci)._style = copy.copy(TM.cell(srcrow[kind], ci + TSHIFT)._style)
        if kind in ('hdr', 'out_hdr', 'pt_hdr'):
            _, _, a, b, c_, d_, e = spec_
            for col, v in zip([SL, SBULL, SBASE, SBEAR, SNOTE], [a, b, c_, d_, e]):
                if v is not None: ws[f'{col}{r}'] = v
        elif kind == 'out':
            _, _, lab, hi, mid, lo, note, nf = spec_
            ws[f'{SL}{r}'] = lab
            for col, v in zip([SBULL, SBASE, SBEAR], [hi, mid, lo]):
                ws[f'{col}{r}'] = v; ws[f'{col}{r}'].number_format = NF[nf]; set_color(ws[f'{col}{r}'], BLUE)
            ws[f'{SNOTE}{r}'] = note
        elif kind == 'pt':
            _, _, lab, _, val, _, note, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBASE}{r}'] = val; ws[f'{SBASE}{r}'].number_format = NF[nf]; set_color(ws[f'{SBASE}{r}'], BLUE); ws[f'{SNOTE}{r}'] = note
        elif kind == 'lev':
            _, _, lab, a, b, c_, note = spec_
            ws[f'{SL}{r}'] = lab
            for col, v in zip([SBULL, SBASE, SBEAR], [a, b, c_]):
                ws[f'{col}{r}'] = v; ws[f'{col}{r}'].number_format = NF['x']
            ws[f'{SNOTE}{r}'] = note
        elif kind == 'blk':
            _, _, lab, db, mid, de, note, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBULL}{r}'] = db; ws[f'{SBASE}{r}'] = mid; ws[f'{SBEAR}{r}'] = de; ws[f'{SNOTE}{r}'] = note
            for col in (SBULL, SBEAR):
                ws[f'{col}{r}'].number_format = '\\+#,##0;\\-#,##0' if nf == 'm' else '\\+0.0%;\\-0.0%'
        elif kind == 'yr':
            _, _, lab, hdr, base, clamp, _, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBASE}{r}'] = base
            if clamp:
                ws[f'{SBULL}{r}'] = f'=MAX(0,{SBASE}{r}+${SBULL}${hdr})'; ws[f'{SBEAR}{r}'] = f'=MAX(0,{SBASE}{r}+${SBEAR}${hdr})'
            else:
                ws[f'{SBULL}{r}'] = f'={SBASE}{r}+${SBULL}${hdr}'; ws[f'{SBEAR}{r}'] = f'={SBASE}{r}+${SBEAR}${hdr}'
            for col in (SBULL, SBASE, SBEAR): ws[f'{col}{r}'].number_format = NF[nf]

def write_header(ws):
    for r in range(1, 8):
        clone_row_style(ws, r, r, 2, CI(LASTC))
    ws[f'{NOTE}1'] = 'SCENARIO SWITCH ▼ (Bull / Base / Bear) — drives all scenario-linked forecast drivers'
    ws[f'{NOTE}2'] = 'Base'
    comment(ws[f'{NOTE}2'], 'Scenario toggle. Choose Bull / Base / Bear. Drives: the FY26 revenue guidance end-point (high / mid / low), 2027E–2031E regional revenue growth, '
            'operating margin, content amortization growth, the cash content spend / amortization ratio, M&A spend and the target cash balance for buybacks '
            f'(scenario table {SL}:{SNOTE}).')
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type='list', formula1='"Bull,Base,Bear"', allow_blank=False); ws.add_data_validation(dv); dv.add(f'{NOTE}2')
    ws['B2'] = 'Netflix, Inc. (NASDAQ: NFLX) — 3-Statement Financial Model'
    ws['B3'] = ('US GAAP as reported · USD millions (per-share data in USD; shares in millions, split-adjusted) · Fiscal years ended 31 December · '
                'Source: Netflix Forms 10-K / 10-Q, Form 8-K shareholder letters (Ex. 99.1) and financing / deal 8-Ks, SEC EDGAR (CIK 1065280); XBRL company facts for cross-checks and '
                'shares outstanding; the investor-relations site (ir.netflix.net) is not used — its letters are the same documents as filed on EDGAR')
    ws['B4'] = ('Basis notes: as originally reported (2018 quarters recast for the Q4-2018 cost reclassification). Splits 7-for-1 (Jul-2015) and 10-for-1 (Nov-2025) applied by filing date. '
                'Segments: Domestic / International streaming + Domestic DVD (2011–19, memo); regions UCAN / EMEA / LATAM / APAC (2017+). DVD service closed Sep-2023. '
                'WBD merger agreed Dec-2025, terminated 27-Feb-2026 ($2.8bn fee received Q1-26). Q3/Q4-26E calibrated to the Q3-26 and FY26 guidance (16-Jul-2026).')
    ws[f'{CAGR[0]}5'] = 'Forecast'
    for c in CAGR[1:]: ws[f'{c}5'] = 'Historical'
    ws['B6'] = '(USDm)'
    for col in COLS: ws[f'{col.c}6'] = col.label
    ws[f'{NOTE}6'] = 'Modelling Notes'
    for c, a, b in zip(CAGR, ['5Y CAGR', '5Y CAGR', '10Y CAGR', '19Y CAGR'], ["26-'31", "20-'25", "15-'25", "06-'25"]):
        ws[f'{c}6'] = a; ws[f'{c}7'] = b

def set_columns(ws):
    for k in list(ws.column_dimensions.keys()): del ws.column_dimensions[k]
    ws.column_dimensions['A'].width = TM.column_dimensions['A'].width
    ws.column_dimensions['B'].width = TM.column_dimensions['B'].width
    wid = {'A': 13.0, 'Q': 9.5703125, 'QE': 11.42578125, 'E26': 14.140625, 'AE': 13.0}
    for col in COLS:
        cd = ws.column_dimensions[col.c]; cd.width = wid[col.kind]
        if col.q:
            cd.outlineLevel = 1; cd.hidden = True
    for ci in range(CI(NOTE), CI(LASTC) + 1):
        ws.column_dimensions[L(ci)].width = TM.column_dimensions[L(ci + TSHIFT)].width or 13.0
    ws.sheet_format.outlineLevelCol = 1

def main():
    import openpyxl
    wb = openpyxl.load_workbook(TEMPLATE)
    if 'DCF_old' in wb.sheetnames: del wb['DCF_old']
    ws = wb['Model']
    blank_style = copy.copy(ws['A1']._style)
    for rowc in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for c in rowc:
            c.value = None; c.comment = None; c._style = copy.copy(blank_style)
    for k in list(ws.row_dimensions.keys()):
        rd = ws.row_dimensions[k]; rd.outlineLevel = 0; rd.hidden = False
    ws.data_validations.dataValidation = []
    scen_layout(V0)
    build_spec(V0)
    n1 = len(fw.SPEC)
    build_spec2(V0)
    for i, row in enumerate(fw.SPEC): row.spec2 = i >= n1
    last = number_rows(8)
    for r in range(1, last + 1): ws.row_dimensions[r].height = 15.0
    ws.row_dimensions[2].height = 21.0
    defs_setup(); comments_setup()
    write_header(ws)
    for row in fw.SPEC: write_row(ws, row)
    write_scen(ws)
    set_columns(ws)
    ws.freeze_panes = 'C8'
    ws.sheet_view.zoomScale = 70
    import other_sheets as osh
    osh.do_all(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    osh.inject_charts(OUT)
    print('rows', fw.SPEC[-1].r, 'saved', OUT)
    return wb

RECALC = os.environ.get('XLSX_RECALC', '/mnt/skills/public/xlsx/scripts/recalc.py')
def recalc(path, timeout=600):
    import subprocess
    out = subprocess.run(['python3', RECALC, path, str(timeout)], capture_output=True, text=True, cwd=os.path.dirname(RECALC))
    try:
        return json.loads(out.stdout)
    except Exception:
        return {'raw': out.stdout[-2000:], 'err': out.stderr[-2000:]}

def scenario_snapshots():
    import openpyxl
    import other_sheets as osh
    snaps = {}
    for scen in ('Bull', 'Bear'):
        p = os.path.join(BASE, f'NFLX_{scen}.xlsx')
        wb = openpyxl.load_workbook(OUT); wb['Model'][f'{NOTE}2'] = scen; wb.save(p)
        res = recalc(p); print(scen, res.get('status'), res.get('total_errors'))
        v = openpyxl.load_workbook(p, data_only=True)['Bull-Base-Bear']
        snaps[scen] = {(r, c): v[f'{c}{r}'].value for r in range(9, 51) for c in 'CDEFGHIJKL' if v[f'{c}{r}'].value is not None}
        os.remove(p)
    wb = openpyxl.load_workbook(OUT)
    osh.snapshot_values(wb, snaps['Bull'], snaps['Bear'])
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    osh.inject_charts(OUT)

NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
def inject_values(src_xlsx, recalc_xlsx):
    """Store cached formula results (from a LibreOffice-recalculated copy) in the openpyxl file, so the workbook reads correctly in any
    viewer; fullCalcOnLoad stays set so Excel recalculates on open. Same approach as the LECO / ESAB finishing step."""
    import zipfile, shutil, openpyxl
    from lxml import etree
    wbv = openpyxl.load_workbook(recalc_xlsx, data_only=True)
    zin = zipfile.ZipFile(src_xlsx)
    wbxml = etree.fromstring(zin.read('xl/workbook.xml'))
    rels = etree.fromstring(zin.read('xl/_rels/workbook.xml.rels'))
    rid2t = {r.get('Id'): r.get('Target') for r in rels}
    sheets = {}
    for s in wbxml.iter('{%s}sheet' % NS):
        rid = s.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
        sheets['xl/' + rid2t[rid].lstrip('/').replace('xl/', '')] = s.get('name')
    tmp = src_xlsx + '.tmp'
    zout = zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED)
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename in sheets:
            ws = wbv[sheets[item.filename]]
            root = etree.fromstring(data)
            for c in root.iter('{%s}c' % NS):
                f = c.find('{%s}f' % NS)
                if f is None: continue
                v = ws[c.get('r')].value
                old = c.find('{%s}v' % NS)
                if old is not None: c.remove(old)
                if v is None: continue
                ve = etree.SubElement(c, '{%s}v' % NS)
                if isinstance(v, bool):
                    c.set('t', 'b'); ve.text = '1' if v else '0'
                elif isinstance(v, (int, float)):
                    if 't' in c.attrib: del c.attrib['t']
                    ve.text = repr(float(v)) if isinstance(v, float) else str(v)
                elif isinstance(v, str) and v.startswith('#'):
                    c.set('t', 'e'); ve.text = v
                else:
                    c.set('t', 'str'); ve.text = str(v)
            data = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
        zout.writestr(item, data)
    zout.close(); zin.close()
    shutil.move(tmp, src_xlsx)

def store_values():
    import shutil
    p = os.path.join(BASE, 'NFLX_recalc.xlsx')
    shutil.copy(OUT, p)
    res = recalc(p); print('recalc', res.get('status'), res.get('total_errors'))
    inject_values(OUT, p)
    os.remove(p)

if __name__ == '__main__':
    main()
    if '--snap' in sys.argv: scenario_snapshots()
    if '--recalc' in sys.argv: store_values()
