"""Build the Karman (KRMN) model workbook on the TDY template conventions."""
import sys
import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import column_index_from_string as ci, get_column_letter as gl
from lib import *
from data import H
from rows import build_rows, SCENARIO_TABLE

STYLE = {'SEC': 8, 'BLK': 10, 'TOPV': 11, 'PCT': 12, 'VAL': 13, 'INROW': 15, 'KEY': 16, 'MEMO': 17, 'CHK': 18, 'RATIO': 19,
         'DRV': 138, 'GP': 156, 'TOT': 200, 'MULT': 107, 'SH': 178, 'EPSK': 181, 'EPSM': 182, 'EPSB': 243, 'EPSV': 302, 'DEF': 197,
         'SCHED': 371, 'SCHEDC': 376, 'LEV': 390, 'LEVK': 406, 'KEYP': 406, 'DSO': 359, 'PX': 432, 'EVX': 439}


def build_model(wb):
    ws = wb.create_sheet('Model')
    S = Sheet(ws)
    build_rows(S)

    # ---------- column widths / sheet view ----------
    ws.column_dimensions['A'].width = TM.column_dimensions['A'].width
    ws.column_dimensions['B'].width = 70
    for col in DATA_COLS:
        ws.column_dimensions[col].width = 12.5
    for col, w in {'AA': 70, 'AB': 3, 'AC': 11, 'AD': 11, 'AE': 3, 'AF': 3, 'AG': 58, 'AH': 11, 'AI': 11, 'AJ': 11, 'AK': 70}.items():
        ws.column_dimensions[col].width = w
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 60
    ws.freeze_panes = 'C8'

    def style_from(tr, col, cell):
        src_col = 'B' if col in ('A', 'B') else CSTYLE.get(col)
        if src_col is None:
            return
        copy_style(TM['%s%d' % (src_col, tr)], cell)

    # ---------- header block (rows 1-7) ----------
    for r in range(1, 8):
        ws.row_dimensions[r].height = TM.row_dimensions[r].height
        for col in ['B'] + DATA_COLS + ['AA', 'AB', 'AC', 'AD']:
            style_from(r, col, ws['%s%d' % (col, r)])
    ws['AA1'] = 'SCENARIO SWITCH ▼ (Bull / Base / Bear) — drives all scenario-linked forecast drivers'
    ws['AA2'] = 'Base'
    ws['AA2'].comment = note_comment('Scenario toggle. Choose Bull / Base / Bear. Drives: FY26 outlook end-points (revenue and Adj. EBITDA — high / mid / low), 2027E–30E end-market growth, Adj. EBITDA margin, acquisition spend (M&A lever), buybacks and the leverage floor (scenario input table to the right, columns AG:AK). Bull / Bear = Base + Δ.')
    dv = DataValidation(type='list', formula1='"Bull,Base,Bear"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add('AA2')
    ws['B2'] = 'Karman Holdings Inc. (NYSE: KRMN) — 3-Statement Financial Model'
    ws['B3'] = ('US GAAP as reported · USD millions (per-share data in USD) · Source: Karman Forms 10-K / 10-Q, IPO prospectus (Form 424B4, 13-Feb-2025) and Form 8-K earnings '
                'releases (Ex. 99.1), SEC EDGAR CIK 0002040127 (investor-relations releases at investors.karman-sd.com are the same documents as filed on EDGAR)')
    ws['B4'] = ('Basis notes: calendar fiscal year. IPO 14-Feb-2025 with Corporate Conversion (TCFIII Spaceco Holdings LLC → Karman Holdings Inc.) — pre-IPO per-unit data (159–167m units) '
                'not comparable with ~132m post-IPO shares. One reportable segment; revenue by end market (2024 quarters on the 2025 recast basis; Maritime Defense Systems disclosed from Q1-26). '
                'Quarterly data from Q1/24 (first quarters reported, as comparatives in the 2025 10-Qs); FY2022–23 annual only. Acquisitions: RMS (Feb-24), MTI (Apr-25), ISP (May-25), Five Axis (Oct-25), '
                'Seemann Composites & Materials Sciences (Q1-26), Walker Precision Engineering (28-Aug-26, after the FY26 outlook). Q3/26E–Q4/26E = FY26 revenue / Adj. EBITDA outlook less 1H/26 actuals.')
    ws['B6'] = '(USDm)'
    for col in DATA_COLS:
        per = CPER[col]
        ws['%s6' % col] = int(per) if per.isdigit() else per
    ws['AA6'] = 'Modelling Notes'
    ws['AC5'], ws['AD5'] = 'Forecast', 'Historical'
    ws['AC6'], ws['AD6'] = '5Y CAGR', '3Y CAGR'
    ws['AC7'], ws['AD7'] = "25-'30", "22-'25"
    for col in ['AC', 'AD']:
        for r in (5, 6, 7):
            copy_style(TM['%s%d' % ('CV' if col == 'AC' else 'CW', r)], ws['%s%d' % (col, r)])

    # ---------- body ----------
    for rnum, row in S.rows:
        tr = STYLE[row.style]
        ws.row_dimensions[rnum].height = TM.row_dimensions[tr].height
        lab = ws['B%d' % rnum]
        style_from(tr, 'B', lab)
        lab.value = row.label
        if row.comment:
            lab.comment = note_comment(row.comment)
        for col in DATA_COLS + ['AA', 'AB', 'AC', 'AD']:
            style_from(tr, col, ws['%s%d' % (col, rnum)])
        if row.note:
            ws['AA%d' % rnum] = row.note
        if row.style in ('SEC', 'BLK'):
            continue
        for col in DATA_COLS:
            t = CTYPE[col]
            per = CPER[col]
            val = None
            if col in row.cells:
                val = row.cells[col]
            elif col in FC_COLS and col in row.fc:
                val = row.fc[col]
            elif t in ('QE', 'AE', 'F') and t in row.fc:
                val = row.fc[t]
            elif row.all is not None:
                val = row.all
            elif t in ('A', 'Q'):
                if row.hist is not None:
                    d = H.get(row.hist, {}) if isinstance(row.hist, str) else row.hist
                    if row.h == 'stock' and per.startswith('Q4'):
                        val = '=IF(ISNUMBER(<fy>[@]),<fy>[@],"")'
                    else:
                        val = d.get(per)
                elif row.histf is not None:
                    val = row.histf
            elif t == 'H':
                if row.h == 'sum':
                    val = '=<q1>[@]+<q2>[@]'
                elif row.h == 'stock':
                    val = '=IF(ISNUMBER(<q2>[@]),<q2>[@],"")'
                elif row.h == 'avg':
                    val = '=AVERAGE(<q1>[@],<q2>[@])'
                elif row.histf is not None:
                    val = row.histf
            elif t in ('QE', 'AE', 'F'):
                val = row.fc.get(t)
            if isinstance(val, str) and val.startswith('='):
                val = S.render(val.replace('<per>', per), col, rnum)
            cell = ws['%s%d' % (col, rnum)]
            cell.value = val
            if row.fmt:
                cell.number_format = row.fmt
            elif cell.number_format == NUM0:
                cell.number_format = NUM1
            # colour convention: blue = hard-coded input, black = formula (keep TDY grey for historical y/y)
            fc = cell.font.color.rgb if (cell.font.color is not None and cell.font.color.type == 'rgb') else None
            if val is None:
                continue
            if fc in (None, BLUE, BLACK):
                if isinstance(val, (int, float)):
                    set_color(cell, BLUE, bold=True if t in ('QE', 'AE', 'F') else None)
                else:
                    set_color(cell, BLACK)
        # CAGR / average columns
        if row.cagr == 'g':
            ws['AC%d' % rnum] = '=IFERROR(IF(AND(Z{0}>0,P{0}>0),(Z{0}/P{0})^(1/5)-1,"n/m"),"n/m")'.format(rnum)
            ws['AD%d' % rnum] = '=IFERROR(IF(AND(P{0}>0,C{0}>0),(P{0}/C{0})^(1/3)-1,"n/m"),"n/m")'.format(rnum)
        elif row.cagr == 'avg':
            ws['AC%d' % rnum] = '=IFERROR(AVERAGE(V{0},W{0},X{0},Y{0},Z{0}),"n/m")'.format(rnum)
            ws['AD%d' % rnum] = '=IFERROR(AVERAGE(C{0},D{0},J{0},P{0}),"n/m")'.format(rnum)
        for col in ('AC', 'AD'):
            if ws['%s%d' % (col, rnum)].value is not None:
                copy_style(TM['%s%d' % ('CV' if col == 'AC' else 'CW', 11 if row.cagr == 'g' else 19)], ws['%s%d' % (col, rnum)])
                if row.cagr == 'avg':
                    ws['%s%d' % (col, rnum)].number_format = '0.0%;\\(0.0%\\);"n/m"'

    # ---------- scenario table (AG:AK) ----------
    tmap = {'title': 9, 'ohead': 11, 'oval': 12, 'phead': 15, 'pval': 17, 'pval%': 16, 'lhead%': 20, 'lhead': 68, 'lval%': 21, 'lval': 69, 'levrow': 81}
    for (r, label, bull, base, bear, note, kind) in SCENARIO_TABLE:
        tr = tmap[kind]
        for col, src, v in [('AG', 'DB', label), ('AH', 'DC', bull), ('AI', 'DD', base), ('AJ', 'DE', bear), ('AK', 'DF', note)]:
            c = ws['%s%d' % (col, r)]
            copy_style(TM['%s%d' % (src, tr)], c)
            c.value = v
        if kind == 'oval':
            for col in ('AH', 'AI', 'AJ'):
                ws['%s%d' % (col, r)].number_format = NUM1
        if kind == 'lhead':
            for col in ('AH', 'AJ'):
                ws['%s%d' % (col, r)].number_format = '\\+#,##0;\\-#,##0'
        if kind in ('lval',):
            for col in ('AH', 'AI', 'AJ'):
                ws['%s%d' % (col, r)].number_format = NUM0
    ws['AG11'].comment = note_comment('FY26 outlook per Karman Q2-2026 earnings release (Form 8-K Ex. 99.1, 6-Aug-2026): revenue $730–745m and non-GAAP Adjusted EBITDA $215.0–222.5m, '
                                      'excluding the impact of any future acquisitions (raised from $720–735m / $208.5–219.5m in the Q1-26 release of 12-May-2026).')
    return S


def main(out, scen='Base'):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    S = build_model(wb)
    wb['Model']['AA2'] = scen
    import sheets
    sheets.build_bbb(wb, S)
    sheets.build_dcf(wb, S)
    sheets.build_charts(wb, S)
    wb.save(out)
    return S


if __name__ == '__main__':
    main(sys.argv[1] if len(sys.argv) > 1 else 'KRMN_Model.xlsx')
