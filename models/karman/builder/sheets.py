"""Bull-Base-Bear, DCF and Charts sheets — same cell layout and styles as the TDY template."""
from copy import copy
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from lib import TDY, copy_style, note_comment, set_color, BLUE, BLACK, GREEN

SNAP = {}   # filled by the driver: {'Bull': {(row, colidx): value}, 'Bear': {...}}


def clone_layout(wb, name, title_override=None):
    src = TDY[name]
    ws = wb.create_sheet(name)
    for row in src.iter_rows():
        for c in row:
            if c.has_style:
                copy_style(c, ws[c.coordinate])
    for k, d in src.column_dimensions.items():
        ws.column_dimensions[k].width = d.width
    for k, d in src.row_dimensions.items():
        if d.height:
            ws.row_dimensions[k].height = d.height
    ws.sheet_view.showGridLines = src.sheet_view.showGridLines
    ws.sheet_view.zoomScale = src.sheet_view.zoomScale
    return ws


def M(S, key, col):
    return "Model!%s%d" % (col, S.r(key))


HIST_COLS = ['C', 'D', 'J', 'P']
YRS = ['2025A', '2026E', '2027E', '2028E', '2029E', '2030E']
MCOLS = ['P', 'V', 'W', 'X', 'Y', 'Z']


def bbb_rows(S):
    # (row, label, model key, kind, hist pct key for LT stats)
    return [
        (9, 'Revenue', 'rev', 'v', None),
        (10, '   Revenue y/y %', 9, 'yy', None),
        (11, '   Organic growth % (where disclosed / modelled)', 'gm_org', 'pct', None),
        (12, 'Adjusted EBITDA (company definition)', 'adjebitda', 'v', None),
        (13, '   Adjusted EBITDA margin %', 12, 'm', 'gm_adj'),
        (15, 'Operating Income (GAAP)', 'oi', 'v', None),
        (16, '   Operating income y/y %', 15, 'yy', None),
        (17, '   Operating margin (GAAP) %', 15, 'm', 'gm_oim'),
        (19, 'Adjusted Net Income (company definition, not tax-effected)', 'adjni', 'v', None),
        (21, 'FDSO (m)', 'sh_d', 'v', None),
        (23, 'Adjusted EPS — Diluted ($, company definition)', 'adjeps', 'v', None),
        (24, '   EPS y/y %', 23, 'yy', None),
        (31, 'Free Cash Flow (CFO − capex)', 'fcf', 'v', None),
        (32, '   FCF %', 31, 'm', 'fcf_m'),
        (33, '   FCF / share ($)', (31, 21), 'div', None),
        (35, 'Adjusted EBITA (model, ex-amortization)', 'ebita', 'v', None),
        (36, '   Adjusted EBITA margin %', 35, 'm', 'gm_ebitam'),
        (38, 'Net Debt (incl. finance leases)', 'nd', 'v0', None),
        (39, '   Net Debt / Adjusted EBITDA (x)', (38, 12), 'div', None),
        (41, 'ROIC (Adj. EBITA NOPAT / book IC)', 'roic', 'link', 'roic'),
        (42, 'RONTA', 'ronta', 'link', 'ronta'),
    ]


def build_bbb(wb, S):
    ws = clone_layout(wb, 'Bull-Base-Bear')
    ws['B2'] = 'Karman Holdings (KRMN) — Bull / Base / Bear Case P&L Summary'
    ws['B3'] = ('Company non-GAAP definitions (Adj. EBITDA, Adj. EPS) plus model Adj. EBITA · USD millions · Live panel linked to Model — toggle scenario via Model!AA2. '
                'Bull and Bear panels are value snapshots generated from the Model with the switch set to each case.')
    ws['J2'], ws['K2'] = 'Share price ($):', '=Model!$V$%d' % S.r('px')
    ws['J3'], ws['K3'] = 'As of:', '=TODAY()'
    ws['B5'], ws['C5'] = 'Scenario:', '=Model!AA2'
    ws['N5'], ws['O5'], ws['P5'] = 'Scenario:', 'Bull', 'Snapshot (values) — Model run with switch = Bull'
    ws['Z5'], ws['AA5'], ws['AB5'] = 'Scenario:', 'Bear', 'Snapshot (values) — Model run with switch = Bear'
    panels = [(2, 'live'), (14, 'Bull'), (26, 'Bear')]   # first data column index (C=3) minus 1 for label col
    hdr = ['(USDm, except EPS and shares)'] + YRS + ["25-'30 CAGR", 'LT Hist Avg', 'LT Hist Min', 'LT Hist Max']
    for off, kind in panels:
        for i, h in enumerate(hdr):
            ws.cell(7, off + i, h)
    rows = bbb_rows(S)
    for off, kind in panels:
        lab = off           # label column index
        c0 = off + 1        # 2025A column
        cl = lambda j: ws.cell(1, c0 + j).column_letter
        for (r, label, key, k, lt) in rows:
            ws.cell(r, lab, label)
            for j, mc in enumerate(MCOLS):
                col = cl(j)
                if kind == 'live':
                    if k in ('v', 'v0', 'pct', 'link'):
                        f = '=%s' % M(S, key, mc)
                    elif k == 'yy':
                        f = None if j == 0 else '=IFERROR(%s%d/%s%d-1,"")' % (col, key, cl(j - 1), key)
                    elif k == 'm':
                        f = '=IFERROR(%s%d/%s9,"")' % (col, key, col)
                    elif k == 'div':
                        f = '=IFERROR(%s%d/%s%d,"")' % (col, key[0], col, key[1])
                    ws['%s%d' % (col, r)] = f
                else:
                    v = SNAP.get(kind, {}).get((r, 3 + j))
                    ws['%s%d' % (col, r)] = v
            # CAGR
            ccol = cl(6)
            if k in ('v',):
                if kind == 'live':
                    ws['%s%d' % (ccol, r)] = '=IFERROR(IF(AND({a}{r}>0,{b}{r}>0),({a}{r}/{b}{r})^(1/5)-1,"n/m"),"n/m")'.format(a=cl(5), b=cl(0), r=r)
                else:
                    ws['%s%d' % (ccol, r)] = SNAP.get(kind, {}).get((r, 9))
            if lt:
                refs = ','.join('Model!%s%d' % (hc, S.r(lt)) for hc in HIST_COLS)
                for jj, fn in ((7, 'AVERAGE'), (8, 'MIN'), (9, 'MAX')):
                    if kind == 'live':
                        ws['%s%d' % (cl(jj), r)] = '=IFERROR(%s(%s),"")' % (fn, refs)
                    else:
                        ws['%s%d' % (cl(jj), r)] = SNAP.get(kind, {}).get((r, 3 + jj))
        # valuation block (formulas in every panel)
        pe = {'live': (40, 40), 'Bull': (50, 50), 'Bear': (28, 28)}[kind]
        ws.cell(26, lab, 'P / E (x) — assumed (Adjusted EPS, company definition)')
        ws.cell(27, lab, 'Price Target ($)')
        ws.cell(28, lab, '   Upside / (Downside) vs current')
        ws.cell(29, lab, 'IRR (to Dec-Y-1)')
        ws.cell(34, lab, '   FCF Yield (on PT)')
        for j, yr in ((4, 2029), (5, 2030)):
            col = cl(j)
            ws['%s26' % col] = pe[j - 4]
            ws['%s27' % col] = '=%s26*%s23' % (col, col)
            ws['%s28' % col] = '=IFERROR(%s27/$K$2-1,"")' % col
            ws['%s29' % col] = '=IFERROR((%s27/$K$2)^(1/((DATE(%d-1,12,31)-$K$3)/365.25))-1,"")' % (col, yr)
        for j in (2, 3, 4, 5):
            col = cl(j)
            ws['%s34' % col] = '=IFERROR(%s33/%s$27,"")' % (col, col)
        ws.cell(44, lab, 'Implied Market Cap (USDm)')
        ws.cell(45, lab, '(+) Net Debt (incl. finance leases)')
        ws.cell(46, lab, 'Implied Enterprise Value (USDm)')
        ws.cell(47, lab, 'Implied EV / Adjusted EBITDA (x)')
        for j in (4, 5):
            col = cl(j)
            ws['%s44' % col] = '=IFERROR(%s27*%s21,"")' % (col, col)
            ws['%s45' % col] = '=%s38' % col
            ws['%s46' % col] = '=IFERROR(%s44+%s45,"")' % (col, col)
            ws['%s47' % col] = '=IFERROR(%s46/%s12,"")' % (col, col)
    # TDY had P/E inputs in E26/F26 cleared? keep only 2029E / 2030E (G,H)
    for a in ('E26', 'F26', 'Q26', 'R26', 'AC26', 'AD26'):
        ws[a] = None
    ws['N26'].comment = note_comment('Bull P/E on Adjusted EPS (company definition — pre-tax add-backs, no amortization add-back). KRMN has traded at very high multiples since the Feb-2025 IPO '
                                     '(52-week range $31.66–$118.38); 50x assumes the growth premium is sustained. Base 40x, Bear 28x.')
    for a in ('C9', 'C12'):
        pass
    return ws


def build_dcf(wb, S):
    ws = clone_layout(wb, 'DCF')
    put = lambda a, v: ws.__setitem__(a, v)
    put('B2', 'DCF VALUATION — KARMAN HOLDINGS (KRMN)')
    put('B4', '1. WACC BUILD')
    rows = [
        (5, 'Risk-free rate (US 10Y govt)', 0.045, 'Assumption (~US 10Y Treasury yield); update to market'),
        (6, 'Equity risk premium', 0.05, 'US implied ERP ~4.5–5.5% (Damodaran)'),
        (7, 'Levered beta', 1.3, 'Short trading history since the Feb-2025 IPO; high-growth defense-tech peers (AVAV, KTOS, MRCY) ~1.2–1.5'),
        (8, 'Cost of equity', '=C5+C7*C6', None),
        (9, 'Pre-tax cost of debt', 0.06, 'Citi Term Loan B at SOFR + 2.25% (Fifth Amendment, 3-Aug-26) ≈ 5.9–6.0%; finance leases ~8.2%'),
        (10, 'Tax rate', 0.245, 'Model 2027E+ effective tax rate'),
        (11, 'After-tax cost of debt', '=C9*(1-C10)', None),
        (12, 'Target D/V', 0.2, 'Net debt incl. leases ~$0.95bn vs ~$4.5bn market cap at $34.14; ~4x Adj. EBITDA, deleveraging with growth'),
        (13, 'WACC', '=C8*(1-C12)+C11*C12', None),
        (15, 'Valuation date (live)', '=TODAY()', 'Periods = (Dec-31-Year − today) / 365.25; cash flows assumed at year-end'),
    ]
    for r, b, c, d in rows:
        put('B%d' % r, b); put('C%d' % r, c); put('D%d' % r, d)
    put('B16', '2. UNLEVERED FREE CASH FLOW BUILD (USDm)')
    yrs = ['2026E', '2027E', '2028E', '2029E', '2030E', '2031E', '2032E', '2033E', '2034E', '2035E']
    cols = list('DEFGHIJKLM')
    mcol = dict(zip('DEFGH', 'VWXYZ'))
    for c, y in zip(cols, yrs):
        put('%s17' % c, y)
    put('N17', 'Terminal'); put('P17', 'CAGR 26–30E'); put('Q17', 'CAGR 26–35E')
    put('B18', 'Adjusted EBITA (Model; before acquired-intangible amortization)')
    put('B19', '   y/y %')
    put('B20', '(−) SBC, transaction, integration & other items excluded from Adj. EBITA')
    put('B21', '× (1 − tax rate)')
    put('B22', 'NOPAT  =  (Adj. EBITA − excluded costs) × (1 − t)')
    put('B23', 'plus depreciation of PP&E and finance-lease ROU amortization')
    put('B24', '− Capex (PP&E) and acquisitions (M&A lever, 2027E+)')
    put('B25', '− Δ Working capital (CF statement)')
    put('B26', 'Unlevered FCF')
    put('B27', '   y/y %')
    put('B28', 'Terminal Value (Gordon Growth)')
    put('B29', '   memo: PP&E capex only (fade years carry no new acquisitions)')
    for i, c in enumerate(cols):
        p = cols[i - 1] if i else None
        if c in mcol:
            m = mcol[c]
            put('%s18' % c, '=%s' % M(S, 'ebita', m))
            put('%s20' % c, '=-(%s-%s)' % (M(S, 'b1_tot', m), M(S, 'b1_amort', m)))
            put('%s23' % c, '=%s+%s' % (M(S, 'dep', m), M(S, 'rou', m)))
            put('%s24' % c, '=%s29' % c if c == 'D' else '=%s29+%s' % (c, M(S, 'acq', m)))
            put('%s25' % c, '=%s' % M(S, 'cf_wc', m))
            put('%s29' % c, '=%s' % M(S, 'capex', m))
        else:
            for rr in (18, 20, 23, 25, 29):
                put('%s%d' % (c, rr), '=%s%d*(1+%s$31)' % (p, rr, c))
            put('%s24' % c, '=%s29' % c)
        if i:
            put('%s19' % c, '=IFERROR(%s18/%s18-1,"")' % (c, p))
            put('%s27' % c, '=IFERROR(%s26/%s26-1,"")' % (c, p))
        put('%s21' % c, '=1-$C$10')
        put('%s22' % c, '=(%s18+%s20)*%s21' % (c, c, c))
        put('%s26' % c, '=%s22+%s23+%s24+%s25' % (c, c, c, c))
        put('%s34' % c, '=(DATE(%s,12,31)-$C$15)/365.25' % yrs[i][:4])
        put('%s35' % c, '=1/(1+$C$13)^%s34' % c)
        put('%s36' % c, '=%s26*%s35' % (c, c))
        put('%s39' % c, '=%s17' % c)
        put('%s40' % c, '=%s22' % c)
        put('%s41' % c, '=%s40-%s26' % (c, c))
        put('%s42' % c, ('=%s' % M(S, 'ic', 'P')) if c == 'D' else '=%s43' % p)
        put('%s43' % c, '=%s42+%s41' % (c, c))
        put('%s44' % c, '=IFERROR(%s40/%s42,"")' % (c, c))
        put('%s45' % c, '=IFERROR(%s41/%s40,"")' % (c, c))
        if i:
            put('%s46' % c, '=IFERROR((%s40-%s40)/%s41,"")' % (c, p, p))
    for rr in (18, 26):
        put('P%d' % rr, '=IFERROR(IF(AND($H{0}>0,$D{0}>0),($H{0}/$D{0})^(1/(VALUE(LEFT($H$17,4))-VALUE(LEFT($D$17,4))))-1,"n/m"),"n/m")'.format(rr))
        put('Q%d' % rr, '=IFERROR(IF(AND($M{0}>0,$D{0}>0),($M{0}/$D{0})^(1/(VALUE(LEFT($M$17,4))-VALUE(LEFT($D$17,4))))-1,"n/m"),"n/m")'.format(rr))
    put('N26', '=M26*(1+$C$32)'); put('N28', '=N26/($C$13-$C$32)')
    put('B31', 'Fade-period growth (2031E–35E)'); put('I31', 0.15)
    for c in 'JKL':
        put('%s31' % c, '=%s31+($M$31-$I$31)/(COLUMNS($I$31:$M$31)-1)' % chr(ord(c) - 1))
    put('M31', 0.06)
    put('B32', 'Terminal growth rate'); put('C32', 0.04)
    put('B34', 'Discount period (years)'); put('N34', '=M34')
    put('B35', 'Discount factor'); put('N35', '=1/(1+$C$13)^N34')
    put('B36', 'PV of UFCF'); put('N36', '=N28*N35')
    put('D37', 'Adjusted EBITA follows Karman’s Adj. EBITDA definition, which excludes SBC and transaction / integration costs; those are deducted again here (row 20) so NOPAT bears them. '
               'Acquisition spend from the M&A lever is deducted in 2027E–30E so that acquired sales / profit in the Model are paid for; the fade period and terminal value assume no new acquisitions. '
               'Finance-lease liabilities are treated as debt (net debt) and lease ROU amortization is added back with depreciation; lease interest sits below EBITA.')
    put('B39', 'Implied ROIC & incremental ROIC (USDm)'); put('N39', 'Terminal')
    put('B40', 'NOPAT (row 22)'); put('N40', '=M40*(1+$C$32)')
    put('B41', 'Net investment  =  NOPAT − UFCF'); put('N41', '=N40-N26')
    put('B42', 'Opening invested capital'); put('N42', '=M43')
    put('B43', 'Closing invested capital')
    put('B44', 'Implied ROIC  =  NOPAT / opening IC'); put('N44', '=IFERROR(N40/N42,"")')
    put('B45', 'Reinvestment rate  =  net inv. / NOPAT'); put('N45', '=IFERROR(N41/N40,"")')
    put('B46', 'Incremental ROIC  =  Δ NOPAT / prior-yr net inv.'); put('N46', '=IFERROR($C$32/N45,"")')
    put('B47', 'Cumulative incremental ROIC 2026E–35E'); put('C47', '=IFERROR((M40-D40)/SUM(D41:L41),"")')
    put('D48', 'Opening IC = Model 2025A book IC (equity + net debt incl. finance leases), rolled forward with DCF net investment (capex + M&A − D&A + ΔWC). Terminal incremental ROIC = g ÷ reinvestment rate. '
               'Cumulative = (NOPAT 2035E − 2026E) / Σ net investment 2026E–34E.')
    put('B49', '3. VALUATION SUMMARY')
    vrows = [(50, 'Sum of PV (explicit + fade)', '=SUM(D36:M36)'), (51, 'PV of Terminal Value', '=N36'), (52, 'Enterprise Value', '=C50+C51'),
             (53, 'Less: net debt incl. finance leases (Model 2026E YE)', '=%s' % M(S, 'nd', 'V')), (54, 'Less: non-controlling interests (none)', 0),
             (55, 'Equity Value', '=C52-C53-C54'), (56, 'Diluted shares (m, Model 2026E)', '=%s' % M(S, 'sh_d', 'V')), (57, 'Implied share price ($)', '=C55/C56'),
             (58, 'Current share price ($)', '=%s' % M(S, 'px', 'V')), (59, 'Implied upside / (downside)', '=C57/C58-1')]
    for r, b, c in vrows:
        put('B%d' % r, b); put('C%d' % r, c)
    put('B62', '4. SENSITIVITY: IMPLIED PRICE ($) — WACC × Terminal Growth')
    put('B63', 'WACC step (across)'); put('C63', 0.005)
    put('B64', 'Terminal growth step (down)'); put('C64', 0.005)
    put('B65', 'g (down) / WACC'); put('D65', '=E65-$C$63'); put('E65', '=F65-$C$63'); put('F65', '=$C$13'); put('G65', '=F65+$C$63'); put('H65', '=G65+$C$63')
    put('B66', '=B67-$C$64'); put('B67', '=B68-$C$64'); put('B68', '=$C$32'); put('B69', '=B68+$C$64'); put('B70', '=B69+$C$64')
    for r in range(66, 71):
        for c in 'DEFGH':
            put('%s%d' % (c, r), '=IFERROR((SUMPRODUCT($D$26:$M$26/(1+%s$65)^$D$34:$M$34)+($M$26*(1+$B%d)/(%s$65-$B%d))/(1+%s$65)^$M$34-$C$53-$C$54)/$C$56,"")' % (c, r, c, r, c))
    put('B73', '5. REVERSE DCF')
    put('B74', 'To solve for the market-implied path: Data ▸ What-If Analysis ▸ Goal Seek — set C57 (implied price) to the current price (C58) by changing C32 (terminal growth) or the Model scenario levers '
               '(end-market growth, Adj. EBITDA margin, M&A spend).')
    put('B76', '6. MAUBOUSSIN EV / NOPAT — TWO-STAGE FRAMEWORK')
    put('B77', 'EV/NOPAT  =  1/r  +  [g × (1 − r/ROIIC) / (r × (r − g))] × [1 − ((1+g)/(1+r))^N]')
    put('B78', 'Formula source: Mauboussin & Rappaport, "Expectations Investing" (2021); Counterpoint Global, "What Does a Price-Earnings Multiple Mean?" (2020).')
    put('B80', 'Input'); put('C80', 'Value'); put('D80', 'Rationale (Karman characteristics)')
    mrows = [(81, 'r  (Cost of capital / WACC)', '=$C$13', 'Linked from Section 1 WACC build'),
             (82, 'g  (NOPAT growth, Stage 1)', 0.12, 'Model Adj. EBITA CAGR 2025–30E well above this; 12% assumes growth normalises after the munitions / interceptor ramp'),
             (83, 'ROIIC  (Return on incremental invested capital)', 0.25, 'Organic growth needs contract-asset and capacity investment (capex ~4–5% of sales); acquisitions at ~10x EBITDA earn ~7–8% after tax'),
             (84, 'N  (Years of value creation / CAP)', 12, 'IP-rich, sole-source subsystem positions on long-lived missile, launch and submarine programs; high qualification barriers'),
             (86, 'Calculation', 'Value', None), (87, 'Steady-state component:  1 / r', '=1/C81', None), (88, 'Value-creation factor: 1 − r / ROIIC', '=1-C81/C83', None),
             (89, 'Growth multiplier: g × (1 − r/ROIIC) / (r × (r − g))', '=C82*C88/(C81*(C81-C82))', None), (90, 'Time decay: 1 − ((1+g)/(1+r))^N', '=1-((1+C82)/(1+C81))^C84', None),
             (91, 'PVGO component (growth value)', '=C89*C90', None), (92, 'Mauboussin EV / NOPAT (steady-state + PVGO)', '=C87+C91', None),
             (94, 'VALUATION  —  Applied to 2028E NOPAT', 'USDm', None), (95, '2028E Adjusted EBITA (Model)', '=%s' % M(S, 'ebita', 'X'), None),
             (96, '(−) 2028E SBC, transaction & integration costs', '=-(%s-%s)' % (M(S, 'b1_tot', 'X'), M(S, 'b1_amort', 'X')), None), (97, 'Tax rate (from WACC build)', '=$C$10', None),
             (98, '2028E NOPAT  =  (EBITA − excluded costs) × (1 − t)', '=(C95+C96)*(1-C97)', None), (99, '× Mauboussin EV/NOPAT multiple', '=C92', None),
             (100, 'Implied Enterprise Value (2028E basis)', '=C98*C99', None), (101, 'Less: YE2027 net debt incl. finance leases (Model)', '=%s' % M(S, 'nd', 'W'), None),
             (102, 'Implied Equity Value', '=C100-C101', None), (103, 'Diluted shares 2028E (m, Model)', '=%s' % M(S, 'sh_d', 'X'), None),
             (104, 'Implied share price 2028 ($)', '=C102/C103', None), (106, 'Current share price ($)', '=C58', None), (107, 'Implied upside / (downside)', '=C104/C106-1', None),
             (109, 'MARKET-IMPLIED EV/NOPAT TODAY', 'USDm', None), (110, 'Current market cap (spot × diluted shares)', '=C58*C56', None),
             (111, 'Plus: YE2026 net debt incl. finance leases', '=C53+C54', None), (112, 'Implied Enterprise Value (today)', '=C110+C111', None),
             (113, "KRMN's market-implied EV/NOPAT (on 2028E NOPAT)", '=C112/C98', None), (115, 'Multiple gap (Mauboussin − Market)', '=C92-C113', None),
             (116, 'Implied re-rating headroom', '=C92/C113-1', None)]
    for r, b, c, d in mrows:
        put('B%d' % r, b); put('C%d' % r, c)
        if d: put('D%d' % r, d)
    # colours: inputs blue, links green
    for row in ws.iter_rows():
        for c in row:
            if c.value is None or c.column < 3:
                continue
            if isinstance(c.value, (int, float)):
                set_color(c, BLUE)
            elif isinstance(c.value, str) and c.value.startswith('=') and 'Model!' in c.value:
                set_color(c, GREEN)
    return ws


def build_charts(wb, S):
    ws = clone_layout(wb, 'Charts')
    ws['B2'] = 'Karman Holdings — Operating Performance (2022A–2030E)'
    ws['B3'] = ('USDm · Linked to Model (Base / live scenario). 2022–25 actual, 2026E–30E forecast. M&A % = disclosed / derived incremental acquisition sales (FY24 RMS) and the M&A lever; '
                'organic split not disclosed for 2022–23 and 2025. Adj. EBITDA = company definition; Adj. EBITA = model.')
    mc = ['C', 'D', 'J', 'P', 'V', 'W', 'X', 'Y', 'Z']
    labels = ['2022', '2023', '2024', '2025', '2026E', '2027E', '2028E', '2029E', '2030E']
    cc = [chr(ord('C') + i) for i in range(9)]
    for c, l in zip(cc, labels):
        for r in (7, 11, 17):
            ws['%s%d' % (c, r)] = l
    spec = [(8, 'Organic %', 'gm_org', True), (9, 'M&A %', 'gm_acq', True), (12, 'Adjusted EBITDA', 'adjebitda', False), (13, 'Operating Income (GAAP)', 'oi', False),
            (14, 'Adj. EBITDA Margin %', 'gm_adj', False), (15, 'GAAP Op. Margin %', 'gm_oim', False), (18, 'ROIC', 'roic', 'na'), (19, 'RONTA', 'ronta', 'na'),
            (20, 'Adj. EBITDA margin %', 'gm_adj', 'na'), (21, 'FCF margin %', 'fcf_m', 'na')]
    for r, lab, key, mode in spec:
        ws['B%d' % r] = lab
        for c, m in zip(cc, mc):
            ref = M(S, key, m)
            if mode is True:
                ws['%s%d' % (c, r)] = '=IFERROR(%s+0,0)' % ref
            elif mode == 'na':
                ws['%s%d' % (c, r)] = '=IFERROR(%s+0,NA())' % ref
            else:
                ws['%s%d' % (c, r)] = '=%s' % ref
    # blank out TDY columns beyond K
    for r in range(7, 22):
        for c in [chr(ord('L') + i) for i in range(11)]:
            ws['%s%d' % (c, r)] = None
    cats = Reference(ws, min_col=3, max_col=11, min_row=7)
    pal = ['2F3E46', '84A98C', '52796F', 'CAD2C5']
    ch1 = BarChart(); ch1.type = 'col'; ch1.grouping = 'stacked'; ch1.overlap = 100
    ch1.title = 'Revenue Growth Decomposition: Organic vs M&A'
    for r in (8, 9):
        ch1.add_data(Reference(ws, min_col=2, max_col=11, min_row=r), titles_from_data=True, from_rows=True)
    ch1.set_categories(cats); ch1.y_axis.number_format = '0%'; ch1.width, ch1.height = 22, 7.5
    ch2 = BarChart(); ch2.type = 'col'; ch2.grouping = 'clustered'
    ch2.title = 'Adjusted EBITDA & GAAP Operating Income (USDm) with Margins (%)'
    for r in (12, 13):
        ch2.add_data(Reference(ws, min_col=2, max_col=11, min_row=r), titles_from_data=True, from_rows=True)
    ch2.set_categories(cats)
    ln = LineChart()
    for r in (14, 15):
        ln.add_data(Reference(ws, min_col=2, max_col=11, min_row=r), titles_from_data=True, from_rows=True)
    ln.y_axis.axId = 200; ln.y_axis.number_format = '0%'; ln.y_axis.crosses = 'max'
    ch2 += ln; ch2.width, ch2.height = 22, 7.5
    ch3 = LineChart(); ch3.title = 'Returns Profile: ROIC, RONTA, Adj. EBITDA Margin & FCF Margin'
    for r in (18, 19, 20, 21):
        ch3.add_data(Reference(ws, min_col=2, max_col=11, min_row=r), titles_from_data=True, from_rows=True)
    ch3.set_categories(cats); ch3.y_axis.number_format = '0%'; ch3.width, ch3.height = 22, 7.5
    for ch in (ch1, ch2, ch3):
        for i, s in enumerate(ch.series):
            s.graphicalProperties = GraphicalProperties(solidFill=pal[i % 4])
            s.graphicalProperties.line = LineProperties(solidFill=pal[i % 4])
    for i, s in enumerate(ln.series):
        s.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=['C00000', 'E9A23B'][i], w=28575))
    for i, s in enumerate(ch3.series):
        s.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=pal[i % 4], w=28575))
    ws.add_chart(ch1, 'B23'); ws.add_chart(ch2, 'B44'); ws.add_chart(ch3, 'B65')
    for row in ws.iter_rows(min_row=8, max_row=21, min_col=3, max_col=11):
        for c in row:
            if c.value is not None:
                set_color(c, GREEN)
    return ws
