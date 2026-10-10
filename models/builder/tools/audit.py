import openpyxl, sys, json
from collections import Counter
f=sys.argv[1]
wb=openpyxl.load_workbook(f,data_only=True)
ws=wb['Model']
errs=Counter(); ex={}
for row in ws.iter_rows():
    for c in row:
        if isinstance(c.value,str) and c.value.startswith('#') or (isinstance(c.value,str) and c.value in ('#VALUE!','#REF!','#NAME?','#DIV/0!','#N/A','Err:502','Err:504')):
            errs[ws.cell(c.row,2).value]+=1; ex.setdefault(ws.cell(c.row,2).value,c.coordinate)
print('ERRORS', len(errs)); 
for k,v in errs.most_common(60): print('  ',v,k,ex[k])
# checks
for r in range(1,ws.max_row+1):
    lab=ws.cell(r,2).value
    if isinstance(lab,str) and ('Check' in lab or 'tie-out' in lab):
        bad=[(ws.cell(6,c).value, ws.cell(r,c).value) for c in range(3,ws.max_column+1) if isinstance(ws.cell(r,c).value,(int,float)) and abs(ws.cell(r,c).value)>(0.011 if 'EPS' in lab else 0.15)]
        print(('OK  ' if not bad else 'BAD ')+lab[:80], bad[:12], len(bad))
for n in ['DCF','Bull-Base-Bear','Charts']:
    w=wb[n]; e=[(c.coordinate,c.value) for row in w.iter_rows() for c in row if isinstance(c.value,str) and (c.value.startswith('#') or c.value.startswith('Err'))]
    print(n,'errors',len(e),e[:10])
