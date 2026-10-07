"""Shared finishing steps: Bull/Bear snapshots and cached-value injection."""
import os, sys, zipfile, shutil, re
import openpyxl
from lxml import etree
sys.path.insert(0, os.path.dirname(__file__))
from recalc import recalc

NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'


def live_panel(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb['Bull-Base-Bear']
    out = {}
    for r in range(8, ws.max_row + 1):
        out[r] = [ws.cell(r, c).value for c in range(3, 13)]   # C..L
    return out


def write_snapshots(wb, tmpl_bbb, bull, bear):
    ws = wb['Bull-Base-Bear']
    for vals, start in ((bull, 15), (bear, 27)):           # O.., AA..
        for r, lst in vals.items():
            if r == 26:      # P/E assumptions are panel-specific inputs
                continue
            for k, v in enumerate(lst):
                tv = tmpl_bbb.cell(r, start + k).value
                bull_tv = tmpl_bbb.cell(r, 15 + k).value
                fill_bear = start == 27 and r in (40, 41, 42) and tmpl_bbb.cell(r, 26).value and isinstance(bull_tv, (int, float))
                if (isinstance(tv, (int, float)) and not isinstance(tv, bool)) or fill_bear:
                    cell = ws.cell(r, start + k)
                    cell.value = v if isinstance(v, (int, float)) else None


def inject_values(src_xlsx, recalc_xlsx, out_xlsx):
    """Copy cached formula results from the LibreOffice-recalculated file into the openpyxl file."""
    wbv = openpyxl.load_workbook(recalc_xlsx, data_only=True)
    zin = zipfile.ZipFile(src_xlsx)
    # map sheet name -> xml path
    wbxml = etree.fromstring(zin.read('xl/workbook.xml'))
    rels = etree.fromstring(zin.read('xl/_rels/workbook.xml.rels'))
    rid2t = {r.get('Id'): r.get('Target') for r in rels}
    sheets = {}
    for s in wbxml.iter('{%s}sheet' % NS):
        rid = s.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
        tgt = rid2t[rid]
        sheets['xl/' + tgt.lstrip('/').replace('xl/', '')] = s.get('name')
    tmp = out_xlsx + '.tmp'
    zout = zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED)
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename in sheets:
            ws = wbv[sheets[item.filename]]
            root = etree.fromstring(data)
            for c in root.iter('{%s}c' % NS):
                f = c.find('{%s}f' % NS)
                if f is None:
                    continue
                v = ws[c.get('r')].value
                old = c.find('{%s}v' % NS)
                if old is not None:
                    c.remove(old)
                if v is None:
                    continue
                ve = etree.SubElement(c, '{%s}v' % NS)
                if isinstance(v, bool):
                    c.set('t', 'b'); ve.text = '1' if v else '0'
                elif isinstance(v, (int, float)):
                    if 't' in c.attrib:
                        del c.attrib['t']
                    ve.text = repr(float(v)) if isinstance(v, float) else str(v)
                elif isinstance(v, str) and v.startswith('#'):
                    c.set('t', 'e'); ve.text = v
                else:
                    c.set('t', 'str'); ve.text = str(v)
            data = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
        zout.writestr(item, data)
    zout.close()
    shutil.move(tmp, out_xlsx)
