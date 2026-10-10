"""Generic builder that lays out a company model in the MLM template format.

Styles are cloned from prototype rows of the template's Model sheet so that fonts,
fills, borders and number formats match the template exactly.

Formula mini-language (used in row hist/fc callables via M.f(...)):
  [key]          same column, row of `key`
  [key@py]       prior-year same period (quarter -> same quarter last year; FY -> prior FY)
  [key@pq]       previous quarter column (quarter cols only; crosses year boundary)
  [key@$LABEL]   absolute reference to the column whose header is LABEL (e.g. $Q2/26)
  [key@LABEL]    relative-free reference to column LABEL (no $)
  [key@+1]/[key@-1] next / previous physical column
  {c}            current column letter
  SUMQ[key]      SUM of the four quarter columns of the current FY column
  SCEN(skey)     CHOOSE(MATCH($SW,{"Bull","Base","Bear"},0),bull,base,bear) for scenario-table row skey
"""
import re, copy
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

NF = {
    'num': '#,##0;\\(#,##0\\);\\-',
    'num1': '#,##0.0;\\(#,##0.0\\);\\-',
    'num2': '0.00;\\(0.00\\);\\-',
    'eps': '0.00;\\(0.00\\);\\-',
    'pct': '0.0%;\\(0.0%\\);\\-',
    'pct0': '0%;\\(0%\\);"n/m"',
    'x': '0.0\\x;\\(0.0"x)";\\-',
    'x2': '0.00\\x;\\(0.00"x)";\\-',
    'days': '0\\d;\\(0")d";\\-',
    'int': '0',
    'gen': 'General',
}
BLUE = 'FF0000FF'; BLACK = 'FF000000'; GREY = 'FFA6A6A6'; GREEN = 'FF008000'
CREAM = 'FFFFFDE7'; SAGE = 'FFCAD2C5'

# prototype rows in the template Model sheet
PROTO = {
    'section': 157, 'block': 113, 'total': 114, 'sub': 162, 'line': 161, 'pct': 12,
    'pctline': 20, 'check': 214, 'pub': 213, 'def': 215, 'blank': 9, 'label': 184,
    'val': 433, 'input': 25,
}


class Col:
    def __init__(self, idx, label, kind, year, q=None):
        self.idx = idx; self.letter = L(idx); self.label = label
        self.kind = kind  # 'A' hist annual, 'Q' hist quarter, 'QE' fc quarter, 'AE' fc annual (FY with fc quarters or pure fc)
        self.year = year; self.q = q

    @property
    def is_fc(self):
        return self.kind in ('QE', 'AE')

    @property
    def is_q(self):
        return self.kind in ('Q', 'QE')

    @property
    def is_a(self):
        return self.kind in ('A', 'AE')

    def __repr__(self):
        return f'Col({self.letter},{self.label},{self.kind})'


class Layout:
    def __init__(self, annual_only_years, q_years, cur_year, cur_actual_q, fc_years, first_col=3):
        cols = []; i = first_col
        for y in annual_only_years:
            cols.append(Col(i, y, 'A', y)); i += 1
        for y in q_years:
            for q in range(1, 5):
                cols.append(Col(i, f'Q{q}/{str(y)[2:]}', 'Q', y, q)); i += 1
            cols.append(Col(i, y, 'A', y)); i += 1
        for q in range(1, 5):
            if q <= cur_actual_q:
                cols.append(Col(i, f'Q{q}/{str(cur_year)[2:]}', 'Q', cur_year, q))
            else:
                cols.append(Col(i, f'Q{q}/{str(cur_year)[2:]}E', 'QE', cur_year, q))
            i += 1
        cols.append(Col(i, f'{cur_year}E', 'AE', cur_year)); i += 1
        for y in fc_years:
            cols.append(Col(i, f'{y}E', 'AE', y)); i += 1
        self.cols = cols
        self.cur_year = cur_year; self.cur_actual_q = cur_actual_q
        self.last_col = i - 1
        self.notes = i            # notes column
        self.gap1 = i + 1
        self.cagr = [i + 2, i + 3, i + 4, i + 5]
        self.gap2 = i + 6; self.gap3 = i + 7
        self.scen_label = i + 8; self.scen = [i + 9, i + 10, i + 11]; self.scen_note = i + 12
        self.by_label = {str(c.label): c for c in cols}
        self.fy = {c.year: c for c in cols if c.is_a}
        self.qcols = {}
        for c in cols:
            if c.is_q:
                self.qcols.setdefault(c.year, []).append(c)
        self.qseq = [c for c in cols if c.is_q]
        self.hist_years = [c.year for c in cols if c.kind == 'A']
        self.fc_years = [c.year for c in cols if c.kind == 'AE']

    def py(self, c):
        if c.is_q:
            for x in self.qcols.get(c.year - 1, []):
                if x.q == c.q:
                    return x
            return None
        return self.fy.get(c.year - 1)

    def pq(self, c):
        i = self.qseq.index(c)
        return self.qseq[i - 1] if i > 0 else None

    def col(self, label):
        return self.by_label[str(label)]


class Model:
    def __init__(self, wb, tmpl_ws, ws, layout, switch_cell='CF2'):
        self.wb = wb; self.t = tmpl_ws; self.ws = ws; self.lay = layout
        self.rows = {}        # key -> row number
        self.specs = []       # list of row specs
        self.r = 8
        self.scen_rows = {}   # scenario key -> row number in scenario table
        self.switch = '$' + L(layout.notes) + '$2'
        self.comments = []

    # ---------- formula helpers ----------
    def scen(self, skey):
        r = self.scen_rows[skey]
        a, b, c = [L(x) for x in self.lay.scen]
        return f'CHOOSE(MATCH({self.switch},{{"Bull","Base","Bear"}},0),${a}${r},${b}${r},${c}${r})'

    def f(self, s, c):
        lay = self.lay

        def ref(m):
            key, mod = m.group(1), m.group(2)
            if key not in self.rows:
                raise KeyError(f'unknown row key {key} in {s}')
            r = self.rows[key]
            if mod is None:
                return f'{c.letter}{r}'
            if mod == 'py':
                p = lay.py(c)
                return f'{p.letter}{r}' if p else 'NA()'
            if mod == 'pq':
                p = lay.pq(c)
                return f'{p.letter}{r}'
            if mod.startswith('+') or (mod.startswith('-') and mod[1:].isdigit()):
                return f'{L(c.idx + int(mod))}{r}'
            if mod.startswith('$'):
                x = lay.col(mod[1:])
                return f'${x.letter}${r}'
            if mod.startswith('#'):
                return f'{mod[1:]}{r}'
            x = lay.col(mod)
            return f'{x.letter}{r}'

        def sumq(m):
            key = m.group(1); r = self.rows[key]
            qs = lay.qcols[c.year]
            return 'SUM(' + ','.join(f'{x.letter}{r}' for x in qs) + ')'

        s = re.sub(r'SUMQ\[([A-Za-z0-9_]+)\]', sumq, s)
        s = re.sub(r'SCEN\(([A-Za-z0-9_]+)\)', lambda m: self.scen(m.group(1)), s)
        s = re.sub(r'\[([A-Za-z0-9_]+)(?:@([^\]]+))?\]', ref, s)
        s = s.replace('{c}', c.letter)
        return s

    # ---------- row registration ----------
    def add(self, key=None, label='', style='line', fmt='num', hist=None, fc=None, note=None,
            cagr=None, comment=None, outline=0, qfc=None, afc=None, hist_q=None, hist_a=None,
            label_color=None, bold=None, link_cols=None, skip_hist_formula_in_q=False):
        """hist / fc: callable(col)->value|formula|None, or dict col_label->value, or constant.
        qfc / afc override fc for quarter / annual forecast columns. hist_q / hist_a override hist."""
        spec = dict(key=key, label=label, style=style, fmt=fmt, hist=hist, fc=fc, note=note, cagr=cagr,
                    comment=comment, outline=outline, qfc=qfc, afc=afc, hist_q=hist_q, hist_a=hist_a,
                    row=self.r, label_color=label_color, bold=bold)
        if key:
            if key in self.rows:
                raise KeyError('dup key ' + key)
            self.rows[key] = self.r
        self.specs.append(spec)
        self.r += 1
        return spec

    def blank(self, n=1):
        for _ in range(n):
            self.add(style='blank')

    # ---------- rendering ----------
    def _clone(self, proto_row, col_src, dst):
        src = self.t.cell(proto_row, col_src)
        dst._style = copy.copy(src._style)

    def _proto_col(self, c, kind):
        """map our column to a template prototype column index"""
        # template: C hist annual, J hist quarter, BY fc quarter, CA fc annual w/ quarters, CB fc annual
        if kind == 'label':
            return 2
        if c.kind == 'A':
            return 74   # BV hist annual (quarterly era)
        if c.kind == 'Q':
            return 75   # BW
        if c.kind == 'QE':
            return 77   # BY
        return 80       # CB

    def _resolve(self, v, c):
        if v is None:
            return None
        if callable(v):
            v = v(c)
        elif isinstance(v, dict):
            v = v.get(str(c.label), v.get(c.label))
        if isinstance(v, str) and v.startswith('='):
            v = self.f(v, c)
        return v

    def render(self):
        ws = self.ws; lay = self.lay; t = self.t
        for sp in self.specs:
            r = sp['row']; st = sp['style']
            proto = PROTO.get(st, 161)
            # label
            lab = ws.cell(r, 2)
            self._clone(proto, 2, lab)
            lab.value = sp['label'] if sp['label'] != '' else None
            if sp['label_color']:
                f = copy.copy(lab.font); f.color = sp['label_color']; lab.font = f
            if st in ('section', 'block'):
                for c in lay.cols:
                    self._clone(proto, self._proto_col(c, 'v'), ws.cell(r, c.idx))
                self._clone(proto, 84, ws.cell(r, lay.notes))
                if sp['note']:
                    ws.cell(r, lay.notes).value = sp['note']
                continue
            if st == 'blank':
                for c in lay.cols:
                    self._clone(proto, self._proto_col(c, 'v'), ws.cell(r, c.idx))
                self._clone(161, 84, ws.cell(r, lay.notes))
                continue
            for c in lay.cols:
                cell = ws.cell(r, c.idx)
                self._clone(proto, self._proto_col(c, 'v'), cell)
                if c.is_fc:
                    v = sp['fc']
                    if c.kind == 'QE' and sp['qfc'] is not None:
                        v = sp['qfc']
                    if c.kind == 'AE' and sp['afc'] is not None:
                        v = sp['afc']
                else:
                    v = sp['hist']
                    if c.kind == 'Q' and sp['hist_q'] is not None:
                        v = sp['hist_q']
                    if c.kind == 'A' and sp['hist_a'] is not None:
                        v = sp['hist_a']
                val = self._resolve(v, c)
                cell.value = val
                cell.number_format = NF[sp['fmt']]
                # colour convention: inputs blue, formulas black, y/y grey in history
                f = copy.copy(cell.font)
                if val is None:
                    pass
                elif isinstance(val, str) and val.startswith('='):
                    if '!' in val:
                        f.color = GREEN
                    elif st == 'pct' and not c.is_fc:
                        f.color = GREY
                    else:
                        f.color = BLACK
                elif isinstance(val, (int, float)):
                    f.color = BLUE
                else:
                    f.color = BLACK
                if sp['bold'] is not None:
                    f.b = sp['bold']
                cell.font = f
                if c.is_fc and st not in ('total',):
                    cell.fill = PatternFill('solid', fgColor=CREAM)
                al = copy.copy(cell.alignment); al.horizontal = 'right'; cell.alignment = al
            # notes
            nc = ws.cell(r, lay.notes)
            self._clone(161, 84, nc)
            if sp['note']:
                nc.value = sp['note']
                if isinstance(sp['note'], str) and sp['note'].startswith('='):
                    nc.value = sp['note']
            # CAGR / averages
            if sp['cagr'] and sp['key']:
                self._cagr(sp)
            if sp['comment']:
                lab.comment = Comment(sp['comment'], 'Model')
                lab.comment.width = 400; lab.comment.height = 160
            if sp['outline']:
                ws.row_dimensions[r].outlineLevel = sp['outline']

    def _cagr(self, sp):
        ws = self.ws; lay = self.lay; r = sp['row']
        last_hist = max(lay.hist_years); last_fc = max(lay.fc_years)
        first_hist = min(lay.hist_years)
        A = lambda y: f'{lay.fy[y].letter}{r}'
        spans = [(last_hist, last_fc, 5), (last_hist - 5, last_hist, 5), (last_hist - 10, last_hist, 10),
                 (first_hist, last_hist, last_hist - first_hist)]
        for k, (a, b, n) in enumerate(spans):
            cell = ws.cell(r, lay.cagr[k])
            src = self.t.cell(PROTO['total'] if sp['style'] in ('total', 'sub') else 11, 86)
            cell._style = copy.copy(src._style)
            if a not in lay.fy or b not in lay.fy:
                continue
            if sp['cagr'] == 'growth':
                cell.value = f'=IFERROR(({A(b)}/{A(a)})^(1/{n})-1,"n/m")'
                cell.number_format = NF['pct0']
            else:  # average of annual values
                ys = [y for y in range(a + (1 if k == 0 else 0), b + 1) if y in lay.fy]
                if k == 0:
                    ys = [y for y in lay.fc_years]
                cell.value = '=IFERROR(AVERAGE(' + ','.join(A(y) for y in ys) + '),"n/m")'
                cell.number_format = NF['pct0'] if sp['fmt'] in ('pct',) else NF[sp['fmt']]

    # ---------- sheet-level formatting ----------
    def frame(self, title, subtitle, basis, scen_title_note):
        ws = self.ws; lay = self.lay; t = self.t
        for col in range(1, lay.scen_note + 1):
            pass
        ws['B2'] = title; ws['B2']._style = copy.copy(t['B2']._style)
        ws['B3'] = subtitle; ws['B3']._style = copy.copy(t['B3']._style)
        ws['B4'] = basis; ws['B4']._style = copy.copy(t['B4']._style)
        # header row 6
        hdr = ws.cell(6, 2); hdr._style = copy.copy(t['B6']._style); hdr.value = '(USDm)'
        for c in lay.cols:
            x = ws.cell(6, c.idx); x._style = copy.copy(t['C6']._style); x.value = c.label
        x = ws.cell(6, lay.notes); x._style = copy.copy(t['CF6']._style); x.value = 'Modelling Notes'
        last_h = max(lay.hist_years); last_f = max(lay.fc_years); first_h = min(lay.hist_years)
        heads = [('Forecast', '5Y CAGR', f"{str(last_h)[2:]}-'{str(last_f)[2:]}"),
                 ('Historical', '5Y CAGR', f"{str(last_h-5)[2:]}-'{str(last_h)[2:]}"),
                 ('Historical', '10Y CAGR', f"{str(last_h-10)[2:]}-'{str(last_h)[2:]}"),
                 ('Historical', f'{last_h-first_h}Y CAGR', f"{str(first_h)[2:]}-'{str(last_h)[2:]}")]
        for k, (a, b, cc) in enumerate(heads):
            for rr, v, src in ((5, a, 'CH5'), (6, b, 'CH6'), (7, cc, 'CH7')):
                x = ws.cell(rr, lay.cagr[k]); x._style = copy.copy(t[src]._style); x.value = v
        # scenario switch
        a = ws.cell(1, lay.notes); a._style = copy.copy(t['CF1']._style)
        a.value = 'SCENARIO SWITCH ▼ (Bull / Base / Bear) — drives all scenario-linked forecast drivers'
        s = ws.cell(2, lay.notes); s._style = copy.copy(t['CF2']._style); s.value = 'Base'
        dv = DataValidation(type='list', formula1='"Bull,Base,Bear"', allow_blank=False)
        ws.add_data_validation(dv); dv.add(s.coordinate)
        s.comment = Comment('Scenario toggle. Choose Bull / Base / Bear. Drives the FY26 outlook end-points used in the Q3/Q4-26E calibration and every scenario lever in the table to the right (columns '
                            + L(lay.scen_label) + '–' + L(lay.scen_note) + ').', 'Model')
        # widths
        ws.column_dimensions['A'].width = t.column_dimensions['A'].width
        ws.column_dimensions['B'].width = 62.0
        for c in lay.cols:
            cd = ws.column_dimensions[c.letter]
            if c.kind == 'A':
                cd.width = 14.140625
            elif c.kind == 'Q':
                cd.width = 9.5703125
            elif c.kind == 'QE':
                cd.width = 11.42578125
            else:
                cd.width = 14.140625
        ws.column_dimensions[L(lay.notes)].width = 86.85546875
        ws.column_dimensions[L(lay.gap1)].width = 2.0
        for k, w in zip(lay.cagr, (14.14, 12.43, 13.43, 13.43)):
            ws.column_dimensions[L(k)].width = w
        ws.column_dimensions[L(lay.gap2)].width = 2.0
        ws.column_dimensions[L(lay.gap3)].width = 6.0
        ws.column_dimensions[L(lay.scen_label)].width = 46.0
        for k in lay.scen:
            ws.column_dimensions[L(k)].width = 11.57
        ws.column_dimensions[L(lay.scen_note)].width = 60.0
        # group quarterly columns (collapsed like the template)
        for y, qs in lay.qcols.items():
            a, b = qs[0].letter, qs[-1].letter
            ws.column_dimensions.group(a, b, outline_level=1, hidden=True)
        ws.sheet_properties.outlinePr.summaryRight = True
        ws.sheet_properties.outlinePr.summaryBelow = True
        ws.freeze_panes = 'C8'
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 70

    # ---------- scenario table ----------
    def scen_table(self, start_row, title, blocks):
        """blocks: list of dicts. kind 'hdr' (title, bull, base, bear labels, note), 'pt' (key,label,bull,base,bear,note,fmt),
        'lever' (key,label,delta_bull,delta_bear,base_by_year{year:val},note,fmt)."""
        ws = self.ws; lay = self.lay; t = self.t
        lc, (cb, cm, cbr), nc = lay.scen_label, lay.scen, lay.scen_note
        r = start_row

        def sty(cell, src):
            cell._style = copy.copy(t[src]._style)
        sty(ws.cell(r, lc), 'CN9'); ws.cell(r, lc).value = title
        for k, v in zip((cb, cm, cbr), ('Bull', 'Base', 'Bear')):
            sty(ws.cell(r, k), 'CO9'); ws.cell(r, k).value = v
        sty(ws.cell(r, nc), 'CR9'); ws.cell(r, nc).value = 'Bull / Bear = Base + Δ (Δ inputs on each lever header). Outlook block = high / mid / low end of the FY26 ranges.'
        r += 2
        for b in blocks:
            if b['kind'] == 'hdr':
                sty(ws.cell(r, lc), 'CN11'); ws.cell(r, lc).value = b['label']
                for k, v in zip((cb, cm, cbr), b.get('cols', ('High', 'Mid', 'Low'))):
                    sty(ws.cell(r, k), 'CO11'); ws.cell(r, k).value = v
                sty(ws.cell(r, nc), 'CR11'); ws.cell(r, nc).value = b.get('note')
                r += 1
            elif b['kind'] == 'pt':
                self.scen_rows[b['key']] = r
                sty(ws.cell(r, lc), 'CN12'); ws.cell(r, lc).value = b['label']
                for k, v in zip((cb, cm, cbr), (b['bull'], b['base'], b['bear'])):
                    sty(ws.cell(r, k), 'CP12'); ws.cell(r, k).value = v
                    ws.cell(r, k).number_format = NF[b.get('fmt', 'num')]
                    if v is None:
                        ws.cell(r, k).value = None
                sty(ws.cell(r, nc), 'CR12'); ws.cell(r, nc).value = b.get('note')
                r += 1
            elif b['kind'] == 'lever':
                sty(ws.cell(r, lc), 'CN30'); ws.cell(r, lc).value = b['label']
                sty(ws.cell(r, cb), 'CO30'); ws.cell(r, cb).value = b['dbull']
                sty(ws.cell(r, cm), 'CP30'); ws.cell(r, cm).value = 'Δ →'
                sty(ws.cell(r, cbr), 'CQ30'); ws.cell(r, cbr).value = b['dbear']
                for k in (cb, cbr):
                    ws.cell(r, k).number_format = NF[b.get('fmt', 'pct')]
                sty(ws.cell(r, nc), 'CR30'); ws.cell(r, nc).value = b.get('note')
                hr = r; r += 1
                for y, v in b['base'].items():
                    self.scen_rows[f"{b['key']}_{y}"] = r
                    sty(ws.cell(r, lc), 'CN31'); ws.cell(r, lc).value = f'  {y}E'
                    sty(ws.cell(r, cb), 'CO31'); ws.cell(r, cb).value = f'={L(cm)}{r}+${L(cb)}${hr}'
                    sty(ws.cell(r, cm), 'CP31'); ws.cell(r, cm).value = v
                    sty(ws.cell(r, cbr), 'CQ31'); ws.cell(r, cbr).value = f'={L(cm)}{r}+${L(cbr)}${hr}'
                    for k in (cb, cm, cbr):
                        ws.cell(r, k).number_format = NF[b.get('fmt', 'pct')]
                    r += 1
        return r
