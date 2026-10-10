import openpyxl, sys, json
f=sys.argv[1]; rows=json.load(open(sys.argv[2])); keys=sys.argv[3].split(',')
cols=sys.argv[4].split(',') if len(sys.argv)>4 else ['2006','2015','2019','2024','2025','Q1/26','Q2/26','Q3/26E','Q4/26E','2026E','2027E','2028E','2029E','2030E']
wb=openpyxl.load_workbook(f,data_only=True); ws=wb['Model']
hdr={str(ws.cell(6,c).value):c for c in range(3,ws.max_column+1) if ws.cell(6,c).value is not None}
print(f"{'key':14s}"+''.join(f'{c:>10s}' for c in cols))
for k in keys:
    r=rows[k]; out=[]
    for c in cols:
        v=ws.cell(r,hdr[c]).value
        if isinstance(v,float): out.append(f'{v:10.3f}' if abs(v)<5 else f'{v:10.1f}')
        elif v is None: out.append(f'{"":>10s}')
        else: out.append(f'{str(v)[:9]:>10s}')
    print(f'{k:14s}'+''.join(out))
