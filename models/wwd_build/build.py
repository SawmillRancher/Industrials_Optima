"""Build WWD_Model.xlsx from the MLM template (same layout, formatting and process as the prior models)."""
import sys, copy, os, json
from openpyxl.styles import Font
from fw import *
import fw
from spec import build_spec, scen_layout, SCEN_ROWS, SC
from spec2 import build_spec2

OUT = os.path.join(os.path.dirname(BASE), 'WWD_Model.xlsx')
G26 = MAN['guidance_fy26']

# ====================================================================== inputs (V0)
def wc_inputs():
    q = DATA['2025']['x']
    rev, cogs = q['rev'], q['cogs']
    return {'dso': round(q['bs_ar'] / rev * 365), 'dio': round(q['bs_inv'] / cogs * 365), 'dpo': round(q['bs_ap'] / cogs * 365),
            'ocl': round(q['bs_ocl'] / rev, 3)}
WC = wc_inputs()
WC['note'] = (f'Forecast inputs at fiscal-year-end (30-Sep-2025) levels on FY2025 net sales / cost of goods sold (DSO {WC["dso"]}d, inventory {WC["dio"]}d, '
              f'payables {WC["dpo"]}d, other current liabilities {WC["ocl"]:.1%} of sales); year ends are compared with year ends.')
AM = MAN['amort_sched']

V0 = {
    'out_hdr': 'FY26 guidance — Q3 FY26 release (29-Jul-2026)',
    'out_note': ('Sales growth +20–23% (unchanged), adjusted EPS $9.30–9.50 (raised from $9.15–9.45), FCF $300–350m, capex ~$290m, ~61.5m shares, adjusted tax rate ~22.5%; '
                 'Aerospace sales +21–23% at ~23.5% segment margin; Industrial sales +19–21% at ~19%.'),
    'outlook': [
        ('out_g', '  FY26 net sales growth', tuple(G26['sales_growth']), 'pct', 'Guided +20–23%. Drives the Q4/26E sales calibration.'),
        ('out_eps', '  FY26 adjusted EPS ($)', tuple(G26['adj_eps']), 'ps', 'Guided $9.30–9.50. Drives the Q4/26E segment-margin calibration.'),
        ('out_ag', '  FY26 Aerospace sales growth — memo', tuple(G26['aero_growth']), 'pct', 'Guided +21–23% (information row; compare the model).'),
        ('out_ig', '  FY26 Industrial sales growth — memo', tuple(G26['ind_growth']), 'pct', 'Guided +19–21% (information row).'),
        ('out_fcf', '  FY26 free cash flow ($m) — memo', tuple(G26['fcf']), 'm', 'Guided $300–350m (information row).'),
        ('bb26', '  Q4/26E share repurchases ($m)', (300.0, 200.0, 100.0), 'm', '9M FY26 $553.4m ($198m in Q3); $1,286m left on the $1.8bn 2026 Authorization at 30-Jun-2026.'),
    ],
    'points': [
        ('pt_int', '  Q4/26E interest expense ($m)', 15.5, 'm1', 'Not guided; Q3 $14.8m; full quarter of the $250m term loan and the revolver (~$1.35bn debt at ~4.6%).'),
        ('pt_intinc', '  Q4/26E interest income ($m)', 0.6, 'm1', 'Q3 $0.6m.'),
        ('pt_othinc', '  Q4/26E other income, net ($m)', 20.0, 'm1', '9M FY26 $60.3m (~$20m a quarter; no divestiture gains).'),
        ('pt_amort', '  Q4/26E amortization of intangibles ($m)', AM['2026_rem'], 'm1', 'Q3 FY26 10-Q schedule: $8.4m remaining in FY26.'),
        ('pt_dep', '  Q4/26E depreciation ($m)', 23.0, 'm1', 'Q3 $22.5m; rising with capex.'),
        ('pt_etr', '  Q4/26E tax rate (GAAP & adjusted)', 0.24, 'pct', 'FY26 adjusted rate guided ~22.5%; 9M FY26 adjusted 21.9% → Q4 ~24% (Q3 24.2%).'),
        ('pt_capex', '  FY26 capital expenditures ($m)', G26['capex'], 'm1', 'Guided ~$290m (9M $156.3m); Spartanburg, SC campus.'),
        ('pt_restr', '  Q4/26E restructuring charges ($m)', 6.0, 'm1', 'China OH wind-down complete; servo-valve transition (~$15m total) and the start of the Santa Clarita plan (8-K 21-Sep-2026).'),
        ('pt_am', '  FY26 Aerospace segment margin guidance', G26['aero_margin'], 'pct', 'Guided ~23.5% (information row).'),
        ('pt_im', '  FY26 Industrial segment margin guidance', G26['ind_margin'], 'pct', 'Guided ~19% (information row).'),
    ],
    'lev': (1.75, 1.25, 0.75),
    'lev_note': ('Net debt $867m at 30-Jun-26 ≈ 1.0x adjusted EBITDA (FY2021–25 0.5–1.1x); company "EBITDA leverage" (total debt / TTM EBITDA) 1.6x. Target = leverage up to which surplus cash is returned '
                 'via buybacks ($1.8bn three-year authorization to Nov-2028); below it free cash flow repays the revolver. Bull = more buybacks (higher target).'),
    'blocks': [
        ('AERO_COEM', 'AERO_COEM (Commercial OEM y/y)', 0.04, -0.06, [0.14, 0.11, 0.09, 0.07, 0.06],
         '9M FY26 +29% (airframer rate increases after the FY25 Boeing work stoppage / inventory normalisation); 737 / A320 / 787 / A350 ramps, A350 spoiler actuation from Spartanburg (2027).', 'pct', False),
        ('AERO_CSERV', 'AERO_CSERV (Commercial services y/y)', 0.03, -0.06, [0.09, 0.08, 0.07, 0.06, 0.06],
         '9M FY26 +36% (FY25 +29%, partly advance purchases); high utilisation of legacy fleets, price; normalising toward high-single digits.', 'pct', False),
        ('AERO_DOEM', 'AERO_DOEM (Defense OEM y/y)', 0.03, -0.04, [0.07, 0.06, 0.05, 0.05, 0.05], '9M FY26 +7% (FY25 +38%: smart-defense / guided-weapons demand, price).', 'pct', False),
        ('AERO_DSERV', 'AERO_DSERV (Defense services y/y)', 0.02, -0.04, [0.05, 0.05, 0.04, 0.04, 0.04], '9M FY26 +9%; FY25 −2% (mix).', 'pct', False),
        ('IND_PG', 'IND_PG (Power generation y/y)', 0.04, -0.06, [0.12, 0.10, 0.08, 0.06, 0.05],
         '9M FY26 +11%; FY25 +10%: gas turbines and reciprocating gensets (data-centre and grid demand).', 'pct', False),
        ('IND_TRANS', 'IND_TRANS (Transportation y/y)', 0.04, -0.08, [0.00, 0.04, 0.04, 0.03, 0.03],
         '9M FY26 +42% (marine, rail, China OH last-time buys); China OH wound down by end of FY26 → flat FY27.', 'pct', False),
        ('IND_OG', 'IND_OG (Oil and gas y/y)', 0.03, -0.06, [0.04, 0.04, 0.03, 0.03, 0.03], '9M FY26 +18%; FY25 +14%: gas compression / midstream.', 'pct', False),
        ('AERO_M', 'AERO_M (Aerospace segment margin)', 0.01, -0.02, [0.240, 0.245, 0.250, 0.253, 0.255],
         'FY26 guided ~23.5% (FY25 21.9%, FY19 peak 20.7%); price and volume leverage vs Spartanburg start-up costs (2027).', 'pct', False),
        ('IND_M', 'IND_M (Industrial segment margin)', 0.01, -0.03, [0.190, 0.192, 0.195, 0.197, 0.200],
         'FY26 guided ~19% (FY25 14.6%, FY24 17.7%); China OH exit accretive; cyclical downside (FY18 6.2%).', 'pct', False),
        ('MNA', 'M&A_spend (bolt-on acquisitions, $m)', 100.0, -100.0, [100.0, 100.0, 100.0, 100.0, 100.0],
         'Woodward spent $40m (Safran EMA, FY25) and $121m (Valve Research, FY26) on bolt-ons besides L\'Orange ($771m, FY18). Base $100m p.a. (assumption).', 'm', True),
    ],
    'corp_pct': -0.035,
    'dep_pct': {2027: 0.023, 2028: 0.025, 2029: 0.026, 2030: 0.026, 2031: 0.026},
    'sga_pct': {2027: 0.093, 2028: 0.092, 2029: 0.091, 2030: 0.090, 2031: 0.090},
    'rd_pct': 0.041, 'sbc_pct': 0.0085, 'etr': 0.225, 'othinc': 80.0,
    'capex_pct': {2027: 0.050, 2028: 0.035, 2029: 0.032, 2030: 0.030, 2031: 0.030},
    'dps_q': 0.32, 'dps_g': 0.10, 'def_q4': 0.0, 'def_ae': -20.0, 'eq_q4': 10.0, 'eq_ae': 60.0, 'iss_q4': 0.10, 'iss_ae': 0.40, 'dil': 1.7,
    'ae_restr': {2027: 30.0, 2028: 15.0, 2029: 0.0, 2030: 0.0, 2031: 0.0},
    'amort_sched': {2027: AM['2027'], 2028: AM['2028'], 2029: AM['2029'], 2030: AM['2030'], 2031: 30.0},
    'amort_note': AM['src'] + ' FY27 $32.5m, FY28 $32.1m, FY29 $31.2m, FY30 $31.2m; FY31 $30.0m assumed ($245.6m thereafter).',
    'dv_proc': {2027: 180.0, 2028: 0.0, 2029: 0.0, 2030: 0.0, 2031: 0.0},
    'dv_nav': {2027: 16.364, 2028: 0.0, 2029: 0.0, 2030: 0.0, 2031: 0.0},
    'iss_notes': {2026: 450.0, 2027: 85.0, 2028: 90.289, 2029: 225.0, 2030: 225.0, 2031: 300.421},
    'repay': {2026: 46.903, 2027: 85.0, 2028: 90.289, 2029: 225.0, 2030: 225.0, 2031: 300.421},
    'mat': {2026: 85.0, 2027: 90.289, 2028: 225.0, 2029: 225.0, 2030: 300.421, 2031: 0.0},
    'repay_note': ('Sep-2026 Series M (EUR; $46.9m at 30-Sep-25 FX); May-2027 Series Q $85m; Sep-2028 Series N (EUR; $90.3m); FY2029 Series R $75m + U $150m; FY2030 S $75m + V $150m; '
                   'FY2031 term loan $250m (May-2031) + Series O (EUR; $50.4m). EUR notes at 30-Sep-2025 FX (assumption). 2027E+ maturities assumed refinanced with new notes (input row above).'),
    'min_cash': 300.0,
    'r_notes': {2027: 0.046, 2028: 0.046, 2029: 0.047, 2030: 0.048, 2031: 0.049},
    'r_notes_note': ('Weighted coupon ~4.6% after the $450m U/V/W issue (5.34–5.64%), term loan (SOFR + 0.875–1.75%), USD notes 4.35–4.61% and EUR notes 1.31–1.57%; '
                     'drifting up as low-coupon EUR notes mature.'),
    'wc': WC,
    'px': MAN['price']['px'],
    'px_note': ('VALUATION — share price input $382.17 (weighted average price paid for shares repurchased in June 2026, Q3 FY26 10-Q Part II Item 2 — the latest '
                'share price in an SEC filing; ~$23bn market cap); update to market'),
    'mna_note': 'Scenario lever M&A_spend (Base $100m p.a. of bolt-ons). Funded from cash / revolver (debt schedule).',
    'rev_guid_note': 'Company FY26 net sales growth guidance +20–23% (raised 29-Apr-2026, unchanged 29-Jul-2026). Drives the sales calibration.',
    'eps_guid_note': 'Company FY26 adjusted EPS guidance $9.30–9.50 (raised 29-Jul-2026). Drives the segment-margin calibration.',
}

# ====================================================================== compute engine
def hist_value(row, col):
    kw = row.kw; path = kw.get('data')
    if not path: return None
    if kw.get('ppa'):
        v = MAN['ppa'].get(pkey(col), {}).get(path.split('_', 1)[1])
        return v if kw.get('textdata') else num(v)
    v = xget(pkey(col), path)
    return v if kw.get('textdata') else num(v)

def first(kw, keys):
    for k in keys:
        if k in kw: return k
    return None

def compute(row, col):
    kw = row.kw; k = col.kind
    CTX.col = col
    if kw.get('ann') and col.q: return None, False
    if k in HISTK:
        v = hist_value(row, col)
        if v is not None: return v, True
        for key in ('hist', 'all'):
            if key in kw:
                return kw[key](), False
        return None, False
    if k == 'QE':
        if row.spec2: return None, False
        if 'qe_in' in kw:
            v = kw['qe_in']; return (v[col.q] if isinstance(v, dict) else v), True
        key = first(kw, ('qe', 'fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    if k == 'E26':
        if 'e26_in' in kw: return kw['e26_in'], True
        if 'e26' in kw: return kw['e26'](), False
        if kw.get('avg') and ('qe' in kw or 'qe_in' in kw): return f'=AVERAGE({Q1}{row.r}:{QE}{row.r})', False
        if kw.get('flow') and not row.spec2 and any(x in kw for x in ('qe', 'qe_in', 'fc')): return S26(row.key), False
        key = first(kw, ('fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    if k == 'AE':
        if 'ae_in' in kw:
            v = kw['ae_in']; return (v[col.year] if isinstance(v, dict) else v), True
        key = first(kw, ('ae', 'fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    return None, False

# ====================================================================== writers
CAGR_G = {'is_rev', 'g_rev', 'aero_rev', 'ind_rev', 'aero_coem', 'aero_cserv', 'aero_doem', 'aero_dserv', 'ind_pg', 'ind_trans', 'ind_og', 'mx_com', 'mx_def',
          'o_adj', 'g_adj', 'r_adj', 'r_ebitda', 'is_ebit', 'd_dda', 'e_adjni', 'e_eps', 'cf_cfo', 'fcf', 'bs_ta', 'bs_teq', 'sh_dil', 'is_gp', 'is_sga',
          'd_amort', 'd_dep', 'aero_oi', 'ind_oi', 'is_ni', 'eps_dil', 'dps'}
CAGR_A = {'aero_m', 'ind_m', 'g_adjm', 'o_adj_m', 'r_adj_m', 'is_gm', 'is_ebit_m', 'r_ebit_m', 'gm_rev', 'gm_aero', 'gm_ind', 'gm_gm', 'gm_ebitda_m',
          'gm_op_m', 'gm_net_m', 'gm_sga', 'gm_rd', 'gm_adjop_m', 'fcf_m', 'fcf_conv', 'capex_pct', 'ra_roic', 'rn_ronta', 'roe', 'lv_x',
          'd_sga', 'd_dep_pct', 'd_sbc_pct', 'wc_nwc_p', 'is_etr', 'e_rate', 'fcf_eb', 'corp_pct', 'aero_aft', 'mx_com_p', 'mx_def_p'}
A = {y: ANN[y].c for y in ANN}

def style_cagr(ws, r, on):
    src = 16 if on else 12
    for c in CAGR:
        ws[f'{c}{r}']._style = copy.copy(TM[f'{L(CI(c) + TSHIFT)}{src}']._style)

def write_row(ws, row):
    r = row.r; arch = ARCH[row.arch]
    clone_row_style(ws, arch, r)
    ws.cell(r, 2).value = row.label or None
    if row.kw.get('note'): ws[f'{NOTE}{r}'] = row.kw['note']
    style_cagr(ws, r, row.key in CAGR_G or row.key in CAGR_A)
    if row.arch in ('section', 'block', 'blank', 'sub') and not row.kw.get('data'):
        return
    if row.key in DEFS:
        write_defs(ws, row); return
    nf = NF[row.nf]
    for col in COLS:
        v, inp = compute(row, col)
        if v is None: continue
        cell = ws[f'{col.c}{r}']
        cell.value = v
        if not row.kw.get('textdata'): cell.number_format = nf
        if inp:
            set_color(cell, BLUE, bold=True if col.kind in FCK else None)
            lk = row.kw.get('labkey')
            if lk and col.kind in HISTK and isinstance(v, (int, float)) and abs(v) > 1e-9:
                t = xget(pkey(col), lk)
                if t: comment(cell, f'Company items ({pkey(col)}): {t}', 380, 110)
        elif row.arch == 'growth' and col.kind in HISTK:
            set_color(cell, GRAY)
        else:
            set_color(cell, BLACK)
    rc = CELL_COMMENTS.get(row.key, {})
    for c, t in rc.items(): comment(ws[f'{c}{r}'], t)
    if row.key in CAGR_G:
        ws[f'{CAGR[0]}{r}'] = f'=IFERROR(({A[2031]}{r}/{A[2026]}{r})^(1/5)-1,"n/m")'
        ws[f'{CAGR[1]}{r}'] = f'=IFERROR(({A[2025]}{r}/{A[2020]}{r})^(1/5)-1,"n/m")'
        ws[f'{CAGR[2]}{r}'] = f'=IFERROR(({A[2025]}{r}/{A[2015]}{r})^(1/10)-1,"n/m")'
        ws[f'{CAGR[3]}{r}'] = f'=IFERROR(({A[2025]}{r}/{A[2011]}{r})^(1/14)-1,"n/m")'
        for c in CAGR: ws[f'{c}{r}'].number_format = '0%;\\(0%\\);"n/m"'
    elif row.key in CAGR_A:
        ws[f'{CAGR[0]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2027, 2032))}),"n/m")'
        ws[f'{CAGR[1]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2021, 2026))}),"n/m")'
        ws[f'{CAGR[2]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2016, 2026))}),"n/m")'
        ws[f'{CAGR[3]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2011, 2026))}),"n/m")'
        for c in CAGR: ws[f'{c}{r}'].number_format = NF['pct']
    t = LABEL_SOURCES.get(row.key)
    if t: comment(ws.cell(r, 2), t)

REL = 'Source: Woodward earnings releases (Form 8-K Item 2.02, Ex. 99.1) — '
LABEL_SOURCES = {
    'aero_coem': ('Source: Woodward Forms 10-Q / 10-K, revenue note "Disaggregation of revenue" (sales by primary market, from FY2019; Q4 = fiscal year (10-K) − nine months (Q3 10-Q)). '
                  'Commercial / defense "aftermarket" renamed "services" in FY2025.'),
    'ind_pg': ('Source: Forms 10-Q / 10-K revenue note. Power generation / transportation / oil & gas introduced in the FY2023 10-K: FY2021–22 annual from the FY2023 10-K, FY2023 quarters from '
               'the FY2024 10-Q comparatives (latest filing presenting the period); FY2023+ as reported. FY2021–22 quarters on the legacy basis only (memo rows).'),
    'aero_oi': REL + 'segment net sales and earnings schedule (as originally reported; FY2018 restated in the FY2019 release for the pension-cost reclassification — memo row).',
    'ns_gaap': REL + 'segment schedule ("Nonsegment expenses"); FY2018+ adjusting items from the "Reconciliation of Net Earnings to Adjusted Net Earnings" tables.',
    'is_rev': REL + 'condensed consolidated statements of earnings, as originally reported (FY2016 and FY2018 cost-line reclassifications in the following year\'s release are not recast).',
    'd_dda': REL + 'Reconciliation of Net Earnings to EBITDA (depreciation expense; amortization of intangible assets).',
    'o_pub': REL + 'Reconciliation of Net Earnings to EBIT and Adjusted EBIT (published from Q2 FY2018).',
    'e_pub': REL + 'Reconciliation of Net Earnings / EPS to Adjusted Net Earnings / Adjusted EPS (published from Q2 FY2018).',
    'r_gebit_pub': REL + 'Reconciliation of Net Earnings to EBIT and EBITDA (published throughout).',
    'cf_ni': ('Source: release condensed cash-flow statements (year to date; discrete quarters derived, Q4 = FY − 9M); D&A from the EBITDA reconciliation; '
              'share-based compensation and deferred income taxes from the 10-K / 10-Q (XBRL, as first filed).'),
    'bs_cash': REL + 'condensed consolidated balance sheets (period ends); shares outstanding = shares issued − treasury shares (XBRL).',
    'ppa_cons': 'Source: FY2018 10-K Note 4 (L\'Orange, preliminary); Q3 FY26 10-Q Note 10 (Safran electro-mechanical actuation, Jul-2025; Valve Research, Apr-2026).',
    'sh_basic': REL + 'statements of earnings (weighted average shares). No stock splits since the 3-for-1 split of Feb-2008.',
    'ds_total': 'Source: release balance sheets; FY2025 10-K Note 15 and Q3 FY26 10-Q Note 15 (debt); 8-Ks of 28-May-2026 (revolver, term loan) and 21-Aug-2026 (Series U/V/W notes).',
    'dv_hdr': 'Source: Q3 FY26 10-Q Note 10 (pilot controls held for sale; ONTIC agreement 15-Apr-2026); 8-K 21-Sep-2026 (Santa Clarita, Item 2.05).',
}
CELL_COMMENTS = {}
def comments_setup():
    a18 = A[2018]
    CELL_COMMENTS['aero_oi'] = {a18: 'FY2018 as originally reported ($301.8m). Restated to $308.6m in the FY2019 release (pension-cost reclassification, ASU 2017-07) — see memo row.'}
    CELL_COMMENTS['ind_oi'] = {a18: 'FY2018 as originally reported ($47.9m); restated to $49.9m in the FY2019 release (ASU 2017-07). Includes L\'Orange purchase-accounting charges (adjusted out).'}
    CELL_COMMENTS['is_cogs'] = {A[2016]: 'FY2016 as originally reported; the FY2017 release shows $1,484.0m (cost of goods sold) / $174.0m (SG&A) after a reclassification.',
                                A[2023]: 'FY2023 total from the Q4 FY23 release; Σ quarters as first reported differs by $1.4m (COGS ↔ R&D reclassification within the year).'}
    CELL_COMMENTS['ind_pg'] = {A[2021]: 'Recast basis from the FY2023 10-K (prior-year comparatives).', A[2022]: 'Recast basis from the FY2023 10-K.'}

DEFS = {}
def defs_setup():
    q = lambda y, n: QC[(y, n)].c
    DEFS['o_def'] = {
        'C': ('FY2011–17: no adjusted measures published (EBIT / EBITDA / FCF only) — adjusted = GAAP EBIT; segments Aerospace / Energy (Energy renamed Industrial in FY2016)',
              'Woodward\'s releases in FY2011–17 reconcile net earnings to EBIT and EBITDA but publish no adjusted EBIT or EPS (FY2016 special charges of $16m were discussed in the text only).'),
        A[2018]: ('FY2018+: adjusted EBIT / EBITDA / EPS published (from Q2 FY2018): restructuring, L\'Orange acquisition & purchase accounting, Duarte move, US tax reform',
                  'Q2 FY2018 release onward; FY2018 annual from the Q4 FY2018 release.'),
        A[2020]: ('FY2020: gains on property / business sales, cross-currency swap gain, Hexcel merger costs and COVID restructuring adjusted out', None),
        q(2022, 1): ('FY2022+: business development, non-recurring matters, product rationalization, inventory / collections charges adjusted out', None),
        q(2026, 1): ('FY2026: restructuring (China OH wind-down, servo-valve transition) only', None),
    }
    DEFS['e_def'] = {
        'C': ('FY2011–17: no adjusted EPS published — adjusted = GAAP diluted EPS (flag)', None),
        A[2018]: ('FY2018+: adjusted EPS = GAAP diluted EPS + items net of tax (company definition); FY2018–19 incl. US tax-reform transition impact', None),
        q(2025, 1): ('FY2025: product rationalization gains and German corporate tax-rate change adjusted out', None),
    }
def write_defs(ws, row):
    for c, (txt, cm) in DEFS[row.key].items():
        cell = ws[f'{c}{row.r}']; cell.value = txt
        cell.font = Font(name='Calibri', sz=10, b=True, i=True, color='FFC00000')
        if cm: comment(cell, cm)

def write_scen(ws):
    srcrow = {'hdr': 9, 'out_hdr': 11, 'pt_hdr': 15, 'out': 12, 'pt': 16, 'lev': 28, 'blk': 30, 'yr': 31}
    for spec_ in SCEN_ROWS:
        r, kind = spec_[0], spec_[1]
        for ci in range(CI(SL), CI(SNOTE) + 1):
            ws.cell(r, ci)._style = copy.copy(TM.cell(srcrow[kind], ci + TSHIFT)._style)
        if kind in ('hdr', 'out_hdr', 'pt_hdr'):
            _, _, a, b, c_, d_, e = spec_
            for col, v in zip([SL, SBULL, SBASE, SBEAR, SNOTE], [a, b, c_, d_, e]):
                if v is not None: ws[f'{col}{r}'] = v
        elif kind == 'out':
            _, _, lab, hi, mid, lo, note, nf = spec_
            ws[f'{SL}{r}'] = lab
            for col, v in zip([SBULL, SBASE, SBEAR], [hi, mid, lo]):
                ws[f'{col}{r}'] = v; ws[f'{col}{r}'].number_format = NF[nf]; set_color(ws[f'{col}{r}'], BLUE)
            ws[f'{SNOTE}{r}'] = note
        elif kind == 'pt':
            _, _, lab, _, val, _, note, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBASE}{r}'] = val; ws[f'{SBASE}{r}'].number_format = NF[nf]; set_color(ws[f'{SBASE}{r}'], BLUE); ws[f'{SNOTE}{r}'] = note
        elif kind == 'lev':
            _, _, lab, a, b, c_, note = spec_
            ws[f'{SL}{r}'] = lab
            for col, v in zip([SBULL, SBASE, SBEAR], [a, b, c_]):
                ws[f'{col}{r}'] = v; ws[f'{col}{r}'].number_format = NF['x']
            ws[f'{SNOTE}{r}'] = note
        elif kind == 'blk':
            _, _, lab, db, mid, de, note, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBULL}{r}'] = db; ws[f'{SBASE}{r}'] = mid; ws[f'{SBEAR}{r}'] = de; ws[f'{SNOTE}{r}'] = note
            for col in (SBULL, SBEAR):
                ws[f'{col}{r}'].number_format = '\\+#,##0;\\-#,##0' if nf == 'm' else '\\+0.0%;\\-0.0%'
        elif kind == 'yr':
            _, _, lab, hdr, base, clamp, _, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBASE}{r}'] = base
            if clamp:
                ws[f'{SBULL}{r}'] = f'=MAX(0,{SBASE}{r}+${SBULL}${hdr})'; ws[f'{SBEAR}{r}'] = f'=MAX(0,{SBASE}{r}+${SBEAR}${hdr})'
            else:
                ws[f'{SBULL}{r}'] = f'={SBASE}{r}+${SBULL}${hdr}'; ws[f'{SBEAR}{r}'] = f'={SBASE}{r}+${SBEAR}${hdr}'
            for col in (SBULL, SBASE, SBEAR): ws[f'{col}{r}'].number_format = NF[nf]

def write_header(ws):
    for r in range(1, 8):
        clone_row_style(ws, r, r, 2, CI(LASTC))
    ws[f'{NOTE}1'] = 'SCENARIO SWITCH ▼ (Bull / Base / Bear) — drives all scenario-linked forecast drivers'
    ws[f'{NOTE}2'] = 'Base'
    comment(ws[f'{NOTE}2'], 'Scenario toggle. Choose Bull / Base / Bear. Drives: FY26 sales-growth and adjusted-EPS guidance end-points (high / mid / low), '
            'Q4/26E buybacks, 2027E–2031E primary-market sales growth, segment margins, M&A spend and the leverage target for 2027E+ buybacks '
            f'(scenario table {SL}:{SNOTE}).')
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type='list', formula1='"Bull,Base,Bear"', allow_blank=False); ws.add_data_validation(dv); dv.add(f'{NOTE}2')
    ws['B2'] = 'Woodward, Inc. (NASDAQ: WWD) — 3-Statement Financial Model'
    ws['B3'] = ('US GAAP as reported · USD millions (per-share data in USD; shares in millions) · Fiscal years ended 30 September (2026 = Oct-2025 – Sep-2026) · '
                'Source: Woodward Forms 10-K / 10-Q, Form 8-K earnings releases (Ex. 99.1) and deal / financing 8-Ks, SEC EDGAR (CIK 108312); XBRL company facts; '
                'investor-relations releases at woodward.com are the same documents as filed on EDGAR')
    ws['B4'] = ('Basis notes: as originally reported (FY2016 and FY2018 reclassifications not recast; FY2018 restated segment earnings in memo rows). Segments: Aerospace, Industrial '
                '(Energy to FY2015). Acquisitions: L\'Orange Jun-18, Safran EMA Jul-25, Valve Research Apr-26; divestitures: renewable power systems May-20, product lines FY25, '
                'pilot controls (ONTIC, pending FY27). Q4/26E calibrated to FY26 sales-growth and adjusted-EPS guidance (29-Jul-2026).')
    ws[f'{CAGR[0]}5'] = 'Forecast'
    for c in CAGR[1:]: ws[f'{c}5'] = 'Historical'
    ws['B6'] = '(USDm)'
    for col in COLS: ws[f'{col.c}6'] = col.label
    ws[f'{NOTE}6'] = 'Modelling Notes'
    for c, a, b in zip(CAGR, ['5Y CAGR', '5Y CAGR', '10Y CAGR', '14Y CAGR'], ["26-'31", "20-'25", "15-'25", "11-'25"]):
        ws[f'{c}6'] = a; ws[f'{c}7'] = b

def set_columns(ws):
    for k in list(ws.column_dimensions.keys()): del ws.column_dimensions[k]
    ws.column_dimensions['A'].width = TM.column_dimensions['A'].width
    ws.column_dimensions['B'].width = TM.column_dimensions['B'].width
    wid = {'A': 13.0, 'Q': 9.5703125, 'QE': 11.42578125, 'E26': 14.140625, 'AE': 13.0}
    for col in COLS:
        cd = ws.column_dimensions[col.c]; cd.width = wid[col.kind]
        if col.q:
            cd.outlineLevel = 1; cd.hidden = True
    for ci in range(CI(NOTE), CI(LASTC) + 1):
        ws.column_dimensions[L(ci)].width = TM.column_dimensions[L(ci + TSHIFT)].width or 13.0
    ws.sheet_format.outlineLevelCol = 1

def main():
    import openpyxl
    wb = openpyxl.load_workbook(TEMPLATE)
    if 'DCF_old' in wb.sheetnames: del wb['DCF_old']
    ws = wb['Model']
    blank_style = copy.copy(ws['A1']._style)
    for rowc in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for c in rowc:
            c.value = None; c.comment = None; c._style = copy.copy(blank_style)
    for k in list(ws.row_dimensions.keys()):
        rd = ws.row_dimensions[k]; rd.outlineLevel = 0; rd.hidden = False
    ws.data_validations.dataValidation = []
    scen_layout(V0)
    build_spec(V0)
    n1 = len(fw.SPEC)
    build_spec2(V0)
    for i, row in enumerate(fw.SPEC): row.spec2 = i >= n1
    last = number_rows(8)
    for r in range(1, last + 1): ws.row_dimensions[r].height = 15.0
    ws.row_dimensions[2].height = 21.0
    defs_setup(); comments_setup()
    write_header(ws)
    for row in fw.SPEC: write_row(ws, row)
    write_scen(ws)
    set_columns(ws)
    ws.freeze_panes = 'C8'
    ws.sheet_view.zoomScale = 70
    import other_sheets as osh
    osh.do_all(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    osh.inject_charts(OUT)
    print('rows', fw.SPEC[-1].r, 'saved', OUT)
    return wb

RECALC = os.environ.get('XLSX_RECALC', '/mnt/skills/public/xlsx/scripts/recalc.py')
def recalc(path, timeout=600):
    import subprocess
    out = subprocess.run(['python3', RECALC, path, str(timeout)], capture_output=True, text=True, cwd=os.path.dirname(RECALC))
    try:
        return json.loads(out.stdout)
    except Exception:
        return {'raw': out.stdout[-2000:], 'err': out.stderr[-2000:]}

def scenario_snapshots():
    import openpyxl
    import other_sheets as osh
    snaps = {}
    for scen in ('Bull', 'Bear'):
        p = os.path.join(BASE, f'WWD_{scen}.xlsx')
        wb = openpyxl.load_workbook(OUT); wb['Model'][f'{NOTE}2'] = scen; wb.save(p)
        res = recalc(p); print(scen, res.get('status'), res.get('total_errors'))
        v = openpyxl.load_workbook(p, data_only=True)['Bull-Base-Bear']
        snaps[scen] = {(r, c): v[f'{c}{r}'].value for r in range(9, 51) for c in 'CDEFGHIJKL' if v[f'{c}{r}'].value is not None}
        os.remove(p)
    wb = openpyxl.load_workbook(OUT)
    osh.snapshot_values(wb, snaps['Bull'], snaps['Bear'])
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    osh.inject_charts(OUT)

NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
def inject_values(src_xlsx, recalc_xlsx):
    """Store cached formula results (from a LibreOffice-recalculated copy) in the openpyxl file, so the workbook reads correctly in any
    viewer; fullCalcOnLoad stays set so Excel recalculates on open. Same approach as the LECO / ESAB finishing step."""
    import zipfile, shutil, openpyxl
    from lxml import etree
    wbv = openpyxl.load_workbook(recalc_xlsx, data_only=True)
    zin = zipfile.ZipFile(src_xlsx)
    wbxml = etree.fromstring(zin.read('xl/workbook.xml'))
    rels = etree.fromstring(zin.read('xl/_rels/workbook.xml.rels'))
    rid2t = {r.get('Id'): r.get('Target') for r in rels}
    sheets = {}
    for s in wbxml.iter('{%s}sheet' % NS):
        rid = s.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
        sheets['xl/' + rid2t[rid].lstrip('/').replace('xl/', '')] = s.get('name')
    tmp = src_xlsx + '.tmp'
    zout = zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED)
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename in sheets:
            ws = wbv[sheets[item.filename]]
            root = etree.fromstring(data)
            for c in root.iter('{%s}c' % NS):
                f = c.find('{%s}f' % NS)
                if f is None: continue
                v = ws[c.get('r')].value
                old = c.find('{%s}v' % NS)
                if old is not None: c.remove(old)
                if v is None: continue
                ve = etree.SubElement(c, '{%s}v' % NS)
                if isinstance(v, bool):
                    c.set('t', 'b'); ve.text = '1' if v else '0'
                elif isinstance(v, (int, float)):
                    if 't' in c.attrib: del c.attrib['t']
                    ve.text = repr(float(v)) if isinstance(v, float) else str(v)
                elif isinstance(v, str) and v.startswith('#'):
                    c.set('t', 'e'); ve.text = v
                else:
                    c.set('t', 'str'); ve.text = str(v)
            data = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
        zout.writestr(item, data)
    zout.close(); zin.close()
    shutil.move(tmp, src_xlsx)

def store_values():
    import shutil
    p = os.path.join(BASE, 'WWD_recalc.xlsx')
    shutil.copy(OUT, p)
    res = recalc(p); print('recalc', res.get('status'), res.get('total_errors'))
    inject_values(OUT, p)
    os.remove(p)

if __name__ == '__main__':
    main()
    if '--snap' in sys.argv: scenario_snapshots()
    if '--recalc' in sys.argv: store_values()
