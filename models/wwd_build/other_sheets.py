"""Rewire the Bull-Base-Bear, DCF and Charts tabs of the MLM template (current layout: DCF with live valuation date, fade period,
implied ROIC, target-IRR entry price, reverse DCF and Mauboussin EV/NOPAT; Bull-Base-Bear with buyback rows) to the WWD Model rows /
columns; keep the template chart XML. The template's DCF_old tab is not carried.

Column mapping: the template's 2025A base column becomes WWD FY2026E (fiscal year ended 30-Sep-2026, a week before the valuation date;
Q1-Q3 actual, Q4 calibrated to guidance) and the template's 2026E-2030E become WWD FY2027E-FY2031E, so the DCF's explicit years start
in FY2027E and net debt is taken at 30-Sep-2026."""
import re, zipfile, shutil
from openpyxl.utils import get_column_letter as L
from fw import R, TEMPLATE, ANN, NOTE

# template Model row -> WWD row key
MAP = {158: 'is_rev', 164: 'r_cash', 166: 'is_ebit', 182: 'is_dda', 186: 'sh_dil', 211: 'r_adj', 212: 'r_adj_m', 219: 'r_ebit', 221: 'r_adj',
       229: 'e_adjni', 231: 'e_eps', 238: 'gm_aero', 245: 'gm_ebitda_m', 247: 'gm_op_m', 260: 'cf_wc', 262: 'cf_capex', 263: 'cf_acq',
       280: 'fcf', 282: 'fcf_m', 290: 'bb_px', 291: 'bb_cash', 292: 'bb_sh', 331: 'bs_nci', 408: 'ra_ic', 411: 'ra_roic', 422: 'rn_ronta',
       428: 'lv_nd', 433: 'v_px', 434: 'v_sh'}
CMAP = {'BV': ANN[2026].c, 'CA': ANN[2027].c, 'CB': ANN[2028].c, 'CC': ANN[2029].c, 'CD': ANN[2030].c, 'CE': ANN[2031].c, 'CF': NOTE}
REF = re.compile(r"Model!(\$?)([A-Z]{1,3})(\$?)(\d+)")
HIST = [ANN[y].c for y in range(2011, 2026)]

def remap(f):
    def rep(m):
        c, r = m.group(2), int(m.group(4))
        c2 = CMAP.get(c, c)
        r2 = R[MAP[r]] if r in MAP else r
        if c == 'CF' and r == 2: r2 = 2
        return f'Model!{m.group(1)}{c2}{m.group(3)}{r2}'
    return REF.sub(rep, f)

def hist_stat(f):
    """LT historical AVERAGE / MIN / MAX over the template's annual columns -> WWD FY2011-FY2025 annual columns."""
    m = re.match(r'=IFERROR\((AVERAGE|MIN|MAX)\(Model!C(\d+),', f)
    if not m: return None
    fn, r = m.group(1), int(m.group(2))
    rr = R[MAP[r]]
    return f'=IFERROR({fn}({",".join(f"Model!{c}{rr}" for c in HIST)}),"")'

def fy_dates(f):
    f = re.sub(r'DATE\((\d{4}),12,31\)', lambda m: f'DATE({int(m.group(1)) + 1},9,30)', f)
    f = re.sub(r'DATE\((\d{4})-1,12,31\)', lambda m: f'DATE({int(m.group(1)) + 1}-1,9,30)', f)
    return f

TEXT = {
    'Bull-Base-Bear': {
        'B2': 'Woodward, Inc. (WWD) — Bull / Base / Bear Case P&L Summary',
        'B3': ('Adjusted figures (company definitions: adjusted EBIT, adjusted EBITDA, adjusted EPS) · USD millions · fiscal years ended 30 September · Live panel linked to Model — '
               'toggle scenario via Model!AV2. Bull and Bear panels are value snapshots generated from the Model with the switch set to each case. FY2026E = Q1–Q3 actual + Q4 '
               'calibrated to the FY26 guidance (sales +20–23%, adjusted EPS $9.30–9.50).'),
        'G26': 32, 'H26': 32, 'S26': 36, 'T26': 36, 'AD26': 25, 'AE26': 25, 'AF26': 25,
    },
    'DCF': {
        'B2': 'DCF VALUATION — WOODWARD, INC. (WWD)',
        'C5': 0.0425, 'D5': 'Assumption (~US 10Y Treasury yield, Oct-2026; government rate sources not reachable from the build environment); update to market',
        'D6': 'Assumption: US implied ERP ~4.5–5.5%',
        'C7': 1.05, 'D7': 'Assumption: aerospace & defense components peers (HEI, TDG, CW, MOG, PH aerospace) ~0.9–1.2; ~30% of sales Industrial (cyclical power generation / transportation / oil & gas)',
        'C9': 0.055, 'D9': 'Aug-2026 Series U / V / W private-placement notes at 5.34% / 5.39% / 5.64% (8-K 21-Aug-2026); revolver at SOFR + 0.875–1.75% (4.73% at 30-Jun-2026)',
        'C10': 0.225, 'D10': 'Model FY2027E+ effective tax rate (FY26 adjusted-rate guidance ~22.5%)',
        'C12': 0.06, 'D12': 'Net debt ~$0.9bn at 30-Jun-2026 vs ~$23bn market cap ($382.17 × ~61m diluted shares); ~1.0x adjusted EBITDA',
        'D15': 'Periods = (30-Sep of fiscal year − today) / 365.25; cash flows assumed at fiscal year-end (30 September); FY2026E (ended 30-Sep-2026) is the base year',
        'B18': 'Adj. EBIT (full-cost) less acquisition, integration & restructuring costs',
        'B21': 'plus D&A (depreciation + amortization of intangibles)',
        'B22': '− Capex (PP&E) and acquisitions (bolt-ons; M&A lever 2027E–31E)',
        'B26': 'Fade-period growth (2032E–36E; 7.0% → 5.0% linear)', 'I26': 0.07, 'M26': 0.05, 'C27': 0.04,
        'D32': ("Woodward's adjusted EBIT already deducts amortization of intangibles and share-based compensation (neither is adjusted out); D&A is added back in row 21. "
                'Restructuring and acquisition / business-development costs (cash; adjusted out by the company, incl. the China OH, servo-valve and Santa Clarita actions) are deducted '
                'inside row 18. Row 22 includes bolt-on acquisitions in FY2027E–31E; the fade period and terminal value assume no new acquisitions. The pending $180m pilot-controls '
                'sale (FY27) is excluded from UFCF (the business is kept in the build).'),
        'D43': 'Opening IC = Model FY2026E book IC (equity + net debt at 30-Sep-2026), rolled forward with capex + M&A − D&A + ΔWC.',
        'B48': 'Less: Net debt + NCI (Model FY2026E year end, 30-Sep-2026)',
        'D48': 'WWD: no noncontrolling interests; net debt at the base-year end (≈ valuation date) so that FY2027E+ cash flows are not double counted.',
        'B50': 'Diluted shares (m, Model FY2026E period end)',
        'D76': 'Rationale (Woodward characteristics)',
        'C78': 0.10, 'D78': 'Model adjusted EBIT CAGR FY2026–31E: high-single-digit sales growth (commercial OEM rate ramps, aftermarket, power generation) with margin expansion → ~10% NOPAT growth',
        'C79': 0.20, 'D79': 'Sole-source, certified content on long-lived aircraft and engine platforms drives 20–30-year aftermarket annuities at high incremental returns; Industrial returns are cyclical; L\'Orange (2018) earned below WACC for years',
        'C80': 15, 'D80': 'Certification barriers and installed-base aftermarket (commercial services ~25% of sales) sustain a long competitive-advantage period; exposure to aircraft production rates and Industrial cyclicality',
        'B90': 'VALUATION  —  Applied to FY2029E NOPAT',
        'B91': 'FY2029E Adj. EBIT less acquisition, integration & restructuring costs',
        'B93': 'FY2029E NOPAT',
        'B95': 'Implied Enterprise Value (FY2029E basis)',
        'B96': 'Less: FY2028E year-end net debt + NCI (Model)',
        'B98': 'Diluted shares FY2029E (m, Model)',
        'B99': 'Implied share price FY2029 ($)',
        'B106': 'Plus: FY2026E year-end net debt + NCI',
        'B108': 'WWD market-implied EV/NOPAT (on FY2029E NOPAT)',
    },
    'Charts': {
        'B2': 'Woodward, Inc. — Operating Performance (FY2012–FY2025) and Model Forecast (FY2026E–FY2031E)',
        'B3': ('USDm · fiscal years ended 30 September · Linked to Model. Bars = contribution of each segment to total net-sales growth (pp). Adjusted EBITDA and EBIT = company '
               'definitions (adjusted measures published from FY2018; earlier years = EBITDA / EBIT). ROIC / RONTA on adjusted EBIT (full-cost).'),
        'B8': 'Aerospace contribution to growth (pp)', 'B9': 'Industrial contribution to growth (pp)', 'B12': 'Adjusted EBITDA', 'B13': 'EBIT (GAAP)',
        'B14': 'Adj. EBITDA Margin %', 'B15': 'EBIT Margin %',
    },
}
BBB_ROWS = {9: 'Net sales', 10: '   Net sales y/y %', 11: '   Aerospace segment sales y/y %', 12: 'Adjusted EBITDA',
            15: 'EBIT (GAAP)', 16: '   EBIT y/y %', 17: '   EBIT margin (GAAP) %', 19: 'Adjusted Net Earnings',
            35: 'Adjusted EBITDA (leverage basis)', 36: '   Adjusted EBITDA % of net sales'}
CHART_TITLES = {
    'Aggregates Revenue Drivers: Shipment Volume vs. Price (ASP), y/y %': 'Net Sales Growth: Aerospace vs. Industrial Contribution, pp',
    'Adjusted EBITDA &amp; Earnings from Operations (USDm)': 'Adjusted EBITDA &amp; EBIT (USDm)',
}

def build_charts(ws):
    for i, y in enumerate(range(2012, 2032)):
        c = L(3 + i); mc = ANN[y].c
        ws[f'{c}7'] = y if y <= 2025 else f'{y}E'
        ws[f'{c}8'] = f'=IFERROR(Model!{mc}{R["gc_aero"]}+0,0)'
        ws[f'{c}9'] = f'=IFERROR(Model!{mc}{R["gc_ind"]}+0,0)'
        ws[f'{c}11'] = f'={c}7'
        ws[f'{c}12'] = f'=Model!{mc}{R["r_adj"]}'
        ws[f'{c}13'] = f'=Model!{mc}{R["is_ebit"]}'
        ws[f'{c}14'] = f'=Model!{mc}{R["gm_ebitda_m"]}'
        ws[f'{c}15'] = f'=Model!{mc}{R["gm_op_m"]}'
        ws[f'{c}17'] = f'={c}7'
        ws[f'{c}18'] = f'=IFERROR(Model!{mc}{R["ra_roic"]}+0,NA())'
        ws[f'{c}19'] = f'=IFERROR(Model!{mc}{R["rn_ronta"]}+0,NA())'
        ws[f'{c}20'] = f'=IFERROR(Model!{mc}{R["gm_ebitda_m"]}+0,NA())'
        ws[f'{c}21'] = f'=IFERROR(Model!{mc}{R["fcf_m"]}+0,NA())'

def do_all(wb):
    for name in ('Bull-Base-Bear', 'DCF', 'Charts'):
        ws = wb[name]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and 'Model!' in c.value:
                    h = hist_stat(c.value) if name == 'Bull-Base-Bear' else None
                    c.value = h or remap(c.value)
                if isinstance(c.value, str) and 'DATE(' in c.value:
                    c.value = fy_dates(c.value)
        for addr, t in TEXT.get(name, {}).items():
            ws[addr] = t
    b = wb['Bull-Base-Bear']
    for r, t in BBB_ROWS.items():
        for c in ('B', 'N', 'Z'): b[f'{c}{r}'] = t
    for c0 in ('C', 'O', 'AA'):
        i = b[f'{c0}7'].column
        b.cell(7, i).value = '2026E'
        for k in range(1, 6): b.cell(7, i + k).value = f'{2026 + k}E'
        b.cell(7, i + 6).value = "26-'31 CAGR"
    from openpyxl.comments import Comment
    cm = Comment('P/E on adjusted EPS (company definition). WWD at $382.17 trades at ~41x FY2026E adjusted EPS ($9.40) and ~36x FY2027E. Exit multiples are assumptions: '
                 'Base 32x, Bull 36x, Bear 25x applied to FY2030E–31E (FY2029E–31E for Bear), converging toward aerospace-components peers as growth normalises.', 'WWD model')
    cm.width, cm.height = 420, 120; b['N26'].comment = cm
    cm2 = Comment(f'Snapshot (values) of Model rows {R["bb_cash"]}, {R["bb_px"]} and {R["bb_sh"]} (Share Buyback Schedule) with Model!{NOTE}2 = Bull / Bear. '
                  'Base panel (D40:H42) is live-linked.', 'WWD model')
    cm2.width, cm2.height = 420, 90; b['P40'].comment = cm2
    b['J2'] = 'Share price ($):'
    for c in ('B', 'N', 'Z'): b[f'{c}29'] = 'IRR (to 30-Sep of Y-1)'
    for c in ('B', 'N', 'Z'):
        if b[f'{c}40'].value: b[f'{c}40'] = '   Share repurchases ($m; FY2026E = Q4)'
    d = wb['DCF']
    for i, col in enumerate('DEFGHIJKLM'):
        d[f'{col}17'] = f'{2027 + i}E'
    a26, a27 = ANN[2026].c, ANN[2027].c
    d['D22'] = f'=Model!{a27}{R["cf_capex"]}+Model!{a27}{R["cf_acq"]}'
    d['D34'] = f'=Model!{a26}{R["ra_ic"]}'
    d['C48'] = f'=Model!{a26}{R["lv_nd"]}+Model!{a26}{R["bs_nci"]}'
    d['C50'] = f'=Model!{a26}{R["v_sh"]}'
    d['C52'] = f'=Model!{a26}{R["v_px"]}'
    d['P17'] = 'CAGR 27–31E'; d['Q17'] = 'CAGR 27–36E'
    build_charts(wb['Charts'])

def snapshot_values(wb, bull, bear):
    ws = wb['Bull-Base-Bear']
    cols = 'CDEFGHIJKL'
    for snap, off in ((bull, 'OPQRSTUVWX'), (bear, ['AA', 'AB', 'AC', 'AD', 'AE', 'AF', 'AG', 'AH', 'AI', 'AJ'])):
        for (r, c), v in snap.items():
            tgt = f'{off[cols.index(c)]}{r}'
            cur = ws[tgt].value
            if isinstance(cur, str) and cur.startswith('='): continue
            if r in (7, 26): continue
            ws[tgt] = v

def inject_charts(path):
    """Replace openpyxl-rewritten chart parts with the template chart XML (combo charts, colours, layout), retitled, ranges kept
    (Charts!C:V = FY2012-FY2031E) and cached values removed so Excel redraws from the WWD cells."""
    tz = zipfile.ZipFile(TEMPLATE)
    tmpl = {n: tz.read(n) for n in tz.namelist() if n.startswith('xl/charts/chart')}
    tmp = path + '.tmp'
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in tmpl:
                t = tmpl[item.filename].decode('utf-8')
                for a, b2 in CHART_TITLES.items(): t = t.replace(a, b2)
                t = re.sub(r'<c:numCache>.*?</c:numCache>', '', t, flags=re.S)
                t = re.sub(r'<c:strCache>.*?</c:strCache>', '', t, flags=re.S)
                data = t.encode('utf-8')
            zout.writestr(item, data)
    shutil.move(tmp, path)
