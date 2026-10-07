"""QA a recalculated model built on the MLM template: list formula errors on every sheet and every check row that is not zero.

Usage: python3 inspect_model.py <recalculated.xlsx> [--tol 0.001]

Check rows are Model-sheet rows whose label (column B) contains "Check" or "tie-out". Only the period columns are scanned:
column C up to the column before the "Modelling Notes" header in row 6. The DuPont "Margin × Turnover × Tax burden = ROIC" row
shows ROIC by design and is skipped. Recalculate first (LibreOffice via the xlsx skill's
recalc.py, or open and save in Excel); a workbook without cached values shows no results.
"""
import sys
from collections import defaultdict

import openpyxl
from openpyxl.utils import get_column_letter as L


def period_columns(ws):
    last = ws.max_column
    for c in range(3, ws.max_column + 1):
        v = ws.cell(6, c).value
        if isinstance(v, str) and 'notes' in v.lower():
            last = c - 1
            break
    return range(3, last + 1)


def main(path, tol):
    wb = openpyxl.load_workbook(path, data_only=True)
    errs = defaultdict(list)
    for name in wb.sheetnames:
        for row in wb[name].iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith('#'):
                    errs[name].append(c.coordinate)
    for k, v in errs.items():
        print('ERR', k, len(v), v[:40])

    ws = wb['Model']
    hdr = {c: ws.cell(6, c).value for c in range(1, ws.max_column + 1)}
    cols = period_columns(ws)
    n_bad = 0
    for r in range(8, ws.max_row + 1):
        lab = ws.cell(r, 2).value
        if not (isinstance(lab, str) and ('Check' in lab or 'tie-out' in lab)):
            continue
        if 'Margin × Turnover' in lab:      # DuPont row shows ROIC by design, not a zero check
            continue
        bad = []
        for c in cols:
            v = ws.cell(r, c).value
            if isinstance(v, (int, float)) and not isinstance(v, bool) and abs(v) > tol:
                bad.append(f'{hdr.get(c) or L(c)}:{round(v, 3)}')
        if bad:
            n_bad += 1
            print(r, lab.strip()[:70], len(bad), bad[:25])
    print(f'{sum(len(v) for v in errs.values())} error cells; {n_bad} check rows not zero')
    return 1 if errs or n_bad else 0


if __name__ == '__main__':
    args = sys.argv[1:]
    tol = 0.001
    if '--tol' in args:
        i = args.index('--tol')
        tol = float(args[i + 1])
        del args[i:i + 2]
    if len(args) != 1:
        sys.exit(__doc__)
    sys.exit(main(args[0], tol))
