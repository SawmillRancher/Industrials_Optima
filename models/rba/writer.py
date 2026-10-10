"""Writes the Model tab from Row definitions (layout & styling mirror the HII model)."""
from __future__ import annotations

from openpyxl.styles import Alignment
from openpyxl.worksheet.datavalidation import DataValidation

from .framework import (FMT_NUM, FMT_NUM1, BLACK, BLUE, CREAM, DARK, FMT_CAGR, GREY, LIGHT, MID, RED, SAGE, WHITE, YELLOW,
                        Col, Cols, Ctx, Row, THIN_TOP, fill, font, put_comment)

FIRST_ROW = 8
HIST_GREY = "A6A6A6"


class Data:
    """period -> {"is.total_rev": v, ...} with source notes."""

    def __init__(self, periods: dict):
        self.p = periods

    def get(self, period, key):
        v = self.p.get(period, {}).get(key)
        return v

    def has(self, period, key):
        return self.get(period, key) is not None


def assign_rows(rows):
    m = {}
    for i, r in enumerate(rows):
        r.n = FIRST_ROW + i
        if r.key:
            if r.key in m:
                raise ValueError(f"duplicate row key {r.key}")
            m[r.key] = r.n
    return m


def _agg_formula(cols: Cols, c: Col, row: Row, rn: int, data: Data):
    """1H or FY aggregation formula from quarter cells."""
    if c.kind == "H":
        q1, q2 = cols.quarter(c.year, 1), cols.quarter(c.year, 2)
        lenient = (row.src.split(".")[0] in ("adjE", "adjN", "is") and "published" not in row.src
                   and row.agg == "sum")
        if lenient and row.agg == "sum" and (data.has(q1.key, row.src) or data.has(q2.key, row.src)):
            if data.has(q1.key, row.src) and data.has(q2.key, row.src):
                return f"={q1.L}{rn}+{q2.L}{rn}"
            return f"=N({q1.L}{rn})+N({q2.L}{rn})"
        if not (data.has(q1.key, row.src) and data.has(q2.key, row.src)):
            return None
        if row.agg == "sum":
            return f"={q1.L}{rn}+{q2.L}{rn}"
        if row.agg == "last":
            return f"={q2.L}{rn}"
        if row.agg == "avg":
            return f"=AVERAGE({q1.L}{rn},{q2.L}{rn})"
        return None
    if c.kind == "Y":
        qs = [cols.quarter(c.year, i) for i in range(1, 5)]
        if not all(data.has(q.key, row.src) for q in qs):
            return None
        if row.agg == "sum":
            h = cols.half(c.year)
            return f"={h.L}{rn}+{qs[2].L}{rn}+{qs[3].L}{rn}"
        if row.agg == "last":
            return f"={qs[3].L}{rn}"
        if row.agg == "avg":
            return f"=AVERAGE({','.join(q.L + str(rn) for q in qs)})"
    return None


def _default_e26(cols: Cols, c: Col, row: Row, rn: int):
    q1, q2 = cols.quarter(c.year, 1), cols.quarter(c.year, 2)
    q3, q4 = cols.quarter(c.year, 3), cols.quarter(c.year, 4)
    if row.agg == "sum":
        return f"=SUM({q1.L}{rn},{q2.L}{rn},{q3.L}{rn},{q4.L}{rn})"
    if row.agg == "last":
        return f"={q4.L}{rn}"
    if row.agg == "avg":
        return f"=AVERAGE({q1.L}{rn},{q2.L}{rn},{q3.L}{rn},{q4.L}{rn})"
    return None


def cell_content(cols: Cols, rowmap, row: Row, c: Col, data: Data):
    """Return (value, kind) where kind in {'data','formula','input',None}."""
    rn = row.n
    ctx = Ctx(cols, rowmap, c)
    is_h26 = c.kind == "H" and c.year == 2026
    if row.annual_only and c.kind in ("Q", "QE"):
        return None, None
    if row.annual_only and c.kind == "H" and not is_h26:
        return None, None
    if not row.quarters_ok and c.kind in ("Q", "H", "QE"):
        return None, None

    if c.kind in ("A", "Q", "H", "Y"):
        if row.hist is not None:
            v = row.hist(ctx)
            return (v, "formula") if v is not None else (None, None)
        if row.hist_q is not None and c.kind in ("Q", "H"):
            v = row.hist_q(ctx)
            return (v, "formula") if v is not None else (None, None)
        if c.kind == "Y" and row.fy_formula is not None:
            v = row.fy_formula(ctx)
            if v is not None:
                return v, "formula"
        if row.src:
            if c.kind == "H" and not (row.annual_only and is_h26):
                if row.h1_override is not None:
                    v = row.h1_override(ctx)
                    return (v, "formula") if v is not None else (None, None)
                v = data.get(c.key, row.src)
                f = _agg_formula(cols, c, row, rn, data)
                if f:
                    return f, "formula"
                return (v, "data") if v is not None else (None, None)
            v = data.get(c.key, row.src)
            if v is not None:
                return v, "data"
            if c.kind == "Y":
                f = _agg_formula(cols, c, row, rn, data)
                if f:
                    return f, "formula"
        return None, None

    # forecast
    fn = None
    if c.kind == "QE":
        fn = row.qe or row.fc
    elif c.kind == "E26":
        fn = row.e26
        flow = (row.agg == "sum" and not row.annual_only and row.style not in ("pct", "growth", "check")
                and row.fmt in (FMT_NUM, FMT_NUM1) and (row.qe is not None or row.fc is not None))
        if fn is None and flow:
            v = _default_e26(cols, c, row, rn)
            return (v, "formula") if v else (None, None)
        fn = fn if fn is not None else row.fc
        if fn is None and (row.qe is not None) and not row.annual_only:
            v = _default_e26(cols, c, row, rn)
            return (v, "formula") if v else (None, None)
    elif c.kind == "E":
        fn = row.e or row.fc
    if fn is None:
        return None, None
    v = fn(ctx) if callable(fn) else fn
    if v is None:
        return None, None
    if isinstance(v, str) and v.startswith("="):
        return v, "formula"
    return v, "input"


def style_label(cell, row: Row):
    st = row.style
    if st == "section":
        cell.font = font(12, True, color=WHITE)
        cell.fill = fill(DARK)
    elif st == "block":
        cell.font = font(11, True, color=WHITE)
        cell.fill = fill(MID)
    elif st == "total":
        cell.font = font(11, True, color=WHITE)
        cell.fill = fill(DARK)
        cell.border = THIN_TOP
    elif st == "sub":
        cell.font = font(12, True)
        cell.border = THIN_TOP
    elif st in ("pct", "growth"):
        cell.font = font(12, color="404040")
    elif st == "check":
        cell.font = font(12, True, color="595959")
    elif st == "memo":
        cell.font = font(12, True, True, color="666666")
    elif st == "sched":
        cell.font = font(12, color="404040")
        cell.fill = fill(GREY)
    elif st == "deftext":
        cell.font = font(11, True, True, color=RED)
    elif st == "note":
        cell.font = font(10, False, True, color=RED)
    else:
        cell.font = font(12)


def style_value(cell, row: Row, c: Col, kind):
    st = row.style
    bold = st in ("total", "sub")
    if kind in ("data", "input"):
        color = BLUE
        bold = bold or c.is_forecast or kind == "input"
    else:
        color = BLACK
        if st == "growth" and not c.is_forecast:
            color = HIST_GREY
    cell.font = font(12, bold, color=color)
    cell.number_format = row.fmt
    cell.alignment = Alignment(horizontal="right")
    if st in ("total", "sub"):
        cell.border = THIN_TOP
    if st == "total" and (c.is_forecast or c.kind in ("A", "Y")):
        cell.fill = fill(LIGHT)
    elif c.is_forecast:
        cell.fill = fill(CREAM)


def write_header(ws, cols: Cols, title, sub, basis):
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 60
    ws["B2"] = title
    ws["B2"].font = font(16, True, color="1F3864")
    ws["B3"] = sub
    ws["B3"].font = font(10, color="595959")
    ws["B4"] = basis
    ws["B4"].font = font(10, color=RED)
    ws.row_dimensions[2].height = 21
    N = cols.NOTES
    ws[f"{N}1"] = "SCENARIO SWITCH ▼ (Bull / Base / Bear) — drives all scenario-linked forecast drivers"
    ws[f"{N}1"].font = font(10, True, color=RED)
    ws[f"{N}1"].fill = fill(CREAM)
    ws[f"{N}2"] = "Base"
    from openpyxl.styles import Border, Side
    ws[f"{N}2"].font = font(12, True, color=BLUE)
    ws[f"{N}2"].fill = fill(YELLOW)
    ws[f"{N}2"].alignment = Alignment(horizontal="center")
    ws[f"{N}2"].border = Border(top=Side("medium"), bottom=Side("thin"), left=Side("medium"), right=Side("medium"))
    dv = DataValidation(type="list", formula1='"Bull,Base,Bear"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"{N}2")
    put_comment(ws[f"{N}2"], "Scenario toggle. Choose Bull / Base / Bear. Drives: FY26 outlook end-points (GTV growth, "
                "adjusted EBITDA), 2027E–30E sector GTV growth, adjusted EBITDA margin and share repurchases "
                "(scenario input table to the right). Bull / Bear = Base + Δ.")
    # row 6 headers
    hdr_font, hdr_fill = font(12, True, color=WHITE), fill(DARK)
    ws["B6"] = "(USDm)"
    ws["B6"].font, ws["B6"].fill = hdr_font, hdr_fill
    ws["B6"].alignment = Alignment(horizontal="left")
    for c in cols.list:
        cell = ws[f"{c.L}6"]
        cell.value = c.label
        cell.font, cell.fill = hdr_font, hdr_fill
        cell.alignment = Alignment(horizontal="center")
        if c.is_forecast:
            ws[f"{c.L}5"].fill = fill(CREAM)
    ws[f"{N}6"] = "Modelling Notes"
    ws[f"{N}6"].font, ws[f"{N}6"].fill = hdr_font, hdr_fill
    ws[f"{N}6"].alignment = Alignment(horizontal="center")
    labels = [("Forecast", "5Y CAGR", "25-'30"), ("Historical", "5Y CAGR", "20-'25"),
              ("Historical", "10Y CAGR", "15-'25"), ("Historical", "19Y CAGR", "06-'25")]
    for L, (a, b, cc) in zip(cols.CAGR, labels):
        for r, v in ((5, a), (6, b)):
            ws[f"{L}{r}"] = v
            ws[f"{L}{r}"].font = font(12, True, color=WHITE)
            ws[f"{L}{r}"].fill = fill(SAGE)
            ws[f"{L}{r}"].alignment = Alignment(horizontal="center")
        ws[f"{L}7"] = cc
        ws[f"{L}7"].font = font(9, color="666666")
        ws[f"{L}7"].fill = fill(CREAM)
        ws[f"{L}7"].alignment = Alignment(horizontal="center")


def cagr_formulas(cols: Cols, row: Row):
    rn = row.n
    a = cols.annual
    if row.cagr == "cagr":
        spec = [(2025, 2030, 5), (2020, 2025, 5), (2015, 2025, 10), (2006, 2025, 19)]
        return [f'=IFERROR(({a(e).L}{rn}/{a(s).L}{rn})^(1/{n})-1,"n/m")' for s, e, n in spec]
    if row.cagr == "avg":
        spec = [(2026, 2030), (2021, 2025), (2016, 2025), (2006, 2025)]
        out = []
        for s, e in spec:
            refs = ",".join(f"{a(y).L}{rn}" for y in range(s, e + 1))
            out.append(f'=IFERROR(AVERAGE({refs}),"n/m")')
        return out
    return None


def write_model(ws, cols: Cols, rows, data: Data, title, sub, basis):
    rowmap = assign_rows(rows)
    write_header(ws, cols, title, sub, basis)
    last_idx = cols.last.idx
    from openpyxl.utils import column_index_from_string
    notes_idx = column_index_from_string(cols.NOTES)
    for row in rows:
        rn = row.n
        lab = ws.cell(rn, 2)
        lab.value = row.label if row.label else None
        style_label(lab, row)
        put_comment(lab, row.comment)
        if row.style in ("section", "block"):
            band = fill(DARK if row.style == "section" else MID)
            fnt = font(12 if row.style == "section" else 11, True, color=WHITE)
            for ci in range(3, notes_idx + 1):
                cell = ws.cell(rn, ci)
                cell.fill, cell.font = band, fnt
            if row.note:
                ws[f"{cols.NOTES}{rn}"] = row.note
                ws[f"{cols.NOTES}{rn}"].font = font(12, True, color=WHITE)
            continue
        if row.style == "deftext":
            for c in cols.list:
                cell = ws.cell(rn, c.idx)
                if c.is_forecast:
                    cell.fill = fill(CREAM)
                t = (row.texts or {}).get(c.key)
                if t:
                    head, _, body = t.partition("||")
                    cell.value = head
                    cell.font = font(10, True, True, color=RED)
                    put_comment(cell, body or head)
            ws[f"{cols.NOTES}{rn}"].fill = fill(CREAM)
            ws[f"{cols.NOTES}{rn}"].value = row.note or ("Definition changes are flagged in the column where they take "
                                                          "effect (cell comment gives the detail).")
            ws[f"{cols.NOTES}{rn}"].font = font(10)
            for L in cols.CAGR:
                ws[f"{L}{rn}"].fill = fill(CREAM)
            continue
        if row.style in ("blank", "note") or not row.key:
            for c in cols.list:
                if c.is_forecast:
                    ws.cell(rn, c.idx).fill = fill(CREAM)
            nc = ws[f"{cols.NOTES}{rn}"]
            nc.fill = fill(CREAM)
            if row.note:
                nc.value = row.note
                nc.font = font(10)
            continue
        for c in cols.list:
            v, kind = cell_content(cols, rowmap, row, c, data)
            cell = ws.cell(rn, c.idx)
            if v is not None:
                cell.value = v
            style_value(cell, row, c, kind)
        nc = ws[f"{cols.NOTES}{rn}"]
        nc.fill = fill(CREAM)
        nc.font = font(10)
        if row.note:
            nc.value = row.note
        cf = cagr_formulas(cols, row)
        for i, L in enumerate(cols.CAGR):
            cell = ws[f"{L}{rn}"]
            if cf:
                cell.value = cf[i]
                cell.font = font(12, True)
                cell.fill = fill(LIGHT)
                cell.number_format = FMT_CAGR if row.cagr == "cagr" else row.fmt
            else:
                cell.fill = fill(CREAM)
    return rowmap


def set_widths(ws, cols: Cols):
    ws.column_dimensions["A"].width = 2.5
    ws.column_dimensions["B"].width = 62
    for c in cols.list:
        cd = ws.column_dimensions[c.L]
        if c.kind in ("A", "Y", "E26", "E"):
            cd.width = 14.13
        elif c.kind == "H":
            cd.width = 10.88
        else:
            cd.width = 9.63 if c.kind == "Q" else 11.5
        if c.hidden:
            cd.hidden = True
            cd.outlineLevel = 1
        elif c.kind in ("Q", "H", "QE"):
            cd.outlineLevel = 1
    ws.column_dimensions[cols.NOTES].width = 86.88
    ws.column_dimensions[cols.GAP1].width = 2
    for L, w in zip(cols.CAGR, (14.13, 12.38, 13.5, 13.5)):
        ws.column_dimensions[L].width = w
    ws.column_dimensions[cols.GAP2].width = 2
    ws.column_dimensions[cols.SPACER].width = 6
    ws.column_dimensions[cols.SCN_LABEL].width = 40
    for L in cols.SCN.values():
        ws.column_dimensions[L].width = 11.63
    ws.column_dimensions[cols.SCN_NOTE].width = 40
    ws.freeze_panes = "C7"
    ws.sheet_properties.outlinePr.summaryRight = True
