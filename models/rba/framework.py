"""Sheet-building framework for the RB Global model (mirrors the HII model layout / styling).

The Model tab is described as an ordered list of Row objects. Each row knows how to fill
every period column: historical hard-codes (blue) from the extracted SEC data, formulas
(black) for derived lines, and forecast formulas / inputs for Q3/26E–Q4/26E, 2026E and
2027E–2030E.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from typing import Callable, Optional

from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------------------
# Palette / formats (taken from the HII model)
# --------------------------------------------------------------------------------------
DARK = "2F3E46"
MID = "52796F"
SAGE = "84A98C"
LIGHT = "CAD2C5"
CREAM = "FFFDE7"
GREY = "E9ECEC"
YELLOW = "FFFF00"
BLUE = "0000FF"
BLACK = "000000"
GREEN = "008000"
RED = "C00000"
NAVY = "1F3864"
WHITE = "FFFFFF"

FMT_NUM = '#,##0;\\(#,##0\\);\\-'
FMT_NUM1 = '#,##0.0;\\(#,##0.0\\);\\-'
FMT_PCT = '0.0%;\\(0.0%\\);\\-'
FMT_PCT0 = '0%;\\(0%\\);\\-'
FMT_EPS = '0.00;\\(0.00\\);\\-'
FMT_SHR = '#,##0.0;\\(#,##0.0\\);\\-'
FMT_X = '0.00\\x;\\(0.00"x)";\\-'
FMT_DAYS = '0\\d;\\(0")d";\\-'
FMT_CAGR = '0%;\\(0%\\);"n/m"'
FMT_BPS = '#,##0"bps";\\(#,##0"bps"\\);\\-'

FONT = "Calibri"


def font(sz=12, b=False, i=False, color=None):
    return Font(name=FONT, sz=sz, b=b, i=i, color=color)


def fill(rgb):
    return PatternFill("solid", fgColor=rgb)


THIN_TOP = Border(top=Side(style="thin"))

# --------------------------------------------------------------------------------------
# Columns
# --------------------------------------------------------------------------------------
ANNUAL_ONLY_YEARS = list(range(2006, 2013))   # 2006–2012: annual only
QUARTER_YEARS = list(range(2013, 2026))       # 2013–2025: Q1, Q2, 1H, Q3, Q4, FY
FC_YEAR = 2026                                # Q1/26, Q2/26, 1H/26, Q3/26E, Q4/26E, 2026E
OUT_YEARS = [2027, 2028, 2029, 2030]


@dataclass
class Col:
    key: str          # data period key: "2006", "Q1/13", "1H/13", "2013", "Q3/26E", "2026E", "2027E"
    label: object     # header label
    kind: str         # A (annual-only hist) | Q (hist quarter) | H (hist 1H) | Y (hist FY, quarter era)
                      # QE (forecast quarter) | E26 (2026E) | E (2027E+)
    year: int
    q: Optional[int] = None
    idx: int = 0
    hidden: bool = False

    @property
    def L(self):
        return get_column_letter(self.idx)

    @property
    def is_forecast(self):
        return self.kind in ("QE", "E26", "E")

    @property
    def is_annual(self):
        return self.kind in ("A", "Y", "E26", "E")


def build_columns(first_idx=3):
    cols = []
    for y in ANNUAL_ONLY_YEARS:
        cols.append(Col(str(y), y, "A", y))
    for y in QUARTER_YEARS + [FC_YEAR]:
        yy = f"{y % 100:02d}"
        fc = y == FC_YEAR
        cols.append(Col(f"Q1/{yy}", f"Q1/{yy}", "Q", y, 1, hidden=not fc))
        cols.append(Col(f"Q2/{yy}", f"Q2/{yy}", "Q", y, 2, hidden=not fc))
        cols.append(Col(f"1H/{yy}", f"1H/{yy}", "H", y, None, hidden=not fc))
        if fc:
            cols.append(Col(f"Q3/{yy}E", f"Q3/{yy}E", "QE", y, 3))
            cols.append(Col(f"Q4/{yy}E", f"Q4/{yy}E", "QE", y, 4))
            cols.append(Col(f"{y}E", f"{y}E", "E26", y))
        else:
            cols.append(Col(f"Q3/{yy}", f"Q3/{yy}", "Q", y, 3, hidden=True))
            cols.append(Col(f"Q4/{yy}", f"Q4/{yy}", "Q", y, 4, hidden=True))
            cols.append(Col(str(y), y, "Y", y))
    for y in OUT_YEARS:
        cols.append(Col(f"{y}E", f"{y}E", "E", y))
    for i, c in enumerate(cols):
        c.idx = first_idx + i
    return cols


class Cols:
    def __init__(self):
        self.list = build_columns()
        self.by_key = {c.key: c for c in self.list}
        self.last = self.list[-1]
        n = self.last.idx
        self.NOTES = get_column_letter(n + 1)
        self.GAP1 = get_column_letter(n + 2)
        self.CAGR = [get_column_letter(n + 3 + i) for i in range(4)]   # 25-30, 20-25, 15-25, 06-25
        self.GAP2 = get_column_letter(n + 7)
        self.SPACER = get_column_letter(n + 8)
        self.SCN_LABEL = get_column_letter(n + 9)
        self.SCN = {s: get_column_letter(n + 10 + i) for i, s in enumerate(["Bull", "Base", "Bear"])}
        self.SCN_NOTE = get_column_letter(n + 13)
        self.SWITCH = f"${self.NOTES}$2"

    def get(self, key):
        return self.by_key[key]

    def annual(self, year):
        """Column for the fiscal year (A/Y hist, E26 or E)."""
        for k in (str(year), f"{year}E"):
            if k in self.by_key:
                return self.by_key[k]
        return None

    def quarter(self, year, q):
        yy = f"{year % 100:02d}"
        return self.by_key.get(f"Q{q}/{yy}") or self.by_key.get(f"Q{q}/{yy}E")

    def half(self, year):
        return self.by_key.get(f"1H/{year % 100:02d}")

    def prior_same(self, c: Col):
        """Same period one year earlier."""
        if c.kind in ("A", "Y", "E26", "E"):
            return self.annual(c.year - 1)
        if c.kind in ("Q", "QE"):
            return self.quarter(c.year - 1, c.q)
        if c.kind == "H":
            return self.half(c.year - 1)
        return None

    def prev_period(self, c: Col):
        """Immediately preceding period of the same frequency (quarter→quarter, year→year)."""
        if c.kind in ("A", "Y", "E26", "E"):
            return self.annual(c.year - 1)
        if c.kind in ("Q", "QE"):
            if c.q == 1:
                return self.quarter(c.year - 1, 4)
            return self.quarter(c.year, c.q - 1)
        if c.kind == "H":
            return self.annual(c.year - 1)   # stock at prior year end
        return None

    def fy_annual_hist(self):
        return [c for c in self.list if c.kind in ("A", "Y")]


# --------------------------------------------------------------------------------------
# Row definitions
# --------------------------------------------------------------------------------------
@dataclass
class Row:
    key: Optional[str]
    label: str = ""
    style: str = "line"         # section | block | total | sub | line | pct | growth | check | memo | sched | blank | note
    fmt: str = FMT_NUM
    src: Optional[str] = None   # data key for historical hard-codes, e.g. "is.total_rev"
    agg: str = "sum"            # how 1H (and missing FY) aggregate: sum | last | avg | none
    hist: Optional[Callable] = None       # formula for all historical columns (overrides src)
    hist_q: Optional[Callable] = None     # formula only for quarter/1H hist columns
    qe: Optional[Callable] = None         # Q3/26E, Q4/26E
    e26: Optional[Callable] = None        # 2026E (default: sum of quarters for flows / Q4 for stocks)
    e: Optional[Callable] = None          # 2027E–2030E
    fc: Optional[Callable] = None         # shortcut: same formula for QE, E26 and E
    annual_only: bool = False   # CF/BS style: only annual cols + 1H/26 + 2026E+
    quarters_ok: bool = True    # False → leave Q/H cells empty (annual-only data rows)
    note: str = ""              # notes column text
    comment: str = ""           # cell comment on the label
    cagr: Optional[str] = None  # "cagr" | "avg" | None
    input_fc: bool = False      # forecast cells are inputs (blue bold)
    h1_override: Optional[Callable] = None
    fy_formula: Optional[Callable] = None  # formula for hist FY (Y) cols instead of data
    keep_blank_hist: bool = False
    texts: Optional[dict] = None   # deftext rows: colkey -> text


class Ctx:
    """Formula context for one column."""

    def __init__(self, cols: Cols, rows: dict, c: Col):
        self.cols, self.rows, self.c = cols, rows, c

    def r(self, key):
        return self.rows[key]

    def a(self, key, col: Col = None):
        col = col or self.c
        return f"{col.L}{self.rows[key]}"

    def py(self, key):
        p = self.cols.prior_same(self.c)
        return None if p is None else f"{p.L}{self.rows[key]}"

    def pp(self, key):
        p = self.cols.prev_period(self.c)
        return None if p is None else f"{p.L}{self.rows[key]}"

    def at(self, key, colkey):
        return f"{self.cols.get(colkey).L}{self.rows[key]}"

    def abs_at(self, key, colkey):
        return f"${self.cols.get(colkey).L}${self.rows[key]}"

    def q(self, key, n, year=None):
        col = self.cols.quarter(year or self.c.year, n)
        return f"{col.L}{self.rows[key]}"

    def h(self, key, year=None):
        return f"{self.cols.half(year or self.c.year).L}{self.rows[key]}"

    def fy(self, key, year):
        return f"{self.cols.annual(year).L}{self.rows[key]}"

    def scn(self, scn_row):
        s = self.cols.SCN
        return (f'CHOOSE(MATCH({self.cols.SWITCH},{{"Bull","Base","Bear"}},0),'
                f'${s["Bull"]}${scn_row},${s["Base"]}${scn_row},${s["Bear"]}${scn_row})')

    def base(self, scn_row):
        return f'${self.cols.SCN["Base"]}${scn_row}'


def yoy(key):
    return lambda x: f'=IFERROR({x.a(key)}/{x.py(key)}-1,"")' if x.py(key) else None


def ratio(num, den):
    return lambda x: f'=IFERROR({x.a(num)}/{x.a(den)},"")'


def put_comment(cell, text):
    if text:
        cm = Comment(text, "Model")
        cm.width, cm.height = 420, 160
        cell.comment = cm
