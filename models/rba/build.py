"""Build models/RBA_Model.xlsx — RB Global (NYSE/TSX: RBA) 3-statement model on the HII template.

Usage:  python -m models.rba.build   (from the repo root)
"""
from __future__ import annotations

import json
import os
import shutil
import subprocess
import sys

import openpyxl

from .assumptions import ASSUMP, DEF_EBITDA, DEF_NI, SCENARIOS
from .data import CATS, load
from .framework import Cols, put_comment
from .other_tabs import BBB_ROWS, write_bbb, write_charts, write_dcf
from .rows import FC_INPUT_CATS, build_rows
from .scenario import write_scenarios
from .writer import Data, set_widths, write_model

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
OUT = os.path.join(ROOT, "models", "RBA_Model.xlsx")
RECALC = os.environ.get("RECALC_SCRIPT", "")

TITLE = "RB Global, Inc. (NYSE / TSX: RBA) — 3-Statement Financial Model"
SUB = ("US GAAP as reported (2015+) · USD millions (per-share data in USD) · Source: RB Global / Ritchie Bros. Auctioneers "
       "Forms 40-F (2006–14), 6-K (2013–15 quarters), 10-K / 10-Q (2015+) and Form 8-K earnings releases (Ex. 99.1), "
       "SEC EDGAR CIK 0001046102")
BASIS = ("Basis notes: FY2006–10 Canadian GAAP and FY2011–12 / 2013–14 quarters IFRS (40-F / 6-K), FY2013–14 annuals US GAAP "
         "as recast in the FY2015 10-K; 2015+ US GAAP. Pre-2017 revenue on the legacy net (agency) basis — ASC 606 "
         "(2018, full retrospective) recast 2017 to gross inventory sales; 2016 shown with a recast memo. IronPlanet "
         "from 31-May-2017; IAA from 20-Mar-2023 (single reportable segment since 2023). Q3/26E–Q4/26E = FY26 outlook less "
         "1H/26 actuals.")


def used_cats(P):
    used_e, used_n = set(FC_INPUT_CATS), set(FC_INPUT_CATS) | {"AMORT"}
    for v in P.values():
        for k, val in v.items():
            if not isinstance(val, (int, float)) or abs(val) < 1e-9:
                continue
            if k.startswith("adjE."):
                used_e.add(k[5:])
            elif k.startswith("adjN."):
                used_n.add(k[5:])
    cats_e = [(k, lab) for k, lab, e, n in CATS if e and k in used_e and k != "AMORT"]
    cats_n = [(k, lab) for k, lab, e, n in CATS if n and k in used_n]
    return cats_e, cats_n


def add_item_comments(ws, cols, rowmap, labels):
    for c in cols.list:
        for key, labs in labels.get(c.key, {}).items():
            pre, cat = key.split(".")
            rk = ("a_" if pre == "adjE" else "n_") + cat
            if rk in rowmap and cat not in ("SBC", "AMORT", "TAX") and labs:
                cell = ws.cell(rowmap[rk], c.idx)
                if cell.value is not None:
                    put_comment(cell, "Company wording: " + " | ".join(labs))


def build(path=OUT, scenario="Base"):
    P, labels, notes, supp, conflicts = load()
    if conflicts:
        print("WARNING data conflicts:", conflicts[:10])
    cats_e, cats_n = used_cats(P)
    info = {"assump": ASSUMP, "def_ebitda": DEF_EBITDA, "def_ni": DEF_NI, "cats_e": cats_e, "cats_n": cats_n}
    cols = Cols()
    rows = build_rows(info)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model"
    rowmap = write_model(ws, cols, rows, Data(P), TITLE, SUB, BASIS)
    ws[f"{cols.NOTES}2"] = scenario
    add_item_comments(ws, cols, rowmap, labels)
    set_widths(ws, cols)
    write_scenarios(ws, cols, SCENARIOS)
    write_bbb(wb, cols, rowmap)
    write_dcf(wb, cols, rowmap)
    write_charts(wb, cols, rowmap)
    wb.create_sheet("Sheet1")
    wb.save(path)
    return cols, rowmap


def recalc(path):
    if not RECALC:
        print("RECALC_SCRIPT not set — skipping recalculation")
        return None
    out = subprocess.run([sys.executable, RECALC, path, "180"], capture_output=True, text=True)
    try:
        return json.loads(out.stdout)
    except Exception:
        print(out.stdout, out.stderr)
        return None


def snapshot(path, cols, rowmap):
    """Run the model in Bull and Bear and paste values into the Bull-Base-Bear snapshot panels."""
    vals = {}
    for scn in ("Bull", "Bear"):
        tmp = path.replace(".xlsx", f"_{scn}.xlsx")
        wb = openpyxl.load_workbook(path)
        wb["Model"][f"{cols.NOTES}2"] = scn
        wb.save(tmp)
        recalc(tmp)
        wbv = openpyxl.load_workbook(tmp, data_only=True)
        b = wbv["Bull-Base-Bear"]
        vals[scn] = {(r, i): b.cell(r, 3 + i).value for r in range(9, 48) for i in range(0, 11)}
        os.remove(tmp)
    wb = openpyxl.load_workbook(path)
    b = wb["Bull-Base-Bear"]
    for scn, start in (("Bull", 15), ("Bear", 27)):   # O and AA columns
        for (r, i), v in vals[scn].items():
            if r in (26, 27, 28, 29, 34, 44, 45, 46, 47):
                continue      # keep the live P/E-driven valuation rows as formulas
            cell = b.cell(r, start + i)
            if cell.value is not None and isinstance(cell.value, str) and cell.value.startswith("="):
                cell.value = v
        # valuation rows in snapshots: values for inputs (EPS etc.) are now static; keep PT formulas
    wb.save(path)


if __name__ == "__main__":
    cols, rowmap = build()
    res = recalc(OUT)
    print(json.dumps(res, indent=1)[:3000] if res else "no recalc")
    if res and res.get("status") in ("success", "errors_found"):
        snapshot(OUT, cols, rowmap)
        res = recalc(OUT)
        print(json.dumps(res, indent=1)[:2000])
