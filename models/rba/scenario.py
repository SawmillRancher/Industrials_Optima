"""Bull / Base / Bear scenario input table (right of the Model notes column) — mirrors HII CH:CK."""
from openpyxl.styles import Alignment

from .framework import BLACK, BLUE, FMT_NUM, FMT_PCT, GREY, LIGHT, MID, WHITE, fill, font, put_comment
from .rows import S_AUTO, S_BUYB, S_CAPEX, S_EBITDA, S_GTV_G, S_HET, S_MARGIN, S_OTH, S_TAX


def write_scenarios(ws, cols, S):
    """S: dict with guidance and per-lever base values / deltas."""
    L, NB, NBa, NBe, NOTE = cols.SCN_LABEL, cols.SCN["Bull"], cols.SCN["Base"], cols.SCN["Bear"], cols.SCN_NOTE

    def hdr(r, text, a="Bull", b="Base", c="Bear"):
        ws[f"{L}{r}"] = text
        ws[f"{L}{r}"].font = font(12, True, color=WHITE)
        ws[f"{L}{r}"].fill = fill(MID)
        for col, v in ((NB, a), (NBa, b), (NBe, c)):
            if v is not None:
                ws[f"{col}{r}"] = v
            ws[f"{col}{r}"].font = font(12, True, color=WHITE)
            ws[f"{col}{r}"].fill = fill(MID)
            ws[f"{col}{r}"].alignment = Alignment(horizontal="center")

    ws[f"{L}9"] = "SCENARIO INPUT TABLE — driver assumptions per Bull / Base / Bear"
    ws[f"{L}9"].font = font(11, True, color=BLACK)
    for col, v in ((NB, "Bull"), (NBa, "Base"), (NBe, "Bear")):
        ws[f"{col}9"] = v
        ws[f"{col}9"].font = font(12, True, color=BLACK)
        ws[f"{col}9"].fill = fill(LIGHT)
        ws[f"{col}9"].alignment = Alignment(horizontal="center")
    ws[f"{NOTE}9"] = ("Bull / Bear = Base + Δ (Δ inputs on each block header). Outlook block = high / mid / low end "
                      "of FY26 ranges.")
    ws[f"{NOTE}9"].font = font(10)

    def val(cell, v, fmt):
        ws[cell] = v
        is_formula = isinstance(v, str) and v.startswith("=")
        ws[cell].font = font(12, color=BLACK if is_formula else BLUE)
        ws[cell].number_format = fmt

    hdr(11, "FY26 outlook — Q2-26 release (4-Aug-26)", "High", "Mid", "Low")
    put_comment(ws[f"{L}11"], "FY26 outlook per RB Global Q2-2026 earnings release (Form 8-K Ex. 99.1, 4-Aug-2026): GTV "
                "growth 9%–11% (raised from 6%–9%); adjusted EBITDA $1,495m–$1,545m (raised low end from $1,485m); "
                "full-year tax rate (GAAP and adjusted) 23%–25%; capital expenditures $350m–$400m (PP&E net of "
                "disposal proceeds + intangible additions).")
    for r, lab, vals, fmt in [(S_GTV_G, "  GTV growth %", S["gtv_g"], FMT_PCT),
                              (S_EBITDA, "  Adjusted EBITDA ($m)", S["ebitda"], FMT_NUM)]:
        ws[f"{L}{r}"] = lab
        ws[f"{L}{r}"].font = font(12)
        for col, v in zip((NB, NBa, NBe), vals):
            val(f"{col}{r}", v, fmt)
    hdr(15, "FY26 outlook — point estimates (all cases)", None, None, None)
    for r, lab, v, fmt in [(S_TAX, "  Tax rate (GAAP & adjusted)", S["tax"], FMT_PCT),
                           (S_CAPEX, "  Capital expenditures ($m)", S["capex"], FMT_NUM)]:
        ws[f"{L}{r}"] = lab
        ws[f"{L}{r}"].font = font(12)
        val(f"{NBa}{r}", v, fmt)

    def lever(base_row, name, base_vals, d_bull, d_bear, fmt=FMT_PCT, years=(2027, 2028, 2029, 2030), floor=False,
              dfmt='\\+0.0%;\\-0.0%', note=""):
        hdr(base_row, name, None, "Δ →", None)
        for col, d in ((NB, d_bull), (NBe, d_bear)):
            ws[f"{col}{base_row}"] = d
            ws[f"{col}{base_row}"].font = font(12, True, color=BLUE)
            ws[f"{col}{base_row}"].fill = fill(GREY)
            ws[f"{col}{base_row}"].number_format = dfmt
        if note:
            ws[f"{NOTE}{base_row}"] = note
            ws[f"{NOTE}{base_row}"].font = font(10)
        for i, (y, bv) in enumerate(zip(years, base_vals)):
            r = base_row + 1 + i
            ws[f"{L}{r}"] = f"  {y}E" if isinstance(y, int) else f"  {y}"
            ws[f"{L}{r}"].font = font(12)
            val(f"{NBa}{r}", bv, fmt)
            for col in (NB, NBe):
                f = f"={NBa}{r}+${col}${base_row}"
                if floor:
                    f = f"=MAX(0,{NBa}{r}+${col}${base_row})"
                val(f"{col}{r}", f, fmt)

    lever(S_AUTO, "AUTO_gtvGrowth (Automotive GTV y/y)", S["auto"], *S["auto_d"],
          note="Automotive: US salvage volumes (accident frequency × total-loss rate), insurer share gains, "
               "price per vehicle.")
    lever(S_HET, "HET_gtvGrowth (Heavy Equipment & Transportation GTV y/y)", S["het"], *S["het_d"],
          note="HE&T: used-equipment supply (fleet de-fleeting, rental), pricing cycle, J.M. Wood / BigIron annualisation "
               "in 2026–27.")
    lever(S_OTH, "OTH_gtvGrowth (Other GTV y/y)", S["oth"], *S["oth_d"])
    lever(S_MARGIN, "ADJ_EBITDA_margin (% of total revenue)", S["margin"], *S["margin_d"],
          note="FY25 31.6% of revenue; FY26 outlook implies ~30% (inventory-sales mix and acquisitions dilute).")
    lever(S_BUYB, "Buybacks ($m)", S["buyb"], *S["buyb_d"], fmt=FMT_NUM,
          years=("2H/26E", 2027, 2028, 2029, 2030), floor=True, dfmt='\\+#,##0;\\-#,##0',
          note="NCIB approved Q1-26: up to 10.0m shares / $500m; $150m repurchased in Q2-26 ($350m remaining at "
               "30-Jun-26).")
