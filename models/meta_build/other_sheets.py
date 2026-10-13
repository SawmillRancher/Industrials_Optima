"""Rewire the Bull-Base-Bear, DCF and Charts tabs of the MLM template (current layout: DCF with live valuation date, fade period,
implied ROIC, target-IRR entry price, reverse DCF and Mauboussin EV/NOPAT; Bull-Base-Bear with buyback rows) to the META Model rows /
columns; keep the template chart XML. The template's DCF_old tab is not carried.

Column mapping: the template's 2025A base column becomes META 2026E (Q1-Q2 actual, Q3-Q4 calibrated to guidance; net debt is taken at
31-Dec-2026E) and the template's 2026E-2030E become META 2027E-2031E, so the DCF's explicit years are 2027E-2031E, the fade years 2032E-36E,
and 2026 cash flows (in year-end net debt) are not double counted. Every DATE(year,12,31) moves one year forward accordingly."""
import re, zipfile, shutil
from openpyxl.utils import get_column_letter as L
from openpyxl.comments import Comment
from fw import R, TEMPLATE, ANN, NOTE

# template Model row -> META row key
MAP = {158: 'is_rev', 164: 'r_cash', 166: 'is_oi', 182: 'is_dda', 186: 'sh_dil', 211: 'r_adj', 212: 'r_adj_m', 219: 'r_aebit', 221: 'r_adj',
       229: 'e_curni', 231: 'e_cureps', 238: 'gm_adv', 245: 'gm_ebitda_m', 247: 'gm_op_m', 260: 'cf_wc', 262: 'cf_capex_all', 263: 'cf_acq',
       280: 'fcf', 282: 'fcf_m', 290: 'bb_px', 291: 'bb_cash', 292: 'bb_sh', 331: 'v_nci', 408: 'ra_ic', 411: 'ra_roic', 422: 'rn_ronta',
       428: 'lv_nd', 433: 'v_px', 434: 'v_sh'}
CMAP = {'BV': ANN[2026].c, 'CA': ANN[2027].c, 'CB': ANN[2028].c, 'CC': ANN[2029].c, 'CD': ANN[2030].c, 'CE': ANN[2031].c, 'CF': NOTE}
REF = re.compile(r"Model!(\$?)([A-Z]{1,3})(\$?)(\d+)")
HIST = [ANN[y].c for y in range(2009, 2026)]

def remap(f):
    def rep(m):
        c, r = m.group(2), int(m.group(4))
        if c == 'CF' and r == 2: return f'Model!{m.group(1)}{NOTE}{m.group(3)}2'          # scenario switch
        if r not in MAP: raise KeyError(f'unmapped template Model row {r} in {f}')
        return f'Model!{m.group(1)}{CMAP.get(c, c)}{m.group(3)}{R[MAP[r]]}'
    return REF.sub(rep, f)

def hist_stat(f):
    """LT historical AVERAGE / MIN / MAX over the template's annual columns -> META FY2009-FY2025 annual columns."""
    m = re.match(r'=IFERROR\((AVERAGE|MIN|MAX)\(Model!C(\d+),', f)
    if not m: return None
    fn, r = m.group(1), int(m.group(2))
    rr = R[MAP[r]]
    return f'=IFERROR({fn}({",".join(f"Model!{c}{rr}" for c in HIST)}),"")'

def shift_dates(f):
    f = re.sub(r'DATE\((\d{4}),12,31\)', lambda m: f'DATE({int(m.group(1)) + 1},12,31)', f)
    f = re.sub(r'DATE\((\d{4})-1,12,31\)', lambda m: f'DATE({int(m.group(1)) + 1}-1,12,31)', f)
    return f

TEXT = {
    'Bull-Base-Bear': {
        'B2': 'Meta Platforms, Inc. (META) — Bull / Base / Bear Case P&L Summary',
        'B3': ('Adjusted figures (current definition: GAAP + company-quantified one-offs, share-based compensation expensed; adjusted EBITDA = model definition) · '
               'USD millions · Live panel linked to Model — toggle scenario via Model!CD2. Bull and Bear panels are value snapshots generated from the Model with the '
               'switch set to each case. 2026E = Q1–Q2 actual + Q3 calibrated to revenue guidance ($61–64bn) and 2H to FY26 expense guidance ($165–169bn).'),
        'G26': 24, 'H26': 24, 'S26': 28, 'T26': 28, 'AD26': 18, 'AE26': 18, 'AF26': 18,
    },
    'DCF': {
        'B2': 'DCF VALUATION — META PLATFORMS, INC. (META)',
        'C5': 0.0425, 'D5': 'Assumption (~US 10Y Treasury yield, Oct-2026; government rate sources not reachable from the build environment); update to market',
        'D6': 'Assumption: US implied ERP ~4.5–5.5%',
        'C7': 1.20, 'D7': ('Assumption: mega-cap internet and advertising peers (GOOGL, AMZN, MSFT, SNAP, PINS) ~1.0–1.3; advertising is cyclical and the AI '
                           'capex cycle raises operating leverage'),
        'C9': 0.05, 'D9': 'May-2026 notes 4.55%–6.45% (effective 4.60%–6.48%); Nov-2025 notes 4.20%–5.75% (Q2-26 10-Q Note 8); AA- / Aa3 class issuer',
        'C10': 0.16, 'D10': 'Model 2027E+ effective tax rate (2026 guidance 15–17% for the remaining quarters)',
        'C12': 0.05, 'D12': 'Notes ~$84bn at 30-Jun-2026 vs ~$1.9tn market cap ($745 × ~2.57bn diluted shares); net debt ≈ 0 (cash and marketable securities $90.3bn)',
        'D15': 'Periods = (31-Dec of the year − today) / 365.25; cash flows assumed at year-end; 2026E is the base year (its cash flows are in YE2026E net debt)',
        'B18': 'Adj. EBIT (current definition, full-cost) less cash one-off items (legal, restructuring)',
        'B21': 'plus D&A',
        'B22': '− Capex incl. principal payments on finance leases and acquisitions (M&A lever 2027E–31E)',
        'B26': 'Fade-period growth (2032E–36E; 9.0% → 5.0% linear)', 'I26': 0.09, 'M26': 0.05, 'C27': 0.035,
        'D32': ("Meta's adjusted EBIT (current definition) expenses share-based compensation (a real cost: ~11% of revenue) and all D&A; only the one-off legal and "
                'restructuring items the company quantified are added back, and those are deducted again inside row 18 because they are cash costs. D&A is added back in '
                'row 21 (dominated by servers and data centers). Row 22 = purchases of property and equipment + principal payments on finance leases (the company\'s capex '
                'measure) + acquisitions; the fade period and terminal value carry capex (no new acquisitions). Taxes on net share settlement (~90% of SBC) are a financing '
                'flow and are not deducted. Operating and finance lease liabilities are excluded from net debt (lease costs sit in operating costs). '
                '2026E is the base year: its cash flows are in YE2026E net debt.'),
        'D43': 'Opening IC = Model 2026E book IC (equity + net debt at 31-Dec-2026E), rolled forward with capex + M&A − D&A + ΔWC.',
        'B48': 'Less: Net debt + NCI (Model 2026E year end; net debt = notes + CP − cash − marketable securities)',
        'D48': 'META: no noncontrolling interests; non-marketable equity investments ($30bn at 30-Jun-2026) are not added back (conservative).',
        'B50': 'Diluted shares (m, Model 2026E period end)',
        'D76': 'Rationale (Meta characteristics)',
        'C78': 0.10, 'D78': 'Model adjusted EBIT CAGR 2026–31E (~14%): high-teens revenue growth fading to ~10% (impressions + price), with D&A from the 2025–27 capex wave absorbing operating leverage',
        'C79': 0.18, 'D79': ('Historical ROIC 25–35% (2015–25) on an asset-light ad platform; incremental returns on the AI data-center build-out are unproven, so ROIIC is set '
                             'below history (assumption)'),
        'C80': 15, 'D80': 'Network effects across 3.6bn daily people, first-party data and ad-tech scale sustain a long competitive-advantage period; regulatory and youth-safety litigation risk',
        'B90': 'VALUATION  —  Applied to 2029E NOPAT',
        'B91': '2029E Adj. EBIT less cash one-off items',
        'B93': '2029E NOPAT',
        'B95': 'Implied Enterprise Value (2029E basis)',
        'B96': 'Less: YE2028E net debt + NCI (Model)',
        'B98': 'Diluted shares 2029E (m, Model)',
        'B99': 'Implied share price 2029 ($)',
        'B106': 'Plus: YE2026E net debt + NCI',
        'B108': 'META market-implied EV/NOPAT (on 2029E NOPAT)',
    },
    'Charts': {
        'B2': 'Meta Platforms, Inc. — Operating Performance (2012–2025) and Model Forecast (2026E–2031E)',
        'B3': ('USDm · Linked to Model. Bars = contribution of advertising and of other revenue to total revenue growth (pp). Adjusted EBITDA = model definition (Meta '
               'publishes none); income from operations GAAP. ROIC / RONTA on adjusted EBIT (current definition). 2012: no y/y growth in the model columns (FY2011 base).'),
        'B8': 'Advertising contribution to growth (pp)', 'B9': 'Other revenue contribution to growth (pp)', 'B12': 'Adjusted EBITDA (model)', 'B13': 'Income from operations (GAAP)',
        'B14': 'Adj. EBITDA Margin %', 'B15': 'Operating Margin (GAAP) %',
    },
}
BBB_ROWS = {9: 'Revenue', 10: '   Revenue y/y %', 11: '   Advertising revenue y/y %', 12: 'Adjusted EBITDA (model definition)',
            15: 'Income from operations (GAAP)', 16: '   Income from operations y/y %', 17: '   Operating margin (GAAP) %',
            19: 'Adjusted net income (current definition)', 21: 'Diluted shares (m)',
            35: 'Adjusted EBITDA (leverage basis)', 36: '   Adjusted EBITDA % of revenue'}
CHART_TITLES = {
    'Aggregates Revenue Drivers: Shipment Volume vs. Price (ASP), y/y %': 'Revenue Growth: Advertising vs. Other Revenue Contribution, pp',
    'Adjusted EBITDA &amp; Earnings from Operations (USDm)': 'Adjusted EBITDA (model) &amp; Income from Operations (USDm)',
}

def build_charts(ws):
    for i, y in enumerate(range(2012, 2032)):
        c = L(3 + i); mc = ANN[y].c
        ws[f'{c}7'] = y if y <= 2025 else f'{y}E'
        ws[f'{c}8'] = f'=IFERROR(Model!{mc}{R["gc_adv"]}+0,0)'
        ws[f'{c}9'] = f'=IFERROR(Model!{mc}{R["gc_oth"]}+0,0)'
        ws[f'{c}11'] = f'={c}7'
        ws[f'{c}12'] = f'=Model!{mc}{R["r_adj"]}'
        ws[f'{c}13'] = f'=Model!{mc}{R["is_oi"]}'
        ws[f'{c}14'] = f'=Model!{mc}{R["gm_ebitda_m"]}'
        ws[f'{c}15'] = f'=Model!{mc}{R["gm_op_m"]}'
        ws[f'{c}17'] = f'={c}7'
        ws[f'{c}18'] = f'=IFERROR(Model!{mc}{R["ra_roic"]}+0,NA())'
        ws[f'{c}19'] = f'=IFERROR(Model!{mc}{R["rn_ronta"]}+0,NA())'
        ws[f'{c}20'] = f'=IFERROR(Model!{mc}{R["gm_ebitda_m"]}+0,NA())'
        ws[f'{c}21'] = f'=IFERROR(Model!{mc}{R["fcf_m"]}+0,NA())'

def do_all(wb):
    build_charts(wb['Charts'])
    for name in ('Bull-Base-Bear', 'DCF', 'Charts'):
        ws = wb[name]
        if name == 'DCF': ws['D34'] = None                         # template-specific opening IC (acquisition consideration) — rebuilt below
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and 'Model!' in c.value and not (name == 'Charts' and 7 <= c.row <= 21):
                    h = hist_stat(c.value) if name == 'Bull-Base-Bear' else None
                    c.value = h or remap(c.value)
                if isinstance(c.value, str) and 'DATE(' in c.value and name != 'Charts':
                    c.value = shift_dates(c.value)
        for addr, t in TEXT.get(name, {}).items():
            ws[addr] = t
    b = wb['Bull-Base-Bear']
    for r, t in BBB_ROWS.items():
        for c in ('B', 'N', 'Z'): b[f'{c}{r}'] = t
    for c0 in ('C', 'O', 'AA'):
        i = b[f'{c0}7'].column
        b.cell(7, i).value = '2026E'
        for k in range(1, 6): b.cell(7, i + k).value = f'{2026 + k}E'
        b.cell(7, i + 6).value = "26-'31 CAGR"
    cm = Comment('P/E on adjusted EPS (current definition: GAAP + company-quantified one-offs and discrete tax items; SBC expensed). META at $745 (per the user, '
                 '12-Oct-2026) trades at ~26x 2026E and ~23x 2027E adjusted EPS (Model). Exit multiples are assumptions: Base 24x, Bull 28x, Bear 18x, applied to '
                 '2030E–31E (2029E–31E for Bear), in line with mega-cap internet peers through the cycle.', 'META model')
    cm.width, cm.height = 420, 130; b['N26'].comment = cm
    cm2 = Comment(f'Snapshot (values) of Model rows {R["bb_cash"]}, {R["bb_px"]} and {R["bb_sh"]} (Share Buyback Schedule) with Model!{NOTE}2 = Bull / Bear. '
                  'Base panel (D40:H42) is live-linked.', 'META model')
    cm2.width, cm2.height = 420, 90; b['P40'].comment = cm2
    b['J2'] = 'Share price ($):'
    for c in ('B', 'N', 'Z'): b[f'{c}29'] = 'IRR (to 31-Dec of Y-1)'
    for c in ('B', 'N', 'Z'):
        if b[f'{c}40'].value: b[f'{c}40'] = '   Share repurchases ($m; 2026E = 2H)'
    d = wb['DCF']
    for i, col in enumerate('DEFGHIJKLM'):
        d[f'{col}17'] = f'{2027 + i}E'
    a26, a27 = ANN[2026].c, ANN[2027].c
    d['D34'] = f'=Model!{a26}{R["ra_ic"]}'
    d['C48'] = f'=Model!{a26}{R["lv_nd"]}+Model!{a26}{R["v_nci"]}'
    d['C50'] = f'=Model!{a26}{R["v_sh"]}'
    d['C52'] = f'=Model!{a26}{R["v_px"]}'
    d['P17'] = 'CAGR 27–31E'; d['Q17'] = 'CAGR 27–36E'
    for r in range(18, 25):                                   # CAGRs are not meaningful on series that start negative (2027E UFCF)
        for c in ('P', 'Q'):
            v = d[f'{c}{r}'].value
            if isinstance(v, str) and v.startswith('=') and 'IFERROR' not in v:
                d[f'{c}{r}'] = f'=IFERROR(IF(D{r}<=0,"n/m",{v[1:]}),"n/m")'

def snapshot_values(wb, bull, bear):
    ws = wb['Bull-Base-Bear']
    cols = 'CDEFGHIJKL'
    for snap, off in ((bull, 'OPQRSTUVWX'), (bear, ['AA', 'AB', 'AC', 'AD', 'AE', 'AF', 'AG', 'AH', 'AI', 'AJ'])):
        for (r, c), v in snap.items():
            tgt = f'{off[cols.index(c)]}{r}'
            cur = ws[tgt].value
            if isinstance(cur, str) and cur.startswith('='): continue
            if r in (7, 26): continue
            ws[tgt] = v

def inject_charts(path):
    """Replace openpyxl-rewritten chart parts with the template chart XML (combo charts, colours, layout), retitled, ranges kept
    (Charts!C:V = 2012-2031E) and cached values removed so Excel redraws from the META cells."""
    tz = zipfile.ZipFile(TEMPLATE)
    tmpl = {n: tz.read(n) for n in tz.namelist() if n.startswith('xl/charts/chart')}
    tmp = path + '.tmp'
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in tmpl:
                t = tmpl[item.filename].decode('utf-8')
                for a, b2 in CHART_TITLES.items(): t = t.replace(a, b2)
                t = re.sub(r'<c:numCache>.*?</c:numCache>', '', t, flags=re.S)
                t = re.sub(r'<c:strCache>.*?</c:strCache>', '', t, flags=re.S)
                data = t.encode('utf-8')
            zout.writestr(item, data)
    shutil.move(tmp, path)
