"""Build META_Model.xlsx from the MLM template (same layout, formatting and process as the prior models)."""
import sys, copy, os, json
from openpyxl.styles import Font
from fw import *
import fw
from spec import build_spec, scen_layout, SCEN_ROWS, SC
from spec2 import build_spec2

OUT = os.path.join(os.path.dirname(BASE), 'META_Model.xlsx')
G26 = MAN['guidance_2026']

# ====================================================================== inputs (V0)
def wc_inputs():
    q = DATA['2025']['x']
    capx = -q['cf_capex'] - q['cf_lease']
    return {'dso': round(q['bs_ar'] / q['rev'] * 365), 'app': round(q['bs_ap'] / capx, 3), 'ocl': round(q['bs_ocl'] / q['rev'], 3)}
WC = wc_inputs()
WC['note'] = (f'Forecast inputs at fiscal-year-end (31-Dec-2025) levels on FY2025 revenue / capex (DSO {WC["dso"]}d; accounts payable {WC["app"]:.1%} of capex incl. '
              f'finance leases; accrued and other current liabilities {WC["ocl"]:.1%} of revenue); year ends are compared with year ends.')
_x25 = DATA['2025']['x']
NS_R = round(-_x25['cf_netshare'] / _x25['cf_sbc'], 2)

V0 = {
    'out_hdr': 'Outlook — Q2-26 release (8-K Ex. 99.1, 29-Jul-2026)',
    'out_note': ('Q3/26 revenue $61–64bn (FX ~1% headwind); FY26 total expenses $165–169bn (low end raised for the $2.4bn Q2 legal charges); FY26 operating income above 2025; '
                 'FY26 capital expenditures incl. principal payments on finance leases $130–145bn; tax rate 15–17% for the remaining quarters. All from the 8-K (SEC EDGAR).'),
    'outlook': [
        ('out_q3rev', '  Q3/26 total revenue ($m)', tuple(G26['q3_revenue']), 'm', 'Guided $61–64bn. Drives the Q3/26E price-per-ad calibration.'),
        ('out_exp', '  FY26 total expenses ($m) — Bull = low end, Bear = high end', tuple(G26['expenses']), 'm',
         'Guided $165–169bn. Drives the 2H/26E other-operating-cost calibration (lower expenses in the Bull case).'),
        ('bb26', '  2H/26E share repurchases ($m)', (5000.0, 0.0, 0.0), 'm',
         'No repurchases in 1H/26 (Q4/25 none either) while capex is funded; $50bn+ authorization remaining (10-Q). Base / Bear nil, Bull $5bn (assumption).'),
    ],
    'points': [
        ('pt_ddag', '  Q3/Q4-26E D&A growth q/q', 0.07, 'pct', 'Q4/25 → Q1/26 +10.9%, Q1 → Q2/26 +6.0% (servers and data centers placed in service). Assumption.'),
        ('pt_sbc', '  Q3/Q4-26E share-based compensation ($m)', DATA['Q2-26']['x']['cf_sbc'], 'm', 'Q2/26 level held ($7,658m; Q1/26 $6,032m).'),
        ('pt_etr', '  Q3/Q4-26E tax rate', 0.16, 'pct', 'Guided 15–17% for the remaining quarters of 2026 (mid-point).'),
        ('pt_capex', '  FY26 capital expenditures incl. finance-lease principal ($m)', G26['capex_mid'], 'm', 'Guided $130–145bn (narrowed from $125–145bn); mid-point. 1H/26 $50.9bn.'),
        ('pt_yld', '  Q3/Q4-26E interest-income yield on cash and marketable securities', 0.04, 'pct', 'Q2/26 interest income $859m on ~$86bn average (~4.0%, 10-Q).'),
        ('pt_rate', '  Q3/Q4-26E interest rate on debt, net of capitalized interest', 0.042, 'pct',
         'Q2/26 interest expense $783m (notes $754m net of capitalized interest) on ~$71.5bn average notes (~4.2%, 10-Q).'),
        ('pt_dps', '  Q3/Q4-26E dividend per share ($)', 0.525, 'ps', 'Quarterly $0.525 declared in 2025 and 2026 (XBRL).'),
    ],
    'lev': (0.5, 0.25, 0.0),
    'lev_note': ('Meta moved from net cash ($22bn at YE2025, cash and marketable securities $81.6bn vs notes $58.7bn) to ~$6.6bn net cash at 30-Jun-2026 as capex '
                 'outran operating cash flow. Ceiling (assumption) up to which surplus cash is returned via buybacks; above it, cash flow reduces debt. Bull = more buybacks.'),
    'blocks': [
        ('IMPR', 'IMPR (ad impressions y/y)', 0.02, -0.04, [0.10, 0.09, 0.08, 0.07, 0.06],
         '1H/26 +19% / +14% (Q1 / Q2); FY25 +12%, FY24 +11%; Threads, WhatsApp Status and Reels add inventory; DAP growth ~3–4% p.a. Assumption.', 'pct', False),
        ('PRICE', 'PRICE (average price per ad y/y)', 0.02, -0.04, [0.08, 0.06, 0.05, 0.04, 0.04],
         '1H/26 +12% (both quarters); FY25 +9%, FY24 +10%; AI ranking and ad-format gains vs mix shift to lower-monetizing geographies. Assumption.', 'pct', False),
        ('FOTH', 'FOTH (FoA other revenue y/y)', 0.10, -0.15, [0.40, 0.30, 0.25, 0.20, 0.15],
         '1H/26 +73%; FY25 +50% (WhatsApp business messaging, Meta Verified subscriptions). Assumption.', 'pct', False),
        ('RLREV', 'RLREV (Reality Labs revenue y/y)', 0.10, -0.10, [0.10, 0.10, 0.10, 0.10, 0.10],
         '1H/26 +7%; FY25 −3% (Quest declines offset by AI glasses). Assumption.', 'pct', False),
        ('RLLOSS', 'RLLOSS (Reality Labs operating loss, $m)', -3000.0, 4000.0, [20000.0, 19500.0, 19000.0, 18000.0, 17000.0],
         'FY25 loss $19.2bn; 1H/26 $8.6bn (−1% y/y). The company said 2026 Reality Labs losses would be similar to 2025. Assumption: gradual narrowing.', 'm', False),
        ('OPEX', 'OPEX (other operating costs excl. D&A, SBC, one-offs; % of revenue)', -0.02, 0.03, [0.41, 0.40, 0.39, 0.385, 0.38],
         'FY25 39.1%, 1H/26 39.4%, 2026E ~42% (guidance-implied 2H step-up: AI compute, cloud capacity, talent). Assumption: modest operating leverage.', 'pct', False),
        ('CAPEX', 'CAPEX (purchases of property and equipment % of revenue)', -0.04, 0.05, [0.46, 0.38, 0.31, 0.26, 0.23],
         'FY25 34.7% ($69.7bn); 2026E ~53% (guidance mid-point $137.5bn incl. finance leases); 2019–21 16–21%. AI data-center build-out (Hyperion, Prometheus). Assumption: intensity fades toward ~23%, above the pre-AI level.', 'pct', False),
        ('MNA', 'MNA (acquisitions, $m)', 2000.0, -1000.0, [1000.0, 1000.0, 1000.0, 1000.0, 1000.0],
         'Acquisitions of businesses and intangible assets: FY25 $0.1bn, FY24 $0.3bn, FY22 $1.3bn (Within, Kustomer era). Base $1bn p.a. (assumption).', 'm', True),
    ],
    'dep_r': {2027: 0.11, 2028: 0.11, 2029: 0.115, 2030: 0.115, 2031: 0.115},
    'dep_note': ('FY2025 11.8% (D&A $18.6bn on opening net PP&E $121.3bn + ½ × $72.2bn capex incl. finance leases); servers depreciated over 5.5 years from 2025, '
                 'buildings and network over longer lives; 2026E implied ~10.8%. 2027E+ input (assumption).'),
    'sbc_pct': {2027: 0.115, 2028: 0.11, 2029: 0.11, 2030: 0.105, 2031: 0.105},
    'etr': {2027: 0.16, 2028: 0.16, 2029: 0.16, 2030: 0.165, 2031: 0.165},
    'dps_g': 0.05, 'ns_r': NS_R,
    'lease': {2027: -4000.0, 2028: -4000.0, 2029: -4000.0, 2030: -4000.0, 2031: -4000.0},
    'iss_2h': 18.0, 'iss_ae': 30.0,
    'iss_note': ('Net shares issued on RSU vesting after shares withheld for taxes (assumption): 1H/26 +18m with no repurchases (2,530m → 2,548m); '
                 '2H/26E +18m; 2027E+ 30m a year.'),
    'dil': 23.0,
    'repay': {2027: 2750.0, 2028: 1500.0, 2029: 1000.0, 2030: 5000.0, 2031: 0.0},
    'repay_note': ('Future principal payments for the notes (Q2-26 10-Q Note 8): 2027 $2,750m, 2028 $1,500m, 2029 $1,000m, 2030 $5,000m, thereafter $73,750m '
                   '(2031 not shown separately; nil assumed). Face $84.0bn; carrying amount $83.66bn.'),
    'refi': 1.0,
    'min_cash': 15000.0,
    'min_note': 'Cash and cash equivalents $15.5bn at 30-Jun-2026 (marketable securities $74.8bn held separately). Assumption.',
    'r_debt': {2027: 0.045, 2028: 0.046, 2029: 0.047, 2030: 0.048, 2031: 0.048},
    'r_debt_note': ('Notes 3.50%–6.45% (effective 3.63%–6.48%); ~4.2% net of capitalized interest in Q2/26; rising as capitalized interest falls with fewer data centers '
                    'under construction and as new issues price above the 2022–24 coupons. Assumption.'),
    'r_cash': {2027: 0.0375, 2028: 0.035, 2029: 0.035, 2030: 0.035, 2031: 0.035},
    'r_cash_note': 'Q2/26 ~4.0% on cash and marketable securities (10-Q); lower short-term rates assumed.',
    'wc': WC,
    'px': MAN['price']['px'],
    'px_note': ('VALUATION — share price input $745.00 (per the user, 12-Oct-2026); ~$1.9tn market capitalisation; update to market'),
    'mna_note': 'Scenario lever MNA (Base $1bn p.a.). Funded from cash / debt (debt schedule).',
    'rev_guid_note': 'Company Q3/26 revenue guidance $61–64bn (Q2-26 release, 29-Jul-2026). Drives the price-per-ad calibration.',
    'exp_guid_note': 'Company FY26 total expenses guidance $165–169bn (Q2-26 release; low end raised from $162bn for the Q2 legal charges). Drives the 2H cost calibration.',
}

# ====================================================================== compute engine
def hist_value(row, col):
    kw = row.kw; path = kw.get('data')
    if not path: return None
    v = xget(pkey(col), path)
    return v if kw.get('textdata') else num(v)

def first(kw, keys):
    for k in keys:
        if k in kw: return k
    return None

def compute(row, col):
    kw = row.kw; k = col.kind
    CTX.col = col
    if kw.get('ann') and col.q: return None, False
    if k in HISTK:
        v = hist_value(row, col)
        if v is not None: return v, True
        for key in ('hist', 'all'):
            if key in kw:
                return kw[key](), False
        return None, False
    if k == 'QE':
        if row.spec2: return None, False
        if 'qe_in' in kw:
            v = kw['qe_in']; return (v[col.q] if isinstance(v, dict) else v), True
        key = first(kw, ('qe', 'fc', 'all'))
        if not key: return None, False
        f = kw[key]()
        return (f, False) if f is not None else (None, False)
    if k == 'E26':
        if 'e26_in' in kw: return kw['e26_in'], True
        if 'e26' in kw: return kw['e26'](), False
        if kw.get('avg') and ('qe' in kw or 'qe_in' in kw): return f'=AVERAGE({Q1}{row.r}:{Q4E}{row.r})', False
        if kw.get('flow') and not row.spec2 and any(x in kw for x in ('qe', 'qe_in', 'fc')): return S26(row.key), False
        key = first(kw, ('fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    if k == 'AE':
        if 'ae_in' in kw:
            v = kw['ae_in']; return (v[col.year] if isinstance(v, dict) else v), True
        key = first(kw, ('ae', 'fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    return None, False

# ====================================================================== writers
CAGR_G = {'is_rev', 'g_rev', 'adv', 'foa_oth', 'rl_rev', 'oth_rev', 'foa_rev', 'foa_oi', 'o_cur', 'r_adj', 'r_ebitda', 'is_oi', 'd_dda', 'd_sbc', 'd_cash',
          'is_costs', 'e_curni', 'e_cureps', 'cf_cfo', 'fcf', 'bs_ta', 'bs_eq', 'sh_dil', 'is_ni', 'eps_dil', 'capex_all', 'dap', 'hc', 'geo_usc', 'geo_eu', 'geo_ap', 'geo_row'}
CAGR_A = {'adv_g', 'impr_g', 'price_g', 'foa_m', 'o_cur_m', 'r_adj_m', 'is_oim', 'r_aebit_m', 'gm_rev', 'gm_adv', 'gm_cogs', 'gm_rd', 'gm_mm', 'gm_ga',
          'gm_ebitda_m', 'gm_op_m', 'gm_net_m', 'gm_cur_m', 'fcf_m', 'fcf_conv', 'capex_pct', 'capex_all_pct', 'ra_roic', 'rn_ronta', 'roe', 'lv_x', 'd_sbc_pct',
          'd_cash_pct', 'wc_nwc_p', 'is_etr', 'e_rate', 'hr_dda'}
A = {y: ANN[y].c for y in ANN}

def style_cagr(ws, r, on):
    src = 16 if on else 12
    for c in CAGR:
        ws[f'{c}{r}']._style = copy.copy(TM[f'{L(CI(c) + TSHIFT)}{src}']._style)

def write_row(ws, row):
    r = row.r; arch = ARCH[row.arch]
    clone_row_style(ws, arch, r)
    ws.cell(r, 2).value = row.label or None
    if row.kw.get('note'): ws[f'{NOTE}{r}'] = row.kw['note']
    style_cagr(ws, r, row.key in CAGR_G or row.key in CAGR_A)
    if row.arch in ('section', 'block', 'blank', 'sub') and not row.kw.get('data'):
        return
    if row.key in DEFS:
        write_defs(ws, row); return
    nf = NF[row.nf]
    for col in COLS:
        v, inp = compute(row, col)
        if v is None: continue
        cell = ws[f'{col.c}{r}']
        cell.value = v
        if not row.kw.get('textdata'): cell.number_format = nf
        if inp:
            set_color(cell, BLUE, bold=True if col.kind in FCK else None)
            lk = row.kw.get('labkey')
            if lk and col.kind in HISTK and isinstance(v, (int, float)) and abs(v) > 1e-9:
                t = xget(pkey(col), lk)
                if t: comment(cell, f'Company disclosure ({pkey(col)}): {t}', 420, 120)
        elif row.arch == 'growth' and col.kind in HISTK:
            set_color(cell, GRAY)
        else:
            set_color(cell, BLACK)
    rc = CELL_COMMENTS.get(row.key, {})
    for c, t in rc.items(): comment(ws[f'{c}{r}'], t)
    if row.key in CAGR_G:
        ws[f'{CAGR[0]}{r}'] = f'=IFERROR(({A[2031]}{r}/{A[2026]}{r})^(1/5)-1,"n/m")'
        ws[f'{CAGR[1]}{r}'] = f'=IFERROR(({A[2025]}{r}/{A[2020]}{r})^(1/5)-1,"n/m")'
        ws[f'{CAGR[2]}{r}'] = f'=IFERROR(({A[2025]}{r}/{A[2015]}{r})^(1/10)-1,"n/m")'
        ws[f'{CAGR[3]}{r}'] = f'=IFERROR(({A[2025]}{r}/{A[2009]}{r})^(1/16)-1,"n/m")'
        for c in CAGR: ws[f'{c}{r}'].number_format = '0%;\\(0%\\);"n/m"'
    elif row.key in CAGR_A:
        ws[f'{CAGR[0]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2027, 2032))}),"n/m")'
        ws[f'{CAGR[1]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2021, 2026))}),"n/m")'
        ws[f'{CAGR[2]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2016, 2026))}),"n/m")'
        ws[f'{CAGR[3]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2009, 2026))}),"n/m")'
        for c in CAGR: ws[f'{c}{r}'].number_format = NF['pct']
    t = LABEL_SOURCES.get(row.key)
    if t: comment(ws.cell(r, 2), t)

REL = 'Source: Meta / Facebook earnings releases (Form 8-K Item 2.02, Ex. 99.1) — '
LABEL_SOURCES = {
    'adv': (REL + 'reconciliation of GAAP to non-GAAP results ("GAAP advertising revenue"); segment information from Q4-2021. FY2009–11: IPO prospectus (424B4, 18-May-2012) '
            'MD&A revenue table. Each period as originally reported in its own release.'),
    'impr_g': ('Source: releases from Q4-2021 ("Ad impressions" / "Average price per ad" highlights); 2014–Q3 2021: 10-Q / 10-K MD&A ("the number of ads delivered '
               'increased by x%"; "the average price per ad increased by y%"). Published as whole percentages. Q4 figures before 2021 are not published (the 10-K gives the year).'),
    'foa_oth': (REL + 'segment information (Q4-2021 release: 2019–2021 recast; 2022+ releases). Not published for earlier periods.'),
    'dap': REL + 'text highlights ("Family daily active people (DAP)"; "Daily active users (DAUs)"; "Headcount").',
    'geo_usc': ('Source: 10-Q / 10-K revenue note ("Revenue by geography is based on the addresses of our customers"), from 2018. Q4 = fiscal year (10-K) − nine months '
                '(Q1–Q3 10-Qs), derived.'),
    'is_rev': (REL + 'condensed consolidated statements of income, as originally reported; FY2009–11 from the IPO prospectus (424B4). FY2011 marketing & sales / G&A '
               'were reclassified in the FY2012 release ($427m / $280m as first reported → $393m / $314m) — not recast.'),
    'd_dda': REL + 'condensed consolidated statements of cash flows (three-month statements for every quarter).',
    'd_legal': 'Source: release text and footnotes (amounts the company quantified); see manual_items.json "one_offs" for the wording and URL per quarter.',
    'd_intinc': 'Source: XBRL company facts (InvestmentIncomeInterest; InterestExpense / InterestExpenseNonoperating), annual; Q1/Q2-26 from the Q2-26 10-Q MD&A table.',
    'o_pub': REL + 'Reconciliation of Non-GAAP Results to Nearest GAAP Measures (non-GAAP income from operations, IPO to Q4-2016).',
    'e_pub': (REL + 'non-GAAP diluted EPS (IPO to Q4-2016); thereafter the "excluding" EPS the company stated for the quantified items '
              '(Q4-17, FY2017, Q1-19, Q2-19, Q3-20, Q4-22, Q3-25, Q1-26).'),
    'fcf_pub': REL + 'reconciliation of GAAP to non-GAAP results (free cash flow, from Q2-2013).',
    'cf_ni': REL + 'condensed consolidated statements of cash flows (three-month columns); FY2009–11 from the IPO prospectus.',
    'bs_cash': (REL + 'condensed consolidated balance sheets (period ends); FY2010–11 from the IPO prospectus (no 31-Dec-2009 balance sheet). Shares outstanding: '
                '10-Q / 10-K balance-sheet caption (Class A + Class B issued and outstanding).'),
    'sh_basic': (REL + 'statements of income (weighted-average shares). FY2009–11: prospectus Note 2 (two-class EPS: basic = Class A + B; diluted = Class A with Class B '
                 'converted and participating income reallocated). No stock splits; Class B carries 10 votes, same economics.'),
    'ds_total': 'Source: release balance sheets; Q2-26 10-Q Note 8 (long-term debt; future principal payments); 8-Ks of Aug-2022, May-2023, Aug-2024, Nov-2025 and May-2026 (notes issues).',
    'cal_rguid': 'Source: Q2-26 release (8-K Ex. 99.1, 29-Jul-2026), "CFO Outlook Commentary" — SEC EDGAR; no guidance from outside EDGAR is used.',
}
CELL_COMMENTS = {}
def comments_setup():
    CELL_COMMENTS['is_mm'] = {A[2011]: 'FY2011 as originally reported in the prospectus ($427m); the FY2012 release shows $393m after a reclassification with G&A (not recast).'}
    CELL_COMMENTS['is_ga'] = {A[2011]: 'FY2011 as originally reported in the prospectus ($280m); $314m in the FY2012 release (reclassification).'}
    CELL_COMMENTS['is_tax'] = {A[2016]: 'FY2016 from the Q4-16 release, which early-adopted ASU 2016-09 from 1-Jan-2016 ($934m lower provision); Q1–Q3 2016 are as originally reported, so Σ quarters ≠ year.',
                               QC[(2025, 3)].c: 'Includes a one-time, non-cash charge of $15.93bn (valuation allowance on U.S. federal deferred tax assets; OBBBA / CAMT).',
                               QC[(2026, 1)].c: 'Includes an $8.03bn benefit (U.S. Treasury Notice 2026-7), a partial reversal of the Q3-25 charge.'}
    CELL_COMMENTS['geo_usc'] = {QC[(y, 4)].c: 'Q4 = fiscal year (10-K) − nine months (10-Qs), derived.' for y in range(2018, 2026)}
    CELL_COMMENTS['o_pub'] = {A[2014]: 'FY2014 non-GAAP figure applies the amortization exclusion (introduced Q3-2014) to the full year, so Σ quarters ≠ year (−$82m).'}
    CELL_COMMENTS['cf_cfo'] = {A[2018]: 'FY2018 release reclassified $65m between operating cash flow and capex vs Q1–Q3 2018 as first reported.'}
    te = {}
    for y in range(2017, 2027):
        for q in (1, 2, 3, 4):
            p = f'Q{q}-{str(y)[2:]}'; x = DATA.get(p, {}).get('x', {})
            if (x.get('oo_legal') or x.get('oo_restr')) and x.get('oo_taxeff') is None and x.get('oo_eps_pub') is None and (y, q) in QC:
                te[QC[(y, q)].c] = 'Tax effect imputed at the period effective tax rate bounded 15–40% (the company stated no "excluding" EPS for this period).'
            if x.get('oo_taxeff') == 0 and (y, q) in QC:
                te[QC[(y, q)].c] = 'Company: the FTC expense is not expected to be tax-deductible (no tax effect).'
    for y in (2019, 2022, 2023, 2024):
        te[A[y]] = 'Year: imputed / company tax effect on the year\'s items (Σ quarters may differ by rounding of the per-quarter derivations).'
    CELL_COMMENTS['e_tax'] = te

DEFS = {}
def defs_setup():
    q = lambda y, n: QC[(y, n)].c
    DEFS['o_def'] = {
        'C': ('FY2009–11 (pre-IPO): no non-GAAP measures published — adjusted = GAAP (model definition)', None),
        A[2012]: ('IPO–2016: company non-GAAP income from operations excludes share-based compensation and related payroll tax',
                  'Facebook published non-GAAP costs, income from operations, net income and EPS from its first release (Q2-2012) to Q4-2016.'),
        q(2014, 3): ('From Q3-2014: amortization of intangible assets also excluded (WhatsApp, Oculus)', 'Q3-2014 release; the FY2014 figure applies it to the whole year.'),
        q(2017, 1): ('2017+: no company non-GAAP operating measure — model definition: GAAP + company-quantified legal and restructuring items (SBC expensed)',
                     'From Q1-2017 the company reports GAAP results plus FX-neutral revenue and free cash flow only.'),
        q(2019, 1): ('FTC accruals $3.0bn (Q1-19) and $2.0bn (Q2-19) — adjusted out (company-quantified)', None),
        q(2022, 3): ('Restructuring Q3-22–Q4-23 (facilities, layoffs, data centers) and Q2-23 legal $1.87bn — adjusted out', None),
        q(2024, 4): ('Q4-24 favorable $1.55bn decrease in legal accruals — adjusted out (negative item)', None),
        q(2026, 2): ('Q2-26 legal charges $2.40bn and severance $1.18bn — adjusted out', None),
    }
    DEFS['e_def'] = {
        'C': ('FY2009–11: no non-GAAP EPS published (pre-IPO two-class EPS) — model = GAAP', None),
        A[2012]: ('IPO–2016: company non-GAAP EPS (excl. SBC, payroll tax, amortization from Q3-14, and their tax effects; preferred conversion assumed pre-IPO)', None),
        q(2017, 1): ('2017+: model = GAAP + company-quantified one-offs and discrete tax items; checked against the company\'s stated "excluding" EPS', None),
        q(2017, 4): ('Q4-17 Tax Act charge $2.27bn excluded (company: EPS −$0.77)', None),
        q(2019, 2): ('Q2-19 Altera tax charge $1.11bn excluded (with the FTC accrual)', None),
        q(2020, 3): ('Q3-20 R&D tax-election benefit $913m excluded', None),
        q(2025, 3): ('Q3-25 OBBBA valuation-allowance charge $15.93bn excluded (company: EPS $7.25)', None),
        q(2026, 1): ('Q1-26 tax benefit $8.03bn excluded (company: EPS $3.13 lower)', None),
    }
    DEFS['fcf_def'] = {
        q(2013, 2): ('From Q2-2013: FCF = CFO − purchases of P&E − P&E acquired under capital leases', None),
        q(2015, 1): ('2015–2018: FCF = CFO − purchases of property and equipment', None),
        q(2019, 1): ('From Q1-2019: FCF = CFO − purchases of P&E − principal payments on finance leases', None),
        q(2022, 1): ('2022–23: purchases of P&E net of proceeds relating to P&E', None),
    }
def write_defs(ws, row):
    for c, (txt, cm) in DEFS[row.key].items():
        cell = ws[f'{c}{row.r}']; cell.value = txt
        cell.font = Font(name='Calibri', sz=10, b=True, i=True, color='FFC00000')
        if cm: comment(cell, cm)

def write_scen(ws):
    srcrow = {'hdr': 9, 'out_hdr': 11, 'pt_hdr': 15, 'out': 12, 'pt': 16, 'lev': 28, 'blk': 30, 'yr': 31}
    for spec_ in SCEN_ROWS:
        r, kind = spec_[0], spec_[1]
        for ci in range(CI(SL), CI(SNOTE) + 1):
            ws.cell(r, ci)._style = copy.copy(TM.cell(srcrow[kind], ci + TSHIFT)._style)
        if kind in ('hdr', 'out_hdr', 'pt_hdr'):
            _, _, a, b, c_, d_, e = spec_
            for col, v in zip([SL, SBULL, SBASE, SBEAR, SNOTE], [a, b, c_, d_, e]):
                if v is not None: ws[f'{col}{r}'] = v
        elif kind == 'out':
            _, _, lab, hi, mid, lo, note, nf = spec_
            ws[f'{SL}{r}'] = lab
            for col, v in zip([SBULL, SBASE, SBEAR], [hi, mid, lo]):
                ws[f'{col}{r}'] = v; ws[f'{col}{r}'].number_format = NF[nf]; set_color(ws[f'{col}{r}'], BLUE)
            ws[f'{SNOTE}{r}'] = note
        elif kind == 'pt':
            _, _, lab, _, val, _, note, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBASE}{r}'] = val; ws[f'{SBASE}{r}'].number_format = NF[nf]; set_color(ws[f'{SBASE}{r}'], BLUE); ws[f'{SNOTE}{r}'] = note
        elif kind == 'lev':
            _, _, lab, a, b, c_, note = spec_
            ws[f'{SL}{r}'] = lab
            for col, v in zip([SBULL, SBASE, SBEAR], [a, b, c_]):
                ws[f'{col}{r}'] = v; ws[f'{col}{r}'].number_format = NF['x2']; set_color(ws[f'{col}{r}'], BLUE)
            ws[f'{SNOTE}{r}'] = note
        elif kind == 'blk':
            _, _, lab, db, mid, de, note, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBULL}{r}'] = db; ws[f'{SBASE}{r}'] = mid; ws[f'{SBEAR}{r}'] = de; ws[f'{SNOTE}{r}'] = note
            for col in (SBULL, SBEAR):
                ws[f'{col}{r}'].number_format = '\\+#,##0;\\-#,##0' if nf == 'm' else '\\+0.0%;\\-0.0%'
                set_color(ws[f'{col}{r}'], BLUE)
        elif kind == 'yr':
            _, _, lab, hdr, base, clamp, _, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBASE}{r}'] = base; set_color(ws[f'{SBASE}{r}'], BLUE)
            if clamp:
                ws[f'{SBULL}{r}'] = f'=MAX(0,{SBASE}{r}+${SBULL}${hdr})'; ws[f'{SBEAR}{r}'] = f'=MAX(0,{SBASE}{r}+${SBEAR}${hdr})'
            else:
                ws[f'{SBULL}{r}'] = f'={SBASE}{r}+${SBULL}${hdr}'; ws[f'{SBEAR}{r}'] = f'={SBASE}{r}+${SBEAR}${hdr}'
            for col in (SBULL, SBASE, SBEAR): ws[f'{col}{r}'].number_format = NF[nf]

def write_header(ws):
    for r in range(1, 8):
        clone_row_style(ws, r, r, 2, CI(LASTC))
    ws[f'{NOTE}1'] = 'SCENARIO SWITCH ▼ (Bull / Base / Bear) — drives all scenario-linked forecast drivers'
    ws[f'{NOTE}2'] = 'Base'
    comment(ws[f'{NOTE}2'], 'Scenario toggle. Choose Bull / Base / Bear. Drives: Q3/26 revenue and FY26 expense guidance end-points (high / mid / low), '
            '2H/26 buybacks, 2027E–2031E ad impressions and price-per-ad growth, other revenue, Reality Labs losses, other operating costs, capex intensity, '
            f'M&A spend and the leverage ceiling for 2027E+ buybacks (scenario table {SL}:{SNOTE}).')
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type='list', formula1='"Bull,Base,Bear"', allow_blank=False); ws.add_data_validation(dv); dv.add(f'{NOTE}2')
    ws['B2'] = 'Meta Platforms, Inc. (NASDAQ: META) — 3-Statement Financial Model'
    ws['B3'] = ('US GAAP as reported · USD millions (per-share data in USD; shares and people in millions) · Fiscal years ended 31 December · '
                'Source: SEC EDGAR only (CIK 1326801): Forms 10-K / 10-Q, Form 8-K earnings releases (Ex. 99.1), IPO prospectus (424B4, 18-May-2012), notes-issue 8-Ks; '
                'XBRL company facts. Investor-relations releases are the same documents as filed on EDGAR; all guidance is from the 8-K releases (no unverified guidance used)')
    ws['B4'] = ('Basis notes: as originally reported (FY2011 cost reclassification, FY2016 ASU 2016-09 and 2018 restricted-cash re-presentations not recast). Facebook, Inc. '
                'until 28-Oct-2021 (ticker FB until 9-Jun-2022); IPO 18-May-2012 (FY2009–11 from the prospectus; no FY2009 balance sheet). Segments FoA / RL from Q4-21 '
                '(2019–21 recast). Acquisitions: Instagram 2012, WhatsApp and Oculus 2014. Q3/26E calibrated to revenue guidance, 2H/26E to FY26 expense guidance (29-Jul-2026).')
    ws[f'{CAGR[0]}5'] = 'Forecast'
    for c in CAGR[1:]: ws[f'{c}5'] = 'Historical'
    ws['B6'] = '(USDm)'
    for col in COLS: ws[f'{col.c}6'] = col.label
    ws[f'{NOTE}6'] = 'Modelling Notes'
    for c, a, b in zip(CAGR, ['5Y CAGR', '5Y CAGR', '10Y CAGR', '16Y CAGR'], ["26-'31", "20-'25", "15-'25", "09-'25"]):
        ws[f'{c}6'] = a; ws[f'{c}7'] = b

def set_columns(ws):
    for k in list(ws.column_dimensions.keys()): del ws.column_dimensions[k]
    ws.column_dimensions['A'].width = TM.column_dimensions['A'].width
    ws.column_dimensions['B'].width = TM.column_dimensions['B'].width
    wid = {'A': 13.0, 'Q': 9.5703125, 'QE': 11.42578125, 'E26': 14.140625, 'AE': 13.0}
    for col in COLS:
        cd = ws.column_dimensions[col.c]; cd.width = wid[col.kind]
        if col.q:
            cd.outlineLevel = 1; cd.hidden = True
    for ci in range(CI(NOTE), CI(LASTC) + 1):
        ws.column_dimensions[L(ci)].width = TM.column_dimensions[L(ci + TSHIFT)].width or 13.0
    ws.sheet_format.outlineLevelCol = 1

def main():
    import openpyxl
    wb = openpyxl.load_workbook(TEMPLATE)
    if 'DCF_old' in wb.sheetnames: del wb['DCF_old']
    ws = wb['Model']
    blank_style = copy.copy(ws['A1']._style)
    for rowc in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for c in rowc:
            c.value = None; c.comment = None; c._style = copy.copy(blank_style)
    for k in list(ws.row_dimensions.keys()):
        rd = ws.row_dimensions[k]; rd.outlineLevel = 0; rd.hidden = False
    ws.data_validations.dataValidation = []
    scen_layout(V0)
    build_spec(V0)
    n1 = len(fw.SPEC)
    build_spec2(V0)
    for i, row in enumerate(fw.SPEC): row.spec2 = i >= n1
    last = number_rows(8)
    for r in range(1, last + 1): ws.row_dimensions[r].height = 15.0
    ws.row_dimensions[2].height = 21.0
    defs_setup(); comments_setup()
    write_header(ws)
    for row in fw.SPEC: write_row(ws, row)
    write_scen(ws)
    set_columns(ws)
    ws.freeze_panes = 'C8'
    ws.sheet_view.zoomScale = 70
    ws.sheet_view.topLeftCell = 'C8'
    import other_sheets as osh
    osh.do_all(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    osh.inject_charts(OUT)
    print('rows', fw.SPEC[-1].r, 'saved', OUT)
    return wb

RECALC = os.environ.get('XLSX_RECALC', '/mnt/skills/public/xlsx/scripts/recalc.py')
def recalc(path, timeout=600):
    import subprocess
    out = subprocess.run(['python3', RECALC, path, str(timeout)], capture_output=True, text=True, cwd=os.path.dirname(RECALC))
    try:
        return json.loads(out.stdout)
    except Exception:
        return {'raw': out.stdout[-2000:], 'err': out.stderr[-2000:]}

def scenario_snapshots():
    import openpyxl
    import other_sheets as osh
    snaps = {}
    for scen in ('Bull', 'Bear'):
        p = os.path.join(BASE, f'META_{scen}.xlsx')
        wb = openpyxl.load_workbook(OUT); wb['Model'][f'{NOTE}2'] = scen; wb.save(p)
        res = recalc(p); print(scen, res.get('status'), res.get('total_errors'))
        v = openpyxl.load_workbook(p, data_only=True)['Bull-Base-Bear']
        snaps[scen] = {(r, c): v[f'{c}{r}'].value for r in range(9, 51) for c in 'CDEFGHIJKL' if v[f'{c}{r}'].value is not None}
        os.remove(p)
    wb = openpyxl.load_workbook(OUT)
    osh.snapshot_values(wb, snaps['Bull'], snaps['Bear'])
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    osh.inject_charts(OUT)

NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
def inject_values(src_xlsx, recalc_xlsx):
    """Store cached formula results (from a LibreOffice-recalculated copy) in the openpyxl file, so the workbook reads correctly in any
    viewer; fullCalcOnLoad stays set so Excel recalculates on open."""
    import zipfile, shutil, openpyxl
    from lxml import etree
    wbv = openpyxl.load_workbook(recalc_xlsx, data_only=True)
    zin = zipfile.ZipFile(src_xlsx)
    wbxml = etree.fromstring(zin.read('xl/workbook.xml'))
    rels = etree.fromstring(zin.read('xl/_rels/workbook.xml.rels'))
    rid2t = {r.get('Id'): r.get('Target') for r in rels}
    sheets = {}
    for s in wbxml.iter('{%s}sheet' % NS):
        rid = s.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
        sheets['xl/' + rid2t[rid].lstrip('/').replace('xl/', '')] = s.get('name')
    tmp = src_xlsx + '.tmp'
    zout = zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED)
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename in sheets:
            ws = wbv[sheets[item.filename]]
            root = etree.fromstring(data)
            for c in root.iter('{%s}c' % NS):
                f = c.find('{%s}f' % NS)
                if f is None: continue
                v = ws[c.get('r')].value
                old = c.find('{%s}v' % NS)
                if old is not None: c.remove(old)
                if v is None: continue
                ve = etree.SubElement(c, '{%s}v' % NS)
                if isinstance(v, bool):
                    c.set('t', 'b'); ve.text = '1' if v else '0'
                elif isinstance(v, (int, float)):
                    if 't' in c.attrib: del c.attrib['t']
                    ve.text = repr(float(v)) if isinstance(v, float) else str(v)
                elif isinstance(v, str) and v.startswith('#'):
                    c.set('t', 'e'); ve.text = v
                else:
                    c.set('t', 'str'); ve.text = str(v)
            data = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
        zout.writestr(item, data)
    zout.close(); zin.close()
    shutil.move(tmp, src_xlsx)

def store_values():
    import shutil
    p = os.path.join(BASE, 'META_recalc.xlsx')
    shutil.copy(OUT, p)
    res = recalc(p); print('recalc', res.get('status'), res.get('total_errors'))
    inject_values(OUT, p)
    os.remove(p)

if __name__ == '__main__':
    main()
    if '--snap' in sys.argv: scenario_snapshots()
    if '--recalc' in sys.argv: store_values()
