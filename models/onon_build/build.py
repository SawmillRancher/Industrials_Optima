"""Build ONON_Model.xlsx from the MLM template (same layout, formatting and process as the prior models)."""
import sys, copy, os, json
from openpyxl.styles import Font
from fw import *
import fw
from spec import build_spec, scen_layout, SCEN_ROWS, SC
from spec2 import build_spec2

OUT = os.path.join(os.path.dirname(BASE), 'ONON_Model.xlsx')

# ====================================================================== inputs (V0)
def wc_inputs():
    q = DATA['Q2-26']['x']; q1 = DATA['Q1-26']['x']
    rev, cogs = (q1['rev'] + q['rev']) * 2, (q1['cogs'] + q['cogs']) * 2
    return {'dso': round(q['bs_ar'] / rev * 365), 'dio': round(q['bs_inv'] / cogs * 365), 'dpo': round(q['bs_ap'] / cogs * 365),
            'oca': round(q['bs_oca'] / rev, 3), 'ocl': round(q['bs_ocl'] / rev, 3)}
WC = wc_inputs()
WC['note'] = (f'Forecast inputs at 30-Jun-2026 levels on annualised 1H/26 net sales / cost of sales (DSO {WC["dso"]}d, inventory {WC["dio"]}d, '
              f'payables {WC["dpo"]}d, other current assets {WC["oca"]:.1%} and other current operating liabilities {WC["ocl"]:.1%} of sales). '
              'Year-end receivables are seasonally lower than at 30-Jun (FY25 DSO 37d).')

V0 = {
    'out_hdr': 'FY26 outlook — Q2-26 release (11-Aug-26), reiterated at Investor Day (22-Sep-26)',
    'out_note': ('Q2-26: cc growth "low-20% range", CHF 3.47–3.56bn at spot, GM ≥65.0% (raised), Adj. EBITDA margin 19.5–20.0%; all excl. tariff refunds. '
                 'Investor Day: Q3-26 cc growth ~17%; tariff refund up to CHF 53m in Q3-26.'),
    'outlook': [
        ('out_ccg', '  FY26 net sales growth, constant currency', (0.23, 0.215, 0.20), 'pct', '"Low-20% range" → 20–23%. Drives the Q4-26E cc calibration.'),
        ('out_rev', '  FY26 net sales at current spot rates (CHF m)', (3560.0, 3515.0, 3470.0), 'm1', 'Guided CHF 3.47–3.56bn. Drives the 2H FX calibration.'),
        ('out_q3', '  Q3-26 net sales growth, constant currency', (0.18, 0.17, 0.16), 'pct', 'Investor Day: "around 17%". Drives the Q3-26E cc calibration.'),
        ('out_gm', '  FY26 gross margin (excl. tariff refund)', (0.658, 0.654, 0.650), 'pct', 'Guided "at least 65.0%" (1H/26 64.8%, Q2-26 65.4%).'),
        ('out_ebm', '  FY26 Adjusted EBITDA margin (excl. tariff refund)', (0.200, 0.1975, 0.195), 'pct', 'Guided 19.5–20.0% (1H/26 20.3%).'),
        ('out_tar', '  Q3-26 tariff refund in gross profit (CHF m)', (53.0, 53.0, 40.0), 'm1', 'Up to USD 65m / CHF 53m expected in Q3-26 (Investor Day); Bear assumes partial recovery.'),
        ('bb26', '  Q4-26 share repurchases (USD m)', (120.0, 77.0, 40.0), 'm', 'USD 1bn authorization (22-Sep-26 to Dec-29); Base = 1/13 of the programme per quarter.'),
    ],
    'points': [
        ('pt_dna', '  D&A (CHF m)', 150.0, 'm1', 'Not guided; 1H/26 CHF 71.9m (FY25 127.4m; ~62% right-of-use depreciation).'),
        ('pt_sbc', '  Share-based compensation (CHF m)', 66.0, 'm1', 'Not guided; 1H/26 CHF 33.5m (FY25 62.6m).'),
        ('pt_finc', '  Financial income (CHF m)', 36.0, 'm1', '1H/26 CHF 18.3m on ~CHF 1.1bn average cash.'),
        ('pt_fexp', '  Financial expenses (CHF m)', 34.0, 'm1', '1H/26 CHF 16.3m (mostly lease interest).'),
        ('pt_etr', '  Tax rate (Q3/Q4-26E)', 0.16, 'pct', '1H/26 13.9%, Q2-26 16.5%; FY24 13.4%.'),
        ('pt_capex', '  Capital expenditures (CHF m)', 100.0, 'm1', 'Not guided; 1H/26 CHF 47.2m (FY25 78.6m; ~2.8% of sales).'),
        ('pt_lease', '  Lease principal payments (CHF m)', 76.0, 'm1', '1H/26 CHF 35.8m (FY25 69.9m).'),
    ],
    'blocks': [
        ('AM_CC', 'AM_CC (Americas net sales growth, constant currency)', 0.03, -0.06, [0.14, 0.135, 0.13, 0.11],
         'Q2-26 +13.0% cc (wholesale sell-in managed); US premium running share gains, DTC / own stores, new categories (football, golf).', 'pct', False),
        ('EMEA_CC', 'EMEA_CC (EMEA net sales growth, constant currency)', 0.03, -0.06, [0.19, 0.18, 0.17, 0.14],
         '1H/26 +22.8% cc; under-penetrated markets (UK, France, Italy, Nordics) and apparel.', 'pct', False),
        ('APAC_CC', 'APAC_CC (Asia-Pacific net sales growth, constant currency)', 0.04, -0.10, [0.30, 0.27, 0.24, 0.19],
         '1H/26 +58.1% cc; Greater China, Japan, South Korea; >20% of group sales. Fades from very high growth.', 'pct', False),
        ('GM', 'GM (gross margin)', 0.005, -0.012, [0.655, 0.657, 0.66, 0.66],
         'FY26 ≥65.0%; Investor Day: "industry-leading gross profit margin of 65.0%+ throughout the period" (DTC mix, full-price discipline).', 'pct', False),
        ('EBM', 'EBM (Adjusted EBITDA margin)', 0.01, -0.025, [0.205, 0.212, 0.22, 0.225],
         'FY26 19.5–20.0%; Investor Day ambition ≥22% by 2029 from SG&A leverage at scale. Base = target path.', 'pct', False),
        ('BB', 'BB (share repurchases, USD m)', 100.0, -150.0, [308.0, 308.0, 307.0, 350.0],
         'Base: USD 1bn authorization spread evenly Q4-26–2029 (≈USD 308m p.a.); 2030E assumes a renewed programme. Bull +USD 100m p.a., Bear −USD 150m p.a.', 'm', True),
    ],
    'fx_ae': 0.0, 'dtc_ae': {2027: 0.465, 2028: 0.48, 2029: 0.495, 2030: 0.51},
    'app_ae': {2027: 0.075, 2028: 0.085, 2029: 0.095, 2030: 0.10}, 'acc_ae': {2027: 0.016, 2028: 0.017, 2029: 0.018, 2030: 0.019},
    'dna_pct': 0.043, 'sbc_pct': 0.019, 'etr': 0.17, 'e_rate': 0.0, 'capex_pct': 0.029, 'eq_ae': 5.0, 'auth': 1000.0,
    'tg_rev': 5600.0, 'wc': WC,
    'rou_share': 0.62, 'rou_note': 'FY25 20-F: right-of-use depreciation CHF 79.4m of CHF 127.4m D&A (62%); PP&E 38.2m; intangibles ~9.8m.',
    'rou_pct': 0.155, 'll_cur': 0.22, 'll_rate': 0.055, 'cash_yield': 0.03,
    'px': 32.00, 'usdchf': 0.80,
    'px_note': 'VALUATION — share price input USD 32.00 (NYSE: ONON, early Oct-2026 close ~USD 31.9) × USD/CHF 0.80; update both to market',
    'ccg_guid_note': 'Company FY26 constant-currency net sales growth "in the low-20% range" (Q2-26 release, reiterated 22-Sep-26). Drives the Q4-26E cc calibration.',
    'rev_guid_note': 'Company FY26 net sales CHF 3.47–3.56bn "at current spot rates" (Q2-26 release). Drives the 2H FX calibration.',
}

# ====================================================================== compute engine
def hist_value(row, col):
    kw = row.kw; path = kw.get('data')
    if not path: return None
    if kw.get('ppa'):
        v = MAN['ppa'].get(pkey(col), {}).get(path.split('_', 1)[1])
        return v if kw.get('textdata') else num(v)
    v = xget(pkey(col), path)
    return v if kw.get('textdata') else num(v)

def first(kw, keys):
    for k in keys:
        if k in kw: return k
    return None

def compute(row, col):
    kw = row.kw; k = col.kind
    CTX.col = col
    if kw.get('ann') and col.q: return None, False
    if k in HISTK:
        v = hist_value(row, col)
        if v is not None: return v, True
        for key in ('hist', 'all'):
            if key in kw:
                return kw[key](), False
        return None, False
    if k == 'QE':
        if row.spec2: return None, False
        if 'qe_in' in kw:
            v = kw['qe_in']; return (v[col.q] if isinstance(v, dict) else v), True
        key = first(kw, ('qe', 'fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    if k == 'E26':
        if 'e26_in' in kw: return kw['e26_in'], True
        if 'e26' in kw: return kw['e26'](), False
        if kw.get('avg') and ('qe' in kw or 'qe_in' in kw): return f'=AVERAGE({Q1}{row.r},{Q2}{row.r},{Q3}{row.r},{Q4}{row.r})', False
        if kw.get('flow') and not row.spec2 and any(x in kw for x in ('qe', 'qe_in', 'fc')): return S26(row.key), False
        key = first(kw, ('fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    if k == 'AE':
        if 'ae_in' in kw:
            v = kw['ae_in']; return (v[col.year] if isinstance(v, dict) else v), True
        key = first(kw, ('ae', 'fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    return None, False

# ====================================================================== writers
CAGR_G = {'is_rev', 'g_rev', 'am', 'emea', 'apac', 'ch_dtc', 'ch_whs', 'pr_shoes', 'pr_apparel', 'r_adj', 'g_adj', 'r_ebitda', 'is_op', 'is_dda',
          'o_adj', 'e_adjni', 'e_eps', 'cf_cfo', 'fcf', 'bs_ta', 'bs_teq', 'sh_dil', 'is_gp', 'is_sga', 'd_dna', 'd_sbc', 'is_ni'}
CAGR_A = {'m_gm', 'm_ebm', 'r_adj_m', 'is_gm', 'is_opm', 'o_adj_m', 'r_ebit_m', 'gm_rev', 'gm_cc', 'gm_gm', 'gm_ebitda_m', 'gm_op_m', 'gm_net_m', 'gm_sga',
          'gm_adjop_m', 'fcf_m', 'fcf_conv', 'capex_pct', 'ra_roic', 'rn_ronta', 'roe', 'ch_dtc_pct', 'm_opex', 'd_dna_pct', 'd_sbc_pct',
          'wc_nwc_p', 'is_etr', 'fcf_eb', 'g_ccg', 'g_fx', 'am_mix', 'emea_mix', 'apac_mix', 'pr_app_pct'}

def style_cagr(ws, r, on):
    src = 16 if on else 12
    for c in CAGR:
        ws[f'{c}{r}']._style = copy.copy(TM[f'{L(CI(c) + TSHIFT)}{src}']._style)

def write_row(ws, row):
    r = row.r; arch = ARCH[row.arch]
    clone_row_style(ws, arch, r)
    ws.cell(r, 2).value = row.label or None
    if row.kw.get('note'): ws[f'{NOTE}{r}'] = row.kw['note']
    if row.kw.get('group'):
        ws.row_dimensions[r].outlineLevel = 1; ws.row_dimensions[r].hidden = True
    style_cagr(ws, r, row.key in CAGR_G or row.key in CAGR_A)
    if row.arch in ('section', 'block', 'blank', 'sub') and not row.kw.get('data'):
        return
    if row.key in DEFS:
        write_defs(ws, row); return
    nf = NF[row.nf]
    for col in COLS:
        v, inp = compute(row, col)
        if v is None: continue
        cell = ws[f'{col.c}{r}']
        cell.value = v
        if not row.kw.get('textdata'): cell.number_format = nf
        if inp:
            set_color(cell, BLUE, bold=True if col.kind in FCK else None)
        elif row.arch == 'growth' and col.kind in HISTK:
            set_color(cell, GRAY)
        else:
            set_color(cell, BLACK)
    a = {y: ANN[y].c for y in ANN}
    if row.key in CAGR_G:
        ws[f'AO{r}'] = f'=IFERROR(({a[2030]}{r}/{a[2025]}{r})^(1/5)-1,"n/m")'
        ws[f'AP{r}'] = f'=IFERROR(({a[2025]}{r}/{a[2020]}{r})^(1/5)-1,"n/m")'
        ws[f'AQ{r}'] = f'=IFERROR(({a[2025]}{r}/{a[2022]}{r})^(1/3)-1,"n/m")'
        ws[f'AR{r}'] = f'=IFERROR(({a[2025]}{r}/{a[2019]}{r})^(1/6)-1,"n/m")'
        for c in CAGR: ws[f'{c}{r}'].number_format = '0%;\\(0%\\);"n/m"'
    elif row.key in CAGR_A:
        ws[f'AO{r}'] = f'=IFERROR(AVERAGE({",".join(f"{a[y]}{r}" for y in range(2026, 2031))}),"n/m")'
        ws[f'AP{r}'] = f'=IFERROR(AVERAGE({",".join(f"{a[y]}{r}" for y in range(2021, 2026))}),"n/m")'
        ws[f'AQ{r}'] = f'=IFERROR(AVERAGE({",".join(f"{a[y]}{r}" for y in range(2023, 2026))}),"n/m")'
        ws[f'AR{r}'] = f'=IFERROR(AVERAGE({",".join(f"{a[y]}{r}" for y in range(2019, 2026))}),"n/m")'
        for c in CAGR: ws[f'{c}{r}'].number_format = NF['pct']
    t = LABEL_SOURCES.get(row.key)
    if t: comment(ws.cell(r, 2), t)

SRC = 'Source: On Holding earnings releases (Form 6-K Ex. 99.1 / 99.3) — '
LABEL_SOURCES = {
    'am': ('Source: release "Net sales by geography" tables (from Q1-24) and 6-K MD&A exhibits (Ex. 99.2; Q1-23 – Q3-23 with restated 2022 comparatives); '
           'FY2021–23 from Forms 20-F. Constant-currency growth as published from Q1-24.'),
    'ch_whs': 'Source: release / MD&A "Net sales by sales channel"; FY2019–20 IPO prospectus (424B4); FY2021–23 Forms 20-F. DTC = net sales − wholesale.',
    'pr_shoes': 'Source: release / MD&A "Net sales by product"; FY2019–20 prospectus; FY2021–23 Forms 20-F.',
    'ro_eu': 'Source: IPO prospectus, 2021–22 6-K MD&A exhibits and Forms 20-F FY2021–22 (former regional definition).',
    'is_rev': ('Source: release consolidated statements of income (Q3-21+; 2021 quarters from the 2022 releases\' comparatives); FY2019–20 audited statements '
               'in the IPO prospectus (424B4, 16-Sep-2021; CHF thousands → millions). Q4 as printed in the full-year releases.'),
    'd_dna': SRC + '"Adjusted EBITDA and adjusted EBITDA margin" reconciliation table.',
    'r_adj_pub': SRC + 'Adjusted EBITDA reconciliation (quarter and full year); FY2019–20 prospectus reconciliation.',
    'e_sbc': SRC + '"Adjusted net income, adjusted basic EPS and adjusted diluted EPS" tables (Class A + Class B columns summed).',
    'e_pub': SRC + 'adjusted diluted EPS, Class A column. FY2019–20 prospectus (÷1,250 for the 2021 share-capital reorganisation).',
    'sh_basic': SRC + 'weighted shares in the adjusted net income tables (Class A + Class B ÷ 10).',
    'cf_ni': 'Source: release cash-flow statements (year-to-date; discrete quarters derived: Q2 = H1 − Q1, Q3 = 9M − H1, Q4 = FY − 9M). FY2020 from the FY2021 release comparatives.',
    'bs_cash': 'Source: release balance sheets (quarter ends from Q3-21); FY2020 from the FY2021 release comparatives.',
}

DEFS = {}
def defs_setup():
    q = lambda y, n: QC[(y, n)].c
    DEFS['r_def'] = {
        'C': ('FY2019–20: prospectus reconciliation (NI + taxes − financial income + financial expenses − FX result + D&A + SBC)', None),
        q(2021, 2): ('2021 – Q4-21: + IPO / equity transaction costs (CHF 7.6m in FY2021)', 'Prospectus: "IPO transaction cost"; 2021 releases: "Equity transaction costs".'),
        q(2022, 1): ('2022+: NI + taxes − financial income + financial expenses − FX result + D&A + SBC (unchanged since)', None),
        q(2026, 3): ('Q3-26E: tariff refund (CHF ~53m) in reported gross profit and Adjusted EBITDA; FY26 guidance excludes it', None),
    }
    DEFS['e_def'] = {
        'C': ('FY2019–20: prospectus (Class A only; ×1,250 share reorganisation); + SBC + tax effect', None),
        q(2021, 1): ('2021: + SBC + equity transaction costs + tax effect; Class B shares from the Sep-2021 IPO reorganisation', None),
        q(2022, 1): ('2022+: + SBC + tax effect on the deductible portion; Class A + Class B ÷ Class A-equivalent shares', None),
    }

def write_defs(ws, row):
    for c, (txt, cm) in DEFS[row.key].items():
        cell = ws[f'{c}{row.r}']; cell.value = txt
        cell.font = Font(name='Calibri', sz=10, b=True, i=True, color='FFC00000')
        if cm: comment(cell, cm)

def write_scen(ws):
    srcrow = {'hdr': 9, 'out_hdr': 11, 'pt_hdr': 15, 'out': 12, 'pt': 16, 'lev': 28, 'blk': 30, 'yr': 31}
    for spec_ in SCEN_ROWS:
        r, kind = spec_[0], spec_[1]
        for ci in range(CI(SL), CI(SNOTE) + 1):
            ws.cell(r, ci)._style = copy.copy(TM.cell(srcrow[kind], ci + TSHIFT)._style)
        if kind in ('hdr', 'out_hdr', 'pt_hdr'):
            _, _, a, b, c_, d_, e = spec_
            for col, v in zip([SL, SBULL, SBASE, SBEAR, SNOTE], [a, b, c_, d_, e]):
                if v is not None: ws[f'{col}{r}'] = v
        elif kind == 'out':
            _, _, lab, hi, mid, lo, note, nf = spec_
            ws[f'{SL}{r}'] = lab
            for col, v in zip([SBULL, SBASE, SBEAR], [hi, mid, lo]):
                ws[f'{col}{r}'] = v; ws[f'{col}{r}'].number_format = NF[nf]; set_color(ws[f'{col}{r}'], BLUE)
            ws[f'{SNOTE}{r}'] = note
        elif kind == 'pt':
            _, _, lab, _, val, _, note, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBASE}{r}'] = val; ws[f'{SBASE}{r}'].number_format = NF[nf]; set_color(ws[f'{SBASE}{r}'], BLUE); ws[f'{SNOTE}{r}'] = note
        elif kind == 'lev':
            _, _, lab, a, b, c_, note = spec_
            ws[f'{SL}{r}'] = lab
            for col, v in zip([SBULL, SBASE, SBEAR], [a, b, c_]):
                ws[f'{col}{r}'] = v; ws[f'{col}{r}'].number_format = NF['x']
            ws[f'{SNOTE}{r}'] = note
        elif kind == 'blk':
            _, _, lab, db, mid, de, note, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBULL}{r}'] = db; ws[f'{SBASE}{r}'] = mid; ws[f'{SBEAR}{r}'] = de; ws[f'{SNOTE}{r}'] = note
            for col in (SBULL, SBEAR):
                ws[f'{col}{r}'].number_format = '\\+#,##0;\\-#,##0' if nf == 'm' else '\\+0.0%;\\-0.0%'
        elif kind == 'yr':
            _, _, lab, hdr, base, clamp, _, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBASE}{r}'] = base
            if clamp:
                ws[f'{SBULL}{r}'] = f'=MAX(0,{SBASE}{r}+${SBULL}${hdr})'; ws[f'{SBEAR}{r}'] = f'=MAX(0,{SBASE}{r}+${SBEAR}${hdr})'
            else:
                ws[f'{SBULL}{r}'] = f'={SBASE}{r}+${SBULL}${hdr}'; ws[f'{SBEAR}{r}'] = f'={SBASE}{r}+${SBEAR}${hdr}'
            for col in (SBULL, SBASE, SBEAR): ws[f'{col}{r}'].number_format = NF[nf]

def write_header(ws):
    for r in range(1, 8):
        clone_row_style(ws, r, r, 2, CI(LASTC))
    ws[f'{NOTE}1'] = 'SCENARIO SWITCH ▼ (Bull / Base / Bear) — drives all scenario-linked forecast drivers'
    ws[f'{NOTE}2'] = 'Base'
    comment(ws[f'{NOTE}2'], 'Scenario toggle. Choose Bull / Base / Bear. Drives: FY26 outlook end-points (cc growth, CHF net sales, Q3 cc guide, gross margin, '
            'Adjusted EBITDA margin; high / mid / low), tariff refund, Q4-26 buybacks, and 2027E–2030E regional cc growth, gross margin, '
            f'Adjusted EBITDA margin and buybacks (scenario table {SL}:{SNOTE}).')
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type='list', formula1='"Bull,Base,Bear"', allow_blank=False); ws.add_data_validation(dv); dv.add(f'{NOTE}2')
    ws['B2'] = 'On Holding AG (NYSE: ONON) — 3-Statement Financial Model'
    ws['B3'] = ('IFRS as reported · CHF millions (per-share data in CHF; shares in millions, Class A-equivalent) · Source: On Holding Forms 20-F, Form 6-K earnings releases '
                '(Ex. 99.1 / 99.3) and MD&A exhibits (Ex. 99.2), Investor Day releases (Oct-2023, 22-Sep-2026) and the IPO prospectus (424B4, 16-Sep-2021), '
                'SEC EDGAR CIK 0001858985 (investor-relations releases at investors.on-running.com are the same documents as filed on EDGAR)')
    ws['B4'] = ('Basis notes: calendar fiscal year; single segment (sales disclosed by region, channel and product). FY2019–20 from the IPO prospectus; quarters from Q1/21 '
                '(Q1–Q2/21 from the 2022 releases\' comparatives). Regions: EMEA / Americas / APAC from 2023 (FY2021 and 2022 quarters restated); former Europe / '
                'North America / APAC / RoW as memo. Class A-equivalent shares = Class A + Class B ÷ 10. Q3/26E–Q4/26E calibrated to FY26 guidance (excl. tariff refund) '
                'and the Q3 cc guide; 2027E–2030E Base on the Investor Day 2029 targets. Valuation in CHF from the USD share price × USD/CHF.')
    ws[f'{CAGR[0]}5'] = 'Forecast'
    for c in CAGR[1:]: ws[f'{c}5'] = 'Historical'
    ws['B6'] = '(CHFm)'
    for col in COLS: ws[f'{col.c}6'] = col.label
    ws[f'{NOTE}6'] = 'Modelling Notes'
    for c, a, b in zip(CAGR, ['5Y CAGR', '5Y CAGR', '3Y CAGR', '6Y CAGR'], ["25-'30", "20-'25", "22-'25", "19-'25"]):
        ws[f'{c}6'] = a; ws[f'{c}7'] = b

def set_columns(ws):
    from openpyxl.worksheet.dimensions import ColumnDimension
    for k in list(ws.column_dimensions.keys()): del ws.column_dimensions[k]
    ws.column_dimensions['A'].width = TM.column_dimensions['A'].width
    ws.column_dimensions['B'].width = TM.column_dimensions['B'].width
    wid = {'A': 13.0, 'Q': 9.5703125, 'QE': 11.42578125, 'E26': 14.140625, 'AE': 13.0}
    for col in COLS:
        cd = ws.column_dimensions[col.c]; cd.width = wid[col.kind]
        if col.q:
            cd.outlineLevel = 1; cd.hidden = True
    for ci in range(CI(NOTE), CI(LASTC) + 1):
        ws.column_dimensions[L(ci)].width = TM.column_dimensions[L(ci + TSHIFT)].width or 13.0
    ws.sheet_format.outlineLevelCol = 1

def main():
    import openpyxl
    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb['Model']
    blank_style = copy.copy(ws['A1']._style)
    for rowc in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for c in rowc:
            c.value = None; c.comment = None; c._style = copy.copy(blank_style)
    for k in list(ws.row_dimensions.keys()):
        rd = ws.row_dimensions[k]; rd.outlineLevel = 0; rd.hidden = False
    ws.data_validations.dataValidation = []
    scen_layout(V0)
    build_spec(V0)
    n1 = len(fw.SPEC)
    build_spec2(V0)
    for i, row in enumerate(fw.SPEC): row.spec2 = i >= n1
    last = number_rows(8)
    for r in range(1, last + 1): ws.row_dimensions[r].height = 15.0
    ws.row_dimensions[2].height = 21.0
    defs_setup()
    write_header(ws)
    for row in fw.SPEC: write_row(ws, row)
    write_scen(ws)
    set_columns(ws)
    ws.freeze_panes = 'C8'
    import other_sheets as osh
    osh.do_all(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    osh.inject_charts(OUT)
    print('rows', fw.SPEC[-1].r, 'saved', OUT)
    return wb

RECALC = os.environ.get('XLSX_RECALC', '/mnt/skills/public/xlsx/scripts/recalc.py')
def recalc(path, timeout=600):
    import subprocess
    out = subprocess.run(['python3', RECALC, path, str(timeout)], capture_output=True, text=True, cwd=os.path.dirname(RECALC))
    try:
        return json.loads(out.stdout)
    except Exception:
        return {'raw': out.stdout[-2000:], 'err': out.stderr[-2000:]}

def scenario_snapshots():
    import openpyxl
    import other_sheets as osh
    snaps = {}
    for scen in ('Bull', 'Bear'):
        p = os.path.join(BASE, f'ONON_{scen}.xlsx')
        wb = openpyxl.load_workbook(OUT); wb['Model'][f'{NOTE}2'] = scen; wb.save(p)
        res = recalc(p); print(scen, res.get('status'), res.get('total_errors'))
        v = openpyxl.load_workbook(p, data_only=True)['Bull-Base-Bear']
        snaps[scen] = {(r, c): v[f'{c}{r}'].value for r in range(9, 51) for c in 'CDEFGHIJKL' if v[f'{c}{r}'].value is not None}
        os.remove(p)
    wb = openpyxl.load_workbook(OUT)
    osh.snapshot_values(wb, snaps['Bull'], snaps['Bear'])
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    osh.inject_charts(OUT)

if __name__ == '__main__':
    main()
    if '--snap' in sys.argv: scenario_snapshots()
