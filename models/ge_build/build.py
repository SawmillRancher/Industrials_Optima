"""Build GE_Model.xlsx (General Electric Company, operating as GE Aerospace) from the MLM template (same layout, formatting and
process as the prior models)."""
import sys, copy, os, json
from openpyxl.styles import Font
from fw import *
import fw
from spec import build_spec, scen_layout, SCEN_ROWS, SC
from spec2 import build_spec2

OUT = os.path.join(os.path.dirname(BASE), 'GE_Model.xlsx')
G26 = MAN['guidance_fy26']

# ====================================================================== inputs (V0)
def wc_inputs():
    q = DATA['2025']['x']
    rev, cogs = q['adj_rev'], q['cogs']
    return {'dso': round(q['bs_ar'] / rev * 365), 'dio': round(q['bs_inv'] / cogs * 365), 'dpo': round(q['bs_ap'] / cogs * 365),
            'ca': round(q['bs_ctr_a'] / rev, 3), 'adv': round((q['bs_prog'] + q['bs_ctr_l']) / rev, 3), 'sda': round(q['bs_sda'] / rev, 3),
            'ocl': round(q['bs_ocl'] / rev, 3)}
WC = wc_inputs()
WC['note'] = (f'Forecast inputs at the 31-Dec-2025 levels on FY2025 adjusted revenue / cost of equipment and services sold (DSO {WC["dso"]}d, inventory {WC["dio"]}d, '
              f'payables {WC["dpo"]}d, contract assets {WC["ca"]:.1%}, customer advances {WC["adv"]:.1%}, SD&A {WC["sda"]:.1%}, other current liabilities {WC["ocl"]:.1%} '
              'of adjusted revenue); year ends are compared with year ends.')
AM = MAN['amort_sched']
YRS = range(FYC + 1, FYL + 1)

V0 = {
    'out_hdr': 'FY26 guidance — Q2-26 release (16-Jul-2026)',
    'out_note': ('Adjusted revenue growth high-teens (prior LDD); operating profit $10.55–10.75B; adjusted EPS $7.65–7.85; FCF $8.9–9.2B. CES revenue ~+20% '
                 '(services low 20s, equipment ~20%), CES operating profit $10.25–10.35B; DPT revenue LDD, operating profit $1.6–1.7B; Corporate cost & eliminations '
                 '$(1.2)–(1.3)B (22-Jan-2026). Numeric mapping of qualitative growth guides (high-teens, ~20%, low 20s, LDD) = assumption.'),
    'outlook': [
        ('out_g', '  FY26 adjusted revenue growth (high-teens → 19 / 18 / 17%)', tuple(G26['adj_rev_growth']), 'pct',
         'Guided "high-teens" (raised from LDD, 16-Jul-2026); numeric mapping is an assumption. Drives the revenue calibration.'),
        ('out_op', '  FY26 operating profit ($m)', tuple(G26['op']), 'm', 'Guided $10.55–10.75B. Drives the 2H segment-margin calibration.'),
        ('out_eps', '  FY26 adjusted EPS ($) — memo', tuple(G26['adj_eps']), 'ps', 'Guided $7.65–7.85 (information row; compare the model).'),
        ('out_fcf', '  FY26 free cash flow ($m)', tuple(G26['fcf']), 'm', 'Guided $8.9–9.2B. Drives the 2026E customer-advances calibration (FCF bridge section).'),
        ('out_ces_sv', '  FY26 CES services revenue growth (low 20s → 23 / 22 / 21%)', tuple(G26['ces_sv_growth']), 'pct', 'Guided "low 20s" on the FY2025 revised base $24.2B (assumption mapping).'),
        ('out_ces_eq', '  FY26 CES equipment revenue growth (~20% → 21 / 20 / 19%)', tuple(G26['ces_eq_growth']), 'pct', 'Guided "~20%" on the FY2025 revised base $7.7B (assumption mapping).'),
        ('out_dpt', '  FY26 DPT revenue growth (LDD → 12 / 11 / 10%)', tuple(G26['dpt_growth']), 'pct', 'Guided "low double digits" on the FY2025 revised base $12.2B (assumption mapping).'),
        ('out_ces_op', '  FY26 CES operating profit ($m)', tuple(G26['ces_op']), 'm', 'Guided $10.25–10.35B.'),
        ('out_dpt_op', '  FY26 DPT operating profit ($m)', tuple(G26['dpt_op']), 'm', 'Guided $1.6–1.7B.'),
        ('out_corp', '  FY26 Corporate cost & eliminations ($m)', tuple(G26['corp']), 'm', G26['corp_text']),
        ('bb26', '  2H-26E share repurchases ($m)', (5000.0, 4000.0, 3000.0), 'm',
         'H1-26 $4,527m (Q2 $2,012m at $292.17); ~$18.0bn left on the $20bn authorization (Dec-2025) at 30-Jun-2026. Scenario input (assumption).'),
    ],
    'points': [
        ('pt_int', '  Q3/Q4-26E interest and other financial charges ($m per quarter)', 222.0, 'm', 'H1-26 $444m (run-rate).'),
        ('pt_etr', '  Q3/Q4-26E adjusted effective tax rate', 0.165, 'pct', 'H1-26 15.7% (Q2 16.7%); FY2025 17.3%. Not guided in the releases — assumption.'),
        ('pt_items_rate', '  Tax rate on adjusting items', MAN['tax']['items_rate'], 'pct', 'H1-26 company tax effects ~21% (non-operating benefit, separation, restructuring).'),
        ('pt_capex', '  2H-26E capex % of adjusted revenue', 0.028, 'pct', 'H1-26 2.7%; FY2025 3.0% (not guided).'),
        ('pt_taxnet', '  2H-26E provision less cash taxes paid ($m)', 400.0, 'm', 'H1-26 $517m; FY2025 $666m.'),
        ('pt_pens_cash', '  FY26 pension / OPEB cash contributions & payments ($m)', 350.0, 'm', 'H1-26: pension line −$526m incl. −$354m non-cash benefit income → ~$170m cash per half.'),
    ],
    'lev': (1.5, 1.0, 0.5),
    'lev_note': ('Net debt ~$9.8bn at 30-Jun-2026 (borrowings $19.2bn − cash $9.3bn) ≈ 0.8x adjusted EBITDA (model definition); ~1.4x pro forma for CPP ($11.75bn, 2H-27). '
                 'Target = leverage up to which surplus cash is returned via buybacks (assumption; GE publishes no leverage target in its filings).'),
    'blocks': [
        ('CES_SV', 'CES_SV (CES services revenue y/y)', 0.03, -0.05, [0.14, 0.11, 0.09, 0.08, 0.07],
         'FY26 guided low 20s; H1-26 +32% (shop-visit output, spare parts, price). LEAP shop-visit ramp and a young fleet sustain double-digit growth into 2027–28, normalising to high-single digits (assumption).', 'pct', False),
        ('CES_EQ', 'CES_EQ (CES equipment revenue y/y)', 0.04, -0.08, [0.12, 0.10, 0.08, 0.06, 0.05],
         'FY26 guided ~20% (LEAP deliveries +41% in H1-26); airframer rate increases (A320neo / 737 MAX), GE9X entry into service (assumption).', 'pct', False),
        ('DPT_G', 'DPT_G (DPT revenue y/y)', 0.03, -0.04, [0.08, 0.07, 0.06, 0.06, 0.05],
         'FY26 guided LDD; defense budgets, F110 / T408 / XA102 programs, Avio Aero; P&AT growth (assumption).', 'pct', False),
        ('CES_M', 'CES_M (CES operating profit margin)', 0.01, -0.02, [0.268, 0.272, 0.276, 0.280, 0.282],
         'FY2025 26.6% (as reported; ~27% revised), H1-26 26.9%, FY26 guide ~27%; services mix supports expansion while LEAP / GE9X install engines dilute (assumption).', 'pct', False),
        ('DPT_M', 'DPT_M (DPT operating profit margin)', 0.01, -0.015, [0.130, 0.135, 0.140, 0.145, 0.150],
         'FY2025 12.3% (as reported), H1-26 12.8%, FY26 guide ~12%; volume and price vs next-generation investment (assumption).', 'pct', False),
        ('CPP_SYN', 'CPP_SYN (CPP net-synergy realisation multiplier)', 0.25, -0.5, [1.0, 1.0, 1.0, 1.0, 1.0],
         'Deck: ~$200m net synergies by year 3, >2x by year 6. Multiplier on the phased run-rate (assumption).', 'x2', True),
        ('MNA', 'M&A_spend (bolt-ons beyond CPP, $m)', 500.0, -500.0, [0.0, 0.0, 500.0, 500.0, 500.0],
         'No acquisitions disclosed beyond CPP; Base assumes no bolt-ons during CPP integration and $0.5bn a year from 2029 (assumption).', 'm', True),
    ],
    'elim_p': -0.042, 'ins_g': -0.03, 'ins_m': 0.22, 'te_pre': -280.0, 'te_tax': 320.0,
    'cpp_cl': {y: (1 if y == 2027 else 0) for y in YRS}, 'cpp_share': {y: (0.25 if y == 2027 else 1.0) for y in YRS},
    'cpp_g': 0.08, 'cpp_ic': 0.25, 'cpp_m': 0.226, 'cpp_phase': {2027: 0.0, 2028: 0.25, 2029: 0.6, 2030: 1.0, 2031: 1.2},
    'cpp_dep': 0.05, 'cpp_intp': 0.35, 'cpp_life': 15, 'cpp_nwc': 0.05, 'cpp_ppe': 0.15,
    'mna_note': 'Scenario lever M&A_spend (bolt-ons beyond CPP). Funded from cash / revolver (debt schedule).',
    'corp_pct': -0.025,
    'dep_pct': 0.019,
    'amort26': AM['2026'],
    'amort_sched': {2027: AM['2027'], 2028: AM['2028'], 2029: AM['2029'], 2030: AM['2030'], 2031: 360.0},
    'amort_note': AM['src'] + ' 2031E $360m assumed.',
    'sga_pct': 0.090, 'rd_pct': 0.037, 'eq_share': 0.285, 'cogs_eq': 1.03, 'adjoi_pct': 0.026,
    'nonop': {2027: 680.0, 2028: 660.0, 2029: 640.0, 2030: 620.0, 2031: 600.0},
    'restr': {2027: 100.0, 2028: 250.0, 2029: 100.0, 2030: 100.0, 2031: 100.0},
    'sep': {2027: 50.0, 2028: 0.0, 2029: 0.0, 2030: 0.0, 2031: 0.0},
    'etr': 0.17,
    'dps_q': 0.47, 'dps_g': 0.15,
    'capex_pct': {2027: 0.030, 2028: 0.032, 2029: 0.031, 2030: 0.030, 2031: 0.030},
    'pens_cash': 350.0, 'taxnet': 400.0,
    'iss_h2': 2.0, 'iss_ae': 4.0, 'dil': 7.0,
    'iss_notes': {2026: 0.0, 2027: 1693.0, 2028: 480.0, 2029: 1639.0, 2030: 1700.0, 2031: 1500.0},
    'repay': {2026: 635.0, 2027: 1693.0, 2028: 480.0, 2029: 1639.0, 2030: 1700.0, 2031: 1500.0},
    'mat': {2026: 1693.0, 2027: 480.0, 2028: 1639.0, 2029: 1700.0, 2030: 1500.0, 2031: 1500.0},
    'repay_note': ('FY2025 10-K Note 10 maturities: 2026 $1,661m (H1-26 repayments $1,026m → $635m in 2H), 2027 $1,693m, 2028 $480m, 2029 $1,639m, 2030 $1,700m, '
                   'thereafter $13,296m (2031E $1.5bn assumed). Maturities 2027E+ assumed refinanced with new notes (input row above).'),
    'min_cash': 8000.0,
    'r_notes': {2027: 0.042, 2028: 0.043, 2029: 0.043, 2030: 0.044, 2031: 0.044},
    'r_notes_note': ('Senior notes average 4.00% and subordinated 4.40% at 31-Dec-2025 (10-K Note 10); 2025 issues at 4.3% / 4.9%; CPP acquisition debt assumed ~5.0%, '
                     'refinancing drifts the average up (assumption).'),
    'fees': 100.0,
    'wc': WC,
    'px': MAN['price']['px'],
    'px_note': ('VALUATION — share price input $308.00 provided by the user (12-Oct-2026; ~$320bn market cap); the Q2-26 10-Q average repurchase price was $292.17. '
                'Update to market.'),
    'rev_guid_note': 'Company FY26 adjusted revenue growth guidance "high-teens" (raised from LDD, 16-Jul-2026) on FY2025 adjusted revenue $42,322m. Drives the revenue calibration.',
    'op_guid_note': 'Company FY26 operating profit guidance $10.55–10.75B (raised from $9.85–10.25B, 16-Jul-2026). Drives the margin calibration.',
}

# ====================================================================== compute engine
def hist_value(row, col):
    kw = row.kw; path = kw.get('data')
    if not path: return None
    if kw.get('ppa'):
        v = MAN['ppa'].get(pkey(col), {}).get(path.split('_', 1)[1])
        return v if kw.get('textdata') else num(v)
    v = xget(pkey(col), path)
    return v if kw.get('textdata') else num(v)

def first(kw, keys):
    for k in keys:
        if k in kw: return k
    return None

def compute(row, col):
    kw = row.kw; k = col.kind
    CTX.col = col
    if kw.get('ann') and col.q: return None, False
    if k in HISTK:
        v = hist_value(row, col)
        if v is not None: return v, True
        for key in ('hist', 'all'):
            if key in kw:
                return kw[key](), False
        return None, False
    if k == 'QE':
        if row.spec2: return None, False
        if 'qe_in' in kw:
            v = kw['qe_in']; return (v[col.q] if isinstance(v, dict) else v), True
        key = first(kw, ('qe', 'fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    if k == 'E26':
        if 'e26_in' in kw: return kw['e26_in'], True
        if 'e26' in kw: return kw['e26'](), False
        if kw.get('avg') and ('qe' in kw or 'qe_in' in kw): return f'=AVERAGE({Q1}{row.r}:{QE}{row.r})', False
        if kw.get('flow') and not row.spec2 and any(x in kw for x in ('qe', 'qe_in', 'fc')): return S26(row.key), False
        key = first(kw, ('fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    if k == 'AE':
        if 'ae_in' in kw:
            v = kw['ae_in']; return (v[col.year] if isinstance(v, dict) else v), True
        key = first(kw, ('ae', 'fc', 'all'))
        return (kw[key](), False) if key else (None, False)
    return None, False

# ====================================================================== writers
CAGR_G = {'g_rev', 'ces', 'ces_sv', 'ces_eq', 'dpt', 'ces_op', 'dpt_op', 'g_op', 'o_op', 'r_adj', 'r_ebitda', 'is_ebt', 'd_dda', 'e_adjni', 'e_eps', 'cf_cfo',
          'fcf', 'bs_ta', 'bs_teq', 'sh_dil', 'is_rev', 'is_ni_cc', 'eps_dil', 'dps', 'ins_rev', 'd_amort', 'd_dep'}
CAGR_A = {'ces_m', 'dpt_m', 'g_opm', 'o_opm', 'r_adj_m', 'is_gm', 'is_ebt_m', 'r_ebit_m', 'gm_rev', 'gm_ces', 'gm_dpt', 'gm_opm', 'gm_ebitda_m',
          'gm_op_m', 'gm_net_m', 'gm_sga', 'gm_rd', 'fcf_m', 'fcf_conv', 'capex_pct', 'ra_roic', 'rn_ronta', 'roe', 'lv_x', 'd_sga', 'd_rd',
          'd_dep_pct', 'wc_nwc_p', 'is_etr', 'e_rate', 'corp_pct', 'elim_p', 'ins_m'}
A = {y: ANN[y].c for y in ANN}

def style_cagr(ws, r, on):
    src = 16 if on else 12
    for c in CAGR:
        ws[f'{c}{r}']._style = copy.copy(TM[f'{L(CI(c) + TSHIFT)}{src}']._style)

def write_row(ws, row):
    r = row.r; arch = ARCH[row.arch]
    clone_row_style(ws, arch, r)
    ws.cell(r, 2).value = row.label or None
    if row.kw.get('note'): ws[f'{NOTE}{r}'] = row.kw['note']
    style_cagr(ws, r, row.key in CAGR_G or row.key in CAGR_A)
    if row.arch in ('section', 'block', 'blank', 'sub') and not row.kw.get('data'):
        return
    if row.key in DEFS:
        write_defs(ws, row); return
    nf = NF[row.nf]
    for col in COLS:
        v, inp = compute(row, col)
        if v is None: continue
        cell = ws[f'{col.c}{r}']
        cell.value = v
        if not row.kw.get('textdata'): cell.number_format = nf
        if inp:
            set_color(cell, BLUE, bold=True if col.kind in FCK else None)
            lk = row.kw.get('labkey')
            if lk and col.kind in HISTK and isinstance(v, (int, float)) and abs(v) > 1e-9:
                t = xget(pkey(col), lk)
                if t: comment(cell, f'Company items ({pkey(col)}): {t}', 380, 110)
        elif row.arch == 'growth' and col.kind in HISTK:
            set_color(cell, GRAY)
        else:
            set_color(cell, BLACK)
    rc = CELL_COMMENTS.get(row.key, {})
    for c, t in rc.items(): comment(ws[f'{c}{r}'], t)
    if row.key in CAGR_G:
        ws[f'{CAGR[0]}{r}'] = f'=IFERROR(({A[2031]}{r}/{A[2026]}{r})^(1/5)-1,"n/m")'
        ws[f'{CAGR[1]}{r}'] = f'=IFERROR(({A[2025]}{r}/{A[2022]}{r})^(1/3)-1,"n/m")'
        ws[f'{CAGR[2]}{r}'] = f'=IFERROR(({A[2025]}{r}/{A[2023]}{r})^(1/2)-1,"n/m")'
        ws[f'{CAGR[3]}{r}'] = f'=IFERROR({A[2026]}{r}/{A[2025]}{r}-1,"n/m")'
        for c in CAGR: ws[f'{c}{r}'].number_format = '0%;\\(0%\\);"n/m"'
    elif row.key in CAGR_A:
        ws[f'{CAGR[0]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2027, 2032))}),"n/m")'
        ws[f'{CAGR[1]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2023, 2026))}),"n/m")'
        ws[f'{CAGR[2]}{r}'] = f'=IFERROR(AVERAGE({",".join(f"{A[y]}{r}" for y in range(2024, 2026))}),"n/m")'
        ws[f'{CAGR[3]}{r}'] = f'=IFERROR({A[2026]}{r}+0,"n/m")'
        for c in CAGR: ws[f'{c}{r}'].number_format = NF['pct']
    t = LABEL_SOURCES.get(row.key)
    if t: comment(ws.cell(r, 2), t)

REL = 'Source: GE Aerospace earnings releases (Form 8-K Item 2.02, Ex. 99.1) — '
QK = 'Source: General Electric Company Forms 10-Q / 10-K (SEC EDGAR, CIK 40545) — '
LABEL_SOURCES = {
    'ces_sv': (QK + 'MD&A segment tables (CES equipment / services revenue, segment profit). 2023 quarters and FY2022–23 from the recast comparatives in the 2024 10-Qs / 10-K '
               '(Q1-23 = H1-23 − Q2-23; Q4 = FY − 9M); Q1-24 from the Q1-25 10-Q. Jan-2026 recast (Aeroderivative CES → DPT, T&O to segments): Q1–Q2/25 from the 2026 10-Qs; '
               'FY2025 revised ($bn estimates) from the Q4-25 release guidance table.'),
    'dpt': QK + 'MD&A segment tables (DPT revenue by Defense & Systems / Propulsion & Additive Technologies and by equipment / services; segment profit).',
    'u_ce': QK + 'MD&A "Sales in units" (commercial engines incl. LEAP; internal shop visit revenue growth; defense engines). Q4 units = FY − 9M.',
    'ord_ces': REL + 'segment tables (orders, three months). Q1-23 = H1-23 − Q2-23 (Q2-24 release).',
    'elim': QK + 'MD&A Corporate & Other table ("Eliminations and other" revenue).',
    'ins_rev': QK + 'statements of operations (insurance revenue); insurance costs and expenses from the release operating-profit reconciliation.',
    'te_pre': REL + 'adjusted earnings / net income reconciliation (U.S. tax equity, pre-tax and tax effect; disclosed from the Q2-24 release with 2023 comparatives).',
    'corp': QK + 'MD&A Corporate & Other table: adjusted Corporate & Other operating costs (non-GAAP) = Corporate & Other profit (costs) + eliminations.',
    'cpp_hdr': 'Source: Form 8-K 8-Sep-2026 (Items 7.01 / 8.01): Ex. 99.1 investor presentation, Ex. 99.2 press release. Price $11.75bn; $7bn cash + new debt; ~18x / ~26x 2027 EBITDA with / without synergies; ~$2.0bn 2027E revenue; ~$200m net synergies by year 3.',
    'g_rev': REL + '"Adjusted revenue, operating profit and profit margin (non-GAAP)" reconciliation; checked against total revenue − insurance revenue.',
    'cal_rev_tgt': 'Source: Q2-26 release (8-K 16-Jul-2026) "GE Aerospace Full-Year 2026 Guidance" and segment guidance text; Q4-25 release (22-Jan-2026) for the FY2025 revised segment base and Corporate guide.',
    'd_dep': QK + 'statements of cash flows (depreciation and amortization of PP&E; amortization of intangible assets), continuing operations; quarters derived from year to date.',
    'd_adjoi': REL + 'operating-profit reconciliation ("Adjusted other income (loss)").',
    'd_restr': REL + 'operating-profit reconciliation (restructuring & other; separation costs; goodwill impairments; gains; noncontrolling interests).',
    'is_rev_eq': (QK + 'statements of operations (earnings), as first presented on the GE Aerospace basis: FY2022–23 and 2023 quarters from the 2024 10-Qs / 10-K (GE Vernova and '
                  'GE HealthCare in discontinued operations); Q1-24 from the Q1-25 10-Q; Q4 = FY − 9M.'),
    'is_ni_cc_pub': QK + 'EPS notes (net income from continuing operations attributable to common shareholders; average equivalent shares). Q4 shares = 4 × FY − (Q1 + Q2 + Q3); Q1-23 from the 11-Apr-2024 8-K.',
    'o_pub': REL + 'operating profit (non-GAAP) as published; 2023 quarters from the 2024 releases (recast), Q1-24 from the Q1-25 release.',
    'e_pub': REL + 'adjusted earnings / adjusted net income and adjusted EPS (non-GAAP); Q1-23 adjusted EPS $0.63 from the 11-Apr-2024 8-K (supplemental GE Aerospace information).',
    'f_pub': REL + 'free cash flow (non-GAAP) as published in each period\'s own release (definition in force; see the definition row).',
    'cf_ni': (QK + 'statements of cash flows (year to date; discrete quarters derived). Q1-23 / Q2-23 not presented on the GE Aerospace basis (H1-23 only), so 2023 quarters start in Q3. '
              'Working capital = sum of the operating working-capital lines; "all other operating" absorbs the rest.'),
    'bs_cash': QK + 'statements of financial position (31-Dec-2023 is the first GE Aerospace balance sheet; 31-Mar-2024 still consolidated GE Vernova and is not used).',
    'ds_total': 'Source: FY2025 10-K Note 10 (borrowings, maturities, average rates, $3.0bn revolver); Q2-26 10-Q Note 10; 8-K 8-Sep-2026 (CPP financing).',
    'sh_basic': QK + 'EPS notes (average equivalent shares, basic / diluted). No stock splits in the history window (1-for-8 reverse split Aug-2021).',
    'dps': 'Source: XBRL CommonStockDividendsPerShareDeclared (10-Q / 10-K): $0.08 a quarter 2023; $0.28 2024 (two declarations in Q2); $0.36 2025; $0.47 2026.',
    'v_px': MAN['price']['src'],
}
CELL_COMMENTS = {}
def comments_setup():
    q = lambda y, n: QC[(y, n)].c
    CELL_COMMENTS['is_rev_eq'] = {q(2023, 1): 'Q1-23 = H1-23 (Q2-24 10-Q comparative) − Q2-23; GE Aerospace basis (recast).',
                                  q(2023, 4): 'Q4-23 = FY2023 (FY2024 10-K) − 9M-23 (Q3-24 10-Q comparative).',
                                  q(2024, 1): 'Q1-24 from the Q1-25 10-Q comparative (the Q1-24 10-Q consolidated GE Vernova, spun 2-Apr-2024).'}
    CELL_COMMENTS['is_debt'] = {A[2022]: 'FY2022 debt extinguishment costs ($465m) — outside segment profit (FY2024 10-K).'}
    CELL_COMMENTS['d_gw'] = {q(2024, 3): 'Q3-24: $251m goodwill impairment, Colibrium Additive reporting unit (excluded from operating profit and adjusted EPS).'}
    CELL_COMMENTS['is_nci'] = {q(2024, 1): 'Q1-24: the Q1-25 10-Q statement shows total NCI of $27m incl. GE Vernova; continuing NCI ($3m) derived as NI from continuing operations − EPS-note numerator.'}
    CELL_COMMENTS['ces_sv_rc'] = {A[2025]: 'FY2025 revised for the Jan-2026 segment changes: CES services $24.2B, equipment $7.7B, CES $31.9B, DPT $12.2B, CES operating profit $8.6B, DPT $1.5B, Corporate $(1.0)B ($bn, unaudited estimates; Q4-25 release).'}
    CELL_COMMENTS['f_pub'] = {A[2025]: 'FY2025 $7,694m on the Q3-25 definition (incl. dispositions of PP&E); Q1-25 ($1,441m) and Q2-25 ($2,105m) as first published excluded them (re-presented later as $1,452m / $2,120m), so the quarters do not sum to the year.',
                              A[2024]: 'FY2024 as first published ($6,089m, Q4-24 release); re-presented as $6,203m in the Q3-25 release after dispositions of PP&E were added to the definition.'}
    CELL_COMMENTS['eps_pub'] = {q(2023, 1): 'Q1-23 continuing EPS $5.98 (revised, GE Aerospace basis) from the 11-Apr-2024 8-K Ex. 99.1; includes the $5.9bn gain on retained ownership interests (GE HealthCare, AerCap).'}
    CELL_COMMENTS['sh_dil'] = {q(2023, 4): 'Q4 shares = 4 × FY − (Q1 + Q2 + Q3) (not published).', q(2024, 4): 'Q4 shares = 4 × FY − (Q1 + Q2 + Q3).',
                               q(2025, 4): 'Q4 shares = 4 × FY − (Q1 + Q2 + Q3).'}

DEFS = {}
def defs_setup():
    q = lambda y, n: QC[(y, n)].c
    DEFS['o_def'] = {
        'C': ('FY2022: operating profit not published on the GE Aerospace basis (segment build only: segment profit + adjusted Corporate & Other costs, FY2024 10-K)', None),
        q(2023, 1): ('2023: GE Aerospace operating profit as re-presented in the 2024 releases (GE Vernova in discontinued operations); Q1-23 = H1-23 − Q2-23',
                     'The 2023 releases reported GE consolidated "adjusted profit" (incl. GE Vernova, before U.S. tax equity was carved out); the model uses the GE Aerospace comparatives published in 2024.'),
        q(2024, 1): ('Q1-24: from the Q1-25 release (the Q1-24 release reported GE consolidated incl. GE Vernova)', None),
        q(2024, 2): ('Q2-24+: operating profit = adjusted revenue − adjusted costs + adjusted other income (excl. insurance, U.S. tax equity, interest, non-operating benefit, restructuring, separation, gains, impairments, NCI)', None),
        q(2026, 1): ('Q1-26: Jan-2026 segment recast (Aeroderivative CES → DPT; T&O site revenue / cost to segments) — total operating profit unchanged',
                     'Recast Q1–Q2/25 comparatives in the 2026 10-Qs; FY2025 revised segment figures in the Q4-25 release guidance table (memo rows).'),
    }
    DEFS['e_def'] = {
        'C': ('FY2022: adjusted EPS not published on the GE Aerospace basis', None),
        q(2023, 1): ('2023: adjusted earnings excl. insurance, U.S. tax equity, non-operating benefit, gains, restructuring, separation, Russia & Ukraine, preferred excise / accretion (re-presented in the 2024 releases)',
                     'Q1-23 adjusted EPS $0.63 from the 11-Apr-2024 8-K (GE Aerospace revised); adjusted net income Q1-23 = H1-23 − Q2-23 (Q2-24 release).'),
        q(2024, 3): ('Q3-24: goodwill impairment ($251m, Colibrium Additive) excluded', None),
        q(2026, 1): ('Q1-26: renamed "adjusted net income" (definition unchanged)', None),
    }
    DEFS['f_def'] = {
        q(2023, 3): ('2023 (re-presented in the 2024 releases): CFOA − gross PP&E additions + separation and Corporate restructuring cash',
                     'The 2023 releases also excluded insurance CFOA and taxes related to business sales (GE consolidated); the 2024 comparatives do not.'),
        q(2025, 3): ('Q3-25+: + dispositions of PP&E (incl. sale-leasebacks); prior periods re-presented in that release (FY2024 $6,203m vs $6,089m)', None),
    }
def write_defs(ws, row):
    for c, (txt, cm) in DEFS[row.key].items():
        cell = ws[f'{c}{row.r}']; cell.value = txt
        cell.font = Font(name='Calibri', sz=10, b=True, i=True, color='FFC00000')
        if cm: comment(cell, cm)

def write_scen(ws):
    srcrow = {'hdr': 9, 'out_hdr': 11, 'pt_hdr': 15, 'out': 12, 'pt': 16, 'lev': 28, 'blk': 30, 'yr': 31}
    for spec_ in SCEN_ROWS:
        r, kind = spec_[0], spec_[1]
        for ci in range(CI(SL), CI(SNOTE) + 1):
            ws.cell(r, ci)._style = copy.copy(TM.cell(srcrow[kind], ci + TSHIFT)._style)
        if kind in ('hdr', 'out_hdr', 'pt_hdr'):
            _, _, a, b, c_, d_, e = spec_
            for col, v in zip([SL, SBULL, SBASE, SBEAR, SNOTE], [a, b, c_, d_, e]):
                if v is not None: ws[f'{col}{r}'] = v
        elif kind == 'out':
            _, _, lab, hi, mid, lo, note, nf = spec_
            ws[f'{SL}{r}'] = lab
            for col, v in zip([SBULL, SBASE, SBEAR], [hi, mid, lo]):
                ws[f'{col}{r}'] = v; ws[f'{col}{r}'].number_format = NF[nf]; set_color(ws[f'{col}{r}'], BLUE)
            ws[f'{SNOTE}{r}'] = note
        elif kind == 'pt':
            _, _, lab, _, val, _, note, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBASE}{r}'] = val; ws[f'{SBASE}{r}'].number_format = NF[nf]; set_color(ws[f'{SBASE}{r}'], BLUE); ws[f'{SNOTE}{r}'] = note
        elif kind == 'lev':
            _, _, lab, a, b, c_, note = spec_
            ws[f'{SL}{r}'] = lab
            for col, v in zip([SBULL, SBASE, SBEAR], [a, b, c_]):
                ws[f'{col}{r}'] = v; ws[f'{col}{r}'].number_format = NF['x']
            ws[f'{SNOTE}{r}'] = note
        elif kind == 'blk':
            _, _, lab, db, mid, de, note, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBULL}{r}'] = db; ws[f'{SBASE}{r}'] = mid; ws[f'{SBEAR}{r}'] = de; ws[f'{SNOTE}{r}'] = note
            for col in (SBULL, SBEAR):
                ws[f'{col}{r}'].number_format = '\\+#,##0;\\-#,##0' if nf == 'm' else '\\+0.0%;\\-0.0%'
        elif kind == 'yr':
            _, _, lab, hdr, base, clamp, _, nf = spec_
            ws[f'{SL}{r}'] = lab; ws[f'{SBASE}{r}'] = base
            if clamp:
                ws[f'{SBULL}{r}'] = f'=MAX(0,{SBASE}{r}+${SBULL}${hdr})'; ws[f'{SBEAR}{r}'] = f'=MAX(0,{SBASE}{r}+${SBEAR}${hdr})'
            else:
                ws[f'{SBULL}{r}'] = f'={SBASE}{r}+${SBULL}${hdr}'; ws[f'{SBEAR}{r}'] = f'={SBASE}{r}+${SBEAR}${hdr}'
            for col in (SBULL, SBASE, SBEAR): ws[f'{col}{r}'].number_format = NF[nf]

def write_header(ws):
    for r in range(1, 8):
        clone_row_style(ws, r, r, 2, CI(LASTC))
    ws[f'{NOTE}1'] = 'SCENARIO SWITCH ▼ (Bull / Base / Bear) — drives all scenario-linked forecast drivers'
    ws[f'{NOTE}2'] = 'Base'
    comment(ws[f'{NOTE}2'], 'Scenario toggle. Choose Bull / Base / Bear. Drives: FY26 guidance end-points (high / mid / low: adjusted revenue growth, operating profit, '
            'segment growth and operating-profit guides, Corporate cost), 2H-26E buybacks, 2027E–2031E CES services / equipment and DPT growth, segment margins, '
            f'CPP synergy realisation, M&A spend and the leverage target for 2027E+ buybacks (scenario table {SL}:{SNOTE}).')
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type='list', formula1='"Bull,Base,Bear"', allow_blank=False); ws.add_data_validation(dv); dv.add(f'{NOTE}2')
    ws['B2'] = 'General Electric Company, operating as GE Aerospace (NYSE: GE) — 3-Statement Financial Model'
    ws['B3'] = ('US GAAP · USD millions (per-share data in USD; shares in millions) · Calendar fiscal years · Source: General Electric Company Forms 10-K / 10-Q, '
                'Form 8-K earnings releases (Ex. 99.1), the 11-Apr-2024 supplemental GE Aerospace financial information (8-K Ex. 99.1) and the CPP deal 8-K (8-Sep-2026), '
                'SEC EDGAR (CIK 40545); XBRL company facts (cross-check, dividends per share); investor-relations releases are the same documents as filed on EDGAR')
    ws['B4'] = ('Basis notes: GE Aerospace continuing operations — GE HealthCare spun 3-Jan-2023 and GE Vernova spun 2-Apr-2024 (both discontinued operations). FY2022–23 and the 2023 '
                'quarters are the recast comparatives first presented on that basis (2024 10-Qs / 10-K / releases; Q1-23 = H1-23 − Q2-23, Q4 = FY − 9M); Q1-24 from the Q1-25 10-Q; '
                'Q2-24+ as originally reported. Balance sheet from 31-Dec-2023; cash flow from Q3-23. Segments CES / DPT (Jan-2026 recast: Aeroderivative to DPT; recast 2025 in memo rows). '
                'Run-off insurance and U.S. tax equity outside operating profit. Pending: CPP ($11.75bn, close 2H-27, toggle). Q3/Q4-26E calibrated to the 16-Jul-2026 FY26 guidance.')
    ws[f'{CAGR[0]}5'] = 'Forecast'
    for c in CAGR[1:]: ws[f'{c}5'] = 'Historical'
    ws['B6'] = '(USDm)'
    for col in COLS: ws[f'{col.c}6'] = col.label
    ws[f'{NOTE}6'] = 'Modelling Notes'
    for c, a, b in zip(CAGR, ['5Y CAGR', '3Y CAGR', '2Y CAGR', '1Y / level'], ["26-'31", "22-'25", "23-'25", "'25-'26E"]):
        ws[f'{c}6'] = a; ws[f'{c}7'] = b

def set_columns(ws):
    for k in list(ws.column_dimensions.keys()): del ws.column_dimensions[k]
    ws.column_dimensions['A'].width = TM.column_dimensions['A'].width
    ws.column_dimensions['B'].width = TM.column_dimensions['B'].width
    wid = {'A': 13.0, 'Q': 9.5703125, 'QE': 11.42578125, 'E26': 14.140625, 'AE': 13.0}
    for col in COLS:
        cd = ws.column_dimensions[col.c]; cd.width = wid[col.kind]
        if col.q:
            cd.outlineLevel = 1; cd.hidden = True
    for ci in range(CI(NOTE), CI(LASTC) + 1):
        ws.column_dimensions[L(ci)].width = TM.column_dimensions[L(ci + TSHIFT)].width or 13.0
    ws.sheet_format.outlineLevelCol = 1

def main():
    import openpyxl
    wb = openpyxl.load_workbook(TEMPLATE)
    if 'DCF_old' in wb.sheetnames: del wb['DCF_old']
    ws = wb['Model']
    blank_style = copy.copy(ws['A1']._style)
    for rowc in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for c in rowc:
            c.value = None; c.comment = None; c._style = copy.copy(blank_style)
    for k in list(ws.row_dimensions.keys()):
        rd = ws.row_dimensions[k]; rd.outlineLevel = 0; rd.hidden = False
    ws.data_validations.dataValidation = []
    scen_layout(V0)
    build_spec(V0)
    n1 = len(fw.SPEC)
    build_spec2(V0)
    for i, row in enumerate(fw.SPEC): row.spec2 = i >= n1
    last = number_rows(8)
    for r in range(1, last + 1): ws.row_dimensions[r].height = 15.0
    ws.row_dimensions[2].height = 21.0
    defs_setup(); comments_setup()
    write_header(ws)
    for row in fw.SPEC: write_row(ws, row)
    write_scen(ws)
    set_columns(ws)
    ws.freeze_panes = 'C8'
    ws.sheet_view.zoomScale = 70
    import other_sheets as osh
    osh.do_all(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    osh.inject_charts(OUT)
    print('rows', fw.SPEC[-1].r, 'saved', OUT)
    return wb

RECALC = os.environ.get('XLSX_RECALC', '/mnt/skills/public/xlsx/scripts/recalc.py')
def recalc(path, timeout=600):
    import subprocess
    out = subprocess.run(['python3', RECALC, path, str(timeout)], capture_output=True, text=True, cwd=os.path.dirname(RECALC))
    try:
        return json.loads(out.stdout)
    except Exception:
        return {'raw': out.stdout[-2000:], 'err': out.stderr[-2000:]}

def scenario_snapshots():
    import openpyxl
    import other_sheets as osh
    snaps = {}
    for scen in ('Bull', 'Bear'):
        p = os.path.join(BASE, f'GE_{scen}.xlsx')
        wb = openpyxl.load_workbook(OUT); wb['Model'][f'{NOTE}2'] = scen; wb.save(p)
        res = recalc(p); print(scen, res.get('status'), res.get('total_errors'))
        v = openpyxl.load_workbook(p, data_only=True)['Bull-Base-Bear']
        snaps[scen] = {(r, c): v[f'{c}{r}'].value for r in range(9, 51) for c in 'CDEFGHIJKL' if v[f'{c}{r}'].value is not None}
        os.remove(p)
    wb = openpyxl.load_workbook(OUT)
    osh.snapshot_values(wb, snaps['Bull'], snaps['Bear'])
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    osh.inject_charts(OUT)

NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
def inject_values(src_xlsx, recalc_xlsx):
    """Store cached formula results (from a LibreOffice-recalculated copy) in the openpyxl file, so the workbook reads correctly in any
    viewer; fullCalcOnLoad stays set so Excel recalculates on open. Same approach as the LECO / ESAB finishing step."""
    import zipfile, shutil, openpyxl
    from lxml import etree
    wbv = openpyxl.load_workbook(recalc_xlsx, data_only=True)
    zin = zipfile.ZipFile(src_xlsx)
    wbxml = etree.fromstring(zin.read('xl/workbook.xml'))
    rels = etree.fromstring(zin.read('xl/_rels/workbook.xml.rels'))
    rid2t = {r.get('Id'): r.get('Target') for r in rels}
    sheets = {}
    for s in wbxml.iter('{%s}sheet' % NS):
        rid = s.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
        sheets['xl/' + rid2t[rid].lstrip('/').replace('xl/', '')] = s.get('name')
    tmp = src_xlsx + '.tmp'
    zout = zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED)
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename in sheets:
            ws = wbv[sheets[item.filename]]
            root = etree.fromstring(data)
            for c in root.iter('{%s}c' % NS):
                f = c.find('{%s}f' % NS)
                if f is None: continue
                v = ws[c.get('r')].value
                old = c.find('{%s}v' % NS)
                if old is not None: c.remove(old)
                if v is None: continue
                ve = etree.SubElement(c, '{%s}v' % NS)
                if isinstance(v, bool):
                    c.set('t', 'b'); ve.text = '1' if v else '0'
                elif isinstance(v, (int, float)):
                    if 't' in c.attrib: del c.attrib['t']
                    ve.text = repr(float(v)) if isinstance(v, float) else str(v)
                elif isinstance(v, str) and v.startswith('#'):
                    c.set('t', 'e'); ve.text = v
                else:
                    c.set('t', 'str'); ve.text = str(v)
            data = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
        zout.writestr(item, data)
    zout.close(); zin.close()
    shutil.move(tmp, src_xlsx)

def store_values():
    import shutil
    p = os.path.join(BASE, 'GE_recalc.xlsx')
    shutil.copy(OUT, p)
    res = recalc(p); print('recalc', res.get('status'), res.get('total_errors'))
    inject_values(OUT, p)
    os.remove(p)

if __name__ == '__main__':
    main()
    if '--snap' in sys.argv: scenario_snapshots()
    if '--recalc' in sys.argv: store_values()
