"""Rewire the Bull-Base-Bear, DCF and Charts tabs of the MLM template (current layout: DCF with live valuation date, fade period,
implied ROIC, target-IRR entry price, reverse DCF and Mauboussin EV/NOPAT; Bull-Base-Bear with buyback rows) to the GE Aerospace Model
rows / columns; keep the template chart XML. The template's DCF_old tab is not carried.

Column mapping: the template's 2025A base column becomes GE FY2026E (Q1-Q2 actual, Q3-Q4 calibrated to guidance; the valuation date
falls in Q4-26) and the template's 2026E-2030E become GE FY2027E-FY2031E, so the DCF's explicit years start in FY2027E and net debt
is taken at 31-Dec-2026E (the FY2026E cash flow sits in that net debt)."""
import re, zipfile, shutil
from openpyxl.utils import get_column_letter as L
from fw import R, TEMPLATE, ANN, NOTE, MAN

# template Model row -> GE row key
MAP = {158: 'g_rev', 164: 'r_cash', 166: 'is_ebt', 182: 'd_dda', 186: 'sh_dil', 211: 'r_adj', 212: 'r_adj_m', 219: 'r_ebit', 221: 'r_adj_pf',
       229: 'e_adjni', 231: 'e_eps', 238: 'gm_ces', 245: 'gm_ebitda_m', 247: 'gm_op_m', 260: 'cf_wc', 262: 'cf_capex', 263: 'cf_acq',
       280: 'fcf', 282: 'fcf_m', 290: 'bb_px', 291: 'bb_cash', 292: 'bb_sh', 331: 'bs_nci', 408: 'ra_ic', 411: 'ra_roic', 422: 'rn_ronta',
       428: 'lv_nd', 433: 'v_px', 434: 'v_sh', 353: 'ra_ic', 355: 'ra_ic'}
CMAP = {'BV': ANN[2026].c, 'CA': ANN[2027].c, 'CB': ANN[2028].c, 'CC': ANN[2029].c, 'CD': ANN[2030].c, 'CE': ANN[2031].c, 'CF': NOTE}
REF = re.compile(r"Model!(\$?)([A-Z]{1,3})(\$?)(\d+)")
HIST = [ANN[y].c for y in range(2022, 2026)]

def remap(f):
    def rep(m):
        c, r = m.group(2), int(m.group(4))
        c2 = CMAP.get(c, c)
        r2 = R[MAP[r]] if r in MAP else r
        if c == 'CF' and r == 2: r2 = 2
        return f'Model!{m.group(1)}{c2}{m.group(3)}{r2}'
    return REF.sub(rep, f)

def hist_stat(f):
    """LT historical AVERAGE / MIN / MAX over the template's annual columns -> GE FY2022-FY2025 annual columns."""
    m = re.match(r'=IFERROR\((AVERAGE|MIN|MAX)\(Model!C(\d+),', f)
    if not m: return None
    fn, r = m.group(1), int(m.group(2))
    rr = R[MAP[r]]
    return f'=IFERROR({fn}({",".join(f"Model!{c}{rr}" for c in HIST)}),"")'

def fy_dates(f):
    """template discount dates (DATE(2026,12,31) for its 2026E) -> GE year mapping (+1 year; calendar year ends)."""
    f = re.sub(r'DATE\((\d{4}),12,31\)', lambda m: f'DATE({int(m.group(1)) + 1},12,31)', f)
    f = re.sub(r'DATE\((\d{4})-1,12,31\)', lambda m: f'DATE({int(m.group(1)) + 1}-1,12,31)', f)
    return f

PX = MAN['price']['px']
TEXT = {
    'Bull-Base-Bear': {
        'B2': 'General Electric Company, operating as GE Aerospace (GE) — Bull / Base / Bear Case P&L Summary',
        'B3': ('Company non-GAAP definitions (adjusted revenue, operating profit, adjusted EPS, FCF); adjusted EBITDA = operating profit + D&A (model) · USD millions · '
               'Live panel linked to Model — toggle scenario via Model!AC2. Bull and Bear panels are value snapshots generated from the Model with the switch set to each case. '
               'FY2026E = H1 actual + 2H calibrated to the FY26 guidance (adjusted revenue high-teens, operating profit $10.55–10.75B). CPP ($11.75bn) from 2H-27.'),
        'G26': 36, 'H26': 36, 'S26': 42, 'T26': 42, 'AD26': 28, 'AE26': 28, 'AF26': 28,
    },
    'DCF': {
        'B2': 'DCF VALUATION — GENERAL ELECTRIC COMPANY, OPERATING AS GE AEROSPACE (GE)',
        'C5': 0.0425, 'D5': 'Assumption (~US 10Y Treasury yield, Oct-2026; government rate sources not reachable from the build environment); update to market',
        'D6': 'Assumption: US implied ERP ~4.5–5.5%',
        'C7': 1.10, 'D7': 'Assumption: aero-engine and aerospace peers (RTX, Safran, Rolls-Royce, MTU, HEI, TDG) ~1.0–1.3; long-cycle services annuity vs airline-traffic and production-rate cyclicality',
        'C9': 0.05, 'D9': 'A2 / A- ratings (Q2-26 10-Q); 2025 senior notes at 4.3% (2030) and 4.9% (2036); CPP acquisition debt (2027E) assumed ~5%',
        'C10': 0.17, 'D10': 'Model 2027E+ adjusted effective tax rate (FY2025 adjusted 17.3%; H1-26 15.7%)',
        'C12': 0.05, 'D12': f'Net debt ~$10bn at 30-Jun-2026 (~$15bn pro forma for CPP) vs ~$320bn market cap (${PX:.2f} × ~1,043m diluted shares)',
        'D15': 'Periods = (31-Dec of the fiscal year − today) / 365.25; cash flows assumed at year end; FY2026E (ends 31-Dec-2026, ~3 months after the valuation date) is the base year',
        'B18': 'Operating profit (full-cost) less restructuring & separation costs',
        'B21': 'plus D&A (PP&E depreciation + amortization of intangibles, incl. CPP PPA)',
        'B22': '− Capex (PP&E) and acquisitions (CPP $11.75bn in 2027E; M&A lever)',
        'B26': 'Fade-period growth (2032E–36E; 9.0% → 5.0% linear)', 'I26': 0.09, 'M26': 0.05, 'C27': 0.035,
        'D32': ("GE Aerospace operating profit already deducts amortization of intangibles and stock-based compensation (neither is adjusted out); D&A is added back in row 21. "
                'Restructuring and separation costs (cash; excluded from operating profit) are deducted in row 18. Run-off insurance, U.S. tax equity and non-operating pension '
                'income are excluded from NOPAT, and the insurance balance sheet from net debt. Row 22 includes the $11.75bn CPP consideration in FY2027E (its operating profit '
                'is in NOPAT from closing) and bolt-ons from the M&A lever; the fade period and terminal value assume no further acquisitions.'),
        'D43': 'Opening IC = Model FY2026E book IC (equity + net debt at 31-Dec-2026E), rolled forward with capex + acquisitions − D&A + ΔWC.',
        'B48': 'Less: Net debt + NCI (Model FY2026E year end, 31-Dec-2026)',
        'D48': 'GE Aerospace: NCI ~$0.2bn; net debt at the base-year end so that FY2027E+ cash flows are not double counted. Run-off insurance (investment securities vs insurance liabilities) not deducted (assumed self-funding).',
        'B50': 'Diluted shares (m, Model FY2026E period end)',
        'D76': 'Rationale (GE Aerospace characteristics)',
        'C78': 0.11, 'D78': 'Model operating profit CAGR FY2026–31E: services growth on a ~50,000-engine commercial installed base, LEAP shop-visit ramp, DPT, plus CPP → ~11% NOPAT growth',
        'C79': 0.35, 'D79': 'Installed-base aftermarket (CES services ~60% of revenue) earns very high incremental returns; new engines are sold near cost and repaid by decades of services',
        'C80': 20, 'D80': 'Duopoly / joint-venture positions (CFM with Safran) on narrowbody and widebody platforms, certification barriers and 25–30-year engine lives sustain a long competitive-advantage period',
        'B90': 'VALUATION  —  Applied to FY2029E NOPAT',
        'B91': 'FY2029E operating profit less restructuring & separation costs',
        'B93': 'FY2029E NOPAT',
        'B95': 'Implied Enterprise Value (FY2029E basis)',
        'B96': 'Less: FY2028E year-end net debt + NCI (Model)',
        'B98': 'Diluted shares FY2029E (m, Model)',
        'B99': 'Implied share price FY2029 ($)',
        'B106': 'Plus: FY2026E year-end net debt + NCI',
        'B108': 'GE market-implied EV/NOPAT (on FY2029E NOPAT)',
    },
    'Charts': {
        'B2': 'General Electric Company (GE Aerospace) — Operating Performance (FY2023–FY2025) and Model Forecast (FY2026E–FY2031E)',
        'B3': ('USDm · calendar years · Linked to Model. Bars = contribution of each segment to adjusted revenue growth (pp; segments as reported in each year, Jan-2026 structure '
               'from 2026). Adjusted EBITDA = operating profit + D&A (model). ROIC / RONTA on operating profit (full-cost); charts start in 2023, the first year with a GE Aerospace balance sheet.'),
        'B8': 'CES contribution to growth (pp)', 'B9': 'DPT contribution to growth (pp)', 'B12': 'Adjusted EBITDA', 'B13': 'Profit before tax (GAAP)',
        'B14': 'Adj. EBITDA Margin %', 'B15': 'Profit Margin (GAAP) %',
    },
}
BBB_ROWS = {9: 'Adjusted revenue', 10: '   Adjusted revenue y/y %', 11: '   CES revenue y/y %', 12: 'Adjusted EBITDA (operating profit + D&A)',
            15: 'Profit before tax (GAAP)', 16: '   Profit before tax y/y %', 17: '   Profit margin (GAAP) %', 19: 'Adjusted Net Income',
            35: 'Adjusted EBITDA (pro forma CPP full year in 2027E)', 36: '   Adjusted EBITDA (pro forma) % of adjusted revenue'}
CHART_TITLES = {
    'Aggregates Revenue Drivers: Shipment Volume vs. Price (ASP), y/y %': 'Adjusted Revenue Growth: CES vs. DPT Contribution, pp',
    'Adjusted EBITDA &amp; Earnings from Operations (USDm)': 'Adjusted EBITDA &amp; Profit before Tax (USDm)',
}
YEARS = list(range(2023, 2032))
LASTCH = L(2 + len(YEARS))                      # Charts columns C.. (one per year)

def build_charts(ws):
    for i in range(20):
        c = L(3 + i)
        for r in (7, 8, 9, 11, 12, 13, 14, 15, 17, 18, 19, 20, 21): ws[f'{c}{r}'] = None
    for i, y in enumerate(YEARS):
        c = L(3 + i); mc = ANN[y].c
        ws[f'{c}7'] = y if y <= 2025 else f'{y}E'
        ws[f'{c}8'] = f'=IFERROR(Model!{mc}{R["gc_ces"]}+0,0)'
        ws[f'{c}9'] = f'=IFERROR(Model!{mc}{R["gc_dpt"]}+0,0)'
        ws[f'{c}11'] = f'={c}7'
        ws[f'{c}12'] = f'=Model!{mc}{R["r_adj"]}'
        ws[f'{c}13'] = f'=Model!{mc}{R["is_ebt"]}'
        ws[f'{c}14'] = f'=Model!{mc}{R["gm_ebitda_m"]}'
        ws[f'{c}15'] = f'=Model!{mc}{R["gm_op_m"]}'
        ws[f'{c}17'] = f'={c}7'
        ws[f'{c}18'] = f'=IFERROR(Model!{mc}{R["ra_roic"]}+0,NA())'
        ws[f'{c}19'] = f'=IFERROR(Model!{mc}{R["rn_ronta"]}+0,NA())'
        ws[f'{c}20'] = f'=IFERROR(Model!{mc}{R["gm_ebitda_m"]}+0,NA())'
        ws[f'{c}21'] = f'=IFERROR(Model!{mc}{R["fcf_m"]}+0,NA())'

def do_all(wb):
    for name in ('Bull-Base-Bear', 'DCF', 'Charts'):
        ws = wb[name]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and 'Model!' in c.value:
                    h = hist_stat(c.value) if name == 'Bull-Base-Bear' else None
                    c.value = h or remap(c.value)
                if isinstance(c.value, str) and 'DATE(' in c.value:
                    c.value = fy_dates(c.value)
        for addr, t in TEXT.get(name, {}).items():
            ws[addr] = t
    b = wb['Bull-Base-Bear']
    for r, t in BBB_ROWS.items():
        for c in ('B', 'N', 'Z'): b[f'{c}{r}'] = t
    for c0 in ('C', 'O', 'AA'):
        i = b[f'{c0}7'].column
        b.cell(7, i).value = '2026E'
        for k in range(1, 6): b.cell(7, i + k).value = f'{2026 + k}E'
        b.cell(7, i + 6).value = "26-'31 CAGR"
    from openpyxl.comments import Comment
    cm = Comment(f'P/E on adjusted EPS (company definition). GE at ${PX:.2f} trades at ~40x FY2026E adjusted EPS (guide $7.65–7.85) and ~34x FY2027E. Exit multiples are '
                 'assumptions: Base 36x, Bull 42x, Bear 28x applied to FY2030E–31E (FY2029E–31E for Bear), converging toward aerospace-aftermarket peers as growth normalises.', 'GE model')
    cm.width, cm.height = 420, 120; b['N26'].comment = cm
    cm2 = Comment(f'Snapshot (values) of Model rows {R["bb_cash"]}, {R["bb_px"]} and {R["bb_sh"]} (Share Buyback Schedule) with Model!{NOTE}2 = Bull / Bear. '
                  'Base panel (D40:H42) is live-linked.', 'GE model')
    cm2.width, cm2.height = 420, 90; b['P40'].comment = cm2
    b['J2'] = 'Share price ($):'
    for c in ('B', 'N', 'Z'): b[f'{c}29'] = 'IRR (to 31-Dec of Y-1)'
    for c in ('B', 'N', 'Z'):
        if b[f'{c}40'].value: b[f'{c}40'] = '   Share repurchases ($m; FY2026E = 2H)'
    # implied valuation rows only where a P/E is assumed (template computed a zero market cap in years without one)
    for row in b.iter_rows(min_row=47, max_row=49):
        for c in row:
            v = c.value
            if isinstance(v, str) and v.startswith('=') and not v.startswith('=IF(ISNUMBER('):
                col = c.column_letter
                c.value = f'=IF(ISNUMBER({col}27),{v[1:]},"")'
    d = wb['DCF']
    for i, col in enumerate('DEFGHIJKLM'):
        d[f'{col}17'] = f'{2027 + i}E'
    a26, a27 = ANN[2026].c, ANN[2027].c
    d['D22'] = f'=Model!{a27}{R["cf_capex"]}+Model!{a27}{R["cf_acq"]}'
    d['D34'] = f'=Model!{a26}{R["ra_ic"]}'
    d['C48'] = f'=Model!{a26}{R["lv_nd"]}+Model!{a26}{R["bs_nci"]}'
    d['C50'] = f'=Model!{a26}{R["v_sh"]}'
    d['C52'] = f'=Model!{a26}{R["v_px"]}'
    d['P17'] = 'CAGR 27–31E'; d['Q17'] = 'CAGR 27–36E'
    for r in range(18, 32):                        # CAGRs across a negative FY2027E UFCF (CPP consideration) are not meaningful
        for c in ('P', 'Q'):
            v = d[f'{c}{r}'].value
            if isinstance(v, str) and v.startswith('=') and not v.startswith('=IFERROR('): d[f'{c}{r}'] = f'=IFERROR({v[1:]},"n/m")'
    build_charts(wb['Charts'])

def snapshot_values(wb, bull, bear):
    ws = wb['Bull-Base-Bear']
    cols = 'CDEFGHIJKL'
    for snap, off in ((bull, 'OPQRSTUVWX'), (bear, ['AA', 'AB', 'AC', 'AD', 'AE', 'AF', 'AG', 'AH', 'AI', 'AJ'])):
        for (r, c), v in snap.items():
            tgt = f'{off[cols.index(c)]}{r}'
            cur = ws[tgt].value
            if isinstance(cur, str) and cur.startswith('='): continue
            if r in (7, 26): continue
            ws[tgt] = v

def inject_charts(path):
    """Replace openpyxl-rewritten chart parts with the template chart XML (combo charts, colours, layout), retitled, ranges cut to the
    GE years (Charts!C:K = FY2023-FY2031E) and cached values removed so Excel redraws from the GE cells."""
    tz = zipfile.ZipFile(TEMPLATE)
    tmpl = {n: tz.read(n) for n in tz.namelist() if n.startswith('xl/charts/chart')}
    tmp = path + '.tmp'
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in tmpl:
                t = tmpl[item.filename].decode('utf-8')
                for a, b2 in CHART_TITLES.items(): t = t.replace(a, b2)
                t = re.sub(r"(Charts'?!\$[A-Z]+\$\d+:\$)V(\$\d+)", lambda m: m.group(1) + LASTCH + m.group(2), t)
                t = re.sub(r'<c:numCache>.*?</c:numCache>', '', t, flags=re.S)
                t = re.sub(r'<c:strCache>.*?</c:strCache>', '', t, flags=re.S)
                data = t.encode('utf-8')
            zout.writestr(item, data)
    shutil.move(tmp, path)
