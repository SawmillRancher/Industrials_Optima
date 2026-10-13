"""Rewire the Bull-Base-Bear, DCF and Charts tabs of the MLM template (current layout: DCF with live valuation date, fade period,
implied ROIC, target-IRR entry price, reverse DCF and Mauboussin EV/NOPAT; Bull-Base-Bear with buyback rows) to the RBC Model rows /
columns; keep the template chart XML. The template's DCF_old tab is not carried.

Column mapping: the template's 2025A base column becomes RBC FY2027E (fiscal year ending 3-Apr-2027; Q1 actual, Q2 calibrated to the
outlook, Q3-Q4 carried) and the template's 2026E-2030E become RBC FY2028E-FY2032E, so the DCF's explicit years start in FY2028E and
net debt is taken at the FY2027E year end (no double counting of the current year's cash flow)."""
import re, zipfile, shutil
from openpyxl.utils import get_column_letter as L
from fw import R, TEMPLATE, ANN, NOTE

# template Model row -> RBC row key
MAP = {158: 'is_rev', 164: 'r_cash', 166: 'is_oi', 182: 'd_dxa', 186: 'sh_dil', 211: 'r_adj', 212: 'r_adj_m', 219: 'r_ebita', 221: 'r_adj',
       229: 'e_adjnic', 231: 'e_eps', 238: 'gm_aero', 245: 'gm_ebitda_m', 247: 'gm_op_m', 260: 'cf_wc', 262: 'cf_capex', 263: 'cf_acq',
       280: 'fcf', 282: 'fcf_m', 290: 'bb_px', 291: 'bb_cash', 292: 'bb_sh', 331: 'v_nci', 408: 'ra_ic', 411: 'ra_roic', 422: 'rn_ronta',
       428: 'lv_nd', 433: 'v_px', 434: 'v_sh'}
CMAP = {'BV': ANN[2027].c, 'CA': ANN[2028].c, 'CB': ANN[2029].c, 'CC': ANN[2030].c, 'CD': ANN[2031].c, 'CE': ANN[2032].c, 'CF': NOTE}
REF = re.compile(r"Model!(\$?)([A-Z]{1,3})(\$?)(\d+)")
HIST = [ANN[y].c for y in range(2006, 2027)]

def remap(f):
    def rep(m):
        c, r = m.group(2), int(m.group(4))
        c2 = CMAP.get(c, c)
        r2 = R[MAP[r]] if r in MAP else r
        if c == 'CF' and r == 2: r2 = 2
        return f'Model!{m.group(1)}{c2}{m.group(3)}{r2}'
    return REF.sub(rep, f)

def hist_stat(f):
    """LT historical AVERAGE / MIN / MAX over the template's annual columns -> RBC FY2006-FY2026 annual columns."""
    m = re.match(r'=IFERROR\((AVERAGE|MIN|MAX)\(Model!C(\d+),', f)
    if not m: return None
    fn, r = m.group(1), int(m.group(2))
    rr = R[MAP[r]]
    return f'=IFERROR({fn}({",".join(f"Model!{c}{rr}" for c in HIST)}),"")'

def fy_dates(f):
    """template calendar year ends (31-Dec of year Y, Y = 2026 for the first explicit year) -> RBC fiscal year ends (~31-Mar of Y + 2)."""
    f = re.sub(r'DATE\((\d{4}),12,31\)', lambda m: f'DATE({int(m.group(1)) + 2},3,31)', f)
    f = re.sub(r'DATE\((\d{4})-1,12,31\)', lambda m: f'DATE({int(m.group(1)) + 2}-1,3,31)', f)
    return f

TEXT = {
    'Bull-Base-Bear': {
        'B2': 'RBC Bearings Incorporated (RBC) — Bull / Base / Bear Case P&L Summary',
        'B3': ('Adjusted figures (company definitions: adjusted operating income, adjusted EBITDA, adjusted EPS — current definition) · USD millions · fiscal years ending ~31 March · '
               'Live panel linked to Model — toggle scenario via Model!CL2. Bull and Bear panels are value snapshots generated from the Model with the switch set to each case. '
               'FY2027E = Q1 actual + Q2 calibrated to the Q2 FY27 outlook (sales $505–515m, gross margin 45.5–45.75%) + Q3–Q4 with the same Δ.'),
        'G26': 32, 'H26': 32, 'S26': 36, 'T26': 36, 'AD26': 25, 'AE26': 25, 'AF26': 25,
    },
    'DCF': {
        'B2': 'DCF VALUATION — RBC BEARINGS INCORPORATED (RBC)',
        'C5': 0.0425, 'D5': 'Assumption (~US 10Y Treasury yield, Oct-2026; government rate sources are not reachable from the build environment); update to market',
        'D6': 'Assumption: US implied ERP ~4.5–5.5%',
        'C7': 1.05, 'D7': 'Assumption: precision-components / aerospace peers (TDG, HEI, CW, TKR, KMT) ~0.9–1.3; ~58% of sales Industrial (cyclical: mining, construction, oil & gas)',
        'C9': 0.055, 'D9': 'Assumption: 4.375% senior notes due 2029 and the revolver at SOFR + 0.85% (~5.0%); marginal cost of new notes ~5.5%',
        'C10': 0.23, 'D10': 'Model FY2028E+ effective tax rate (FY2026 22.1%; Q1 FY27 22.1%)',
        'C12': 0.05, 'D12': 'Net debt ~$0.7bn at 27-Jun-2026 vs ~$16bn market cap ($510 × ~31.8m diluted shares); ~1.0x adjusted EBITDA',
        'D15': 'Periods = (31-Mar of fiscal year − today) / 365.25; cash flows assumed at fiscal year end; FY2027E (ends 3-Apr-2027) is the base year — its cash flow is in the net debt below',
        'B18': 'Adj. EBITA (after stock comp; before M&A amortization) less transaction & restructuring costs',
        'B21': 'plus D&A excluding M&A-related amortization (depreciation & software)',
        'B22': '− Capex (PP&E) and acquisitions (M&A lever FY2028E–32E)',
        'B26': 'Fade-period growth (FY2033E–37E; 7.0% → 5.0% linear)', 'I26': 0.07, 'M26': 0.05, 'C27': 0.04,
        'D32': ("NOPAT basis (playbook 5.7): RBC's adjusted EPS and adjusted EBITDA add back stock compensation; here it is deducted (row 18 starts from adjusted operating income, "
                'which is after stock compensation) and M&A-related amortization of acquired intangibles is added back (EBITA), so the tax is charged on pre-amortization profit; '
                'row 21 adds back depreciation and non-acquired amortization only. Transaction, transition and restructuring costs (cash; adjusted out by the company) are '
                'deducted in row 18. Row 22 includes acquisitions from the M&A lever in FY2028E–32E; the fade period and terminal value assume no new acquisitions. '
                'The cash tax shield on acquired-intangible amortization (e.g. VACCO, Section 338(h)(10) election per the FY2026 10-K) is not modelled separately.'),
        'D43': 'Opening IC = Model FY2027E book IC (equity + net debt at the FY2027E year end), rolled forward with capex + M&A − D&A excl. M&A amortization + ΔWC.',
        'B48': 'Less: Net debt + NCI (Model FY2027E year end)',
        'D48': 'RBC: no noncontrolling interests; net debt at the base-year end (3-Apr-2027) so that FY2028E+ cash flows are not double counted; operating leases not treated as debt.',
        'B50': 'Diluted shares (m, Model FY2027E period end)',
        'D76': 'Rationale (RBC Bearings characteristics)',
        'C78': 0.09, 'D78': 'Model adjusted EBITA CAGR FY2027–32E: high-single-digit sales growth (defense / space / commercial aerospace, Dodge industrial) with stable margins → ~9% NOPAT growth',
        'C79': 0.15, 'D79': 'Sole-source, certified bearings on long-lived aircraft and defense platforms plus Dodge replacement demand earn high incremental returns; the Dodge goodwill depresses reported ROIC (~8–10%)',
        'C80': 15, 'D80': 'Qualification barriers (OEM approvals, military specs) and a large aftermarket sustain a long competitive-advantage period; Industrial exposure is cyclical',
        'B90': 'VALUATION  —  Applied to FY2030E NOPAT',
        'B91': 'FY2030E Adj. EBITA less transaction & restructuring costs',
        'B93': 'FY2030E NOPAT',
        'B95': 'Implied Enterprise Value (FY2030E basis)',
        'B96': 'Less: FY2029E year-end net debt + NCI (Model)',
        'B98': 'Diluted shares FY2030E (m, Model)',
        'B99': 'Implied share price FY2030 ($)',
        'B106': 'Plus: FY2027E year-end net debt + NCI',
        'B108': 'RBC market-implied EV/NOPAT (on FY2030E NOPAT)',
    },
    'Charts': {
        'B2': 'RBC Bearings — Operating Performance (FY2007–FY2026) and Model Forecast (FY2027E–FY2032E)',
        'B3': ('USDm · fiscal years ending ~31 March · Linked to Model. Bars = contribution of each segment to total net-sales growth (pp; Aerospace & Defense / Industrial from FY2022, '
               'total growth before). Adjusted EBITDA = company definition (adjusted operating income + D&A + stock compensation). ROIC / RONTA on adjusted EBIT (full-cost).'),
        'B8': 'Aerospace & Defense contribution to growth (pp; total before FY2022)', 'B9': 'Industrial contribution to growth (pp)', 'B12': 'Adjusted EBITDA', 'B13': 'Operating income (GAAP)',
        'B14': 'Adj. EBITDA Margin %', 'B15': 'Operating Margin %',
    },
}
BBB_ROWS = {9: 'Net sales', 10: '   Net sales y/y %', 11: '   Aerospace & Defense sales y/y %', 12: 'Adjusted EBITDA',
            15: 'Operating income (GAAP)', 16: '   Operating income y/y %', 17: '   Operating margin (GAAP) %', 19: 'Adjusted Net Income (to common)',
            35: 'Adjusted EBITDA (leverage basis)', 36: '   Adjusted EBITDA % of net sales'}
CHART_TITLES = {
    'Aggregates Revenue Drivers: Shipment Volume vs. Price (ASP), y/y %': 'Net Sales Growth: Aerospace &amp; Defense vs. Industrial Contribution, pp',
    'Adjusted EBITDA &amp; Earnings from Operations (USDm)': 'Adjusted EBITDA &amp; Operating Income (USDm)',
}

def build_charts(ws):
    for i, y in enumerate(range(2013, 2033)):
        c = L(3 + i); mc = ANN[y].c
        ws[f'{c}7'] = y if y <= 2026 else f'{y}E'
        ws[f'{c}8'] = f'=IFERROR(Model!{mc}{R["gc_aero"]}+0,IFERROR(Model!{mc}{R["gc_tot"]}+0,0))'
        ws[f'{c}9'] = f'=IFERROR(Model!{mc}{R["gc_ind"]}+0,0)'
        ws[f'{c}11'] = f'={c}7'
        ws[f'{c}12'] = f'=Model!{mc}{R["r_adj"]}'
        ws[f'{c}13'] = f'=Model!{mc}{R["is_oi"]}'
        ws[f'{c}14'] = f'=Model!{mc}{R["gm_ebitda_m"]}'
        ws[f'{c}15'] = f'=Model!{mc}{R["gm_op_m"]}'
        ws[f'{c}17'] = f'={c}7'
        ws[f'{c}18'] = f'=IFERROR(Model!{mc}{R["ra_roic"]}+0,NA())'
        ws[f'{c}19'] = f'=IFERROR(Model!{mc}{R["rn_ronta"]}+0,NA())'
        ws[f'{c}20'] = f'=IFERROR(Model!{mc}{R["gm_ebitda_m"]}+0,NA())'
        ws[f'{c}21'] = f'=IFERROR(Model!{mc}{R["fcf_m"]}+0,NA())'

def do_all(wb):
    for name in ('Bull-Base-Bear', 'DCF', 'Charts'):
        ws = wb[name]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and 'Model!' in c.value:
                    h = hist_stat(c.value) if name == 'Bull-Base-Bear' else None
                    c.value = h or remap(c.value)
                if isinstance(c.value, str) and 'DATE(' in c.value:
                    c.value = fy_dates(c.value)
        for addr, t in TEXT.get(name, {}).items():
            ws[addr] = t
    b = wb['Bull-Base-Bear']
    for r, t in BBB_ROWS.items():
        for c in ('B', 'N', 'Z'): b[f'{c}{r}'] = t
    for c0 in ('C', 'O', 'AA'):
        i = b[f'{c0}7'].column
        b.cell(7, i).value = '2027E'
        for k in range(1, 6): b.cell(7, i + k).value = f'{2027 + k}E'
        b.cell(7, i + 6).value = "27-'32 CAGR"
    from openpyxl.comments import Comment
    cm = Comment('P/E on adjusted EPS (company current definition: adds back M&A amortization and stock compensation). RBC at $510 (user price, 13-Oct-2026) trades at '
                 '~34x FY2027E Base adjusted EPS. Exit multiples are assumptions: Base 32x, Bull 36x, Bear 25x, converging toward aerospace-components peers as growth normalises.', 'RBC model')
    cm.width, cm.height = 420, 120; b['N26'].comment = cm
    cm2 = Comment(f'Snapshot (values) of Model rows {R["bb_cash"]}, {R["bb_px"]} and {R["bb_sh"]} (Share Buyback Schedule) with Model!{NOTE}2 = Bull / Bear. '
                  'Base panel (D40:H42) is live-linked.', 'RBC model')
    cm2.width, cm2.height = 420, 90; b['P40'].comment = cm2
    b['J2'] = 'Share price ($):'
    for c in ('B', 'N', 'Z'): b[f'{c}29'] = 'IRR (to ~31-Mar of Y-1)'
    for c in ('B', 'N', 'Z'):
        if b[f'{c}40'].value: b[f'{c}40'] = '   Share repurchases ($m; open market, ex withholding)'
    d = wb['DCF']
    for i, col in enumerate('DEFGHIJKLM'):
        d[f'{col}17'] = f'{2028 + i}E'
    a26, a27 = ANN[2027].c, ANN[2028].c
    d['D22'] = f'=Model!{a27}{R["cf_capex"]}+Model!{a27}{R["cf_acq"]}'
    d['D34'] = f'=Model!{a26}{R["ra_ic"]}'
    d['C48'] = f'=Model!{a26}{R["lv_nd"]}+Model!{a26}{R["v_nci"]}'
    d['C50'] = f'=Model!{a26}{R["v_sh"]}'
    d['C52'] = f'=Model!{a26}{R["v_px"]}'
    d['P17'] = 'CAGR 28–32E'; d['Q17'] = 'CAGR 28–37E'
    build_charts(wb['Charts'])

def snapshot_values(wb, bull, bear):
    ws = wb['Bull-Base-Bear']
    cols = 'CDEFGHIJKL'
    for snap, off in ((bull, 'OPQRSTUVWX'), (bear, ['AA', 'AB', 'AC', 'AD', 'AE', 'AF', 'AG', 'AH', 'AI', 'AJ'])):
        for (r, c), v in snap.items():
            tgt = f'{off[cols.index(c)]}{r}'
            cur = ws[tgt].value
            if isinstance(cur, str) and cur.startswith('='): continue
            if r in (7, 26): continue
            ws[tgt] = v

def inject_charts(path):
    """Replace openpyxl-rewritten chart parts with the template chart XML (combo charts, colours, layout), retitled, ranges kept
    (Charts!C:V = FY2013-FY2032E) and cached values removed so Excel redraws from the RBC cells."""
    tz = zipfile.ZipFile(TEMPLATE)
    tmpl = {n: tz.read(n) for n in tz.namelist() if n.startswith('xl/charts/chart')}
    tmp = path + '.tmp'
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in tmpl:
                t = tmpl[item.filename].decode('utf-8')
                for a, b2 in CHART_TITLES.items(): t = t.replace(a, b2)
                t = re.sub(r'<c:numCache>.*?</c:numCache>', '', t, flags=re.S)
                t = re.sub(r'<c:strCache>.*?</c:strCache>', '', t, flags=re.S)
                data = t.encode('utf-8')
            zout.writestr(item, data)
    shutil.move(tmp, path)
